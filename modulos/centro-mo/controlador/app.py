# -*- coding: utf-8 -*-
"""Centro MO — módulo SGO7 (servido bajo /centro-mo/).

Nace de la herramienta local `centro_mo.py`, que corría en el PC del operador. Aquí vive TODO lo que no necesita Outlook: el store de pendientes, las
reglas de completitud, la generación de la planilla con la librería de nuevo/ (Balance), las
lecturas de SGS (liquidación previa, líderes, fecha de bandejas, materiales), el embudo hacia
liquidador2 (mismo servidor, :5001) y la interfaz.

Lo que SÍ necesita Outlook lo hace el **agente del PC** (`agentes-pc/centro-mo/agente_centro_mo.py`), la
única pieza que ve el buzón del operador — mismo patrón que el agente de correos ANS:
  - vigila la Bandeja de entrada (categoría MO / Orden Pendiente), analiza los correos y los
    entrega aquí (`POST /agente/ingesta`);
  - atiende **pedidos** que deja esta app (`GET /agente/estado`): reanalizar un correo, abrir el
    borrador de consulta al supervisor, abrir el correo en Outlook, traer la planilla de un link
    de SharePoint (perfil de Microsoft del PC);
  - vigila las respuestas de los supervisores a las consultas abiertas.
El agente entra por `/centro-mo/agente/*` con la cabecera `X-Agente-Token` (nginx no pide la
puerta en esa ruta y borra cualquier X-SGO7-* que traiga). Todo lo demás pasa por la puerta
única de SGO7 (rol desarrollador) y queda en Registros (`sgo7_registro`).

Un solo worker de gunicorn (`-w 1 --threads 8`): el estado vive en memoria y se persiste en
`data/centro/pendientes.json` (mismo formato que en el PC) y `data/centro/pedidos.json`.
"""
import base64
import hashlib
import hmac
import json
import os
import re
import sys
import threading
import time
import uuid
from datetime import datetime
from functools import wraps
from urllib.parse import urlparse

from flask import Flask, jsonify, request, send_file

# sgo7_registro ANTES que el puente: el puente mete nuevo/ en sys.path[0] y nuevo/ trae su propia
# copia de sgo7_registro (idéntica, pero mejor no depender de eso).
try:
    import sgo7_registro
except Exception as _e_reg:          # pragma: no cover — en el PC no hace falta
    sgo7_registro = None
    print("sgo7_registro no disponible:", _e_reg)

import balance_mo_bridge as bridge

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

app = Flask(__name__)

PUERTO = int(os.environ.get("CENTRO_MO_PUERTO", "5088"))
POLL_SEGUNDOS = int(os.environ.get("CENTRO_MO_POLL", "180"))       # cada cuánto el agente revisa el buzón
DIAS_ATRAS = int(os.environ.get("CENTRO_MO_DIAS_ATRAS", "120"))     # ventana de detección
AGENTE_TIMEOUT = int(os.environ.get("CENTRO_MO_AGENTE_TIMEOUT", "150"))   # sin latido en N s = agente apagado
PEDIDO_TIMEOUT = 15 * 60                                              # un pedido tomado sin respuesta caduca
TOKEN_AGENTE = os.environ.get("CENTRO_MO_AGENTE_TOKEN", "")

# ── Cierre mensual de SGS ─────────────────────────────────────────────────────
# En la herramienta local era una constante que se actualizaba a mano cada mes. En sgo7 el Liquidador
# aprende el rango de producción solo (sgs_cierre.json, regla autorizada el 8 oct 2026) y
# aquí se lee de ahí; si no hay cierre vigente se usa el último valor conocido.
FECHA_MINIMA_LIQ_DEFECTO = "2026-08-05"


def fecha_minima_liq() -> str:
    try:
        import sgs_cierre
        f = sgs_cierre.fecha_minima()
        if f:
            return f
    except Exception:
        pass
    return FECHA_MINIMA_LIQ_DEFECTO


INCLUIR_MAT = True
AUTO_GENERAR = True
# Embudo hacia liquidador2: mismo servidor. La URL pública (lo que abre el navegador) va con
# el prefijo de nginx; la interna, directa al puerto.
LIQUIDADOR2_BASE = os.environ.get("LIQUIDADOR2_BASE", "http://127.0.0.1:5001/liquidador2")
LIQUIDADOR2_PUBLIC = os.environ.get("LIQUIDADOR2_PUBLIC", "/liquidador2")
BALANCE_PUBLIC = os.environ.get("BALANCE_PUBLIC", "/balance")

DIR_SCRIPT = os.path.dirname(os.path.abspath(__file__))
DIR_DATA = os.environ.get("CENTRO_MO_DATA") or os.path.join(DIR_SCRIPT, "data", "centro")
RUTA_STORE = os.path.join(DIR_DATA, "pendientes.json")
RUTA_PEDIDOS = os.path.join(DIR_DATA, "pedidos.json")
# Las planillas (generadas aquí o traídas por el agente) viven con los registros de SGO7.
DIR_PLANILLAS = os.environ.get("CENTRO_MO_PLANILLAS") or os.path.join(DIR_SCRIPT, "data", "planillas")

os.makedirs(DIR_DATA, exist_ok=True)
os.makedirs(DIR_PLANILLAS, exist_ok=True)

# ── Store persistente ─────────────────────────────────────────────────────────
_lock = threading.RLock()
_store: dict = {}            # id → registro
_pedidos: dict = {}          # id → pedido para el agente del PC
_agente = {"ultimo_latido": 0.0, "ultimo_scan": "", "scaneando": False, "error": "",
           "version": "", "scan_ahora": False, "nuevos_ultimo_scan": 0}


def _ahora():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _cargar_store():
    global _store
    try:
        with open(RUTA_STORE, encoding="utf-8") as f:
            _store = json.load(f) or {}
    except Exception:
        _store = {}


def _guardar_store():
    with _lock:
        tmp = RUTA_STORE + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(_store, f, ensure_ascii=False, indent=1)
        os.replace(tmp, RUTA_STORE)


def _cargar_pedidos():
    global _pedidos
    try:
        with open(RUTA_PEDIDOS, encoding="utf-8") as f:
            _pedidos = json.load(f) or {}
    except Exception:
        _pedidos = {}


def _guardar_pedidos():
    with _lock:
        tmp = RUTA_PEDIDOS + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(_pedidos, f, ensure_ascii=False, indent=1)
        os.replace(tmp, RUTA_PEDIDOS)


# ── Trazabilidad SGO7 ─────────────────────────────────────────────────────────
if sgo7_registro:
    sgo7_registro.instalar(app, modulo="centro-mo", acciones={
        r"/api/generar/": "centro_mo_generar_planilla",
        r"/api/enviar-liq2/": "centro_mo_a_liquidador2",
        r"/api/consulta/": "centro_mo_consulta_supervisor",
        r"/api/editar/": "centro_mo_editar",
        r"/api/subir-excel/": "centro_mo_subir_planilla",
        r"/api/traer-link/": "centro_mo_traer_link",
        r"/api/aplicar-lideres/": "centro_mo_lideres",
        r"/api/mat-aprobar/": "centro_mo_materiales",
        r"/api/estado/": "centro_mo_estado",
        r"/api/reanalizar": "centro_mo_reanalizar",
        r"/api/descargar/": "descargar",
    }, ignorar=("/agente/", "/api/pedido/", "/api/lideres-sgs/", "/api/fecha-bandejas/",
                "/api/mat-preanalisis/", "/api/abrir-outlook/", "/api/sp-login", "/api/items"))


def _registrar(accion, reg=None, detalle=None, archivo=None, estado="ok"):
    """Registro explícito. Dentro de una petición toma el usuario de la puerta; desde el canal
    del agente (sin X-SGO7-*) o desde un hilo, firma como agente-pc."""
    if not sgo7_registro:
        return
    try:
        det = dict(detalle or {})
        if reg:
            det.setdefault("codigo", reg.get("codigo"))
            det.setdefault("asunto", (reg.get("asunto") or "")[:120])
            det.setdefault("remitente", reg.get("remitente"))
        usuario = rol = None
        try:
            usuario = request.headers.get("X-SGO7-User") or None
            rol = request.headers.get("X-SGO7-Rol") or None
        except Exception:
            pass
        if not usuario:
            usuario, rol = "agente-pc", "agente"
        sgo7_registro.registrar(accion, detalle=det, estado=estado, archivo=archivo,
                                sesion=(reg or {}).get("id"), usuario=usuario, rol=rol)
    except Exception:
        pass


def _sha256(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _nombre_seguro(nombre: str, defecto="planilla.xlsx") -> str:
    base = os.path.basename((nombre or "").replace("\\", "/")) or defecto
    base = re.sub(r"[^A-Za-z0-9._ ()\-]+", "_", base)[:140]
    return base or defecto


def _ruta_planilla(rid: str, nombre: str) -> str:
    d = os.path.join(DIR_PLANILLAS, rid)
    os.makedirs(d, exist_ok=True)
    return os.path.join(d, _nombre_seguro(nombre))


# ── Reglas de completitud (idénticas al PC) ───────────────────────────────────

def recalcular(reg: dict):
    cod = (reg.get("codigo") or "").upper()
    if cod.startswith("DF"):
        reg["tipo"] = "MI"
        reg["pdl"] = ""            # regla: MI no lleva PDL
    elif cod.startswith("LCL") or (cod.isdigit() and len(cod) >= 8):
        reg["tipo"] = "Programados"
    faltantes = []
    if not cod:
        faltantes.append("código (DF o LCL)")
    if reg.get("tipo") == "Programados" and not reg.get("pdl"):
        faltantes.append("PDL")
    if not reg.get("filas_mo"):
        faltantes.append("planilla de prestaciones" + (" (hay link externo)" if reg.get("link_externo") else ""))
    if not reg.get("lideres"):
        faltantes.append("líder")
    if not reg.get("fecha_ejecucion"):
        faltantes.append("fecha de ejecución")
    else:
        # Regla del cierre mensual: fechas antes del mínimo se corren
        f = reg["fecha_ejecucion"]
        minima = fecha_minima_liq()
        if f < minima:
            reg["fecha_original"] = f
            reg["fecha_ejecucion"] = minima
            nota = (f"Fecha de ejecución {f} corrida a {minima} "
                    f"(cierre mensual: SGS no deja liquidar antes)")
            if nota not in reg.get("notas", []):
                reg.setdefault("notas", []).append(nota)
    reg["faltantes"] = faltantes
    if reg.get("estado") not in ("planilla_lista", "completado", "descartado", "esperando_supervisor"):
        reg["carril"] = "super" if not faltantes else "semi"
        reg["estado"] = "listo_auto" if not faltantes else "falta_info"
    elif reg.get("estado") == "esperando_supervisor" and not faltantes:
        reg["carril"] = "super"
        reg["estado"] = "listo_auto"
    if (reg.get("planilla_nativa") and reg.get("excel_ruta") and not faltantes
            and reg.get("estado") not in ("completado", "descartado")):
        reg["carril"] = "super"
        reg["estado"] = "planilla_lista"


_PRIORIDAD_FUENTE = {"planilla": 0, "asunto": 1, "cuerpo": 2, "imagen": 3, "cadena": 4,
                     "excel": 5}


def _prio_fuente(f: str) -> int:
    return _PRIORIDAD_FUENTE.get((f or "").split(" ")[0], 9)


# Instrucción de MATERIAL del correo (la calcula el agente del PC, instruccion_mat.py).
# Regla (22 oct 2026): sin instrucción en contra, la planilla estaba habilitada para MAT.
_IM_VALORES = {"no_mat", "parcial", "pendiente_almacen", "con_mat", "solo_menciona_mo", "sin_mencion"}
_IM_EXPLICITOS = {"no_mat", "parcial", "pendiente_almacen", "con_mat"}
_IM_CAMPOS = ("valor", "incluye", "frase", "fuente", "en", "version")


def _merge_analisis(base: dict, nuevo: dict):
    """Completa en base lo que el nuevo análisis encontró y base no tenía (reglas del PC)."""
    gana = nuevo.get("codigo_validado") and nuevo.get("codigo") and (
        not base.get("codigo_validado")
        or _prio_fuente(nuevo.get("codigo_fuente")) < _prio_fuente(base.get("codigo_fuente"))
        or (nuevo.get("codigo_coherencia") and not base.get("codigo_coherencia")))
    if gana:
        viejo = base.get("codigo")
        if viejo and nuevo["codigo"] != viejo:
            base.setdefault("notas", []).append(
                f"⚠ CÓDIGO CORREGIDO: {viejo} → {nuevo['codigo']} "
                f"(fuente {nuevo.get('codigo_fuente')}) — la planilla/sesión "
                f"anteriores eran del proyecto equivocado y se descartaron")
            for k in ("excel_ruta", "sid", "liq2_url", "liq_previa"):
                base.pop(k, None)
            if base.get("estado") == "planilla_lista":
                base["estado"] = "falta_info"
            if nuevo.get("pdl"):
                base["pdl"] = nuevo["pdl"]
        base["codigo"] = nuevo["codigo"]
        base["codigo_validado"] = True
        base["codigo_fuente"] = nuevo.get("codigo_fuente", "")
        base["codigo_coherencia"] = bool(nuevo.get("codigo_coherencia"))
        if nuevo["codigo"].startswith("LCL"):
            base["lcl"] = nuevo["codigo"]
        if nuevo.get("tipo") and nuevo["tipo"] != "Desconocido":
            base["tipo"] = nuevo["tipo"]
        if nuevo.get("pdl") and not base.get("pdl"):
            base["pdl"] = nuevo["pdl"]
    for k in ("codigo", "pdl", "lcl", "adjunto", "fecha_ejecucion"):
        if nuevo.get(k) and not base.get(k):
            base[k] = nuevo[k]
    if nuevo.get("filas_mo"):
        viejas = base.get("filas_mo") or []
        if not viejas:
            base["filas_mo"] = nuevo["filas_mo"]
        elif nuevo["filas_mo"] != viejas:
            base["filas_mo"] = nuevo["filas_mo"]
            if base.get("excel_ruta") and not base.get("planilla_nativa"):
                base.setdefault("notas", []).append(
                    f"⚠ PRESTACIONES ACTUALIZADAS: {len(viejas)} → "
                    f"{len(nuevo['filas_mo'])} filas — la planilla anterior "
                    "se generó con datos viejos y se descartó")
                for k in ("excel_ruta", "sid", "liq2_url", "liq_previa"):
                    base.pop(k, None)
                if base.get("estado") == "planilla_lista":
                    base["estado"] = "falta_info"
    if nuevo.get("filas_mat"):
        base["filas_mat"] = nuevo["filas_mat"]
    if nuevo.get("instruccion_lideres"):
        base["instruccion_lideres"] = nuevo["instruccion_lideres"]
    _im = nuevo.get("instruccion_mat")
    if isinstance(_im, dict) and _im.get("valor") in _IM_VALORES:
        # lo explícito del correo más reciente manda; «no dice nada» no borra lo dicho antes.
        if _im["valor"] in _IM_EXPLICITOS or not (base.get("instruccion_mat") or {}).get("valor"):
            base["instruccion_mat"] = _im
    for lid in nuevo.get("lideres", []):
        if lid.upper() not in {x.upper() for x in base.get("lideres", [])}:
            base.setdefault("lideres", []).append(lid)
    for ld in nuevo.get("lideres_det", []):
        if ld["nombre"].upper() not in {x["nombre"].upper() for x in base.get("lideres_det", [])}:
            base.setdefault("lideres_det", []).append(ld)
    base["link_externo"] = base.get("link_externo") or nuevo.get("link_externo")
    if nuevo.get("links_externos") and not base.get("links_externos"):
        base["links_externos"] = nuevo["links_externos"]
    if nuevo.get("planilla_nativa"):
        base["planilla_nativa"] = True
        if nuevo.get("excel_ruta") and not base.get("excel_ruta"):
            base["excel_ruta"] = nuevo["excel_ruta"]
    for k in ("entry_id", "store_id", "carpeta", "conversacion", "remitente_email"):
        if nuevo.get(k) and not base.get(k):
            base[k] = nuevo[k]
    base.setdefault("notas", []).extend(n for n in nuevo.get("notas", []) if n not in base.get("notas", []))
    recalcular(base)


def _clave_nodo(n) -> str:
    s = re.sub(r"\s+", "", str(n).upper()).replace("NODO", "")
    s = re.sub(r"^N(?=\d)", "", s)
    s = re.sub(r"[-_]N(?=\d)", "-", s)
    return s


def _verificar_liq(reg: dict):
    """Actualiza reg['liq_previa'] contra SGS (solo lectura)."""
    cod = reg.get("codigo")
    if not cod:
        return
    filas = reg.get("filas_mo", [])
    filtro = reg.get("nodos_filtro") or []
    if filtro:
        claves = [_clave_nodo(x) for x in filtro]
        filas = [f for f in filas if _clave_nodo(f["nodo"]) in claves]
    resumen, err = bridge.verificar_liquidacion_previa(cod, filas or None, filtro or None)
    with _lock:
        if err:
            reg.setdefault("notas", []).append(f"Verificación liquidación previa: {err}")
            return
        resumen["comparado"] = bool(filas)
        if not filas:
            resumen["conflictos"] = []
        reg["liq_previa"] = resumen


# ── Generación de la planilla (carril super) ─────────────────────────────────

def _inyectar_fecha_excel(excel: bytes, fecha: str) -> bytes:
    import io
    from copy import copy
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(excel))
        ws = wb["MO_MANUAL"]
        fila = next((i for i in range(2, 12)
                     if str(ws.cell(i, 1).value or "").startswith("Generado")), None)
        if not fila:
            return excel
        ws.insert_rows(fila)
        from openpyxl.worksheet.cell_range import CellRange
        afectados = [r for r in list(ws.merged_cells.ranges) if r.min_row >= fila]
        for r in afectados:
            ws.merged_cells.ranges.remove(r)
        for r in afectados:
            ws.merged_cells.ranges.add(CellRange(
                min_col=r.min_col, min_row=r.min_row + 1,
                max_col=r.max_col, max_row=r.max_row + 1))
        arriba = ws.cell(fila - 1, 1)
        cl = ws.cell(fila, 1)
        cl.value = "Fecha ejecución:"
        cl.font = copy(arriba.font)
        cl.fill = copy(arriba.fill)
        ws.cell(fila, 2).value = fecha
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=9)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except Exception:
        return excel


def _tokens_nombre(s: str) -> set:
    plano = bridge._sin_tildes(str(s or "").upper())
    return {t for t in re.split(r"[^A-Z0-9]+", plano) if len(t) >= 3}


def _score_lider(toks: set, cand_nombre: str):
    from difflib import SequenceMatcher
    ct = _tokens_nombre(cand_nombre)
    hits, suma = 0, 0.0
    for t in toks:
        mejor = max((SequenceMatcher(None, t, c).ratio() for c in ct), default=0.0)
        if mejor >= 0.8:
            hits += 1
            suma += mejor
    return hits, suma


def _mejor_lider(nombre: str, pool: list):
    toks = _tokens_nombre(nombre)
    if not toks:
        return None
    minimo = 2 if len(toks) >= 2 else 1
    puntuados = []
    for c in pool:
        if not str(c.get("cedula") or "").strip():
            continue
        hits, suma = _score_lider(toks, c["nombre"])
        if hits >= minimo:
            puntuados.append((hits, suma, c))
    if not puntuados:
        return None
    puntuados.sort(key=lambda x: (x[0], x[1]), reverse=True)
    if (len(puntuados) > 1 and puntuados[0][0] == puntuados[1][0]
            and abs(puntuados[0][1] - puntuados[1][1]) < 0.001
            and puntuados[0][2]["cedula"] != puntuados[1][2]["cedula"]):
        return None
    return puntuados[0][2]


def _resolver_cedulas(reg: dict):
    det = [dict(d) for d in (reg.get("lideres_det") or [])]
    if not det:
        nombres = reg.get("lideres", [])
        n = len(nombres)
        if n:
            base_pct = round(100.0 / n, 1)
            pcts = [base_pct] * n
            pcts[-1] = round(100.0 - base_pct * (n - 1), 1)
            det = [{"cedula": "", "nombre": nom, "pct": p}
                   for nom, p in zip(nombres, pcts)]
    pendientes = [d for d in det if not str(d.get("cedula") or "").strip()]
    if not pendientes:
        return det, []

    res, _err = bridge.lideres_de_bandejas(reg.get("codigo", ""), reg.get("nodos_filtro"))
    bandejas = res or []
    faltan = []
    for d in pendientes:
        m = _mejor_lider(d["nombre"], bandejas)
        origen = "bandejas del proyecto"
        if not m:
            pool, _e2 = bridge.buscar_lideres_por_nombre(d["nombre"])
            m = _mejor_lider(d["nombre"], pool or [])
            origen = "directorio SGS (búsqueda por nombre)"
        if m:
            nota = (f"Cédula del líder resuelta ({origen}): {d['nombre']} → "
                    f"{m['nombre']} — CC {m['cedula']}")
            d["cedula"] = m["cedula"]
            d["nombre"] = m["nombre"]
            if nota not in reg.get("notas", []):
                reg.setdefault("notas", []).append(nota)
        else:
            faltan.append(d["nombre"])
    unicos = {}
    for d in det:
        clave = str(d.get("cedula") or "").strip() or ("~" + d["nombre"].upper())
        if clave in unicos:
            unicos[clave]["pct"] = round((unicos[clave].get("pct") or 0) +
                                         (d.get("pct") or 0), 1)
        else:
            unicos[clave] = d
    det = list(unicos.values())
    if not faltan:
        reg["lideres_det"] = det
    return det, faltan


def generar_planilla(reg: dict, solo_pendientes: bool = False) -> str:
    """Crea la sesión en nuevo/ (misma carpeta de sesiones que el Balance del servidor) y
    escribe balance_sgs_manual_*.xlsx junto a los registros de SGO7. Solo MO (+MAT aprobado).
    Solo lecturas contra SGS. Devuelve "" si todo bien, o el error."""
    if not reg.get("fecha_ejecucion"):
        return ("Falta la fecha real de ejecución — complétala (✎) o "
                "consúltala al supervisor ()")
    minima = fecha_minima_liq()
    if reg["fecha_ejecucion"] < minima:
        reg["fecha_original"] = reg["fecha_ejecucion"]
        reg["fecha_ejecucion"] = minima
        reg.setdefault("notas", []).append(
            f"Fecha corrida a {minima} por el cierre mensual de SGS")
    filas = reg.get("filas_mo", [])
    filtro = [_clave_nodo(x) for x in (reg.get("nodos_filtro") or [])]
    if filtro:
        filas = [f for f in filas if _clave_nodo(f["nodo"]) in filtro]
        if not filas:
            return (f"El filtro de nodos {filtro} no dejó ninguna fila — revisa "
                    f"los identificadores (la planilla trae otros nodos)")

    resumen, err_liq = bridge.verificar_liquidacion_previa(
        reg.get("codigo", ""), filas, reg.get("nodos_filtro"))
    if err_liq:
        reg.setdefault("notas", []).append(
            f"No se pudo verificar liquidación previa ({err_liq}) — revisar a mano en SGS")
    elif resumen:
        resumen["comparado"] = True
        reg["liq_previa"] = resumen
        if resumen["conflictos"] and not solo_pendientes:
            lista = ", ".join(f"{c['qp']} nodo {c['nodo']} (avance {c['advance']:g})"
                              for c in resumen["conflictos"][:6])
            extra = "…" if len(resumen["conflictos"]) > 6 else ""
            nota = (f"⚠ AVANCE PREVIO EN SGS (no bloquea): {lista}{extra} — "
                    f"liquidador2 calcula el delta y tú decides allá; "
                    f"' solo pendientes' sigue disponible si quieres excluirlo")
            if nota not in reg.get("notas", []):
                reg.setdefault("notas", []).append(nota)
        if resumen["conflictos"] and solo_pendientes:
            completas = {(c["qp"].upper(), _clave_nodo(c["nodo"]))
                         for c in resumen["conflictos"] if c.get("balance", 1) <= 0.01}
            antes = len(filas)
            filas = [f for f in filas
                     if (str(f["codigo"]).upper(), _clave_nodo(f["nodo"])) not in completas]
            if not filas:
                return ("Todo lo de esta planilla ya está liquidado en SGS — "
                        "no queda nada pendiente por liquidar")
            reg.setdefault("notas", []).append(
                f"Planilla SOLO PENDIENTES: {antes - len(filas)} filas ya liquidadas "
                f"excluidas, quedan {len(filas)} por liquidar")

    det, sin_cedula = _resolver_cedulas(reg)
    if sin_cedula:
        return ("Falta la CÉDULA del líder: " + ", ".join(sin_cedula) +
                " — sin cédula liquidador2 no puede crear el líder en el equipo. "
                "Usa 'líderes SGS' para elegirlo con cédula o complétala en ✎")
    info_mat = {}
    mat_ok = INCLUIR_MAT and bool(reg.get("mat_aprobado"))
    sid, err = bridge.crear_sesion_nuevo({
        "codigo": reg.get("codigo", ""),
        "pdl": reg.get("pdl", ""),
        "filas_mo": filas,
        "incluir_mat": mat_ok,
        "mat_nodos_sel": reg.get("mat_nodos_sel") or [],
    }, info_mat)
    if err:
        return err
    from manual_balance import get_sesion, save_sesion, generar_excel_manual, ahora_co
    s = get_sesion(sid)
    if det:
        s["lideres"] = [{"cedula": d.get("cedula", ""), "nombre": d["nombre"],
                         "pct": d.get("pct") or round(100.0 / len(det), 1)} for d in det]
    s["supervisor"] = reg.get("remitente", "")
    save_sesion(sid)
    try:
        excel = generar_excel_manual(s)
    except Exception as e:
        return f"Generando Excel: {e}"
    excel = _inyectar_fecha_excel(excel, reg["fecha_ejecucion"])
    ts = ahora_co().strftime("%Y%m%d_%H%M")
    nombre = f"balance_sgs_manual_{s.get('lcl_odm', reg.get('codigo', 'balance'))}_{ts}.xlsx"
    ruta = _ruta_planilla(reg["id"], nombre)
    with open(ruta, "wb") as f:
        f.write(excel)
    reg["sid"] = sid
    reg["excel_ruta"] = ruta
    reg["estado"] = "planilla_lista"
    reg["carril"] = "super"
    reg.setdefault("notas", []).append(f"Planilla generada automáticamente: {nombre}")
    if info_mat.get("mat_nodos"):
        reg["mat_nodos"] = info_mat["mat_nodos"]
        reg["notas"].append(
            f"Materiales incluidos: balance SGS de {info_mat['mat_nodos']} nodo(s) "
            "en hojas N<nodo> — en liquidador2 decides qué MAT liquidar")
    elif mat_ok and info_mat.get("mat_error"):
        reg["notas"].append(
            f"⚠ Materiales NO incluidos ({info_mat['mat_error']}) — planilla solo MO")
    elif mat_ok:
        reg["notas"].append(
            "⚠ Materiales: ningún nodo SGS calzó con los nodos aprobados — solo MO")
    elif INCLUIR_MAT and reg.get("mat_aprobado") is None:
        nota = ("MAT pendiente de tu visto bueno: usa 'pre-análisis MAT' "
                "(planilla física vs SGS, revisa devoluciones) — por ahora solo MO")
        if nota not in reg.get("notas", []):
            reg["notas"].append(nota)
    _registrar("centro_mo_planilla_generada", reg, archivo=nombre,
               detalle={"sha256": _sha256(ruta), "bytes": len(excel), "filas_mo": len(filas),
                        "fecha_ejecucion": reg.get("fecha_ejecucion"), "solo_pendientes": solo_pendientes,
                        "mat": bool(info_mat.get("mat_nodos")), "sid": sid})
    return ""


def _intentar_generar(reg: dict):
    if reg.get("planilla_nativa"):
        return
    if not AUTO_GENERAR or reg.get("estado") != "listo_auto":
        return
    err = generar_planilla(reg)
    if err:
        reg.setdefault("notas", []).append(f"Autogeneración falló: {err}")
        reg["carril"] = "semi"
        reg["estado"] = "falta_info"
        reg.setdefault("faltantes", []).append(f"revisión manual ({err})")


# ── Pedidos al agente del PC ─────────────────────────────────────────────────

def _agente_conectado() -> bool:
    return (time.time() - _agente["ultimo_latido"]) < AGENTE_TIMEOUT


def _visto_hace() -> int:
    return int(time.time() - _agente["ultimo_latido"]) if _agente["ultimo_latido"] else -1


def _crear_pedido(tipo: str, reg: dict = None, **datos) -> dict:
    pid = uuid.uuid4().hex[:10]
    p = {"id": pid, "tipo": tipo, "rid": (reg or {}).get("id"), "creado": _ahora(),
         "estado": "pendiente", "resultado": None, "error": None, "datos": datos}
    if reg:
        p["entry_id"] = reg.get("entry_id")
        p["store_id"] = reg.get("store_id")
        p["carpeta"] = reg.get("carpeta", "")
    with _lock:
        _pedidos[pid] = p
        _guardar_pedidos()
    return p


def _caducar_pedidos():
    """Pedidos tomados por el agente y nunca cerrados (el agente murió a mitad) → error.
    Pedidos cerrados se olvidan a las 6 h."""
    ahora = time.time()
    cambio = False
    with _lock:
        for p in list(_pedidos.values()):
            t = p.get("ts_tomado") or 0
            if p["estado"] == "tomado" and t and ahora - t > PEDIDO_TIMEOUT:
                p["estado"] = "error"
                p["error"] = "el agente del PC tomó el pedido pero no respondió"
                cambio = True
            elif p["estado"] in ("hecho", "error") and ahora - (p.get("ts_fin") or ahora) > 6 * 3600:
                _pedidos.pop(p["id"], None)
                cambio = True
        if cambio:
            _guardar_pedidos()


def _respuesta_pedido(p: dict, extra: dict = None):
    d = {"ok": True, "pedido": p["id"], "tipo": p["tipo"]}
    d.update(extra or {})
    return jsonify(d)


def _exige_agente():
    if not _agente_conectado():
        return jsonify({"error": "El agente del PC no está conectado — arranca el agente del PC "
                                 "(agentes-pc/centro-mo) y reintenta"}), 503
    return None


def _agente_auth(f):
    @wraps(f)
    def w(*a, **k):
        tok = request.headers.get("X-Agente-Token", "")
        if not TOKEN_AGENTE or not hmac.compare_digest(tok, TOKEN_AGENTE):
            return jsonify({"error": "no autorizado"}), 401
        _agente["ultimo_latido"] = time.time()
        return f(*a, **k)
    return w


# ── Canal del agente ──────────────────────────────────────────────────────────

@app.route("/agente/estado")
@_agente_auth
def agente_estado():
    """Lo que el agente necesita para una vuelta: qué correos ya conocemos (para no reanalizar),
    qué hilos vigilar (consultas abiertas), pedidos pendientes y si hay que escanear ya."""
    _caducar_pedidos()
    _agente["version"] = request.args.get("version", "") or _agente["version"]
    with _lock:
        rids = {rid: r.get("estado", "") for rid, r in _store.items()}
        vigilar = [{"id": r["id"], "entry_id": r.get("entry_id"), "conversacion": r.get("conversacion"),
                    "asunto": r.get("asunto"), "consulta_en": r.get("consulta_en", "")}
                   for r in _store.values()
                   if r.get("consulta_en") and r.get("estado") in ("esperando_supervisor", "falta_info")]
        pend = [p for p in _pedidos.values() if p["estado"] == "pendiente"]
        for p in pend:
            p["estado"] = "tomado"
            p["ts_tomado"] = time.time()
        if pend:
            _guardar_pedidos()
        scan_ahora = _agente["scan_ahora"]
    return jsonify({"rids": rids, "vigilar": vigilar, "pedidos": pend, "scan_ahora": scan_ahora,
                    "config": {"poll_segundos": POLL_SEGUNDOS, "dias_atras": DIAS_ATRAS}})


@app.route("/agente/scan", methods=["POST"])
@_agente_auth
def agente_scan():
    d = request.get_json(force=True) or {}
    if d.get("fase") == "inicio":
        _agente["scaneando"] = True
        _agente["scan_ahora"] = False
        _agente["error"] = ""
    else:
        _agente["scaneando"] = False
        _agente["ultimo_scan"] = datetime.now().strftime("%H:%M:%S")
        _agente["error"] = (d.get("error") or "")[:300]
        _agente["nuevos_ultimo_scan"] = int(d.get("nuevos") or 0)
    return jsonify({"ok": True})


def _guardar_nativa(reg: dict) -> None:
    """La planilla del sistema que venía adjunta en el correo llega en base64 desde el agente."""
    b64 = reg.pop("excel_nativa_b64", None)
    nombre = reg.pop("excel_nativa_nombre", None) or "balance_sgs_manual.xlsx"
    reg.pop("excel_nativa_ruta", None)
    if not b64:
        return
    try:
        datos = base64.b64decode(b64)
        ruta = _ruta_planilla(reg["id"], nombre)
        with open(ruta, "wb") as f:
            f.write(datos)
        reg["excel_ruta"] = ruta
        reg["planilla_nativa"] = True
        _registrar("centro_mo_planilla_del_correo", reg, archivo=os.path.basename(ruta),
                   detalle={"sha256": _sha256(ruta), "bytes": len(datos)})
    except Exception as e:
        reg.setdefault("notas", []).append(f"Planilla del sistema: no se pudo guardar ({e})")


@app.route("/agente/ingesta", methods=["POST"])
@_agente_auth
def agente_ingesta():
    """Correos analizados por el agente (nuevos, reabiertos o reanalizados) y respuestas de
    supervisores. Aquí corren las reglas, la verificación SGS y la autogeneración, igual que
    hacía el hilo vigilante en el PC."""
    d = request.get_json(force=True) or {}
    nuevos = reabiertos = fusionados = respondidos = 0
    for reg in d.get("registros") or []:
        if not isinstance(reg, dict) or not reg.get("id"):
            continue
        motivo = reg.pop("motivo", "nuevo")
        _guardar_nativa(reg)
        with _lock:
            previo = _store.get(reg["id"])
        if previo is None:
            reg["estado"] = "listo_auto" if reg.get("carril") == "super" else "falta_info"
            reg["detectado_en"] = _ahora()
            with _lock:
                recalcular(reg)
            _verificar_liq(reg)
            _intentar_generar(reg)
            with _lock:
                _store[reg["id"]] = reg
            nuevos += 1
            _registrar("centro_mo_correo_detectado", reg,
                       detalle={"carril": reg.get("carril"), "estado": reg.get("estado"),
                                "faltantes": reg.get("faltantes"), "fecha_correo": reg.get("fecha")})
        elif previo.get("estado") == "descartado" and motivo != "reanalisis":
            with _lock:
                previo["estado"] = "falta_info"
                _merge_analisis(previo, reg)
                previo.setdefault("notas", []).append("Reabierto: la etiqueta volvió a aparecer en Outlook")
            _verificar_liq(previo)
            with _lock:
                _intentar_generar(previo)
            reabiertos += 1
        else:
            with _lock:
                _merge_analisis(previo, reg)
            _verificar_liq(previo)
            with _lock:
                _intentar_generar(previo)
            fusionados += 1
    for resp in d.get("respuestas") or []:
        with _lock:
            reg = _store.get(resp.get("id"))
        if not reg or not isinstance(resp.get("analisis"), dict):
            continue
        nuevo = resp["analisis"]
        _guardar_nativa(nuevo)
        f_h = resp.get("f_h") or _ahora()
        with _lock:
            reg["estado"] = "falta_info"
            _merge_analisis(reg, nuevo)
            reg["consulta_en"] = f_h
            reg.setdefault("notas", []).append(
                f"Respuesta detectada de {resp.get('remitente') or '?'} ({f_h}) — reanalizado")
        _verificar_liq(reg)
        with _lock:
            _intentar_generar(reg)
        respondidos += 1
        _registrar("centro_mo_respuesta_supervisor", reg,
                   detalle={"remitente": resp.get("remitente"), "fecha": f_h, "estado": reg.get("estado")})
    with _lock:
        _guardar_store()
    return jsonify({"ok": True, "nuevos": nuevos, "reabiertos": reabiertos,
                    "fusionados": fusionados, "respuestas": respondidos})


def _aplicar_datos_excel(reg: dict, datos: dict, etiqueta: str):
    with _lock:
        if datos.get("filas"):
            reg["filas_mo"] = [{"nodo": n, "codigo": c, "cantidad": q} for n, c, q in datos["filas"]]
            reg["adjunto"] = etiqueta
        for lid in datos.get("lideres", []):
            if lid.upper() not in {x.upper() for x in reg.get("lideres", [])}:
                reg.setdefault("lideres", []).append(lid)
        for ld in datos.get("lideres_det", []):
            if ld["nombre"].upper() not in {x["nombre"].upper() for x in reg.get("lideres_det", [])}:
                reg.setdefault("lideres_det", []).append(ld)
        meta = (datos.get("codigo") or "").upper()
        if meta:
            if not reg.get("codigo"):
                reg["codigo"] = meta
            elif meta != reg["codigo"] and not reg.get("pdl"):
                reg["pdl"] = meta
        if datos.get("fecha_ejecucion") and not reg.get("fecha_ejecucion"):
            reg["fecha_ejecucion"] = datos["fecha_ejecucion"]
        reg.setdefault("notas", []).extend(datos.get("errores", []))
        recalcular(reg)
        _intentar_generar(reg)
        _guardar_store()


@app.route("/agente/excel/<rid>", methods=["POST"])
@_agente_auth
def agente_excel(rid):
    """Planilla traída por el agente desde el link de SharePoint/Drive del correo."""
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    f = request.files.get("archivo")
    if not f or not (f.filename or "").lower().endswith((".xlsx", ".xlsm")):
        return jsonify({"error": "Se esperaba un .xlsx o .xlsm"}), 400
    ruta = _ruta_planilla(rid, f.filename)
    f.save(ruta)
    try:
        datos = bridge.extraer_de_excel(ruta)
    except Exception as e:
        return jsonify({"error": f"Descargó pero no se pudo leer: {e}"}), 500
    _aplicar_datos_excel(reg, datos, os.path.basename(ruta) + " (del link)")
    _registrar("centro_mo_planilla_del_link", reg, archivo=os.path.basename(ruta),
               detalle={"sha256": _sha256(ruta), "filas": len(reg.get("filas_mo", []))})
    return jsonify({"ok": True, "archivo": os.path.basename(ruta), "filas": len(reg.get("filas_mo", []))})


@app.route("/agente/instruccion-mat", methods=["POST"])
@_agente_auth
def agente_instruccion_mat():
    """Barrido de la instrucción de MATERIAL de correos ya conocidos (el agente relee el
    cuerpo en Outlook). Solo escribe `instruccion_mat`; no reanaliza ni regenera nada."""
    d = request.get_json(force=True) or {}
    items = d.get("items") or {}
    puestos, ignorados = 0, []
    with _lock:
        for rid, im in items.items():
            reg = _store.get(rid)
            if not reg or not isinstance(im, dict) or im.get("valor") not in _IM_VALORES:
                ignorados.append(rid)
                continue
            reg["instruccion_mat"] = {k: im.get(k) for k in _IM_CAMPOS}
            puestos += 1
        _guardar_store()
    _registrar("centro_mo_instruccion_mat", None,
               detalle={"puestos": puestos, "ignorados": len(ignorados)})
    return jsonify({"ok": True, "puestos": puestos, "ignorados": ignorados[:20]})


@app.route("/agente/pedido/<pid>", methods=["POST"])
@_agente_auth
def agente_pedido(pid):
    """El agente cierra un pedido: hecho (con resultado) o error. Las consecuencias sobre el
    registro se aplican aquí (p. ej. consulta abierta → esperando al supervisor)."""
    d = request.get_json(force=True) or {}
    with _lock:
        p = _pedidos.get(pid)
        if not p:
            return jsonify({"error": "pedido desconocido"}), 404
        p["estado"] = "hecho" if d.get("estado") == "hecho" else "error"
        p["resultado"] = d.get("resultado")
        p["error"] = (d.get("error") or None)
        p["ts_fin"] = time.time()
        reg = _store.get(p.get("rid") or "")
        if reg is not None:
            if p["estado"] == "hecho" and p["tipo"] == "consulta":
                reg["estado"] = "esperando_supervisor"
                reg["carril"] = "semi"
                reg["consulta_en"] = _ahora()
                reg.setdefault("notas", []).append(
                    "Borrador de consulta abierto en Outlook (PC) — el envío es manual. "
                    "El agente reanaliza cuando llegue respuesta en la conversación.")
            elif p["estado"] == "error" and p["tipo"] != "abrir":
                reg.setdefault("notas", []).append(f"⚠ {p['tipo']} (agente del PC): {p['error']}")
            _guardar_store()
        _guardar_pedidos()
    if p["estado"] == "hecho" and p["tipo"] == "consulta":
        _registrar("centro_mo_consulta_abierta", reg, detalle={"pedido": pid})
    return jsonify({"ok": True})


# ── API de la interfaz (tras la puerta única) ─────────────────────────────────

@app.route("/api/items")
def api_items():
    _caducar_pedidos()
    with _lock:
        items = sorted(_store.values(), key=lambda r: r.get("fecha", ""), reverse=True)
    return jsonify({"items": items,
                    "agente": {"conectado": _agente_conectado(), "visto_hace_s": _visto_hace(),
                               "ultimo_scan": _agente["ultimo_scan"], "scaneando": _agente["scaneando"],
                               "error": _agente["error"], "version": _agente["version"],
                               "cada_s": POLL_SEGUNDOS},
                    "cfg": {"balance": BALANCE_PUBLIC, "liquidador2": LIQUIDADOR2_PUBLIC,
                            "fecha_minima": fecha_minima_liq()}})


@app.route("/api/pedido/<pid>")
def api_pedido(pid):
    with _lock:
        p = _pedidos.get(pid)
    if not p:
        return jsonify({"error": "pedido desconocido"}), 404
    return jsonify({k: p.get(k) for k in ("id", "tipo", "estado", "resultado", "error", "creado")})


@app.route("/api/reanalizar-todo", methods=["POST"])
def api_reanalizar_todo():
    _agente["scan_ahora"] = True

    def _chequear_liq_todos():
        with _lock:
            regs = [r for r in _store.values()
                    if r.get("estado") not in ("completado", "descartado") and r.get("codigo")]
        for r in regs:
            try:
                _verificar_liq(r)
            except Exception:
                pass
        with _lock:
            _guardar_store()

    threading.Thread(target=_chequear_liq_todos, daemon=True).start()
    return jsonify({"ok": True, "agente": _agente_conectado()})


@app.route("/api/reanalizar/<rid>", methods=["POST"])
def api_reanalizar(rid):
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    if not reg.get("entry_id"):
        return jsonify({"error": "Este registro no tiene referencia al correo de Outlook"}), 400
    falta = _exige_agente()
    if falta:
        return falta
    return _respuesta_pedido(_crear_pedido("reanalizar", reg))


@app.route("/api/editar/<rid>", methods=["POST"])
def api_editar(rid):
    d = request.get_json(force=True) or {}
    with _lock:
        reg = _store.get(rid)
        if not reg:
            return jsonify({"error": "No existe"}), 404
        if d.get("codigo"):
            reg["codigo"] = d["codigo"].strip().upper()
        if d.get("pdl"):
            reg["pdl"] = d["pdl"].strip().upper()
        if d.get("lider") is not None:
            lid = d["lider"].strip()
            reg["lideres"] = [lid] if lid else []
            reg["lideres_det"] = []
        if d.get("nodos") is not None:
            reg["nodos_filtro"] = [x.strip().upper() for x in d["nodos"].replace(";", ",").split(",") if x.strip()]
        if d.get("fecha"):
            f = d["fecha"].strip()
            m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$", f)
            if m:
                f = f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
            reg["fecha_ejecucion"] = f
        if d.get("borrar_fecha"):
            reg["fecha_ejecucion"] = ""
            reg.pop("fecha_original", None)
            reg.setdefault("notas", []).append("Fecha de ejecución borrada — vuelve a ser faltante")
        recalcular(reg)
        _intentar_generar(reg)
        _guardar_store()
        return jsonify({"ok": True})


@app.route("/api/subir-excel/<rid>", methods=["POST"])
def api_subir_excel(rid):
    """Caso SharePoint/Drive: la planilla llega a mano desde el navegador."""
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    f = request.files.get("archivo")
    if not f or not (f.filename or "").lower().endswith((".xlsx", ".xlsm")):
        return jsonify({"error": "Sube un .xlsx o .xlsm"}), 400
    ruta = _ruta_planilla(rid, f.filename)
    try:
        f.save(ruta)
        datos = bridge.extraer_de_excel(ruta)
    except Exception as e:
        return jsonify({"error": f"No se pudo leer: {e}"}), 400
    _aplicar_datos_excel(reg, datos, f.filename + " (subida manual)")
    return jsonify({"ok": True, "filas": len(reg.get("filas_mo", []))})


@app.route("/api/traer-link/<rid>", methods=["POST"])
def api_traer_link(rid):
    """El link de SharePoint solo se puede abrir con el perfil de Microsoft del PC: pedido al agente."""
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    links = reg.get("links_externos") or []
    if not links:
        return jsonify({"error": "Este correo no tiene link de SharePoint/Drive detectado"}), 400
    falta = _exige_agente()
    if falta:
        return falta
    return _respuesta_pedido(_crear_pedido("traer_link", reg, url=links[0]))


@app.route("/api/sp-login", methods=["POST"])
def api_sp_login():
    falta = _exige_agente()
    if falta:
        return falta
    return _respuesta_pedido(_crear_pedido("sp_login"))


@app.route("/api/generar/<rid>", methods=["POST"])
def api_generar(rid):
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    solo_pend = bool((request.get_json(silent=True) or {}).get("solo_pendientes"))
    err = generar_planilla(reg, solo_pendientes=solo_pend)
    with _lock:
        _guardar_store()
    if err:
        return jsonify({"error": err}), 400
    return jsonify({"ok": True, "archivo": os.path.basename(reg.get("excel_ruta", ""))})


@app.route("/api/descargar/<rid>")
def api_descargar(rid):
    with _lock:
        reg = _store.get(rid)
    ruta = (reg or {}).get("excel_ruta", "")
    if not ruta or not os.path.exists(ruta):
        return jsonify({"error": "La planilla no existe en el servidor — genérala de nuevo ()"}), 404
    return send_file(ruta, as_attachment=True, download_name=os.path.basename(ruta))


def _nombre_de_pila(nombre_completo: str) -> str:
    partes = [p for p in (nombre_completo or "").split() if p]
    if len(partes) >= 4:
        return " ".join(partes[:2]).title()
    return (partes[0].title() if partes else "").strip()


_PEDIDOS_FALTANTE = [
    ("planilla", "adjuntar la planilla de prestaciones (mano de obra)"),
    ("líder", "confirmarme el líder responsable de la ejecución"),
    ("fecha", "indicarme la fecha real de ejecución de los trabajos"),
    ("PDL", "indicarme el número de PDL"),
    ("código", "indicarme el código del proyecto (DF o LCL)"),
]


def _redactar_pedidos(faltantes: list) -> list:
    pedidos = []
    for f in faltantes:
        for clave, texto in _PEDIDOS_FALTANTE:
            if clave.lower() in f.lower():
                if texto not in pedidos:
                    pedidos.append(texto)
                break
        else:
            pedidos.append(f"confirmarme: {f}")
    return pedidos


@app.route("/api/consulta/<rid>", methods=["POST"])
def api_consulta(rid):
    """Borrador de respuesta al supervisor pidiendo lo que falta. Se redacta aquí y el agente
    lo abre en Outlook (ReplyAll). NO envía: el envío sigue siendo manual en el PC."""
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    if not reg.get("entry_id"):
        return jsonify({"error": "Este registro no tiene referencia al correo de Outlook"}), 400
    falta = _exige_agente()
    if falta:
        return falta
    pedidos = _redactar_pedidos(reg.get("faltantes") or ["la información del proyecto"])
    nombre = _nombre_de_pila(reg.get("remitente", ""))
    saludo = f"Buen día {nombre}," if nombre else "Buen día,"
    if len(pedidos) == 1:
        cuerpo = (f"<p>{saludo}</p>"
                  f"<p>Para proceder correctamente con la liquidación de la mano de obra de este "
                  f"proyecto, le agradezco {pedidos[0]}.</p>"
                  f"<p>Quedo atento a sus comentarios.</p>"
                  f"<p>Cordial saludo,</p>")
    else:
        items_html = "".join(f"<li>{p}</li>" for p in pedidos)
        cuerpo = (f"<p>{saludo}</p>"
                  f"<p>Para proceder correctamente con la liquidación de la mano de obra de este "
                  f"proyecto, le agradezco:</p><ul>{items_html}</ul>"
                  f"<p>Quedo atento a sus comentarios.</p>"
                  f"<p>Cordial saludo,</p>")
    return _respuesta_pedido(_crear_pedido("consulta", reg, cuerpo=cuerpo))


@app.route("/api/lideres-sgs/<rid>", methods=["POST"])
def api_lideres_sgs(rid):
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    if not reg.get("codigo"):
        return jsonify({"error": "Primero necesita un código DF/LCL validado"}), 400
    det, err = bridge.lideres_de_bandejas(reg["codigo"], reg.get("nodos_filtro"))
    if err:
        return jsonify({"error": err}), 502
    if not det:
        return jsonify({"error": "SGS no tiene bandejas con líder para este "
                                 "proyecto (aún sin despachos de materiales)"}), 404
    return jsonify({"ok": True, "lideres": det})


@app.route("/api/fecha-bandejas/<rid>", methods=["POST"])
def api_fecha_bandejas(rid):
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    if not reg.get("codigo"):
        return jsonify({"error": "Primero necesita un código DF/LCL validado"}), 400
    res, err = bridge.fecha_sugerida_bandejas(reg["codigo"], reg.get("nodos_filtro"))
    if err:
        return jsonify({"error": err}), 502
    return jsonify({"ok": True, **res})


@app.route("/api/aplicar-lideres/<rid>", methods=["POST"])
def api_aplicar_lideres(rid):
    d = request.get_json(force=True) or {}
    sel = [x for x in (d.get("lideres") or []) if x.get("nombre")]
    if not sel:
        return jsonify({"error": "No elegiste ningún líder"}), 400
    with _lock:
        reg = _store.get(rid)
        if not reg:
            return jsonify({"error": "No existe"}), 404
        n = len(sel)
        base_pct = round(100.0 / n, 1)
        for i, x in enumerate(sel):
            x["pct"] = round(100.0 - base_pct * (n - 1), 1) if i == n - 1 else base_pct
            x.setdefault("cedula", "")
        reg["lideres"] = [x["nombre"] for x in sel]
        reg["lideres_det"] = sel
        reg.setdefault("notas", []).append(
            "Líderes elegidos de las bandejas de SGS: " + ", ".join(reg["lideres"]))
        recalcular(reg)
        _intentar_generar(reg)
        _guardar_store()
    return jsonify({"ok": True, "lideres": sel})


@app.route("/api/mat-preanalisis/<rid>", methods=["POST"])
def api_mat_preanalisis(rid):
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    cod = reg.get("codigo")
    if not cod:
        return jsonify({"error": "La tarjeta aún no tiene código de proyecto"}), 400
    nodos_obj = reg.get("nodos_filtro") or [f["nodo"] for f in reg.get("filas_mo", [])]
    sgs, err = bridge.cruce_materiales(cod, nodos_obj or None)
    if err:
        return jsonify({"error": err}), 502
    fisica = []
    for f in (reg.get("filas_mat") or []):
        fisica.append({"nodos": bridge._expandir_nodos_planilla([f["nodo"]]),
                       "nodo_txt": str(f["nodo"]), "cod": str(f["codigo"]),
                       "cant": f.get("cantidad") or 0, "usada": False})
    for bloque in sgs:
        n_norm = _clave_nodo(bloque["nodo"])
        for fila in bloque["filas"]:
            total = None
            for ff in fisica:
                if ff["cod"] == fila["cod"] and n_norm in ff["nodos"]:
                    total = (total or 0) + ff["cant"]
                    ff["usada"] = True
            fila["fisica"] = total
        bloque["alerta_devol"] = bloque["devoluciones"] > 0
    solo_fisica = [{"nodo": ff["nodo_txt"], "cod": ff["cod"], "cant": ff["cant"]}
                   for ff in fisica if not ff["usada"]]
    return jsonify({"nodos": sgs, "solo_fisica": solo_fisica,
                    "fisica_filas": len(fisica),
                    "mat_aprobado": reg.get("mat_aprobado"),
                    "mat_nodos_sel": reg.get("mat_nodos_sel") or []})


@app.route("/api/mat-aprobar/<rid>", methods=["POST"])
def api_mat_aprobar(rid):
    d = request.get_json(force=True) or {}
    aprobar = bool(d.get("aprobar"))
    nodos = [str(x).strip() for x in (d.get("nodos") or []) if str(x).strip()]
    with _lock:
        reg = _store.get(rid)
        if not reg:
            return jsonify({"error": "No existe"}), 404
        reg["mat_aprobado"] = aprobar
        reg["mat_nodos_sel"] = nodos if aprobar else []
        reg.setdefault("notas", []).append(
            f"MAT APROBADO (nodos: {', '.join(nodos) if nodos else 'los de la planilla'}) "
            "— planilla regenerada con materiales" if aprobar
            else "MAT descartado — planilla solo MO")
        err = ""
        if not reg.get("planilla_nativa"):
            err = generar_planilla(reg)
            if err:
                reg["notas"].append(f"Regeneración tras decisión MAT falló: {err}")
        recalcular(reg)
        _guardar_store()
    return jsonify({"ok": not err, "error": err or None})


@app.route("/api/enviar-liq2/<rid>", methods=["POST"])
def api_enviar_liq2(rid):
    """Embudo directo: sube la planilla a liquidador2 (mismo servidor) con su fecha de ejecución.
    Se le pasa el usuario de la puerta (X-SGO7-*) para que la sesión quede a su nombre y
    liquide con SU cuenta de SGS. La liquidación real sigue siendo un acto manual allá."""
    import requests as rq
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    ruta = reg.get("excel_ruta", "")
    if not ruta or not os.path.exists(ruta):
        return jsonify({"error": "Primero genera la planilla ()"}), 400
    cabeceras = {k: v for k, v in (("X-SGO7-User", request.headers.get("X-SGO7-User")),
                                   ("X-SGO7-Rol", request.headers.get("X-SGO7-Rol")),
                                   ("X-SGO7-Nombre", request.headers.get("X-SGO7-Nombre")),
                                   ("X-Real-IP", request.headers.get("X-Real-IP"))) if v}
    try:
        with open(ruta, "rb") as f:
            r = rq.post(
                f"{LIQUIDADOR2_BASE}/subir",
                files={"excel": (os.path.basename(ruta), f,
                                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                data={"fecha_planilla": reg.get("fecha_ejecucion", "")},
                headers=cabeceras, allow_redirects=False, timeout=180)
    except rq.exceptions.ConnectionError:
        return jsonify({"error": "liquidador2 no responde (proceso `liquidador` en PM2) — reintenta en un momento"}), 502
    except Exception as e:
        return jsonify({"error": f"Error enviando: {str(e)[:120]}"}), 500

    loc = r.headers.get("Location", "")
    if r.status_code in (301, 302, 303, 307) and "verificar" in loc:
        path = urlparse(loc).path if loc.startswith("http") else loc
        prefijo = urlparse(LIQUIDADOR2_BASE).path.rstrip("/")
        if prefijo and path.startswith(prefijo):
            path = path[len(prefijo):]
        url = LIQUIDADOR2_PUBLIC.rstrip("/") + path
        with _lock:
            reg["liq2_url"] = url
            reg.setdefault("notas", []).append(
                f"Enviada a liquidador2 (fecha {reg.get('fecha_ejecucion')}) — verificar y liquidar allá")
            _guardar_store()
        _registrar("centro_mo_enviada_liquidador2", reg, archivo=os.path.basename(ruta),
                   detalle={"url": url, "fecha_ejecucion": reg.get("fecha_ejecucion")})
        return jsonify({"ok": True, "url": url})
    return jsonify({"error": f"liquidador2 no aceptó la planilla "
                             f"(HTTP {r.status_code} → {loc or 'sin redirección'}) — "
                             f"revisa el mensaje en su página de subir"}), 500


@app.route("/api/abrir-outlook/<rid>", methods=["POST"])
def api_abrir_outlook(rid):
    with _lock:
        reg = _store.get(rid)
    if not reg:
        return jsonify({"error": "No existe"}), 404
    if not reg.get("entry_id"):
        return jsonify({"error": "Este registro no tiene referencia al correo de Outlook"}), 400
    falta = _exige_agente()
    if falta:
        return falta
    return _respuesta_pedido(_crear_pedido("abrir", reg))


@app.route("/api/estado/<rid>", methods=["POST"])
def api_estado(rid):
    nuevo = (request.get_json(force=True) or {}).get("estado", "")
    if nuevo not in ("completado", "descartado", "falta_info"):
        return jsonify({"error": "Estado inválido"}), 400
    with _lock:
        reg = _store.get(rid)
        if not reg:
            return jsonify({"error": "No existe"}), 404
        reg["estado"] = nuevo
        if nuevo == "completado":
            reg["completado_en"] = _ahora()
        if nuevo == "falta_info":
            recalcular(reg)
        _guardar_store()
    return jsonify({"ok": True})


@app.route("/health")
def health():
    with _lock:
        n = len(_store)
    return jsonify({"ok": True, "registros": n, "agente": _agente_conectado()})


# ── UI ────────────────────────────────────────────────────────────────────────
with open(os.path.join(DIR_SCRIPT, "ui.html"), encoding="utf-8") as _f:
    PAGINA = _f.read()


@app.route("/")
def index():
    return PAGINA


def _arranque():
    _cargar_store()
    _cargar_pedidos()
    with _lock:
        for _r in _store.values():
            if _r.get("estado") not in ("completado", "descartado"):
                recalcular(_r)
        _guardar_store()
    print(f"[Centro MO] {len(_store)} registros · planillas en {DIR_PLANILLAS} · "
          f"agente {'con' if TOKEN_AGENTE else 'SIN'} token", flush=True)


# Gestión General (2026-10-22): el Centro MO responde como árbitro de SGS (solo lectura) por canal interno.
from gestion_arbitro import registrar as _registrar_gestion   # noqa: E402
_registrar_gestion(app, lambda: _store, _lock)

_arranque()

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=PUERTO, debug=False)
