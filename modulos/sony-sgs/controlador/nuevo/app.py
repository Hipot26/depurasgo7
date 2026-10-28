#!/usr/bin/env python3
"""
/nuevo — Balance Integrado SIPWEBORA + SGS
==========================================
App Flask independiente. No modifica /pruebas ni sgs_api.

Uso (local):
  cd nuevo
  python app.py

Luego abre: http://localhost:5002
"""

import io
import json as _json
import os
import sys
import tempfile
import uuid
from datetime import datetime, timezone, timedelta
from functools import wraps

from flask import Flask, render_template, request, send_file, redirect, url_for, session
from flask import render_template_string
# Asegurar que nuevo/ esté en sys.path para que gunicorn pueda importar integracion.py
_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

app = Flask(__name__, template_folder=os.path.join(_HERE, "templates"))

# ── SGO7: trazabilidad — todo archivo subido, proceso o descarga queda registrado en sabogal ──
try:
    import sgo7_registro
    sgo7_registro.instalar(app, modulo="balance", acciones={
            r"/casos/\d+/mandar-a-liquidar": "balance_mandar_liquidar",
            r"/casos/\d+/estado": "balance_caso_estado",
            r"/generar-sgs": "balance_sgs",
            r"/generar-final": "balance_final",
            r"/generar-equivalencia": "balance_equivalencia",
            r"/generar": "balance_integrado",
            r"/sgs-libre": "balance_sgs_lcl",
            r"/manual/importar-excel": "balance_importar_excel",
            r"/manual": "balance_manual",
            r"/rg03/analizar": "rg03_analizar",
            r"/rg03": "rg03",
            r"/materiales/anexo": "materiales_anexo",
            r"/materiales/buscar": "materiales_buscar",
            r"/materiales": "materiales_cargue",
            r"/reiniciar": "reiniciar_servidor",
            r"/descargar": "descargar"
    })
except Exception as _e_sgo7:
    print("sgo7_registro no instalado:", _e_sgo7)
app.secret_key = os.environ.get("NUEVO_SECRET", "")

# ── Hora Colombia (UTC-5, sin horario de verano) ───────────────────────────────
# El servidor corre en UTC; todo timestamp visible al usuario va en hora CO.
TZ_CO = timezone(timedelta(hours=-5))


def ahora_co() -> datetime:
    return datetime.now(TZ_CO).replace(tzinfo=None)

# ── Credenciales (variables de entorno obligatorias en producción) ─────────────
_AUTH_USER = os.environ.get("NUEVO_USER", "")
_AUTH_PASS = os.environ.get("NUEVO_PASS", "")
# Perfiles adicionales (2026-09-16, Módulo Almacén): NUEVO_USERS="user:clave:rol;..."
# rol admin = todo · rol almacen = SOLO el tab Materiales Sony (su cargue SIPWEBORA).
_RUTAS_ALMACEN = ("/materiales", "/descargar/", "/logout", "/login", "/health")


def _usuarios_extra() -> dict:
    out = {}
    for item in os.environ.get("NUEVO_USERS", "").split(";"):
        partes = item.strip().split(":", 2)
        if len(partes) == 3 and partes[0]:
            out[partes[0].strip()] = {"clave": partes[1], "rol": partes[2].strip().lower() or "almacen"}
    return out


def _check_auth(usuario: str, password: str):
    """Devuelve el rol si las credenciales son válidas, si no None."""
    import hmac
    if _AUTH_USER and usuario == _AUTH_USER and hmac.compare_digest(password or "", _AUTH_PASS):
        return "admin"
    u = _usuarios_extra().get(usuario)
    if u and hmac.compare_digest(password or "", u["clave"]):
        return u["rol"]
    return None


def _rol() -> str:
    return session.get("rol", "admin") if session.get("autenticado") else ""


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("autenticado"):
            return redirect(url_for("login", next=request.path))
        if _rol() == "almacen" and not request.path.startswith(_RUTAS_ALMACEN) \
                and request.path != "/":
            return redirect(url_for("index") + "?tab=materiales")
        return f(*args, **kwargs)
    return decorated


# Almacén de archivos generados: guardado en disco para sobrevivir entre workers Gunicorn.
_ARCHIVOS_DIR = os.path.join(tempfile.gettempdir(), "nuevo_archivos")
os.makedirs(_ARCHIVOS_DIR, exist_ok=True)


def _save_archivo(token: str, excel_bytes: bytes, filename: str, advertencias: list = None):
    with open(os.path.join(_ARCHIVOS_DIR, f"{token}.xlsx"), "wb") as _f:
        _f.write(excel_bytes)
    with open(os.path.join(_ARCHIVOS_DIR, f"{token}.json"), "w", encoding="utf-8") as _f:
        _json.dump({"filename": filename, "advertencias": advertencias or []}, _f)


def _get_archivo(token: str):
    xls  = os.path.join(_ARCHIVOS_DIR, f"{token}.xlsx")
    meta = os.path.join(_ARCHIVOS_DIR, f"{token}.json")
    if not os.path.isfile(xls) or not os.path.isfile(meta):
        return None
    with open(meta, encoding="utf-8") as _f:
        m = _json.load(_f)
    with open(xls, "rb") as _f:
        return {"bytes": _f.read(), "filename": m["filename"], "advertencias": m.get("advertencias", [])}


# ── Store de previsualización (datos del cruce entre Fase 1 y la generación) ────
import pickle as _pickle

_PREVIEW_DIR = os.path.join(tempfile.gettempdir(), "nuevo_preview")
os.makedirs(_PREVIEW_DIR, exist_ok=True)


def _save_preview(sid: str, data: dict):
    with open(os.path.join(_PREVIEW_DIR, f"{sid}.pkl"), "wb") as _f:
        _pickle.dump(data, _f)


def _get_preview(sid: str):
    p = os.path.join(_PREVIEW_DIR, f"{sid}.pkl")
    if not os.path.isfile(p):
        return None
    try:
        with open(p, "rb") as _f:
            return _pickle.load(_f)
    except Exception:
        return None


# ── Render de hojas xlsx → grilla HTML (para la previsualización fiel) ──────────

def _argb_to_css(argb) -> str:
    """Convierte un color ARGB de openpyxl a CSS #RRGGBB. Blanco/negro → vacío."""
    if not argb or not isinstance(argb, str) or len(argb) < 6:
        return ""
    s = argb[-6:].upper()
    if s in ("FFFFFF", "000000"):
        return ""
    return "#" + s


def _sheet_to_grid(ws, max_rows: int = 6000) -> list:
    """
    Convierte una hoja openpyxl en una lista de filas para render HTML.
    Cada celda: {v, bg, b (bold), s (colspan)}. Respeta merges horizontales
    y descarta filas totalmente vacías (sin valor ni color).
    """
    spans, skip = {}, set()
    for mr in ws.merged_cells.ranges:
        spans[(mr.min_row, mr.min_col)] = mr.max_col - mr.min_col + 1
        for rr in range(mr.min_row, mr.max_row + 1):
            for cc in range(mr.min_col, mr.max_col + 1):
                if (rr, cc) != (mr.min_row, mr.min_col):
                    skip.add((rr, cc))

    maxcol = ws.max_column
    rows = []
    for r in range(1, min(ws.max_row, max_rows) + 1):
        cells, c = [], 1
        while c <= maxcol:
            if (r, c) in skip:
                c += 1
                continue
            cell = ws.cell(row=r, column=c)
            bg = ""
            f = cell.fill
            if f is not None and f.patternType == "solid":
                bg = _argb_to_css(getattr(f.fgColor, "rgb", None))
            v = "" if cell.value is None else str(cell.value)
            cells.append({
                "v": v, "bg": bg,
                "b": bool(cell.font and cell.font.bold),
                "s": spans.get((r, c), 1),
            })
            c += 1
        if any(cl["v"] or cl["bg"] for cl in cells):
            rows.append(cells)
    return rows


def _xlsx_a_hojas(excel_bytes: bytes) -> list:
    """Lee el Excel base y devuelve [{titulo, filas}] para la previsualización."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False)
    hojas = []
    for ws in wb.worksheets:
        hojas.append({"titulo": ws.title, "filas": _sheet_to_grid(ws)})
    return hojas

# ── Cliente reutilizable (Balance Integrado) ───────────────────────────────────

_balance_cliente = None


def _get_cliente():
    global _balance_cliente
    from integracion import BalanceIntegrado, cfg_desde_archivo
    if _balance_cliente is not None:
        try:
            c = _balance_cliente._cliente
            c._call_kw("res.users", "read", [[c.uid]], {"fields": ["id"]})
            return _balance_cliente
        except Exception:
            _balance_cliente = None

    cfg = cfg_desde_archivo()
    _balance_cliente = BalanceIntegrado(**cfg)
    return _balance_cliente


# ── Cliente SGS puro (sin SIPWEBORA) ──────────────────────────────────────────
# Reutiliza el _cliente interno de BalanceIntegrado, que ya maneja auto-detección de DB.

def _get_sgs_cliente():
    return _get_cliente()._cliente


# ── Rutas ──────────────────────────────────────────────────────────────────────

# ── SGO7 (puerta única): nginx valida la sesión de sabogal y entrega el usuario en
#    cabeceras X-SGO7-* que solo él puede poner (esta app escucha en 127.0.0.1). ─────
_SGO7_ADMIN = ("desarrollador", "supervisor", "directivo", "facturacion")   # facturación desde 2026-10-10


@app.before_request
def _sgo7_sso():
    u = request.headers.get("X-SGO7-User")
    if not u:
        return
    if session.get("usuario") != u or not session.get("autenticado"):
        session.clear()
        session["autenticado"] = True
        session["usuario"] = u
        session["rol"] = "admin" if request.headers.get("X-SGO7-Rol") in _SGO7_ADMIN else "almacen"
        session.permanent = True


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("autenticado"):
        return redirect(url_for("index"))
    error = None
    if request.method == "POST":
        rol = _check_auth(request.form.get("usuario", ""), request.form.get("password", ""))
        if rol:
            session["autenticado"] = True
            session["rol"] = rol
            session["usuario"] = request.form.get("usuario", "")
            session.permanent = True
            next_url = request.form.get("next") or ("/?tab=materiales" if rol == "almacen" else "/")
            return redirect(next_url)
        error = "Usuario o contraseña incorrectos."
    return render_template("login.html", error=error, next=request.args.get("next", ""))


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def index():
    return render_template("index.html", anexo_info=_anexo_info_safe(),
                           tab="materiales" if _rol() == "almacen" else request.args.get("tab", ""),
                           # SGO7 facturación: «Reconsultar» desde el panel trae los datos y envía solo
                           id_planilla=request.args.get("id_planilla", "")[:120],
                           proyecto_id=request.args.get("proyecto_id", "")[:120],
                           auto=request.args.get("auto", "") == "1")


@app.route("/generar", methods=["POST"])
@login_required
def generar():
    id_planilla   = request.form.get("id_planilla",   "").strip()
    proyecto_id   = request.form.get("proyecto_id",   "").strip() or None

    ids_sip = [x.strip() for x in id_planilla.replace(";", ",").split(",") if x.strip()]
    if not ids_sip or not all(x.isdigit() for x in ids_sip):
        return render_template(
            "index.html",
            error="El ID de planilla debe ser un número, o varios números separados por coma.",
            id_planilla=id_planilla,
            proyecto_id=proyecto_id or "",
        )

    try:
        cliente = _get_cliente()
        # Fase 1: correr el cruce SIN generar Excel todavía (eso es la fase final).
        datos = cliente.preparar_cruce(id_planilla, proyecto_id_sgs=proyecto_id)

        # Excel base (sin fecha/líderes) — sirve para renderizar la previsualización
        # fiel a lo que será el archivo. Los Excel definitivos se generan en la fase 2.
        base_xlsx = cliente._generar_excel(
            datos["sip_datos"], datos["sgs_data"], datos["id_label"],
            advertencias=datos["advertencias"],
        )

        sid = uuid.uuid4().hex[:12]
        datos["_base_xlsx"]  = base_xlsx
        datos["id_planilla"] = id_planilla
        datos["proyecto_id"] = proyecto_id or ""
        # SGO7 facturación: la consulta queda guardada como caso con su semáforo (OK/descuadre/alertas)
        try:
            import sgo7_casos
            datos["caso"] = sgo7_casos.registrar_consulta(datos, base_xlsx, sid=sid)
        except Exception as _e:
            print("sgo7_casos: consulta no registrada:", _e)
            datos["caso"] = None
        _save_preview(sid, datos)
        return redirect(url_for("preview", sid=sid))

    except RuntimeError as e:
        return render_template(
            "index.html",
            error=str(e),
            id_planilla=id_planilla,
            proyecto_id=proyecto_id or "",
        )
    except Exception as e:
        return render_template(
            "index.html",
            error=f"Error inesperado: {e}",
            id_planilla=id_planilla,
            proyecto_id=proyecto_id or "",
        )


@app.route("/preview/<sid>")
@login_required
def preview(sid):
    """Fase 2: muestra las hojas del balance con buscador/filtros y pide
    fecha real de ejecución + líderes a liquidar antes de generar los Excel."""
    datos = _get_preview(sid)
    if not datos:
        return render_template(
            "index.html",
            error="La previsualización expiró. Genera el balance de nuevo.",
        ), 404

    hojas = _xlsx_a_hojas(datos["_base_xlsx"])
    caso_error = datos.pop("caso_error", None)   # SGO7 facturación: aviso de una sola vez
    if caso_error:
        _save_preview(sid, datos)
    return render_template(
        "preview.html",
        sid=sid,
        hojas=hojas,
        metadata=datos["metadata"],
        sgs_padre=datos["sgs_padre"],
        nodos=datos["nodos_resumen"],
        equivalencias=(datos.get("sgs_data") or {}).get("_equivalencias") or [],
        advertencias=datos.get("advertencias", []),
        caso=datos.get("caso"), caso_error=caso_error,   # SGO7 facturación
    )


# ── Equivalencias de nodo: la facturadora decide si es el mismo tramo ─────
_REENVIO_EQ = """<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Recalculando el cruce</title></head>
<body style="font-family:Segoe UI,system-ui,sans-serif;padding:40px;color:#1F3864">
<p><b>{{ msg }}</b></p><p>Recalculando el cruce con la decisión aplicada (mismas planillas y proyecto)…</p>
<form id="f" method="POST" action="/generar">
<input type="hidden" name="id_planilla" value="{{ id_planilla }}">
<input type="hidden" name="proyecto_id" value="{{ proyecto_id }}">
<noscript><button>Continuar</button></noscript></form>
<script>document.getElementById('f').submit();</script></body></html>"""


@app.route("/generar-equivalencia/<sid>", methods=["POST"])
@login_required
def generar_equivalencia(sid):
    """Un nodo de la planilla que no cruza por nombre y un nodo SGS del mismo LCL con los mismos
    materiales. La decisión queda guardada por LCL (data/equivalencias.json) y se aplica en toda consulta
    siguiente (también la reconsulta automática); aquí se vuelve a correr el cruce."""
    datos = _get_preview(sid)
    if not datos:
        return render_template("index.html", error="La previsualización expiró. Genera el balance de nuevo."), 404
    sip = (request.form.get("sip") or "")[:40]
    decision = request.form.get("decision") or ""
    sgs_data = datos.get("sgs_data") or {}
    eq = next((e for e in sgs_data.get("_equivalencias") or [] if e.get("sip") == sip), None)
    if (eq is None or decision not in ("aprobada", "rechazada", "pendiente")
            or (decision == "aprobada" and eq.get("ambigua"))):
        return render_template("index.html",
                               error="No se guardó la decisión: la equivalencia no existe o es ambigua."), 400
    import integracion as _integ
    usuario = request.headers.get("X-SGO7-User") or session.get("usuario") or ""
    _integ.equiv_decidir(sgs_data.get("_equiv_clave") or "", eq, decision, usuario=usuario)
    msg = {"aprobada": f"Aprobado: {eq['sip']} es el mismo tramo que {eq['sgs']} ({eq.get('pro', '')}).",
           "rechazada": f"Rechazado: {eq['sip']} no es {eq['sgs']}.",
           "pendiente": f"Decisión deshecha para {eq['sip']}."}[decision]
    return render_template_string(_REENVIO_EQ, msg=msg, id_planilla=datos.get("id_planilla", ""),
                                  proyecto_id=datos.get("proyecto_id", ""))


@app.route("/generar-final/<sid>", methods=["POST"])
@login_required
def generar_final(sid):
    """Fase final: con fecha + líderes confirmados, genera los 2 Excel
    (balance integrado y RG11 supervisor) y muestra la página de descarga."""
    datos = _get_preview(sid)
    if not datos:
        return render_template(
            "index.html",
            error="La sesión expiró. Genera el balance de nuevo.",
        ), 404

    fecha   = request.form.get("fecha_ejecucion", "").strip()
    cedulas = request.form.getlist("lider_cedula")
    nombres = request.form.getlist("lider_nombre")
    pcts    = request.form.getlist("lider_pct")
    lideres = [
        {"cedula": c.strip(), "nombre": n.strip(), "pct": p}
        for c, n, p in zip(cedulas, nombres, pcts) if (c.strip() or n.strip())
    ]

    cliente = _get_cliente()
    sip, sgs = datos["sip_datos"], datos["sgs_data"]
    ts = ahora_co().strftime("%Y%m%d_%H%M")
    padre   = datos["sgs_padre"]
    odm_sgs = padre.get("order_number") or padre.get("odm") or datos["parent_id"]

    # Excel 1: balance integrado + hoja LIQUIDACION (fecha y líderes)
    bal_bytes = cliente._generar_excel(
        sip, sgs, datos["id_label"], advertencias=datos.get("advertencias", []),
        fecha_ejecucion=fecha, lideres=lideres,
    )
    bal_name = f"balance_integrado_{datos['id_label']}_{odm_sgs}_{ts}.xlsx"
    tok_bal = uuid.uuid4().hex
    _save_archivo(tok_bal, bal_bytes, bal_name, datos.get("advertencias", []))

    # Excel 2: RG11 supervisor
    import rg11
    rg11_bytes, rg11_adv = rg11.generar_rg11(sip, sgs, meta=datos["metadata"],
                                             fecha=fecha, lideres=lideres)
    nombre_proy = (padre.get("name") or odm_sgs).replace("/", "-")
    rg11_name = f"RG11 {nombre_proy}.xlsx"
    tok_rg11 = uuid.uuid4().hex
    _save_archivo(tok_rg11, rg11_bytes, rg11_name, rg11_adv)

    # SGO7 facturación: el Excel queda guardado en el caso (copia + huella) y la página final se sirve
    # por GET (/final/<sid>) para poder volver a ella tras «Mandar a liquidar».
    datos["_final"] = {"fecha": fecha, "lideres": lideres, "tok_bal": tok_bal, "bal_name": bal_name,
                       "tok_rg11": tok_rg11, "rg11_name": rg11_name, "rg11_adv": rg11_adv}
    try:
        import sgo7_casos
        if datos.get("caso"):
            datos["caso"] = sgo7_casos.guardar_excel_caso(datos["caso"], bal_bytes, bal_name, rg11_bytes, rg11_name,
                                                          fecha=fecha, lideres=lideres, sid=sid) or datos["caso"]
    except Exception as _e:
        print("sgo7_casos: excel del caso no guardado:", _e)
    _save_preview(sid, datos)
    return redirect(url_for("final", sid=sid))


@app.route("/final/<sid>")
@login_required
def final(sid):
    """Página de descarga (fase final) + bloque «Mandar a liquidar» del caso."""
    datos = _get_preview(sid)
    if not datos or not datos.get("_final"):
        return render_template("index.html", error="La sesión expiró. Genera el balance de nuevo."), 404
    f = datos["_final"]
    caso_error = datos.pop("caso_error", None)
    if caso_error:
        _save_preview(sid, datos)
    return render_template(
        "resultado_final.html",
        sid=sid, caso=datos.get("caso"), caso_error=caso_error, orden=request.args.get("orden", ""),
        metadata=datos["metadata"], sgs_padre=datos["sgs_padre"],
        fecha=f["fecha"], lideres=f["lideres"],
        tok_bal=f["tok_bal"], bal_name=f["bal_name"],
        tok_rg11=f["tok_rg11"], rg11_name=f["rg11_name"], rg11_adv=f["rg11_adv"],
        advertencias=datos.get("advertencias", []),
    )


@app.route("/casos/<int:cid>/estado", methods=["POST"])
@login_required
def caso_estado(cid):
    """Facturador: «Dejar en espera de almacén» / «Quitar de espera» desde la previsualización."""
    import sgo7_casos
    sid = request.form.get("sid", "")[:40]
    caso, err = sgo7_casos.cambiar_estado(cid, estado=request.form.get("estado", "") or None,
                                          nota=(request.form.get("nota", "").strip() or None))
    datos = _get_preview(sid) if sid else None
    if datos is not None:
        if caso:
            datos["caso"] = caso
        else:
            datos["caso_error"] = err
        _save_preview(sid, datos)
        return redirect(url_for("preview", sid=sid))
    return redirect(url_for("index"))


@app.route("/casos/<int:cid>/mandar-a-liquidar", methods=["POST"])
@login_required
def caso_mandar_liquidar(cid):
    """Facturador: crea la orden para el Liquidador Automático (aparece en el panel del Liquidador)."""
    import sgo7_casos
    sid = request.form.get("sid", "")[:40]
    caso, err = sgo7_casos.mandar_a_liquidar(cid, nota=request.form.get("nota", "").strip())
    datos = _get_preview(sid) if sid else None
    if datos is not None:
        if caso:
            datos["caso"] = caso
        else:
            datos["caso_error"] = err
        _save_preview(sid, datos)
        return redirect(url_for("final", sid=sid, orden="ok" if caso else "error"))
    return redirect(url_for("index"))


@app.route("/descargar/<token>")
@login_required
def descargar(token):
    archivo = _get_archivo(token)
    if not archivo:
        return render_template("index.html",
                               error="El archivo expiró o no está disponible. Genera el balance nuevamente.",
                               tab="sgs"), 404
    return send_file(
        io.BytesIO(archivo["bytes"]),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=archivo["filename"],
    )


@app.route("/generar-sgs", methods=["POST"])
@login_required
def generar_sgs():
    from generar_balance import generar_excel

    # Paso 1 → viene de index.html (solo proyecto_id_sgs)
    # Paso 2 → viene de seleccionar_nodos.html (parent_id + nodo_codes[])
    proyecto_id  = request.form.get("proyecto_id_sgs", "").strip()
    parent_id_hd = request.form.get("parent_id", "").strip()   # hidden en paso 2
    nodos_sel    = request.form.getlist("nodo_codes")           # lista de PRO-XXXXXX seleccionados
    es_paso2     = bool(parent_id_hd)                          # True cuando viene de seleccionar_nodos

    pid_raw = parent_id_hd or proyecto_id
    if not pid_raw:
        return render_template("index.html", error_sgs="Ingresa un ID o código de proyecto SGS.", tab="sgs")

    try:
        cliente = _get_sgs_cliente()

        if pid_raw.isdigit():
            parent_id = int(pid_raw)
        else:
            proyecto = cliente.buscar_por_codigo(pid_raw)
            if not proyecto:
                return render_template(
                    "index.html",
                    error_sgs=f"Proyecto '{pid_raw}' no encontrado en SGS.",
                    proyecto_id_sgs=pid_raw,
                    tab="sgs",
                )
            parent_id = proyecto["id"]

        data  = cliente.balance_completo(parent_id)
        padre = data["padre"]
        nodos = data["nodos"]

        nodos_resumen = [
            {
                "pro":        n.get("code", ""),
                "nodo":       n.get("nodo", "—"),
                "status":     (n.get("status", ["", ""])[1]
                               if isinstance(n.get("status"), list)
                               else str(n.get("status", ""))),
                "bandejas":   len(n.get("bandejas", [])),
                "materiales": len(n.get("liq_materiales", [])),
            }
            for n in nodos
        ]

        # ── Paso 1: mostrar pantalla de selección de nodos ────────────────────
        if not es_paso2:
            return render_template(
                "seleccionar_nodos.html",
                padre=padre,
                parent_id=str(parent_id),
                nodos=nodos_resumen,
            )

        # ── Paso 2: filtrar nodos seleccionados y generar Excel ───────────────
        if not nodos_sel:
            return render_template(
                "seleccionar_nodos.html",
                padre=padre,
                parent_id=str(parent_id),
                nodos=nodos_resumen,
                error="Selecciona al menos un nodo.",
            )

        nodos_filtrados = [n for n in nodos if n.get("code", "") in nodos_sel]
        data_filtrada   = dict(data, nodos=nodos_filtrados)

        odm      = padre.get("order_number") or padre.get("odm") or str(parent_id)
        ts       = ahora_co().strftime("%Y%m%d_%H%M")
        sufijo   = f"_{len(nodos_filtrados)}nodos" if len(nodos_filtrados) < len(nodos) else ""
        filename = f"balance_sgs_{odm}{sufijo}_{ts}.xlsx"

        excel_bytes = generar_excel(data_filtrada)

        token = str(uuid.uuid4())
        _save_archivo(token, excel_bytes, filename)

        nodos_resultado = [nr for nr in nodos_resumen if nr["pro"] in nodos_sel]

        return render_template(
            "resultado_sgs.html",
            token=token,
            filename=filename,
            padre=padre,
            odm=odm,
            nodos=nodos_resultado,
        )

    except Exception as e:
        msg = str(e)
        try:
            return render_template("index.html", error_sgs=msg, proyecto_id_sgs=pid_raw, tab="sgs")
        except Exception:
            return f"<pre style='color:red;padding:20px'>ERROR SGS:\n{msg}</pre>", 500


@app.route("/sgs-libre", methods=["POST"])
@login_required
def sgs_libre():
    """
    Opción B — búsqueda por LCL / DF / ODM.
    Encuentra el proyecto padre en SGS y abre el wizard de 4 pasos.
    """
    from manual_balance import nueva_sesion

    codigo_raw = request.form.get("lcl_codigo", "").strip()
    pdl_raw    = request.form.get("pdl", "").strip()
    if not codigo_raw:
        return render_template("index.html", error_sgs_libre="Ingresa un código LCL, DF u ODM.", tab="sgs", pdl=pdl_raw)

    codigo = codigo_raw.upper()
    if codigo.startswith("LCL"):
        codigo = codigo[3:].lstrip("_- ")

    try:
        cliente  = _get_sgs_cliente()
        proyecto = cliente.buscar_por_codigo(codigo)
        if not proyecto:
            return render_template(
                "index.html",
                error_sgs_libre=f"No se encontró ningún proyecto con código '{codigo_raw}' en SGS.",
                lcl_codigo=codigo_raw,
                tab="sgs",
            )

        parent_id = proyecto["id"]
        padre     = cliente.leer_proyecto(parent_id)
        sid       = nueva_sesion(padre or proyecto, parent_id)
        if pdl_raw:
            from manual_balance import get_sesion, save_sesion
            s = get_sesion(sid)
            s["pdl"] = pdl_raw
            save_sesion(sid)
        # Lanzar fetch SGS en background para que esté listo en paso 3
        from manual_balance import iniciar_fetch_sgs
        iniciar_fetch_sgs(sid, _get_sgs_cliente)
        return redirect(url_for("manual_paso2", sid=sid))

    except Exception as e:
        return render_template(
            "index.html",
            error_sgs_libre=f"Error consultando SGS: {e}",
            lcl_codigo=codigo_raw,
            tab="sgs",
        )


# ── Wizard Manual — Balance SGS Puro ──────────────────────────────────────────

@app.route("/manual/mo/<sid>", methods=["GET", "POST"])
@login_required
def manual_mo(sid):
    import json as _json
    from manual_balance import get_sesion, parsear_mo, get_precios, save_sesion
    sesion = get_sesion(sid)
    if not sesion:
        return render_template("index.html", tab="sgs",
                               error_sgs_libre="Sesión expirada — ingresa el código del proyecto nuevamente.")

    precios_json = _json.dumps(get_precios(), ensure_ascii=False)

    if request.method == "POST":
        from manual_balance import enriquecer_mo
        texto = request.form.get("texto_mo", "").strip()
        rows, errors = parsear_mo(texto)
        enriquecer_mo(rows)
        sesion["rows_mo"] = rows
        sesion["errs_mo"] = errors
        save_sesion(sid)
        if not errors and rows:
            return redirect(url_for("manual_nodos", sid=sid))
        if not rows:
            errors = errors or ["Pega al menos una fila con Nodo y QP."]
        return render_template("manual_mo.html", sesion=sesion, sid=sid,
                               texto=texto, errors=errors,
                               precios_json=precios_json)

    return render_template("manual_mo.html", sesion=sesion, sid=sid,
                           texto="", errors=[], precios_json=precios_json)


@app.route("/manual/paso2/<sid>", methods=["GET"])
@login_required
def manual_paso2(sid):
    import json as _json
    from manual_balance import get_sesion, get_precios, resolver_nodos
    sesion = get_sesion(sid)
    if not sesion:
        return render_template("index.html", tab="sgs",
                               error_sgs_libre="Sesión expirada — ingresa el código del proyecto nuevamente.")
    precios_json = _json.dumps(get_precios(), ensure_ascii=False)
    sgs_ready = sesion.get("sgs_ready")
    nodos_resueltos = None
    if sgs_ready is True and sesion.get("sgs_data"):
        mo_nodos = [r["nodo"] for r in sesion.get("rows_mo", [])]
        nodos_resueltos = resolver_nodos(sesion["sgs_data"]["nodos"], mo_nodos)
    # Supervisor: use saved override only — responsible_id is always the API user, not the real supervisor
    sup_val = sesion.get("supervisor") or ""
    return render_template("manual_paso2.html",
                           sesion=sesion, sid=sid,
                           precios_json=precios_json,
                           sgs_ready=sgs_ready,
                           nodos_resueltos=nodos_resueltos,
                           texto_mo=sesion.get("texto_mo", ""),
                           texto_nodos=sesion.get("texto_nodos", ""),
                           supervisor=sup_val,
                           pros_importados=sesion.get("pros_importados", []))


@app.route("/api/guardar-mo/<sid>", methods=["POST"])
@login_required
def api_guardar_mo(sid):
    from manual_balance import get_sesion, parsear_mo, enriquecer_mo, save_sesion, resolver_nodos
    sesion = get_sesion(sid)
    if not sesion:
        return {"ok": False, "error": "Sesión expirada"}, 404
    data = request.get_json(silent=True) or {}
    texto = data.get("texto", "").strip()
    rows, errors = parsear_mo(texto)
    if not errors and rows:
        enriquecer_mo(rows)
        sesion["rows_mo"]  = rows
        sesion["errs_mo"]  = []
        sesion["texto_mo"] = texto
        save_sesion(sid)
        # Verificar nodos contra SGS si ya está listo
        resueltos = None
        if sesion.get("sgs_ready") is True and sesion.get("sgs_data"):
            mo_nodos = list(dict.fromkeys(r["nodo"] for r in rows))
            resueltos = resolver_nodos(sesion["sgs_data"]["nodos"], mo_nodos)
        return {"ok": True, "count": len(rows), "resueltos": resueltos,
                "sgs_ready": sesion.get("sgs_ready")}
    return {"ok": False, "errors": errors or ["Sin filas válidas"], "count": len(rows)}


@app.route("/api/sgs-status/<sid>")
@login_required
def api_sgs_status(sid):
    from manual_balance import get_sesion, resolver_nodos
    sesion = get_sesion(sid)
    if not sesion:
        return {"ready": None, "error": "Sesión expirada"}, 404
    ready = sesion.get("sgs_ready")
    error = sesion.get("sgs_error")
    if ready is True and not error:
        sgs_data = sesion.get("sgs_data") or {}
        nodos = sgs_data.get("nodos", [])
        mo_nodos = [r["nodo"] for r in sesion.get("rows_mo", [])]
        resueltos = resolver_nodos(nodos, mo_nodos)
        return {"ready": True, "error": None, "nodos_count": len(nodos), "resueltos": resueltos}
    return {"ready": ready, "error": error, "nodos_count": 0}


@app.route("/api/buscar-nodos/<sid>", methods=["POST"])
@login_required
def api_buscar_nodos(sid):
    from manual_balance import get_sesion, save_sesion, resolver_nodos, detectar_nodos_solos
    sesion = get_sesion(sid)
    if not sesion:
        return {"ok": False, "error": "Sesión expirada"}, 404
    if sesion.get("sgs_ready") is not True:
        return {"ok": False, "sgs_loading": True, "error": "SGS aún cargando, espera un momento"}
    sgs_data = sesion.get("sgs_data")
    if not sgs_data:
        return {"ok": False, "error": "Datos SGS no disponibles"}
    data = request.get_json(silent=True) or {}
    texto = data.get("texto", "").strip()
    sesion["texto_nodos"] = texto
    save_sesion(sid)
    nodos_solos = detectar_nodos_solos(texto)
    if not nodos_solos:
        return {"ok": False, "error": "Formato no reconocido — pon un nodo por línea"}
    resueltos = resolver_nodos(sgs_data["nodos"], nodos_solos)
    return {"ok": True, "resueltos": resueltos}


@app.route("/manual/paso2/<sid>/confirmar", methods=["POST"])
@login_required
def manual_paso2_confirmar(sid):
    from manual_balance import get_sesion, save_sesion, detectar_nodos_solos, resolver_nodos
    sesion = get_sesion(sid)
    if not sesion:
        return render_template("index.html", tab="sgs",
                               error_sgs_libre="Sesión expirada.")

    # Red de seguridad: si rows_mo se perdió por race condition, recuperar del formulario
    texto_mo_form = request.form.get("texto_mo_form", "").strip()
    if texto_mo_form and not sesion.get("rows_mo"):
        from manual_balance import parsear_mo, enriquecer_mo
        rows, errors = parsear_mo(texto_mo_form)
        if not errors and rows:
            enriquecer_mo(rows)
            sesion["rows_mo"]  = rows
            sesion["texto_mo"] = texto_mo_form
            sesion["errs_mo"]  = []

    omitir = request.form.get("omitir_mat") == "1"
    pros_sel   = request.form.getlist("pro_sel")
    pros_extra = [v.strip() for v in request.form.getlist("pro_extra") if v.strip()]
    todos = list(dict.fromkeys(pros_sel + pros_extra))

    # texto_nodos viene del form (enviado por JS) o de la sesión previa (cuando clickaron "Buscar")
    texto_nodos_form = (request.form.get("texto_nodos") or "").strip()
    if texto_nodos_form:
        sesion["texto_nodos"] = texto_nodos_form

    sgs_data = sesion.get("sgs_data")
    nodos_cache = (sgs_data or {}).get("nodos", [])

    sesion["rows_mat"] = []
    sesion["modo_mat"] = "nodo_solo"

    # Asegurar que sgs_data esté disponible (puede no estar en disco si el worker cambió)
    sgs_data = sesion.get("sgs_data")
    if not sgs_data and sesion.get("sgs_ready") is True:
        try:
            cli = _get_sgs_cliente()
            sgs_data = cli.balance_completo(sesion["parent_id"])
            sesion["sgs_data"] = sgs_data
            save_sesion(sid)
        except Exception:
            pass

    if omitir:
        # "Solo MO" — sin materiales
        sesion["nodos_mat"] = []
        sesion["nodos_sel"] = []

    elif not todos:
        # Fallback 1: PROs importados desde un Excel previo
        pros_importados = sesion.get("pros_importados", [])
        if pros_importados and nodos_cache:
            pros_en_cache = {n.get("code") for n in nodos_cache}
            pros_faltantes = [p for p in pros_importados if p not in pros_en_cache]
            nodos_extra = []
            if pros_faltantes:
                try:
                    cli = _get_sgs_cliente()
                    for pro in pros_faltantes:
                        p = cli.buscar_por_codigo(pro)
                        if p:
                            d = cli.balance_completo(p["id"])
                            nodos_extra.extend(d.get("nodos", []))
                except Exception:
                    pass
            nodos_filtrados = [n for n in nodos_cache if n.get("code") in set(pros_importados)] + nodos_extra
            sesion["nodos_mat"] = nodos_filtrados
            sesion["nodos_sel"] = [n.get("code") for n in nodos_filtrados]
        else:
            # Fallback 2: resolver desde el texto del textarea de materiales
            texto_nodos = (texto_nodos_form or sesion.get("texto_nodos", "")).strip()
            if texto_nodos and nodos_cache:
                nodos_solos = detectar_nodos_solos(texto_nodos)
                if nodos_solos:
                    resueltos = resolver_nodos(nodos_cache, nodos_solos)
                    pros_ok = [r["pro"] for r in resueltos.get("encontrados", [])]
                    nodos_filtrados = [n for n in nodos_cache if n.get("code") in set(pros_ok)]
                    sesion["nodos_mat"] = nodos_filtrados
                    sesion["nodos_sel"] = [n.get("code") for n in nodos_filtrados]
                else:
                    sesion["nodos_mat"] = []
                    sesion["nodos_sel"] = []
            else:
                sesion["nodos_mat"] = []
                sesion["nodos_sel"] = []

    else:
        # PROs seleccionados explícitamente para materiales
        pros_en_cache = {n.get("code") for n in nodos_cache}
        pros_faltantes = [p for p in todos if p not in pros_en_cache]
        nodos_extra = []
        if pros_faltantes:
            try:
                cliente = _get_sgs_cliente()
                for pro in pros_faltantes:
                    p = cliente.buscar_por_codigo(pro)
                    if p:
                        d = cliente.balance_completo(p["id"])
                        nodos_extra.extend(d.get("nodos", []))
            except Exception:
                pass  # best effort — usar lo que haya en cache
        nodos_filtrados = [n for n in nodos_cache if n.get("code") in set(todos)] + nodos_extra
        sesion["nodos_mat"] = nodos_filtrados
        sesion["nodos_sel"] = [n.get("code") for n in nodos_filtrados]

    sesion["nodos_pendientes_sgs"] = request.form.getlist("pendiente_sgs")
    cedulas = request.form.getlist("lider_cedula")
    nombres = request.form.getlist("lider_nombre")
    pcts    = request.form.getlist("lider_pct")
    sesion["lideres"] = [
        {"cedula": c.strip(), "nombre": n.strip(), "pct": p}
        for c, n, p in zip(cedulas, nombres, pcts)
        if c.strip() or n.strip()
    ]
    sup_form = request.form.get("supervisor", "").strip()
    if sup_form and sup_form != "—":
        sesion["supervisor"] = sup_form
    save_sesion(sid)
    return redirect(url_for("manual_preview", sid=sid))


@app.route("/manual/importar-excel", methods=["POST"])
@login_required
def manual_importar_excel():
    """
    Opción C — importar un Excel ya generado por el sistema.
    Lee MO, líderes, supervisor y nodos de materiales del Excel,
    crea una sesión pre-poblada y lanza fetch SGS fresco.
    """
    from manual_balance import nueva_sesion, get_sesion, save_sesion, iniciar_fetch_sgs, leer_excel_sesion

    archivo = request.files.get("excel_file")
    if not archivo or not archivo.filename:
        return render_template("index.html", tab="sgs",
                               error_importar="Selecciona un archivo Excel (.xlsx).")

    try:
        datos = leer_excel_sesion(archivo)
    except ValueError as e:
        return render_template("index.html", tab="sgs", error_importar=str(e))
    except Exception as e:
        return render_template("index.html", tab="sgs",
                               error_importar=f"No se pudo leer el Excel: {e}")

    codigo_padre = datos["codigo_padre"]
    if not codigo_padre:
        return render_template("index.html", tab="sgs",
                               error_importar="El Excel no contiene el Código padre del proyecto.")

    try:
        cliente  = _get_sgs_cliente()
        proyecto = cliente.buscar_por_codigo(codigo_padre)
        if not proyecto:
            return render_template("index.html", tab="sgs",
                                   error_importar=f"No se encontró '{codigo_padre}' en SGS.")
        parent_id = proyecto["id"]
        padre     = cliente.leer_proyecto(parent_id) or proyecto
    except Exception as e:
        return render_template("index.html", tab="sgs",
                               error_importar=f"Error consultando SGS: {e}")

    sid = nueva_sesion(padre, parent_id)
    sesion = get_sesion(sid)

    # Pre-poblar desde el Excel
    sesion["lcl_odm"]        = datos["lcl_odm"] or sesion.get("lcl_odm", "")
    sesion["supervisor"]     = datos["supervisor"]
    sesion["pdl"]            = datos.get("pdl", "")
    sesion["texto_mo"]       = datos["texto_mo"]
    sesion["rows_mo"]        = datos["rows_mo"]
    sesion["lideres"]        = datos["lideres"]
    sesion["texto_nodos"]    = datos.get("texto_nodos", "")
    sesion["pros_importados"] = datos.get("pros_importados", [])
    sesion["nodos_mat"]      = []
    sesion["nodos_sel"]      = []
    sesion["rows_mat"]       = []
    sesion["modo_mat"]       = "nodo_solo"
    save_sesion(sid)

    # Fetch SGS en background (datos frescos — devoluciones actualizadas)
    iniciar_fetch_sgs(sid, _get_sgs_cliente)
    return redirect(url_for("manual_paso2", sid=sid))


@app.route("/api/sesiones-guardadas")
@login_required
def api_sesiones_guardadas():
    """AJAX: lista los balances manuales guardados (panel Retomar del index)."""
    from manual_balance import listar_sesiones
    try:
        return {"ok": True, "sesiones": listar_sesiones()}
    except Exception as e:
        return {"ok": False, "error": str(e)}, 500


@app.route("/manual/retomar/<sid>")
@login_required
def manual_retomar(sid):
    """
    Retomar un balance guardado.
    Reconsulta SGS fresco (devoluciones al día) y vuelve al paso 2 con la
    MO, líderes, supervisor y nodos tal como quedaron.
    """
    from manual_balance import get_sesion, save_sesion, iniciar_fetch_sgs
    sesion = get_sesion(sid)
    if not sesion:
        return render_template("index.html", tab="sgs",
                               error_sgs_libre="Ese balance guardado ya no existe — inicia de nuevo con el código del proyecto.")
    # Descartar el SGS viejo: el almacén pudo haber hecho devoluciones desde entonces
    sesion["sgs_data"] = None
    save_sesion(sid)
    iniciar_fetch_sgs(sid, _get_sgs_cliente)
    return redirect(url_for("manual_paso2", sid=sid))


@app.route("/manual/sesion/<sid>/eliminar", methods=["POST"])
@login_required
def manual_sesion_eliminar(sid):
    from manual_balance import eliminar_sesion
    return {"ok": eliminar_sesion(sid)}


@app.route("/manual/nodos/<sid>", methods=["GET", "POST"])
@login_required
def manual_nodos(sid):
    from manual_balance import get_sesion, save_sesion, resolver_nodos

    sesion = get_sesion(sid)
    if not sesion:
        return render_template("index.html", tab="sgs",
                               error_sgs_libre="Sesión expirada — ingresa el código del proyecto nuevamente.")

    # POST: el usuario confirmó la selección de nodos (checkboxes + PROs manuales)
    if request.method == "POST":
        pros_sel   = request.form.getlist("pro_sel")        # PROs seleccionados del fetch
        pros_extra = [v.strip() for v in request.form.getlist("pro_extra") if v.strip()]
        todos = list(dict.fromkeys(pros_sel + pros_extra))  # deduplicar, preservar orden

        if not todos:
            # re-mostrar con error
            pass
        else:
            sgs_data = sesion.get("sgs_data") or {}
            nodos_cache = sgs_data.get("nodos", [])

            # Los PROs extra que no estén en el cache → consultarlos ahora
            pros_en_cache = {n.get("code") for n in nodos_cache}
            pros_faltantes = [p for p in todos if p not in pros_en_cache]

            nodos_extra = []
            errores_pro = []
            if pros_faltantes:
                try:
                    cliente = _get_sgs_cliente()
                    for pro in pros_faltantes:
                        proy = cliente.buscar_por_codigo(pro)
                        if proy:
                            nodo_data = cliente.balance_completo(proy["id"])
                            nodos_extra.extend(nodo_data.get("nodos", []))
                        else:
                            errores_pro.append(f"PRO '{pro}' no encontrado en SGS")
                except Exception as e:
                    errores_pro.append(f"Error consultando SGS: {e}")

            if errores_pro:
                # Re-mostrar pantalla con errores
                sgs_ready = sesion.get("sgs_ready")
                nodos_resueltos = resolver_nodos(nodos_cache, [r["nodo"] for r in sesion.get("rows_mo", [])])
                return render_template("manual_nodos.html", sesion=sesion, sid=sid,
                                       sgs_ready=sgs_ready, nodos_resueltos=nodos_resueltos,
                                       errores_pro=errores_pro)

            # Filtrar nodos del cache por PROs seleccionados + agregar extras
            nodos_filtrados = [n for n in nodos_cache if n.get("code") in set(todos)]
            nodos_filtrados += nodos_extra

            padre_data = (sesion.get("sgs_data") or {}).get("padre") or sesion["padre"]
            sesion["sgs_data"]  = {"padre": padre_data, "nodos": nodos_filtrados}
            sesion["nodos_sel"] = [n.get("code") for n in nodos_filtrados]
            sesion["modo_mat"]  = "nodo_solo"
            sesion["rows_mat"]  = []
            save_sesion(sid)
            return redirect(url_for("manual_preview", sid=sid))

    # GET: mostrar pantalla de selección
    sgs_ready = sesion.get("sgs_ready")  # None | False | True
    sgs_error = sesion.get("sgs_error")
    nodos_resueltos = None

    if sgs_ready is True:
        # sgs_data puede ser None si el fetch terminó en otro worker y el disco no lo guardó
        if sesion.get("sgs_data") is None and not sgs_error:
            try:
                cliente  = _get_sgs_cliente()
                data     = cliente.balance_completo(sesion["parent_id"])
                sesion["sgs_data"] = data
                save_sesion(sid)
            except Exception as e:
                sgs_error = str(e)
        if sesion.get("sgs_data") and not sgs_error:
            sgs_data   = sesion.get("sgs_data") or {}
            nodos_sgs  = sgs_data.get("nodos", [])
            mo_nodos   = [r["nodo"] for r in sesion.get("rows_mo", [])]
            nodos_resueltos = resolver_nodos(nodos_sgs, mo_nodos)

    return render_template("manual_nodos.html", sesion=sesion, sid=sid,
                           sgs_ready=sgs_ready, sgs_error=sgs_error,
                           nodos_resueltos=nodos_resueltos, errores_pro=[])


@app.route("/api/lideres-sugeridos/<sid>")
@login_required
def api_lideres_sugeridos(sid):
    """Detecta líderes desde el partner_id de las bandejas de cada nodo
    (el Equipo del proyecto en SGS) y enriquece con cédula via res.partner.vat."""
    import re as _re
    from manual_balance import get_sesion
    # El sid puede venir de la sesión manual o del store de previsualización.
    sesion = get_sesion(sid)
    sgs_src = sesion.get("sgs_data") if (sesion and sesion.get("sgs_data")) else None
    if sgs_src is None:
        prev = _get_preview(sid)
        sgs_src = prev.get("sgs_data") if prev else None
    if not sgs_src:
        return {"ok": True, "lideres": []}

    cedula_re = _re.compile(r"^C\d+$", _re.IGNORECASE)

    # ── Fuente principal: partner_id de las bandejas ─────────────────────────
    # Cada bandeja tiene el partner que recibió/usó los materiales = el líder real
    ban_partners: dict = {}  # partner_id → nombre
    for nodo in sgs_src.get("nodos", []):
        for ban in nodo.get("bandejas", []):
            partner = ban.get("partner_id")
            if isinstance(partner, list) and len(partner) >= 2 and partner[0]:
                ban_partners[partner[0]] = partner[1]

    # ── Fallback: responsible_id del nodo si no hay bandejas ──────────────────
    responsibles: dict = {}  # user_id → nombre
    if not ban_partners:
        for nodo in sgs_src.get("nodos", []):
            resp = nodo.get("responsible_id")
            if isinstance(resp, list) and len(resp) >= 2 and resp[0]:
                responsibles[resp[0]] = resp[1]

    if not ban_partners and not responsibles:
        return {"ok": True, "lideres": []}

    try:
        cliente = _get_sgs_cliente()

        # Inicializar entradas desde bandejas (son res.partner directamente)
        pid_to_ldr: dict = {
            pid: {"cedula": "", "nombre": nombre, "pct": 0, "cedula_valida": False}
            for pid, nombre in ban_partners.items()
        }

        # Para el fallback: resolver res.users → partner_id
        if responsibles:
            try:
                users = cliente.read("res.users", list(responsibles.keys()), ["id", "partner_id"]) or []
                for u in users:
                    pid_raw = u.get("partner_id")
                    if isinstance(pid_raw, list) and pid_raw and pid_raw[0] not in pid_to_ldr:
                        pid_to_ldr[pid_raw[0]] = {
                            "cedula": "", "nombre": responsibles.get(u["id"], ""),
                            "pct": 0, "cedula_valida": False,
                        }
            except Exception:
                pass

        # Enriquecer todos con cédula desde res.partner.vat
        if pid_to_ldr:
            partners = cliente.read("res.partner", list(pid_to_ldr.keys()), ["id", "name", "vat"]) or []
            for p in partners:
                pid = p["id"]
                if pid not in pid_to_ldr:
                    continue
                cedula = _re.sub(r"^[Cc]{2}(\d)", r"C\1", (p.get("vat") or "")).upper()
                pid_to_ldr[pid]["cedula"] = cedula
                pid_to_ldr[pid]["cedula_valida"] = bool(cedula_re.match(cedula))
                if p.get("name"):
                    pid_to_ldr[pid]["nombre"] = p["name"]

        # Deduplicar por nombre y devolver (con y sin cédula)
        seen: set = set()
        lideres = []
        for ldr in pid_to_ldr.values():
            if not ldr["nombre"] or ldr["nombre"] in seen:
                continue
            seen.add(ldr["nombre"])
            lideres.append(ldr)

        return {"ok": True, "lideres": lideres}
    except Exception as e:
        return {"ok": False, "error": str(e)}, 500


@app.route("/api/buscar-lider")
@login_required
def api_buscar_lider():
    """AJAX: busca líderes/socios por nombre o cédula en SGS (res.partner)."""
    import re as _re
    import unicodedata as _ud
    q = request.args.get("q", "").strip()
    if len(q) < 2:
        return {"ok": True, "results": []}
    try:
        cliente = _get_sgs_cliente()
        # Normalizar: CC11111 → C11111
        q_norm = _re.sub(r"^[Cc]{2}(\d)", r"C\1", q)

        # En SGS los nombres van "APELLIDOS NOMBRES"; un ilike de la frase
        # completa no encuentra "JAVIER CASTILLO". Se busca cada palabra por
        # separado (AND implícito del dominio Odoo), sin tildes.
        sin_tildes = _ud.normalize("NFKD", q).encode("ascii", "ignore").decode("ascii")
        palabras = [p for p in sin_tildes.split() if p]
        dominio = [["name", "ilike", p] for p in palabras] or [["name", "ilike", q]]

        by_name = cliente.search_read(
            "res.partner", dominio,
            ["id", "name", "vat"], limit=10)
        by_vat = cliente.search_read(
            "res.partner", [["vat", "ilike", q_norm]],
            ["id", "name", "vat"], limit=5)

        seen = set()
        candidatos = []
        for r in (by_name + by_vat):
            if r["id"] in seen:
                continue
            seen.add(r["id"])
            cedula = r.get("vat") or ""
            cedula = _re.sub(r"^[Cc]{2}(\d)", r"C\1", cedula)
            candidatos.append({
                "id":     r["id"],
                "nombre": r.get("name", ""),
                "cedula": cedula,
            })

        # SGS tiene socios duplicados (mismo nombre con y sin cédula C válida):
        # dejar uno por nombre, prefiriendo el registro con cédula usable.
        def _rango(c):
            if _re.fullmatch(r"C\d+", c["cedula"], _re.I):
                return 0
            return 1 if c["cedula"] else 2
        mejores = {}
        for c in candidatos:
            clave = " ".join(c["nombre"].upper().split())
            if clave not in mejores or _rango(c) < _rango(mejores[clave]):
                mejores[clave] = c
        results = sorted(mejores.values(), key=_rango)

        return {"ok": True, "results": results[:10]}
    except Exception as e:
        return {"ok": False, "error": str(e)}, 500


@app.route("/api/validar-pro")
@login_required
def api_validar_pro():
    """AJAX: verifica si un código PRO existe en SGS. Retorna JSON."""
    pro = request.args.get("pro", "").strip().upper()
    if not pro:
        return {"ok": False, "error": "Código vacío"}, 400
    try:
        cliente = _get_sgs_cliente()
        proy = cliente.buscar_por_codigo(pro)
        if proy:
            return {"ok": True, "name": proy.get("name", ""), "code": proy.get("code", pro)}
        return {"ok": False, "error": f"'{pro}' no encontrado en SGS"}
    except Exception as e:
        return {"ok": False, "error": str(e)}, 500


@app.route("/manual/mat/<sid>", methods=["GET", "POST"])
@login_required
def manual_mat(sid):
    from manual_balance import get_sesion, parsear_mat, detectar_nodos_solos, filtrar_nodos_sgs, save_sesion
    sesion = get_sesion(sid)
    if not sesion:
        return render_template("index.html", tab="sgs",
                               error_sgs_libre="Sesión expirada — ingresa el código del proyecto nuevamente.")

    if request.method == "POST":
        texto  = request.form.get("texto_mat", "").strip()
        omitir = request.form.get("omitir_mat") == "1"

        if omitir or not texto:
            sesion["rows_mat"] = []
            sesion["errs_mat"] = []
            sesion["modo_mat"] = "nodo_solo"
            # Nodos de la MO como fallback
            sesion["nodos_sel"] = list({r["nodo"] for r in sesion["rows_mo"]})
        else:
            nodos_solos = detectar_nodos_solos(texto)
            if nodos_solos:
                sesion["modo_mat"]  = "nodo_solo"
                sesion["nodos_sel"] = nodos_solos
                sesion["rows_mat"]  = []
                sesion["errs_mat"]  = []
            else:
                rows, errors = parsear_mat(texto)
                sesion["rows_mat"]  = rows
                sesion["errs_mat"]  = errors
                sesion["modo_mat"]  = "manual_sap"
                sesion["nodos_sel"] = list({r["nodo"] for r in rows})
                if errors:
                    return render_template("manual_mat.html", sesion=sesion, sid=sid,
                                           texto=texto, errors=errors)

        # Obtener datos SGS para los nodos seleccionados
        try:
            cliente   = _get_sgs_cliente()
            sgs_full  = cliente.balance_completo(sesion["parent_id"])
            if sesion["nodos_sel"]:
                nodos_f = filtrar_nodos_sgs(sgs_full["nodos"], sesion["nodos_sel"])
            else:
                nodos_f = sgs_full["nodos"]
            sesion["sgs_data"] = {"padre": sgs_full["padre"], "nodos": nodos_f}
        except Exception as e:
            return render_template("manual_mat.html", sesion=sesion, sid=sid,
                                   texto=texto, errors=[], error_sgs=f"Error consultando SGS: {e}")

        save_sesion(sid)
        return redirect(url_for("manual_preview", sid=sid))

    return render_template("manual_mat.html", sesion=sesion, sid=sid,
                           texto="", errors=[])


@app.route("/manual/preview/<sid>")
@login_required
def manual_preview(sid):
    from manual_balance import get_sesion, save_sesion, preparar_preview_materiales
    sesion = get_sesion(sid)
    if not sesion:
        return redirect(url_for("index") + "?tab=sgs")
    if not sesion.get("sgs_data"):
        sesion["sgs_data"] = {"padre": sesion.get("padre", {}), "nodos": []}
        save_sesion(sid)
    preview_materiales = []
    try:
        nodos_mat = sesion.get("nodos_mat", [])
        preview_materiales = preparar_preview_materiales({"nodos": nodos_mat})
    except Exception:
        pass
    return render_template("manual_preview.html", sesion=sesion, sid=sid,
                           preview_materiales=preview_materiales,
                           reanalisis=request.args.get("reanalisis"),
                           reanalisis_error=request.args.get("reanalisis_error"))


@app.route("/manual/guardar/<sid>", methods=["POST"])
@login_required
def manual_guardar(sid):
    """AJAX: fuerza el guardado de la sesión completa (MO + MAT + líderes) y
    confirma qué quedó guardado — el usuario puede cerrar el navegador."""
    from manual_balance import get_sesion, save_sesion
    sesion = get_sesion(sid)
    if not sesion:
        return {"ok": False, "error": "Sesión expirada"}, 404
    save_sesion(sid)
    n_mo  = len(sesion.get("rows_mo") or [])
    n_nod = len(sesion.get("nodos_mat") or [])
    n_lid = len(sesion.get("lideres") or [])
    return {"ok": True,
            "resumen": f"{n_mo} actividades MO, {n_nod} nodo(s) de materiales y {n_lid} líder(es)"}


@app.route("/manual/reanalizar/<sid>", methods=["POST"])
@login_required
def manual_reanalizar(sid):
    """
    Reconsulta SGS en caliente desde la vista previa (sin repetir el flujo):
    baja bandejas/despachos/devoluciones frescos y recalcula los nodos de
    materiales con la misma selección de PROs que ya tenía la sesión.
    """
    from manual_balance import get_sesion, save_sesion
    sesion = get_sesion(sid)
    if not sesion:
        return render_template("index.html", tab="sgs",
                               error_sgs_libre="Sesión expirada — ingresa el código del proyecto nuevamente.")
    try:
        cliente = _get_sgs_cliente()
        data = cliente.balance_completo(sesion["parent_id"])
        sesion["sgs_data"]  = data
        sesion["sgs_ready"] = True
        sesion["sgs_error"] = None

        sel = set(sesion.get("nodos_sel") or [])
        if sel:
            nodos_cache = data.get("nodos", [])
            en_cache = {n.get("code") for n in nodos_cache}
            nodos_extra = []
            for pro in [p for p in sel if p not in en_cache]:
                try:
                    p = cliente.buscar_por_codigo(pro)
                    if p:
                        nodos_extra.extend(cliente.balance_completo(p["id"]).get("nodos", []))
                except Exception:
                    pass  # best effort — igual que en paso2/confirmar
            sesion["nodos_mat"] = [n for n in nodos_cache if n.get("code") in sel] + nodos_extra

        save_sesion(sid)
        return redirect(url_for("manual_preview", sid=sid, reanalisis="ok"))
    except Exception as e:
        save_sesion(sid)
        return redirect(url_for("manual_preview", sid=sid, reanalisis_error=str(e)[:200]))


@app.route("/manual/generar/<sid>", methods=["POST"])
@login_required
def manual_generar(sid):
    from manual_balance import get_sesion, generar_excel_manual
    sesion = get_sesion(sid)
    if not sesion:
        return redirect(url_for("index") + "?tab=sgs")

    excel_bytes = generar_excel_manual(sesion)
    ts       = ahora_co().strftime("%Y%m%d_%H%M")
    lcl      = sesion.get("lcl_odm", "balance")
    filename = f"balance_sgs_manual_{lcl}_{ts}.xlsx"

    token = str(uuid.uuid4())
    _save_archivo(token, excel_bytes, filename)
    return redirect(url_for("descargar", token=token))


@app.route("/test-sgs")
@login_required
def test_sgs():
    """Diagnóstico: muestra estado de conexión SGS y busca PRO pasado por query string."""
    from integracion import cfg_desde_archivo
    from sgs_client import SGSClient
    import traceback

    resultado = {}
    pro = request.args.get("pro", "").strip()

    try:
        cfg = cfg_desde_archivo()
        resultado["config"] = {
            "url":   cfg.get("sgs_url", ""),
            "db":    cfg.get("sgs_db",  "(vacío — se auto-detecta)"),
            "login": cfg.get("sgs_login", ""),
        }
    except Exception as e:
        return {"error": f"cfg_desde_archivo falló: {e}"}, 500

    try:
        if not cfg["sgs_db"]:
            detected = SGSClient.detectar_db(cfg["sgs_url"])
            resultado["db_detectada"] = detected or "ninguna"
        cliente = _get_sgs_cliente()
        resultado["auth"] = "OK"
        resultado["uid"]  = cliente.uid
    except Exception as e:
        resultado["auth"] = f"FALLÓ: {e}"
        resultado["traceback"] = traceback.format_exc()
        return resultado, 200

    if pro:
        try:
            p = cliente.buscar_por_codigo(pro)
            resultado["proyecto"] = {"id": p["id"], "name": p.get("name"), "code": p.get("code")} if p else "no encontrado"
        except Exception as e:
            resultado["proyecto"] = f"error: {e}"

    return resultado, 200


@app.route("/health")
def health():
    return {"ok": True, "service": "balance-integrado"}


@app.route("/reiniciar", methods=["POST"])
def reiniciar():
    import threading
    def _exit():
        import time
        import os
        time.sleep(0.5)
        os._exit(3)  # código 3 → launcher.py reinicia el proceso
    threading.Thread(target=_exit, daemon=True).start()
    return ("", 204)


@app.route("/debug")
@login_required
def debug():
    """Verifica qué versión de integracion.py está cargada."""
    import sys
    import inspect
    from integracion import BalanceIntegrado, _normalizar_desc
    src = inspect.getsource(BalanceIntegrado._generar_excel)
    hojas = [l.strip() for l in src.splitlines() if "hoja" in l]
    modulo = sys.modules.get("integracion")
    return {
        "integracion_file": modulo.__file__ if modulo else "no cargado",
        "hojas": hojas,
        "tiene_normalizar_desc": hasattr(modulo, "_normalizar_desc"),
        "tiene_construir_desc_idx": hasattr(modulo, "_construir_desc_idx"),
        "test_normalizar": _normalizar_desc("Fusible Dual 1.0A 15 kV"),
    }


# ══════════════════════════════════════════════════════════════════════════════
# Módulo Formato RG02 · RG03 — registro fotográfico con descripciones IA
# (funcionalidad independiente: rg03_fotos.py + rg02_parser.py + rg03_ia.py)
# ══════════════════════════════════════════════════════════════════════════════

_RG03_DIR = os.path.join(tempfile.gettempdir(), "nuevo_rg03")
os.makedirs(_RG03_DIR, exist_ok=True)


def _rg03_save(sid: str, data: dict):
    with open(os.path.join(_RG03_DIR, f"{sid}.pkl"), "wb") as _f:
        _pickle.dump(data, _f)


def _rg03_get(sid: str):
    p = os.path.join(_RG03_DIR, f"{sid}.pkl")
    if not os.path.isfile(p):
        return None
    try:
        with open(p, "rb") as _f:
            return _pickle.load(_f)
    except Exception:
        return None


@app.route("/rg03/analizar", methods=["POST"])
@login_required
def rg03_analizar():
    from rg02_parser import combinar_rg02, parsear_pdt, parsear_rg02
    from rg03_fotos import parsear_rg03

    f_rg03 = request.files.get("rg03")
    f_rg02s = [f for f in request.files.getlist("rg02") if f and f.filename]
    f_pdt = request.files.get("pdt")

    if not f_rg03 or not f_rg03.filename:
        return render_template("index.html", tab="rg03", error_rg03="Debes subir el archivo RG03 (.xlsm).")
    if not f_rg02s:
        return render_template("index.html", tab="rg03", error_rg03="Debes subir al menos una planilla RG02 (PDF).")

    rg03_bytes = f_rg03.read()
    try:
        res = parsear_rg03(rg03_bytes)
    except Exception as e:
        return render_template("index.html", tab="rg03", error_rg03=f"No se pudo leer el RG03: {e}")

    advertencias = list(res["advertencias"])
    parsed = []
    for f in f_rg02s:
        try:
            parsed.append(parsear_rg02(f.read()))
        except Exception as e:
            return render_template("index.html", tab="rg03", error_rg03=f"No se pudo leer el RG02 «{f.filename}»: {e}")
    rg02 = combinar_rg02(parsed)
    advertencias.extend(rg02.pop("advertencias", []))

    if f_pdt and f_pdt.filename:
        try:
            pdt = parsear_pdt(f_pdt.read())
            rg02["pdt_texto"] = pdt["texto"]
            if pdt["ticket"] and rg02.get("ticket") and pdt["ticket"] != rg02["ticket"]:
                advertencias.append(f"El ticket del PDT ({pdt['ticket']}) no coincide con el del RG02 ({rg02['ticket']}).")
        except Exception as e:
            advertencias.append(f"No se pudo leer el PDT (se continúa sin él): {e}")

    # cruce RG03 ↔ RG02
    h = res["header"]
    if h.get("lcl") and rg02.get("lcl") and h["lcl"] != rg02["lcl"]:
        advertencias.append(f"LCL distinto: RG03={h['lcl']} vs RG02={rg02['lcl']} — ¿son del mismo proyecto?")
    if h.get("ticket") and rg02.get("ticket") and h["ticket"] != rg02["ticket"]:
        advertencias.append(f"Ticket distinto: RG03={h['ticket']} vs RG02={rg02['ticket']} — ¿son del mismo proyecto?")
    for b in res["bloques"]:
        if b["nodo"] and b["nodo"] not in rg02["nodos"]:
            advertencias.append(f"Foto {b['n']}: el nodo «{b['nodo']}» no tiene prestaciones en la planilla RG02.")

    sid = uuid.uuid4().hex[:12]
    _rg03_save(sid, {
        "filename": f_rg03.filename,
        "rg03_bytes": rg03_bytes,
        "header": h,
        "bloques": res["bloques"],
        "rg02": rg02,
        "advertencias": advertencias,
        "textos": {},
    })
    return redirect(url_for("rg03_ver", sid=sid))


@app.route("/rg03/<sid>")
@login_required
def rg03_ver(sid):
    ses = _rg03_get(sid)
    if ses is None:
        return redirect(url_for("index"))
    import rg03_ia
    bloques = [{k: v for k, v in b.items() if k != "jpg"} for b in ses["bloques"]]
    nodos_qp = {n: [f"{p['qp']} ×{p['cantidad']}" for p in ps] for n, ps in ses["rg02"]["nodos"].items()}
    # nodos en orden de aparición, con sus fotos (para generar por secuencia)
    nodos_orden = []
    for b in bloques:
        entrada = next((x for x in nodos_orden if x["nodo"] == b["nodo"]), None)
        if entrada is None:
            entrada = {"nodo": b["nodo"], "fotos": []}
            nodos_orden.append(entrada)
        entrada["fotos"].append(b["n"])
    return render_template(
        "rg03.html", sid=sid, filename=ses["filename"], header=ses["header"],
        bloques=bloques, textos=ses["textos"], nodos_qp=nodos_qp,
        nodos_orden=nodos_orden,
        rg02=ses["rg02"], advertencias=ses["advertencias"],
        hay_api_key=rg03_ia.hay_api_key(),
    )


@app.route("/rg03/foto/<sid>/<int:n>")
@login_required
def rg03_foto(sid, n):
    ses = _rg03_get(sid)
    if ses is None:
        return "", 404
    for b in ses["bloques"]:
        if b["n"] == n and b.get("jpg"):
            mime = "image/png" if b["jpg"][:4] == b"\x89PNG" else "image/jpeg"
            return send_file(io.BytesIO(b["jpg"]), mimetype=mime)
    return "", 404


@app.route("/rg03/api/generar/<sid>/<int:n>", methods=["POST"])
@login_required
def rg03_generar(sid, n):
    """Regenera UNA foto — envía toda la secuencia del nodo como contexto
    visual (flujo de sucesos) pero devuelve solo la observación pedida."""
    ses = _rg03_get(sid)
    if ses is None:
        return {"error": "Sesión no encontrada — vuelve a subir los archivos."}, 404
    bloque = next((b for b in ses["bloques"] if b["n"] == n), None)
    if bloque is None:
        return {"error": f"No existe la foto {n} en esta sesión."}, 404
    bloques_nodo = [b for b in ses["bloques"] if b["nodo"] == bloque["nodo"]]
    from rg03_ia import generar_nodo
    try:
        salida = generar_nodo(bloques_nodo, ses["rg02"], solo_foto=n)
    except RuntimeError as e:
        return {"error": str(e)}, 502
    except Exception as e:
        return {"error": f"Error inesperado: {e}"}, 500
    ses["textos"][n] = salida[n]
    _rg03_save(sid, ses)
    return {"texto": salida[n]}


@app.route("/rg03/api/generar-nodo/<sid>/<nodo>", methods=["POST"])
@login_required
def rg03_generar_nodo(sid, nodo):
    """Genera TODAS las fotos de un nodo en una sola llamada multi-imagen,
    manteniendo la coherencia de la secuencia."""
    ses = _rg03_get(sid)
    if ses is None:
        return {"error": "Sesión no encontrada — vuelve a subir los archivos."}, 404
    bloques_nodo = [b for b in ses["bloques"] if b["nodo"] == nodo]
    if not bloques_nodo:
        return {"error": f"No hay fotos del nodo {nodo} en esta sesión."}, 404
    from rg03_ia import generar_nodo
    try:
        salida = generar_nodo(bloques_nodo, ses["rg02"])
    except RuntimeError as e:
        return {"error": str(e)}, 502
    except Exception as e:
        return {"error": f"Error inesperado: {e}"}, 500
    ses["textos"].update(salida)
    _rg03_save(sid, ses)
    return {"textos": {str(k): v for k, v in salida.items()}}


@app.route("/rg03/descargar/<sid>", methods=["POST"])
@login_required
def rg03_descargar(sid):
    ses = _rg03_get(sid)
    if ses is None:
        return redirect(url_for("index"))
    from rg03_fotos import escribir_observaciones
    try:
        textos = {int(k): v for k, v in _json.loads(request.form.get("textos", "{}")).items()}
    except Exception:
        textos = {}
    ses["textos"].update({k: v for k, v in textos.items() if str(v).strip()})
    _rg03_save(sid, ses)
    if not ses["textos"]:
        return redirect(url_for("rg03_ver", sid=sid))
    salida = escribir_observaciones(ses["rg03_bytes"], ses["textos"])
    base = os.path.splitext(ses["filename"])[0]
    return send_file(
        io.BytesIO(salida),
        as_attachment=True,
        download_name=f"{base} - DILIGENCIADO.xlsm",
        mimetype="application/vnd.ms-excel.sheet.macroEnabled.12",
    )


# ══════════════════════════════════════════════════════════════════════════════
# Módulo Materiales Sony — cargue de instalaciones desde bandejas SGS
# (funcionalidad independiente: materiales_sony.py)
# ══════════════════════════════════════════════════════════════════════════════

def _anexo_info_safe():
    try:
        from materiales_sony import anexo_status
        return anexo_status()
    except Exception:
        return {"cargado": False}


def _render_index_mat(error=None, ok=None, **kw):
    return render_template("index.html", tab="materiales", anexo_info=_anexo_info_safe(),
                           error_materiales=error, ok_materiales=ok, **kw)


@app.route("/materiales/anexo", methods=["POST"])
@login_required
def materiales_anexo():
    """Sube o reemplaza el ANEXO Nº 14 (persistente entre reinicios)."""
    from materiales_sony import guardar_anexo
    archivo = request.files.get("anexo_file")
    if not archivo or not archivo.filename:
        return _render_index_mat(error="Selecciona el archivo ANEXO Nº 14 (.xlsx).")
    try:
        st = guardar_anexo(archivo.read(), archivo.filename)
    except ValueError as e:
        return _render_index_mat(error=str(e))
    except Exception as e:
        return _render_index_mat(error=f"No se pudo leer el ANEXO: {e}")
    return _render_index_mat(ok=f"ANEXO cargado: {st['total']} materiales de aportación.")


@app.route("/materiales/anexo/eliminar", methods=["POST"])
@login_required
def materiales_anexo_eliminar():
    from materiales_sony import eliminar_anexo
    eliminar_anexo()
    return _render_index_mat(ok="ANEXO Nº 14 eliminado.")


@app.route("/materiales/buscar", methods=["POST"])
@login_required
def materiales_buscar():
    """Paso 1 — cualquier PRO (hijo o padre) → resuelve el padre y sus nodos."""
    from materiales_sony import nueva_sesion_mat

    pro_raw = request.form.get("pro_materiales", "").strip()
    if not pro_raw:
        return _render_index_mat(error="Ingresa un código PRO (hijo o padre).")

    try:
        cliente = _get_sgs_cliente()
        if pro_raw.isdigit():
            proyecto = cliente.leer_proyecto(int(pro_raw))
            if proyecto is None:
                return _render_index_mat(error=f"Proyecto id {pro_raw} no encontrado en SGS.",
                                         pro_materiales=pro_raw)
        else:
            proyecto = cliente.buscar_por_codigo(pro_raw)
            if not proyecto:
                return _render_index_mat(error=f"'{pro_raw}' no encontrado en SGS.",
                                         pro_materiales=pro_raw)

        # PRO hijo → subir al padre automáticamente
        parent_field = proyecto.get("parent_id")
        if isinstance(parent_field, list) and parent_field:
            parent_id = parent_field[0]
            pro_origen = proyecto.get("code", pro_raw)
        else:
            parent_id = proyecto["id"]
            pro_origen = None

        data  = cliente.balance_completo(parent_id)
        padre = data["padre"]
        nodos = data["nodos"]

        sesion = {
            "padre":      padre,
            "parent_id":  parent_id,
            "nodos":      nodos,
            "pro_origen": pro_origen,   # el hijo que digitó el usuario (para resaltarlo)
        }
        sid = nueva_sesion_mat(sesion)
        return redirect(url_for("materiales_nodos", sid=sid))

    except Exception as e:
        return _render_index_mat(error=f"Error consultando SGS: {e}", pro_materiales=pro_raw)


@app.route("/materiales/nodos/<sid>", methods=["GET", "POST"])
@login_required
def materiales_nodos(sid):
    """Paso 2 — checklist de nodos hijos del padre."""
    from materiales_sony import get_sesion_mat, save_sesion_mat
    sesion = get_sesion_mat(sid)
    if not sesion:
        return _render_index_mat(error="La sesión expiró — busca el PRO nuevamente."), 404

    if request.method == "POST":
        sel = request.form.getlist("nodo_codes")
        if not sel:
            pass  # cae al render con error
        else:
            sesion["nodos_sel"] = sel
            save_sesion_mat(sid, sesion)
            return redirect(url_for("materiales_bandejas", sid=sid))

    nodos_resumen = []
    for n in sesion["nodos"]:
        bans = n.get("bandejas", [])
        nodos_resumen.append({
            "pro":      n.get("code", ""),
            "nodo":     n.get("nodo") or "—",
            "name":     n.get("name", ""),
            "status":   (n.get("status")[1] if isinstance(n.get("status"), list)
                         else str(n.get("status") or "—")),
            "despachos":     sum(1 for b in bans if b.get("type", "deliver") != "devolution"),
            "devoluciones":  sum(1 for b in bans if b.get("type", "deliver") == "devolution"),
            "es_origen": n.get("code") == sesion.get("pro_origen"),
        })
    error = "Selecciona al menos un nodo." if request.method == "POST" else None
    return render_template("materiales_nodos.html", sid=sid, padre=sesion["padre"],
                           nodos=nodos_resumen, pro_origen=sesion.get("pro_origen"),
                           error=error)


@app.route("/materiales/bandejas/<sid>", methods=["GET", "POST"])
@login_required
def materiales_bandejas(sid):
    """Paso 3 — bandejas de los nodos elegidos, vista tipo SGS."""
    from materiales_sony import get_sesion_mat, save_sesion_mat, presentar_bandeja, construir_cuadre
    sesion = get_sesion_mat(sid)
    if not sesion:
        return _render_index_mat(error="La sesión expiró — busca el PRO nuevamente."), 404

    nodos_sel = sesion.get("nodos_sel", [])
    nodos = [n for n in sesion["nodos"] if n.get("code") in nodos_sel]

    if request.method == "POST":
        ban_ids = {int(x) for x in request.form.getlist("ban_ids") if x.isdigit()}
        if ban_ids:
            sesion["ban_ids_sel"] = sorted(ban_ids)
            cuadre = construir_cuadre(nodos, ban_ids)
            sesion["cuadre"] = cuadre
            save_sesion_mat(sid, sesion)
            return redirect(url_for("materiales_cuadre", sid=sid))

    grupos = []
    for n in nodos:
        bans = sorted(n.get("bandejas", []), key=lambda b: str(b.get("date") or ""))
        grupos.append({
            "pro":      n.get("code", ""),
            "nodo":     n.get("nodo") or "—",
            "name":     n.get("name", ""),
            "bandejas": [presentar_bandeja(b, n) for b in bans],
        })
    error = ("Selecciona al menos una bandeja de despacho."
             if request.method == "POST" else None)
    return render_template("materiales_bandejas.html", sid=sid, padre=sesion["padre"],
                           grupos=grupos, error=error)


@app.route("/materiales/cuadre/<sid>")
@login_required
def materiales_cuadre(sid):
    """Paso 4 — cuadre editable: despachos + devoluciones + aportación."""
    from materiales_sony import get_sesion_mat, get_prefs
    sesion = get_sesion_mat(sid)
    if not sesion or not sesion.get("cuadre"):
        return _render_index_mat(error="La sesión expiró — busca el PRO nuevamente."), 404

    cuadre = sesion["cuadre"]
    prefs  = get_prefs()
    hoy    = ahora_co().strftime("%Y-%m-%d")
    # Hojas fotografiadas/validadas en el módulo de devoluciones → columna "Dev. hoja"
    # (se consulta FRESCO en cada carga: almacén puede validar y recargar)
    from materiales_sony import aplicar_hojas_devoluciones
    info_hojas = aplicar_hojas_devoluciones(cuadre)
    return render_template("materiales_cuadre.html", sid=sid, padre=sesion["padre"],
                           cuadre=cuadre, usuario_def=prefs.get("usuario_cargue", ""),
                           hoy=hoy, info_hojas=info_hojas)


@app.route("/materiales/generar/<sid>", methods=["POST"])
@login_required
def materiales_generar(sid):
    """Paso final — arma las filas desde el formulario y genera el xlsx plantilla."""
    from materiales_sony import (get_sesion_mat, generar_cargue, save_prefs,
                                 _parse_fecha, _num)
    sesion = get_sesion_mat(sid)
    if not sesion or not sesion.get("cuadre"):
        return _render_index_mat(error="La sesión expiró — busca el PRO nuevamente."), 404

    form    = request.form
    usuario = form.get("usuario", "").strip()
    if usuario:
        save_prefs(usuario_cargue=usuario)
    fecha_asig = ahora_co()

    advertencias = list(sesion["cuadre"].get("advertencias", []))
    filas = []

    for nc in sesion["cuadre"]["nodos_cuadre"]:
        pro       = nc["pro"]
        id_nodo   = form.get(f"idnodo_{pro}", "").strip()
        pdi       = form.get(f"pdi_{pro}", "").strip()
        obs       = form.get(f"obs_{pro}", "").strip()
        reintegro = _parse_fecha(form.get(f"reintegro_{pro}", ""))
        if not id_nodo:
            advertencias.append(f"{pro}: sin Id_nodo — las filas salen con la celda vacía.")

        for desp in nc["despachos"]:
            f_ent = _parse_fecha(form.get(f"fban_{desp['ban_id']}", "")) \
                    or _parse_fecha(desp["fecha"])
            for fila in desp["filas"]:
                i = fila["idx"]
                if not form.get(f"r{i}_incl"):
                    continue
                qty = _num(form.get(f"r{i}_qty", fila["qty"]))
                dev = _num(form.get(f"r{i}_dev", 0))
                if dev > qty:
                    advertencias.append(
                        f"{pro} / {desp['ban_name']} / {fila['cod']}: devolución ({dev:g}) "
                        f"mayor que lo entregado ({qty:g}).")
                filas.append({
                    "id_nodo":          id_nodo,
                    "usuario":          usuario,
                    "fecha_asignacion": fecha_asig,
                    "codigo":           fila["cod"],
                    "qty":              qty,
                    "dev":              dev,
                    "fecha_entrega":    f_ent,
                    "pdi":              pdi,
                    "consecutivo":      desp["ban_name"],
                    "obs":              obs,
                    "k":                form.get(f"r{i}_k", str(fila["k"])),
                    "fecha_reintegro":  reintegro,
                })

    if not filas:
        return _render_index_mat(error="No quedó ninguna fila incluida — revisa el cuadre.")

    excel_bytes = generar_cargue(filas)
    padre = sesion["padre"]
    odm   = padre.get("order_number") or padre.get("odm") or padre.get("code") or "proyecto"
    ts    = ahora_co().strftime("%Y%m%d_%H%M")
    filename = f"cargue_instalaciones_{odm}_{ts}.xlsx"
    token = uuid.uuid4().hex
    _save_archivo(token, excel_bytes, filename, advertencias)

    resumen = {
        "n_filas":      len(filas),
        "n_nodos":      len(sesion["cuadre"]["nodos_cuadre"]),
        "total_ent":    sum(float(f["qty"] or 0) for f in filas),
        "total_dev":    sum(float(f["dev"] or 0) for f in filas),
        "n_aportacion": sum(1 for f in filas if str(f["k"]) == "1"),
    }
    return render_template("materiales_resultado.html", sid=sid, padre=padre,
                           token=token, filename=filename, filas=filas,
                           resumen=resumen, advertencias=advertencias)


# ── Main ───────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    # Añadir /nuevo al path para que `from integracion import ...` funcione
    if _HERE not in sys.path:
        sys.path.insert(0, _HERE)

    port = int(os.environ.get("PORT", 5002))
    print(f"[/nuevo] Iniciando en http://localhost:{port}")
    app.run(host="0.0.0.0", port=port, debug=False)
