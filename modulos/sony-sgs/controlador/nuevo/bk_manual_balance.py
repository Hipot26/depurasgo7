"""
manual_balance.py — Flujo de Balance Manual SGS Puro (4 pasos).

Sesiones: LCL search (index.html) → MO → Materiales → Preview + Excel
"""
import re
import os
import uuid
import io
import json
import time
import tempfile
import threading
from datetime import datetime, timezone, timedelta
from collections import defaultdict
from pathlib import Path

# ── Hora Colombia (UTC-5, sin horario de verano) ─────────────────────────────
# El servidor corre en UTC; todo timestamp visible al usuario va en hora CO.
TZ_CO = timezone(timedelta(hours=-5))


def ahora_co() -> datetime:
    """datetime naive en hora Colombia — para timestamps visibles al usuario."""
    return datetime.now(TZ_CO).replace(tzinfo=None)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from generar_balance import _cruce_nodo
from sgs_client import parsear_producto

# ── Baremo de precios MDO&MAT ─────────────────────────────────────────────────
_BAREMO_PATHS = [
    Path(__file__).parent / "precios_mdo_mat.json",
]
_baremo_mem: dict = {}


def get_precios() -> dict:
    """Carga {QP_code: {precio, unidad, descripcion}} desde JSON cache."""
    global _baremo_mem
    if _baremo_mem:
        return _baremo_mem
    for p in _BAREMO_PATHS:
        if p.exists():
            try:
                with open(p, encoding="utf-8") as f:
                    _baremo_mem = json.load(f)
                return _baremo_mem
            except Exception:
                pass
    return _baremo_mem


def enriquecer_mo(rows: list) -> list:
    """Agrega desc, unidad, precio_unit, subtotal a cada fila de MO."""
    precios = get_precios()
    for r in rows:
        info = precios.get(r["qp"], {})
        r["desc"]       = info.get("descripcion", "")
        r["unidad"]     = info.get("unidad", "C/U")
        r["precio_unit"] = info.get("precio", 0.0)
        r["subtotal"]   = round(r["precio_unit"] * r["cantidad"], 2)
        r["en_baremo"]  = bool(info)
    return rows


# ── Colores (esquema verde SGS) ───────────────────────────────────────────────
C_HDR    = "375623"   # verde oscuro — títulos principales
C_SUBHDR = "E2EFDA"   # verde claro  — cabecera columnas
C_PAR    = "F2F2F2"   # gris         — fila par alternada
C_TOTAL  = "FFE699"   # amarillo     — fila total
C_OK     = "C6EFCE"   # verde        — OK
C_AMBAR  = "FFEB9C"   # ámbar        — pendiente
C_ROJO   = "FFC7CE"   # rojo         — error / descuadre
C_GRIS   = "D9D9D9"   # gris oscuro  — sin movimiento
C_DEVOL  = "DEEBF7"   # azul claro   — devolución completa (no liquidar)
BLANCO   = "FFFFFF"

# ── Sesiones ─────────────────────────────────────────────────────────────────
_sm: dict = {}
# Directorio persistente: antes vivía en el tempdir del SO y las sesiones se
# perdían con las limpiezas de /tmp — imposible "retomar" un balance días después.
_SESSION_DIR = os.environ.get("NUEVO_SESSIONS_DIR") or os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "data", "sessions")
os.makedirs(_SESSION_DIR, exist_ok=True)

# Migración única: adoptar las sesiones que quedaron en el viejo dir temporal
_OLD_SESSION_DIR = os.path.join(tempfile.gettempdir(), "nuevo_sessions")
if os.path.isdir(_OLD_SESSION_DIR) and os.path.realpath(_OLD_SESSION_DIR) != os.path.realpath(_SESSION_DIR):
    import shutil as _shutil
    for _f in os.listdir(_OLD_SESSION_DIR):
        if _f.startswith("s_") and _f.endswith(".json") and \
           not os.path.exists(os.path.join(_SESSION_DIR, _f)):
            try:
                _shutil.copy2(os.path.join(_OLD_SESSION_DIR, _f), os.path.join(_SESSION_DIR, _f))
            except OSError:
                pass


def _session_path(sid: str) -> str:
    return os.path.join(_SESSION_DIR, f"s_{sid}.json")


def save_sesion(sid: str):
    if sid not in _sm:
        return
    # Marca de tiempo (reloj de pared, compartido entre workers del mismo host):
    # permite que get_sesion detecte cuándo otro worker guardó cambios más nuevos.
    _sm[sid]["_saved_at"] = time.time()
    try:
        with open(_session_path(sid), "w", encoding="utf-8") as f:
            json.dump(_sm[sid], f, ensure_ascii=False, default=str)
    except Exception:
        # sgs_data puede no ser serializable — guardar sin él para al menos preservar el estado
        try:
            slim = {k: v for k, v in _sm[sid].items() if k != "sgs_data"}
            with open(_session_path(sid), "w", encoding="utf-8") as f:
                json.dump(slim, f, ensure_ascii=False, default=str)
        except Exception:
            pass


def nueva_sesion(padre: dict, parent_id: int) -> str:
    sid = uuid.uuid4().hex[:8]
    odm = padre.get("order_number") or padre.get("odm") or str(parent_id)
    _sm[sid] = {
        "padre":      padre,
        "parent_id":  parent_id,
        "lcl_odm":    odm,
        "rows_mo":    [],
        "errs_mo":    [],
        "rows_mat":   [],
        "errs_mat":   [],
        "modo_mat":   None,
        "nodos_sel":  [],
        "sgs_data":   None,
        "sgs_ready":  None,
        "sgs_error":  None,
        "texto_mo":            "",
        "texto_nodos":         "",
        "nodos_pendientes_sgs": [],
        "lideres":              [],
        "creado_at":  ahora_co().strftime("%Y-%m-%d %H:%M"),
    }
    save_sesion(sid)
    return sid


def get_sesion(sid: str):
    if sid in _sm:
        s = _sm[sid]
        # Reconciliación multi-worker: si OTRO worker guardó cambios más nuevos en
        # disco (líderes, MO, materiales, nodos, fin del fetch…), adoptarlos. Sin
        # esto, el worker que corrió el fetch se queda con su copia en memoria y
        # nunca ve lo que otros workers escribieron (p.ej. los líderes del paso 2).
        path = _session_path(sid)
        if os.path.isfile(path):
            try:
                with open(path, encoding="utf-8") as f:
                    disk = json.load(f)
            except Exception:
                disk = None
            if disk is not None and disk.get("_saved_at", 0) > s.get("_saved_at", 0):
                # Preservar sgs_data en memoria si el disco no lo tiene (es grande
                # y a veces se guarda en modo "slim" sin él)
                if disk.get("sgs_data") is None and s.get("sgs_data") is not None:
                    disk["sgs_data"] = s["sgs_data"]
                _sm[sid] = disk
                return disk
        return s
    # Fallback: cargar desde archivo si el worker se reinició
    path = _session_path(sid)
    if os.path.isfile(path):
        try:
            with open(path, encoding="utf-8") as f:
                data = json.load(f)
            _sm[sid] = data
            return data
        except Exception:
            pass
    return None


def listar_sesiones(limit: int = 12) -> list:
    """Resumen de los balances guardados en disco, más recientes primero."""
    try:
        archivos = [f for f in os.listdir(_SESSION_DIR)
                    if f.startswith("s_") and f.endswith(".json")]
    except OSError:
        return []
    archivos.sort(key=lambda f: os.path.getmtime(os.path.join(_SESSION_DIR, f)),
                  reverse=True)
    out = []
    vistos = set()
    for f in archivos[:60]:
        if len(out) >= limit:
            break
        try:
            with open(os.path.join(_SESSION_DIR, f), encoding="utf-8") as fh:
                s = json.load(fh)
        except Exception:
            continue
        # Mismo código diligenciado varias veces → mostrar solo la versión
        # más reciente (los archivos vienen ordenados por mtime descendente)
        clave = (s.get("lcl_odm") or "").strip().upper() or f[2:-5]
        if clave in vistos:
            continue
        vistos.add(clave)
        padre = s.get("padre") or {}
        n_mo  = len(s.get("rows_mo") or [])
        n_lid = len(s.get("lideres") or [])
        con_mat = bool(s.get("nodos_sel") or s.get("rows_mat"))
        if con_mat:
            paso = "Lista para generar"
        elif n_mo:
            paso = "MO diligenciada"
        else:
            paso = "Recién creada"
        saved_at = s.get("_saved_at")
        actualizado = (datetime.fromtimestamp(saved_at, TZ_CO).strftime("%Y-%m-%d %H:%M")
                       if saved_at else s.get("creado_at", ""))
        out.append({
            "sid":            f[2:-5],
            "lcl_odm":        s.get("lcl_odm", ""),
            "proyecto":       padre.get("name") or padre.get("code") or "",
            "pdl":            s.get("pdl", ""),
            "creado":         s.get("creado_at", ""),
            "actualizado":    actualizado,
            "n_mo":           n_mo,
            "n_lideres":      n_lid,
            "con_materiales": con_mat,
            "paso":           paso,
        })
    return out


def eliminar_sesion(sid: str) -> bool:
    """Borra un balance guardado (memoria + disco)."""
    _sm.pop(sid, None)
    try:
        path = _session_path(sid)
        if os.path.isfile(path):
            os.remove(path)
        return True
    except OSError:
        return False


def iniciar_fetch_sgs(sid: str, get_cliente_fn):
    """Lanza un hilo daemon que corre balance_completo() en background."""
    sesion = get_sesion(sid)
    if sesion is None:
        return
    sesion["sgs_ready"] = False
    sesion["sgs_error"] = None
    save_sesion(sid)

    def _fetch():
        s = _sm.get(sid)
        if s is None:
            path = _session_path(sid)
            if os.path.isfile(path):
                try:
                    with open(path, encoding="utf-8") as f:
                        s = json.load(f)
                    _sm[sid] = s
                except Exception:
                    return
            else:
                return
        try:
            cliente  = get_cliente_fn()
            data     = cliente.balance_completo(s["parent_id"])
            sgs_ok, sgs_err = True, None
        except Exception as e:
            data, sgs_ok, sgs_err = None, True, str(e)

        # Re-leer disco para preservar cambios del usuario hechos mientras SGS cargaba
        # (fix: evita que el hilo de Worker A sobreescriba rows_mo/lideres guardados por Worker B)
        path = _session_path(sid)
        try:
            with open(path, encoding="utf-8") as f:
                latest = json.load(f)
        except Exception:
            latest = _sm.get(sid) or s
        if data is not None:
            latest["sgs_data"] = data
        latest["sgs_ready"] = sgs_ok
        latest["sgs_error"] = sgs_err
        _sm[sid] = latest
        save_sesion(sid)

    threading.Thread(target=_fetch, daemon=True).start()


def resolver_nodos(sgs_nodos: list, mo_nodo_ids: list) -> dict:
    """
    Cruza los nodos de la MO con los nodos SGS usando ID canónico.
    - Nodos simples ('127') hacen match contra cables que los contienen ('127AL122').
    - Se deduplicam nodos repetidos de la MO (mismo nodo con distinto QP).
    """
    # Índice exacto: id_canónico → nodo SGS
    idx = {}
    # Índice primario: número base → nodo SGS cable (ej. "127" → nodo "127AL122")
    primary_idx = {}
    for n in sgs_nodos:
        cid = _nodo_id_canonico(n.get("nodo"))
        if not cid:
            continue
        idx[cid] = n
        # Indexar también por el número primario del cable
        m = re.match(r'^(\d+)', cid)
        if m:
            pnum = m.group(1)
            if pnum not in primary_idx:
                primary_idx[pnum] = n

    # Deduplicar nodos de MO (mismo nodo puede tener varios QP)
    seen: set = set()
    unique_ids = []
    for mid in mo_nodo_ids:
        k = str(mid)
        if k not in seen:
            seen.add(k)
            unique_ids.append(mid)

    encontrados = []
    no_encontrados = []

    for mid in unique_ids:
        cid = _nodo_id_canonico(str(mid))
        nodo_sgs = idx.get(cid)
        # Fallback: nodo simple "127" → buscar cable que empiece con ese número
        if not nodo_sgs and re.match(r'^\d+$', cid):
            nodo_sgs = primary_idx.get(cid)
        if nodo_sgs:
            encontrados.append({
                "num":      mid,
                "pro":      nodo_sgs.get("code", ""),
                "nodo_raw": nodo_sgs.get("nodo", mid),
                "sgs_id":   nodo_sgs.get("id"),
            })
        else:
            no_encontrados.append(mid)

    return {"encontrados": encontrados, "no_encontrados": no_encontrados}


def preparar_preview_materiales(sgs_data: dict) -> list:
    """Prepara filas de materiales para la vista previa (misma lógica que Excel)."""

    def _n(v):
        """Convierte float entero a int para display limpio."""
        if isinstance(v, float) and v.is_integer():
            return int(v)
        return v

    resultado = []
    for nodo in sgs_data.get("nodos", []):
        num   = _label_nodo(nodo.get("nodo"))
        code  = nodo.get("code", "")
        name  = nodo.get("name", "") or f"NODO {num}"
        sap_map   = _dom_a_sap(nodo)
        filas_raw = _cruce_nodo(nodo)

        nd_tray = nd_dev = nd_neto = nd_cons = 0.0
        filas = []
        for r in filas_raw:
            sap_cod = sap_map.get(r["cod"]) or (r["cod"].split("/")[-1] if "/" in r["cod"] else r["cod"])
            filas.append({
                "sap":          sap_cod,
                "desc":         r["desc"],
                "ban_desp":     _fmt_ban_names(r["ban_delivers"]),
                "ban_dev":      _fmt_ban_names(r["ban_devs"]),
                "tray":         _n(r["tray"]),
                "devoluciones": _n(r["devoluciones"]),
                "neto":         _n(r["desp_neto"]),
                "consumido":    _n(r["consumido"]),
                "estado":       r["estado"],
                "color":        r["color"],
            })
            nd_tray += r["tray"]
            nd_dev  += r["devoluciones"]
            nd_neto += r["desp_neto"]
            nd_cons += r["consumido"]

        resultado.append({
            "num":             num,
            "code":            code,
            "nombre_completo": name,
            "filas":           filas,
            "totales": {
                "tray": _n(nd_tray),
                "dev":  _n(nd_dev),
                "neto": _n(nd_neto),
                "cons": _n(nd_cons),
            },
        })
    return resultado


# ── Parseo de MO ──────────────────────────────────────────────────────────────

def parsear_mo(texto: str):
    """
    Nodo TAB QP_Code TAB Cantidad (desde Excel del supervisor).
    Retorna (rows, errors).
    """
    rows, errors = [], []
    for i, linea in enumerate(texto.strip().splitlines(), 1):
        partes = re.split(r"\t| {2,}", linea.strip())
        partes = [p.strip() for p in partes if p.strip()]
        if not partes:
            continue
        if len(partes) < 2:
            errors.append(f"Fila {i}: necesitas Nodo y QP (ej: 74&Tab;QP1630)")
            continue
        nodo_raw = partes[0].strip()
        s_up = re.sub(r"\s+", "", nodo_raw.upper())
        # Quitar "NODO " o "N" inicial si va seguido de dígito
        s_up = re.sub(r"^NODO\s*", "", s_up)
        if s_up.startswith("N") and len(s_up) > 1 and s_up[1].isdigit():
            s_up = s_up[1:]
        if re.match(r"^\d+[A-Za-z]?$", s_up):
            # Nodo simple: 74, 6A, 2B, etc.
            nodo = s_up.upper()
        else:
            # Nodo cable/tramo: N127_N126, 127_126, N1 AL N2, etc.
            m_tramo = re.match(r"^(\d+[A-Z]?)(?:_|AL|-)N?(\d+[A-Z]?)$", s_up)
            if m_tramo:
                nodo = f"{m_tramo.group(1)}AL{m_tramo.group(2)}"
            else:
                errors.append(f"Fila {i}: nodo inválido '{nodo_raw}'")
                continue
        qp = partes[1].strip().upper()
        if not re.match(r"^QP\d+", qp):
            errors.append(f"Fila {i}: código QP inválido '{qp}'")
            continue
        try:
            cantidad = float(partes[2].replace(",", ".")) if len(partes) >= 3 else 1.0
        except ValueError:
            errors.append(f"Fila {i}: cantidad inválida '{partes[2]}'")
            continue
        rows.append({"nodo": nodo, "qp": qp, "cantidad": cantidad})

    # Deduplicar por (nodo, qp) sumando cantidades
    dedup: dict = {}
    for r in rows:
        k = (r["nodo"], r["qp"])
        if k in dedup:
            dedup[k]["cantidad"] += r["cantidad"]
        else:
            dedup[k] = dict(r)
    return list(dedup.values()), errors


# ── Parseo de Materiales ──────────────────────────────────────────────────────

def detectar_nodos_solos(texto: str):
    """Retorna lista de identificadores de nodo si cada línea tiene UN SOLO token.
    Acepta números (133, 133.0) y nodos cable (N127_N126, N17_N29).
    Si una línea tiene TAB o dos o más espacios → es modo SAP, retorna None."""
    nodos = []
    for linea in texto.strip().splitlines():
        tok = linea.strip()
        if not tok:
            continue
        if re.search(r'\t| {2,}', tok):
            return None  # Múltiples columnas → modo SAP
        # Normalizar floats de Excel: "133.0" → "133"
        try:
            f = float(tok)
            n = int(f)
            tok = str(n) if abs(f - n) < 0.001 else tok
        except ValueError:
            pass
        nodos.append(tok)
    return list(dict.fromkeys(nodos)) if nodos else None


def parsear_mat(texto: str):
    """
    Nodo TAB SAP [TAB Descripción] TAB Cantidad.
    Retorna (rows, errors).
    """
    rows, errors = [], []
    for i, linea in enumerate(texto.strip().splitlines(), 1):
        partes = re.split(r"\t| {2,}", linea.strip())
        partes = [p.strip() for p in partes if p.strip()]
        if not partes:
            continue
        if len(partes) < 2:
            errors.append(f"Fila {i}: necesitas Nodo y Código SAP")
            continue
        try:
            nodo = str(int(float(partes[0])))
        except ValueError:
            errors.append(f"Fila {i}: nodo inválido '{partes[0]}'")
            continue
        sap = partes[1].strip()
        if len(partes) >= 4:
            desc = partes[2]
            try:
                cantidad = float(partes[3].replace(",", "."))
            except ValueError:
                errors.append(f"Fila {i}: cantidad inválida '{partes[3]}'")
                continue
        elif len(partes) == 3:
            try:
                cantidad = float(partes[2].replace(",", "."))
                desc = ""
            except ValueError:
                desc = partes[2]
                cantidad = 1.0
        else:
            desc, cantidad = "", 1.0
        rows.append({"nodo": nodo, "sap": sap, "desc": desc, "cantidad": cantidad})

    dedup: dict = {}
    for r in rows:
        k = (r["nodo"], r["sap"])
        if k in dedup:
            dedup[k]["cantidad"] += r["cantidad"]
        else:
            dedup[k] = dict(r)
    return list(dedup.values()), errors


# ── Helpers SGS ───────────────────────────────────────────────────────────────

def _num_nodo(nodo_val) -> str:
    """'NODO 74' → '74', '74' → '74', False/None → ''"""
    if not nodo_val:
        return ""
    m = re.match(r"^(?:NODO\s+)?(\d+[A-Z]?)", str(nodo_val).strip().upper())
    return m.group(1) if m else str(nodo_val).strip()


def _label_nodo(nodo_val) -> str:
    """
    Etiqueta completa del nodo para hojas y cabeceras Excel.
    Preserva el formato cable con ambos números.
    '127'       → '127'
    'NODO 127'  → '127'
    '127_126'   → '127_N126'
    'N127_N126' → '127_N126'
    '127AL126'  → '127_N126'
    """
    if not nodo_val:
        return ""
    s = str(nodo_val).strip().upper()
    s = re.sub(r"^NODO\s+", "", s)
    if s.startswith("N") and len(s) > 1 and s[1].isdigit():
        s = s[1:]   # N127_N126 → 127_N126
    m_cable = re.match(r"^(\d+[A-Z]?)(?:_|-|AL)N?(\d+[A-Z]?)$", s)
    if m_cable:
        return f"{m_cable.group(1)}_N{m_cable.group(2)}"
    m_alnum = re.match(r"^(\d+[A-Z]?)$", s)
    if m_alnum:
        return m_alnum.group(1)
    m_num = re.match(r"^(\d+)", s)
    return m_num.group(1) if m_num else s


def _nodo_id_canonico(nodo_val) -> str:
    """
    Normaliza cualquier representación de nodo al mismo ID canónico para comparación.
    '74'            → '74'
    'NODO 74'       → '74'
    '127AL126'      → '127AL126'
    'N127_N126'     → '127AL126'
    'NODO N119_N123'→ '119AL123'
    """
    if not nodo_val:
        return ""
    s = str(nodo_val).strip().upper()
    s = re.sub(r"^NODO\s+", "", s)          # quitar "NODO " inicial
    if s.startswith("N") and len(s) > 1 and s[1].isdigit():
        s = s[1:]                            # quitar N inicial (N127_N126 → 127_N126)
    m_cable = re.match(r"^(\d+[A-Z]?)(?:_|-|AL)N?(\d+[A-Z]?)$", s)
    if m_cable:
        return f"{m_cable.group(1)}AL{m_cable.group(2)}"
    # Nodo simple con sufijo letra (6A, 2B) — preservar completo
    m_alnum = re.match(r"^(\d+[A-Z]?)$", s)
    if m_alnum:
        return m_alnum.group(1)
    m_num = re.match(r"^(\d+)", s)
    return m_num.group(1) if m_num else s


def filtrar_nodos_sgs(nodos: list, numeros: list) -> list:
    """Filtra nodos SGS por número. Los nodos cable/tramo se ignoran.
    Si no hay coincidencia (p.ej. números de poste físico), devuelve todos."""
    nums_simples = {n for n in numeros if re.match(r"^\d+[A-Za-z]?$", str(n))}
    if not nums_simples:
        return nodos  # Solo nodos cable → todos
    resultado = [n for n in nodos if _num_nodo(n.get("nodo")) in nums_simples]
    return resultado if resultado else nodos  # Sin coincidencia → todos


def _dom_a_sap(nodo: dict) -> dict:
    """
    Construye {cod_dom: cod_sap} desde bandejas + liq_materiales.
    parsear_producto retorna (dom_prefix, sap_num, desc) — solo guardamos
    los pares donde sap_num no está vacío.
    """
    idx: dict = {}
    for ban in nodo.get("bandejas", []):
        for mat in ban.get("materiales", []):
            dom, sap, _ = parsear_producto(mat.get("product_id"))
            if dom and sap:
                idx[dom] = sap
    for m in nodo.get("liq_materiales", []):
        dom, sap, _ = parsear_producto(m.get("product_id"))
        if dom and sap:
            idx[dom] = sap
    return idx


# ── Excel: helpers de escritura ───────────────────────────────────────────────

def _fill(hex_color: str):
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, italic=False, color="000000", size=10):
    return Font(bold=bold, italic=italic, color=color, size=size, name="Calibri")

def _align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def _border():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)

def _w(ws, row, col, value="", bold=False, italic=False, bg=None, fg="000000",
       h="left", size=10, wrap=False, borders=True, num_fmt=None, span=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = _font(bold=bold, italic=italic, color=fg, size=size)
    c.alignment = _align(h=h, wrap=wrap)
    if bg:
        c.fill = _fill(bg)
    if borders:
        c.border = _border()
    if num_fmt:
        c.number_format = num_fmt
    if span:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=span)
    return c


def _get_sup(padre: dict) -> str:
    sup = padre.get("responsible_id")
    if isinstance(sup, list) and len(sup) >= 2:
        return sup[1]
    return str(sup) if sup else "—"


def _fmt_ban_names(lista):
    """[(ban_name, qty), ...] → 'BAN-103222, BAN-103993'  or  '—'"""
    if not lista:
        return "—"
    seen = []
    for name, _ in lista:
        if name not in seen:
            seen.append(name)
    return ", ".join(seen)


# ── Hoja 1: MO_MANUAL ─────────────────────────────────────────────────────────

def _hoja_mo(ws, rows_mo: list, lcl_odm: str, padre: dict, pendientes_sgs: list = None, lideres: list = None, supervisor: str = None, pdl: str = None):
    pendientes_set = set(pendientes_sgs or [])
    NCOLS = 9
    titulo = f"MO — MANO DE OBRA MANUAL  |  {lcl_odm}"
    _w(ws, 1, 1, titulo, bold=True, bg=C_HDR, fg=BLANCO, size=12, span=NCOLS)
    ws.row_dimensions[1].height = 22

    meta = [
        ("Proyecto:",    padre.get("name", "—")),
        ("Código padre:", padre.get("code", "—")),
        ("OdM / LCL:",   lcl_odm),
        ("PDL:",         pdl or "—"),
        ("Supervisor:",  supervisor or _get_sup(padre)),
        ("Generado:",    ahora_co().strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (label, val) in enumerate(meta, 2):
        _w(ws, i, 1, label, bold=True, borders=False, bg=C_SUBHDR)
        _w(ws, i, 2, val, borders=False, span=NCOLS)

    sep = len(meta) + 2
    _w(ws, sep, 1, "MANO DE OBRA", bold=True, bg=C_HDR, fg=BLANCO, span=NCOLS)

    hdr = sep + 1
    for j, h in enumerate(["#", "Nodo", "Código QP", "Descripción", "UM", "Precio Unit.", "Cantidad", "Subtotal ($)", "Observaciones"], 1):
        _w(ws, hdr, j, h, bold=True, bg=C_SUBHDR, h="center")

    NF_NUM = "0"
    NF_PES = "0"
    total_cant = 0.0
    total_sub  = 0.0
    for idx, r in enumerate(rows_mo, 1):
        dr  = hdr + idx
        es_pendiente = r["nodo"] in pendientes_set
        bg  = C_AMBAR if es_pendiente else (BLANCO if idx % 2 else C_PAR)
        _w(ws, dr, 1, idx,                     bg=bg, h="center")
        _w(ws, dr, 2, r["nodo"],               bg=bg, h="center")
        _w(ws, dr, 3, r["qp"],                 bg=bg, h="center")
        _w(ws, dr, 4, r.get("desc", ""),       bg=bg, wrap=True)
        _w(ws, dr, 5, r.get("unidad", "C/U"),  bg=bg, h="center")
        _w(ws, dr, 6, r.get("precio_unit", 0), bg=bg, h="right", num_fmt=NF_PES)
        _w(ws, dr, 7, r["cantidad"],            bg=bg, h="center", num_fmt=NF_NUM)
        _w(ws, dr, 8, r.get("subtotal", 0),    bg=bg, h="right",  num_fmt=NF_PES, bold=True)
        _w(ws, dr, 9, "PENDIENTE CREAR SGS" if es_pendiente else "",
           bg=bg, h="center", bold=es_pendiente, fg="7a5800" if es_pendiente else "000000")
        total_cant += r["cantidad"]
        total_sub  += r.get("subtotal", 0)

    tr = hdr + len(rows_mo) + 1
    _w(ws, tr, 1, "TOTAL MANO DE OBRA", bold=True, bg=C_TOTAL, span=7)
    _w(ws, tr, 8, total_sub,  bold=True, bg=C_TOTAL, h="right",  num_fmt=NF_PES)
    _w(ws, tr, 9, f"{len(pendientes_set)} pend. SGS" if pendientes_set else "",
       bold=True, bg=C_TOTAL, h="center")

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 48
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 11
    ws.column_dimensions["H"].width = 16
    ws.column_dimensions["I"].width = 22

    # ── Sección Líderes ────────────────────────────────────────────────────────
    if lideres:
        ldr_start = tr + 2
        _w(ws, ldr_start, 1, "DISTRIBUCIÓN DE LIQUIDACIÓN — LÍDERES RESPONSABLES",
           bold=True, bg=C_HDR, fg=BLANCO, size=11, span=NCOLS)
        ldr_hdr = ldr_start + 1
        for j, h in enumerate(["#", "Cédula", "Nombre del Líder", "% Liquidación", "", "", "", "", ""], 1):
            _w(ws, ldr_hdr, j, h, bold=bool(h), bg=C_SUBHDR, h="center")
        for idx_l, ldr in enumerate(lideres, 1):
            dr = ldr_hdr + idx_l
            bg = BLANCO if idx_l % 2 else C_PAR
            _w(ws, dr, 1, idx_l,                       bg=bg, h="center")
            _w(ws, dr, 2, ldr.get("cedula", ""),       bg=bg, h="center")
            _w(ws, dr, 3, ldr.get("nombre", ""),       bg=bg)
            _w(ws, dr, 4, f"{ldr.get('pct', 0)}%",    bg=bg, h="center", bold=True)
            for col in range(5, NCOLS + 1):
                _w(ws, dr, col, "", bg=bg)


# ── Hoja 2: SGS_BALANCE ───────────────────────────────────────────────────────

def _hoja_sgs_balance(wb, sgs_data: dict, lcl_odm: str):
    """Crea una hoja por nodo — freeze fijo en A3 (header nodo + columnas)."""
    nodos  = sgs_data["nodos"]
    pdl    = sgs_data.get("pdl", "") or ""
    NF     = "0"
    NCOLS  = 10

    COLS_HDR = ["#", "Código SAP", "Descripción", "Bandejas Desp.", "Bandejas Dev.",
                "Despachos", "Devoluciones", "Bal. Neto", "Consumido", "Estado"]
    widths   = [5, 14, 36, 18, 18, 11, 13, 11, 11, 16]

    for nodo in nodos:
        num  = _label_nodo(nodo.get("nodo")) or ""
        code = nodo.get("code", "")

        # Nombre de la pestaña: N127, N127_N126, … (max 31 chars)
        sheet_name = (f"N{num}" if num else code)[:31]
        ws = wb.create_sheet(sheet_name)
        ws.sheet_view.showGridLines = False
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w

        # ── Fila 1: header del nodo (siempre visible — frozen) ────────────────
        hdr_txt = f"Nodo {num}  —  {code}  |  {lcl_odm}"
        if pdl:
            hdr_txt += f"  |  PDL: {pdl}"
        _w(ws, 1, 1, hdr_txt, bold=True, bg=C_HDR, fg=BLANCO, size=11, span=NCOLS)
        ws.row_dimensions[1].height = 18

        # ── Fila 2: cabecera de columnas (siempre visible — frozen) ───────────
        for j, h in enumerate(COLS_HDR, 1):
            _w(ws, 2, j, h, bold=True, bg=C_SUBHDR, h="center")
        ws.row_dimensions[2].height = 14

        # Freeze: header nodo + columnas siempre al tope
        ws.freeze_panes = "A3"

        # ── Filas de datos desde fila 3 ───────────────────────────────────────
        sap_map    = _dom_a_sap(nodo)
        filas_nodo = _cruce_nodo(nodo)
        nd_tray = nd_dalm = nd_neto = nd_cons = nd_pend = 0.0
        seq  = 1
        fila = 3

        for r in filas_nodo:
            color   = r["color"]
            sap_cod = sap_map.get(r["cod"]) or r["cod"]
            bg_dev  = C_DEVOL if r["devoluciones"] > 0 else color
            _w(ws, fila,  1, seq,                               bg=color,  h="center")
            _w(ws, fila,  2, sap_cod,                           bg=color,  h="center")
            _w(ws, fila,  3, r["desc"],                         bg=color,  wrap=True)
            _w(ws, fila,  4, _fmt_ban_names(r["ban_delivers"]), bg=color,  h="left",   size=9, wrap=True)
            _w(ws, fila,  5, _fmt_ban_names(r["ban_devs"]),     bg=color,  h="left",   size=9, wrap=True)
            _w(ws, fila,  6, r["tray"],                         bg=color,  h="center", num_fmt=NF)
            _w(ws, fila,  7, r["devoluciones"],                 bg=bg_dev, h="center", num_fmt=NF,
               bold=(r["devoluciones"] > 0))
            _w(ws, fila,  8, r["desp_neto"],                    bg=color,  h="center", num_fmt=NF, bold=True)
            _w(ws, fila,  9, r["consumido"],                    bg=color,  h="center", num_fmt=NF)
            _w(ws, fila, 10, r["estado"],                       bg=color,  h="center")
            ws.row_dimensions[fila].height = 22
            nd_tray += r["tray"];      nd_dalm += r["devoluciones"]
            nd_neto += r["desp_neto"]; nd_cons += r["consumido"]
            nd_pend += r["pendiente"]
            fila += 1
            seq  += 1

        # ── Subtotal del nodo ─────────────────────────────────────────────────
        bg_sp      = C_OK if abs(nd_pend) < 0.01 else (C_AMBAR if nd_pend > 0 else C_ROJO)
        bg_dev_tot = C_DEVOL if nd_dalm > 0 else C_TOTAL
        _w(ws, fila, 1, f"Subtotal Nodo {num}", bold=True, bg=C_TOTAL, span=5)
        _w(ws, fila, 6, nd_tray,  bold=True, bg=C_TOTAL,    h="center", num_fmt=NF)
        _w(ws, fila, 7, nd_dalm,  bold=True, bg=bg_dev_tot, h="center", num_fmt=NF)
        _w(ws, fila, 8, nd_neto,  bold=True, bg=bg_sp,      h="center", num_fmt=NF)
        _w(ws, fila, 9, nd_cons,  bold=True, bg=C_TOTAL,    h="center", num_fmt=NF)
        _w(ws, fila, 10, "",      bold=True, bg=C_TOTAL)


# ── Hoja 3: CRUCE_MAT ────────────────────────────────────────────────────────

def _sap_desde_cod(cod: str) -> str:
    """'DOM-028733-12/274287' → '274287'  |  '274287' → '274287'"""
    return cod.split("/")[-1] if "/" in cod else cod


def _hoja_cruce_mat(ws, rows_mat: list, sgs_data: dict, modo_mat: str):
    NCOLS = 7

    if modo_mat == "manual_sap":
        titulo = "CRUCE — Materiales Manual vs SGS"
        hdrs   = ["#", "Nodo", "Código SAP", "Descripción", "Manual Cant.", "SGS Neto", "Estado"]
    else:
        titulo = "CRUCE — Balance SGS (Neto vs Consumido)"
        hdrs   = ["#", "Nodo", "Código", "Descripción", "SGS Neto", "Consumido", "Estado"]

    _w(ws, 1, 1, titulo, bold=True, bg=C_HDR, fg=BLANCO, size=12, span=NCOLS)
    ws.row_dimensions[1].height = 22

    leyenda = [
        (C_OK,    "OK — cantidades coinciden o listo"),
        (C_AMBAR, "PENDIENTE — sin liquidar o parcial"),
        (C_GRIS,  "SIN MOV — neto cero"),
        (C_ROJO,  "DESCUADRE — cantidades distintas"),
    ]
    for i, (color, txt) in enumerate(leyenda, 2):
        _w(ws, i, 1, "", bg=color, borders=False)
        _w(ws, i, 2, txt, bg=color, borders=False, span=NCOLS - 1)

    hdr_row = len(leyenda) + 2
    for j, h in enumerate(hdrs, 1):
        _w(ws, hdr_row, j, h, bold=True, bg=C_SUBHDR, h="center")
    ws.freeze_panes = f"A{hdr_row + 1}"

    fila = hdr_row + 1
    seq  = 1
    NF   = "0"

    if modo_mat == "manual_sap":
        # Indexar SGS por nodo → SAP → {neto, desc}
        sgs_idx: dict = defaultdict(dict)
        for nodo in sgs_data["nodos"]:
            num     = _label_nodo(nodo.get("nodo"))
            sap_map = _dom_a_sap(nodo)
            for r in _cruce_nodo(nodo):
                sap = sap_map.get(r["cod"]) or r["cod"]
                sgs_idx[num][sap] = {"neto": r["desp_neto"], "desc": r["desc"]}

        vistos: set = set()
        for r in rows_mat:
            nodo, sap, manual_q = r["nodo"], r["sap"], r["cantidad"]
            sgs_info = sgs_idx.get(nodo, {}).get(sap)
            if sgs_info:
                sgs_neto = sgs_info["neto"]
                desc     = sgs_info["desc"] or r.get("desc", "")
                diff     = manual_q - sgs_neto
                if abs(diff) < 0.001:
                    estado, color = "OK", C_OK
                else:
                    estado, color = f"DIFER {diff:+.2f}", C_ROJO
            else:
                sgs_neto = "—"
                desc     = r.get("desc", "")
                estado, color = "SOLO MANUAL", C_AMBAR

            _w(ws, fila, 1, seq,      bg=color, h="center")
            _w(ws, fila, 2, nodo,     bg=color, h="center")
            _w(ws, fila, 3, sap,      bg=color, h="center")
            _w(ws, fila, 4, desc,     bg=color, wrap=True)
            _w(ws, fila, 5, manual_q, bg=color, h="center", num_fmt=NF)
            _w(ws, fila, 6, sgs_neto, bg=color, h="center", num_fmt=NF if sgs_neto != "—" else None)
            _w(ws, fila, 7, estado,   bg=color, h="center")
            vistos.add((nodo, sap))
            fila += 1; seq += 1

        # Materiales SOLO en SGS (no ingresados manualmente)
        for nodo in sgs_data["nodos"]:
            num     = _label_nodo(nodo.get("nodo"))
            sap_map = _dom_a_sap(nodo)
            for r in _cruce_nodo(nodo):
                sap = sap_map.get(r["cod"]) or r["cod"]
                if (num, sap) in vistos or r["desp_neto"] == 0:
                    continue
                _w(ws, fila, 1, seq,             bg=C_GRIS, h="center")
                _w(ws, fila, 2, num,             bg=C_GRIS, h="center")
                _w(ws, fila, 3, sap,             bg=C_GRIS, h="center")
                _w(ws, fila, 4, r["desc"],       bg=C_GRIS, wrap=True)
                _w(ws, fila, 5, "—",             bg=C_GRIS, h="center")
                _w(ws, fila, 6, r["desp_neto"],  bg=C_GRIS, h="center", num_fmt=NF)
                _w(ws, fila, 7, "SOLO SGS",      bg=C_GRIS, h="center")
                fila += 1; seq += 1

    else:
        # Nodo solo: SGS neto vs consumido
        for nodo in sgs_data["nodos"]:
            num = _label_nodo(nodo.get("nodo"))
            for r in _cruce_nodo(nodo):
                neto = r["desp_neto"]; cons = r["consumido"]
                if neto == 0 and cons == 0:
                    color = C_GRIS
                elif cons >= neto > 0:
                    color = C_OK
                elif cons > 0:
                    color = C_AMBAR
                elif neto > 0:
                    color = C_AMBAR
                else:
                    color = C_GRIS
                sap_cod = _sap_desde_cod(r["cod"])
                _w(ws, fila, 1, seq,      bg=color, h="center")
                _w(ws, fila, 2, num,      bg=color, h="center")
                _w(ws, fila, 3, sap_cod,  bg=color, h="center")
                _w(ws, fila, 4, r["desc"], bg=color, wrap=True)
                _w(ws, fila, 5, neto,     bg=color, h="center", num_fmt=NF, bold=True)
                _w(ws, fila, 6, cons,     bg=color, h="center", num_fmt=NF)
                _w(ws, fila, 7, r["estado"], bg=color, h="center")
                fila += 1; seq += 1

    widths = [5, 8, 14, 40, 12, 12, 16]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


# ── Generador principal ────────────────────────────────────────────────────────

def generar_excel_manual(sesion: dict) -> bytes:
    wb = Workbook()

    ws_mo = wb.active;  ws_mo.title = "MO_MANUAL"

    sgs_data_base = sesion.get("sgs_data") or {"padre": sesion.get("padre", {}), "nodos": []}

    # sgs_data para materiales: solo los nodos que el usuario seleccionó explícitamente
    nodos_mat = sesion.get("nodos_mat", [])
    sgs_data_mat = {
        "padre": sgs_data_base.get("padre") or sesion.get("padre", {}),
        "nodos": nodos_mat,
        "pdl":   sesion.get("pdl", ""),
    }

    _hoja_mo(ws_mo, sesion["rows_mo"], sesion["lcl_odm"], sesion["padre"],
             sesion.get("nodos_pendientes_sgs", []),
             sesion.get("lideres", []),
             sesion.get("supervisor"),
             sesion.get("pdl", ""))
    # Una hoja por nodo de materiales seleccionados; vacío si el usuario no eligió ninguno
    _hoja_sgs_balance(wb, sgs_data_mat, sesion["lcl_odm"])
    ws_cruce = wb.create_sheet("CRUCE_MAT")
    _hoja_cruce_mat(ws_cruce, sesion["rows_mat"], sgs_data_mat,
                    sesion.get("modo_mat") or "nodo_solo")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Importar Excel anterior ────────────────────────────────────────────────────

def leer_excel_sesion(fileobj) -> dict:
    """
    Lee un Excel generado por este sistema y extrae todo lo necesario para
    reutilizarlo: MO, líderes, supervisor, LCL, código padre, nodos de materiales.

    Returns dict con claves:
      codigo_padre, lcl_odm, supervisor, texto_mo, rows_mo, lideres, texto_nodos
    """
    from openpyxl import load_workbook as _load_wb

    wb = _load_wb(fileobj, read_only=True, data_only=True)

    if "MO_MANUAL" not in wb.sheetnames:
        raise ValueError("El Excel no tiene hoja 'MO_MANUAL' — no parece ser un balance generado por este sistema.")

    ws = wb["MO_MANUAL"]
    rows = list(ws.iter_rows(values_only=True))

    def _str(row_idx, col_idx):
        try:
            v = rows[row_idx][col_idx]
            return str(v).strip() if v is not None else ""
        except (IndexError, TypeError):
            return ""

    # Meta por etiqueta (col A) — tolerante a Excels viejos (sin fila PDL)
    # y nuevos (con ella): no depende de posiciones fijas.
    meta_map = {}
    for i in range(1, min(10, len(rows))):
        lbl = _str(i, 0).rstrip(":").strip().lower()
        if lbl:
            meta_map[lbl] = _str(i, 1)
    codigo_padre = meta_map.get("código padre", "") or _str(2, 1)
    lcl_odm      = meta_map.get("odm / lcl", "")   or _str(3, 1)
    supervisor   = meta_map.get("supervisor", "")  or _str(4, 1)
    pdl          = meta_map.get("pdl", "")
    if pdl in ("—", "-"):
        pdl = ""

    # Filas de MO: arrancan 2 filas después del separador "MANO DE OBRA"
    # (posición variable según cuántas filas tenga la cabecera)
    mo_start = 8
    for i, row in enumerate(rows[:15]):
        a = str(row[0]).strip().upper() if row and row[0] is not None else ""
        if a == "MANO DE OBRA":
            mo_start = i + 2
            break

    # Col A=idx (numérico), B=nodo, C=qp, G=cantidad
    rows_raw = []
    in_lideres = False
    lideres_raw = []
    ldr_header_passed = False

    for row in rows[mo_start:]:
        a = row[0]
        a_str = str(a).strip() if a is not None else ""

        # Detectar fin de MO / inicio de líderes
        if "TOTAL" in a_str.upper():
            continue
        if "DISTRIBUCI" in a_str.upper():
            in_lideres = True
            ldr_header_passed = False
            continue
        if in_lideres:
            # Saltar fila de cabecera (#, Cédula, Nombre, %)
            if not ldr_header_passed:
                ldr_header_passed = True
                continue
            # Fila de líder: col A=idx (int), B=cedula, C=nombre, D="XX%"
            try:
                idx_l = int(float(str(a))) if a is not None else None
            except (ValueError, TypeError):
                idx_l = None
            if idx_l is not None:
                cedula = str(row[1]).strip() if row[1] is not None else ""
                nombre = str(row[2]).strip() if row[2] is not None else ""
                pct_raw = str(row[3]).strip() if row[3] is not None else "0"
                pct = pct_raw.replace("%", "").strip()
                if nombre:
                    lideres_raw.append({"cedula": cedula, "nombre": nombre, "pct": pct})
            continue

        # Fila de MO: col A debe ser número entero (secuencia)
        if a is None:
            continue
        try:
            int(float(str(a)))
        except (ValueError, TypeError):
            continue

        nodo    = str(row[1]).strip() if row[1] is not None else ""
        qp      = str(row[2]).strip() if row[2] is not None else ""
        try:
            cantidad = int(float(str(row[6]))) if row[6] is not None else 0
        except (ValueError, TypeError):
            cantidad = 0
        if nodo and qp and cantidad > 0:
            rows_raw.append({"nodo": nodo, "qp": qp, "cantidad": cantidad})

    texto_mo = "\n".join(f"{r['nodo']}\t{r['qp']}\t{r['cantidad']}" for r in rows_raw)
    rows_mo  = enriquecer_mo(rows_raw)

    # Nodos de materiales: leer PRO code de la fila 1 de cada hoja de nodo
    # Formato: "Nodo 127  —  PRO-056358  |  6300995703  |  PDL: ..."
    _pro_re = re.compile(r'PRO-\d+', re.IGNORECASE)
    pros_importados = []
    for sname in wb.sheetnames:
        if sname in ("MO_MANUAL", "CRUCE_MAT"):
            continue
        try:
            ws_n = wb[sname]
            row1 = next(ws_n.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if row1 and row1[0]:
                m = _pro_re.search(str(row1[0]))
                if m:
                    code = m.group(0).upper()
                    if code not in pros_importados:
                        pros_importados.append(code)
        except Exception:
            pass

    return {
        "codigo_padre":   codigo_padre,
        "lcl_odm":        lcl_odm,
        "supervisor":     supervisor,
        "pdl":            pdl,
        "texto_mo":       texto_mo,
        "rows_mo":        rows_mo,
        "lideres":        lideres_raw,
        "texto_nodos":    "",
        "pros_importados": pros_importados,
    }
