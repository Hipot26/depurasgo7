"""sgo7_casos.py — Línea de facturación SGO7 desde el Balance Integrado (módulo Flask `nuevo`).

Cada consulta de Balance Integrado queda guardada en sabogal como un «caso» (planilla SIPWEBORA +
proyecto SGS) con su semáforo: cuántas filas cuadran, cuántas descuadran y qué alertas hay. El caso
tiene estado (descuadre, en espera de almacén, cuadra, enviado a liquidar…), se reconsulta solo
desde el cron (`reconsultar_casos.py`) y, cuando cuadra y se genera el Excel, el facturador lo manda
a liquidar con un botón: la orden aparece en el panel del Liquidador Automático.

Habla con el backend por el canal interno (token SGO7_INTERNAL_TOKEN, nginx lo niega desde fuera).
Nunca rompe el Balance: si el backend no responde, devuelve None y el flujo sigue igual que antes.
"""
import hashlib
import json
import os
import re
import urllib.error
import urllib.request
from datetime import datetime, timedelta, timezone

API = os.environ.get("SGO7_API_URL", "http://127.0.0.1:3008/api/sgo7")
TOKEN = os.environ.get("SGO7_INTERNAL_TOKEN", "")
CASOS_DIR = os.environ.get("SGO7_CASOS_DIR", "/var/www/sgo7/registros/casos")
TZ_CO = timezone(timedelta(hours=-5))

ESTADOS = {
    "consultado": "Consultado", "descuadre": "Con descuadre", "espera_almacen": "En espera de almacén",
    "cuadra": "Cuadra", "enviado": "Enviado a liquidar", "en_liquidacion": "En liquidación",
    "liquidado": "Liquidado", "con_fallos": "Liquidado con fallos", "archivado": "Archivado",
}
ORDEN_ESTADOS = {"pendiente": "Pendiente en el Liquidador", "en_liquidacion": "En liquidación",
                 "liquidada": "Liquidada", "con_fallos": "Liquidada con fallos", "cancelada": "Cancelada"}
MENSAJES = {
    "descuadre": "El balance no cuadra. Si ya lo reportaste a almacén, déjalo en espera: el sistema lo reconsulta solo cada 2 horas y te avisa en el panel cuando cuadre.",
    "espera_almacen": "En espera de almacén. Se reconsulta solo cada 2 horas; cuando cuadre pasará a «Cuadra» en el panel.",
    "cuadra": "El balance cuadra. Genera el Excel (fecha + líderes) y mándalo a liquidar.",
    "consultado": "Caso registrado.",
    "enviado": "Ya fue enviado a liquidar. El Liquidador lo tiene en su panel.",
    "en_liquidacion": "El Liquidador está trabajando este caso.",
    "liquidado": "Este caso ya fue liquidado en SGS.",
    "con_fallos": "La liquidación terminó con fallos; revisa el panel del Liquidador.",
    "archivado": "Caso archivado. Una consulta nueva lo reabre.",
}


class CasoError(Exception):
    pass


def _api(path, payload=None, method=None, timeout=12):
    data = json.dumps(payload, ensure_ascii=False, default=str).encode("utf-8") if payload is not None else None
    req = urllib.request.Request(API + path, data=data, method=method or ("POST" if data is not None else "GET"),
                                 headers={"Content-Type": "application/json", "X-SGO7-Interno": TOKEN})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            cuerpo = r.read().decode("utf-8")
            return json.loads(cuerpo) if cuerpo else None
    except urllib.error.HTTPError as e:
        try:
            det = json.loads(e.read().decode("utf-8")).get("detail")
        except Exception:
            det = None
        raise CasoError(det or f"HTTP {e.code}")


def _usuario(usuario=None, rol=None):
    if usuario:
        return usuario, rol
    try:
        from flask import has_request_context, request
        if has_request_context():
            return (request.headers.get("X-SGO7-User") or None), (request.headers.get("X-SGO7-Rol") or None)
    except Exception:
        pass
    return None, None


def _co(iso):
    """ISO (UTC) → «10/10 14:32» hora de Colombia."""
    if not iso:
        return ""
    try:
        d = datetime.fromisoformat(str(iso).replace("Z", "+00:00"))
        if d.tzinfo is None:
            d = d.replace(tzinfo=timezone.utc)
        return d.astimezone(TZ_CO).strftime("%d/%m %H:%M")
    except Exception:
        return str(iso)[:16]


def enriquecer(caso):
    """Añade etiquetas y fechas en hora CO para las plantillas. Devuelve el mismo dict."""
    if not isinstance(caso, dict):
        return caso
    res = caso.get("resumen") or {}
    caso["estado_label"] = ESTADOS.get(caso.get("estado"), caso.get("estado"))
    caso["orden_estado_label"] = ORDEN_ESTADOS.get(caso.get("orden_estado") or "", caso.get("orden_estado"))
    caso["mensaje"] = MENSAJES.get(caso.get("estado"), "")
    caso["rojos"] = int(res.get("rojos") or 0)
    caso["ok"] = int(res.get("ok") or 0)
    caso["n_alertas"] = len(res.get("alertas") or []) + len(res.get("advertencias") or [])
    caso["ultima_co"] = _co(caso.get("ultima_consulta_en"))
    caso["orden_enviada_co"] = _co(caso.get("orden_enviada_en"))
    caso["excel_co"] = _co(caso.get("excel_generado_en"))
    return caso


# ── Semáforo: se lee de la hoja CRUCE_FACTURAS del Excel base (la misma que ve el facturador) ──
def resumen_cruce(xlsx_bytes) -> dict:
    import io
    import openpyxl
    out = {"ok": 0, "descuadre": 0, "error_clasif": 0, "solo": 0, "chatarra": 0, "compra_c": 0,
           "filas": 0, "filas_rojas": [], "alertas": [], "resumen_txt": ""}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    except Exception as e:
        out["error"] = f"no se pudo leer el Excel: {e}"
        return out
    ws = None
    for nombre in ("CRUCE_FACTURAS", "CRUCE_SIP_SGS"):
        if nombre in wb.sheetnames:
            ws = wb[nombre]
            break
    if ws is None:
        out["error"] = "sin hoja CRUCE_FACTURAS"
        return out
    cont = {"ok": 0, "descuadre": 0, "error_clasif": 0, "solo": 0, "chatarra": 0}
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        a = row[0]
        if isinstance(a, (int, float)) and len(row) >= 10 and row[9]:
            est = str(row[9]).strip()
            u = est.upper()
            if u.startswith("OK"):
                clase = "ok"
            elif u.startswith("DIFER") or u.startswith("SOLO SIP") or u.startswith("SOLO SGS (DESPACHADO"):
                clase = "descuadre"
            elif u.startswith("! ERROR CLASIFICACION") or u.startswith("ERROR CLASIFICACION"):
                clase = "error_clasif"
            elif u.startswith("SOLO SGS"):
                clase = "solo"
            elif u.startswith("CHATARRA"):
                clase = "chatarra"
            else:
                clase = "solo"
            cont[clase] += 1
            out["filas"] += 1
            if clase in ("descuadre", "error_clasif") and len(out["filas_rojas"]) < 80:
                out["filas_rojas"].append({
                    "nodo": str(row[1] or ""), "codigo": str(row[2] or ""), "desc": str(row[3] or "")[:70],
                    "sip": row[4] if isinstance(row[4], (int, float)) else (None if row[4] in (None, "—") else str(row[4])),
                    "sgs": row[5] if isinstance(row[5], (int, float)) else (None if row[5] in (None, "—") else str(row[5])),
                    "dif": row[6] if isinstance(row[6], (int, float)) else None, "estado": est[:60]})
        elif isinstance(a, str):
            s = a.strip()
            if s.startswith("RESUMEN"):
                out["resumen_txt"] = s[:300]
                for clave, rx in (("error_clasif", r"Errores clasificacion SIP:\s*(\d+)"), ("chatarra", r"Chatarra confirmada:\s*(\d+)"),
                                  ("compra_c", r"Compra Dominion \(C\):\s*(\d+)"), ("ok", r"OK:\s*(\d+)"),
                                  ("descuadre", r"Descuadre:\s*(\d+)"), ("solo", r"Solo un sistema:\s*(\d+)")):
                    m = re.search(rx, s)
                    if m:
                        out[clave] = int(m.group(1))
            elif s.startswith("CRITICO") or s.startswith("REVISAR") or s.startswith("EQUIVALENCIA"):
                if len(out["alertas"]) < 40:
                    out["alertas"].append(s[:220])
    # Si la fila RESUMEN no apareció, valen los conteos por fila
    if not out["resumen_txt"]:
        out.update(cont)
    out["rojos"] = int(out["descuadre"]) + int(out["error_clasif"])
    return out


# ── Operaciones ──────────────────────────────────────────────────────────────────────
def registrar_consulta(datos, base_xlsx, sid=None, usuario=None, rol=None, origen="manual", proyecto_id=None):
    """Tras preparar_cruce: manda el semáforo al backend y devuelve el caso (dict enriquecido) o None."""
    u, r = _usuario(usuario, rol)
    meta = datos.get("metadata") or {}
    padre = datos.get("sgs_padre") or {}
    res = resumen_cruce(base_xlsx)
    res["nodos"] = len(datos.get("nodos_resumen") or [])
    res["total_mo"] = meta.get("total_mo")
    res["total_mat"] = meta.get("total_mat")
    try:
        pid = int(datos.get("parent_id")) if datos.get("parent_id") else None
    except Exception:
        pid = None
    payload = {
        "usuario": u, "rol": r, "origen": origen, "sid": sid,
        "id_planilla": datos.get("id_label") or datos.get("id_planilla") or "",
        "proyecto_id": proyecto_id if proyecto_id is not None else (datos.get("proyecto_id") or None),
        "parent_id": pid,
        "pro_code": padre.get("code"), "proyecto_nombre": padre.get("name"),
        "lcl": meta.get("lcl"), "odm": meta.get("odm"), "ticket": meta.get("ticket"),
        "zona": meta.get("zona_nombre") or meta.get("zona_cd"), "empresa": meta.get("empresa"),
        "resumen": res, "advertencias": datos.get("advertencias") or [],
    }
    try:
        return enriquecer(_api("/interno/casos/consulta", payload))
    except Exception as e:
        print("sgo7_casos: consulta no registrada:", e)
        return None


def guardar_excel_caso(caso, bal_bytes, bal_name, rg11_bytes=None, rg11_name=None, fecha=None, lideres=None,
                       usuario=None, rol=None, sid=None):
    """Copia los Excel generados a registros/casos/<id>/ y los ata al caso. Devuelve el caso o None."""
    if not caso or not caso.get("id"):
        return None
    u, r = _usuario(usuario, rol)
    cid = int(caso["id"])
    carpeta = os.path.join(CASOS_DIR, str(cid))
    os.makedirs(carpeta, exist_ok=True)
    ts = datetime.now(TZ_CO).strftime("%Y%m%d-%H%M%S")

    def _guardar(data, nombre):
        base = re.sub(r"[^A-Za-z0-9._-]+", "_", nombre or "archivo.xlsx")[:140]
        ruta = os.path.join(carpeta, f"{ts}_{base}")
        with open(ruta, "wb") as f:
            f.write(data)
        os.chmod(ruta, 0o640)
        return ruta

    ruta_bal = _guardar(bal_bytes, bal_name)
    ruta_rg11 = _guardar(rg11_bytes, rg11_name) if rg11_bytes else None
    payload = {"usuario": u, "rol": r, "sid": sid, "excel_archivo": bal_name, "excel_ruta": ruta_bal,
               "excel_sha256": hashlib.sha256(bal_bytes).hexdigest(), "excel_bytes": len(bal_bytes),
               "fecha_ejecucion": fecha or None, "lideres": lideres or [], "rg11_archivo": rg11_name, "rg11_ruta": ruta_rg11}
    try:
        return enriquecer(_api(f"/interno/casos/{cid}/excel", payload))
    except Exception as e:
        print("sgo7_casos: excel no atado al caso:", e)
        return None


def mandar_a_liquidar(caso_id, nota="", usuario=None, rol=None):
    """Crea la orden para el Liquidador. Devuelve (caso, None) o (None, mensaje de error)."""
    u, r = _usuario(usuario, rol)
    try:
        return enriquecer(_api(f"/interno/casos/{int(caso_id)}/orden", {"usuario": u, "rol": r, "nota": nota or ""})), None
    except CasoError as e:
        return None, str(e)
    except Exception as e:
        return None, f"El panel no respondió: {e}"


def cambiar_estado(caso_id, estado=None, nota=None, reconsulta_auto=None, usuario=None, rol=None):
    u, r = _usuario(usuario, rol)
    payload = {"usuario": u, "rol": r}
    if estado:
        payload["estado"] = estado
    if nota is not None:
        payload["nota"] = nota
    if reconsulta_auto is not None:
        payload["reconsulta_auto"] = bool(reconsulta_auto)
    try:
        return enriquecer(_api(f"/interno/casos/{int(caso_id)}/estado", payload)), None
    except CasoError as e:
        return None, str(e)
    except Exception as e:
        return None, f"El panel no respondió: {e}"


def pendientes_reconsulta(limite=8):
    try:
        return _api(f"/interno/casos/reconsultar?limite={int(limite)}") or []
    except Exception as e:
        print("sgo7_casos: no se pudo leer la cola de reconsulta:", e)
        return []


def reconsulta_error(caso_id, error):
    try:
        _api(f"/interno/casos/{int(caso_id)}/reconsulta-error", {"error": str(error)[:400]})
    except Exception as e:
        print("sgo7_casos: error de reconsulta no registrado:", e)
