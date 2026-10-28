#!/usr/bin/env python3
"""
Módulo Formato RG02/RG03 — parser y escritor del RG03 (.xlsm).

La hoja "FOTOS  " del RG03 está organizada en bloques de 11 filas a partir
de la fila 13. Por bloque N (fila base r = 13 + 11*(N-1)):

    H(r)   = Foto N°          L(r)   = Nodo N°
    K(r+1) = CD / DME         K(r+2) = Punto Físico
    G(r+5) = Latitud          L(r+5) = Longitud
    G(r+7) = OBSERVACIONES  ← celda que este módulo rellena

Las fotos NO son accesibles vía openpyxl (ws._images = 0): están ancladas en
xl/drawings/drawingN.xml. Se mapean foto↔bloque leyendo el XML del zip.

La escritura se hace por "cirugía de zip": se reescribe únicamente el XML de
la hoja FOTOS (celdas de observación como inlineStr) y el resto del archivo
se copia byte a byte — el proyecto VBA, los drawings y las imágenes quedan
intactos. NUNCA re-guardar el xlsm con openpyxl: destruye shapes y macros.
"""
import io
import re
import zipfile

from lxml import etree

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL_DOC = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

FILA_BASE = 13      # primera fila de bloque
FILAS_BLOQUE = 11   # alto de cada bloque
COL_OBS = "G"       # columna de la celda OBSERVACIONES (fila base + 7)


# ─────────────────────────────────────────────────────────────────────────────
# Localización de la hoja FOTOS dentro del paquete
# ─────────────────────────────────────────────────────────────────────────────

def _localizar_hoja_fotos(z: zipfile.ZipFile):
    """Devuelve (nombre_hoja, ruta_sheet_xml) de la hoja cuyo nombre sea FOTOS."""
    wb_xml = etree.fromstring(z.read("xl/workbook.xml"))
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rid_a_target = {
        rel.get("Id"): rel.get("Target")
        for rel in rels.findall(f"{{{NS_PKG_REL}}}Relationship")
    }
    for sheet in wb_xml.findall(f"{{{NS_MAIN}}}sheets/{{{NS_MAIN}}}sheet"):
        nombre = (sheet.get("name") or "").strip().upper()
        if nombre == "FOTOS":
            rid = sheet.get(f"{{{NS_REL_DOC}}}id")
            target = rid_a_target.get(rid, "")
            if target.startswith("/"):
                target = target.lstrip("/")
            elif not target.startswith("xl/"):
                target = "xl/" + target
            return sheet.get("name"), target
    raise ValueError("El archivo no tiene una hoja 'FOTOS' — ¿es un RG03 válido?")


def _mapear_fotos(z: zipfile.ZipFile, sheet_path: str) -> dict:
    """Mapea nº de bloque → ruta xl/media/... leyendo el drawing de la hoja."""
    base = sheet_path.rsplit("/", 1)[-1]
    rels_path = f"xl/worksheets/_rels/{base}.rels"
    if rels_path not in z.namelist():
        return {}
    rels = etree.fromstring(z.read(rels_path))
    drawing_path = None
    for rel in rels.findall(f"{{{NS_PKG_REL}}}Relationship"):
        if rel.get("Type", "").endswith("/drawing"):
            drawing_path = "xl/" + rel.get("Target").replace("../", "")
            break
    if not drawing_path or drawing_path not in z.namelist():
        return {}

    d_rels_path = "xl/drawings/_rels/" + drawing_path.rsplit("/", 1)[-1] + ".rels"
    rid_a_media = {}
    if d_rels_path in z.namelist():
        d_rels = etree.fromstring(z.read(d_rels_path))
        for rel in d_rels.findall(f"{{{NS_PKG_REL}}}Relationship"):
            rid_a_media[rel.get("Id")] = "xl/" + rel.get("Target").replace("../", "")

    bloque_a_media = {}
    root = etree.fromstring(z.read(drawing_path))
    anclas = root.findall(f"{{{NS_XDR}}}twoCellAnchor") + root.findall(f"{{{NS_XDR}}}oneCellAnchor")
    for anc in anclas:
        frm = anc.find(f"{{{NS_XDR}}}from")
        blip = anc.find(f".//{{{NS_A}}}blip")
        if frm is None or blip is None:
            continue
        fila = int(frm.findtext(f"{{{NS_XDR}}}row", "0")) + 1  # 1-based
        rid = blip.get(f"{{{NS_REL_DOC}}}embed")
        media = rid_a_media.get(rid)
        if not media:
            continue
        n = round((fila - FILA_BASE) / FILAS_BLOQUE) + 1
        if n < 1:      # logos de cabecera y adornos fuera de los bloques
            continue
        bloque_a_media.setdefault(n, media)
    return bloque_a_media


# ─────────────────────────────────────────────────────────────────────────────
# Parser
# ─────────────────────────────────────────────────────────────────────────────

def parsear_rg03(xlsm_bytes: bytes) -> dict:
    """Lee el RG03 y devuelve encabezado + bloques con foto.

    Retorna::

        {
          "hoja": "FOTOS  ",
          "header": {...},
          "bloques": [ {n, nodo, cd, pf, lat, lon, obs_actual, media, jpg}, ... ],
          "advertencias": [str, ...],
        }
    """
    import openpyxl

    advertencias = []
    with zipfile.ZipFile(io.BytesIO(xlsm_bytes)) as z:
        nombre_hoja, sheet_path = _localizar_hoja_fotos(z)
        bloque_a_media = _mapear_fotos(z, sheet_path)
        media_bytes = {m: z.read(m) for m in set(bloque_a_media.values())}

    wb = openpyxl.load_workbook(io.BytesIO(xlsm_bytes), data_only=True)
    ws = wb[nombre_hoja]

    def _celda(fila, col):
        v = ws.cell(row=fila, column=col).value
        return str(v).strip() if v is not None else ""

    header = {
        "codigo": _celda(1, 10),                # J1 — código RG del formato
        "empresa": _celda(5, 1),                # A5
        "contrato": _celda(5, 3),               # C5
        "responsable": _celda(5, 6),            # F5
        "fecha": "/".join(x for x in (_celda(5, 12), _celda(5, 15), _celda(5, 16)) if x),
        "lcl": _celda(7, 1),                    # A7
        "circuito": _celda(7, 5),               # E7
        "cd": _celda(7, 8),                     # H7
        "ticket": _celda(9, 1),                 # A9
        "zona": _celda(9, 8),                   # H9
        "municipio": _celda(11, 1),             # A11
        "direccion": _celda(11, 5),             # E11
        "vereda": _celda(11, 8),                # H11
    }

    bloques = []
    r = FILA_BASE
    while r <= ws.max_row:
        foto_n = ws.cell(row=r, column=8).value
        if foto_n is None:
            r += FILAS_BLOQUE
            continue
        try:
            n = int(foto_n)
        except (TypeError, ValueError):
            r += FILAS_BLOQUE
            continue
        media = bloque_a_media.get(n)
        if media:  # solo interesan los bloques que tienen foto real
            nodo_raw = _celda(r, 12)
            nodo = re.sub(r"(?i)^\s*nodo\s*", "", nodo_raw).strip()
            bloques.append({
                "n": n,
                "fila": r,
                "nodo": nodo,
                "nodo_raw": nodo_raw,
                "cd": _celda(r + 1, 11),
                "pf": _celda(r + 2, 11),
                "lat": _celda(r + 5, 7),
                "lon": _celda(r + 5, 12),
                "obs_actual": _celda(r + 7, 7),
                "media": media,
                "jpg": media_bytes.get(media, b""),
            })
        r += FILAS_BLOQUE

    if not bloques:
        advertencias.append("No se detectó ninguna foto embebida en la hoja FOTOS.")

    nodos = {b["nodo"] for b in bloques if b["nodo"]}
    if len(nodos) > 1:
        advertencias.append(f"Las fotos referencian varios nodos: {sorted(nodos)} — verifica que sea correcto.")

    return {"hoja": nombre_hoja, "header": header, "bloques": bloques, "advertencias": advertencias}


# ─────────────────────────────────────────────────────────────────────────────
# Writer (cirugía de zip: solo se reescribe el XML de la hoja FOTOS)
# ─────────────────────────────────────────────────────────────────────────────

def _col_a_indice(ref: str) -> int:
    idx = 0
    for ch in re.match(r"[A-Z]+", ref).group(0):
        idx = idx * 26 + (ord(ch) - 64)
    return idx


def _poner_texto(sheet_root, fila: int, col: str, texto: str):
    """Deja la celda {col}{fila} como inlineStr con `texto`, creándola si falta."""
    ref = f"{col}{fila}"
    sheet_data = sheet_root.find(f"{{{NS_MAIN}}}sheetData")
    row_el = None
    for row in sheet_data.findall(f"{{{NS_MAIN}}}row"):
        if row.get("r") == str(fila):
            row_el = row
            break
    if row_el is None:
        row_el = etree.Element(f"{{{NS_MAIN}}}row")
        row_el.set("r", str(fila))
        siguiente = None
        for rw in sheet_data.findall(f"{{{NS_MAIN}}}row"):
            if int(rw.get("r", "0")) > fila:
                siguiente = rw
                break
        if siguiente is not None:
            siguiente.addprevious(row_el)
        else:
            sheet_data.append(row_el)

    cell = None
    estilo_prestado = None
    for c in row_el.findall(f"{{{NS_MAIN}}}c"):
        if c.get("r") == ref:
            cell = c
            break
    if cell is None:
        cell = etree.Element(f"{{{NS_MAIN}}}c")
        cell.set("r", ref)
        objetivo = _col_a_indice(ref)
        insertado = False
        for i, c in enumerate(row_el.findall(f"{{{NS_MAIN}}}c")):
            if _col_a_indice(c.get("r", "A")) > objetivo:
                c.addprevious(cell)
                insertado = True
                break
        if not insertado:
            row_el.append(cell)
        # heredar estilo de la misma celda de otro bloque (mismo formato visual)
        if estilo_prestado is None:
            for row2 in sheet_data.findall(f"{{{NS_MAIN}}}row"):
                for c2 in row2.findall(f"{{{NS_MAIN}}}c"):
                    r2 = c2.get("r", "")
                    if r2.startswith(col) and r2 != ref and c2.get("s"):
                        try:
                            f2 = int(r2[len(col):])
                        except ValueError:
                            continue
                        if (f2 - FILA_BASE - 7) % FILAS_BLOQUE == 0:
                            estilo_prestado = c2.get("s")
                            break
                if estilo_prestado:
                    break
        if estilo_prestado:
            cell.set("s", estilo_prestado)

    # limpiar contenido previo y dejar inlineStr
    for child in list(cell):
        cell.remove(child)
    cell.set("t", "inlineStr")
    is_el = etree.SubElement(cell, f"{{{NS_MAIN}}}is")
    t_el = etree.SubElement(is_el, f"{{{NS_MAIN}}}t")
    t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t_el.text = texto


def escribir_observaciones(xlsm_bytes: bytes, textos: dict) -> bytes:
    """Escribe `textos` ({nº de foto: observación}) en la hoja FOTOS.

    Devuelve el .xlsm completo con TODO lo demás copiado byte a byte.
    """
    src = zipfile.ZipFile(io.BytesIO(xlsm_bytes))
    _, sheet_path = _localizar_hoja_fotos(src)

    sheet_root = etree.fromstring(src.read(sheet_path))
    for n, texto in sorted(textos.items()):
        if not str(texto or "").strip():
            continue
        fila = FILA_BASE + FILAS_BLOQUE * (int(n) - 1) + 7
        _poner_texto(sheet_root, fila, COL_OBS, str(texto).strip())

    sheet_xml = etree.tostring(sheet_root, xml_declaration=True,
                               encoding="UTF-8", standalone=True)

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = sheet_xml if item.filename == sheet_path else src.read(item.filename)
            dst.writestr(item, data)
    src.close()
    return out.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# Prueba manual:  python rg03_fotos.py archivo.xlsm
# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    import sys

    with open(sys.argv[1], "rb") as f:
        data = f.read()
    res = parsear_rg03(data)
    print("Hoja:", repr(res["hoja"]))
    print("Header:", res["header"])
    for b in res["bloques"]:
        print(f"  Foto {b['n']:>2} | nodo {b['nodo']:<4} | PF {b['pf']:<12} | "
              f"{len(b['jpg'])//1024:>4} KB | obs: {b['obs_actual'][:60]!r}")
    for a in res["advertencias"]:
        print("  ⚠", a)

    # round-trip de prueba: escribir textos dummy y re-leer
    textos = {b["n"]: f"[PRUEBA] Observación generada para la foto {b['n']}." for b in res["bloques"]}
    nuevo = escribir_observaciones(data, textos)
    res2 = parsear_rg03(nuevo)
    ok = all(b["obs_actual"].startswith("[PRUEBA]") for b in res2["bloques"])
    print(f"\nRound-trip: {'OK' if ok else 'FALLO'} — {len(nuevo)//1024} KB "
          f"(original {len(data)//1024} KB)")
    if len(sys.argv) > 2:
        with open(sys.argv[2], "wb") as f:
            f.write(nuevo)
        print("Escrito:", sys.argv[2])
