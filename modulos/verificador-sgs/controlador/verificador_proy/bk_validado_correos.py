"""
validado_correos.py — Exporte validado por correos (Control Mensual de Liquidación, 2026-10-22)
=============================================================================================
Origen: material despachado y nunca liquidado en un LCL. Objetivo: mirar los presupuestos enviados por
correo, mirar los nodos, ver con ese dato qué no se liquidó y justificar con el presupuesto que indica
liquidar. El exporte, inspirado en «Plantilla materiales», entrega lo que sí se mandó a liquidar y no se
ha liquidado, por mes.

Fuentes (todo lectura):
  · Centro MO (`GET 127.0.0.1:5088/api/items`): cada correo con su código (LCL/DF), su
    presupuesto de materiales (`filas_mat`: nodo, código SAP, cantidad) y lo que el correo dice
    del material (`instruccion_mat`, 2026-10-22).
  · SGS (cuenta de servicio, solo lectura): TODOS los proyectos padre del LCL (SGS puede
    partir un LCL en varios padres) y sus nodos; por material: neto despachado (bandejas
    aprobadas), consumido y pendiente — la misma cuenta que `sgs_checker` («Plantilla
    materiales»).

Regla por fila (correo × nodo × material):
  · pendiente en SGS > 0 y el presupuesto lo trae  → «MANDADO Y NO LIQUIDADO».
        objetivo = min(neto despachado, cantidad del presupuesto)
        a liquidar = objetivo − consumido (si > 0).  La hoja CRUCE_LIQUIDADOR lleva el OBJETIVO en
        «SGS Neto»: Liquidador 2 liquida «SGS Neto − consumido en vivo», así que nunca pasa del
        presupuesto ni de lo despachado, y si algo se liquidó después del exporte se corrige solo.
  · correo «solo MO / no liquidar material» (no_mat) → OMITIDO.
  · correo «a excepción de…» (parcial)                → REVISAR (sin marcar por defecto).
  · nodo del presupuesto sin nodo SGS del mismo nombre → se busca por materiales iguales
    (mismo código y misma cantidad despachada); si no aparece, «SIN NODO EN SGS».
  · el material no está proyectado en el nodo SGS      → «NO ESTÁ EN SGS».
  · proyectado pero sin despacho aprobado              → «SIN DESPACHO DE ALMACÉN».
  · ya consumido lo presupuestado                      → «LIQUIDADO» (solo se cuenta).
Regla: sin instrucción en contra, la planilla estaba habilitada para liquidar MAT.

El mes es el del correo (cuando se mandó a liquidar). «todos» recorre los 140+ correos.
"""

import io
import json
import os
import re
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import datetime

import openpyxl
from openpyxl.utils import get_column_letter

from verificador_proy import sgs_checker as sc
from verificador_proy.export_auto import (
    _set, _num, _just_bandejas, C_HDR_BG, C_HDR_FG, C_PADRE_BG, C_WARN, C_MUTED,
)
from verificador_proy.export_plantilla_mat import HOJA_LIQUIDADOR, _hoja_cruce, _autoancho

CENTRO_MO_ITEMS = os.environ.get("CENTRO_MO_ITEMS", "http://127.0.0.1:5088/api/items")
CENTRO_MO_BASE = os.environ.get("CENTRO_MO_BASE", "http://127.0.0.1:5088")
BACKEND_BASE = os.environ.get("SGO7_BACKEND", "http://127.0.0.1:3008")
MAX_PADRES = 6
_RE_ID = re.compile(r"^[0-9a-f]{6,16}$")

INSTR_TXT = {
    "no_mat": "Solo MO: no liquidar material",
    "parcial": "Material con excepciones",
    "pendiente_almacen": "Material pendiente de almacén",
    "con_mat": "MO + material",
    "solo_menciona_mo": "Habla de MO; no excluye material",
    "sin_mencion": "No menciona el material",
    "": "Instrucción no leída",
}

ESTADOS = {
    "LIQUIDAR":    "Mandado y no liquidado",
    "REVISAR":     "Revisar: el correo trae excepciones",
    "OMITIDO":     "Omitido: el correo dice solo MO",
    "SOBRANTE":    "Despachado de más; lo presupuestado ya está liquidado",
    "LIQUIDADO":   "Liquidado",
    "SIN_DESPACHO": "Sin despacho de almacén",
    "NO_EN_SGS":   "No está proyectado en el nodo SGS",
    "SIN_NODO":    "Sin nodo en SGS",
}


# ── Centro MO ────────────────────────────────────────────────────────────────

def _items_centro_mo() -> list:
    with urllib.request.urlopen(CENTRO_MO_ITEMS, timeout=60) as r:
        return json.loads(r.read().decode("utf-8")).get("items") or []


# ── «Ver correo»: cuerpo guardado por Gestión General + abrir en el Outlook del PC ──────

def _http(url: str, method: str = "GET", headers: dict | None = None, timeout: int = 30):
    """(json, código HTTP). Un error de red devuelve 502 con el motivo."""
    import urllib.error
    req = urllib.request.Request(url, method=method, headers=headers or {}, data=b"{}" if method == "POST" else None)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read().decode("utf-8") or "{}"), r.status
    except urllib.error.HTTPError as e:
        try:
            return json.loads(e.read().decode("utf-8") or "{}"), e.code
        except Exception:
            return {"error": f"HTTP {e.code}"}, e.code
    except Exception as e:
        return {"error": str(e)[:200]}, 502


def correo_de(rid: str, bearer: str = "") -> tuple[dict, int]:
    """El correo de un registro de Centro MO: cabecera de Centro MO + cuerpo (HTML saneado) de Gestión
    General por EntryID. Se pide por el canal interno del backend (token interno), así lo ve quien entra
    al Control Mensual aunque no sea desarrollador. Solo correos que el Centro MO tiene registrados."""
    if not _RE_ID.match(rid or ""):
        return {"ok": False, "error": "id inválido"}, 400
    try:
        reg = next((r for r in _items_centro_mo() if r.get("id") == rid), None)
    except Exception as e:
        return {"ok": False, "error": f"No pude leer Centro MO: {e}"}, 502
    if not reg:
        return {"ok": False, "error": "Ese correo no está en Centro MO"}, 404
    im = reg.get("instruccion_mat") or {}
    out = {"ok": True, "id": rid, "asunto": reg.get("asunto", ""), "remitente": reg.get("remitente", ""),
           "remitente_email": reg.get("remitente_email", ""), "fecha": reg.get("fecha", ""),
           "carpeta": reg.get("carpeta", ""), "codigo": reg.get("codigo", ""), "adjunto": reg.get("adjunto", ""),
           "instruccion_mat": {k: im.get(k) for k in ("valor", "frase", "fuente")} if im else {},
           "tiene_entry": bool(reg.get("entry_id")), "encontrado": False,
           "de": "", "de_correo": "", "para": "", "html": "", "texto": "", "adjuntos": [], "aviso": ""}
    if not reg.get("entry_id"):
        out["aviso"] = "Centro MO no guardó la referencia al correo de Outlook"
        return out, 200
    tok = os.getenv("SGO7_INTERNAL_TOKEN", "")
    if not tok and not bearer:
        out["aviso"] = "sin token interno para pedir el cuerpo a Gestión General"
        return out, 200
    q = urllib.parse.quote(reg["entry_id"], safe="")
    if tok:
        d, code = _http(f"{BACKEND_BASE}/api/sgo7/interno/correo?entry_id={q}", headers={"X-SGO7-Interno": tok})
    else:
        d, code = _http(f"{BACKEND_BASE}/api/sgo7/correo?entry_id={q}", headers={"Authorization": f"Bearer {bearer}"})
    if code == 200:
        out.update(encontrado=True, html=d.get("html") or "", texto=d.get("texto") or "",
                   adjuntos=d.get("adjuntos") or [], de=d.get("de") or "", de_correo=d.get("de_correo") or "",
                   para=d.get("para") or "", gestion_id=d.get("id"), gestion_proyecto=d.get("proyecto_id"))
    elif code == 404:
        out["aviso"] = "Gestión General no tiene guardado el cuerpo de este correo"
    else:
        out["aviso"] = f"Gestión General respondió {code}: {d.get('detail') or d.get('error') or ''}"
    return out, 200


def abrir_en_outlook(rid: str, cabeceras: dict) -> tuple[dict, int]:
    """Pedido «abrir» al agente del PC por Centro MO (el agente hace Display() del correo; nunca envía).
    Se reenvían las X-SGO7-* del usuario para que Centro MO lo registre a su nombre."""
    if not _RE_ID.match(rid or ""):
        return {"ok": False, "error": "id inválido"}, 400
    h = {k: v for k, v in (cabeceras or {}).items() if v}
    h["Content-Type"] = "application/json"
    return _http(f"{CENTRO_MO_BASE}/api/abrir-outlook/{rid}", method="POST", headers=h)


def pedido_estado(pid: str) -> tuple[dict, int]:
    if not _RE_ID.match(pid or ""):
        return {"error": "pedido inválido"}, 400
    return _http(f"{CENTRO_MO_BASE}/api/pedido/{pid}")


def _clave_codigo(codigo: str) -> str:
    c = re.sub(r"\s+", "", str(codigo or "")).upper()
    m = re.match(r"^(?:LCL)?(63\d{8})$", c)
    return f"LCL{m.group(1)}" if m else c


def _mes(fecha: str) -> str:
    return str(fecha or "")[:7]


# ── Nodos: misma clave en presupuesto, SIP y SGS ─────────────────────────────

def clave_nodo(n) -> str:
    """'NODO 28' → '28' · 'N26-N28' / 'N26_N28' / '26_28' / '26AL28' → '26-28' · '25AAL48' → '25A-48'."""
    s = re.sub(r"\s+", "", str(n or "").upper())
    s = s.replace("NODO", "")
    s = re.sub(r"(\d[A-Z]?)AL(?=N?\d)", r"\1-", s)
    s = s.replace("_", "-")
    s = re.sub(r"(?<![A-Z0-9])N(?=\d)", "", s)
    s = re.sub(r"-N(?=\d)", "-", s)
    s = re.sub(r"(?<!\d)0+(?=\d)", "", s)
    return s.strip("-.")


def _clave_inversa(k: str) -> str:
    partes = k.split("-")
    return "-".join(reversed(partes)) if len(partes) == 2 else ""


def _nodo_sgs(dom: dict) -> str:
    n = str(dom.get("nodo") or "").strip()
    if n and n.lower() != "false":
        return n
    nombre = str(dom.get("name") or "")
    m = re.search(r"(?:NODO\s*)?(N?\d+[A-Z]?(?:\s*(?:AL|[_\-])\s*N?\d+[A-Z]?)?)\s*$", nombre, re.I)
    return m.group(1) if m else ""


# ── SGS: familia completa del código (todos los padres) ──────────────────────

def _padres(c, clave: str) -> list:
    padres = {}
    m = re.match(r"^LCL(63\d{8})$", clave)
    if m:
        num = m.group(1)
        for on in (num, clave):
            for p in c.search_read("project.dominion",
                                   [["order_number", "=", on], ["parent_id", "=", False]],
                                   sc._CAMPOS_DOM, limit=20):
                padres[p["id"]] = p
    elif re.match(r"^DF\d{10,14}$", clave):
        for dom in ([["pdl", "=", clave], ["parent_id", "=", False]],
                    [["name", "ilike", clave], ["parent_id", "=", False]]):
            for p in c.search_read("project.dominion", dom, sc._CAMPOS_DOM, limit=20):
                padres[p["id"]] = p
    if not padres:
        p = sc._buscar_proyecto(c, clave)
        if p:
            raiz, _ = sc._obtener_familia(c, p)
            padres[raiz["id"]] = raiz
    out = sorted(padres.values(), key=lambda p: p["id"])
    return out[:MAX_PADRES]


def _familias_sgs(c, claves: list) -> dict:
    """{clave: {"padres": [dom], "nodos": [dom], "error": str}}."""
    out = {}
    padres_por_clave = {}
    for k in claves:
        try:
            padres_por_clave[k] = _padres(c, k)
        except Exception as e:
            out[k] = {"padres": [], "nodos": [], "error": f"SGS: {str(e)[:120]}"}
    ids_padres = sorted({p["id"] for ps in padres_por_clave.values() for p in ps})
    hijos_por_padre = defaultdict(list)
    for chunk in sc._chunks(ids_padres, 200):
        for h in sc._search_global(c, "project.dominion", [["parent_id", "in", chunk]],
                                   sc._CAMPOS_DOM, 20000):
            hijos_por_padre[sc._pid(h.get("parent_id"))].append(h)
    for k, ps in padres_por_clave.items():
        nodos = []
        for p in ps:
            nodos.append({**p, "_padre_code": p.get("code", "")})
            nodos += [{**h, "_padre_code": p.get("code", "")} for h in hijos_por_padre.get(p["id"], [])]
        out[k] = {"padres": ps, "nodos": nodos,
                  "error": "" if ps else "no aparece en SGS"}
    return out


def _materiales_por_nodo(c, node_ids: list) -> tuple[dict, dict]:
    """({nid: {codigo: {"advance", "product_id"}}}, {product_id: costo})."""
    por_nodo = defaultdict(dict)
    prods = set()
    for chunk in sc._chunks(node_ids, 300):
        for m in sc._search_global(c, "project.material", [["project_id", "in", chunk]],
                                   ["id", "project_id", "product_code", "product_id", "advance"], 50000):
            nid = sc._pid(m.get("project_id"))
            cod = str(m.get("product_code") or "").strip()
            pid = sc._pid(m.get("product_id"))
            if nid and cod:
                prev = por_nodo[nid].get(cod)
                adv = float(m.get("advance") or 0) + (prev["advance"] if prev else 0)
                por_nodo[nid][cod] = {"advance": adv, "product_id": pid}
                if pid:
                    prods.add(pid)
    costos = {}
    try:
        for chunk in sc._chunks(sorted(prods), 300):
            for p in c.search_read("product.product", [["id", "in", chunk]],
                                   ["id", "standard_price"], limit=len(chunk) + 5):
                costos[p["id"]] = float(p.get("standard_price") or 0)
    except Exception:
        pass            # el valor es informativo; sin precio la fila sale en $0
    return por_nodo, costos


# ── Presupuestos de los correos ──────────────────────────────────────────────

def _presupuesto(correos: list) -> dict:
    """{(clave_nodo, codigo): fila} — el correo MÁS RECIENTE que trae ese nodo y material manda."""
    out = {}
    for r in sorted(correos, key=lambda r: r.get("fecha") or ""):
        acum = defaultdict(float)
        nombre = {}
        for f in r.get("filas_mat") or []:
            cod = str(f.get("codigo") or "").strip()
            if not cod or cod.startswith("970"):        # chatarra: SGS nunca la maneja
                continue
            k = clave_nodo(f.get("nodo"))
            if not k:
                continue
            try:
                acum[(k, cod)] += float(f.get("cantidad") or 0)
            except (TypeError, ValueError):
                continue
            nombre[(k, cod)] = str(f.get("nodo") or "")
        im = r.get("instruccion_mat") or {}
        for key, q in acum.items():
            if q <= 0:
                continue
            out[key] = {"nodo_presu": nombre[key], "clave": key[0], "codigo": key[1], "qty": q,
                        "correo_id": r.get("id"), "correo_fecha": r.get("fecha", ""),
                        "correo_asunto": r.get("asunto", ""), "remitente": r.get("remitente", ""),
                        "adjunto": r.get("adjunto", ""),
                        "instr": im.get("valor", ""), "frase": im.get("frase", ""),
                        "instr_fuente": im.get("fuente", "")}
    return out


# ── «Por qué» de cada fila: de dónde sale lo que se manda a liquidar ─────

def _m2o(v) -> str:
    return str(v[1]) if isinstance(v, (list, tuple)) and len(v) > 1 else ("" if v in (None, False) else str(v))


def _bandejas_txt(lst, signo: str) -> str:
    out = []
    for b in lst or []:
        f = b.get("fecha_aprob") or b.get("fecha") or ""
        out.append(f"{b.get('nombre', '')} {signo}{_num(b.get('qty'))}" + (f" ({f})" if f else ""))
    return ", ".join(out)


def _porque(x: dict) -> str:
    """Una frase con la cadena completa: correo → presupuesto → SGS (bandejas, neto, consumo, estado del nodo)."""
    fuente = f"Presupuesto «{x['adjunto']}»" if x.get("adjunto") else "Presupuesto"
    pide = (f"{fuente} del correo del {str(x.get('correo_fecha') or '')[:10]} ({x.get('remitente') or '?'}) "
            f"pide {_num(x['qty'])} en {x['nodo_presu']}")
    est = x["estado"]
    if est == "SIN_NODO":
        return f"{pide}. En SGS no hay un nodo {x['nodo_presu']} en ningún padre del LCL ni otro con los mismos materiales."
    sgs = f"SGS {x['pro']} «{x['nodo_sgs']}»" + (f" ({x['estado_nodo']})" if x.get("estado_nodo") else "")
    if est == "NO_EN_SGS":
        return f"{pide}. {sgs} no tiene ese material proyectado."
    if est == "SIN_DESPACHO":
        return f"{pide}. {sgs} lo tiene proyectado pero sin despacho aprobado de almacén."
    if est == "LIQUIDADO":
        return f"{pide}. {sgs} ya lo tiene consumido ({_num(x['advance'])})."
    mov = []
    if x.get("desp"):
        mov.append("despacho " + _bandejas_txt(x["desp"], "+"))
    if x.get("dev"):
        mov.append("devolución " + _bandejas_txt(x["dev"], "−"))
    base = (f"{pide}. {sgs}" + (": " + " · ".join(mov) if mov else "")
            + f" → neto {_num(x['neto'])}, consumido {_num(x['advance'])}")
    if est == "SOBRANTE":
        return base + ". Lo presupuestado ya está consumido; lo demás es devolución."
    if est == "OMITIDO":
        return base + f". No entra: el correo dice «{x.get('frase') or ''}»."
    txt = base + f". A liquidar {_num(x['a_liquidar'])} (lo menor entre neto y presupuesto, menos lo consumido)."
    if x.get("instr") == "pendiente_almacen" and x.get("frase"):
        txt += f" El correo dice «{x['frase']}»."
    if est == "REVISAR":
        txt += f" Revisar: el correo trae excepciones («{x.get('frase') or ''}»)."
    return txt


# ── Análisis ─────────────────────────────────────────────────────────────────

def analizar(mes: str = "") -> dict:
    """mes = 'AAAA-MM' (mes del correo) o '' / 'todos'."""
    mes = (mes or "").strip()
    todos = mes in ("", "todos")
    try:
        items = _items_centro_mo()
    except Exception as e:
        return {"ok": False, "error": f"No pude leer Centro MO: {e}"}

    correos = [r for r in items
               if r.get("estado") != "descartado" and r.get("codigo")
               and (todos or _mes(r.get("fecha")) == mes)]
    por_codigo = defaultdict(list)
    for r in correos:
        por_codigo[_clave_codigo(r["codigo"])].append(r)
    if not por_codigo:
        return {"ok": True, "mes": mes or "todos", "familias": [], "totales": {},
                "n_correos": 0, "generado": datetime.now().strftime("%Y-%m-%d %H:%M")}

    try:
        c = sc._conectar()
    except Exception as e:
        return {"ok": False, "error": f"No se pudo conectar a SGS: {e}"}

    fams = _familias_sgs(c, sorted(por_codigo))
    todos_nodos = sorted({n["id"] for f in fams.values() for n in f["nodos"]})
    estado = sc._estado_materiales_por_nodo(c, todos_nodos) if todos_nodos else {}
    mats, costos = _materiales_por_nodo(c, todos_nodos) if todos_nodos else ({}, {})

    familias = []
    for clave, regs in sorted(por_codigo.items()):
        fam = fams.get(clave) or {"padres": [], "nodos": [], "error": "sin datos"}
        familias.append(_cruzar(clave, regs, fam, estado, mats, costos))

    familias.sort(key=lambda f: (-f["resumen"]["valor_liquidar"], f["codigo"]))
    tot = defaultdict(float)
    for f in familias:
        for k, v in f["resumen"].items():
            if isinstance(v, (int, float)):
                tot[k] += v
    return {"ok": True, "mes": mes or "todos", "familias": familias,
            "totales": {k: round(v, 2) for k, v in tot.items()},
            "n_correos": len(correos), "generado": datetime.now().strftime("%Y-%m-%d %H:%M")}


def _cruzar(clave: str, regs: list, fam: dict, estado: dict, mats: dict, costos: dict) -> dict:
    nodos = fam["nodos"]
    por_clave = defaultdict(list)
    for n in nodos:
        k = clave_nodo(_nodo_sgs(n))
        if k:
            por_clave[k].append(n)

    def pend_de(n, cod):
        for p in (estado.get(n["id"], {}).get("pendientes") or []):
            if str(p.get("codigo") or "") == cod:
                return p
        return None

    presu = _presupuesto(regs)
    # nodos del presupuesto → candidatos SGS (nombre, nombre invertido, o por materiales)
    claves_presu = sorted({v["clave"] for v in presu.values()})
    cand, via = {}, {}
    for k in claves_presu:
        if por_clave.get(k):
            cand[k], via[k] = por_clave[k], "nombre"
        elif _clave_inversa(k) and por_clave.get(_clave_inversa(k)):
            cand[k], via[k] = por_clave[_clave_inversa(k)], "nombre invertido"
        else:
            # por materiales: nodo SGS con el mismo código y la misma cantidad despachada
            filas_k = [v for v in presu.values() if v["clave"] == k]
            mejores = []
            for n in nodos:
                hits = 0
                for v in filas_k:
                    p = pend_de(n, v["codigo"])
                    if p and abs(float(p.get("neto") or 0) - v["qty"]) <= max(0.5, 0.005 * v["qty"]):
                        hits += 1
                if hits:
                    mejores.append((hits, n))
            mejores.sort(key=lambda t: -t[0])
            if mejores and (len(mejores) == 1 or mejores[0][0] > mejores[1][0]):
                cand[k], via[k] = [mejores[0][1]], "materiales iguales"
            else:
                cand[k], via[k] = [], ""

    filas, usados = [], set()
    for (k, cod), v in sorted(presu.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        base = {**v, "via": via.get(k, ""), "nodo_sgs": "", "pro": "", "pro_id": None,
                "padre": "", "neto": 0.0, "advance": 0.0, "objetivo": 0.0,
                "a_liquidar": 0.0, "costo": 0.0, "valor": 0.0, "nombre": "",
                "lider": "", "bandejas": "", "ult_salida": "",
                "estado_nodo": "", "desp": [], "dev": []}
        cs = cand.get(k) or []
        if not cs:
            filas.append({**base, "estado": "SIN_NODO", "incluir": False})
            continue
        n_pend = next((n for n in cs if pend_de(n, cod)), None)
        n_mat = n_pend or next((n for n in cs if cod in mats.get(n["id"], {})), None)
        if not n_mat:
            n0 = cs[0]
            filas.append({**base, "nodo_sgs": _nodo_sgs(n0), "pro": n0.get("code", ""),
                          "pro_id": n0["id"], "padre": n0.get("_padre_code", ""),
                          "estado_nodo": _m2o(n0.get("status")),
                          "estado": "NO_EN_SGS", "incluir": False})
            continue
        info = mats.get(n_mat["id"], {}).get(cod, {})
        costo = costos.get(info.get("product_id"), 0.0)
        base.update(nodo_sgs=_nodo_sgs(n_mat), pro=n_mat.get("code", ""), pro_id=n_mat["id"],
                    padre=n_mat.get("_padre_code", ""), costo=costo, estado_nodo=_m2o(n_mat.get("status")))
        p = pend_de(n_mat, cod)
        if p:
            usados.add((n_mat["id"], cod))
            neto, adv = float(p.get("neto") or 0), float(p.get("advance") or 0)
            objetivo = round(min(neto, v["qty"]), 4)
            a_liq = round(objetivo - adv, 4)
            base.update(neto=neto, advance=adv, objetivo=objetivo, nombre=p.get("nombre", ""),
                        lider=p.get("lider", ""), bandejas=_just_bandejas(p),
                        ult_salida=p.get("fecha_aprob_ultima", ""),
                        desp=p.get("bandejas_despacho") or [], dev=p.get("bandejas_devolucion") or [])
            if a_liq <= 0.001:
                filas.append({**base, "estado": "SOBRANTE", "incluir": False})
                continue
            base.update(a_liquidar=a_liq, valor=round(a_liq * costo, 2))
            if v["instr"] == "no_mat":
                filas.append({**base, "estado": "OMITIDO", "incluir": False})
            elif v["instr"] == "parcial":
                filas.append({**base, "estado": "REVISAR", "incluir": False})
            else:
                filas.append({**base, "estado": "LIQUIDAR", "incluir": True})
        else:
            adv = float(info.get("advance") or 0)
            base.update(advance=adv)
            filas.append({**base, "estado": "LIQUIDADO" if adv > 0.001 else "SIN_DESPACHO",
                          "incluir": False})

    # Pendiente en la familia que ningún presupuesto respalda (contexto, no entra a la plantilla)
    sin_presu = []
    for n in nodos:
        for p in (estado.get(n["id"], {}).get("pendientes") or []):
            cod = str(p.get("codigo") or "")
            if (n["id"], cod) in usados or float(p.get("pendiente") or 0) <= 0.001:
                continue
            info = mats.get(n["id"], {}).get(cod, {})
            pend = float(p.get("pendiente") or 0)
            sin_presu.append({"nodo_sgs": _nodo_sgs(n), "pro": n.get("code", ""),
                              "padre": n.get("_padre_code", ""), "codigo": cod,
                              "nombre": p.get("nombre", ""), "neto": float(p.get("neto") or 0),
                              "advance": float(p.get("advance") or 0), "pendiente": pend,
                              "valor": round(pend * costos.get(info.get("product_id"), 0.0), 2),
                              "lider": p.get("lider", ""), "ult_salida": p.get("fecha_aprob_ultima", "")})

    for i, f in enumerate(filas):
        f["id"] = f"{clave}|{f['clave']}|{f['codigo']}|{i}"
        f["porque"] = _porque(f)

    cuenta = defaultdict(int)
    for f in filas:
        cuenta[f["estado"]] += 1
    liq = [f for f in filas if f["estado"] == "LIQUIDAR"]
    correos = [{"id": r.get("id"), "fecha": r.get("fecha", ""), "asunto": r.get("asunto", ""),
                "remitente": r.get("remitente", ""), "estado": r.get("estado", ""),
                "adjunto": r.get("adjunto", ""),
                "n_mat": len(r.get("filas_mat") or []),
                "instr": (r.get("instruccion_mat") or {}).get("valor", ""),
                "frase": (r.get("instruccion_mat") or {}).get("frase", "")}
               for r in sorted(regs, key=lambda r: r.get("fecha") or "")]
    return {
        "codigo": clave,
        "tipo": "DF" if clave.startswith("DF") else ("6300" if clave.startswith("LCL") else "Otro"),
        "mes": _mes(max((r.get("fecha") or "") for r in regs)),
        "padres": [p.get("code", "") for p in fam["padres"]],
        "error": fam.get("error", ""),
        "correos": correos,
        "sin_presupuesto_leido": all(not r.get("filas_mat") for r in regs),
        "filas": filas,
        "sin_presupuesto": sorted(sin_presu, key=lambda x: (x["padre"], x["nodo_sgs"], x["codigo"])),
        "resumen": {
            "n_liquidar": len(liq),
            "cant_liquidar": round(sum(f["a_liquidar"] for f in liq), 2),
            "valor_liquidar": round(sum(f["valor"] for f in liq), 2),
            "n_revisar": cuenta["REVISAR"], "n_omitido": cuenta["OMITIDO"],
            "n_sin_nodo": cuenta["SIN_NODO"], "n_no_en_sgs": cuenta["NO_EN_SGS"],
            "n_sin_despacho": cuenta["SIN_DESPACHO"], "n_liquidado": cuenta["LIQUIDADO"],
            "n_sobrante": cuenta["SOBRANTE"], "n_sin_presupuesto": len(sin_presu),
            "valor_sin_presupuesto": round(sum(x["valor"] for x in sin_presu), 2),
        },
    }


# ── Excel ────────────────────────────────────────────────────────────────────

def _encabezado(ws, fila: int, cols: list):
    _autoancho(ws, [w for _, w in cols])
    for c, (name, _) in enumerate(cols, start=1):
        _set(ws, fila, c, name, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, align="center", wrap=True)


def _titulo(ws, texto: str, ncols: int, nota: str = ""):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    _set(ws, 1, 1, texto, bold=True, fg=C_HDR_FG, bg=C_PADRE_BG, size=11)
    if nota:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        _set(ws, 2, 1, nota, fg=C_MUTED, size=9, wrap=True)
        ws.row_dimensions[2].height = 30


_COLS_FILA = [("Código (LCL/DF)", 17), ("Nodo presupuesto", 15), ("Nodo SGS", 12), ("Cómo se halló", 15),
              ("PRO nodo", 14), ("PRO padre", 14), ("Código SAP", 11), ("Descripción", 40),
              ("Presupuesto", 11), ("Neto despachado", 12), ("Consumido", 11), ("Objetivo", 10),
              ("A liquidar", 10), ("Valor", 13), ("Estado", 26), ("Líder (bandejas)", 28),
              ("Última salida almacén", 14), ("Correo (fecha)", 16), ("Remitente", 22),
              ("Instrucción del correo", 26), ("Frase literal", 50), ("Asunto del correo", 50),
              ("Bandejas", 50), ("Adjunto del correo (presupuesto)", 40), ("Estado del nodo SGS", 22),
              ("Por qué", 90)]


def _fila_excel(ws, r: int, cod: str, f: dict, resaltar: bool = False):
    vals = [cod, f["nodo_presu"], f["nodo_sgs"], f["via"], f["pro"], f["padre"], f["codigo"],
            f["nombre"], f["qty"], f["neto"], f["advance"], f["objetivo"], f["a_liquidar"],
            f["valor"], ESTADOS.get(f["estado"], f["estado"]), f["lider"], f["ult_salida"],
            f["correo_fecha"], f["remitente"], INSTR_TXT.get(f["instr"], f["instr"]),
            f["frase"], f["correo_asunto"], f["bandejas"], f.get("adjunto", ""), f.get("estado_nodo", ""),
            f.get("porque") or _porque(f)]
    for c, v in enumerate(vals, start=1):
        num = c in (9, 10, 11, 12, 13)
        _set(ws, r, c, v, align="right" if num else "left",
             num_fmt='"$" #,##0' if c == 14 else ("#,##0.##" if num else None),
             bold=(c == 13 and resaltar), fg=C_WARN if (c == 13 and resaltar) else None,
             wrap=c in (8, 21, 22, 23, 24, 26), size=9 if c in (21, 22, 23, 24, 26) else 10)


def generar_excel(familias: list, mes: str = "") -> tuple[bytes, str]:
    """`familias` = las del análisis, con `filas[*].incluir` como quedó en pantalla."""
    familias = familias or []
    liquidar = [(f, x) for f in familias for x in f.get("filas") or [] if x.get("incluir")
                and x.get("estado") in ("LIQUIDAR", "REVISAR") and float(x.get("a_liquidar") or 0) > 0.001
                and x.get("pro")]
    if not liquidar:
        raise ValueError("No hay filas marcadas para liquidar. Marca filas «Mandado y no liquidado».")

    # Bloques para CRUCE_LIQUIDADOR: un bloque por PRO de nodo, «SGS Neto» = OBJETIVO
    bloques = {}
    for f, x in liquidar:
        b = bloques.setdefault(x["pro"], {"nodo": str(x["nodo_sgs"] or "0"), "pro_code": x["pro"],
                                          "pdl_code": "", "materiales": []})
        b["materiales"].append({"codigo": x["codigo"], "nombre": x["nombre"], "neto": x["objetivo"]})
    bloques = sorted(bloques.values(), key=lambda b: b["pro_code"])

    n_cant = sum(float(x["a_liquidar"]) for _, x in liquidar)
    n_valor = sum(float(x["valor"]) for _, x in liquidar)
    gen = datetime.now().strftime("%d/%m/%Y %H:%M")
    periodo = mes if mes and mes != "todos" else "todos los meses"
    titulo = (f"Exporte validado por correos — {periodo} · {len({f['codigo'] for f, _ in liquidar})} "
              f"código(s) · {len(bloques)} nodo(s) · {len(liquidar)} material(es) · "
              f"{_num(n_cant)} ud · $ {n_valor:,.0f} · generado {gen}")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _hoja_cruce(wb, bloques, titulo)          # la que lee Liquidador 2 (queda activa)

    # RESUMEN por código
    ws = wb.create_sheet("RESUMEN")
    cols = [("Código", 17), ("Tipo", 7), ("Mes del correo", 11), ("Padres SGS", 26), ("Correos", 8),
            ("Instrucciones", 40), ("A liquidar (filas)", 11), ("Cantidad", 11), ("Valor", 14),
            ("Revisar", 8), ("Omitidos", 9), ("Sin nodo", 9), ("No en SGS", 9), ("Sin despacho", 10),
            ("Liquidado", 10), ("Pendiente sin presupuesto", 13), ("Nota", 40)]
    _titulo(ws, titulo, len(cols), "Una fila por código (LCL o DF) con correos en Centro MO en el periodo. "
            "Mandado = el presupuesto del correo trae ese nodo y material.")
    _encabezado(ws, 3, cols)
    r = 4
    for f in familias:
        s = f["resumen"]
        instr = "; ".join(sorted({INSTR_TXT.get(c_["instr"], c_["instr"]) for c_ in f["correos"]}))
        nota = f.get("error") or ("Ningún correo trae presupuesto de materiales leído"
                                  if f.get("sin_presupuesto_leido") else "")
        vals = [f["codigo"], f["tipo"], f["mes"], " + ".join(f["padres"]), len(f["correos"]), instr,
                s["n_liquidar"], s["cant_liquidar"], s["valor_liquidar"], s["n_revisar"], s["n_omitido"],
                s["n_sin_nodo"], s["n_no_en_sgs"], s["n_sin_despacho"], s["n_liquidado"],
                s["n_sin_presupuesto"], nota]
        for c, v in enumerate(vals, start=1):
            _set(ws, r, c, v, align="right" if c in (5, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16) else "left",
                 num_fmt='"$" #,##0' if c == 9 else None, wrap=c in (6, 17), size=9 if c in (6, 17) else 10,
                 bold=c in (1, 9), fg=C_WARN if (c == 7 and s["n_liquidar"]) else None)
        r += 1
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{max(r - 1, 3)}"

    # DETALLE: todas las filas, con su justificación
    ws = wb.create_sheet("DETALLE")
    _titulo(ws, titulo, len(_COLS_FILA), "Cada fila: correo (fecha, remitente, asunto, frase literal sobre el "
            "material) × nodo del presupuesto × material, contra SGS leído en vivo. «Objetivo» = mínimo entre lo "
            "despachado y lo presupuestado; «A liquidar» = Objetivo − Consumido.")
    _encabezado(ws, 3, _COLS_FILA)
    r = 4
    orden = {"LIQUIDAR": 0, "REVISAR": 1, "OMITIDO": 2, "SIN_NODO": 3, "NO_EN_SGS": 4,
             "SIN_DESPACHO": 5, "SOBRANTE": 6, "LIQUIDADO": 7}
    for f in familias:
        for x in sorted(f.get("filas") or [], key=lambda x: (orden.get(x["estado"], 9), x["clave"], x["codigo"])):
            _fila_excel(ws, r, f["codigo"], x, resaltar=x["estado"] == "LIQUIDAR")
            if x["estado"] == "LIQUIDAR" and not x.get("incluir"):
                _set(ws, r, 15, "Mandado y no liquidado (desmarcado en pantalla)", fg=C_MUTED)
            r += 1
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(_COLS_FILA))}{max(r - 1, 3)}"

    # OMITIDOS / REVISAR / SIN NODO aparte, para que nada desaparezca en silencio
    for hoja, estados, nota in (
        ("OMITIDOS", ("OMITIDO",), "El correo dice «solo MO» o «no liquidar material»: no entran a la plantilla."),
        ("REVISAR", ("REVISAR",), "El correo trae excepciones («a excepción de…»). Entran solo si se marcaron en pantalla."),
        ("SIN_NODO", ("SIN_NODO", "NO_EN_SGS", "SIN_DESPACHO"),
         "El presupuesto los manda pero SGS no tiene dónde liquidarlos: nodo inexistente, material no "
         "proyectado o sin despacho aprobado de almacén."),
    ):
        ws = wb.create_sheet(hoja)
        _titulo(ws, titulo, len(_COLS_FILA), nota)
        _encabezado(ws, 3, _COLS_FILA)
        r = 4
        for f in familias:
            for x in f.get("filas") or []:
                if x["estado"] in estados:
                    _fila_excel(ws, r, f["codigo"], x)
                    r += 1
        ws.freeze_panes = "C4"

    # Pendiente en SGS que ningún presupuesto respalda (contexto)
    ws = wb.create_sheet("PENDIENTE_SIN_PRESUPUESTO")
    cols = [("Código", 17), ("PRO padre", 14), ("Nodo SGS", 12), ("PRO nodo", 14), ("Código SAP", 11),
            ("Descripción", 40), ("Neto despachado", 12), ("Consumido", 11), ("Pendiente", 11),
            ("Valor", 13), ("Líder", 28), ("Última salida almacén", 14)]
    _titulo(ws, titulo, len(cols), "Material despachado y sin liquidar en la familia del código que el "
            "presupuesto de los correos del periodo no trae. No entra a la plantilla: revisar con el supervisor.")
    _encabezado(ws, 3, cols)
    r = 4
    for f in familias:
        for x in f.get("sin_presupuesto") or []:
            vals = [f["codigo"], x["padre"], x["nodo_sgs"], x["pro"], x["codigo"], x["nombre"], x["neto"],
                    x["advance"], x["pendiente"], x["valor"], x["lider"], x["ult_salida"]]
            for c, v in enumerate(vals, start=1):
                _set(ws, r, c, v, align="right" if c in (7, 8, 9, 10) else "left",
                     num_fmt='"$" #,##0' if c == 10 else ("#,##0.##" if c in (7, 8, 9) else None),
                     wrap=c == 6)
            r += 1
    ws.freeze_panes = "B4"

    ws = wb.create_sheet("LEEME")
    _autoancho(ws, [110])
    _set(ws, 1, 1, titulo, bold=True, fg=C_HDR_FG, bg=C_PADRE_BG, size=11)
    leeme = [
        ("Qué es", True),
        ("Lo que los correos de Centro MO mandaron a liquidar (presupuesto del supervisor) y SGS muestra "
         "despachado y sin consumir. Cada fila lleva el correo, la frase literal sobre el material y la lectura "
         "de SGS que la justifica.", False),
        ("Regla", True),
        ("Si el correo no dice «solo MO» o «no liquidar material», la planilla estaba habilitada para liquidar "
         "material (regla del 22 oct 2026). «Solo MO» → hoja OMITIDOS. «A excepción de…» → REVISAR.", False),
        ("Cómo se usa", True),
        (f"Liquidador 2 → Subir plantilla. Solo lee la hoja {HOJA_LIQUIDADOR}: un bloque por nodo con su PRO "
         "explícito, así que liquida en el nodo del supervisor sin adivinar nombres ni crear nodos.", False),
        ("Cantidades", True),
        ("«SGS Neto» en la hoja del Liquidador = OBJETIVO = mínimo entre lo despachado y lo presupuestado. "
         "Liquidador 2 liquida Objetivo − consumido en vivo: nunca pasa del presupuesto ni de lo despachado, y lo "
         "liquidado después de generar el archivo se descuenta solo. No editar esa columna.", False),
        ("Nodos", True),
        ("El nodo del presupuesto se busca en TODOS los padres del LCL en SGS por nombre (N26-N28 = 26_28 = "
         "26AL28), luego con el nombre invertido y por último por materiales iguales (mismo código y misma "
         "cantidad despachada). La columna «Cómo se halló» lo dice.", False),
    ]
    r = 3
    for texto, tit in leeme:
        _set(ws, r, 1, texto, bold=tit, fg=C_PADRE_BG if tit else "000000", size=11 if tit else 10, wrap=not tit)
        if not tit:
            ws.row_dimensions[r].height = 44
        r += 1
    wb.active = 0

    slug = re.sub(r"[^A-Za-z0-9]+", "_", periodo).strip("_").upper()
    fname = f"VALIDADO_CORREOS_{slug}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), fname


if __name__ == "__main__":       # prueba: python -m verificador_proy.validado_correos 2026-10 [salida.xlsx]
    import sys
    res = analizar(sys.argv[1] if len(sys.argv) > 1 else "")
    if not res.get("ok"):
        print("ERROR", res.get("error"))
        sys.exit(1)
    print("mes", res["mes"], "· correos", res["n_correos"], "· códigos", len(res["familias"]))
    print("totales", json.dumps(res["totales"], ensure_ascii=False))
    for f in res["familias"]:
        s = f["resumen"]
        print(f"  {f['codigo']:16s} padres={'+'.join(f['padres']):28s} liquidar={s['n_liquidar']:3d} "
              f"cant={s['cant_liquidar']:9.1f} ${s['valor_liquidar']:>13,.0f} rev={s['n_revisar']} "
              f"omit={s['n_omitido']} sin_nodo={s['n_sin_nodo']} no_sgs={s['n_no_en_sgs']} "
              f"sin_desp={s['n_sin_despacho']} liq={s['n_liquidado']} sin_presu={s['n_sin_presupuesto']} {f['error']}")
    if len(sys.argv) > 2:
        data, nombre = generar_excel(res["familias"], res["mes"])
        open(sys.argv[2], "wb").write(data)
        print("excel", sys.argv[2], len(data), "bytes", nombre)
