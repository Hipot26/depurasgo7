"""
export_auto.py — Exportación a Excel del modo "Automático por mes" del Verificador
==================================================================================
Genera un .xlsx con dos hojas a partir de `analisis_mes`:

  1. "Proyectos"             — una fila por familia (PRO padre), filtrable.
  2. "Materiales pendientes" — vista jerárquica PRO padre → PRO hijos → cada
                               material pendiente, con su LÍDER A LIQUIDAR y la
                               ANTIGÜEDAD (días desde el despacho).

Modos:
  * Por mes  → generar_excel_auto(año, mes)
  * Por año  → generar_excel_auto(año)        (recorre los 12 meses con datos)

El líder a liquidar y la antigüedad provienen de `sgs_checker._estado_materiales_por_nodo`
(técnica del liquidador2: el líder sale del partner_id de la bandeja de despacho).
"""

import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from verificador_proy.sgs_checker import (
    analisis_mes, MESES_ES, _conectar, CONTRATO_2026_ID, CONTRATO_2026_NOMBRE,
)

# ── Paleta ───────────────────────────────────────────────────────────────────
C_HDR_BG   = "1F3864"   # azul oscuro encabezado
C_HDR_FG   = "FFFFFF"
C_PADRE_BG = "1A237E"   # índigo — fila de PRO padre
C_PADRE_FG = "FFFFFF"
C_NODO_BG  = "D6E4F0"   # azul claro — fila de PRO hijo / nodo
C_OK       = "1B5E20"   # verde
C_WARN     = "B26A00"   # ámbar
C_MUTED    = "777777"
C_ANTIG_HI = "FDECEA"   # rojo claro — antigüedad alta (>30 días)
C_BORDE    = "BDC3C7"

_THIN = Side(style="thin", color=C_BORDE)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _clasif_txt(c: str) -> str:
    return {"completo": "Completo", "pendiente": "Pendiente",
            "sin_materiales": "Sin materiales"}.get(c, c)


def _tipo_txt(t: str) -> str:
    return {"programado": "6300", "mayor_impacto": "DF", "otro": "Otro"}.get(t, t)


def _antig_txt(d) -> str:
    if d is None:
        return ""
    return f"{d} día" + ("" if d == 1 else "s")


def _num(q) -> str:
    """Formatea cantidad sin decimales innecesarios (720 / 35.5)."""
    q = float(q or 0)
    return str(int(q)) if abs(q - round(q)) < 0.005 else f"{q:g}"


def _just_bandejas(p: dict) -> str:
    """
    Texto que justifica la cantidad de un material a partir de sus bandejas:
      'Despacho: BAN-A (+720), BAN-B (+100) | Devolución: DEV-X (-100) → neto 720'
    """
    desp = p.get("bandejas_despacho", []) or []
    dev  = p.get("bandejas_devolucion", []) or []
    partes = []
    if desp:
        partes.append("Despacho: " + ", ".join(
            f"{b['nombre']} (+{_num(b['qty'])})" for b in desp))
    if dev:
        partes.append("Devolución: " + ", ".join(
            f"{b['nombre']} (-{_num(b['qty'])})" for b in dev))
    if not partes:
        return ""
    return "  |  ".join(partes) + f"  → neto {_num(p.get('neto', 0))}"


def _set(ws, row, col, value, *, bold=False, fg=None, bg=None,
         align="left", num_fmt=None, wrap=False, size=10):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold, color=fg or "000000", size=size)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)
    if num_fmt:
        cell.number_format = num_fmt
    cell.border = _BORDER
    return cell


# ── Recolección de datos ──────────────────────────────────────────────────────

def _meses_a_consultar(año: int, mes: int | None) -> list[int]:
    if mes:
        return [mes]
    hoy = datetime.now()
    tope = hoy.month if año == hoy.year else 12
    return list(range(1, tope + 1))


def _recolectar(año: int, mes: int | None) -> tuple[list, list]:
    """
    Devuelve (familias_anotadas, errores). Cada familia lleva su 'mes'.
    Reutiliza UNA conexión SGS y un caché de estado de materiales entre meses:
    el estado es una foto actual, recalcularlo por mes solo duplica consultas
    (era la causa del 504 del Excel del año).
    """
    familias = []
    errores = []
    try:
        c = _conectar()
    except Exception as e:
        return [], [f"No se pudo conectar a SGS: {e}"]
    estado_cache: dict = {}
    # El AÑO completo es el consolidado del contrato vigente: se excluye el
    # contrato viejo (ya cerró). El Excel de UN mes no se filtra.
    contrato = CONTRATO_2026_ID if not mes else None
    for m in _meses_a_consultar(año, mes):
        try:
            d = analisis_mes(año, m, c=c, estado_cache=estado_cache,
                             solo_contrato_id=contrato)
        except Exception as e:
            errores.append(f"{MESES_ES.get(m, m)}: {e}")
            continue
        if not d.get("ok"):
            errores.append(f"{MESES_ES.get(m, m)}: {d.get('error', 'error')}")
            continue
        for f in d.get("familias", []):
            f["_mes"] = m
            familias.append(f)
    return familias, errores


# ── Hoja 1: Proyectos (resumen filtrable) ─────────────────────────────────────

_COLS_PROY = [
    ("Mes", 12), ("Tipo", 8), ("Estado mat.", 13), ("PDL / OT", 28), ("LCL", 14),
    ("Cód. padre", 14), ("Proyecto padre", 40), ("Supervisor", 18),
    ("MO liquidada", 15), ("Fecha MO", 12),
    ("Antig. (bandeja)", 14), ("Antig. (aprob.)", 14),
    ("Líder a liquidar", 34), ("Nodos pend.", 11),
]


def _lcl_txt(f) -> str:
    lcl = f.get("lcl")
    return lcl if lcl and lcl != "N/A" else "—"


def _hoja_proyectos(wb, familias, titulo, avisos=None):
    ws = wb.create_sheet("Proyectos")
    _set(ws, 1, 1, titulo, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_COLS_PROY))
    ws.row_dimensions[1].height = 22
    if avisos:
        _set(ws, 2, 1, "⚠ AVISOS — meses no incluidos o incompletos: " + " · ".join(avisos),
             bold=True, fg="B00020", wrap=True, size=9)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(_COLS_PROY))
        ws.row_dimensions[2].height = 26

    r = 3
    for i, (name, w) in enumerate(_COLS_PROY, start=1):
        _set(ws, r, i, name, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, align="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{r + 1}"
    ws.auto_filter.ref = f"A{r}:{get_column_letter(len(_COLS_PROY))}{r}"

    r += 1
    for f in familias:
        clasif = f.get("clasificacion", "")
        col_estado = (C_WARN if clasif == "pendiente"
                      else C_OK if clasif == "completo" else C_MUTED)
        antig = f.get("antiguedad_dias")
        antig_ap = f.get("antiguedad_aprob_dias")
        bg_antig = C_ANTIG_HI if (antig is not None and antig > 30) else None
        bg_antig_ap = C_ANTIG_HI if (antig_ap is not None and antig_ap > 30) else None
        _set(ws, r, 1, MESES_ES.get(f.get("_mes"), "").capitalize())
        _set(ws, r, 2, _tipo_txt(f.get("tipo", "")), align="center")
        _set(ws, r, 3, _clasif_txt(clasif), bold=True, fg=col_estado, align="center")
        _set(ws, r, 4, f.get("pdl_code", ""), wrap=True)
        _set(ws, r, 5, _lcl_txt(f), align="center")
        _set(ws, r, 6, f.get("padre_code", ""))
        _set(ws, r, 7, f.get("padre_name", ""), wrap=True)
        _set(ws, r, 8, f.get("supervisor_full") or f.get("supervisor", ""))
        _set(ws, r, 9, f.get("mo_value", 0) or 0, align="right",
             num_fmt='"$"#,##0')
        _set(ws, r, 10, f.get("fecha_mo", ""), align="center")
        _set(ws, r, 11, _antig_txt(antig), align="center", bg=bg_antig)
        _set(ws, r, 12, _antig_txt(antig_ap), align="center", bg=bg_antig_ap)
        _set(ws, r, 13, f.get("lider_liquidar", ""), wrap=True)
        _set(ws, r, 14, f.get("n_pendientes", 0), align="center")
        r += 1
    return ws


# ── Hoja 2: Materiales pendientes (jerárquica) ────────────────────────────────

_COLS_MAT = [
    ("Mes / Nivel", 16), ("PDL / Nodo", 26), ("Código", 16),
    ("Material", 16), ("Nombre material", 42),
    ("Cant. pendiente", 15), ("Líder a liquidar", 34),
    ("Fecha despacho (bandeja)", 16), ("Antig. (bandeja)", 14),
    ("Fecha aprob. almacén", 16), ("Antig. (aprob.)", 14),
    ("Supervisor", 18),
    ("Bandejas (despacho / devolución)", 54),
]


def _hoja_materiales(wb, familias, titulo):
    ws = wb.create_sheet("Materiales pendientes")
    ncol = len(_COLS_MAT)
    _set(ws, 1, 1, titulo, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.row_dimensions[1].height = 22

    r = 3
    for i, (name, w) in enumerate(_COLS_MAT, start=1):
        _set(ws, r, i, name, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, align="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{r + 1}"
    r += 1

    # Solo familias con materiales pendientes. El estado de materiales es una
    # foto ACTUAL: una familia liquidada en varios meses aparecería duplicada
    # con los mismos materiales → conservar solo su aparición más reciente.
    por_padre: dict = {}
    for f in familias:
        if f.get("clasificacion") == "pendiente":
            por_padre[f.get("padre_id") or f.get("padre_code")] = f  # meses ascendentes
    pend_fams = sorted(por_padre.values(),
                       key=lambda f: (f.get("_mes", 0), f.get("pdl_code", "")))

    if not pend_fams:
        _set(ws, r, 1, "No hay materiales pendientes de liquidar en el periodo.",
             fg=C_MUTED)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        return ws

    for f in pend_fams:
        mes_lbl = MESES_ES.get(f.get("_mes"), "").capitalize()
        # ── Fila PRO padre ──
        _lcl = _lcl_txt(f)
        _pdl_lbl = f"{f.get('pdl_code', '')}" + (f"  ·  LCL {_lcl}" if _lcl != "—" else "")
        _set(ws, r, 1, mes_lbl, bold=True, fg=C_PADRE_FG, bg=C_PADRE_BG)
        _set(ws, r, 2, _pdl_lbl, bold=True,
             fg=C_PADRE_FG, bg=C_PADRE_BG)
        _set(ws, r, 3, f.get("padre_code", ""), bold=True, fg=C_PADRE_FG, bg=C_PADRE_BG)
        _set(ws, r, 4, f.get("padre_name", ""), fg=C_PADRE_FG, bg=C_PADRE_BG)
        for c in range(5, 12):
            _set(ws, r, c, "", bg=C_PADRE_BG)
        _set(ws, r, 12, f.get("supervisor_full") or f.get("supervisor", ""), fg=C_PADRE_FG, bg=C_PADRE_BG)
        _set(ws, r, 13, "", bg=C_PADRE_BG)
        padre_row = r
        r += 1
        grupo_inicio = r

        # ── Nodos (PRO hijos) con material pendiente ──
        nodos_pend = [n for n in f.get("nodos_mat", [])
                      if n.get("estado_mat") == "pendiente" and n.get("pendientes")]
        nodos_pend.sort(key=lambda n: (str(n.get("nodo") or ""), n.get("pro_code", "")))
        for n in nodos_pend:
            nodo_lbl = f"N{n['nodo']}" if n.get("nodo") else (n.get("pro_code") or "—")
            _set(ws, r, 1, "  └ Nodo", fg="333333", bg=C_NODO_BG)
            _set(ws, r, 2, nodo_lbl, bold=True, bg=C_NODO_BG)
            _set(ws, r, 3, n.get("pro_code", ""), bg=C_NODO_BG)
            for c in range(4, 14):
                _set(ws, r, c, "", bg=C_NODO_BG)
            r += 1
            # ── Materiales pendientes del nodo ──
            for p in n.get("pendientes", []):
                antig = p.get("antiguedad_dias")
                antig_ap = p.get("antiguedad_aprob_dias")
                bg_antig = C_ANTIG_HI if (antig is not None and antig > 30) else None
                bg_antig_ap = C_ANTIG_HI if (antig_ap is not None and antig_ap > 30) else None
                _set(ws, r, 1, "")
                _set(ws, r, 2, "")
                _set(ws, r, 3, "")
                _set(ws, r, 4, p.get("codigo", ""))
                _set(ws, r, 5, p.get("nombre", ""), wrap=True)
                _set(ws, r, 6, p.get("pendiente", 0), align="right", num_fmt="#,##0.##")
                _set(ws, r, 7, p.get("lider", ""), wrap=True)
                _set(ws, r, 8, p.get("fecha_despacho", ""), align="center")
                _set(ws, r, 9, _antig_txt(antig), align="center", bg=bg_antig)
                _set(ws, r, 10, p.get("fecha_aprobacion", ""), align="center")
                _set(ws, r, 11, _antig_txt(antig_ap), align="center", bg=bg_antig_ap)
                _set(ws, r, 12, "")
                _set(ws, r, 13, _just_bandejas(p), wrap=True, size=9)
                r += 1
        # Agrupar (outline) las filas de la familia bajo el padre
        if r > grupo_inicio:
            for rr in range(grupo_inicio, r):
                ws.row_dimensions[rr].outline_level = 1
        ws.row_dimensions[padre_row].height = 18
    return ws


# ── API pública ───────────────────────────────────────────────────────────────

def generar_excel_auto(año: int, mes: int | None = None):
    """
    Genera el Excel del modo automático. Devuelve (bytes, filename).
    Si `mes` es None o 0 → exporta el AÑO completo (meses con datos).
    Lanza ValueError si no hay datos en el periodo.
    """
    mes = mes or None
    familias, errores = _recolectar(año, mes)

    if mes:
        periodo = f"{MESES_ES.get(mes, mes).capitalize()} {año}"
        fname = f"Verificador_AUTO_{MESES_ES.get(mes, mes)}_{año}.xlsx"
    else:
        periodo = f"Año {año} · Contrato {CONTRATO_2026_NOMBRE}"
        fname = f"Verificador_AUTO_{año}_CONTRATO_2026.xlsx"

    if not familias:
        detalle = (" · ".join(errores)) if errores else "sin proyectos del equipo"
        raise ValueError(f"No hay datos para {periodo} ({detalle}).")

    gen = datetime.now().strftime("%d/%m/%Y %H:%M")
    # Conteos por familia ÚNICA (en el año la misma familia puede aparecer en
    # varios meses; contar filas duplicaría proyectos y pendientes).
    uniq      = {f.get("padre_id") or f.get("padre_code") for f in familias}
    uniq_pend = {f.get("padre_id") or f.get("padre_code") for f in familias
                 if f.get("clasificacion") == "pendiente"}
    detalle_filas = "" if mes else f" · {len(familias)} filas mes-proyecto"
    titulo = (f"Verificador SGS — Automático · {periodo}   "
              f"({len(uniq)} proyectos · {len(uniq_pend)} con material pendiente"
              f"{detalle_filas} · generado {gen})")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # quitar hoja por defecto
    _hoja_proyectos(wb, familias, titulo, avisos=errores)
    _hoja_materiales(wb, familias, titulo)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue(), fname
