"""
generar_excel.py — Generador del reporte BALANCE FINAL SGS (multi-hoja)
========================================================================
Produce un .xlsx con 4 hojas:

  1. Consolidado       — balance mensual (formato original + líderes %)
  2. Prestaciones      — una fila por prestación individual
  3. Análisis PDL      — historial de nodos, detección secuencial vs simultáneo
  4. Gráficas          — tablas de resumen + gráficas dinámicas

Compatibilidad exportar_comparacion.py: las funciones helper públicas
(_escribir_fila_datos, _estilo_borde, _fill, _calcular_totales_por_supervisor,
COLOR_*) permanecen con la misma firma.
"""

import os
import sys
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import (
        Font, Alignment, PatternFill, Border, Side,
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, PieChart, Reference
except ImportError:
    print("ERROR: openpyxl no instalado. Ejecuta: pip install openpyxl")
    sys.exit(1)

import config as cfg

# ── Paleta de colores ──────────────────────────────────────────────────────────
COLOR_HEADER_BG   = "1F3864"
COLOR_HEADER_FG   = "FFFFFF"
COLOR_TOTALES_BG  = "D6E4F0"
COLOR_RESALTADO_BG   = "FFF2CC"
COLOR_MI_BG       = "F2F2F2"
COLOR_PG_BG       = "FFFFFF"
COLOR_BORDE       = "BDC3C7"

# Supervisor resuelto por mapa manual (datos SGS incorrectos)
COLOR_SUP_MANUAL  = "FFF3CD"   # amarillo tipo advertencia

# Prestaciones sheet
COLOR_PREST_HDR   = "1E3A5F"
COLOR_PREST_MI    = "EBF5FB"
COLOR_PREST_PG    = "FDFEFE"
COLOR_PREST_TOT   = "1F3864"

# Análisis PDL
COLOR_PDL_HDR     = "1A237E"
COLOR_PDL_SINGLE  = "E8F5E9"   # verde claro — nodo único
COLOR_PDL_SIMULT  = "FFF9C4"   # amarillo — simultáneo
COLOR_PDL_SEQ     = "FCE4EC"   # rosa — secuencial

# Gráficas
COLOR_GRAF_HDR    = "263238"


def _estilo_borde():
    thin = Side(style="thin", color=COLOR_BORDE)
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _estilo_header():
    return Font(bold=True, color=COLOR_HEADER_FG, size=10)


def _fill(hex_color: str):
    return PatternFill("solid", fgColor=hex_color)


def _fmt_pesos(val) -> str:
    try:
        return f"${float(val):,.2f}"
    except (TypeError, ValueError):
        return str(val or "")


def _calcular_totales_por_supervisor(proyectos: list) -> dict:
    totales = {s: 0.0 for s in cfg.SUPERVISORES_EQUIPO}
    for p in proyectos:
        sup = p.get("supervisor", "").strip().upper()
        for s in cfg.SUPERVISORES_EQUIPO:
            if sup.startswith(s.split()[0]):
                totales[s] += float(p.get("man_obra") or 0)
                break
    return totales


def _lider_display(p: dict) -> str:
    """Devuelve el texto del líder para mostrar en el consolidado.
    Si hay múltiples líderes, muestra 'NOMBRE (pct%) / OTRO (pct%)'."""
    lideres_list = p.get("lideres_list", [])
    mo_total = sum(l["mo"] for l in lideres_list if l["mo"] > 0)
    if not lideres_list or mo_total <= 0:
        return p.get("lider", "")
    if len(lideres_list) == 1:
        return lideres_list[0]["nombre"]
    partes = []
    for l in lideres_list:
        if l["mo"] <= 0:
            continue
        pct = l["mo"] / mo_total * 100
        # Abreviar nombre largo: primera palabra
        nombre_corto = l["nombre"].split()[0]
        partes.append(f"{nombre_corto} ({pct:.0f}%)")
    return " / ".join(partes)


# ── Hoja 1: Consolidado ────────────────────────────────────────────────────────

def _escribir_consolidado(ws, proyectos_ord, mes_str, marcar_supervisor):
    borde = _estilo_borde()
    a_centro = Alignment(horizontal="center", vertical="center", wrap_text=False)
    a_izq    = Alignment(horizontal="left",   vertical="center", wrap_text=False)
    a_der    = Alignment(horizontal="right",  vertical="center", wrap_text=False)

    HEADERS = [
        "FECHA EJECUCIÓN", "PDL/PI/TP/DF/VYP", "LIDER",
        "LCL", "PROCESO", "# COD CONF SGS",
        "$ MAN OBRA", "SUPERVISOR", "ESTADO",
    ]
    # Columnas de totales por supervisor: 10 .. 10+n-1; columna de total general: 10+n
    n_sups    = len(cfg.SUPERVISORES_EQUIPO)
    col_total = 10 + n_sups

    for col_idx, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font      = _estilo_header()
        cell.fill      = _fill(COLOR_HEADER_BG)
        cell.alignment = a_centro
        cell.border    = borde
    for col_idx in range(10, col_total + 1):
        cell = ws.cell(row=1, column=col_idx, value=None)
        cell.fill   = _fill(COLOR_HEADER_BG)
        cell.border = borde

    # Preprocesar proyectos para mostrar líderes con %
    def _prep(p):
        pc = dict(p)
        pc["lider"] = _lider_display(p)
        return pc

    fila2_datos = _prep(proyectos_ord[0]) if proyectos_ord else {}
    _escribir_fila_datos(ws, 2, fila2_datos, marcar_supervisor, borde, a_centro, a_izq, a_der)

    totales_sup   = _calcular_totales_por_supervisor(proyectos_ord)
    total_general = sum(totales_sup.values())

    for i, sup in enumerate(cfg.SUPERVISORES_EQUIPO):
        col = 10 + i
        cell = ws.cell(row=2, column=col, value=sup)
        cell.font      = Font(bold=True, size=9)
        cell.fill      = _fill(COLOR_TOTALES_BG)
        cell.alignment = a_centro
        cell.border    = borde

    fila3_datos = _prep(proyectos_ord[1]) if len(proyectos_ord) > 1 else {}
    _escribir_fila_datos(ws, 3, fila3_datos, marcar_supervisor, borde, a_centro, a_izq, a_der)

    for i, sup in enumerate(cfg.SUPERVISORES_EQUIPO):
        col = 10 + i
        total = totales_sup.get(sup, 0)
        cell = ws.cell(row=3, column=col, value=total if total else 0)
        cell.number_format = '#,##0.00'
        cell.font          = Font(bold=True, size=9)
        cell.fill          = _fill(COLOR_TOTALES_BG)
        cell.alignment     = a_der
        cell.border        = borde

    cell_total = ws.cell(row=3, column=col_total, value=total_general)
    cell_total.number_format = '#,##0.00'
    cell_total.font          = Font(bold=True, size=9, color="C0392B")
    cell_total.fill          = _fill("FFE0CC")
    cell_total.alignment     = a_der
    cell_total.border        = borde

    for idx, p in enumerate(proyectos_ord[2:], start=4):
        _escribir_fila_datos(ws, idx, _prep(p), marcar_supervisor, borde, a_centro, a_izq, a_der)

    for fila_idx in range(4, 4 + max(0, len(proyectos_ord) - 2)):
        for col_idx in range(10, col_total + 1):
            cell = ws.cell(row=fila_idx, column=col_idx)
            if not cell.value:
                cell.border = borde

    fila_resumen = 4 + max(0, len(proyectos_ord) - 2) + 2
    total_general_real = sum(float(p.get("man_obra") or 0) for p in proyectos_ord)
    total_carlos       = totales_sup.get(cfg.SUPERVISOR_YO, 0)

    c1 = ws.cell(row=fila_resumen, column=1, value="TOTAL GENERAL")
    c1.font = Font(bold=True, color="FFFFFF")
    c1.fill = _fill(COLOR_HEADER_BG)
    c7 = ws.cell(row=fila_resumen, column=7, value=total_general_real)
    c7.number_format = '#,##0.00'
    c7.font = Font(bold=True, color="FFFFFF")
    c7.fill = _fill(COLOR_HEADER_BG)

    fila_carlos = fila_resumen + 1
    cy1 = ws.cell(row=fila_carlos, column=1, value=f"TOTAL {cfg.SUPERVISOR_YO}")
    cy1.fill = _fill(COLOR_RESALTADO_BG)
    cy7 = ws.cell(row=fila_carlos, column=7, value=total_carlos)
    cy7.number_format = '#,##0.00'
    cy7.font = Font(bold=True, size=10)
    cy7.fill = _fill(COLOR_RESALTADO_BG)

    anchos = {1:14, 2:22, 3:30, 4:15, 5:9, 6:16, 7:18, 8:14, 9:14}
    for col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = ancho
    for col in range(10, col_total):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions[get_column_letter(col_total)].width = 20
    ws.freeze_panes = "A2"


# ── Hoja 2: Prestaciones individuales ─────────────────────────────────────────

def _escribir_prestaciones(ws, detalle: list, mes_str: str):
    thin = Side(style="thin", color="2E3354")
    brd  = Border(bottom=thin, left=thin, right=thin, top=thin)
    a_c  = Alignment(horizontal="center", vertical="center")
    a_l  = Alignment(horizontal="left",   vertical="center")
    a_r  = Alignment(horizontal="right",  vertical="center")

    headers = [
        "FECHA", "# COD SGS", "PDL / OT (historial)", "NOMBRE PROYECTO",
        "PROCESO", "SUPERVISOR", "LÍDER", "$ PRESTACIÓN",
        "LIQUIDADO POR", "ESTADO", "PDL VIGENTE",
    ]
    widths = [14, 20, 30, 48, 10, 20, 28, 18, 24, 16, 20]

    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    hdr_fill = _fill(COLOR_PREST_HDR)

    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = a_c
        cell.border    = brd
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22

    total_mo = 0.0
    for row_idx, d in enumerate(detalle, start=2):
        proceso = d.get("proceso", "")
        bg = COLOR_PREST_MI if proceso == "MI" else COLOR_PREST_PG
        fill = _fill(bg)

        mo = float(d.get("man_obra") or 0)
        total_mo += mo

        fecha = d.get("fecha", "")
        if hasattr(fecha, "strftime"):
            fecha = fecha.strftime("%Y-%m-%d")

        vals = [
            fecha,
            d.get("pro_code", ""),
            d.get("pdl_code", ""),
            d.get("nombre_pro", ""),
            proceso,
            d.get("supervisor", ""),
            d.get("lider", ""),
            mo,
            d.get("liquidador", ""),
            d.get("state", ""),
            d.get("pdl_ultimo", ""),
        ]
        for col, val in enumerate(vals, start=1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill   = fill
            cell.border = brd
            if col == 8:
                cell.number_format = '"$"#,##0.00'
                cell.alignment = a_r
            elif col in (1, 5, 6, 10):
                cell.alignment = a_c
            else:
                cell.alignment = a_l

    # Fila de total
    total_row = len(detalle) + 2
    tot_fill = _fill(COLOR_PREST_TOT)
    lbl = ws.cell(row=total_row, column=7, value="TOTAL GENERAL")
    lbl.font = Font(bold=True, color="FFFFFF", size=11)
    lbl.fill = tot_fill
    lbl.alignment = a_r
    tot = ws.cell(row=total_row, column=8, value=total_mo)
    tot.font = Font(bold=True, color="6EE7B7", size=11)
    tot.fill = tot_fill
    tot.number_format = '"$"#,##0.00'
    tot.alignment = a_r

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.freeze_panes = "A2"


# ── Hoja 3: Análisis PDL ───────────────────────────────────────────────────────

def _escribir_analisis_pdl(ws, proyectos: list):
    """Análisis de nodos PDL: historia temporal, detección simultáneo vs secuencial."""
    thin = Side(style="thin", color="90A4AE")
    brd  = Border(left=thin, right=thin, top=thin, bottom=thin)
    a_c  = Alignment(horizontal="center", vertical="center", wrap_text=False)
    a_l  = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    a_r  = Alignment(horizontal="right",  vertical="center")

    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    hdr_fill = _fill(COLOR_PDL_HDR)

    HEADERS = [
        "TIPO", "PDL / DF / LCL (historial)", "# NODOS", "SUPERVISOR",
        "LIDER(ES)", "PROCESO", "FECHA PRIMERA PREST.",
        "FECHA ÚLTIMA PREST.", "$ MO TOTAL", "# PRESTACIONES",
        "TIPO LIQUIDACIÓN", "DETALLE NODOS", "PDL VIGENTE",
    ]
    WIDTHS = [8, 30, 9, 14, 36, 10, 20, 20, 18, 14, 22, 70, 20]

    for col, (h, w) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = a_c
        cell.border    = brd
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22

    # Agrupar proyectos por clave PDL (para unificar el análisis)
    # Los proyectos ya vienen mergeados, pero pueden tener múltiples nodos_grupo
    row_idx = 2
    for p in proyectos:
        nodos = p.get("nodos_grupo", [{"pro_code": p.get("pro_code", ""),
                                        "nodo_id": p.get("nodo_id", ""),
                                        "pdl_code": p.get("pdl_code", ""),
                                        "fecha_primera": p.get("fecha_primera"),
                                        "fecha_ultima":  p.get("fecha"),
                                        "mo": p.get("man_obra", 0),
                                        "n_avances": p.get("n_avances", 0),
                                        "lider": p.get("lider", "")}])

        n_nodos = len(nodos)

        # Determinar tipo de liquidación
        if n_nodos == 1:
            tipo_liq = "Nodo único"
            bg = COLOR_PDL_SINGLE
        else:
            fechas_primeras = [n["fecha_primera"] for n in nodos if n.get("fecha_primera")]
            if len(fechas_primeras) >= 2:
                dias_dif = (max(fechas_primeras) - min(fechas_primeras)).days
                if dias_dif <= 1:
                    tipo_liq = "Multi-nodo simultáneo"
                    bg = COLOR_PDL_SIMULT
                else:
                    tipo_liq = f"Multi-nodo secuencial (+{dias_dif}d)"
                    bg = COLOR_PDL_SEQ
            else:
                tipo_liq = "Multi-nodo (sin fechas)"
                bg = COLOR_PDL_SIMULT

        fill = _fill(bg)

        # Líderes con porcentaje
        lideres_list = p.get("lideres_list", [])
        mo_total_lids = sum(l["mo"] for l in lideres_list if l["mo"] > 0)
        if mo_total_lids > 0 and len(lideres_list) > 1:
            lider_txt = " / ".join(
                f"{l['nombre'].split()[0]} ({l['mo']/mo_total_lids*100:.0f}%)"
                for l in lideres_list if l["mo"] > 0
            )
        elif lideres_list:
            lider_txt = lideres_list[0]["nombre"]
        else:
            lider_txt = p.get("lider", "")

        # Detalle de nodos (texto compacto)
        detalle_nodos_parts = []
        for n in nodos:
            fpr = n["fecha_primera"].strftime("%Y-%m-%d") if n.get("fecha_primera") else "?"
            ful = n["fecha_ultima"].strftime("%Y-%m-%d") if n.get("fecha_ultima") else "?"
            nid = n.get("nodo_id") or n.get("pro_code", "")
            pdl_n = n.get("pdl_ultimo") or ""
            detalle_nodos_parts.append(
                f"{nid}: {fpr}→{ful} ${n['mo']:,.0f} ({n['n_avances']} prest.)"
                + (f" PDL {pdl_n}" if pdl_n else "")
            )
        detalle_nodos = " | ".join(detalle_nodos_parts)

        fecha_primera_p = p.get("fecha_primera")
        fecha_ultima_p  = p.get("fecha")
        fp_str = fecha_primera_p.strftime("%Y-%m-%d") if fecha_primera_p else ""
        fu_str = fecha_ultima_p.strftime("%Y-%m-%d")  if fecha_ultima_p  else ""

        vals = [
            p.get("proceso", ""),
            p.get("pdl_code", ""),
            n_nodos,
            p.get("supervisor", ""),
            lider_txt,
            p.get("proceso", ""),
            fp_str,
            fu_str,
            float(p.get("man_obra") or 0),
            p.get("n_avances", 0),
            tipo_liq,
            detalle_nodos,
            p.get("pdl_ultimo", ""),
        ]

        aligns = [a_c, a_l, a_c, a_c, a_l, a_c, a_c, a_c, a_r, a_c, a_c, a_l, a_l]

        for col, (val, aln) in enumerate(zip(vals, aligns), start=1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill      = fill
            cell.border    = brd
            cell.alignment = aln
            if col == 9:
                cell.number_format = '"$"#,##0.00'
            if col == 3 and n_nodos > 1:
                cell.font = Font(bold=True)

        ws.row_dimensions[row_idx].height = 15
        row_idx += 1

    # Fila de totales
    tot_fill = _fill(COLOR_PDL_HDR)
    t1 = ws.cell(row=row_idx + 1, column=1, value="TOTAL")
    t1.font = Font(bold=True, color="FFFFFF")
    t1.fill = tot_fill
    total_mo = sum(float(p.get("man_obra") or 0) for p in proyectos)
    t9 = ws.cell(row=row_idx + 1, column=9, value=total_mo)
    t9.number_format = '"$"#,##0.00'
    t9.font = Font(bold=True, color="FFFFFF")
    t9.fill = tot_fill
    t9.alignment = a_r

    # Leyenda
    leyenda_row = row_idx + 3
    ws.cell(row=leyenda_row, column=1, value="LEYENDA:").font = Font(bold=True, size=9)
    leyenda = [
        (COLOR_PDL_SINGLE, "Nodo único — un solo PRO-code en el DF/LCL"),
        (COLOR_PDL_SIMULT, "Multi-nodo simultáneo — todos los nodos liquidados en la misma fecha (±1 día)"),
        (COLOR_PDL_SEQ,    "Multi-nodo secuencial — nodos liquidados en fechas distintas (PDL agregado después)"),
    ]
    for i, (color, texto) in enumerate(leyenda):
        cell_color = ws.cell(row=leyenda_row + 1 + i, column=1, value="  ")
        cell_color.fill = _fill(color)
        ws.cell(row=leyenda_row + 1 + i, column=2, value=texto).font = Font(size=9)

    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
    ws.freeze_panes = "A2"


# ── Hoja 4: Gráficas ──────────────────────────────────────────────────────────

def _escribir_graficas(ws, proyectos: list, wb):
    """Escribe tablas de datos y gráficas dinámicas."""
    a_c = Alignment(horizontal="center", vertical="center")
    a_r = Alignment(horizontal="right",  vertical="center")
    a_l = Alignment(horizontal="left",   vertical="center")
    thin = Side(style="thin", color="B0BEC5")
    brd  = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _hdr(ws, row, col, val, bg=COLOR_GRAF_HDR, fg="FFFFFF", bold=True, size=10):
        c = ws.cell(row=row, column=col, value=val)
        c.font = Font(bold=bold, color=fg, size=size)
        c.fill = _fill(bg)
        c.alignment = a_c
        c.border = brd
        return c

    def _dat(ws, row, col, val, bg="F5F5F5", num_fmt=None, aln=None):
        c = ws.cell(row=row, column=col, value=val)
        c.fill = _fill(bg)
        c.border = brd
        c.alignment = aln or a_c
        if num_fmt:
            c.number_format = num_fmt
        return c

    # ── TABLA 1: MO por Supervisor (cols A-C, desde fila 1) ────────────────────
    ws.cell(row=1, column=1, value="MO POR SUPERVISOR").font = Font(bold=True, size=11)

    _hdr(ws, 2, 1, "SUPERVISOR")
    _hdr(ws, 2, 2, "$ MAN OBRA")
    _hdr(ws, 2, 3, "# PROY")

    acum_sup = {}
    for p in proyectos:
        s = p.get("supervisor", "?")
        acum_sup.setdefault(s, {"mo": 0.0, "n": 0})
        acum_sup[s]["mo"] += float(p.get("man_obra") or 0)
        acum_sup[s]["n"]  += 1

    sups_ordenados = sorted(acum_sup.items(), key=lambda x: x[1]["mo"], reverse=True)
    t1_start = 3
    t1_end   = t1_start + len(sups_ordenados) - 1

    row_idx_t1 = t1_start
    for sup, datos in sups_ordenados:
        bg = "FFF9C4" if sup == cfg.SUPERVISOR_YO else "F5F5F5"
        _dat(ws, row_idx_t1, 1, sup,         bg=bg, aln=a_l)
        _dat(ws, row_idx_t1, 2, datos["mo"], bg=bg, num_fmt='"$"#,##0.00', aln=a_r)
        _dat(ws, row_idx_t1, 3, datos["n"],  bg=bg)
        row_idx_t1 += 1

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 10

    # ── TABLA 2: MO por Proceso (cols E-G, desde fila 1) ──────────────────────
    ws.cell(row=1, column=5, value="MI vs PG").font = Font(bold=True, size=11)
    _hdr(ws, 2, 5, "PROCESO")
    _hdr(ws, 2, 6, "$ MAN OBRA")
    _hdr(ws, 2, 7, "# PROY")

    acum_proc = {}
    for p in proyectos:
        pr = p.get("proceso", "?")
        acum_proc.setdefault(pr, {"mo": 0.0, "n": 0})
        acum_proc[pr]["mo"] += float(p.get("man_obra") or 0)
        acum_proc[pr]["n"]  += 1

    t2_start = 3
    row_idx_t2 = t2_start
    for proc, datos in sorted(acum_proc.items()):
        _dat(ws, row_idx_t2, 5, proc,        aln=a_c)
        _dat(ws, row_idx_t2, 6, datos["mo"], num_fmt='"$"#,##0.00', aln=a_r)
        _dat(ws, row_idx_t2, 7, datos["n"])
        row_idx_t2 += 1
    t2_end = row_idx_t2 - 1

    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 10

    # ── TABLA 3: MO por Fecha (cols I-K, desde fila 1) ────────────────────────
    ws.cell(row=1, column=9, value="EVOLUCIÓN TEMPORAL").font = Font(bold=True, size=11)
    _hdr(ws, 2, 9,  "FECHA")
    _hdr(ws, 2, 10, "$ MAN OBRA")
    _hdr(ws, 2, 11, "# PROY")

    acum_fecha = {}
    for p in proyectos:
        f = p.get("fecha")
        if f:
            fs = f.strftime("%Y-%m-%d") if hasattr(f, "strftime") else str(f)[:10]
            acum_fecha.setdefault(fs, {"mo": 0.0, "n": 0})
            acum_fecha[fs]["mo"] += float(p.get("man_obra") or 0)
            acum_fecha[fs]["n"]  += 1

    t3_start = 3
    row_idx_t3 = t3_start
    for fecha_s in sorted(acum_fecha.keys()):
        datos = acum_fecha[fecha_s]
        _dat(ws, row_idx_t3, 9,  fecha_s,     aln=a_c)
        _dat(ws, row_idx_t3, 10, datos["mo"], num_fmt='"$"#,##0.00', aln=a_r)
        _dat(ws, row_idx_t3, 11, datos["n"])
        row_idx_t3 += 1
    t3_end = row_idx_t3 - 1

    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 20
    ws.column_dimensions["K"].width = 10

    # ── TABLA 4: Top Líderes por MO (cols M-N, desde fila 1) ──────────────────
    ws.cell(row=1, column=13, value="TOP LÍDERES").font = Font(bold=True, size=11)
    _hdr(ws, 2, 13, "LÍDER")
    _hdr(ws, 2, 14, "$ MAN OBRA")

    acum_lider = {}
    for p in proyectos:
        for ll in p.get("lideres_list", []):
            nombre = ll["nombre"]
            acum_lider[nombre] = acum_lider.get(nombre, 0.0) + ll["mo"]
    # Si no hay lideres_list, usar lider simple
    if not acum_lider:
        for p in proyectos:
            nombre = p.get("lider") or "?"
            acum_lider[nombre] = acum_lider.get(nombre, 0.0) + float(p.get("man_obra") or 0)

    t4_start = 3
    row_idx_t4 = t4_start
    for nombre, mo in sorted(acum_lider.items(), key=lambda x: x[1], reverse=True)[:15]:
        _dat(ws, row_idx_t4, 13, nombre, aln=a_l)
        _dat(ws, row_idx_t4, 14, mo,     num_fmt='"$"#,##0.00', aln=a_r)
        row_idx_t4 += 1
    t4_end = row_idx_t4 - 1

    ws.column_dimensions["M"].width = 35
    ws.column_dimensions["N"].width = 20

    # ── GRÁFICA 1: BarChart MO por Supervisor ─────────────────────────────────
    chart1 = BarChart()
    chart1.type        = "bar"          # horizontal
    chart1.grouping    = "clustered"
    chart1.title       = "$ MAN OBRA por Supervisor"
    chart1.y_axis.title = "Supervisor"
    chart1.x_axis.title = "$ MAN OBRA"
    chart1.width  = 20
    chart1.height = 12
    chart1.style  = 10

    data_c1   = Reference(ws, min_col=2, min_row=2, max_row=t1_end)
    cats_c1   = Reference(ws, min_col=1, min_row=t1_start, max_row=t1_end)
    chart1.add_data(data_c1, titles_from_data=True)
    chart1.set_categories(cats_c1)
    chart1.series[0].graphicalProperties.solidFill = "1F3864"

    anchor_row = max(t1_end, t2_end, t3_end, t4_end) + 3
    ws.add_chart(chart1, f"A{anchor_row}")

    # ── GRÁFICA 2: PieChart MI vs PG ──────────────────────────────────────────
    if t2_end >= t2_start:
        chart2 = PieChart()
        chart2.title  = "Distribución MI vs PG"
        chart2.width  = 14
        chart2.height = 10
        chart2.style  = 10

        data_c2 = Reference(ws, min_col=6, min_row=2, max_row=t2_end)
        cats_c2 = Reference(ws, min_col=5, min_row=t2_start, max_row=t2_end)
        chart2.add_data(data_c2, titles_from_data=True)
        chart2.set_categories(cats_c2)
        ws.add_chart(chart2, f"L{anchor_row}")

    # ── GRÁFICA 3: BarChart Evolución Temporal ────────────────────────────────
    if t3_end >= t3_start:
        chart3 = BarChart()
        chart3.type        = "col"
        chart3.grouping    = "clustered"
        chart3.title       = "MAN OBRA por Fecha"
        chart3.x_axis.title = "Fecha"
        chart3.y_axis.title = "$ MAN OBRA"
        chart3.width  = 24
        chart3.height = 12
        chart3.style  = 10

        data_c3 = Reference(ws, min_col=10, min_row=2, max_row=t3_end)
        cats_c3 = Reference(ws, min_col=9,  min_row=t3_start, max_row=t3_end)
        chart3.add_data(data_c3, titles_from_data=True)
        chart3.set_categories(cats_c3)
        chart3.series[0].graphicalProperties.solidFill = "2874A6"
        anchor_row2 = anchor_row + 22
        ws.add_chart(chart3, f"A{anchor_row2}")

    # ── GRÁFICA 4: BarChart Top Líderes ───────────────────────────────────────
    if t4_end >= t4_start:
        chart4 = BarChart()
        chart4.type        = "bar"
        chart4.grouping    = "clustered"
        chart4.title       = "MAN OBRA por Líder (Top 15)"
        chart4.y_axis.title = "Líder"
        chart4.x_axis.title = "$ MAN OBRA"
        chart4.width  = 20
        chart4.height = 12
        chart4.style  = 10

        data_c4 = Reference(ws, min_col=14, min_row=2, max_row=t4_end)
        cats_c4 = Reference(ws, min_col=13, min_row=t4_start, max_row=t4_end)
        chart4.add_data(data_c4, titles_from_data=True)
        chart4.set_categories(cats_c4)
        chart4.series[0].graphicalProperties.solidFill = "1A5276"
        anchor_row2 = anchor_row + 22
        ws.add_chart(chart4, f"L{anchor_row2}")

    ws.freeze_panes = None


# ── Función principal ──────────────────────────────────────────────────────────

def generar_excel(
    proyectos: list,
    mes_str: str = None,
    detalle: list = None,
    incluir_todos: bool = True,
    marcar_supervisor: str = None,
) -> str:
    """
    Genera el Excel multi-hoja de balance y lo guarda en OUTPUT_DIR.

    Args:
        proyectos:        lista consolidada devuelta por leer_proyectos_mes()["consolidado"]
        mes_str:          nombre del mes (ej: "MAYO 2026")
        detalle:          lista de prestaciones individuales (opcional)
        incluir_todos:    parámetro de compatibilidad (no usado)
        marcar_supervisor: supervisor a resaltar (default: SUPERVISOR_YO)

    Returns:
        Ruta al archivo generado.
    """
    if not proyectos:
        print("  AVISO: Lista de proyectos vacía — no se generará el Excel")
        return ""

    if mes_str is None:
        hoy = datetime.now()
        MESES = {1:"ENERO",2:"FEBRERO",3:"MARZO",4:"ABRIL",5:"MAYO",
                 6:"JUNIO",7:"JULIO",8:"AGOSTO",9:"SEPTIEMBRE",
                 10:"OCTUBRE",11:"NOVIEMBRE",12:"DICIEMBRE"}
        mes_str = f"{MESES[hoy.month]} {hoy.year}"

    if marcar_supervisor is None:
        marcar_supervisor = cfg.SUPERVISOR_YO

    def sort_key(p):
        f = p.get("fecha")
        return (f.toordinal() if f else 0, p.get("supervisor", "zzz"))

    proyectos_ord = sorted(proyectos, key=sort_key)

    wb = openpyxl.Workbook()

    # ── Hoja 1: Consolidado ───────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Consolidado"
    _escribir_consolidado(ws1, proyectos_ord, mes_str, marcar_supervisor)

    # ── Hoja 2: Prestaciones individuales ────────────────────────────────────
    ws2 = wb.create_sheet("Prestaciones")
    _escribir_prestaciones(ws2, detalle or [], mes_str)

    # ── Hoja 3: Análisis PDL ─────────────────────────────────────────────────
    ws3 = wb.create_sheet("Análisis PDL")
    _escribir_analisis_pdl(ws3, proyectos_ord)

    # ── Hoja 4: Gráficas ─────────────────────────────────────────────────────
    ws4 = wb.create_sheet("Gráficas")
    _escribir_graficas(ws4, proyectos_ord, wb)

    # ── Guardar ───────────────────────────────────────────────────────────────
    os.makedirs(cfg.OUTPUT_DIR, exist_ok=True)
    timestamp     = datetime.now().strftime("%Y%m%d_%H%M")
    nombre_archivo = f"BALANCE_FINAL_SGS_{mes_str.replace(' ', '_')}_{timestamp}.xlsx"
    ruta           = os.path.join(cfg.OUTPUT_DIR, nombre_archivo)
    wb.save(ruta)

    total_general = sum(float(p.get("man_obra") or 0) for p in proyectos_ord)
    total_carlos  = sum(float(p.get("man_obra") or 0) for p in proyectos_ord
                        if p.get("supervisor", "").upper().startswith(cfg.SUPERVISOR_YO.upper().split()[0]))
    print(f"\n  Excel generado: {ruta}")
    print(f"  Proyectos: {len(proyectos_ord)}  |  Prestaciones: {len(detalle or [])}")
    print(f"  Total MO: ${total_general:>20,.2f}")
    print(f"  Total {cfg.SUPERVISOR_YO}: ${total_carlos:>18,.2f}")

    return ruta


# ── Helper público mantenido por compatibilidad (usado por exportar_comparacion) ──

def _escribir_fila_datos(ws, fila_num, p, marcar_sup, borde, a_centro, a_izq, a_der):
    """Escribe una fila de datos del proyecto en la hoja."""
    if not p:
        return

    sup = p.get("supervisor", "").strip().upper()
    es_resaltado = sup.startswith(marcar_sup.upper().split()[0])

    bg_color = COLOR_RESALTADO_BG if es_resaltado else (
        COLOR_MI_BG if p.get("proceso") == "MI" else COLOR_PG_BG
    )

    fecha = p.get("fecha")
    cell_fecha = ws.cell(row=fila_num, column=1,
                         value=fecha.strftime("%Y-%m-%d") if fecha else "")
    cell_fecha.alignment = a_centro
    cell_fecha.fill      = _fill(bg_color)
    cell_fecha.border    = borde

    # PDL: vigente primero y el historial completo detrás — la casilla de SGS acumula
    # todos los PDL del trabajo y cualquiera puede ser el que busque el ING (Ctrl+F).
    pdl_raw = str(p.get("pdl_code", "") or "")
    pdl_vig = str(p.get("pdl_ultimo") or "")
    if pdl_vig and pdl_vig != pdl_raw:
        otros = [s.strip().strip("/") for s in pdl_raw.replace(":", ";").replace(",", ";").split(";")]
        otros = [s for s in otros if s and s != pdl_vig]
        pdl_txt = f"{pdl_vig}  |  hist: {'; '.join(otros)}" if otros else pdl_vig
    else:
        pdl_txt = pdl_raw
    cell_pdl = ws.cell(row=fila_num, column=2, value=pdl_txt)
    cell_pdl.alignment = a_izq
    cell_pdl.fill      = _fill(bg_color)
    cell_pdl.border    = borde

    cell_lider = ws.cell(row=fila_num, column=3, value=p.get("lider", ""))
    cell_lider.alignment = a_izq
    cell_lider.fill      = _fill(bg_color)
    cell_lider.border    = borde

    cell_lcl = ws.cell(row=fila_num, column=4, value=p.get("lcl", "N/A"))
    cell_lcl.alignment = a_centro
    cell_lcl.fill      = _fill(bg_color)
    cell_lcl.border    = borde

    cell_proc = ws.cell(row=fila_num, column=5, value=p.get("proceso", ""))
    cell_proc.alignment = a_centro
    cell_proc.fill      = _fill(bg_color)
    cell_proc.border    = borde

    cell_cod = ws.cell(row=fila_num, column=6, value=p.get("pro_code", ""))
    cell_cod.alignment = a_centro
    cell_cod.fill      = _fill(bg_color)
    cell_cod.border    = borde
    if es_resaltado:
        cell_cod.font = Font(bold=True, size=10)

    mo = float(p.get("man_obra") or 0)
    cell_mo = ws.cell(row=fila_num, column=7, value=mo)
    cell_mo.number_format = '#,##0.00'
    cell_mo.alignment     = a_der
    cell_mo.fill          = _fill(bg_color)
    cell_mo.border        = borde
    if es_resaltado:
        cell_mo.font = Font(bold=True, size=10)

    cell_sup = ws.cell(row=fila_num, column=8, value=p.get("supervisor", ""))
    cell_sup.alignment = a_centro
    cell_sup.fill      = _fill(bg_color)
    cell_sup.border    = borde
    if es_resaltado:
        cell_sup.font = Font(bold=True, size=10)
    raw = p.get("supervisor_raw", "")
    if raw and raw != p.get("supervisor", ""):
        try:
            from openpyxl.comments import Comment as _Comment
            nota = f"SGS responsible_id: {raw}\n(asignado por MAPA_PROYECTO_SUPERVISOR)"
            cell_sup.comment = _Comment(nota, "Sistema")
            # Fondo especial si fue resuelto por mapeo manual
            if "LIQUIDADOR" in raw or p.get("supervisor_manual"):
                cell_sup.fill = _fill(COLOR_SUP_MANUAL)
        except Exception:
            pass

    cell_est = ws.cell(row=fila_num, column=9, value=p.get("estado", ""))
    cell_est.alignment = a_centro
    cell_est.fill      = _fill(bg_color)
    cell_est.border    = borde

    for col in range(10, 10 + len(cfg.SUPERVISORES_EQUIPO) + 1):
        c = ws.cell(row=fila_num, column=col)
        c.fill   = _fill(bg_color)
        c.border = borde


# ── Ejecución directa: datos de prueba ────────────────────────────────────────
if __name__ == "__main__":
    print("\n  Generando Excel de PRUEBA con datos de ejemplo...")

    datos_prueba = [
        {
            "pro_code": "PRO-052835", "pro_id": 1,
            "fecha": datetime(2026, 5, 7), "fecha_primera": datetime(2026, 5, 5),
            "pdl_code": "DF202664341631", "nodo_id": "NODO-001",
            "lider": "LIDER UNO",
            "lideres_list": [{"nombre": "LIDER UNO", "mo": 2707667.35}],
            "lcl": "N/A", "proceso": "MI", "supervisor": "SUPERVISOR_1",
            "man_obra": 2707667.35, "estado": "ok", "fuente_mo": "prueba",
            "n_avances": 1, "liquidadores": [], "liquidado_por": "LIQUIDADOR",
            "nodos_grupo": [{"pro_code": "PRO-052835", "nodo_id": "NODO-001",
                             "pdl_code": "DF202664341631",
                             "fecha_primera": datetime(2026, 5, 5),
                             "fecha_ultima":  datetime(2026, 5, 7),
                             "mo": 2707667.35, "n_avances": 1, "lider": "LIDER UNO"}],
        },
        {
            "pro_code": "DF202664341631", "pro_id": 2,
            "fecha": datetime(2026, 5, 14), "fecha_primera": datetime(2026, 5, 7),
            "pdl_code": "DF202664341631", "nodo_id": "NODO-002",
            "lider": "LIDER TRES",
            "lideres_list": [
                {"nombre": "LIDER TRES",   "mo": 2800000.0},
                {"nombre": "LIDER CUATRO", "mo": 1598909.12},
            ],
            "lcl": "N/A", "proceso": "MI", "supervisor": "SUPERVISOR_1",
            "man_obra": 4398909.12, "estado": "ok", "fuente_mo": "prueba",
            "n_avances": 2, "liquidadores": [], "liquidado_por": "LIQUIDADOR",
            "nodos_grupo": [
                {"pro_code": "PRO-052801", "nodo_id": "NODO-002A",
                 "pdl_code": "DF202664341631",
                 "fecha_primera": datetime(2026, 5, 7),
                 "fecha_ultima":  datetime(2026, 5, 7),
                 "mo": 2800000.0, "n_avances": 1, "lider": "LIDER TRES"},
                {"pro_code": "PRO-052802", "nodo_id": "NODO-002B",
                 "pdl_code": "DF202664341631",
                 "fecha_primera": datetime(2026, 5, 14),
                 "fecha_ultima":  datetime(2026, 5, 14),
                 "mo": 1598909.12, "n_avances": 1, "lider": "LIDER CUATRO"},
            ],
        },
        {
            "pro_code": "PRO-050219", "pro_id": 3,
            "fecha": datetime(2026, 5, 3), "fecha_primera": datetime(2026, 5, 3),
            "pdl_code": "CDOPP2026R00887", "nodo_id": "",
            "lider": "LIDER DOS",
            "lideres_list": [{"nombre": "LIDER DOS", "mo": 2707628.68}],
            "lcl": "6300963765", "proceso": "PG", "supervisor": "SUPERVISOR_2",
            "man_obra": 2707628.68, "estado": "ok", "fuente_mo": "prueba",
            "n_avances": 1, "liquidadores": [], "liquidado_por": "LIQUIDADOR",
            "nodos_grupo": [{"pro_code": "PRO-050219", "nodo_id": "",
                             "pdl_code": "CDOPP2026R00887",
                             "fecha_primera": datetime(2026, 5, 3),
                             "fecha_ultima":  datetime(2026, 5, 3),
                             "mo": 2707628.68, "n_avances": 1, "lider": "LIDER DOS"}],
        },
    ]

    detalle_prueba = [
        {"fecha": datetime(2026, 5, 5), "pro_code": "PRO-052835",
         "pdl_code": "DF202664341631", "nombre_pro": "CIRCUITO NORTE",
         "proceso": "MI", "supervisor": "SUPERVISOR_1", "lider": "LIDER UNO",
         "man_obra": 2707667.35, "liquidador": "LIQUIDADOR EQUIPO", "state": "done"},
        {"fecha": datetime(2026, 5, 7), "pro_code": "PRO-052801",
         "pdl_code": "DF202664341631", "nombre_pro": "CIRCUITO NORTE NODO B",
         "proceso": "MI", "supervisor": "SUPERVISOR_1", "lider": "LIDER TRES",
         "man_obra": 2800000.0, "liquidador": "LIQUIDADOR EQUIPO", "state": "done"},
        {"fecha": datetime(2026, 5, 14), "pro_code": "PRO-052802",
         "pdl_code": "DF202664341631", "nombre_pro": "CIRCUITO NORTE NODO C",
         "proceso": "MI", "supervisor": "SUPERVISOR_1", "lider": "LIDER CUATRO",
         "man_obra": 1598909.12, "liquidador": "LIQUIDADOR EQUIPO", "state": "done"},
        {"fecha": datetime(2026, 5, 3), "pro_code": "PRO-050219",
         "pdl_code": "CDOPP2026R00887", "nombre_pro": "REPOSICION REDES PG",
         "proceso": "PG", "supervisor": "SUPERVISOR_2", "lider": "LIDER DOS",
         "man_obra": 2707628.68, "liquidador": "LIQUIDADOR EQUIPO", "state": "done"},
    ]

    ruta = generar_excel(datos_prueba, "MAYO 2026", detalle=detalle_prueba)
    print(f"\n  Archivo guardado en: {ruta}")
