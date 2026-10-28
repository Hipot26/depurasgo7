"""
balance_proy/blueprint.py — Blueprint Flask del módulo BALANCE FINAL SGS
=========================================================================
Rutas montadas en / (root):
  GET  /                          → dashboard principal
  POST /api/consultar             → consulta proyectos desde SGS API
  POST /api/consultar-reporte     → lee proyectos desde Excel de producción
  POST /api/subir-reporte         → sube Excel de informe de producción
  GET  /api/reporte-info          → info del reporte disponible
  POST /api/generar-excel         → genera y descarga Excel multi-hoja
  POST /api/subir-presupuesto     → sube Excel de presupuesto (Wilerson)
  GET  /api/presupuesto-info      → info del presupuesto disponible
  POST /api/comparar              → compara SGS vs presupuesto
  POST /api/exportar-comparacion  → descarga Excel de comparación
  POST /api/exportar-detalle      → descarga Excel de prestaciones individuales
"""

import os
import sys
import json
import pickle
import uuid
from datetime import datetime
from functools import wraps

# ── Path setup — asegura que los módulos del proyecto sean encontrables ────────
_HERE    = os.path.dirname(os.path.abspath(__file__))
_APP_DIR = os.path.dirname(_HERE)           # raíz del proyecto (balance_proy/..)
if _APP_DIR not in sys.path:
    sys.path.insert(0, _APP_DIR)

# ── Auto-detect templates (local dev usa templates_balance, VPS usa templates) ─
_TPL_LOCAL = os.path.join(_APP_DIR, "templates_balance")
_TPL_VPS   = os.path.join(_APP_DIR, "templates")
_TPL_DIR   = _TPL_LOCAL if os.path.isdir(_TPL_LOCAL) else _TPL_VPS

from flask import (
    Blueprint, render_template, request, jsonify,
    send_file, redirect, url_for, session,
)

import config as cfg


# ── Auth compartida con el Verificador ────────────────────────────────────────
# Ambos módulos comparten la misma sesión Flask (mismo secret_key).
# El login en /verificador/login establece session["autenticado"]=True.

def _login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("autenticado"):
            if request.path.startswith("/api/"):
                return jsonify({
                    "ok": False,
                    "error": "Sesión expirada. Recarga e inicia sesión.",
                    "login_required": True,
                }), 401
            return redirect(url_for("verificador_proy.login", next=request.path))
        return f(*args, **kwargs)
    return decorated


# ── Cache persistente por sesión (cross-worker safe) ─────────────────────────
_CACHE_DIR = os.path.join(_APP_DIR, "output", "_cache")


def _save_cache(proyectos: list, detalle: list, mes_str: str) -> str:
    """Serializa cache a archivo pickle y devuelve la clave."""
    os.makedirs(_CACHE_DIR, exist_ok=True)
    key  = str(uuid.uuid4())
    path = os.path.join(_CACHE_DIR, f"{key}.pkl")
    with open(path, "wb") as f:
        pickle.dump({"proyectos": proyectos, "detalle": detalle, "mes_str": mes_str}, f)
    return key


def _load_cache(key: str) -> dict | None:
    """Carga cache desde archivo pickle por clave de sesión."""
    if not key:
        return None
    path = os.path.join(_CACHE_DIR, f"{key}.pkl")
    if os.path.exists(path):
        with open(path, "rb") as f:
            return pickle.load(f)
    return None

bp = Blueprint(
    "balance",
    __name__,
    url_prefix=None,
    template_folder=_TPL_DIR,
)

# ── Directorios de trabajo ─────────────────────────────────────────────────────
INPUT_DIR  = os.path.join(_APP_DIR, "input")
PRESUP_DIR = os.path.join(_APP_DIR, "input_presupuesto")
os.makedirs(INPUT_DIR,  exist_ok=True)
os.makedirs(PRESUP_DIR, exist_ok=True)

# ── Cache de proyectos en memoria (estado del servidor) ───────────────────────
_cache_proyectos:   list = []
_cache_detalle:     list = []
_cache_mes_str:     str  = ""
_cache_comparacion: dict = {}

MESES_ES = {
    1:"ENERO", 2:"FEBRERO", 3:"MARZO", 4:"ABRIL", 5:"MAYO",
    6:"JUNIO", 7:"JULIO", 8:"AGOSTO", 9:"SEPTIEMBRE",
    10:"OCTUBRE", 11:"NOVIEMBRE", 12:"DICIEMBRE",
}


def _proyectos_a_json(proyectos: list) -> list:
    result = []
    for p in proyectos:
        fecha = p.get("fecha")
        result.append({
            "pro_code":      p.get("pro_code", ""),
            "pro_id":        p.get("pro_id", 0),
            "fecha":         fecha.strftime("%Y-%m-%d") if fecha else "",
            "pdl_code":      p.get("pdl_code", ""),
            "lider":         p.get("lider", ""),
            "lcl":           p.get("lcl", "N/A"),
            "proceso":       p.get("proceso", ""),
            "supervisor":    p.get("supervisor", ""),
            "man_obra":      float(p.get("man_obra") or 0),
            "estado":        p.get("estado", ""),
            "fuente_mo":     p.get("fuente_mo", ""),
            "n_detalles":    p.get("n_detalles", p.get("n_avances", 0)),
            "tecnologia":    p.get("tecnologia", ""),
            "liquidado_por": p.get("liquidado_por", ""),
            "liquidadores":  p.get("liquidadores", []),
        })
    return result


def _totales_por_supervisor(proyectos: list):
    totales = {s: 0.0 for s in cfg.SUPERVISORES_EQUIPO}
    otros   = {}
    for p in proyectos:
        sup = p.get("supervisor", "?")
        if sup in totales:
            totales[sup] += float(p.get("man_obra") or 0)
        else:
            otros[sup] = otros.get(sup, 0.0) + float(p.get("man_obra") or 0)
    total_general = sum(totales.values()) + sum(otros.values())
    return totales, otros, total_general


# ── Rutas ──────────────────────────────────────────────────────────────────────

@bp.route("/")
@_login_required
def index():
    hoy = datetime.now()
    reporte_disponible = any(
        f.endswith(".xlsx") or f.endswith(".xls")
        for f in os.listdir(INPUT_DIR)
    ) if os.path.isdir(INPUT_DIR) else False

    reporte_nombre = ""
    if reporte_disponible:
        archivos = [
            (os.path.getmtime(os.path.join(INPUT_DIR, f)), f)
            for f in os.listdir(INPUT_DIR)
            if f.endswith(".xlsx") or f.endswith(".xls")
        ]
        if archivos:
            reporte_nombre = sorted(archivos, reverse=True)[0][1]

    return render_template(
        "index.html",
        año_actual=hoy.year,
        mes_actual=hoy.month,
        meses=MESES_ES,
        supervisores=cfg.SUPERVISORES_EQUIPO,
        liquidador=getattr(cfg, "LIQUIDADOR", cfg.SUPERVISOR_YO if hasattr(cfg, "SUPERVISOR_YO") else "SERGIO DURAN"),
        reporte_disponible=reporte_disponible,
        reporte_nombre=reporte_nombre,
    )


@bp.route("/api/consultar", methods=["POST"])
@_login_required
def api_consultar():
    global _cache_proyectos, _cache_detalle, _cache_mes_str

    data = request.get_json() or {}
    año  = int(data.get("año",  datetime.now().year))
    mes  = int(data.get("mes",  datetime.now().month))

    try:
        import importlib, leer_proyectos_mes as _lpm_mod
        importlib.reload(_lpm_mod)
        resultado = _lpm_mod.leer_proyectos_mes(año=año, mes=mes, verbose=False)
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

    proyectos   = resultado.get("consolidado", [])
    detalle_raw = resultado.get("detalle", [])

    _cache_proyectos = proyectos
    _cache_mes_str   = f"{MESES_ES[mes]} {año}"

    totales, otros, total_general = _totales_por_supervisor(proyectos)

    detalle_json = []
    for d in detalle_raw:
        fecha = d.get("fecha")
        detalle_json.append({
            "fecha":      fecha.strftime("%Y-%m-%d") if fecha else "",
            "pro_code":   d.get("pro_code", ""),
            "pdl_code":   d.get("pdl_code", ""),
            "nombre_pro": d.get("nombre_pro", ""),
            "proceso":    d.get("proceso", ""),
            "supervisor": d.get("supervisor", ""),
            "lider":      d.get("lider", ""),
            "man_obra":   float(d.get("man_obra") or 0),
            "liquidador": d.get("liquidador", ""),
            "state":      d.get("state", ""),
        })
    _cache_detalle = detalle_json

    # Persistir cache a archivo — cross-worker safe
    cache_key = _save_cache(proyectos, detalle_raw, _cache_mes_str)
    session["balance_cache_key"] = cache_key

    return jsonify({
        "ok":            True,
        "mes_str":       _cache_mes_str,
        "proyectos":     _proyectos_a_json(proyectos),
        "detalle":       detalle_json,
        "totales":       totales,
        "otros":         otros,
        "total_general": total_general,
        "n":             len(proyectos),
        "n_detalle":     len(detalle_json),
        "fuente":        "api_sgs",
    })


@bp.route("/api/consultar-reporte", methods=["POST"])
@_login_required
def api_consultar_reporte():
    global _cache_proyectos, _cache_mes_str

    data = request.get_json() or {}
    año  = int(data.get("año",  datetime.now().year))
    mes  = int(data.get("mes",  datetime.now().month))

    try:
        import importlib, leer_reporte_produccion as _lrp_mod
        importlib.reload(_lrp_mod)
        proyectos = _lrp_mod.leer_reporte_produccion(año=año, mes=mes, verbose=False)
    except FileNotFoundError as e:
        return jsonify({"ok": False, "error": str(e)}), 400
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

    _cache_proyectos = proyectos
    _cache_mes_str   = f"{MESES_ES[mes]} {año}"

    totales, otros, total_general = _totales_por_supervisor(proyectos)

    return jsonify({
        "ok":            True,
        "mes_str":       _cache_mes_str,
        "proyectos":     _proyectos_a_json(proyectos),
        "totales":       totales,
        "otros":         otros,
        "total_general": total_general,
        "n":             len(proyectos),
        "fuente":        "reporte_produccion",
    })


@bp.route("/api/subir-reporte", methods=["POST"])
@_login_required
def api_subir_reporte():
    if "archivo" not in request.files:
        return jsonify({"ok": False, "error": "No se recibió ningún archivo"}), 400
    archivo = request.files["archivo"]
    if not archivo.filename:
        return jsonify({"ok": False, "error": "Nombre de archivo vacío"}), 400
    nombre = archivo.filename
    if not (nombre.endswith(".xlsx") or nombre.endswith(".xls")):
        return jsonify({"ok": False, "error": "Solo se aceptan archivos .xlsx o .xls"}), 400
    ruta = os.path.join(INPUT_DIR, nombre)
    archivo.save(ruta)
    return jsonify({"ok": True, "nombre": nombre, "ruta": ruta})


@bp.route("/api/reporte-info", methods=["GET"])
@_login_required
def api_reporte_info():
    archivos = []
    if os.path.isdir(INPUT_DIR):
        for f in os.listdir(INPUT_DIR):
            if f.endswith(".xlsx") or f.endswith(".xls"):
                ruta = os.path.join(INPUT_DIR, f)
                archivos.append({
                    "nombre":   f,
                    "mtime":    os.path.getmtime(ruta),
                    "size_kb":  round(os.path.getsize(ruta) / 1024, 1),
                })
    archivos.sort(key=lambda x: x["mtime"], reverse=True)
    return jsonify({"ok": True, "archivos": archivos})


@bp.route("/api/generar-excel", methods=["POST"])
@_login_required
def api_generar_excel():
    # Intentar cache en memoria primero, luego cache de archivo (cross-worker safe)
    proyectos = _cache_proyectos
    detalle   = _cache_detalle
    mes_str   = _cache_mes_str

    if not proyectos:
        cached = _load_cache(session.get("balance_cache_key"))
        if cached:
            proyectos = cached.get("proyectos", [])
            detalle   = cached.get("detalle", [])
            mes_str   = cached.get("mes_str", "")

    if not proyectos:
        return jsonify({"ok": False, "error": "Consulta primero los proyectos (botón Consultar SGS)."}), 400

    try:
        from generar_excel import generar_excel
        ruta = generar_excel(proyectos, mes_str, detalle=detalle)
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

    return send_file(
        ruta,
        as_attachment=True,
        download_name=os.path.basename(ruta),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/api/subir-presupuesto", methods=["POST"])
@_login_required
def api_subir_presupuesto():
    if "archivo" not in request.files:
        return jsonify({"ok": False, "error": "No se recibió ningún archivo"}), 400
    archivo = request.files["archivo"]
    if not archivo.filename:
        return jsonify({"ok": False, "error": "Nombre de archivo vacío"}), 400
    nombre = archivo.filename
    if not (nombre.endswith(".xlsx") or nombre.endswith(".xls")):
        return jsonify({"ok": False, "error": "Solo se aceptan archivos .xlsx o .xls"}), 400
    ruta = os.path.join(PRESUP_DIR, nombre)
    archivo.save(ruta)
    return jsonify({"ok": True, "nombre": nombre})


@bp.route("/api/presupuesto-info", methods=["GET"])
@_login_required
def api_presupuesto_info():
    archivos = []
    if os.path.isdir(PRESUP_DIR):
        for f in os.listdir(PRESUP_DIR):
            if f.endswith(".xlsx") or f.endswith(".xls"):
                ruta = os.path.join(PRESUP_DIR, f)
                archivos.append({
                    "nombre":  f,
                    "mtime":   os.path.getmtime(ruta),
                    "size_kb": round(os.path.getsize(ruta) / 1024, 1),
                })
    archivos.sort(key=lambda x: x["mtime"], reverse=True)
    return jsonify({"ok": True, "archivos": archivos})


@bp.route("/api/comparar", methods=["POST"])
@_login_required
def api_comparar():
    global _cache_comparacion, _cache_proyectos

    data = request.get_json() or {}
    año  = int(data.get("año",  datetime.now().year))
    mes  = int(data.get("mes",  datetime.now().month))

    archivos_presup = []
    if os.path.isdir(PRESUP_DIR):
        archivos_presup = [
            f for f in os.listdir(PRESUP_DIR)
            if f.endswith(".xlsx") or f.endswith(".xls")
        ]
    if not archivos_presup:
        return jsonify({"ok": False, "error": "No hay Excel de presupuesto. Súbelo primero en la pestaña."}), 400

    archivos_presup.sort(
        key=lambda f: os.path.getmtime(os.path.join(PRESUP_DIR, f)),
        reverse=True,
    )
    archivo = archivos_presup[0]
    ruta    = os.path.join(PRESUP_DIR, archivo)

    try:
        import importlib, leer_proyectos_mes as _lpm_mod
        importlib.reload(_lpm_mod)
        resultado_sgs = _lpm_mod.leer_proyectos_mes(año=año, mes=mes, verbose=False)
    except Exception as e:
        return jsonify({"ok": False, "error": f"Error al consultar SGS: {e}"}), 500

    proyectos_sgs = resultado_sgs.get("consolidado", [])
    if proyectos_sgs:
        _cache_proyectos = proyectos_sgs

    try:
        import importlib, comparar_sgs_presupuesto as _cmp_mod
        importlib.reload(_cmp_mod)
        resultado_cmp = _cmp_mod.comparar_sgs_presupuesto(ruta, proyectos_sgs)
    except Exception as e:
        return jsonify({"ok": False, "error": f"Error al comparar: {e}"}), 500

    mes_str = f"{MESES_ES[mes]} {año}"
    _cache_comparacion = {"mes_str": mes_str, "archivo": archivo, **resultado_cmp}

    return jsonify({
        "ok":      True,
        "mes_str": mes_str,
        "archivo": archivo,
        **resultado_cmp,
    })


@bp.route("/api/exportar-comparacion", methods=["POST"])
@_login_required
def api_exportar_comparacion():
    if not _cache_comparacion:
        return jsonify({"ok": False, "error": "Ejecuta la comparación primero."}), 400
    try:
        import importlib, exportar_comparacion as _exp_mod
        importlib.reload(_exp_mod)
        ruta = _exp_mod.exportar_excel_comparacion(_cache_comparacion, _cache_proyectos)
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500
    return send_file(
        ruta,
        as_attachment=True,
        download_name=os.path.basename(ruta),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/api/exportar-detalle", methods=["POST"])
@_login_required
def api_exportar_detalle():
    payload  = request.get_json() or {}
    detalle  = payload.get("detalle", [])
    mes_str  = payload.get("mes_str", "")

    if not detalle:
        return jsonify({"ok": False, "error": "No hay prestaciones para exportar."}), 400

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Prestaciones individuales"

        hdr_font  = Font(bold=True, color="FFFFFF", size=11)
        hdr_fill  = PatternFill("solid", fgColor="1E3A5F")
        mi_fill   = PatternFill("solid", fgColor="EBF5FB")
        pg_fill   = PatternFill("solid", fgColor="FDFEFE")
        tot_fill  = PatternFill("solid", fgColor="1F3864")
        thin      = Side(style="thin", color="2E3354")
        brd       = Border(bottom=thin, left=thin, right=thin, top=thin)

        headers = ["FECHA", "# COD SGS", "PDL / OT", "NOMBRE PROYECTO",
                   "PROCESO", "SUPERVISOR", "LÍDER", "$ PRESTACIÓN",
                   "LIQUIDADO POR", "ESTADO"]
        widths  = [14, 20, 24, 48, 10, 20, 28, 18, 24, 16]

        for col, (h, w) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font      = hdr_font
            cell.fill      = hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = brd
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.row_dimensions[1].height = 22

        for row_idx, d in enumerate(detalle, start=2):
            fill = mi_fill if d.get("proceso") == "MI" else pg_fill
            valores = [
                d.get("fecha", ""),
                d.get("pro_code", ""),
                d.get("pdl_code", ""),
                d.get("nombre_pro", ""),
                d.get("proceso", ""),
                d.get("supervisor", ""),
                d.get("lider", ""),
                float(d.get("man_obra") or 0),
                d.get("liquidador", ""),
                d.get("state", ""),
            ]
            for col, val in enumerate(valores, start=1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.fill   = fill
                cell.border = brd
                if col == 8:
                    cell.number_format = '"$"#,##0.00'
                    cell.alignment = Alignment(horizontal="right")

        total_row = len(detalle) + 2
        total     = sum(float(d.get("man_obra") or 0) for d in detalle)
        lbl = ws.cell(row=total_row, column=7, value="TOTAL GENERAL")
        lbl.font = Font(bold=True, color="FFFFFF", size=11)
        lbl.fill = tot_fill
        lbl.alignment = Alignment(horizontal="right")
        tot = ws.cell(row=total_row, column=8, value=total)
        tot.font = Font(bold=True, color="6EE7B7", size=11)
        tot.fill = tot_fill
        tot.number_format = '"$"#,##0.00'
        tot.alignment = Alignment(horizontal="right")

        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        ws.freeze_panes = "A2"

        output_dir = os.path.join(_APP_DIR, "output")
        os.makedirs(output_dir, exist_ok=True)
        ts       = datetime.now().strftime("%Y%m%d_%H%M%S")
        mes_slug = mes_str.replace(" ", "_") if mes_str else ts
        nombre   = f"PRESTACIONES_{mes_slug}_{ts}.xlsx"
        ruta     = os.path.join(output_dir, nombre)
        wb.save(ruta)

    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

    return send_file(
        ruta,
        as_attachment=True,
        download_name=nombre,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
