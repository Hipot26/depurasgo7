"""exportar_comparacion.py
Genera un Excel con dos hojas:
  Hoja 1 "SGS vs Presupuesto" — tabla de comparación (verde/rojo/amarillo)
  Hoja 2 "Balance SGS"        — balance general del mes (mismo formato que generar_excel)
"""

import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

import config as cfg
from generar_excel import (
    _estilo_borde, _fill, _calcular_totales_por_supervisor, _escribir_fila_datos,
    COLOR_HEADER_BG, COLOR_TOTALES_BG, COLOR_CARLOS_BG,
)

# ── Paleta hoja comparación ────────────────────────────────────────────────────
FILL_LIQ   = PatternFill('solid', fgColor='C6EFCE')
FILL_PEND  = PatternFill('solid', fgColor='FFC7CE')
FILL_EXTRA = PatternFill('solid', fgColor='FFEB9C')
FILL_HDR   = PatternFill('solid', fgColor='1F3864')
FILL_SUB   = PatternFill('solid', fgColor='D6E4F0')

FONT_LIQ   = Font(color='276221', size=9)
FONT_PEND  = Font(color='9C0006', size=9)
FONT_EXTRA = Font(color='9C6500', size=9)
FONT_HDR   = Font(bold=True, color='FFFFFF', size=10)


def _fmt_pesos(val) -> str:
    try:
        return f"${float(val):,.2f}"
    except (TypeError, ValueError):
        return str(val or "—")


def _parse_fecha_safe(f):
    if isinstance(f, datetime):
        return f
    if isinstance(f, str) and f:
        for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
            try:
                return datetime.strptime(f[:len(fmt)], fmt)
            except ValueError:
                continue
    return None


def exportar_excel_comparacion(cache_cmp: dict, sgs_proyectos: list) -> str:
    mes_str         = cache_cmp.get('mes_str', 'MES')
    filas           = cache_cmp.get('filas', [])
    archivo         = cache_cmp.get('archivo', '')
    total_liquidado = cache_cmp.get('total_liquidado', 0)
    total_pendiente = cache_cmp.get('total_pendiente', 0)
    total_extra_sgs = cache_cmp.get('total_extra_sgs', 0)
    mao_esperado    = cache_cmp.get('mao_esperado', 0.0)
    mao_sgs_total   = cache_cmp.get('mao_sgs_total', 0.0)

    wb = openpyxl.Workbook()

    a_centro = Alignment(horizontal='center', vertical='center', wrap_text=False)
    a_izq    = Alignment(horizontal='left',   vertical='center', wrap_text=False)
    a_der    = Alignment(horizontal='right',  vertical='center', wrap_text=False)
    borde    = _estilo_borde()

    # ── HOJA 1: SGS vs Presupuesto ────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = 'SGS vs Presupuesto'

    N_COLS = 14

    # Fila 1: título
    ws1.merge_cells(f'A1:{get_column_letter(N_COLS)}1')
    c = ws1.cell(row=1, column=1,
                 value=f'SGS vs Presupuesto — {mes_str}  |  Presupuesto: {archivo}')
    c.font = Font(bold=True, color='FFFFFF', size=11)
    c.fill = FILL_HDR
    c.alignment = a_centro

    # Fila 2: resumen
    ws1.merge_cells(f'A2:{get_column_letter(N_COLS)}2')
    txt = (f'Liquidados: {total_liquidado}  |  Pendientes: {total_pendiente}  |  '
           f'Solo SGS: {total_extra_sgs}  |  '
           f'MAO Presup: {_fmt_pesos(mao_esperado)}  |  '
           f'MAO SGS Total: {_fmt_pesos(mao_sgs_total)}')
    c = ws1.cell(row=2, column=1, value=txt)
    c.font = Font(bold=True, size=9, color='1F3864')
    c.fill = FILL_SUB
    c.alignment = a_izq

    # Fila 3: encabezados
    HDRS = [
        'FECHA', 'PROCESO', 'PROYECTO / ANTECEDENTE', 'LCL',
        'PDT / PDL', 'TIPO', 'MUNICIPIO', 'SUPERVISOR',
        '$ PRESUPUESTO', 'ESTADO', 'VÍA MATCH',
        '$ SGS', 'LIQUIDADO POR', 'COD SGS',
    ]
    for ci, h in enumerate(HDRS, start=1):
        c = ws1.cell(row=3, column=ci, value=h)
        c.font = FONT_HDR
        c.fill = FILL_HDR
        c.alignment = a_centro
        c.border = borde

    # Columnas alineadas a la derecha (monetarias) y al centro
    NUM_COLS    = {9, 12}
    CENTER_COLS = {1, 4, 5, 6, 10, 11, 14}

    for ri, f in enumerate(filas, start=4):
        if f.get('sgs_extra'):
            fill, font, estado_txt = FILL_EXTRA, FONT_EXTRA, 'Solo SGS'
        elif f.get('liquidado'):
            fill, font, estado_txt = FILL_LIQ,   FONT_LIQ,   'Liquidado'
        else:
            fill, font, estado_txt = FILL_PEND,  FONT_PEND,  'Pendiente'

        sgs = f.get('sgs') or {}
        vals = [
            f.get('fecha', ''),
            f.get('proceso', ''),
            f.get('proyecto', ''),
            f.get('lcl', ''),
            f.get('pdt', ''),
            f.get('tipo_match', ''),
            f.get('municipio', ''),
            f.get('supervisor_excel', ''),
            f.get('valor_mao') or None,
            estado_txt,
            f.get('match_via', ''),
            sgs.get('man_obra') or None,
            sgs.get('liquidado_por', ''),
            sgs.get('pro_code', ''),
        ]
        for ci, v in enumerate(vals, start=1):
            c = ws1.cell(row=ri, column=ci, value=v)
            c.font   = font
            c.fill   = fill
            c.border = borde
            c.alignment = a_der if ci in NUM_COLS else (a_centro if ci in CENTER_COLS else a_izq)
            if ci in NUM_COLS and isinstance(v, (int, float)):
                c.number_format = '#,##0.00'

    # Fila de totales
    fila_tot = 4 + len(filas) + 1
    for ci in range(1, N_COLS + 1):
        c = ws1.cell(row=fila_tot, column=ci)
        c.fill = FILL_HDR
        c.border = borde
    ws1.cell(row=fila_tot, column=1, value='TOTALES').font = Font(bold=True, color='FFFFFF', size=10)
    for col, val in ((9, mao_esperado), (12, mao_sgs_total)):
        c = ws1.cell(row=fila_tot, column=col, value=val)
        c.number_format = '#,##0.00'
        c.font = Font(bold=True, color='FFFFFF', size=10)
        c.alignment = a_der

    # Anchos hoja 1
    for col, w in {1:12, 2:9, 3:32, 4:14, 5:20, 6:7, 7:14, 8:16,
                   9:18, 10:12, 11:10, 12:18, 13:20, 14:14}.items():
        ws1.column_dimensions[get_column_letter(col)].width = w
    ws1.freeze_panes = 'A4'

    # ── HOJA 2: Balance SGS ───────────────────────────────────────────────────
    ws2 = wb.create_sheet(title='Balance SGS')
    _escribir_balance_sheet(ws2, sgs_proyectos, mes_str, borde, a_centro, a_izq, a_der)

    # ── Guardar ────────────────────────────────────────────────────────────────
    os.makedirs(cfg.OUTPUT_DIR, exist_ok=True)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M')
    nombre = f'COMPARACION_{mes_str.replace(" ", "_")}_{timestamp}.xlsx'
    ruta   = os.path.join(cfg.OUTPUT_DIR, nombre)
    wb.save(ruta)
    return ruta


def _escribir_balance_sheet(ws, proyectos, mes_str, borde, a_centro, a_izq, a_der):
    """Escribe la hoja Balance SGS con el mismo formato que generar_excel."""
    if not proyectos:
        ws.cell(row=1, column=1,
                value='Sin proyectos SGS en caché — ejecuta "Consultar API SGS" antes de exportar.')
        return

    def sort_key(p):
        f = _parse_fecha_safe(p.get('fecha'))
        return (f.toordinal() if f else 0, p.get('supervisor', 'zzz'))

    # Normalizar fecha a datetime para _escribir_fila_datos
    proyectos_ord = []
    for p in sorted(proyectos, key=sort_key):
        pc = dict(p)
        pc['fecha'] = _parse_fecha_safe(pc.get('fecha'))
        proyectos_ord.append(pc)

    totales_sup   = _calcular_totales_por_supervisor(proyectos_ord)
    total_general = sum(totales_sup.values())
    marcar_sup    = cfg.SUPERVISOR_YO

    HEADERS = ['FECHA EJECUCIÓN', 'PDL/PI/TP/DF/VYP', 'LIDER', 'LCL', 'PROCESO',
               '# COD CONF SGS', '$ MAN OBRA', 'SUPERVISOR', 'ESTADO']
    for ci, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font      = Font(bold=True, color='FFFFFF', size=10)
        c.fill      = _fill(COLOR_HEADER_BG)
        c.alignment = a_centro
        c.border    = borde
    for ci in range(10, 19):
        c = ws.cell(row=1, column=ci)
        c.fill   = _fill(COLOR_HEADER_BG)
        c.border = borde

    _escribir_fila_datos(ws, 2, proyectos_ord[0] if proyectos_ord else {},
                         marcar_sup, borde, a_centro, a_izq, a_der)
    for i, sup in enumerate(cfg.SUPERVISORES_EQUIPO):
        c = ws.cell(row=2, column=10 + i, value=sup)
        c.font = Font(bold=True, size=9)
        c.fill = _fill(COLOR_TOTALES_BG)
        c.alignment = a_centro
        c.border = borde

    _escribir_fila_datos(ws, 3, proyectos_ord[1] if len(proyectos_ord) > 1 else {},
                         marcar_sup, borde, a_centro, a_izq, a_der)
    for i, sup in enumerate(cfg.SUPERVISORES_EQUIPO):
        total = totales_sup.get(sup, 0)
        c = ws.cell(row=3, column=10 + i, value=total or 0)
        c.number_format = '#,##0.00'
        c.font = Font(bold=True, size=9)
        c.fill = _fill(COLOR_TOTALES_BG)
        c.alignment = a_der
        c.border = borde
    c_tot = ws.cell(row=3, column=18, value=total_general)
    c_tot.number_format = '#,##0.00'
    c_tot.font = Font(bold=True, size=9, color='C0392B')
    c_tot.fill = _fill('FFE0CC')
    c_tot.alignment = a_der
    c_tot.border = borde

    for idx, p in enumerate(proyectos_ord[2:], start=4):
        _escribir_fila_datos(ws, idx, p, marcar_sup, borde, a_centro, a_izq, a_der)

    fila_resumen = 4 + max(0, len(proyectos_ord) - 2) + 2
    total_real = sum(float(p.get('man_obra') or 0) for p in proyectos_ord)

    c1 = ws.cell(row=fila_resumen, column=1, value='TOTAL GENERAL')
    c1.font = Font(bold=True, color='FFFFFF')
    c1.fill = _fill(COLOR_HEADER_BG)
    c7 = ws.cell(row=fila_resumen, column=7, value=total_real)
    c7.number_format = '#,##0.00'
    c7.font = Font(bold=True, color='FFFFFF')
    c7.fill = _fill(COLOR_HEADER_BG)

    fila_yo = fila_resumen + 1
    total_yo = totales_sup.get(cfg.SUPERVISOR_YO, 0)
    cy1 = ws.cell(row=fila_yo, column=1, value=f'TOTAL {cfg.SUPERVISOR_YO}')
    cy1.fill = _fill(COLOR_CARLOS_BG)
    cy7 = ws.cell(row=fila_yo, column=7, value=total_yo)
    cy7.number_format = '#,##0.00'
    cy7.font = Font(bold=True, size=10)
    cy7.fill = _fill(COLOR_CARLOS_BG)

    for col, w in {1:14, 2:22, 3:22, 4:15, 5:9, 6:16, 7:18, 8:14, 9:14}.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for col in range(10, 18):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions[get_column_letter(18)].width = 20
    ws.freeze_panes = 'A2'
