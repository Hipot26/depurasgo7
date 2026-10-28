"""
parser_excel.py — Extrae MO y MAT del Excel de liquidación.
Solo lectura de archivos locales. No toca SGS.
Lógica extraída de sgs_api/liquidacion.py.
"""

import re
import unicodedata
import openpyxl

# ── Hojas y marcadores ────────────────────────────────────────────────────────
_HOJAS_PRIORITARIAS_MO = [
    "LIQUIDACION", "LIQ", "PRESTACIONES", "ACTIVIDADES",
    "SIP_MANO_OBRA", "MANO_OBRA",
]
_MARCADORES_PRO   = {"PRO", "PROYECTO", "PROJECT", "ODM", "NODO_PRO", "NODO"}
_MARCADORES_COD   = {"CODIGO", "COD", "CODE", "PRESTACION", "ACTIVIDAD", "ACTIVITY"}
_MARCADORES_QTY   = {"CANTIDAD", "CANT", "QTY", "AMOUNT", "UNIDADES", "TOTAL"}
_MARCADORES_LIDER = {"LIDER", "RESPONSABLE", "PARTNER", "TECNICO"}
_EXCLUIR          = {"DESC", "NOMBRE", "OBSERV", "NOTA"}


def _sin_tildes(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )


def _col_match(v: str, marcadores: set, excluir: set = None) -> bool:
    v_norm = _sin_tildes(v)
    if excluir and any(e in v_norm for e in excluir):
        return False
    return any(m in v_norm for m in marcadores)


def _parse_float(val) -> float | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace("$", "").replace(" ", "").replace("\xa0", "")
    if not s:
        return None
    has_dot   = "." in s
    has_comma = "," in s
    if has_dot and has_comma:
        if s.rfind(".") > s.rfind(","):
            s = s.replace(",", "")
        else:
            s = s.replace(".", "").replace(",", ".")
    elif has_comma:
        partes = s.split(",")
        if len(partes) == 2 and len(partes[1]) <= 2:
            s = s.replace(",", ".")
        else:
            s = s.replace(",", "")
    elif has_dot and s.count(".") > 1:
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def _normalizar_codigo_proyecto(val) -> str:
    raw = str(val or "").strip().upper()
    if not raw or raw in {"0", "N/A", "NA", "-", "NONE"}:
        return ""
    m = re.search(r"\bLCL\s*[-_:]?\s*(6300\d+)\b", raw)
    if m:
        return m.group(1)
    m = re.search(r"\b(6300\d{4,})\b", raw)
    if m:
        return m.group(1)
    m = re.search(r"\b(DF[\w\d]+)\b", raw)
    if m:
        return m.group(1)
    # Marcadores de "aún no asignado" ("PENDIENTE", "SIN ASIGNAR") NO son
    # códigos de proyecto: un código siempre lleva al menos un dígito.
    if not re.search(r"\d", raw):
        return ""
    return raw


def _normalizar_pdl(val) -> str:
    raw = str(val or "").strip().upper()
    if not raw or raw in {"-", "0", "N/A", "NA", "NONE"}:
        return ""
    matches = re.findall(r"\b(CDOPPI[\w\d-]+)\b", raw)
    if matches:
        return ";".join(dict.fromkeys(matches))  # orden preservado, sin duplicados
    return raw


# ── Parser MO ─────────────────────────────────────────────────────────────────
def parsear_mo(wb: openpyxl.Workbook) -> tuple[list, list]:
    """
    Extrae filas de mano de obra.
    Devuelve (rows, errores).
    Cada row: {fila, pro, codigo, cantidad, lider, _odm}
    """
    rows, errs = [], []

    ws = None
    for nombre in _HOJAS_PRIORITARIAS_MO:
        if nombre in wb.sheetnames:
            ws = wb[nombre]
            break
    if ws is None:
        ws = wb.active

    col = {"pro": None, "cod": None, "qty": None, "lider": None}
    header_row = None

    for r in range(1, 21):
        vals = [str(ws.cell(r, c).value or "").upper().strip() for c in range(1, 25)]
        row_col = {"pro": None, "cod": None, "qty": None, "lider": None}
        for i, v in enumerate(vals):
            if not row_col["pro"]   and _col_match(v, _MARCADORES_PRO):
                row_col["pro"]   = i + 1
            if not row_col["cod"]   and _col_match(v, _MARCADORES_COD, _EXCLUIR):
                row_col["cod"]   = i + 1
            if not row_col["qty"]   and _col_match(v, _MARCADORES_QTY, _EXCLUIR):
                row_col["qty"]   = i + 1
            if not row_col["lider"] and _col_match(v, _MARCADORES_LIDER):
                row_col["lider"] = i + 1
        if row_col["cod"] and row_col["qty"]:
            col        = row_col
            header_row = r
            break

    if not header_row:
        errs.append("No se detectaron columnas requeridas (CODIGO + CANTIDAD).")
        return rows, errs

    # ODM/DF/LCL del bloque de metadatos
    odm_meta = ""
    if ws.title in ("SIP_MANO_OBRA", "MANO_OBRA"):
        # Programados: preferir LCL/6300 sobre el PDL CDOPPI.
        # Mayor Impacto: si no hay LCL valido, conserva el DF/PDL.
        for grupo in [("LCL",), ("TICKET", "PDL"), ("ODM", "GRAFO")]:
            for r_m in range(1, min(header_row + 1, 20)):
                label = _sin_tildes(str(ws.cell(r_m, 1).value or "").upper().strip())
                if any(lbl in label for lbl in grupo):
                    val = _normalizar_codigo_proyecto(ws.cell(r_m, 2).value)
                    if val:
                        odm_meta = val
                        break
            if odm_meta:
                break
        if odm_meta:
            errs.append(f"[INFO] Modo SIP_MANO_OBRA — DF/LCL: {odm_meta}")

    vistos = set()
    for r in range(header_row + 1, ws.max_row + 1):
        pro     = str(ws.cell(r, col["pro"]).value  or "").strip().upper() if col["pro"]   else ""
        cod     = str(ws.cell(r, col["cod"]).value  or "").strip().upper()
        lid     = str(ws.cell(r, col["lider"]).value or "").strip()        if col["lider"] else ""
        qty_raw = ws.cell(r, col["qty"]).value

        if not cod or qty_raw is None:
            continue
        try:
            qty = float(qty_raw)
        except (TypeError, ValueError):
            continue
        if qty == 0:
            continue

        clave = (pro or odm_meta, cod)
        if clave in vistos:
            errs.append(f"Fila {r}: duplicado ({pro or odm_meta}, {cod}) — ignorado")
            continue
        vistos.add(clave)

        rows.append({
            "fila":     r,
            "pro":      pro,
            "codigo":   cod,
            "cantidad": qty,
            "lider":    lid,
            "_odm":     odm_meta,
        })

    return rows, errs


def extraer_total_mo(wb: openpyxl.Workbook) -> float | None:
    """Lee el total de mano de obra del Excel (fila TOTAL MANO DE OBRA)."""
    for hoja in ("SIP_MANO_OBRA", "MANO_OBRA"):
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        for r in range(1, ws.max_row + 1):
            label = _sin_tildes(str(ws.cell(r, 1).value or "").upper().strip())
            if "TOTAL" in label and ("MANO" in label or "OBRA" in label):
                for c in range(2, 15):
                    val = _parse_float(ws.cell(r, c).value)
                    if val and val > 0:
                        return val
    return None


def extraer_pdl_programado(wb: openpyxl.Workbook) -> str:
    """
    Extrae el valor de Ticket / PDL desde la cabecera del Excel SIP.
    Ejemplo esperado: CDOPPI202678029.
    """
    for hoja in ("SIP_MANO_OBRA", "MANO_OBRA"):
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        for r in range(1, min(ws.max_row + 1, 20)):
            label = _sin_tildes(str(ws.cell(r, 1).value or "").upper().strip())
            if "TICKET" in label or "PDL" in label:
                pdl = _normalizar_pdl(ws.cell(r, 2).value)
                if pdl:
                    return pdl
    return ""


# ── Parser MAT ────────────────────────────────────────────────────────────────
def parsear_mat(wb: openpyxl.Workbook) -> tuple[list, list]:
    """
    Extrae materiales de la hoja CRUCE_LIQUIDADOR (o CRUCE_SERGIO en archivos antiguos).
    Cada row: {fila, nodo, codigo, descripcion, sip_cant, sgs_neto_excel, estado_excel}
    """
    rows, errs = [], []
    hoja = next((h for h in ("CRUCE_LIQUIDADOR", "CRUCE_SERGIO") if h in wb.sheetnames), None)
    if hoja is None:
        errs.append("Hoja 'CRUCE_LIQUIDADOR' no encontrada.")
        return rows, errs

    ws = wb[hoja]
    col = {"nodo": None, "cod": None, "desc": None, "sip": None, "sgs": None, "estado": None}
    header_row = None

    for r in range(1, min(ws.max_row + 1, 15)):
        vals = [_sin_tildes(str(ws.cell(r, c).value or "").upper().strip()) for c in range(1, 20)]
        # Detectar todas las columnas en la misma fila candidata
        tmp = {"nodo": None, "cod": None, "desc": None,
               "sip": None, "sgs": None, "estado": None}
        for i, v in enumerate(vals):
            if v == "NODO"                  and tmp["nodo"]   is None: tmp["nodo"]   = i + 1
            if ("CODIGO" in v or v == "COD") and tmp["cod"]   is None: tmp["cod"]    = i + 1
            if "DESCRIPCION" in v           and tmp["desc"]   is None: tmp["desc"]   = i + 1
            if "SIP" in v and "CANT" in v   and tmp["sip"]    is None: tmp["sip"]    = i + 1
            if "SGS" in v and "NETO" in v   and tmp["sgs"]    is None: tmp["sgs"]    = i + 1
            if v == "ESTADO"                and tmp["estado"] is None: tmp["estado"] = i + 1
        # Solo aceptar esta fila como cabecera si tiene COD + SGS en la misma fila
        if tmp["cod"] and tmp["sgs"]:
            col        = tmp
            header_row = r
            break

    if not header_row:
        errs.append("Cabecera no detectada en la hoja de materiales (busca 'Codigo' y 'SGS Neto').")
        return rows, errs

    current_nodo = "1"
    for r in range(header_row + 1, ws.max_row + 1):
        # Buscar marcador de nodo en cualquier celda de la fila (no solo col1)
        fila_vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, 5)]
        fila_str  = " ".join(fila_vals)
        fila_norm = _sin_tildes(fila_str.upper())

        if "NODO:" in fila_norm:
            m = re.search(r"\d+", fila_str)
            if m:
                current_nodo = m.group()
            continue
        if any(k in fila_norm for k in ("RESUMEN", "VERDE", "ROJO", "AMBAR", "GRIS")):
            continue

        cod = str(ws.cell(r, col["cod"]).value or "").strip()
        if not cod or not any(ch.isdigit() for ch in cod):
            continue

        nodo_v = str(ws.cell(r, col["nodo"]).value or "").strip() if col["nodo"] else current_nodo
        desc   = str(ws.cell(r, col["desc"]).value or "").strip() if col["desc"] else ""
        estado = str(ws.cell(r, col["estado"]).value or "").strip().upper() if col["estado"] else ""

        sip_raw  = ws.cell(r, col["sip"]).value if col["sip"] else None
        sip_cant = None
        if sip_raw is not None and str(sip_raw).strip() not in ("", "-", "?"):
            sip_cant = _parse_float(sip_raw)

        sgs_raw  = ws.cell(r, col["sgs"]).value if col["sgs"] else None
        sgs_neto = _parse_float(sgs_raw) or 0.0

        rows.append({
            "fila":           r,
            "nodo":           nodo_v or current_nodo,
            "codigo":         cod,
            "descripcion":    desc,
            "sip_cant":       sip_cant,
            "sgs_neto_excel": sgs_neto,
            "estado_excel":   estado,
        })

    return rows, errs
