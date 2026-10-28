"""
inyectorfechas/blueprint.py — Inyector de fechas de cierre post-liquidación.

Standalone en /inyectorfechas/. No depende de sesiones del módulo julian ni de liquidador2.

Problema que resuelve:
  El liquidador 2 a veces omite la Fase 3 (fechas + TERMINADO) en algunos nodos.
  El nodo queda liquidado pero con "Fecha de finalización real" vacía y estado
  1.1-REGISTRADO, y aparece sin Fecha Fin en los reportes generales de SGS.

Flujo:
  1. Usuario sube el reporte mensual BALANCE_FINAL_SGS_*.xlsx (módulo Verificador)
  2. Se parsea la hoja "Prestaciones" → fecha de ejecución por PRO a nivel nodo
     (el Consolidado NO sirve: solo trae el primer PRO de cada PDL multi-nodo)
  3. Análisis read-only contra project.dominion: nodos con Fin real vacío,
     fechas que difieren, o solo estado pendiente
  4. Usuario revisa y selecciona → ejecutar
  5. Background: actualizar_fechas_proyecto (7 fechas + estado 1.4, PDL intacto)
     con verificación read-only después de cada escritura (regla «verificar siempre después de escribir»)

La fecha inyectada es la FECHA de las prestaciones del nodo (= Fecha de Avance
de las actividades MO). Si un nodo tiene prestaciones en varias fechas se usa
la última y se marca "multifecha" para revisión.
"""

import io
import logging
import re
import threading
import time
import unicodedata
import uuid
from datetime import datetime, date as _date

import openpyxl
from flask import (Blueprint, render_template, request,
                   redirect, url_for, flash, jsonify)

import os

from sgs_readonly_client import SGSReadOnlyClient
from browser_client import PlaywrightClient

inyectorfechas = Blueprint(
    "inyectorfechas",
    __name__,
    template_folder="../../templates/inyectorfechas",
    url_prefix="/inyectorfechas",
)

log = logging.getLogger("inyectorfechas")

# ── SGS client (propio, solo lectura) ─────────────────────────────────────────
SGS_URL      = os.environ.get("SGS_URL",      "")
SGS_DB       = os.environ.get("SGS_DB",       "")
SGS_LOGIN    = os.environ.get("SGS_LOGIN",    "")
SGS_PASSWORD = os.environ.get("SGS_PASSWORD", "")

_sgs_cache: SGSReadOnlyClient | None = None


def _sgs() -> SGSReadOnlyClient:
    global _sgs_cache
    if _sgs_cache and _sgs_cache.uid:
        try:
            _sgs_cache.search("res.users", [["id", "=", _sgs_cache.uid]], limit=1)
            return _sgs_cache
        except Exception:
            pass
    _sgs_cache = SGSReadOnlyClient(SGS_URL, SGS_DB, SGS_LOGIN, SGS_PASSWORD)
    _sgs_cache.autenticar()
    return _sgs_cache


# Sesión compartida con el browser principal del app
def _obtener_storage_state():
    try:
        import app as _app
        return _app._obtener_storage_state_activo(timeout_s=5.0)
    except Exception:
        return None


# ── Constantes (mismas convenciones que cierre rápido / Fase 3 liquidador2) ──
CIERRE_FECHAS = [
    ("date_assignment",    "Asignación"),
    ("date_expected",      "Inicio prevista"),
    ("date_real",          "Inicio real"),
    ("certification_date", "Certificación"),
    ("date_end_expected",  "Fin prevista"),
    ("date_end_real",      "Fin real"),
    ("last_status_change", "Últ. cambio estado"),
]
ESTADO_TERMINADO = "1.4-TERMINADO/EN BALANCE"

INYECTOR_MAX_WORKERS = 4   # igual que TERMINADO_MAX_WORKERS de liquidador2 Fase 3

_RE_PRO = re.compile(r"^PRO-\d+$")

# ── Sesiones en memoria ───────────────────────────────────────────────────────
_sesiones: dict = {}


# ═══════════════════════════════════════════════════════════════════════════════
# Helpers de fechas
# ═══════════════════════════════════════════════════════════════════════════════

def _fecha_iso(valor) -> str:
    """Cualquier cosa (str/datetime/False) → 'YYYY-MM-DD' o ''."""
    if valor is None or valor is False:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, _date):
        return valor.strftime("%Y-%m-%d")
    s = str(valor or "").strip()
    if not s:
        return ""
    s = s.replace("T", " ")
    candidatos = [
        ("%Y-%m-%d %H:%M:%S", 19),
        ("%Y-%m-%d %H:%M", 16),
        ("%Y-%m-%d", 10),
        ("%d-%m-%Y", 10),
        ("%d/%m/%Y", 10),
    ]
    for fmt, ancho in candidatos:
        try:
            return datetime.strptime(s[:ancho], fmt).strftime("%Y-%m-%d")
        except Exception:
            continue
    return ""


def _fecha_ui(valor) -> str:
    """'YYYY-MM-DD' (u otro) → 'DD-MM-YYYY' para mostrar y para Playwright."""
    iso = _fecha_iso(valor)
    if not iso:
        return "-"
    return datetime.strptime(iso, "%Y-%m-%d").strftime("%d-%m-%Y")


def _norm_header(txt) -> str:
    """Normaliza encabezado: sin acentos, mayúsculas, solo A-Z0-9."""
    s = unicodedata.normalize("NFKD", str(txt or ""))
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"[^A-Z0-9]", "", s.upper())


# ═══════════════════════════════════════════════════════════════════════════════
# Paso 1 — Parser del reporte BALANCE_FINAL (hoja Prestaciones)
# ═══════════════════════════════════════════════════════════════════════════════

# encabezado normalizado → clave interna
_COLS_PRESTACIONES = {
    "FECHA":          "fecha",
    "CODSGS":         "pro",         # "# COD SGS"
    "PDLOT":          "pdl",         # "PDL / OT"
    "NOMBREPROYECTO": "nombre",
    "PROCESO":        "proceso",
    "SUPERVISOR":     "supervisor",
    "LIDER":          "lider",       # "LÍDER"
    "PRESTACION":     "monto",       # "$ PRESTACIÓN"
    "LIQUIDADOPOR":   "liquidador",
    "ESTADO":         "estado",
}


def _parsear_reporte(contenido: bytes) -> dict:
    """
    Parsea la hoja Prestaciones del BALANCE_FINAL.
    Retorna {"grupos": {pro_code: {...}}, "warnings": [...], "rango": (min,max)}.
    """
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    try:
        ws = None
        for hoja in wb.worksheets:
            if _norm_header(hoja.title) == "PRESTACIONES":
                ws = hoja
                break
        if ws is None:
            # Fallback: hoja cuyo encabezado tenga FECHA y # COD SGS
            for hoja in wb.worksheets:
                primera = next(hoja.iter_rows(max_row=1, values_only=True), None)
                headers = {_norm_header(c) for c in (primera or [])}
                if "FECHA" in headers and "CODSGS" in headers:
                    ws = hoja
                    break
        if ws is None:
            raise ValueError(
                "El archivo no tiene hoja 'Prestaciones'. "
                "¿Es realmente un BALANCE_FINAL_SGS del verificador?")

        filas = ws.iter_rows(values_only=True)
        header = next(filas, None) or []
        idx = {}   # clave interna → índice de columna
        for i, celda in enumerate(header):
            clave = _COLS_PRESTACIONES.get(_norm_header(celda))
            if clave and clave not in idx:
                idx[clave] = i

        faltan = [k for k in ("fecha", "pro") if k not in idx]
        if faltan:
            raise ValueError(f"Hoja Prestaciones sin columnas requeridas: {faltan}")

        def _celda(row, clave, default=""):
            i = idx.get(clave)
            if i is None or i >= len(row):
                return default
            v = row[i]
            return default if v is None else v

        grupos: dict = {}
        warnings: list = []
        codigos_raros: dict = {}
        fechas_global: list = []
        n_filas = 0

        for row in filas:
            if not row or all(v is None for v in row):
                continue
            pro_raw   = str(_celda(row, "pro")).strip().upper()
            fecha_iso = _fecha_iso(_celda(row, "fecha"))
            if not pro_raw and not fecha_iso:
                continue
            n_filas += 1

            if not _RE_PRO.match(pro_raw):
                codigos_raros.setdefault(pro_raw or "(vacío)", 0)
                codigos_raros[pro_raw or "(vacío)"] += 1
                continue
            if not fecha_iso:
                warnings.append(f"{pro_raw}: prestación sin fecha válida — fila omitida.")
                continue

            g = grupos.setdefault(pro_raw, {
                "pro_code":     pro_raw,
                "fechas":       set(),
                "pdls":         set(),
                "nombres":      set(),
                "procesos":     set(),
                "supervisores": set(),
                "lideres":      set(),
                "liquidadores": set(),
                "estados":      set(),
                "monto":        0.0,
                "n_prest":      0,
            })
            g["fechas"].add(fecha_iso)
            fechas_global.append(fecha_iso)
            for clave, destino in [("pdl", "pdls"), ("nombre", "nombres"),
                                   ("proceso", "procesos"), ("supervisor", "supervisores"),
                                   ("lider", "lideres"), ("liquidador", "liquidadores"),
                                   ("estado", "estados")]:
                v = str(_celda(row, clave)).strip()
                if v:
                    g[destino].add(v)
            try:
                g["monto"] += float(_celda(row, "monto", 0) or 0)
            except (TypeError, ValueError):
                pass
            g["n_prest"] += 1

        for cod, n in sorted(codigos_raros.items()):
            warnings.append(
                f"Código '{cod}' no es un PRO válido ({n} prestación(es)) — omitido. "
                f"Revísalo manualmente en el reporte.")

        if not grupos:
            raise ValueError("No se encontró ninguna prestación con código PRO en el reporte.")

        # sets → estructuras serializables y fecha objetivo
        for g in grupos.values():
            fechas = sorted(g.pop("fechas"))
            g["fechas"]        = fechas
            g["fecha_objetivo"] = fechas[-1]          # última prestación
            g["multifecha"]    = len(fechas) > 1
            for k in ("pdls", "nombres", "procesos", "supervisores",
                      "lideres", "liquidadores", "estados"):
                g[k] = sorted(g.pop(k))

        rango = (min(fechas_global), max(fechas_global)) if fechas_global else ("", "")
        return {"grupos": grupos, "warnings": warnings,
                "rango": rango, "n_filas": n_filas}
    finally:
        wb.close()


# ═══════════════════════════════════════════════════════════════════════════════
# Paso 2 — Análisis read-only contra SGS
# ═══════════════════════════════════════════════════════════════════════════════

def _armar_analisis(grupos: dict) -> tuple[list, dict]:
    """
    Cruza los PRO del reporte contra project.dominion (SOLO LECTURA).
    Clasifica cada nodo:
      sin_fecha    → Fin real vacío (el bug) — preseleccionado
      difiere      → Fin real existe pero != fecha del reporte — selección manual
      solo_estado  → fechas OK pero estado != 1.4 — selección manual
      ok           → nada que hacer
      no_encontrado→ PRO no existe en SGS
    """
    c = _sgs()
    codes = sorted(grupos.keys())
    fields = (["id", "code", "name", "nodo", "order_number", "status", "pdl"]
              + [campo for campo, _ in CIERRE_FECHAS])
    registros = c.search_read(
        "project.dominion", [["code", "in", codes]], fields,
        limit=len(codes) + 10)
    por_code = {str(r.get("code") or "").strip().upper(): r for r in registros}

    rows = []
    meta = {"total": 0, "sin_fecha": 0, "difiere": 0,
            "solo_estado": 0, "ok": 0, "no_encontrado": 0}

    for pro_code in codes:
        g   = grupos[pro_code]
        reg = por_code.get(pro_code)
        objetivo_iso = g["fecha_objetivo"]

        base = {
            "pro_code":     pro_code,
            "nombre_rep":   ", ".join(g["nombres"])[:80] or "-",
            "pdl_rep":      ", ".join(g["pdls"]) or "-",
            "lider_rep":    ", ".join(g["lideres"]) or "-",
            "proceso_rep":  ", ".join(g["procesos"]) or "-",
            "n_prest":      g["n_prest"],
            "monto":        g["monto"],
            "multifecha":   g["multifecha"],
            "fechas_rep":   g["fechas"],
            "fecha_objetivo_iso": objetivo_iso,
            "fecha_objetivo_ui":  _fecha_ui(objetivo_iso),
        }

        if not reg:
            rows.append({**base, "pid": None, "nombre_sgs": "-", "nodo_sgs": "-",
                         "estado_actual": "-", "fin_real_ui": "-",
                         "fechas_detalle": [], "n_vacias": 0,
                         "categoria": "no_encontrado", "preseleccionado": False,
                         "estado_destino": ""})
            meta["no_encontrado"] += 1
            continue

        estado_actual = ""
        st = reg.get("status")
        if isinstance(st, (list, tuple)):
            estado_actual = str(st[1] if len(st) > 1 else st[0] or "").strip()
        else:
            estado_actual = str(st or "").strip()
        status_ok = estado_actual == ESTADO_TERMINADO

        fechas_detalle, n_vacias, n_coinciden = [], 0, 0
        for campo, label in CIERRE_FECHAS:
            actual_iso = _fecha_iso(reg.get(campo))
            vacia      = not actual_iso
            coincide   = actual_iso == objetivo_iso
            if vacia:
                n_vacias += 1
            if coincide:
                n_coinciden += 1
            fechas_detalle.append({"campo": campo, "label": label,
                                   "actual_ui": _fecha_ui(reg.get(campo)),
                                   "vacia": vacia, "coincide": coincide})

        fin_real_iso = _fecha_iso(reg.get("date_end_real"))

        if not fin_real_iso:
            categoria = "sin_fecha"
        elif fin_real_iso != objetivo_iso:
            categoria = "difiere"
        elif not status_ok:
            categoria = "solo_estado"
        else:
            categoria = "ok"
        meta[categoria] += 1

        rows.append({
            **base,
            "pid":            reg["id"],
            "nombre_sgs":     str(reg.get("name") or "").strip() or "-",
            "nodo_sgs":       str(reg.get("nodo") or "").strip() or "-",
            "estado_actual":  estado_actual or "-",
            "status_ok":      status_ok,
            "fin_real_ui":    _fecha_ui(reg.get("date_end_real")),
            "fechas_detalle": fechas_detalle,
            "n_vacias":       n_vacias,
            "n_coinciden":    n_coinciden,
            "categoria":      categoria,
            "preseleccionado": categoria == "sin_fecha",
            "estado_destino": "" if status_ok else ESTADO_TERMINADO,
        })

    orden = {"sin_fecha": 0, "difiere": 1, "solo_estado": 2,
             "no_encontrado": 3, "ok": 4}
    rows.sort(key=lambda r: (orden.get(r["categoria"], 9), r["pro_code"]))
    meta["total"] = len(rows)
    return rows, meta


# ═══════════════════════════════════════════════════════════════════════════════
# Rutas
# ═══════════════════════════════════════════════════════════════════════════════

@inyectorfechas.route("/")
def inicio():
    return render_template("inyectorfechas/inicio.html")


@inyectorfechas.route("/subir", methods=["POST"])
def subir():
    f = request.files.get("reporte")
    if not f or not f.filename:
        flash("Selecciona el archivo BALANCE_FINAL_SGS_*.xlsx")
        return redirect(url_for("inyectorfechas.inicio"))
    try:
        parsed = _parsear_reporte(f.read())
    except Exception as ex:
        log.error("[INYECTOR] Error parseando %s: %s", f.filename, ex)
        flash(f"No se pudo leer el reporte: {ex}")
        return redirect(url_for("inyectorfechas.inicio"))

    sid = str(uuid.uuid4())
    _sesiones[sid] = {
        "creado_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "filename":  f.filename,
        "grupos":    parsed["grupos"],
        "warnings":  parsed["warnings"],
        "rango":     parsed["rango"],
        "n_filas":   parsed["n_filas"],
        "job":       None,
        "resultados": [],
    }
    log.info("[INYECTOR] Sesión %s — %s: %d nodos, %d filas",
             sid[:8], f.filename, len(parsed["grupos"]), parsed["n_filas"])
    return redirect(url_for("inyectorfechas.analizar", sid=sid))


@inyectorfechas.route("/analizar/<sid>")
def analizar(sid):
    s = _sesiones.get(sid)
    if not s:
        flash("Sesión no encontrada. Sube el reporte de nuevo.")
        return redirect(url_for("inyectorfechas.inicio"))
    try:
        rows, meta = _armar_analisis(s["grupos"])
    except Exception as ex:
        log.error("[INYECTOR] Error analizando sid=%s: %s", sid[:8], ex, exc_info=True)
        flash(f"Error consultando SGS: {ex}")
        return redirect(url_for("inyectorfechas.inicio"))
    s["analisis"] = {r["pro_code"]: r for r in rows}
    return render_template("inyectorfechas/analizar.html",
                           sid=sid, s=s, rows=rows, meta=meta)


@inyectorfechas.route("/ejecutar/<sid>", methods=["POST"])
def ejecutar(sid):
    s = _sesiones.get(sid)
    if not s:
        flash("Sesión no encontrada.")
        return redirect(url_for("inyectorfechas.inicio"))
    if s.get("job") and s["job"].get("estado") == "running":
        flash("Ya hay una inyección en curso para esta sesión.")
        return redirect(url_for("inyectorfechas.esperar", sid=sid))

    analisis = s.get("analisis") or {}
    seleccion = request.form.getlist("pro")
    plan = []
    for pro_code in seleccion:
        r = analisis.get(pro_code)
        if not r or not r.get("pid") or r["categoria"] == "ok":
            continue
        plan.append({
            "pid":            r["pid"],
            "pro_code":       r["pro_code"],
            "nombre":         r["nombre_sgs"],
            "fecha_iso":      r["fecha_objetivo_iso"],
            "fecha_ui":       r["fecha_objetivo_ui"],
            "estado_destino": r["estado_destino"],
            "categoria":      r["categoria"],
        })

    if not plan:
        flash("No seleccionaste ningún nodo inyectable.")
        return redirect(url_for("inyectorfechas.analizar", sid=sid))

    s["plan"] = plan
    s["resultados"] = []
    s["job"] = {
        "estado":   "running",
        "inicio":   datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "logs":     [],
        "progreso": 0,
        "total":    len(plan),
        "ok_count": 0,
        "errores":  0,
    }
    threading.Thread(target=_bg_inyectar, args=(sid,), daemon=True,
                     name=f"inyector-{sid[:8]}").start()
    return redirect(url_for("inyectorfechas.esperar", sid=sid))


@inyectorfechas.route("/esperar/<sid>")
def esperar(sid):
    s = _sesiones.get(sid)
    if not s:
        flash("Sesión no encontrada.")
        return redirect(url_for("inyectorfechas.inicio"))
    return render_template("inyectorfechas/esperar.html", sid=sid, s=s)


@inyectorfechas.route("/api/job_estado/<sid>")
def api_job_estado(sid):
    s = _sesiones.get(sid)
    if not s:
        return jsonify({"estado": "not_found"}), 404
    job = s.get("job") or {}
    return jsonify({
        "estado":     job.get("estado", "idle"),
        "progreso":   job.get("progreso", 0),
        "total":      job.get("total", 0),
        "ok_count":   job.get("ok_count", 0),
        "errores":    job.get("errores", 0),
        "logs":       job.get("logs", []),
        "resultados": s.get("resultados", []),
    })


# ═══════════════════════════════════════════════════════════════════════════════
# Background — inyección con verificación post-escritura
# ═══════════════════════════════════════════════════════════════════════════════

def _bg_inyectar(sid: str) -> None:
    s   = _sesiones.get(sid)
    job = s.get("job") if s else None
    if not s or not job:
        return

    plan       = s.get("plan") or []
    resultados = [None] * len(plan)
    job_lock   = threading.Lock()

    def _log(msg: str) -> None:
        ts = datetime.now().strftime("%H:%M:%S")
        with job_lock:
            job["logs"].append(f"[{ts}] {msg}")
        log.info("[INYECTOR BG] %s", msg)

    # ── Login único / sesión compartida ──────────────────────────────────────
    storage_state = _obtener_storage_state()
    if not storage_state:
        try:
            _log("Sin sesión activa — iniciando login…")
            pw_setup = PlaywrightClient(headless=False, slow_mo=100)
            pw_setup.start()
            pw_setup.login()
            storage_state = pw_setup._page.context.storage_state()
            pw_setup.close()
            _log("Login OK — sesión capturada")
        except Exception as ex:
            _log(f"✗ Error en login: {ex}")
            job["estado"] = "listo"
            return
    else:
        _log("Sesión existente encontrada")

    # ── Workers round-robin (patrón Fase 3 liquidador2) ──────────────────────
    items    = list(enumerate(plan))
    n_workers = min(INYECTOR_MAX_WORKERS, len(items))
    chunks    = [items[i::n_workers] for i in range(n_workers) if items[i::n_workers]]
    _log(f"Inyectando fechas en {len(items)} nodo(s) con {len(chunks)} worker(s)…")

    def _worker(w_id: int, mis_items: list) -> None:
        prefix = f"[W{w_id}]"
        pw = PlaywrightClient(headless=False, slow_mo=200)
        try:
            pw.start(storage_state=storage_state)
            c = _sgs()
            for idx, item in mis_items:
                pro   = item["pro_code"]
                fecha = item["fecha_ui"]
                res_row = {
                    "pro_code":   pro,
                    "nombre":     item["nombre"],
                    "fecha":      fecha,
                    "categoria":  item["categoria"],
                    "ok":         False,
                    "verificado": False,
                    "msg":        "",
                    "error":      None,
                }
                try:
                    _log(f"{prefix} {pro} → {fecha}"
                         + (f" + estado {ESTADO_TERMINADO}" if item["estado_destino"] else ""))
                    res = pw.actualizar_fechas_proyecto(
                        pid=item["pid"],
                        fecha=fecha,
                        pro_code=pro,
                        estado_destino=item["estado_destino"],
                    )
                    res_row["ok"]  = bool(res.get("ok"))
                    res_row["msg"] = res.get("msg", "")
                    if not res_row["ok"]:
                        res_row["error"] = res.get("error") or "actualizar_fechas_proyecto falló"
                        _log(f"{prefix} ✗ {pro}: {res_row['error']}")
                    else:
                        # Regla «verificar siempre después de escribir»: lectura fresca
                        try:
                            time.sleep(1.0)
                            fresco = c.search_read(
                                "project.dominion", [["id", "=", item["pid"]]],
                                ["id", "date_end_real", "status"], limit=1)
                            reg = fresco[0] if fresco else {}
                            fin_ok = _fecha_iso(reg.get("date_end_real")) == item["fecha_iso"]
                            res_row["verificado"] = fin_ok
                            if fin_ok:
                                _log(f"{prefix} ✓ {pro} verificado — Fin real = {fecha}")
                            else:
                                res_row["msg"] += " | ⚠ Fin real no coincide tras guardar — REVISAR"
                                _log(f"{prefix} ⚠ {pro}: guardado pero Fin real no coincide — REVISAR")
                        except Exception as ex_v:
                            res_row["msg"] += f" | ⚠ No se pudo verificar: {ex_v}"
                            _log(f"{prefix} ⚠ {pro}: escrito pero sin verificación ({ex_v})")
                except Exception as ex:
                    res_row["error"] = str(ex)
                    _log(f"{prefix} ✗ {pro}: {ex}")
                    log.error("[INYECTOR W%d] %s: %s", w_id, pro, ex, exc_info=True)
                finally:
                    with job_lock:
                        resultados[idx] = res_row
                        if res_row["ok"]:
                            job["ok_count"] += 1
                        else:
                            job["errores"] += 1
                        s["resultados"] = [r for r in resultados if r is not None]
                        job["progreso"] = int(
                            sum(1 for r in resultados if r is not None)
                            / max(len(plan), 1) * 100)
        except Exception as ex:
            _log(f"{prefix} ✗ Error crítico del worker: {ex}")
            log.error("[INYECTOR W%d] crítico: %s", w_id, ex, exc_info=True)
            with job_lock:
                for idx, item in mis_items:
                    if resultados[idx] is None:
                        resultados[idx] = {
                            "pro_code": item["pro_code"], "nombre": item["nombre"],
                            "fecha": item["fecha_ui"], "categoria": item["categoria"],
                            "ok": False, "verificado": False, "msg": "",
                            "error": f"Worker abortado: {ex}",
                        }
                        job["errores"] += 1
                s["resultados"] = [r for r in resultados if r is not None]
        finally:
            try:
                pw.close()
            except Exception:
                pass

    threads = []
    for w_id, chunk in enumerate(chunks):
        t = threading.Thread(target=_worker, args=(w_id, chunk),
                             daemon=True, name=f"inyector-w{w_id}")
        threads.append(t)
        t.start()
    for t in threads:
        t.join()

    s["resultados"] = [r or {"ok": False, "verificado": False,
                             "error": "Sin resultado — worker interrumpido"}
                       for r in resultados]
    job["estado"]   = "listo"
    job["progreso"] = 100
    ok_v  = sum(1 for r in s["resultados"] if r.get("ok") and r.get("verificado"))
    ok_nv = sum(1 for r in s["resultados"] if r.get("ok") and not r.get("verificado"))
    err   = sum(1 for r in s["resultados"] if not r.get("ok"))
    _log(f"Finalizado — {ok_v} verificados / {ok_nv} escritos sin verificar / {err} errores.")
