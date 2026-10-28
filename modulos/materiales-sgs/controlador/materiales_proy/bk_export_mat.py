"""
export_mat.py — Exportación a Excel del módulo MATERIALES (Zona VII)
===================================================================
Genera un .xlsx con 4 hojas a partir de `sgs_materiales.analizar`:

  1. "Resumen ejecutivo"     — KPIs + totales por supervisor y por zona.
  2. "Proyectos"             — una fila por familia (filtrable) con zona y valor pendiente.
  3. "Materiales pendientes" — jerárquica PRO padre → nodos → material, con líder,
                               antigüedad real de almacén, costo y valor pendiente.
  4. "Por supervisor"        — familias agrupadas por supervisor.
"""

import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from materiales_proy.sgs_materiales import analizar

# ── Paleta ───────────────────────────────────────────────────────────────────
C_HDR_BG, C_HDR_FG     = "1F3864", "FFFFFF"
C_PADRE_BG, C_PADRE_FG = "1A237E", "FFFFFF"
C_NODO_BG              = "D6E4F0"
C_OK, C_WARN, C_MUTED  = "1B5E20", "B26A00", "777777"
C_ANTIG_HI            = "FDECEA"
C_KPI_BG              = "0B5394"
C_BORDE               = "BDC3C7"
_THIN   = Side(style="thin", color=C_BORDE)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _clasif_txt(c):
    return {"completo": "Completo", "pendiente": "Pendiente",
            "sin_materiales": "Sin materiales"}.get(c, c)


def _tipo_txt(t):
    return {"programado": "6300", "mayor_impacto": "DF", "otro": "Otro"}.get(t, t)


def _antig_txt(d):
    return "" if d is None else f"{d} día" + ("" if d == 1 else "s")


def _num(q):
    q = float(q or 0)
    return str(int(q)) if abs(q - round(q)) < 0.005 else f"{q:g}"


def _just_bandejas(p):
    desp = p.get("bandejas_despacho", []) or []
    dev  = p.get("bandejas_devolucion", []) or []
    partes = []
    if desp:
        partes.append("Despacho: " + ", ".join(f"{b['nombre']} (+{_num(b['qty'])})" for b in desp))
    if dev:
        partes.append("Devolución: " + ", ".join(f"{b['nombre']} (-{_num(b['qty'])})" for b in dev))
    if not partes:
        return ""
    return "  |  ".join(partes) + f"  → neto {_num(p.get('neto', 0))}"


def _set(ws, row, col, value, *, bold=False, fg=None, bg=None,
         align="left", num_fmt=None, wrap=False, size=10):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold, color=fg or "000000", size=size)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)
    if num_fmt:
        cell.number_format = num_fmt
    cell.border = _BORDER
    return cell


def _lcl_txt(f):
    lcl = f.get("lcl")
    return lcl if lcl and lcl != "N/A" else "—"


def _zona_label(zid):
    import config as cfg
    for z in cfg.MACRO_ZONAS:
        if z["id"] == zid:
            return z["label"]
    return zid


# ── Hoja 1: Resumen ejecutivo ─────────────────────────────────────────────────
def _hoja_resumen(wb, d, titulo):
    ws = wb.create_sheet("Resumen ejecutivo")
    ws.column_dimensions["A"].width = 28
    for col in "BCDE":
        ws.column_dimensions[col].width = 20
    _set(ws, 1, 1, titulo, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, size=12)
    ws.merge_cells("A1:E1")
    ws.row_dimensions[1].height = 22

    kpis = [
        ("Proyectos", d.get("n_total", 0)),
        ("Con material pendiente", d.get("n_pendiente", 0)),
        ("Completos", d.get("n_completo", 0)),
        ("MO liquidada total", d.get("mo_total", 0)),
        ("Valor pendiente total", d.get("valor_pendiente_total", 0)),
    ]
    r = 3
    for i, (k, v) in enumerate(kpis):
        col = 1 + i
        _set(ws, r, col, k, bold=True, fg="FFFFFF", bg=C_KPI_BG, align="center", wrap=True)
        fmt = '"$"#,##0' if "MO" in k or "Valor" in k else None
        _set(ws, r + 1, col, v, bold=True, align="center", num_fmt=fmt, size=14)
    ws.row_dimensions[r].height = 28
    ws.row_dimensions[r + 1].height = 24

    # Tabla por supervisor
    r = 7
    _set(ws, r, 1, "Por supervisor", bold=True, fg=C_HDR_FG, bg=C_HDR_BG)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 1
    for i, h in enumerate(["Supervisor", "Proyectos", "Pendientes", "MO liquidada", "Valor pendiente"], 1):
        _set(ws, r, i, h, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, align="center")
    r += 1
    for s in d.get("agregados", {}).get("por_supervisor", []):
        _set(ws, r, 1, s["nombre"])
        _set(ws, r, 2, s["n"], align="center")
        _set(ws, r, 3, s["n_pend"], align="center", fg=C_WARN if s["n_pend"] else "000000")
        _set(ws, r, 4, round(s["mo"], 2), align="right", num_fmt='"$"#,##0')
        _set(ws, r, 5, round(s["valor_pend"], 2), align="right", num_fmt='"$"#,##0')
        r += 1

    # Tabla por zona
    r += 1
    _set(ws, r, 1, "Por zona", bold=True, fg=C_HDR_FG, bg=C_HDR_BG)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 1
    for i, h in enumerate(["Zona", "Proyectos", "Pendientes", "MO liquidada", "Valor pendiente"], 1):
        _set(ws, r, i, h, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, align="center")
    r += 1
    for z in d.get("agregados", {}).get("por_zona", []):
        _set(ws, r, 1, z["nombre"])
        _set(ws, r, 2, z["n"], align="center")
        _set(ws, r, 3, z["n_pend"], align="center", fg=C_WARN if z["n_pend"] else "000000")
        _set(ws, r, 4, round(z["mo"], 2), align="right", num_fmt='"$"#,##0')
        _set(ws, r, 5, round(z["valor_pend"], 2), align="right", num_fmt='"$"#,##0')
        r += 1
    return ws


# ── Hoja 2: Proyectos ─────────────────────────────────────────────────────────
_COLS_PROY = [
    ("Tipo", 8), ("Estado mat.", 13), ("Zona", 14), ("PDL / OT", 26), ("LCL", 14),
    ("Cód. padre", 14), ("Proyecto padre", 38), ("Supervisor", 18),
    ("MO liquidada", 14), ("Valor pendiente", 15), ("Fecha MO", 12),
    ("Antig. almacén", 14), ("Líder a liquidar", 32), ("Nodos pend.", 11),
]


def _hoja_proyectos(wb, familias, titulo):
    ws = wb.create_sheet("Proyectos")
    _set(ws, 1, 1, titulo, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_COLS_PROY))
    ws.row_dimensions[1].height = 22
    r = 3
    for i, (name, w) in enumerate(_COLS_PROY, 1):
        _set(ws, r, i, name, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, align="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{r + 1}"
    ws.auto_filter.ref = f"A{r}:{get_column_letter(len(_COLS_PROY))}{r}"
    r += 1
    for f in familias:
        clasif = f.get("clasificacion", "")
        col_estado = C_WARN if clasif == "pendiente" else C_OK if clasif == "completo" else C_MUTED
        antig = f.get("antiguedad_almacen_dias")
        if antig is None:
            antig = f.get("antiguedad_aprob_dias")
        bg_antig = C_ANTIG_HI if (antig is not None and antig > 30) else None
        _set(ws, r, 1, _tipo_txt(f.get("tipo", "")), align="center")
        _set(ws, r, 2, _clasif_txt(clasif), bold=True, fg=col_estado, align="center")
        _set(ws, r, 3, _zona_label(f.get("zona")), align="center")
        _set(ws, r, 4, f.get("pdl_code", ""), wrap=True)
        _set(ws, r, 5, _lcl_txt(f), align="center")
        _set(ws, r, 6, f.get("padre_code", ""))
        _set(ws, r, 7, f.get("padre_name", ""), wrap=True)
        _set(ws, r, 8, f.get("supervisor_full") or f.get("supervisor", ""))
        _set(ws, r, 9, f.get("mo_value", 0) or 0, align="right", num_fmt='"$"#,##0')
        _set(ws, r, 10, f.get("valor_pendiente", 0) or 0, align="right", num_fmt='"$"#,##0')
        _set(ws, r, 11, f.get("fecha_mo", ""), align="center")
        _set(ws, r, 12, _antig_txt(antig), align="center", bg=bg_antig)
        _set(ws, r, 13, f.get("lider_liquidar", ""), wrap=True)
        _set(ws, r, 14, f.get("n_pendientes", 0), align="center")
        r += 1
    return ws


# ── Hoja 3: Materiales pendientes (jerárquica) ────────────────────────────────
_COLS_MAT = [
    ("Nivel", 12), ("PDL / Nodo", 24), ("Código", 16), ("Material", 14),
    ("Nombre material", 40), ("Cant. pendiente", 14), ("Costo unit.", 13),
    ("Valor pendiente", 15), ("Líder a liquidar", 30),
    ("Fecha almacén", 14), ("Antig. almacén", 13), ("Supervisor", 18),
    ("Bandejas (despacho / devolución)", 52),
]


def _hoja_materiales(wb, familias, titulo):
    ws = wb.create_sheet("Materiales pendientes")
    ncol = len(_COLS_MAT)
    _set(ws, 1, 1, titulo, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.row_dimensions[1].height = 22
    r = 3
    for i, (name, w) in enumerate(_COLS_MAT, 1):
        _set(ws, r, i, name, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, align="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{r + 1}"
    r += 1

    pend_fams = [f for f in familias if f.get("clasificacion") == "pendiente"]
    pend_fams.sort(key=lambda f: (f.get("zona", ""), -(f.get("valor_pendiente") or 0)))
    if not pend_fams:
        _set(ws, r, 1, "No hay materiales pendientes en el periodo.", fg=C_MUTED)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
        return ws

    for f in pend_fams:
        _lcl = _lcl_txt(f)
        _pdl_lbl = f"{f.get('pdl_code', '')}" + (f"  ·  LCL {_lcl}" if _lcl != "—" else "")
        _set(ws, r, 1, _zona_label(f.get("zona")), bold=True, fg=C_PADRE_FG, bg=C_PADRE_BG)
        _set(ws, r, 2, _pdl_lbl, bold=True, fg=C_PADRE_FG, bg=C_PADRE_BG)
        _set(ws, r, 3, f.get("padre_code", ""), bold=True, fg=C_PADRE_FG, bg=C_PADRE_BG)
        _set(ws, r, 4, f.get("padre_name", ""), fg=C_PADRE_FG, bg=C_PADRE_BG)
        for c in range(5, 8):
            _set(ws, r, c, "", bg=C_PADRE_BG)
        _set(ws, r, 8, f.get("valor_pendiente", 0) or 0, bold=True, fg=C_PADRE_FG,
             bg=C_PADRE_BG, align="right", num_fmt='"$"#,##0')
        for c in range(9, 12):
            _set(ws, r, c, "", bg=C_PADRE_BG)
        _set(ws, r, 12, f.get("supervisor_full") or f.get("supervisor", ""), fg=C_PADRE_FG, bg=C_PADRE_BG)
        _set(ws, r, 13, "", bg=C_PADRE_BG)
        padre_row = r
        r += 1
        grupo_inicio = r

        nodos_pend = [n for n in f.get("nodos_mat", [])
                      if n.get("estado_mat") == "pendiente" and n.get("pendientes")]
        nodos_pend.sort(key=lambda n: (str(n.get("nodo") or ""), n.get("pro_code", "")))
        for n in nodos_pend:
            nodo_lbl = f"N{n['nodo']}" if n.get("nodo") else (n.get("pro_code") or "—")
            _set(ws, r, 1, "  └ Nodo", fg="333333", bg=C_NODO_BG)
            _set(ws, r, 2, nodo_lbl, bold=True, bg=C_NODO_BG)
            _set(ws, r, 3, n.get("pro_code", ""), bg=C_NODO_BG)
            for c in range(4, 14):
                _set(ws, r, c, "", bg=C_NODO_BG)
            r += 1
            for p in n.get("pendientes", []):
                antig = p.get("antiguedad_almacen_dias")
                if antig is None:
                    antig = p.get("antiguedad_aprob_dias")
                bg_antig = C_ANTIG_HI if (antig is not None and antig > 30) else None
                _set(ws, r, 4, p.get("codigo", ""))
                _set(ws, r, 5, p.get("nombre", ""), wrap=True)
                _set(ws, r, 6, p.get("pendiente", 0), align="right", num_fmt="#,##0.##")
                _set(ws, r, 7, p.get("costo_unit", 0) or 0, align="right", num_fmt='"$"#,##0')
                _set(ws, r, 8, p.get("valor_pendiente", 0) or 0, align="right", num_fmt='"$"#,##0')
                _set(ws, r, 9, p.get("lider", ""), wrap=True)
                _set(ws, r, 10, p.get("fecha_almacen", "") or p.get("fecha_aprobacion", ""), align="center")
                _set(ws, r, 11, _antig_txt(antig), align="center", bg=bg_antig)
                _set(ws, r, 12, "")
                _set(ws, r, 13, _just_bandejas(p), wrap=True, size=9)
                r += 1
        if r > grupo_inicio:
            for rr in range(grupo_inicio, r):
                ws.row_dimensions[rr].outline_level = 1
        ws.row_dimensions[padre_row].height = 18
    return ws


# ── Hoja 4: Por supervisor ────────────────────────────────────────────────────
def _hoja_por_supervisor(wb, familias, titulo):
    ws = wb.create_sheet("Por supervisor")
    cols = [("Supervisor / Proyecto", 40), ("Zona", 14), ("Tipo", 8), ("Estado", 13),
            ("PDL / OT", 24), ("MO liquidada", 14), ("Valor pendiente", 15), ("Antig.", 10)]
    _set(ws, 1, 1, titulo, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    r = 3
    for i, (name, w) in enumerate(cols, 1):
        _set(ws, r, i, name, bold=True, fg=C_HDR_FG, bg=C_HDR_BG, align="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{r + 1}"
    r += 1

    grupos = {}
    for f in familias:
        grupos.setdefault(f.get("supervisor_full") or f.get("supervisor") or "?", []).append(f)
    for sup in sorted(grupos):
        fams = grupos[sup]
        mo = sum(x.get("mo_value", 0) for x in fams)
        vp = sum(x.get("valor_pendiente", 0) for x in fams)
        _set(ws, r, 1, f"{sup}  ({len(fams)} proy.)", bold=True, fg=C_PADRE_FG, bg=C_PADRE_BG)
        for c in range(2, 6):
            _set(ws, r, c, "", bg=C_PADRE_BG)
        _set(ws, r, 6, round(mo, 2), bold=True, fg=C_PADRE_FG, bg=C_PADRE_BG, align="right", num_fmt='"$"#,##0')
        _set(ws, r, 7, round(vp, 2), bold=True, fg=C_PADRE_FG, bg=C_PADRE_BG, align="right", num_fmt='"$"#,##0')
        _set(ws, r, 8, "", bg=C_PADRE_BG)
        r += 1
        for f in sorted(fams, key=lambda x: -(x.get("valor_pendiente") or 0)):
            clasif = f.get("clasificacion", "")
            col_estado = C_WARN if clasif == "pendiente" else C_OK if clasif == "completo" else C_MUTED
            antig = f.get("antiguedad_almacen_dias")
            if antig is None:
                antig = f.get("antiguedad_aprob_dias")
            _set(ws, r, 1, "   " + (f.get("padre_name", "") or f.get("padre_code", "")), wrap=True)
            _set(ws, r, 2, _zona_label(f.get("zona")), align="center")
            _set(ws, r, 3, _tipo_txt(f.get("tipo", "")), align="center")
            _set(ws, r, 4, _clasif_txt(clasif), fg=col_estado, align="center")
            _set(ws, r, 5, f.get("pdl_code", ""), wrap=True)
            _set(ws, r, 6, f.get("mo_value", 0) or 0, align="right", num_fmt='"$"#,##0')
            _set(ws, r, 7, f.get("valor_pendiente", 0) or 0, align="right", num_fmt='"$"#,##0')
            _set(ws, r, 8, _antig_txt(antig), align="center")
            r += 1
    return ws


# ── API pública ───────────────────────────────────────────────────────────────
def generar_excel(fecha_inicio, fecha_fin, zonas=None, supervisores=None):
    """Genera el Excel del módulo MATERIALES. Devuelve (bytes, filename)."""
    d = analizar(fecha_inicio, fecha_fin, zonas, supervisores, use_cache=True)
    if not d.get("ok"):
        raise ValueError(d.get("error", "Error al analizar."))
    familias = d.get("familias", [])
    if not familias:
        raise ValueError(f"No hay datos para el periodo {d.get('periodo', '')}.")

    gen = datetime.now().strftime("%d/%m/%Y %H:%M")
    titulo = (f"Materiales Zona VII · {d.get('periodo', '')}   "
              f"({d.get('n_total', 0)} proyectos · {d.get('n_pendiente', 0)} con pendiente · "
              f"valor pendiente ${d.get('valor_pendiente_total', 0):,.0f} · generado {gen})")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _hoja_resumen(wb, d, titulo)
    _hoja_proyectos(wb, familias, titulo)
    _hoja_materiales(wb, familias, titulo)
    _hoja_por_supervisor(wb, familias, titulo)

    di = fecha_inicio[:10].replace("-", "")
    df = fecha_fin[:10].replace("-", "")
    fname = f"Materiales_ZonaVII_{di}_{df}.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue(), fname
