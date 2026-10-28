"""Módulo de Bolsas: casos por tipo (TUTELA, PQR, IPPE, CUN100, AMP, NNDD, TLC…).
Diseño núcleo común + datos_extra (JSON) → importación automática de cualquier
hoja sin perder columnas. De momento se expone solo en la cuenta dev (gate en UI)."""
import io
import json
from typing import Optional
from datetime import date, datetime
from collections import Counter
from fastapi import APIRouter, Depends, Query, UploadFile, File, Form, HTTPException
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_, or_, func, delete
import openpyxl
from app.db.database import get_db
from app.api.deps import get_current_active_user, assert_puede_escribir
from app.models.models import Bolsa, Usuario
from app.schemas.schemas import BolsaCreate, BolsaUpdate

router = APIRouter(prefix="/bolsas", tags=["bolsas"])

# Tipos canónicos (orden de presentación)
TIPOS_BOLSA = ["TUTELA", "PQR", "BOCATOMA", "IPPE", "CUN100", "CUNBT", "AMP", "NNDD", "TLC", "DRONE", "BP"]

# Hoja del Excel → tipo canónico (lo demás se ignora: Hoja1, Hoja2…)
SHEET_TIPO = {
    "TUTELA": "TUTELA", "PQR": "PQR", "CASOS BOCATOMA": "BOCATOMA", "IPPE": "IPPE",
    "CUN100": "CUN100", "BP": "BP", "AMP": "AMP", "CUNBT": "CUNBT", "NNDD": "NNDD",
    "DRONE": "DRONE", "TLC": "TLC",
}


def _nh(s) -> str:
    s = str(s or "").strip().lower()
    for a, b in [("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ñ", "n")]:
        s = s.replace(a, b)
    return s


# Encabezado normalizado → campo del núcleo
CORE_MAP = {
    "zona": "zona", "zona+": "zona",
    "municipio": "municipio",
    "circuito": "circuito", "cto": "circuito", "codigo cto": "circuito",
    "lcl": "lcl", "lcl asignada": "lcl", "lcl ejecucion": "lcl", "lcl ejecución": "lcl",
    "pdl": "pdl",
    "estado": "estado", "estado dominion": "estado", "estado revisado": "estado",
    "estado ostk": "estado", "estado equipo": "estado",
    "supervisor": "supervisor", "supervisor responsable": "supervisor",
    "responsable": "supervisor", "responsable bra": "supervisor",
    "fecha de prog": "fecha_prog", "fecha de programacion": "fecha_prog",
    "fecha de program": "fecha_prog", "fecha programacion": "fecha_prog",
    "fecha de programa": "fecha_prog", "fecha prgramación": "fecha_prog",
    "fecha entrega lc": "fecha_entrega", "fecha entrega ee": "fecha_entrega",
    "fecha entrega": "fecha_entrega",
    "observacion": "observacion", "observaciones": "observacion", "observación": "observacion",
}

# Columna identificadora (caso) por tipo, normalizada
CASO_POR_TIPO = {
    "TUTELA": "caso", "PQR": "pqrs", "BOCATOMA": "num. caso", "IPPE": "qr",
    "CUN100": "nombre", "AMP": "id", "CUNBT": "cd", "NNDD": "radicado",
    "DRONE": "activo", "TLC": "ticket nuevo", "BP": "caso",
}

_MAXLEN = {"zona": 100, "caso": 400, "municipio": 150, "circuito": 150, "lcl": 100, "pdl": 100, "estado": 100, "supervisor": 150}


def _to_date(v):
    """Acepta datetime/date o texto ISO/DD-MM-YYYY (input del formulario)."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _cell(v):
    """Valor JSON-serializable para datos_extra."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return v
    return str(v).strip()


def mapear_core(tipo: str, fila: dict) -> dict:
    """De {encabezado: valor} → campos del núcleo (zona, caso, estado, fechas…).
    Misma lógica para importar y para crear/editar a mano."""
    core: dict = {}
    caso_h = CASO_POR_TIPO.get(tipo)
    for h, val in fila.items():
        if val is None or str(val).strip() == "":
            continue
        nh = _nh(h)
        cf = CORE_MAP.get(nh)
        if cf and cf not in core:
            if cf in ("fecha_prog", "fecha_entrega"):
                d = _to_date(val)
                if d:
                    core[cf] = d
            else:
                core[cf] = str(val).strip()[:_MAXLEN.get(cf, 200)]
        if caso_h and nh == caso_h and "caso" not in core:
            core["caso"] = str(val).strip()[:400]
    return core


def parse_bolsas(wb) -> list[dict]:
    out: list[dict] = []
    for sheet, tipo in SHEET_TIPO.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = [r for r in ws.iter_rows(values_only=True)]
        # fila de encabezado = primera (de las 6 primeras) con ≥3 celdas de texto
        hidx = None
        for i, r in enumerate(rows[:6]):
            if r and sum(1 for c in r if isinstance(c, str) and c.strip()) >= 3:
                hidx = i
                break
        if hidx is None:
            continue
        headers = [str(c).strip() if c is not None else "" for c in rows[hidx]]

        for r in rows[hidx + 1:]:
            if not r or all(c is None or str(c).strip() == "" for c in r):
                continue
            # {encabezado: valor} de celdas no vacías
            fila = {
                headers[j]: val
                for j, val in enumerate(r)
                if j < len(headers) and headers[j] and val is not None and str(val).strip() != ""
            }
            if not fila:
                continue
            core = mapear_core(tipo, fila)
            core["tipo"] = tipo
            core["datos_extra"] = {h: _cell(v) for h, v in fila.items()}
            out.append(core)
    return out


def _to_dict(b: Bolsa) -> dict:
    return {
        "id": b.id, "unidad": b.unidad, "tipo": b.tipo, "zona": b.zona, "caso": b.caso,
        "municipio": b.municipio, "circuito": b.circuito, "lcl": b.lcl, "pdl": b.pdl,
        "estado": b.estado, "supervisor": b.supervisor,
        "fecha_entrega": b.fecha_entrega.isoformat() if b.fecha_entrega else None,
        "fecha_prog": b.fecha_prog.isoformat() if b.fecha_prog else None,
        "observacion": b.observacion,
        "datos_extra": json.loads(b.datos_extra) if b.datos_extra else {},
    }


# ── Catálogo / resumen ──────────────────────────────────────────────────────────

@router.get("/tipos")
async def tipos_bolsa(_: Usuario = Depends(get_current_active_user)):
    return TIPOS_BOLSA


# Campos por defecto si un tipo aún no tiene datos cargados
CAMPOS_FALLBACK = ["ZONA", "CASO", "MUNICIPIO", "CIRCUITO", "LCL", "PDL",
                   "ESTADO", "SUPERVISOR", "FECHA DE PROG", "FECHA ENTREGA LC", "OBSERVACIÓN"]


@router.get("/campos")
async def campos_tipo(tipo: str, db: AsyncSession = Depends(get_db), _: Usuario = Depends(get_current_active_user)):
    """Columnas que maneja un tipo (deducidas de los datos existentes) para armar
    el formulario dinámico. Si el tipo no tiene datos, usa un set por defecto."""
    rows = (await db.execute(
        select(Bolsa.datos_extra).where(Bolsa.tipo == tipo).limit(300)
    )).scalars().all()
    seen: list[str] = []
    for r in rows:
        if not r:
            continue
        try:
            d = json.loads(r)
        except Exception:
            continue
        for k in d.keys():
            if k not in seen:
                seen.append(k)
    return {"tipo": tipo, "campos": seen or CAMPOS_FALLBACK}


@router.get("/resumen")
async def resumen(db: AsyncSession = Depends(get_db), _: Usuario = Depends(get_current_active_user)):
    total = (await db.execute(select(func.count()).select_from(Bolsa))).scalar() or 0
    por_tipo = [
        {"tipo": t, "total": n}
        for t, n in (await db.execute(
            select(Bolsa.tipo, func.count()).group_by(Bolsa.tipo).order_by(func.count().desc())
        )).all()
    ]
    por_estado = [
        {"estado": (e or "SIN ESTADO"), "total": n}
        for e, n in (await db.execute(
            select(func.upper(func.trim(Bolsa.estado)), func.count())
            .group_by(func.upper(func.trim(Bolsa.estado))).order_by(func.count().desc())
        )).all()
    ]
    return {"total": total, "por_tipo": por_tipo, "por_estado": por_estado}


# ── Filtros compartidos (lista + reporte) ──────────────────────────────────────

class BolsaFiltros:
    def __init__(
        self,
        tipo: Optional[str] = None,
        zona: Optional[str] = None,
        estado: Optional[str] = None,
        supervisor: Optional[str] = None,
        municipio: Optional[str] = None,
        q: Optional[str] = None,
        fecha_inicio: Optional[date] = None,
        fecha_fin: Optional[date] = None,
    ):
        self.tipo = tipo
        self.zona = zona
        self.estado = estado
        self.supervisor = supervisor
        self.municipio = municipio
        self.q = q
        self.fecha_inicio = fecha_inicio
        self.fecha_fin = fecha_fin

    def filtros(self) -> list:
        f: list = []
        if self.tipo:
            f.append(Bolsa.tipo == self.tipo)
        if self.zona:
            f.append(Bolsa.zona == self.zona)
        if self.estado:
            f.append(func.upper(func.trim(Bolsa.estado)) == self.estado.strip().upper())
        if self.supervisor:
            f.append(Bolsa.supervisor.ilike(f"%{self.supervisor}%"))
        if self.municipio:
            f.append(Bolsa.municipio.ilike(f"%{self.municipio}%"))
        if self.q:
            like = f"%{self.q}%"
            f.append(or_(
                Bolsa.caso.ilike(like), Bolsa.municipio.ilike(like),
                Bolsa.observacion.ilike(like), Bolsa.pdl.ilike(like), Bolsa.lcl.ilike(like),
            ))
        if self.fecha_inicio:
            f.append(Bolsa.fecha_prog >= self.fecha_inicio)
        if self.fecha_fin:
            f.append(Bolsa.fecha_prog <= self.fecha_fin)
        return f


# ── Listado ─────────────────────────────────────────────────────────────────────

@router.get("")
async def listar(
    f: BolsaFiltros = Depends(),
    skip: int = Query(0, ge=0),
    limit: int = Query(3000, le=5000),
    db: AsyncSession = Depends(get_db),
    _: Usuario = Depends(get_current_active_user),
):
    filters = f.filtros()
    base = select(Bolsa)
    if filters:
        base = base.where(and_(*filters))
    total = (await db.execute(select(func.count()).select_from(base.subquery()))).scalar() or 0
    items = (await db.execute(base.order_by(Bolsa.tipo, Bolsa.id).offset(skip).limit(limit))).scalars().all()
    return {"total": total, "items": [_to_dict(b) for b in items]}


# ── Reporte gráfico: KPIs y opciones de filtro ──────────────────────────────────

@router.get("/opciones")
async def bolsas_opciones(db: AsyncSession = Depends(get_db), _: Usuario = Depends(get_current_active_user)):
    async def dist(col, limite=None):
        q = (
            select(func.trim(col)).where(col.isnot(None), func.trim(col) != "")
            .group_by(func.trim(col)).order_by(func.count().desc())
        )
        vals = [r[0] for r in (await db.execute(q)).all() if r[0]]
        return vals[:limite] if limite else vals

    eq = (
        select(func.upper(func.trim(Bolsa.estado)))
        .where(Bolsa.estado.isnot(None), func.trim(Bolsa.estado) != "")
        .group_by(func.upper(func.trim(Bolsa.estado))).order_by(func.count().desc())
    )
    estados = [r[0] for r in (await db.execute(eq)).all() if r[0]]
    return {
        "tipos": TIPOS_BOLSA,
        "zonas": await dist(Bolsa.zona),
        "estados": estados,
        "supervisores": await dist(Bolsa.supervisor, 300),
        "municipios": await dist(Bolsa.municipio, 300),
    }


@router.get("/kpis")
async def bolsas_kpis(
    f: BolsaFiltros = Depends(),
    db: AsyncSession = Depends(get_db),
    _: Usuario = Depends(get_current_active_user),
):
    base = f.filtros()

    def W(q):
        return q.where(and_(*base)) if base else q

    total = (await db.execute(W(select(func.count()).select_from(Bolsa)))).scalar() or 0

    async def grupo(col, norm=False, limite=None):
        expr = func.upper(func.trim(col)) if norm else col
        q = W(select(expr.label("k"), func.count().label("n")).group_by(expr).order_by(func.count().desc()))
        rows = (await db.execute(q)).all()
        out = [{"label": (r.k if r.k not in (None, "") else "—"), "total": r.n} for r in rows]
        return out[:limite] if limite else out

    por_tipo = await grupo(Bolsa.tipo)
    por_estado = await grupo(Bolsa.estado, norm=True)
    por_zona = await grupo(Bolsa.zona)
    por_supervisor = await grupo(Bolsa.supervisor, limite=12)
    por_municipio = await grupo(Bolsa.municipio, limite=12)

    # Línea de tiempo por mes de fecha_prog
    mexpr = func.to_char(Bolsa.fecha_prog, "YYYY-MM")
    tq = (
        select(mexpr.label("m"), func.count().label("n"))
        .where(and_(*(base + [Bolsa.fecha_prog.isnot(None)])))
        .group_by(mexpr).order_by(mexpr)
    )
    timeline = [{"periodo": r.m, "total": r.n} for r in (await db.execute(tq)).all()]
    con_fecha = sum(t["total"] for t in timeline)

    return {
        "total": total,
        "por_tipo": por_tipo,
        "por_estado": por_estado,
        "por_zona": por_zona,
        "por_supervisor": por_supervisor,
        "por_municipio": por_municipio,
        "timeline": timeline,
        "con_fecha_prog": con_fecha,
        "sin_fecha_prog": total - con_fecha,
    }


# ── CRUD ────────────────────────────────────────────────────────────────────────

@router.post("")
async def crear(req: BolsaCreate, db: AsyncSession = Depends(get_db), current_user: Usuario = Depends(get_current_active_user)):
    assert_puede_escribir(current_user)
    data = req.model_dump()
    extra = data.pop("datos_extra", None) or {}
    # Núcleo derivado de los campos del tipo (lo explícito tiene prioridad)
    if extra:
        for k, v in mapear_core(data.get("tipo"), extra).items():
            if not data.get(k):
                data[k] = v
    if not data.get("unidad"):
        data["unidad"] = current_user.unidad or "SABOGAL"
    b = Bolsa(**data, datos_extra=json.dumps(extra, ensure_ascii=False) if extra else None)
    db.add(b)
    await db.commit()
    await db.refresh(b)
    return _to_dict(b)


@router.patch("/{bolsa_id}")
async def actualizar(bolsa_id: int, req: BolsaUpdate, db: AsyncSession = Depends(get_db), current_user: Usuario = Depends(get_current_active_user)):
    assert_puede_escribir(current_user)
    b = (await db.execute(select(Bolsa).where(Bolsa.id == bolsa_id))).scalar_one_or_none()
    if not b:
        raise HTTPException(status_code=404, detail="No encontrado")
    data = req.model_dump(exclude_unset=True)
    extra = data.pop("datos_extra", None)
    if extra is not None:
        b.datos_extra = json.dumps(extra, ensure_ascii=False)
        # Recalcular el núcleo a partir de los campos del tipo
        for k, v in mapear_core(data.get("tipo") or b.tipo, extra).items():
            setattr(b, k, v)
    for k, v in data.items():   # campos del núcleo explícitos mandan
        setattr(b, k, v)
    await db.commit()
    await db.refresh(b)
    return _to_dict(b)


@router.delete("/{bolsa_id}")
async def eliminar(bolsa_id: int, db: AsyncSession = Depends(get_db), current_user: Usuario = Depends(get_current_active_user)):
    assert_puede_escribir(current_user)
    b = (await db.execute(select(Bolsa).where(Bolsa.id == bolsa_id))).scalar_one_or_none()
    if not b:
        raise HTTPException(status_code=404, detail="No encontrado")
    await db.delete(b)
    await db.commit()
    return {"ok": True}


# ── Importación masiva del Excel multi-hoja ─────────────────────────────────────

@router.post("/import")
async def importar_bolsas(
    file: UploadFile = File(...),
    reemplazar: bool = Form(False),
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    assert_puede_escribir(current_user)
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="El archivo debe ser .xlsx")
    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo: {e}")

    registros = parse_bolsas(wb)
    uni = current_user.unidad or "SABOGAL"

    if reemplazar:
        await db.execute(delete(Bolsa))
        await db.flush()

    por_tipo = Counter()
    for rec in registros:
        extra = rec.pop("datos_extra", {})
        b = Bolsa(unidad=uni, datos_extra=json.dumps(extra, ensure_ascii=False, default=str), **rec)
        db.add(b)
        por_tipo[rec["tipo"]] += 1
    await db.commit()

    return {"inserted": sum(por_tipo.values()), "por_tipo": dict(por_tipo)}
