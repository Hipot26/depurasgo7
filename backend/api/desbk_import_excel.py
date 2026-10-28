"""
Endpoint de importación de Excel.
Solo accesible para el Jefe de Zona (supervisor).
Lógica inteligente: evita duplicados por (zona, fecha, lider).
"""
from fastapi import APIRouter, Depends, UploadFile, File, Form, HTTPException
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_, delete
from typing import Optional
from datetime import date, datetime
import io
import openpyxl

from app.db.database import get_db
from app.api.deps import get_current_active_user
from app.models.models import Usuario, Programacion, RolEnum, ZonaEnum

router = APIRouter(prefix="/import", tags=["import"])

def _norm(s: str) -> str:
    return (s or "").strip().lower()\
        .replace("á","a").replace("é","e").replace("í","i").replace("ó","o").replace("ú","u")

COL_MAP = {
    "zona":                        "zona",
    "fecha":                       "fecha",
    "lider":                       "lider",
    "tipo grupo":                  "tipo_grupo",
    "tipo recurso":                "tipo_grupo",
    "proceso":                     "proceso",
    "subproceso":                  "subproceso",
    "proyecto o antecedente":      "proyecto_antecedente",
    "pms":                         "pms",
    "circuito":                    "circuito",
    "lcl":                         "lcl",
    "fecha vigencia":              "fecha_vigencia",
    "tipo pdt":                    "tipo_pdt",
    "# pdt":                       "num_pdt",
    "estado":                      "estado",
    "cd afectados":                "cd_afectados",
    "hora":                        "hora",
    "municipio":                   "municipio",
    "vda":                         "vereda",
    "alcance":                     "alcance",
    "supervisor":                  "supervisor",
    "valordia":                    "valor_dia",
    "valor dia":                   "valor_dia",
    "mo proyectada":               "mo_proyectada",
    "moproyectada":                "mo_proyectada",
    "mo real":                     "mo_real",
    "valor dejado de ejecutar":    "valor_dejado_ejecutar",
    "valor material":              "valor_material",
    "valormaterial":               "valor_material",
    "estado pdt":                  "estado_pdt",
    "notas":                       "notas",
    "#proliquidado":               "num_proliquidado",
    "fecha liquidacion sgs":       "fecha_liquidacion_sgs",
    "estado facturcion":           "estado_facturacion",
    "estado facturacion":          "estado_facturacion",
    "fecha envio a factuacion":    "fecha_envio_facturacion",
    "fecha envio a facturacion":   "fecha_envio_facturacion",
    "ot":                          "ot",
    "estado cumplimentacion":      "estado_cumplimentacion",
    "inoperatividad":              "inoperatividad",
}
SKIP            = {"mes", "dia"}
NUMERIC_FIELDS  = {"valor_dia","mo_proyectada","mo_real","valor_dejado_ejecutar","valor_material","inoperatividad"}
DATE_FIELDS     = {"fecha","fecha_vigencia","fecha_liquidacion_sgs","fecha_envio_facturacion"}
VALID_ZONAS     = {"ALMEIDAS","GUAVIO","UBATE","LV&TLC"}
VALID_ESTADOS   = {"PROGRAMADO","EJECUTADO","PARCIAL","CANCELADO","REPROGRAMADO"}
VALID_TIPOS     = {"CAMP5","CAMP3","ARRIME","GRUA","ESTACA","LVL3","MIE"}
VALID_EPDT      = {"PENDIENTE","ENTREGADO_FACTURACION","LIQUIDADO_SGS","RECHAZADO"}


def _to_str(v) -> Optional[str]:
    if v is None: return None
    s = str(v).strip(); return s or None

def _to_date(v) -> Optional[date]:
    if v is None: return None
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date):     return v
    for fmt in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y"):
        try: return datetime.strptime(str(v).strip(), fmt).date()
        except: pass
    return None

def _to_num(v) -> Optional[float]:
    if v is None: return None
    try: return float(v)
    except: return None


def parse_workbook(wb: openpyxl.Workbook) -> tuple[list[dict], list[str]]:
    errors: list[str] = []
    records: list[dict] = []

    # Buscar la hoja principal
    sheet = None
    for name in wb.sheetnames:
        nu = name.upper()
        if any(x in nu for x in ["TELECONTROL","RESUMEN","HOJA1"]):
            continue
        sheet = wb[name]; break

    if sheet is None:
        errors.append("No se encontró hoja principal (se ignoraron TELECONTROL, RESUMEN, Hoja1)")
        return records, errors

    raw_headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    headers = [_norm(h) for h in raw_headers]

    fecha_idx = next((i for i, h in enumerate(headers) if h == "fecha"), None)
    if fecha_idx is None:
        errors.append("No se encontró columna FECHA"); return records, errors

    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if fecha_idx >= len(row) or row[fecha_idx] is None:
            continue
        rec: dict = {}
        for i, val in enumerate(row):
            h = headers[i] if i < len(headers) else None
            if not h or h in SKIP: continue
            field = COL_MAP.get(h)
            if not field: continue
            if field in DATE_FIELDS:    rec[field] = _to_date(val)
            elif field in NUMERIC_FIELDS: rec[field] = _to_num(val)
            else:                        rec[field] = _to_str(val)

        if not rec.get("fecha"): continue

        zona = (rec.get("zona") or "").upper().strip()
        if zona not in VALID_ZONAS:
            errors.append(f"Fila {row_num}: zona '{zona}' inválida — omitida"); continue
        rec["zona"] = zona

        # estado, tipo_grupo y estado_pdt son TEXTO LIBRE (multi-unidad):
        # se conservan tal cual (Sabogal: CAMP5/APROBADO… · Agobardo: MOE2/TAI…)
        rec["estado"]     = (rec.get("estado") or "").strip().upper() or None
        rec["tipo_grupo"] = (rec.get("tipo_grupo") or "").strip().upper() or None
        rec["estado_pdt"] = (rec.get("estado_pdt") or "").strip().upper() or None

        records.append(rec)

    return records, errors


@router.post("/excel")
async def importar_excel(
    file: UploadFile = File(...),
    mode: str = Form("smart"),   # smart | replace_zona | replace_all
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    if current_user.rol != RolEnum.supervisor:
        raise HTTPException(status_code=403, detail="Solo el Jefe de Zona puede importar")

    # Unidad del importador: cada quien solo escribe/borra en SU unidad
    uni = current_user.unidad or "SABOGAL"

    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="El archivo debe ser .xlsx o .xls")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo: {e}")

    records, parse_errors = parse_workbook(wb)
    if not records:
        return {"inserted": 0, "skipped": 0, "errors": parse_errors}

    zonas_en_archivo = {r["zona"] for r in records}

    # ── Modos de limpieza (SIEMPRE acotados a la unidad del importador) ─────────
    if mode == "replace_all":
        await db.execute(delete(Programacion).where(Programacion.unidad == uni))
        await db.flush()

    elif mode == "replace_zona":
        for z in zonas_en_archivo:
            await db.execute(
                delete(Programacion).where(and_(Programacion.zona == z, Programacion.unidad == uni))
            )
        await db.flush()

    # ── Importación inteligente: skip duplicados (zona+fecha+lider) en la unidad ─
    inserted = skipped = 0

    if mode == "smart":
        existing_q = await db.execute(
            select(Programacion.zona, Programacion.fecha, Programacion.lider)
            .where(and_(Programacion.zona.in_(list(zonas_en_archivo)), Programacion.unidad == uni))
        )
        existing_keys = {(r.zona, r.fecha, r.lider) for r in existing_q.all()}
    else:
        existing_keys = set()

    for rec in records:
        key = (rec.get("zona"), rec.get("fecha"), rec.get("lider"))
        if mode == "smart" and key in existing_keys:
            skipped += 1
            continue
        db.add(Programacion(**{k: v for k, v in rec.items() if v is not None},
                            unidad=uni, creado_por=current_user.id))
        inserted += 1

    await db.commit()

    return {
        "inserted": inserted,
        "skipped":  skipped,
        "errors":   parse_errors[:20],   # max 20 errores reportados
    }
