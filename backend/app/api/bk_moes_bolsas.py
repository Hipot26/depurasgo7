"""API del módulo Bolsas MOES (unidad CHOCONTA / Pisos Móviles).

Flujo:
  1. El coordinador de la unidad sube el Excel de Enel (semana 1 = base; semana 2, 3… = actualizaciones).
  2. El sistema lo desbloquea si viene cifrado, lo parsea y CRUZA contra el estado actual:
       • tickets nuevos  → alta
       • tickets que cambian → modificación (con la línea de tiempo de qué cambió)
       • cambios en campos de Enel que no deberían cambiar → inconsistencia
       • mismo ticket repetido dentro del archivo → duplicado (se marca, no se descarta)
  3. Cada Excel queda guardado y descargable (pestaña 'Excel subidos') justificando todo.

Aislado por unidad: la cuenta `moes` (unidad CHOCONTA) solo ve su unidad; dev/directivo ven todo.
"""
import io
import os
import json
import hashlib
from collections import Counter, defaultdict
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, Query, UploadFile, File, Form, HTTPException, Body
from fastapi.responses import StreamingResponse, FileResponse
from sqlalchemy import select, and_, or_, func, delete
from sqlalchemy.ext.asyncio import AsyncSession
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.db.database import get_db
from app.api.deps import get_current_active_user, assert_puede_escribir, unidad_de, ROLES_GLOBALES
from app.models.models import BolsaMoes, BolsaMoesImport, BolsaMoesCambio, Usuario
from app.services import moes

router = APIRouter(prefix="/moes", tags=["bolsas-moes"])

MOES_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "data", "moes_fuentes"))
UNIDAD_MOES_DEFAULT = "CHOCONTA"

# Campos del núcleo que se vuelcan a columnas estructuradas (el resto va en datos_extra)
CORE_FIELDS = ["estado", "antecedente", "lcl", "odm", "responsable_enel", "proyecto",
               "municipio", "cto", "valor_mo", "fecha_max", "fecha_ejecucion",
               "pdl_tck", "se_ejecuto", "pct_ejecutado", "observaciones"]
# Campos editables a mano por la unidad (amarillos)
EDITABLES = {"estado", "fecha_ejecucion", "pdl_tck", "se_ejecuto", "pct_ejecutado", "observaciones"}
_MAXLEN = {"antecedente": 300, "lcl": 100, "odm": 150, "estado": 80, "responsable_enel": 150,
           "proyecto": 200, "municipio": 120, "cto": 120, "pdl_tck": 200, "se_ejecuto": 10}


# ── Helpers ──────────────────────────────────────────────────────────────────────

def _unidad_destino(user: Usuario, unidad: Optional[str]) -> str:
    """Unidad sobre la que opera: la cuenta moes/supervisor usa la suya; dev/directivo
    pueden indicar otra (?unidad=), con CHOCONTA por defecto."""
    if user.rol in ROLES_GLOBALES:
        return unidad or UNIDAD_MOES_DEFAULT
    return unidad_de(user)


def _scope(user: Usuario, unidad: Optional[str]) -> list:
    if user.rol in ROLES_GLOBALES:
        return [BolsaMoes.unidad == unidad] if unidad else []
    return [BolsaMoes.unidad == unidad_de(user)]


def _trunc(field: str, v):
    if isinstance(v, str):
        return v[:_MAXLEN.get(field, 400)]
    return v


def _aplicar_core(b: BolsaMoes, core: dict):
    """Vuelca los campos del núcleo (typed) sobre la fila. Solo los presentes (no vacíos)
    → un Excel que no trae un campo no lo borra."""
    for f in CORE_FIELDS:
        if f in core and core[f] not in (None, ""):
            setattr(b, f, _trunc(f, core[f]))


def _to_dict(b: BolsaMoes) -> dict:
    return {
        "id": b.id, "unidad": b.unidad, "clave": b.clave, "antecedente": b.antecedente,
        "lcl": b.lcl, "odm": b.odm, "estado": b.estado, "responsable_enel": b.responsable_enel,
        "proyecto": b.proyecto, "municipio": b.municipio, "cto": b.cto, "valor_mo": b.valor_mo,
        "fecha_max": b.fecha_max.isoformat() if b.fecha_max else None,
        "fecha_ejecucion": b.fecha_ejecucion.isoformat() if b.fecha_ejecucion else None,
        "pdl_tck": b.pdl_tck, "se_ejecuto": b.se_ejecuto, "pct_ejecutado": b.pct_ejecutado,
        "observaciones": b.observaciones, "duplicado": b.duplicado,
        "import_alta_id": b.import_alta_id, "import_visto_id": b.import_visto_id,
        "datos_extra": json.loads(b.datos_extra) if b.datos_extra else {},
        "actualizado_en": b.actualizado_en.isoformat() if b.actualizado_en else None,
    }


# ── Importación + cruce (corazón del módulo) ─────────────────────────────────────

@router.post("/import")
async def importar(
    file: UploadFile = File(...),
    password: Optional[str] = Form(None),
    reemplazar: bool = Form(False),
    fecha_version: Optional[str] = Form(None),   # fecha de la versión (la indica quien sube)
    forzar: bool = Form(False),                  # aplicar aunque sea más vieja que la última
    unidad: Optional[str] = Form(None),
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    assert_puede_escribir(current_user)
    if not file.filename.lower().endswith((".xlsx", ".xlsm", ".xls")):
        raise HTTPException(status_code=400, detail="El archivo debe ser un Excel (.xlsx)")
    contents = await file.read()
    try:
        wb, cifrado = moes.cargar_libro(contents, password)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo: {e}")

    parsed = moes.parse(wb)
    uni = _unidad_destino(current_user, unidad)
    fv = moes._to_date(fecha_version) or date.today()

    # Guardar el archivo original (procedencia / descargable)
    file_hash = hashlib.sha256(contents).hexdigest()
    nombre = os.path.basename(file.filename)
    carpeta = os.path.join(MOES_DIR, uni)
    os.makedirs(carpeta, exist_ok=True)
    ruta_abs = os.path.join(carpeta, f"{datetime.now():%Y%m%d_%H%M%S}_{nombre}")
    with open(ruta_abs, "wb") as fh:
        fh.write(contents)

    # Recencia POR FECHA DE VERSIÓN (no por orden de carga): la más reciente manda.
    ultima_version = (await db.execute(
        select(func.max(BolsaMoesImport.fecha_version))
        .where(BolsaMoesImport.unidad == uni, BolsaMoesImport.aplicado.is_(True))
    )).scalar()
    # Aplica (actualiza el estado actual) si: reemplaza todo, fuerza, no hay versión previa,
    # o esta versión es igual/más nueva que la última aplicada.
    aplicar = reemplazar or forzar or ultima_version is None or fv >= ultima_version

    imp = BolsaMoesImport(
        unidad=uni, nombre=nombre, ruta=os.path.relpath(ruta_abs, MOES_DIR),
        hoja=parsed["hoja"], tamano=len(contents), hash=file_hash, cifrado=cifrado,
        fecha_version=fv, aplicado=aplicar,
        subido_por_id=current_user.id, subido_por=current_user.nombre,
    )
    db.add(imp)
    await db.flush()  # imp.id

    # Versión más vieja que la última aplicada → se archiva pero NO pisa el estado actual
    if not aplicar:
        imp.filas_total = len(parsed["filas"])
        await db.commit()
        return {
            "import_id": imp.id, "hoja": parsed["hoja"], "cifrado": cifrado, "archivo": nombre,
            "aplicado": False, "fecha_version": fv.isoformat(),
            "ultima_version": ultima_version.isoformat() if ultima_version else None,
            "aviso": (f"Este archivo es de la versión {fv:%d/%m/%Y}, MÁS ANTIGUA que la última "
                      f"cargada ({ultima_version:%d/%m/%Y}). Se guardó como respaldo pero NO actualicé "
                      f"el estado actual para no pisar datos más nuevos. Si de verdad debe mandar, "
                      f"vuelve a subirlo marcando 'forzar'."),
            "tipo_archivo": moes.clasificar_archivo(parsed)["tipo"],
            "filas_total": imp.filas_total, "alta": 0, "modificadas": 0, "sin_cambio": 0,
            "cambios_estado": 0, "valores_a_revisar": 0, "duplicados": 0,
            "columnas_nuevas": [], "ausentes": 0,
        }

    if reemplazar:
        await db.execute(delete(BolsaMoesCambio).where(BolsaMoesCambio.unidad == uni))
        await db.execute(delete(BolsaMoes).where(BolsaMoes.unidad == uni))
        await db.flush()

    existentes = {b.clave: b for b in (
        await db.execute(select(BolsaMoes).where(BolsaMoes.unidad == uni))
    ).scalars().all()}
    headers_conocidos = set()
    for b in existentes.values():
        if b.datos_extra:
            try:
                headers_conocidos.update(moes._nh(h) for h in json.loads(b.datos_extra).keys())
            except Exception:
                pass

    # Dedup dentro del mismo archivo: agrupamos por clave (última fila = estado actual).
    # Un duplicado REAL es la misma clave repetida con el MISMO estado (repetición de verdad).
    # Si se repite con estados distintos es HISTORIAL (copia trabajada), no duplicado.
    ocurrencias: dict = defaultdict(list)
    por_clave: dict = {}
    orden: list = []
    for fila in parsed["filas"]:
        k = fila["clave"]
        if k not in por_clave:
            orden.append(k)
        ocurrencias[k].append(fila)
        por_clave[k] = fila

    def es_dup_real(k: str) -> bool:
        rows = ocurrencias[k]
        if len(rows) <= 1:
            return False
        estados = {moes._norm(r["core"].get("estado")) for r in rows}
        return len(estados) == 1   # mismo estado repetido = duplicado real

    usuario_nombre = current_user.nombre
    alta = mod = igual = cambios_estado = inconsist = 0
    for k in orden:
        fila = por_clave[k]
        core, extra = fila["core"], fila["extra"]
        dup = es_dup_real(k)
        b = existentes.get(k)
        if b is None:
            b = BolsaMoes(unidad=uni, clave=k, datos_extra=json.dumps(extra, ensure_ascii=False, default=str),
                          duplicado=dup, import_alta_id=imp.id, import_visto_id=imp.id, version_aplicada=fv)
            _aplicar_core(b, core)
            db.add(b)
            await db.flush()
            db.add(BolsaMoesCambio(unidad=uni, bolsa_id=b.id, import_id=imp.id, clave=k,
                                   antecedente=b.antecedente, tipo="alta", campo=None,
                                   valor_anterior=None, valor_nuevo=(b.estado or "—"),
                                   usuario_id=current_user.id, usuario_nombre=usuario_nombre))
            alta += 1
        else:
            viejo = json.loads(b.datos_extra) if b.datos_extra else {}
            merged = moes.fusionar_extra(viejo, extra)
            # Solo rastreamos columnas con significado (las calculadas/fórmula se ignoran: son ruido)
            cambios = [c for c in moes.diff_extra(viejo, extra) if not moes.es_calculado(c["campo"])]
            if cambios:
                for c in cambios:
                    inc = moes.es_inconsistencia(c["campo"], c["antes"], c["despues"])
                    if inc:
                        inconsist += 1
                    if moes._nh(c["campo"]) == "estado":
                        cambios_estado += 1
                    db.add(BolsaMoesCambio(
                        unidad=uni, bolsa_id=b.id, import_id=imp.id, clave=k, antecedente=b.antecedente,
                        tipo="modificacion", campo=c["campo"],
                        valor_anterior=(None if c["antes"] is None else str(c["antes"])),
                        valor_nuevo=(None if c["despues"] is None else str(c["despues"])),
                        inconsistencia=inc,
                        usuario_id=current_user.id, usuario_nombre=usuario_nombre))
                _aplicar_core(b, core)
                b.datos_extra = json.dumps(merged, ensure_ascii=False, default=str)
                mod += 1
            else:
                igual += 1
            if dup:
                b.duplicado = True
            b.import_visto_id = imp.id
            b.version_aplicada = fv

    # Columnas que aparecen por primera vez (evolución del formato de Enel)
    nuevas = [h for h in parsed["headers"] if headers_conocidos and moes._nh(h) not in headers_conocidos]
    ausentes = sum(1 for b in existentes.values() if b.import_visto_id != imp.id)
    dup_reales = sum(1 for k in orden if es_dup_real(k))
    clasif = moes.clasificar_archivo(parsed)

    imp.filas_total = len(parsed["filas"])
    imp.filas_alta = alta
    imp.filas_mod = mod
    imp.filas_igual = igual
    imp.duplicados = dup_reales
    imp.columnas_nuevas = json.dumps(nuevas, ensure_ascii=False) if nuevas else None
    await db.commit()

    aviso = None
    if clasif["tipo"] == "copia_trabajada":
        aviso = ("Este archivo parece una COPIA YA TRABAJADA (hoja con historial), no una bolsa "
                 "cruda de Enel. Para una trazabilidad limpia sube siempre la bolsa original de Enel.")

    return {
        "import_id": imp.id, "hoja": parsed["hoja"], "cifrado": cifrado, "archivo": nombre,
        "tipo_archivo": clasif["tipo"], "aviso": aviso,
        "aplicado": True, "fecha_version": fv.isoformat(),
        "ultima_version": ultima_version.isoformat() if ultima_version else None,
        "filas_total": imp.filas_total, "alta": alta, "modificadas": mod, "sin_cambio": igual,
        "cambios_estado": cambios_estado, "valores_a_revisar": inconsist,
        "duplicados": dup_reales, "columnas_nuevas": nuevas, "ausentes": ausentes,
    }


# ── Listado del estado actual ────────────────────────────────────────────────────

@router.get("")
async def listar(
    q: Optional[str] = None,
    estado: Optional[str] = None,
    municipio: Optional[str] = None,
    solo_duplicados: bool = False,
    solo_ausentes: bool = False,
    unidad: Optional[str] = None,
    skip: int = Query(0, ge=0),
    limit: int = Query(3000, le=8000),
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    filtros = _scope(current_user, unidad)
    if estado:
        filtros.append(func.upper(func.trim(BolsaMoes.estado)) == estado.strip().upper())
    if municipio:
        filtros.append(BolsaMoes.municipio.ilike(f"%{municipio}%"))
    if solo_duplicados:
        filtros.append(BolsaMoes.duplicado.is_(True))
    if q:
        like = f"%{q}%"
        filtros.append(or_(
            BolsaMoes.antecedente.ilike(like), BolsaMoes.lcl.ilike(like), BolsaMoes.odm.ilike(like),
            BolsaMoes.municipio.ilike(like), BolsaMoes.proyecto.ilike(like),
            BolsaMoes.pdl_tck.ilike(like), BolsaMoes.observaciones.ilike(like),
            BolsaMoes.datos_extra.ilike(like),
        ))
    base = select(BolsaMoes)
    if filtros:
        base = base.where(and_(*filtros))
    # Última carga de la unidad: un ticket es "ausente" si no vino en ella
    uni = _unidad_destino(current_user, unidad)
    ult = (await db.execute(
        select(func.max(BolsaMoesImport.id)).where(BolsaMoesImport.unidad == uni)
    )).scalar()
    if solo_ausentes and ult:
        base = base.where(or_(BolsaMoes.import_visto_id.is_(None), BolsaMoes.import_visto_id != ult))

    total = (await db.execute(select(func.count()).select_from(base.subquery()))).scalar() or 0
    items = (await db.execute(
        base.order_by(BolsaMoes.estado, BolsaMoes.id).offset(skip).limit(limit)
    )).scalars().all()
    def out(b):
        d = _to_dict(b)
        d["ausente"] = bool(ult and b.import_visto_id != ult)
        return d
    return {"total": total, "items": [out(b) for b in items]}


@router.get("/resumen")
async def resumen(unidad: Optional[str] = None, db: AsyncSession = Depends(get_db),
                  current_user: Usuario = Depends(get_current_active_user)):
    scope = _scope(current_user, unidad)
    def W(qq):
        return qq.where(and_(*scope)) if scope else qq
    total = (await db.execute(W(select(func.count()).select_from(BolsaMoes)))).scalar() or 0
    por_estado = [
        {"estado": (e or "SIN ESTADO"), "total": n}
        for e, n in (await db.execute(W(
            select(func.upper(func.trim(BolsaMoes.estado)), func.count())
            .group_by(func.upper(func.trim(BolsaMoes.estado))).order_by(func.count().desc())
        ))).all()
    ]
    dups = (await db.execute(W(select(func.count()).select_from(BolsaMoes).where(BolsaMoes.duplicado.is_(True))))).scalar() or 0
    uni = _unidad_destino(current_user, unidad)
    n_imports = (await db.execute(select(func.count()).select_from(BolsaMoesImport).where(BolsaMoesImport.unidad == uni))).scalar() or 0
    return {"total": total, "por_estado": por_estado, "duplicados": dups, "imports": n_imports}


@router.get("/opciones")
async def opciones(unidad: Optional[str] = None, db: AsyncSession = Depends(get_db),
                   current_user: Usuario = Depends(get_current_active_user)):
    scope = _scope(current_user, unidad)
    async def dist(col, lim=None):
        qq = select(func.trim(col)).where(col.isnot(None), func.trim(col) != "")
        if scope:
            qq = qq.where(and_(*scope))
        qq = qq.group_by(func.trim(col)).order_by(func.count().desc())
        vals = [r[0] for r in (await db.execute(qq)).all() if r[0]]
        return vals[:lim] if lim else vals
    eq = select(func.upper(func.trim(BolsaMoes.estado))).where(BolsaMoes.estado.isnot(None), func.trim(BolsaMoes.estado) != "")
    if scope:
        eq = eq.where(and_(*scope))
    eq = eq.group_by(func.upper(func.trim(BolsaMoes.estado))).order_by(func.count().desc())
    estados = [r[0] for r in (await db.execute(eq)).all() if r[0]]
    return {"estados": estados, "municipios": await dist(BolsaMoes.municipio, 200)}


# ── Indicadores (pivote estado×mes + cumplimiento económico) ─────────────────────

EST_FIN = {"FINALIZADO", "EJECUTADO", "FINALIZADA"}


@router.get("/indicadores")
async def indicadores(unidad: Optional[str] = None, db: AsyncSession = Depends(get_db),
                      current_user: Usuario = Depends(get_current_active_user)):
    scope = _scope(current_user, unidad)
    base = select(BolsaMoes)
    if scope:
        base = base.where(and_(*scope))
    rows = (await db.execute(base)).scalars().all()

    por_estado: Counter = Counter()
    pivote: dict = defaultdict(Counter)            # mes -> estado -> count
    estados_set: set = set()
    val_presupuestado = val_ejecutado = 0.0
    por_municipio: Counter = Counter()
    for b in rows:
        est = (b.estado or "SIN ESTADO").strip().upper()
        por_estado[est] += 1
        estados_set.add(est)
        f = b.fecha_ejecucion or b.fecha_max
        mes = f.strftime("%Y-%m") if f else "Sin fecha"
        pivote[mes][est] += 1
        if b.municipio:
            por_municipio[b.municipio.strip()] += 1
        v = b.valor_mo or 0
        val_presupuestado += v
        val_ejecutado += v * (b.pct_ejecutado or 0)

    total = len(rows)
    finalizadas = sum(n for e, n in por_estado.items() if e in EST_FIN)
    meses = sorted(pivote.keys())
    estados_orden = [e for e, _ in por_estado.most_common()]
    return {
        "total": total,
        "finalizadas": finalizadas,
        "cumplimiento": round(finalizadas / total, 4) if total else 0,
        "por_estado": [{"estado": e, "total": n} for e, n in por_estado.most_common()],
        "pivote": {
            "meses": meses,
            "estados": estados_orden,
            "datos": {m: {e: pivote[m].get(e, 0) for e in estados_orden} for m in meses},
            "totales_mes": {m: sum(pivote[m].values()) for m in meses},
        },
        "economico": {
            "presupuestado": round(val_presupuestado, 2),
            "ejecutado": round(val_ejecutado, 2),
            "pct": round(val_ejecutado / val_presupuestado, 4) if val_presupuestado else 0,
        },
        "por_municipio": [{"municipio": m, "total": n} for m, n in por_municipio.most_common(15)],
    }


# ── Línea de tiempo / historial de cambios ───────────────────────────────────────

@router.get("/cambios")
async def cambios(
    clave: Optional[str] = None,
    bolsa_id: Optional[int] = None,
    import_id: Optional[int] = None,
    solo_inconsistencias: bool = False,
    unidad: Optional[str] = None,
    limit: int = Query(500, le=3000),
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    filtros = []
    if current_user.rol not in ROLES_GLOBALES:
        filtros.append(BolsaMoesCambio.unidad == unidad_de(current_user))
    elif unidad:
        filtros.append(BolsaMoesCambio.unidad == unidad)
    if clave:
        filtros.append(BolsaMoesCambio.clave == clave)
    if bolsa_id:
        filtros.append(BolsaMoesCambio.bolsa_id == bolsa_id)
    if import_id:
        filtros.append(BolsaMoesCambio.import_id == import_id)
    if solo_inconsistencias:
        filtros.append(BolsaMoesCambio.inconsistencia.is_(True))
    qq = select(BolsaMoesCambio)
    if filtros:
        qq = qq.where(and_(*filtros))
    rows = (await db.execute(qq.order_by(BolsaMoesCambio.id.desc()).limit(limit))).scalars().all()
    # Nombre del Excel de cada cambio (para leer la línea de tiempo por versión)
    ids = {c.import_id for c in rows if c.import_id}
    nombres = {}
    if ids:
        nombres = {i.id: i.nombre for i in (await db.execute(
            select(BolsaMoesImport).where(BolsaMoesImport.id.in_(ids))
        )).scalars().all()}
    return [{
        "id": c.id, "bolsa_id": c.bolsa_id, "import_id": c.import_id, "clave": c.clave,
        "import_nombre": nombres.get(c.import_id), "antecedente": c.antecedente,
        "tipo": c.tipo, "campo": c.campo,
        "valor_anterior": c.valor_anterior, "valor_nuevo": c.valor_nuevo,
        "inconsistencia": c.inconsistencia, "usuario": c.usuario_nombre,
        "creado_en": c.creado_en.isoformat() if c.creado_en else None,
    } for c in rows]


# ── Excel subidos (procedencia) ──────────────────────────────────────────────────

@router.get("/importaciones")
async def importaciones(unidad: Optional[str] = None, db: AsyncSession = Depends(get_db),
                        current_user: Usuario = Depends(get_current_active_user)):
    filtros = []
    if current_user.rol not in ROLES_GLOBALES:
        filtros.append(BolsaMoesImport.unidad == unidad_de(current_user))
    elif unidad:
        filtros.append(BolsaMoesImport.unidad == unidad)
    qq = select(BolsaMoesImport)
    if filtros:
        qq = qq.where(and_(*filtros))
    rows = (await db.execute(
        qq.order_by(BolsaMoesImport.fecha_version.desc().nullslast(), BolsaMoesImport.id.desc())
    )).scalars().all()
    return [{
        "id": i.id, "nombre": i.nombre, "hoja": i.hoja, "tamano": i.tamano, "cifrado": i.cifrado,
        "fecha_version": i.fecha_version.isoformat() if i.fecha_version else None, "aplicado": i.aplicado,
        "filas_total": i.filas_total, "filas_alta": i.filas_alta, "filas_mod": i.filas_mod,
        "filas_igual": i.filas_igual, "duplicados": i.duplicados,
        "columnas_nuevas": json.loads(i.columnas_nuevas) if i.columnas_nuevas else [],
        "subido_por": i.subido_por, "creado_en": i.creado_en.isoformat() if i.creado_en else None,
    } for i in rows]


@router.get("/importaciones/{imp_id}/descargar")
async def descargar_import(imp_id: int, db: AsyncSession = Depends(get_db),
                           current_user: Usuario = Depends(get_current_active_user)):
    i = (await db.execute(select(BolsaMoesImport).where(BolsaMoesImport.id == imp_id))).scalar_one_or_none()
    if not i:
        raise HTTPException(status_code=404, detail="Archivo no encontrado")
    if current_user.rol not in ROLES_GLOBALES and i.unidad != unidad_de(current_user):
        raise HTTPException(status_code=403, detail="Sin acceso a este archivo")
    ruta_abs = os.path.join(MOES_DIR, i.ruta or "")
    if not os.path.exists(ruta_abs):
        raise HTTPException(status_code=404, detail="El archivo ya no está en disco")
    return FileResponse(ruta_abs, filename=i.nombre,
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ── Edición manual de un ticket (campos amarillos) ───────────────────────────────

@router.patch("/{bolsa_id}")
async def actualizar(bolsa_id: int, cambios: dict = Body(...), db: AsyncSession = Depends(get_db),
                     current_user: Usuario = Depends(get_current_active_user)):
    assert_puede_escribir(current_user)
    b = (await db.execute(select(BolsaMoes).where(BolsaMoes.id == bolsa_id))).scalar_one_or_none()
    if not b:
        raise HTTPException(status_code=404, detail="No encontrado")
    if current_user.rol not in ROLES_GLOBALES and b.unidad != unidad_de(current_user):
        raise HTTPException(status_code=403, detail="Sin acceso")
    for campo, valor in cambios.items():
        if campo not in EDITABLES:
            continue
        antes = getattr(b, campo, None)
        if campo == "fecha_ejecucion":
            valor = moes._to_date(valor)
        elif campo == "pct_ejecutado":
            valor = moes._to_float(valor)
        if str(antes if antes is not None else "") == str(valor if valor is not None else ""):
            continue
        setattr(b, campo, _trunc(campo, valor))
        db.add(BolsaMoesCambio(
            unidad=b.unidad, bolsa_id=b.id, import_id=None, clave=b.clave, antecedente=b.antecedente,
            tipo="edicion", campo=moes.ETIQUETAS.get(campo, campo),
            valor_anterior=(None if antes is None else str(antes)),
            valor_nuevo=(None if valor is None else str(valor)),
            usuario_id=current_user.id, usuario_nombre=current_user.nombre))
    await db.commit()
    await db.refresh(b)
    return _to_dict(b)


# ── Export a Excel (resumen + datos + historial) ─────────────────────────────────

_AZUL = "1F4E79"
_AZUL2 = "2E75B6"


@router.get("/export")
async def exportar(unidad: Optional[str] = None, db: AsyncSession = Depends(get_db),
                   current_user: Usuario = Depends(get_current_active_user)):
    scope = _scope(current_user, unidad)
    base = select(BolsaMoes)
    if scope:
        base = base.where(and_(*scope))
    rows = (await db.execute(base.order_by(BolsaMoes.estado, BolsaMoes.id))).scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RESUMEN"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F2")
    ws["A1"] = "BOLSAS MOES — RESUMEN"
    ws["A1"].font = Font(bold=True, size=18, color="FFFFFF")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for row in ws["A1:F2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=_AZUL)
    ws["A3"] = f"Generado: {datetime.now():%d/%m/%Y %H:%M}  ·  {len(rows)} tickets"
    ws["A3"].font = Font(italic=True, color="666666")
    por_estado = Counter((b.estado or "SIN ESTADO").upper() for b in rows)
    ws["A5"] = "Estado"; ws["B5"] = "Tickets"
    ws["A5"].font = ws["B5"].font = Font(bold=True)
    for i, (e, n) in enumerate(por_estado.most_common(), start=6):
        ws.cell(row=i, column=1, value=e)
        ws.cell(row=i, column=2, value=n)

    # Hoja de datos: todas las columnas originales (union de datos_extra)
    cols, parsed = [], []
    for b in rows:
        d = json.loads(b.datos_extra) if b.datos_extra else {}
        parsed.append((b, d))
        for k in d:
            if k not in cols:
                cols.append(k)
    wd = wb.create_sheet("DATOS")
    wd.sheet_view.showGridLines = False
    encabezados = ["Estado", "Clave", "Duplicado"] + cols
    for j, c in enumerate(encabezados, 1):
        cell = wd.cell(row=1, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_AZUL2)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wd.freeze_panes = "A2"
    for i, (b, d) in enumerate(parsed, 2):
        wd.cell(row=i, column=1, value=b.estado)
        wd.cell(row=i, column=2, value=b.clave)
        wd.cell(row=i, column=3, value="SÍ" if b.duplicado else "")
        for j, c in enumerate(cols, 4):
            wd.cell(row=i, column=j, value=d.get(c))
    for j, c in enumerate(encabezados, 1):
        wd.column_dimensions[get_column_letter(j)].width = min(max(len(str(c)) + 2, 12), 40)

    # Historial
    hfiltros = []
    if current_user.rol not in ROLES_GLOBALES:
        hfiltros.append(BolsaMoesCambio.unidad == unidad_de(current_user))
    elif unidad:
        hfiltros.append(BolsaMoesCambio.unidad == unidad)
    hq = select(BolsaMoesCambio)
    if hfiltros:
        hq = hq.where(and_(*hfiltros))
    crows = (await db.execute(hq.order_by(BolsaMoesCambio.id.desc()).limit(5000))).scalars().all()
    wh = wb.create_sheet("HISTORIAL")
    wh.sheet_view.showGridLines = False
    for j, c in enumerate(["Fecha", "Tipo", "Antecedente", "Campo", "Antes", "Después", "Inconsistencia", "Usuario"], 1):
        cell = wh.cell(row=1, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_AZUL2)
    wh.freeze_panes = "A2"
    for i, c in enumerate(crows, 2):
        wh.cell(row=i, column=1, value=c.creado_en.strftime("%d/%m/%Y %H:%M") if c.creado_en else "")
        wh.cell(row=i, column=2, value=c.tipo)
        wh.cell(row=i, column=3, value=c.antecedente)
        wh.cell(row=i, column=4, value=c.campo)
        wh.cell(row=i, column=5, value=c.valor_anterior)
        wh.cell(row=i, column=6, value=c.valor_nuevo)
        wh.cell(row=i, column=7, value="SÍ" if c.inconsistencia else "")
        wh.cell(row=i, column=8, value=c.usuario_nombre)
    for j in range(1, 9):
        wh.column_dimensions[get_column_letter(j)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="bolsas_moes_{datetime.now():%Y%m%d}.xlsx"'},
    )
