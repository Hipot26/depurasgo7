"""
api/sgs.py — Relación con SGS (autocompletar MO real liquidada).
Acceso: el desarrollador o el supervisor de la unidad SABOGAL, vía `require_sgs_access`. SOLO opera sobre la unidad
SABOGAL. `relacionar` es solo lectura; `aplicar` escribe MO real + #Proliquidado
+ F. liq SGS en las filas aprobadas (auditado por el listener before_flush).
SGS Exporte (`/exporte` y `/exporte/excel`, solo consulta) va aparte, por
`require_sgs_exporte`: el desarrollador y las cuentas de USUARIOS_SGS_EXPORTE (2026-09-15).
Senda de meta (`/senda`, `/senda/metas`): avance de la MO liquidada frente a la
meta mensual por senda, con cortes cada 10 días (2026-09-17). Mismo gate.
"""
import calendar
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Optional, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fastapi import APIRouter, Depends, HTTPException
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import select, and_, func, delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.database import get_db
from app.api.deps import require_sgs_access, require_sgs_exporte
from app.models.models import Programacion, SgsMeta, Usuario, UNIDAD_DEFAULT
from app.services import sgs as sgs_svc

router = APIRouter(prefix="/sgs", tags=["sgs"])

UNIDAD = UNIDAD_DEFAULT  # "SABOGAL" — la relación con SGS es exclusiva de esa unidad


class RelacionarReq(BaseModel):
    mes: int
    anio: int
    unidad: Optional[str] = "SABOGAL"
    zona: Optional[str] = None   # None = todas las zonas de la unidad


class CambioReq(BaseModel):
    prog_id: int
    mo_real: Optional[float] = None
    num_proliquidado: Optional[str] = None
    fecha_liquidacion_sgs: Optional[str] = None  # ISO "YYYY-MM-DD"


class AplicarReq(BaseModel):
    cambios: List[CambioReq]


async def _cargar_prog_rows(db: AsyncSession, mes: int, anio: int, zona: Optional[str] = None) -> list:
    filtros = [
        Programacion.unidad == UNIDAD,
        func.extract("month", Programacion.fecha) == mes,
        func.extract("year", Programacion.fecha) == anio,
    ]
    if zona:
        filtros.append(Programacion.zona == zona)
    q = select(Programacion).where(and_(*filtros))
    rows = (await db.execute(q)).scalars().all()
    return [{
        "id":                   r.id,
        "fecha":                r.fecha.isoformat() if r.fecha else None,
        "lider":                r.lider,
        "lcl":                  r.lcl,
        "tipo_pdt":             r.tipo_pdt,
        "num_pdt":              r.num_pdt,
        "proyecto_antecedente": r.proyecto_antecedente,
        "num_proliquidado":     r.num_proliquidado,
        "mo_proyectada":        r.mo_proyectada,
        "mo_real":              r.mo_real,
        "supervisor":           r.supervisor,
        "zona":                 r.zona.value if r.zona else None,
    } for r in rows]


@router.post("/relacionar")
async def relacionar(
    req: RelacionarReq,
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(require_sgs_access),
):
    if not (1 <= req.mes <= 12) or req.anio < 2020:
        raise HTTPException(status_code=400, detail="Mes o año inválido")

    zona = req.zona or None
    zonas = [zona] if zona else None

    prog_rows = await _cargar_prog_rows(db, req.mes, req.anio, zona)

    # Consulta SGS (bloqueante de red) fuera del event loop.
    try:
        prest = await run_in_threadpool(sgs_svc.prestaciones_mes, req.anio, req.mes, zonas)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"No se pudo consultar SGS: {e}")

    propuesta = await run_in_threadpool(
        sgs_svc.cruzar, prog_rows, prest["nodos"], prest["total"])
    propuesta["periodo"] = f"{req.anio}-{req.mes:02d}"
    propuesta["zona"] = zona
    propuesta["zonas_disponibles"] = list(sgs_svc.SABOGAL_SCOPE.keys())
    propuesta["n_prog_total"] = len(prog_rows)
    return propuesta


# ── SGS Exporte — liquidaciones del mes por zona + descarga Excel ─────────────
# Misma fuente que el popup de Relación (prestaciones_mes), pero en modo NO
# estricto (requisito 2026-09-01): universo completo de SGS — sin
# whitelist de liquidador por zona y con los departamentos sin zona como
# "OTROS" — para que el total cuadre con el exporte de SGS. Solo consulta.

ZONAS_EXPORTE = list(sgs_svc.SABOGAL_SCOPE.keys()) + [sgs_svc.ZONA_OTROS]

MESES_ES = ["", "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
            "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]


class ExporteReq(BaseModel):
    mes: int
    anio: int
    zona: Optional[str] = None   # None = todas las zonas de la unidad


def _validar_exporte(req: ExporteReq):
    if not (1 <= req.mes <= 12) or req.anio < 2020:
        raise HTTPException(status_code=400, detail="Mes o año inválido")
    if req.zona and req.zona not in ZONAS_EXPORTE:
        raise HTTPException(status_code=400, detail=f"Zona desconocida: {req.zona}")


async def _prestaciones(req: ExporteReq) -> dict:
    zonas = [req.zona] if req.zona else None
    try:
        return await run_in_threadpool(
            sgs_svc.prestaciones_mes, req.anio, req.mes, zonas, False)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"No se pudo consultar SGS: {e}")


def _totales_zona(consolidado: list, detalle: list) -> list:
    orden = ZONAS_EXPORTE
    acc = {}
    for n in consolidado:
        t = acc.setdefault(n["zona"], {"zona": n["zona"], "total": 0.0, "n": 0, "n_prestaciones": 0})
        t["total"] += n["man_obra"]
        t["n"] += 1
    for d in detalle:
        if d["zona"] in acc:
            acc[d["zona"]]["n_prestaciones"] += 1
    # Sin redondear por zona: los avances traen 3 decimales y redondear cada zona
    # antes de sumar dejaba el TOTAL del RESUMEN un centavo por debajo del detalle.
    return sorted(acc.values(), key=lambda t: orden.index(t["zona"]) if t["zona"] in orden else 99)


def _clasificar_pdl(nodos_grupo: list) -> str:
    """Port de la clasificación PDL del Verificador (historial de nodos)."""
    if not nodos_grupo:
        return "Sin datos"
    if len(nodos_grupo) == 1:
        return "Nodo único"
    fechas = sorted(n["fecha_primera"] for n in nodos_grupo if n.get("fecha_primera"))
    if len(fechas) >= 2:
        d0 = datetime.strptime(fechas[0], "%Y-%m-%d")
        d1 = datetime.strptime(fechas[-1], "%Y-%m-%d")
        dias = (d1 - d0).days
        if dias <= 1:
            return "Multi-nodo simultáneo"
        return f"Multi-nodo secuencial (+{dias}d)"
    return "Multi-nodo (sin fechas)"


@router.post("/exporte")
async def exporte(
    req: ExporteReq,
    current_user: Usuario = Depends(require_sgs_exporte),
):
    _validar_exporte(req)
    prest = await _prestaciones(req)
    consolidado = await run_in_threadpool(sgs_svc.consolidar_nodos, prest["nodos"])
    for c in consolidado:
        c["clasificacion_pdl"] = _clasificar_pdl(c.get("nodos_grupo", []))
    detalle = prest.get("detalle", [])
    return {
        "periodo": f"{req.anio}-{req.mes:02d}",
        "consolidado": consolidado,
        "detalle": detalle,
        "total": prest["total"],
        "n": len(consolidado),
        "n_prestaciones": len(detalle),
        "totales_zona": _totales_zona(consolidado, detalle),
        "zonas_disponibles": ZONAS_EXPORTE,
    }


# Estilos compartidos del Excel de exporte (paleta del Verificador)
_THIN      = Side(style="thin", color="BFBFBF")
_BORDE     = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
_TOT_FILL  = PatternFill("solid", fgColor="DDEBF7")
_MI_FILL   = PatternFill("solid", fgColor="DDEBF7")   # Mayor Impacto → azul claro
_PG_FILL   = PatternFill("solid", fgColor="E2EFDA")   # Programado   → verde claro
_ROJO      = Font(color="C00000")
_FMT_MO    = '"$" #,##0.00'


def _encabezado(ws, cols, fila=1):
    for i, (titulo, ancho) in enumerate(cols, 1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.fill, c.font, c.border = _HEAD_FILL, _HEAD_FONT, _BORDE
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = f"A{fila + 1}"


def _fila_total(ws, fila, col_label, col_mo, total):
    lab = ws.cell(row=fila, column=col_label, value="TOTAL")
    tot = ws.cell(row=fila, column=col_mo, value=round(total, 2))
    tot.number_format = _FMT_MO
    for c in (lab, tot):
        c.font, c.fill, c.border = Font(bold=True), _TOT_FILL, _BORDE


def _hoja_consolidado(wb, nombre: str, filas: list):
    """Hoja estilo 'Consolidado' del Verificador: una fila por proyecto/DF/LCL,
    fondo azul MI / verde PG, estado ⚠ parcial cuando hay reversiones."""
    cols = [("FECHA", 12), ("PDL / OT", 22), ("LÍDER", 26), ("LCL", 13),
            ("PROC", 7), ("# COD ERP", 14), ("$ MAN OBRA", 16),
            ("LIQUIDADO POR", 22), ("ESTADO", 10), ("AVANCES", 9),
            ("NODOS", 26)]
    ws = wb.create_sheet(nombre[:31])
    _encabezado(ws, cols)
    for r, p in enumerate(filas, 2):
        fill = _MI_FILL if p["proceso"] == "MI" else _PG_FILL
        fecha = p.get("fecha_liq", "")
        if p.get("fecha_primera") and p["fecha_primera"] != fecha:
            fecha = f'{p["fecha_primera"]} → {fecha}'
        nodos_txt = " · ".join(n["pro_code"] for n in p.get("nodos_grupo", [])) \
            if len(p.get("nodos_grupo", [])) > 1 else ""
        vals = [fecha, p.get("pdl_code", ""), p.get("lider", ""), p.get("lcl", ""),
                p["proceso"], p["pro_code"], p["man_obra"], p.get("liquidado_por", ""),
                "anulado" if p.get("estado") == "anulado" else "⚠ parcial" if p.get("tiene_reversiones") else "ok",
                p.get("n_avances", 0), nodos_txt]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.fill, c.border = fill, _BORDE
            if i == 7:
                c.number_format = _FMT_MO
                if p["man_obra"] < 0:
                    c.font = _ROJO
            if i == 9 and p.get("tiene_reversiones"):
                c.font = _ROJO
    _fila_total(ws, len(filas) + 2, 6, 7, sum(f["man_obra"] for f in filas))


def _hoja_prestaciones(wb, filas: list):
    """Hoja 'Prestaciones': una fila por avance individual, negativas en rojo."""
    cols = [("FECHA", 12), ("ZONA", 11), ("DEPTO SGS", 22), ("# COD ERP", 14),
            ("PDL / OT", 22), ("NOMBRE PROYECTO", 38), ("LÍDER", 26), ("PROC", 7),
            ("LCL", 13), ("$ PRESTACIÓN", 16), ("LIQUIDADO POR", 26), ("ESTADO", 10)]
    ws = wb.create_sheet("Prestaciones")
    _encabezado(ws, cols)
    for r, d in enumerate(filas, 2):
        vals = [d["fecha"], d["zona"], d.get("nota_zona") or d.get("departamento", ""), d["pro_code"],
                d["pdl_code"], d["nombre_pro"], d["lider"], d["proceso"], d["lcl"],
                d["man_obra"], d["liquidador"], d["state"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = _BORDE
            if i == 10:
                c.number_format = _FMT_MO
                if d["man_obra"] < 0:
                    c.font = _ROJO
    _fila_total(ws, len(filas) + 2, 9, 10, sum(f["man_obra"] for f in filas))


def _hoja_analisis_pdl(wb, filas: list):
    """Hoja 'Análisis PDL': historial de nodos por contrato (clasificación)."""
    cols = [("FECHAS", 24), ("PDL / OT", 22), ("LCL", 13), ("PROC", 7),
            ("ZONA", 11), ("# NODOS", 9), ("CLASIFICACIÓN", 26),
            ("$ MAN OBRA", 16), ("DETALLE NODOS", 60)]
    ws = wb.create_sheet("Análisis PDL")
    _encabezado(ws, cols)
    for r, p in enumerate(filas, 2):
        nodos_g = p.get("nodos_grupo", [])
        fechas_p = sorted(n["fecha_primera"] for n in nodos_g if n.get("fecha_primera"))
        fechas_u = sorted(n["fecha_ultima"] for n in nodos_g if n.get("fecha_ultima"))
        f_min = fechas_p[0] if fechas_p else p.get("fecha_liq", "")
        f_max = fechas_u[-1] if fechas_u else p.get("fecha_liq", "")
        rango = f_min if f_min == f_max else f"{f_min} → {f_max}"
        det = " | ".join(
            f'{n["pro_code"]}{" N" + str(n["nodo_id"]) if n.get("nodo_id") else ""}'
            f' ${n["mo"]:,.0f}{" " + n["fecha_primera"] if n.get("fecha_primera") else ""}'
            for n in nodos_g)
        vals = [rango, p.get("pdl_code", ""), p.get("lcl", ""), p["proceso"], p["zona"],
                len(nodos_g), p.get("clasificacion_pdl", ""), p["man_obra"], det]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = _BORDE
            if i == 8:
                c.number_format = _FMT_MO
                if p["man_obra"] < 0:
                    c.font = _ROJO
    _fila_total(ws, len(filas) + 2, 7, 8, sum(f["man_obra"] for f in filas))


def _hoja_resumen(wb, totales: list, mes: int, anio: int, zona: Optional[str]):
    ws = wb.create_sheet("RESUMEN", 0)
    titulo = f"EXPORTE SGS — {MESES_ES[mes]} {anio}" + (f" — {zona}" if zona else "")
    ws.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=13)
    cols = [("ZONA", 14), ("PROYECTOS", 12), ("PRESTACIONES", 14), ("MO TOTAL", 20)]
    for i, (t, ancho) in enumerate(cols, 1):
        c = ws.cell(row=3, column=i, value=t)
        c.fill, c.font, c.border = _HEAD_FILL, _HEAD_FONT, _BORDE
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for r, t in enumerate(totales, 4):
        for i, v in enumerate([t["zona"], t["n"], t["n_prestaciones"], t["total"]], 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = _BORDE
            if i == 4:
                c.number_format = _FMT_MO
    ft = 4 + len(totales)
    vals = ["TOTAL", sum(t["n"] for t in totales), sum(t["n_prestaciones"] for t in totales),
            round(sum(t["total"] for t in totales), 2)]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=ft, column=i, value=v)
        c.font, c.fill, c.border = Font(bold=True), _TOT_FILL, _BORDE
        if i == 4:
            c.number_format = _FMT_MO


def _excel_exporte(consolidado: list, detalle: list, totales: list,
                   mes: int, anio: int, zona: Optional[str]) -> BytesIO:
    """Workbook estilo Verificador: RESUMEN + Consolidado (por zona u hoja única)
    + Prestaciones + Análisis PDL."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _hoja_resumen(wb, totales, mes, anio, zona)
    if zona:
        _hoja_consolidado(wb, zona, [p for p in consolidado if p["zona"] == zona])
    else:
        for z in ZONAS_EXPORTE:
            _hoja_consolidado(wb, z, [p for p in consolidado if p["zona"] == z])
    _hoja_prestaciones(wb, detalle)
    _hoja_analisis_pdl(wb, consolidado)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.post("/exporte/excel")
async def exporte_excel(
    req: ExporteReq,
    current_user: Usuario = Depends(require_sgs_exporte),
):
    _validar_exporte(req)
    prest = await _prestaciones(req)
    consolidado = await run_in_threadpool(sgs_svc.consolidar_nodos, prest["nodos"])
    for c in consolidado:
        c["clasificacion_pdl"] = _clasificar_pdl(c.get("nodos_grupo", []))
    detalle = prest.get("detalle", [])
    totales = _totales_zona(consolidado, detalle)
    buf = await run_in_threadpool(
        _excel_exporte, consolidado, detalle, totales, req.mes, req.anio, req.zona)
    zona_tag = f"{req.zona.replace('&', '')}_" if req.zona else ""
    filename = f"SGS_{zona_tag}{MESES_ES[req.mes]}_{req.anio}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/aplicar")
async def aplicar(
    req: AplicarReq,
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(require_sgs_access),
):
    if not req.cambios:
        return {"ok": True, "actualizadas": 0}

    ids = [c.prog_id for c in req.cambios]
    rows = (await db.execute(
        select(Programacion).where(Programacion.id.in_(ids)))).scalars().all()
    by_id = {r.id: r for r in rows}

    actualizadas = 0
    errores = []
    for c in req.cambios:
        prog = by_id.get(c.prog_id)
        if not prog:
            errores.append({"prog_id": c.prog_id, "error": "No encontrado"})
            continue
        # Blindaje: nunca tocar otra unidad (CHOCONTA, etc.)
        if (prog.unidad or UNIDAD_DEFAULT) != UNIDAD:
            errores.append({"prog_id": c.prog_id, "error": "Fuera de la unidad SABOGAL"})
            continue
        if c.mo_real is not None:
            prog.mo_real = c.mo_real
        if c.num_proliquidado:
            prog.num_proliquidado = c.num_proliquidado
        if c.fecha_liquidacion_sgs:
            try:
                prog.fecha_liquidacion_sgs = datetime.strptime(
                    c.fecha_liquidacion_sgs[:10], "%Y-%m-%d").date()
            except ValueError:
                pass
        prog.actualizado_por = current_user.id
        actualizadas += 1

    await db.commit()
    return {"ok": True, "actualizadas": actualizadas, "errores": errores}


# ── Senda de meta — avance de la MO liquidada frente a la meta mensual ─────────
# Requisito 2026-09-17: UBATE TP debe llegar a $520 M, LV a $184 M y MOE a
# $36 M; medir cada 10 días (día 10 / 20 / fin de mes) qué % de la meta lleva SGS.
# Las metas NO coinciden con la programación de septiembre 2026 (UBATE $408 M,
# LV&TLC $195 M, cuadrillas MOE $37,7 M), así que quedan como variables por mes en
# `sgs_metas`. «Cuadrillas MOE» = líderes cuyo tipo_grupo en la programación del
# mes empieza por MOE (MOE / MOE2); en SGS liquidan bajo la zona del departamento
# (UBATE, LV&TLC, ALMEIDAS…), por eso UBATE TP y LV se miden SIN esas cuadrillas
# y MOE con ellas en cualquier zona — para no contar dos veces la misma MO.
# Senda esperada: lineal por día calendario (día 10 de 30 → 33,3 %).

CUADRILLAS_OPCIONES = ("todas", "moe", "sin_moe")
PROCESOS_VALIDOS = ("MI", "PG")

SENDA_DEFAULT = [
    {"clave": "ubate_tp", "nombre": "UBATE TP", "meta": 520_000_000.0,
     "zonas": "UBATE", "procesos": "", "cuadrillas": "sin_moe", "orden": 1},
    {"clave": "lv",       "nombre": "LV",       "meta": 184_000_000.0,
     "zonas": "LV&TLC", "procesos": "", "cuadrillas": "sin_moe", "orden": 2},
    {"clave": "moe",      "nombre": "MOE",      "meta": 36_000_000.0,
     "zonas": "", "procesos": "", "cuadrillas": "moe", "orden": 3},
]

# Cortes de la senda: día 10, día 20 y último día del mes.
CORTES_DIAS = (10, 20)


class MetaIn(BaseModel):
    clave: str
    nombre: str
    meta: float
    zonas: List[str] = []
    procesos: List[str] = []
    cuadrillas: str = "todas"


class MetasReq(BaseModel):
    mes: int
    anio: int
    metas: List[MetaIn]


def _meta_dict(m, origen: str, heredado_de: Optional[str] = None) -> dict:
    g = (lambda k: m[k]) if isinstance(m, dict) else (lambda k: getattr(m, k))
    return {
        "clave": g("clave"), "nombre": g("nombre"), "meta": float(g("meta") or 0),
        "zonas": [z for z in (g("zonas") or "").split(",") if z],
        "procesos": [p for p in (g("procesos") or "").split(",") if p],
        "cuadrillas": g("cuadrillas") or "todas",
        "orden": g("orden") or 0,
        "origen": origen,            # guardada | heredada | defecto
        "heredado_de": heredado_de,  # "YYYY-MM" cuando origen = heredada
    }


async def _cargar_metas(db: AsyncSession, mes: int, anio: int) -> list:
    """Metas del mes. Sin filas → hereda del último mes configurado antes de este;
    sin ninguno → defaults del código. Solo se persisten al guardar."""
    q = select(SgsMeta).where(SgsMeta.anio == anio, SgsMeta.mes == mes).order_by(SgsMeta.orden, SgsMeta.id)
    rows = (await db.execute(q)).scalars().all()
    if rows:
        return [_meta_dict(r, "guardada") for r in rows]
    q = (select(SgsMeta.anio, SgsMeta.mes)
         .where((SgsMeta.anio < anio) | ((SgsMeta.anio == anio) & (SgsMeta.mes < mes)))
         .order_by(SgsMeta.anio.desc(), SgsMeta.mes.desc()).limit(1))
    prev = (await db.execute(q)).first()
    if prev:
        q = (select(SgsMeta).where(SgsMeta.anio == prev[0], SgsMeta.mes == prev[1])
             .order_by(SgsMeta.orden, SgsMeta.id))
        rows = (await db.execute(q)).scalars().all()
        return [_meta_dict(r, "heredada", f"{prev[0]}-{prev[1]:02d}") for r in rows]
    return [_meta_dict(m, "defecto") for m in SENDA_DEFAULT]


async def _cuadrillas_moe(db: AsyncSession, mes: int, anio: int) -> list:
    """Líderes de cuadrillas MOE en la programación de la unidad SABOGAL del mes (tipo_grupo MOE%)."""
    q = (select(Programacion.lider, Programacion.tipo_grupo, Programacion.zona,
                func.count(Programacion.id), func.sum(Programacion.mo_proyectada))
         .where(Programacion.unidad == UNIDAD,
                func.extract("month", Programacion.fecha) == mes,
                func.extract("year", Programacion.fecha) == anio,
                Programacion.tipo_grupo.ilike("MOE%"))
         .group_by(Programacion.lider, Programacion.tipo_grupo, Programacion.zona)
         .order_by(Programacion.lider))
    rows = (await db.execute(q)).all()
    return [{"lider": r[0] or "", "tipo_grupo": r[1], "zona": r[2].value if r[2] else None,
             "n": int(r[3] or 0), "mo_proyectada": float(r[4] or 0)} for r in rows if r[0]]


def _tokens(nombre: str) -> list:
    return [t for t in sgs_svc._norm_txt(nombre).replace(".", " ").split() if len(t) >= 3]


def _mismo_lider(corto: str, largo: str) -> bool:
    """La programación escribe «NOMBRE APELLIDO»; SGS «APELLIDO APELLIDO NOMBRE NOMBRE».
    Casan si cada token del nombre corto aparece en el largo (igual o con los
    mismos 5 primeros caracteres: OSPINA ↔ OSPINO)."""
    tc, tl = _tokens(corto), _tokens(largo)
    if not tc or not tl:
        return False
    return all(any(a == b or (len(a) >= 5 and len(b) >= 5 and a[:5] == b[:5]) for b in tl) for a in tc)


def _lider_es_moe(lider_sgs: str, cuadrillas_moe: list, cache: dict) -> Optional[str]:
    """Devuelve el nombre de la cuadrilla MOE (programación) que casa con el líder SGS, o None."""
    if lider_sgs in cache:
        return cache[lider_sgs]
    hit = next((c["lider"] for c in cuadrillas_moe if _mismo_lider(c["lider"], lider_sgs)), None)
    cache[lider_sgs] = hit
    return hit


def _dias_habiles(desde: date, hasta: date) -> int:
    """Días lunes–viernes en [desde, hasta]."""
    n, d = 0, desde
    while d <= hasta:
        if d.weekday() < 5:
            n += 1
        d += timedelta(days=1)
    return n


def _estado_corte(pct: float, esperado: float) -> str:
    if pct >= 100:
        return "cumplido"
    if pct >= esperado:
        return "en_senda"
    if esperado - pct <= 5:
        return "ajustado"
    return "rezago"


def _calcular_senda(meta: dict, detalle: list, cuadrillas_moe: list, cache: dict,
                    mes: int, anio: int, hoy: date) -> dict:
    dias_mes = calendar.monthrange(anio, mes)[1]
    fin_mes = date(anio, mes, dias_mes)
    zonas = set(meta["zonas"])
    procesos = set(meta["procesos"])
    modo = meta["cuadrillas"]

    por_dia = [0.0] * (dias_mes + 1)   # índice = día del mes
    lideres: dict = {}
    for d in detalle:
        if zonas and d["zona"] not in zonas:
            continue
        if procesos and d["proceso"] not in procesos:
            continue
        cuad = _lider_es_moe(d["lider"], cuadrillas_moe, cache) if modo != "todas" else None
        if modo == "moe" and not cuad:
            continue
        if modo == "sin_moe" and cuad:
            continue
        if not d["fecha"]:
            continue
        dia = int(d["fecha"][8:10])
        por_dia[dia] += d["man_obra"]
        l = lideres.setdefault(d["lider"] or "—", {"lider": d["lider"] or "—", "mo": 0.0, "n": 0,
                                                    "zonas": set(), "cuadrilla_moe": cuad})
        l["mo"] += d["man_obra"]
        l["n"] += 1
        l["zonas"].add(d["zona"])

    meta_val = meta["meta"] or 0.0
    pct = (lambda v: round(v / meta_val * 100, 1) if meta_val else 0.0)
    acumulado, serie = 0.0, []
    for dia in range(1, dias_mes + 1):
        acumulado += por_dia[dia]
        f = date(anio, mes, dia)
        serie.append({
            "dia": dia, "fecha": f.isoformat(),
            "acumulado": round(acumulado, 2) if f <= hoy else None,
            "pct": pct(acumulado) if f <= hoy else None,
            "esperado_pct": round(dia / dias_mes * 100, 1),
            "esperado": round(meta_val * dia / dias_mes, 2),
        })

    cortes = []
    for dia in list(CORTES_DIAS) + [dias_mes]:
        pt = serie[dia - 1]
        pendiente = pt["acumulado"] is None
        cortes.append({
            "dia": dia, "fecha": pt["fecha"], "es_fin_mes": dia == dias_mes,
            "acumulado": pt["acumulado"], "pct": pt["pct"],
            "esperado": pt["esperado"], "esperado_pct": pt["esperado_pct"],
            "delta_pp": None if pendiente else round(pt["pct"] - pt["esperado_pct"], 1),
            "estado": "pendiente" if pendiente else _estado_corte(pt["pct"], pt["esperado_pct"]),
        })

    # Situación a hoy (mes en curso) o al cierre (mes pasado)
    if hoy < date(anio, mes, 1):
        corte_hoy = None
    else:
        ref = min(hoy, fin_mes)
        dia_ref = ref.day
        pt = serie[dia_ref - 1]
        acum = pt["acumulado"] or 0.0
        ritmo = acum / dia_ref
        faltante = max(meta_val - acum, 0.0)
        restantes = max(dias_mes - dia_ref, 0)
        habiles = _dias_habiles(ref + timedelta(days=1), fin_mes) if ref < fin_mes else 0
        corte_hoy = {
            "dia": dia_ref, "fecha": ref.isoformat(), "cerrado": hoy > fin_mes,
            "acumulado": round(acum, 2), "pct": pct(acum),
            "esperado_pct": pt["esperado_pct"], "delta_pp": round(pct(acum) - pt["esperado_pct"], 1),
            "estado": _estado_corte(pct(acum), pt["esperado_pct"]),
            "ritmo_dia": round(ritmo, 2),
            "proyeccion": round(ritmo * dias_mes, 2), "proyeccion_pct": pct(ritmo * dias_mes),
            "faltante": round(faltante, 2),
            "dias_restantes": restantes, "habiles_restantes": habiles,
            "requerido_habil": round(faltante / habiles, 2) if habiles else None,
            "requerido_dia": round(faltante / restantes, 2) if restantes else None,
        }

    lid_list = sorted(lideres.values(), key=lambda l: l["mo"], reverse=True)
    for l in lid_list:
        l["mo"] = round(l["mo"], 2)
        l["zonas"] = sorted(l["zonas"])
    return {
        **meta,
        "dias_mes": dias_mes,
        "total": round(acumulado, 2), "pct_total": pct(acumulado),
        "serie": serie, "cortes": cortes, "hoy": corte_hoy,
        "lideres": lid_list, "n_prestaciones": sum(l["n"] for l in lid_list),
    }


@router.get("/senda")
async def senda(
    mes: int, anio: int,
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(require_sgs_exporte),
):
    req = ExporteReq(mes=mes, anio=anio)
    _validar_exporte(req)
    metas = await _cargar_metas(db, mes, anio)
    cuadrillas_moe = await _cuadrillas_moe(db, mes, anio)
    prest = await _prestaciones(req)
    detalle = prest.get("detalle", [])
    hoy = date.today()
    cache: dict = {}
    sendas = [_calcular_senda(m, detalle, cuadrillas_moe, cache, mes, anio, hoy) for m in metas]

    # Cuadrillas MOE de la programación: cuáles casaron con algún líder SGS del mes
    lideres_sgs = {d["lider"] for d in detalle if d["lider"]}
    for c in cuadrillas_moe:
        c["en_sgs"] = sorted({l for l in lideres_sgs if _mismo_lider(c["lider"], l)})

    total_sgs = round(sum(d["man_obra"] for d in detalle), 2)
    total_sendas = round(sum(s["total"] for s in sendas), 2)
    meta_total = round(sum(s["meta"] for s in sendas), 2)
    fechas = sorted({d["fecha"] for d in detalle if d["fecha"]})
    dias_mes = calendar.monthrange(anio, mes)[1]
    return {
        "periodo": f"{anio}-{mes:02d}",
        "hoy": hoy.isoformat(),
        "dias_mes": dias_mes,
        "ultima_fecha_sgs": fechas[-1] if fechas else None,
        "cortes_dias": list(CORTES_DIAS) + [dias_mes],
        "metas_origen": metas[0]["origen"] if metas else "defecto",
        "metas_heredadas_de": metas[0].get("heredado_de") if metas else None,
        "sendas": sendas,
        "cuadrillas_moe": cuadrillas_moe,
        "total_sgs": total_sgs,
        "total_sendas": total_sendas,
        "fuera_de_senda": round(total_sgs - total_sendas, 2),
        "meta_total": meta_total,
        "pct_total": round(total_sendas / meta_total * 100, 1) if meta_total else 0.0,
        "zonas_disponibles": ZONAS_EXPORTE,
        "cuadrillas_opciones": list(CUADRILLAS_OPCIONES),
    }


@router.put("/senda/metas")
async def guardar_metas(
    req: MetasReq,
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(require_sgs_exporte),
):
    """Reemplaza las metas del mes (las variables de la senda). Se auditan por el
    listener before_flush como cualquier escritura."""
    _validar_exporte(ExporteReq(mes=req.mes, anio=req.anio))
    if not req.metas:
        raise HTTPException(status_code=400, detail="Debe quedar al menos una senda")
    claves = set()
    for m in req.metas:
        clave = m.clave.strip().lower()
        if not clave or clave in claves:
            raise HTTPException(status_code=400, detail=f"Clave repetida o vacía: {m.clave!r}")
        claves.add(clave)
        if not m.nombre.strip():
            raise HTTPException(status_code=400, detail="Toda senda necesita nombre")
        if m.meta < 0:
            raise HTTPException(status_code=400, detail=f"Meta negativa en {m.nombre}")
        if m.cuadrillas not in CUADRILLAS_OPCIONES:
            raise HTTPException(status_code=400, detail=f"Cuadrillas inválidas: {m.cuadrillas}")
        for z in m.zonas:
            if z not in ZONAS_EXPORTE:
                raise HTTPException(status_code=400, detail=f"Zona desconocida: {z}")
        for p in m.procesos:
            if p not in PROCESOS_VALIDOS:
                raise HTTPException(status_code=400, detail=f"Proceso desconocido: {p}")

    await db.execute(delete(SgsMeta).where(SgsMeta.anio == req.anio, SgsMeta.mes == req.mes))
    for i, m in enumerate(req.metas, 1):
        db.add(SgsMeta(
            anio=req.anio, mes=req.mes, clave=m.clave.strip().lower(), nombre=m.nombre.strip()[:60],
            meta=float(m.meta), zonas=",".join(m.zonas), procesos=",".join(m.procesos),
            cuadrillas=m.cuadrillas, orden=i, actualizado_por=current_user.email,
        ))
    await db.commit()
    return {"ok": True, "metas": await _cargar_metas(db, req.mes, req.anio)}
