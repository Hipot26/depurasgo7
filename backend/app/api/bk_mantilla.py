"""Importación de la plantilla Mantilla → programación de la unidad CHOCONTA (no programados).
Cruza por código en notas y, si no, por el nombre de líder del encabezado
(CÓDIGO_NOMBRE) acotado a la zona del prefijo; actualiza # PDT, MO REAL e
INOPERATIVIDAD y estampa el código en notas y en el catálogo de líderes.
Devuelve resumen + (si hay errores) el mismo Excel con las celdas sin cruce en rojo."""
import io
import base64
from collections import Counter
from typing import Optional
from fastapi import APIRouter, Depends, UploadFile, File, Form, HTTPException
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
import openpyxl
from openpyxl.styles import PatternFill
from app.db.database import get_db
from app.api.deps import get_current_active_user, assert_puede_escribir, ROLES_GLOBALES, unidad_de
from app.models.models import Programacion, Usuario
from app.services.mantilla import parse_seguimiento, cruzar_mantilla

router = APIRouter(prefix="/mantilla", tags=["mantilla"])


@router.post("/import")
async def importar_mantilla(
    file: UploadFile = File(...),
    unidad: Optional[str] = Form(None),
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    assert_puede_escribir(current_user)
    uni = unidad if (current_user.rol in ROLES_GLOBALES and unidad) else unidad_de(current_user)

    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="El archivo debe ser .xlsx")
    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo: {e}")

    celdas = parse_seguimiento(wb)
    if not celdas:
        raise HTTPException(status_code=400, detail="No se encontró la hoja 'Seguimiento Pisos' o no tiene datos")

    actualizadas, sin_match = await cruzar_mantilla(db, uni, celdas)  # [(celda, motivo|None)]
    await db.commit()

    # ── Motivo específico de cada celda sin cruce ──────────────────────────────
    codigos_sis = set((await db.execute(
        select(Programacion.notas).where(Programacion.unidad == uni, Programacion.notas.isnot(None)).distinct()
    )).scalars().all())
    fechas_sis = set((await db.execute(
        select(Programacion.fecha).where(Programacion.unidad == uni).distinct()
    )).scalars().all())

    def _motivo(c, especifico):
        if especifico:
            return especifico
        if c["fecha"] not in fechas_sis:
            return "No hay programación cargada para esa fecha"
        if c["codigo"] not in codigos_sis:
            return "Código desconocido y el nombre del encabezado no coincide con ningún líder del día"
        return "La cuadrilla no tiene fila ese día (no cruzó ni por código ni por nombre)"

    motivos_count = Counter(_motivo(c, m) for c, m in sin_match)

    # Informe de errores: el MISMO Excel con las celdas sin cruce en rojo + nota del motivo
    informe_b64 = None
    if sin_match:
        wb2 = openpyxl.load_workbook(io.BytesIO(contents))  # copia editable
        sp = wb2["Seguimiento Pisos"]
        rojo = PatternFill("solid", fgColor="FFC7CE")
        for c, m in sin_match:
            for rr in range(c["fila"], c["fila"] + 7):
                for cc in range(c["col"], c["col"] + 3):
                    sp.cell(row=rr, column=cc).fill = rojo
            # comentario con el motivo en la primera celda del bloque
            try:
                from openpyxl.comments import Comment
                sp.cell(row=c["fila"], column=c["col"]).comment = Comment(_motivo(c, m), "Validación")
            except Exception:
                pass
        buf = io.BytesIO()
        wb2.save(buf)
        buf.seek(0)
        informe_b64 = base64.b64encode(buf.read()).decode()

    return {
        "unidad": uni,
        "total_celdas": len(celdas),
        "actualizadas": actualizadas,
        "sin_match": len(sin_match),
        "sin_match_detalle": [
            {"codigo": c["codigo"], "lider": c.get("lider") or "", "fecha": c["fecha"].isoformat(),
             "motivo": _motivo(c, m)}
            for c, m in sin_match[:300]
        ],
        "motivos": [{"motivo": m, "total": n} for m, n in motivos_count.most_common()],
        "informe_b64": informe_b64,
        "informe_nombre": f"mantilla_errores_{uni.lower()}.xlsx" if informe_b64 else None,
    }
