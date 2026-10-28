"""Módulo de Bolsas: casos por tipo (TUTELA, PQR, IPPE, CUN100, AMP, NNDD, TLC…).
Diseño núcleo común + datos_extra (JSON) → importación automática de cualquier
hoja sin perder columnas. De momento se expone solo en la cuenta dev (gate en UI)."""
import io
import re
import json
from typing import Optional
from datetime import date, datetime
from collections import Counter
from fastapi import APIRouter, Depends, Query, UploadFile, File, Form, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_, or_, func, delete
import openpyxl
from openpyxl.chart import PieChart, BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from app.db.database import get_db
from app.api.deps import get_current_active_user, assert_puede_escribir
from app.models.models import Bolsa, Usuario
from app.schemas.schemas import BolsaCreate, BolsaUpdate

router = APIRouter(prefix="/bolsas", tags=["bolsas"])

# Tipos canónicos (orden de presentación)
TIPOS_BOLSA = ["TUTELA", "PQR", "BOCATOMA", "IPPE", "CUN100", "CUNBT", "AMP", "NNDD", "TLC", "DRONE", "BP"]

# Hoja del Excel → tipo canónico (lo demás se ignora: Hoja1, Hoja2…)
SHEET_TIPO = {
    "TUTELA": "TUTELA", "PQR": "PQR", "CASOS BOCATOMA": "BOCATOMA", "IPPE": "IPPE",
    "CUN100": "CUN100", "BP": "BP", "AMP": "AMP", "CUNBT": "CUNBT", "NNDD": "NNDD",
    "DRONE": "DRONE", "TLC": "TLC",
}


def _nh(s) -> str:
    s = str(s or "").strip().lower()
    for a, b in [("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ñ", "n")]:
        s = s.replace(a, b)
    return s


# Encabezado normalizado → campo del núcleo
CORE_MAP = {
    "zona": "zona", "zona+": "zona",
    "municipio": "municipio",
    "circuito": "circuito", "cto": "circuito", "codigo cto": "circuito",
    "lcl": "lcl", "lcl asignada": "lcl", "lcl ejecucion": "lcl", "lcl ejecución": "lcl",
    "pdl": "pdl",
    "estado": "estado", "estado dominion": "estado", "estado revisado": "estado",
    "estado ostk": "estado", "estado equipo": "estado",
    "supervisor": "supervisor", "supervisor responsable": "supervisor",
    "responsable": "supervisor", "responsable bra": "supervisor",
    "fecha de prog": "fecha_prog", "fecha de programacion": "fecha_prog",
    "fecha de program": "fecha_prog", "fecha programacion": "fecha_prog",
    "fecha de programa": "fecha_prog", "fecha prgramación": "fecha_prog",
    "fecha entrega lc": "fecha_entrega", "fecha entrega ee": "fecha_entrega",
    "fecha entrega": "fecha_entrega",
    "observacion": "observacion", "observaciones": "observacion", "observación": "observacion",
}

# Columna identificadora (caso) por tipo, normalizada
CASO_POR_TIPO = {
    "TUTELA": "caso", "PQR": "pqrs", "BOCATOMA": "num. caso", "IPPE": "qr",
    "CUN100": "nombre", "AMP": "id", "CUNBT": "cd", "NNDD": "radicado",
    "DRONE": "activo", "TLC": "ticket nuevo", "BP": "caso",
}

_MAXLEN = {"zona": 100, "caso": 400, "municipio": 150, "circuito": 150, "lcl": 100, "pdl": 100, "estado": 100, "supervisor": 150}


def _to_date(v):
    """Acepta datetime/date o texto ISO/DD-MM-YYYY (input del formulario)."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _cell(v):
    """Valor JSON-serializable para datos_extra."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return v
    return str(v).strip()


def mapear_core(tipo: str, fila: dict) -> dict:
    """De {encabezado: valor} → campos del núcleo (zona, caso, estado, fechas…).
    Misma lógica para importar y para crear/editar a mano."""
    core: dict = {}
    caso_h = CASO_POR_TIPO.get(tipo)
    for h, val in fila.items():
        if val is None or str(val).strip() == "":
            continue
        nh = _nh(h)
        cf = CORE_MAP.get(nh)
        if cf and cf not in core:
            if cf in ("fecha_prog", "fecha_entrega"):
                d = _to_date(val)
                if d:
                    core[cf] = d
            else:
                core[cf] = str(val).strip()[:_MAXLEN.get(cf, 200)]
        if caso_h and nh == caso_h and "caso" not in core:
            core["caso"] = str(val).strip()[:400]
    return core


def parse_bolsas(wb) -> list[dict]:
    out: list[dict] = []
    for sheet, tipo in SHEET_TIPO.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = [r for r in ws.iter_rows(values_only=True)]
        # fila de encabezado = primera (de las 6 primeras) con ≥3 celdas de texto
        hidx = None
        for i, r in enumerate(rows[:6]):
            if r and sum(1 for c in r if isinstance(c, str) and c.strip()) >= 3:
                hidx = i
                break
        if hidx is None:
            continue
        headers = [str(c).strip() if c is not None else "" for c in rows[hidx]]

        for r in rows[hidx + 1:]:
            if not r or all(c is None or str(c).strip() == "" for c in r):
                continue
            # {encabezado: valor} de celdas no vacías
            fila = {
                headers[j]: val
                for j, val in enumerate(r)
                if j < len(headers) and headers[j] and val is not None and str(val).strip() != ""
            }
            if not fila:
                continue
            core = mapear_core(tipo, fila)
            core["tipo"] = tipo
            core["datos_extra"] = {h: _cell(v) for h, v in fila.items()}
            out.append(core)
    return out


def _to_dict(b: Bolsa) -> dict:
    return {
        "id": b.id, "unidad": b.unidad, "tipo": b.tipo, "zona": b.zona, "caso": b.caso,
        "municipio": b.municipio, "circuito": b.circuito, "lcl": b.lcl, "pdl": b.pdl,
        "estado": b.estado, "supervisor": b.supervisor,
        "fecha_entrega": b.fecha_entrega.isoformat() if b.fecha_entrega else None,
        "fecha_prog": b.fecha_prog.isoformat() if b.fecha_prog else None,
        "observacion": b.observacion,
        "datos_extra": json.loads(b.datos_extra) if b.datos_extra else {},
    }


# ── Catálogo / resumen ──────────────────────────────────────────────────────────

@router.get("/tipos")
async def tipos_bolsa(_: Usuario = Depends(get_current_active_user)):
    return TIPOS_BOLSA


# Campos por defecto si un tipo aún no tiene datos cargados
CAMPOS_FALLBACK = ["ZONA", "CASO", "MUNICIPIO", "CIRCUITO", "LCL", "PDL",
                   "ESTADO", "SUPERVISOR", "FECHA DE PROG", "FECHA ENTREGA LC", "OBSERVACIÓN"]


@router.get("/campos")
async def campos_tipo(tipo: str, db: AsyncSession = Depends(get_db), _: Usuario = Depends(get_current_active_user)):
    """Columnas que maneja un tipo (deducidas de los datos existentes) para armar
    el formulario dinámico. Si el tipo no tiene datos, usa un set por defecto."""
    rows = (await db.execute(
        select(Bolsa.datos_extra).where(Bolsa.tipo == tipo).limit(300)
    )).scalars().all()
    seen: list[str] = []
    for r in rows:
        if not r:
            continue
        try:
            d = json.loads(r)
        except Exception:
            continue
        for k in d.keys():
            if k not in seen:
                seen.append(k)
    return {"tipo": tipo, "campos": seen or CAMPOS_FALLBACK}


@router.get("/resumen")
async def resumen(db: AsyncSession = Depends(get_db), _: Usuario = Depends(get_current_active_user)):
    total = (await db.execute(select(func.count()).select_from(Bolsa))).scalar() or 0
    por_tipo = [
        {"tipo": t, "total": n}
        for t, n in (await db.execute(
            select(Bolsa.tipo, func.count()).group_by(Bolsa.tipo).order_by(func.count().desc())
        )).all()
    ]
    por_estado = [
        {"estado": (e or "SIN ESTADO"), "total": n}
        for e, n in (await db.execute(
            select(func.upper(func.trim(Bolsa.estado)), func.count())
            .group_by(func.upper(func.trim(Bolsa.estado))).order_by(func.count().desc())
        )).all()
    ]
    return {"total": total, "por_tipo": por_tipo, "por_estado": por_estado}


# ── Filtros compartidos (lista + reporte) ──────────────────────────────────────

class BolsaFiltros:
    def __init__(
        self,
        tipo: Optional[str] = None,
        zona: Optional[str] = None,
        estado: Optional[str] = None,
        supervisor: Optional[str] = None,
        municipio: Optional[str] = None,
        q: Optional[str] = None,
        fecha_inicio: Optional[date] = None,
        fecha_fin: Optional[date] = None,
    ):
        self.tipo = tipo
        self.zona = zona
        self.estado = estado
        self.supervisor = supervisor
        self.municipio = municipio
        self.q = q
        self.fecha_inicio = fecha_inicio
        self.fecha_fin = fecha_fin

    def filtros(self) -> list:
        f: list = []
        if self.tipo:
            f.append(Bolsa.tipo == self.tipo)
        if self.zona:
            f.append(Bolsa.zona == self.zona)
        if self.estado:
            f.append(func.upper(func.trim(Bolsa.estado)) == self.estado.strip().upper())
        if self.supervisor:
            f.append(Bolsa.supervisor.ilike(f"%{self.supervisor}%"))
        if self.municipio:
            f.append(Bolsa.municipio.ilike(f"%{self.municipio}%"))
        if self.q:
            like = f"%{self.q}%"
            f.append(or_(
                Bolsa.caso.ilike(like), Bolsa.municipio.ilike(like),
                Bolsa.observacion.ilike(like), Bolsa.pdl.ilike(like), Bolsa.lcl.ilike(like),
                Bolsa.circuito.ilike(like), Bolsa.supervisor.ilike(like), Bolsa.zona.ilike(like),
                Bolsa.datos_extra.ilike(like),   # busca también en las columnas propias del tipo
            ))
        if self.fecha_inicio:
            f.append(Bolsa.fecha_prog >= self.fecha_inicio)
        if self.fecha_fin:
            f.append(Bolsa.fecha_prog <= self.fecha_fin)
        return f


# ── Listado ─────────────────────────────────────────────────────────────────────

@router.get("")
async def listar(
    f: BolsaFiltros = Depends(),
    skip: int = Query(0, ge=0),
    limit: int = Query(3000, le=5000),
    db: AsyncSession = Depends(get_db),
    _: Usuario = Depends(get_current_active_user),
):
    filters = f.filtros()
    base = select(Bolsa)
    if filters:
        base = base.where(and_(*filters))
    total = (await db.execute(select(func.count()).select_from(base.subquery()))).scalar() or 0
    items = (await db.execute(base.order_by(Bolsa.tipo, Bolsa.id).offset(skip).limit(limit))).scalars().all()
    return {"total": total, "items": [_to_dict(b) for b in items]}


# ── Reporte gráfico: KPIs y opciones de filtro ──────────────────────────────────

@router.get("/opciones")
async def bolsas_opciones(db: AsyncSession = Depends(get_db), _: Usuario = Depends(get_current_active_user)):
    async def dist(col, limite=None):
        q = (
            select(func.trim(col)).where(col.isnot(None), func.trim(col) != "")
            .group_by(func.trim(col)).order_by(func.count().desc())
        )
        vals = [r[0] for r in (await db.execute(q)).all() if r[0]]
        return vals[:limite] if limite else vals

    eq = (
        select(func.upper(func.trim(Bolsa.estado)))
        .where(Bolsa.estado.isnot(None), func.trim(Bolsa.estado) != "")
        .group_by(func.upper(func.trim(Bolsa.estado))).order_by(func.count().desc())
    )
    estados = [r[0] for r in (await db.execute(eq)).all() if r[0]]
    return {
        "tipos": TIPOS_BOLSA,
        "zonas": await dist(Bolsa.zona),
        "estados": estados,
        "supervisores": await dist(Bolsa.supervisor, 300),
        "municipios": await dist(Bolsa.municipio, 300),
    }


async def _agregaciones(db: AsyncSession, base: list) -> dict:
    def W(q):
        return q.where(and_(*base)) if base else q

    total = (await db.execute(W(select(func.count()).select_from(Bolsa)))).scalar() or 0

    async def grupo(col, norm=False, limite=None):
        expr = func.upper(func.trim(col)) if norm else col
        q = W(select(expr.label("k"), func.count().label("n")).group_by(expr).order_by(func.count().desc()))
        rows = (await db.execute(q)).all()
        out = [{"label": (r.k if r.k not in (None, "") else "—"), "total": r.n} for r in rows]
        return out[:limite] if limite else out

    por_tipo = await grupo(Bolsa.tipo)
    por_estado = await grupo(Bolsa.estado, norm=True)
    por_zona = await grupo(Bolsa.zona)
    por_supervisor = await grupo(Bolsa.supervisor, limite=12)
    por_municipio = await grupo(Bolsa.municipio, limite=12)

    mexpr = func.to_char(Bolsa.fecha_prog, "YYYY-MM")
    tq = (
        select(mexpr.label("m"), func.count().label("n"))
        .where(and_(*(base + [Bolsa.fecha_prog.isnot(None)])))
        .group_by(mexpr).order_by(mexpr)
    )
    timeline = [{"periodo": r.m, "total": r.n} for r in (await db.execute(tq)).all()]
    con_fecha = sum(t["total"] for t in timeline)

    return {
        "total": total, "por_tipo": por_tipo, "por_estado": por_estado, "por_zona": por_zona,
        "por_supervisor": por_supervisor, "por_municipio": por_municipio, "timeline": timeline,
        "con_fecha_prog": con_fecha, "sin_fecha_prog": total - con_fecha,
    }


@router.get("/kpis")
async def bolsas_kpis(
    f: BolsaFiltros = Depends(),
    db: AsyncSession = Depends(get_db),
    _: Usuario = Depends(get_current_active_user),
):
    return await _agregaciones(db, f.filtros())


# ── Export a Excel con gráficas (dashboard + datos por tipo) ────────────────────

_AZUL = "1F4E79"
_AZUL2 = "2E75B6"


def _bucket(por_estado, patron):
    return sum(e["total"] for e in por_estado if re.search(patron, e["label"]))


@router.get("/export")
async def exportar_bolsas(
    f: BolsaFiltros = Depends(),
    db: AsyncSession = Depends(get_db),
    _: Usuario = Depends(get_current_active_user),
):
    base = f.filtros()
    agg = await _agregaciones(db, base)
    rows = (await db.execute(
        (select(Bolsa).where(and_(*base)) if base else select(Bolsa)).order_by(Bolsa.tipo, Bolsa.id)
    )).scalars().all()

    wb = openpyxl.Workbook()

    # ── Hoja oculta con las tablas que alimentan las gráficas ──────────────────
    wt = wb.create_sheet("Datos_graficas")

    def escribir_tabla(col, etiqueta, data, limite=12):
        wt.cell(row=1, column=col, value=etiqueta).font = Font(bold=True)
        wt.cell(row=1, column=col + 1, value="Casos").font = Font(bold=True)
        d = data[:limite]
        for i, x in enumerate(d, start=2):
            wt.cell(row=i, column=col, value=x.get("label") or x.get("periodo"))
            wt.cell(row=i, column=col + 1, value=x["total"])
        return len(d)

    n_tipo = escribir_tabla(1, "Tipo", agg["por_tipo"])
    n_est = escribir_tabla(4, "Estado", agg["por_estado"])
    n_zona = escribir_tabla(7, "Zona", agg["por_zona"])
    n_sup = escribir_tabla(10, "Supervisor", agg["por_supervisor"])
    n_tl = escribir_tabla(13, "Mes", agg["timeline"], limite=60)

    # ── Hoja RESUMEN (dashboard) ───────────────────────────────────────────────
    ws = wb.active
    ws.title = "RESUMEN"
    ws.sheet_view.showGridLines = False
    for c in "ABCDEFGHIJKLMNOP":
        ws.column_dimensions[c].width = 12

    ws.merge_cells("A1:H2")
    ws["A1"] = "REPORTE DE BOLSAS"
    ws["A1"].font = Font(bold=True, size=20, color="FFFFFF")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for row in ws["A1:H2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=_AZUL)
    ws["A3"] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}  ·  {agg['total']} casos en total"
    ws["A3"].font = Font(italic=True, color="666666")

    # KPI cards (filas 5-6)
    kpis = [
        ("TOTAL", agg["total"], _AZUL2),
        ("EJECUTADAS", _bucket(agg["por_estado"], r"EJECUT|SOLVED|CERR|FACTUR"), "548235"),
        ("PENDIENTES", _bucket(agg["por_estado"], r"PTE|PENDIENT|OPEN|VISITA|PROGRAM|GESTION"), "BF8F00"),
        ("ANULADAS", _bucket(agg["por_estado"], r"ANULAD|RECHAZ|CANCEL"), "C00000"),
    ]
    for i, (label, val, color) in enumerate(kpis):
        c1 = i * 2 + 1
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c1 + 1)
        ws.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c1 + 1)
        v = ws.cell(row=5, column=c1, value=val)
        v.font = Font(bold=True, size=22, color="FFFFFF")
        v.alignment = Alignment(horizontal="center", vertical="center")
        lab = ws.cell(row=6, column=c1, value=label)
        lab.font = Font(size=10, color="FFFFFF")
        lab.alignment = Alignment(horizontal="center", vertical="center")
        for rr in (5, 6):
            for cc in (c1, c1 + 1):
                ws.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor=color)

    # Gráfica de torta: por tipo
    pie = PieChart()
    pie.title = "Casos por tipo"
    pie.add_data(Reference(wt, min_col=2, min_row=1, max_row=1 + n_tipo), titles_from_data=True)
    pie.set_categories(Reference(wt, min_col=1, min_row=2, max_row=1 + n_tipo))
    pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
    pie.height = 8; pie.width = 13
    ws.add_chart(pie, "A8")

    def barra(titulo, col_lbl, col_val, ncats, anchor, horizontal=True, color="2E75B6"):
        ch = BarChart()
        ch.type = "bar" if horizontal else "col"
        ch.title = titulo
        ch.legend = None
        ch.add_data(Reference(wt, min_col=col_val, min_row=1, max_row=1 + ncats), titles_from_data=True)
        ch.set_categories(Reference(wt, min_col=col_lbl, min_row=2, max_row=1 + ncats))
        ch.height = 8; ch.width = 13
        ws.add_chart(ch, anchor)

    barra("Casos por estado", 4, 5, n_est, "I8", horizontal=True)
    barra("Casos por zona", 7, 8, n_zona, "A26", horizontal=False)
    barra("Casos por supervisor", 10, 11, n_sup, "I26", horizontal=True)

    line = LineChart()
    line.title = "Casos programados por mes"
    line.legend = None
    line.add_data(Reference(wt, min_col=14, min_row=1, max_row=1 + n_tl), titles_from_data=True)
    line.set_categories(Reference(wt, min_col=13, min_row=2, max_row=1 + n_tl))
    line.height = 8; line.width = 27
    ws.add_chart(line, "A44")

    wt.sheet_state = "hidden"

    # ── Hojas por tipo con sus columnas originales ─────────────────────────────
    by_tipo: dict = {}
    for b in rows:
        by_tipo.setdefault(b.tipo, []).append(b)

    for tipo, bs in by_tipo.items():
        cols: list = []
        parsed = []
        for b in bs:
            d = json.loads(b.datos_extra) if b.datos_extra else {}
            parsed.append(d)
            for k in d:
                if k not in cols:
                    cols.append(k)
        if not cols:
            cols = ["caso", "zona", "estado"]
        wst = wb.create_sheet(tipo[:31])
        wst.sheet_view.showGridLines = False
        for j, c in enumerate(cols, 1):
            cell = wst.cell(row=1, column=j, value=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=_AZUL2)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        wst.freeze_panes = "A2"
        wst.row_dimensions[1].height = 28
        for i, d in enumerate(parsed, 2):
            for j, c in enumerate(cols, 1):
                wst.cell(row=i, column=j, value=d.get(c))
        for j, c in enumerate(cols, 1):
            wst.column_dimensions[get_column_letter(j)].width = min(max(len(str(c)) + 2, 12), 42)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"reporte_bolsas_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ── CRUD ────────────────────────────────────────────────────────────────────────

@router.post("")
async def crear(req: BolsaCreate, db: AsyncSession = Depends(get_db), current_user: Usuario = Depends(get_current_active_user)):
    assert_puede_escribir(current_user)
    data = req.model_dump()
    extra = data.pop("datos_extra", None) or {}
    # Núcleo derivado de los campos del tipo (lo explícito tiene prioridad)
    if extra:
        for k, v in mapear_core(data.get("tipo"), extra).items():
            if not data.get(k):
                data[k] = v
    if not data.get("unidad"):
        data["unidad"] = current_user.unidad or "SABOGAL"
    b = Bolsa(**data, datos_extra=json.dumps(extra, ensure_ascii=False) if extra else None)
    db.add(b)
    await db.commit()
    await db.refresh(b)
    return _to_dict(b)


@router.patch("/{bolsa_id}")
async def actualizar(bolsa_id: int, req: BolsaUpdate, db: AsyncSession = Depends(get_db), current_user: Usuario = Depends(get_current_active_user)):
    assert_puede_escribir(current_user)
    b = (await db.execute(select(Bolsa).where(Bolsa.id == bolsa_id))).scalar_one_or_none()
    if not b:
        raise HTTPException(status_code=404, detail="No encontrado")
    data = req.model_dump(exclude_unset=True)
    extra = data.pop("datos_extra", None)
    if extra is not None:
        b.datos_extra = json.dumps(extra, ensure_ascii=False)
        # Recalcular el núcleo a partir de los campos del tipo
        for k, v in mapear_core(data.get("tipo") or b.tipo, extra).items():
            setattr(b, k, v)
    for k, v in data.items():   # campos del núcleo explícitos mandan
        setattr(b, k, v)
    await db.commit()
    await db.refresh(b)
    return _to_dict(b)


@router.delete("/{bolsa_id}")
async def eliminar(bolsa_id: int, db: AsyncSession = Depends(get_db), current_user: Usuario = Depends(get_current_active_user)):
    assert_puede_escribir(current_user)
    b = (await db.execute(select(Bolsa).where(Bolsa.id == bolsa_id))).scalar_one_or_none()
    if not b:
        raise HTTPException(status_code=404, detail="No encontrado")
    await db.delete(b)
    await db.commit()
    return {"ok": True}


# ── Importación masiva del Excel multi-hoja ─────────────────────────────────────

@router.post("/import")
async def importar_bolsas(
    file: UploadFile = File(...),
    reemplazar: bool = Form(False),
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    assert_puede_escribir(current_user)
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="El archivo debe ser .xlsx")
    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo: {e}")

    registros = parse_bolsas(wb)
    uni = current_user.unidad or "SABOGAL"

    if reemplazar:
        await db.execute(delete(Bolsa))
        await db.flush()

    por_tipo = Counter()
    for rec in registros:
        extra = rec.pop("datos_extra", {})
        b = Bolsa(unidad=uni, datos_extra=json.dumps(extra, ensure_ascii=False, default=str), **rec)
        db.add(b)
        por_tipo[rec["tipo"]] += 1
    # Snapshot de procedencia (indicador 3.1): el Excel que justifica el contenido
    # de las bolsas queda guardado con su hora de subida, descargable desde ANS → Procedencia.
    from app.api.ans import _guardar_fuente
    periodo_carga = datetime.now().strftime("%Y-%m")
    await _guardar_fuente(db, periodo_carga, "bolsas", file.filename, contents, sum(por_tipo.values()))
    await db.commit()

    return {"inserted": sum(por_tipo.values()), "por_tipo": dict(por_tipo)}
