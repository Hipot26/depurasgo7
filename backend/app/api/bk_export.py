from datetime import date
from typing import Optional
from io import BytesIO
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_, func
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.db.database import get_db
from app.api.deps import get_current_active_user, unidad_de, ROLES_GLOBALES
from app.models.models import Programacion, Usuario, ZonaEnum, UNIDAD_DEFAULT
from app.services.festivos import es_no_laborable

router = APIRouter(prefix="/export", tags=["export"])


def _unidad_efectiva(user: Usuario, unidad: Optional[str]) -> str:
    if user.rol in ROLES_GLOBALES:
        return unidad or UNIDAD_DEFAULT
    return unidad_de(user)


# Columnas que NO aplican a cada unidad (la unidad CHOCONTA retiró estas del Excel MOES;
# la unidad SABOGAL pasó a HORA AB/CR y a INOPERATIVIDAD de texto con la plantilla 2026-08)
UNIDAD_OCULTA = {
    "SABOGAL":  {"hora", "inoperatividad"},
    "CHOCONTA": {"pms", "fecha_vigencia", "cd_afectados", "ot", "estado_cumplimentacion",
                 "saidi", "hora_ab", "hora_cr", "restricciones", "se_gestiona",
                 "inoperatividad_tipo", "observaciones_inoperatividad"},
}
# Etiqueta de la columna de tipo según la unidad
UNIDAD_TIPO_LABEL = {"CHOCONTA": "TIPO RECURSO"}

# ── Colores exactos de los Excel originales de la programación ─────────────────
HEADER_BLUE      = "0070C0"   # azul dominante (cols 1-14, 19-25, 28-29, 37)
HEADER_LIME      = "92D050"   # verde lima (# PDT, ESTADO, CD AFECTADOS, HORA, OT, ESTADO CUMPLIM.)
HEADER_LBLUE     = "00B0F0"   # azul claro (MO REAL, #PROLIQUIDADO, FECHA LIQ. SGS, etc.)

# Colores de fondo en celdas de datos (por rol, ayuda visual para el usuario)
CELL_INGENIERO   = "DDEEFF"
CELL_PROGRAMADOR = "E2EFDA"

# Formato contable colombiano ($)
CURRENCY_FMT = r'_-"$"\ * #,##0_-;\-"$"\ * #,##0_-;_-"$"\ * "-"??_-;_-@_-'

# Columnas con cabecera lima en el original
_LIME_HEADERS = {"# PDT", "ESTADO", "CD AFECTADOS", "SAIDI", "HORA", "HORA AB", "HORA CR",
                 "OT", "ESTADO CUMPLIMENTACIÓN"}
# Cabecera gris azulada (plantilla 2026-08): V. dejado, notas, inoperatividad
_SLATE_HEADERS = {"VALOR DEJADO DE EJECUTAR", "NOTAS", "INOPERATIVIDAD",
                  "OBSERVACIONES INOPERATIVIDAD"}
HEADER_SLATE = "737F8F"
# Cabecera ROJA (plantilla 2026-08): SE GESTIONA SI O NO — la editan todos los roles
_RED_HEADERS = {"SE GESTIONA SI O NO"}
HEADER_RED = "FF0000"
# Rojo de día no laborable (domingo/festivo) — mismo rojo de la plantilla del cliente
CELL_NO_LABORABLE = "FF0000"
# Columnas con cabecera azul claro en el original
_LBLUE_HEADERS = {"MO REAL", "MO FACTURADA", "#PROLIQUIDADO", "FECHA LIQUIDACIÓN SGS",
                  "ESTADO FACTURACIÓN", "FECHA ENVIO A FACTURACIÓN"}
# Columnas monetarias que reciben formato de número
_MONEY_HEADERS = {"VALORDIA", "VALOR DIA", "MO PROYECTADA", "MO REAL", "MO FACTURADA",
                  "VALOR DEJADO DE EJECUTAR", "VALOR MATERIAL"}

# ── Definición de columnas ────────────────────────────────────────────────────
# (header, campo_model, color_celda_dato)
# color_celda_dato: None = sin color | CELL_INGENIERO | CELL_PROGRAMADOR
_ING  = CELL_INGENIERO
_PROG = CELL_PROGRAMADOR

COLUMNAS_EXCEL = [
    ("ZONA",                     "zona",                  None),
    ("MES",                      "fecha",                 None),   # sub="mes"
    ("DÍA",                      "fecha",                 None),   # sub="dia"
    ("FECHA",                    "fecha",                 None),
    ("LÍDER",                    "lider",                 None),
    ("TIPO GRUPO",               "tipo_grupo",            None),
    ("PROCESO",                  "proceso",               _ING),
    ("SUBPROCESO",               "subproceso",            _ING),
    ("PROYECTO O ANTECEDENTE",   "proyecto_antecedente",  _ING),
    ("PMS",                      "pms",                   _ING),
    ("CIRCUITO",                 "circuito",              _ING),
    ("LCL",                      "lcl",                   _ING),
    ("FECHA VIGENCIA",           "fecha_vigencia",        _ING),
    ("TIPO PDT",                 "tipo_pdt",              _ING),
    ("# PDT",                    "num_pdt",               _ING),
    ("ESTADO",                   "estado",                _ING),
    ("CD AFECTADOS",             "cd_afectados",          _ING),
    ("SAIDI",                    "saidi",                 _ING),
    ("HORA AB",                  "hora_ab",               _ING),
    ("HORA CR",                  "hora_cr",               _ING),
    ("HORA",                     "hora",                  _ING),
    ("MUNICIPIO",                "municipio",             _ING),
    ("VDA",                      "vereda",                _ING),
    ("ALCANCE",                  "alcance",               _ING),
    ("GESTOR ENEL",              "gestor_cliente",        _PROG),
    ("SUPERVISOR",               "supervisor",            _PROG),
    ("VALOR DIA",                "valor_dia",             _PROG),
    ("MO PROYECTADA",            "mo_proyectada",         _PROG),
    ("MO REAL",                  "mo_real",               _PROG),
    ("MO FACTURADA",             "mo_facturada",          _PROG),
    ("VALOR DEJADO DE EJECUTAR", "valor_dejado_ejecutar", _PROG),
    ("VALOR MATERIAL",           "valor_material",        _PROG),
    ("ESTADO PDT",               "estado_pdt",            None),
    ("NOTAS",                    "notas",                 None),
    ("RESTRICCIONES",            "restricciones",         _ING),
    ("SE GESTIONA SI O NO",      "se_gestiona",           _ING),
    ("#PROLIQUIDADO",            "num_proliquidado",      None),
    ("FECHA LIQUIDACIÓN SGS",    "fecha_liquidacion_sgs", None),
    ("ESTADO FACTURACIÓN",       "estado_facturacion",    None),
    ("FECHA ENVIO A FACTURACIÓN","fecha_envio_facturacion",None),
    ("OT",                       "ot",                    None),
    ("ESTADO CUMPLIMENTACIÓN",   "estado_cumplimentacion",None),
    ("INOPERATIVIDAD",           "inoperatividad",        None),   # $ (CHOCONTA)
    ("INOPERATIVIDAD",           "inoperatividad_tipo",   None),   # recurso (SABOGAL)
    ("OBSERVACIONES INOPERATIVIDAD", "observaciones_inoperatividad", None),
]

# Anchos de columna (en unidades Excel) — mismo orden que COLUMNAS_EXCEL
COL_WIDTHS = [8,7,10,12,21,10,16,16,36,26,15,14,14,10,20,14,18,10,9,9,8,15,15,42,20,20,12,14,14,14,22,14,14,30,24,12,15,14,18,20,12,22,12,14,28]


def cols_y_anchos(unidad: Optional[str]):
    """Devuelve (columnas, anchos) según la unidad: oculta las columnas que no
    aplican y ajusta la etiqueta del tipo (TIPO GRUPO → TIPO RECURSO en CHOCONTA)."""
    ocultas    = UNIDAD_OCULTA.get(unidad, set())
    tipo_label = UNIDAD_TIPO_LABEL.get(unidad)
    cols, anchos = [], []
    for (header, field, color), w in zip(COLUMNAS_EXCEL, COL_WIDTHS):
        if field in ocultas:
            continue
        if field == "tipo_grupo" and tipo_label:
            header = tipo_label
        cols.append((header, field, color))
        anchos.append(w)
    return cols, anchos

MESES_ES = {1:"enero",2:"febrero",3:"marzo",4:"abril",5:"mayo",6:"junio",
            7:"julio",8:"agosto",9:"septiembre",10:"octubre",11:"noviembre",12:"diciembre"}
DIAS_ES  = {0:"lunes",1:"martes",2:"miércoles",3:"jueves",4:"viernes",5:"sábado",6:"domingo"}


def _get_val(prog: Programacion, campo: str, sub: Optional[str] = None):
    if sub == "mes":
        return MESES_ES.get(prog.fecha.month, "") if prog.fecha else ""
    if sub == "dia":
        return DIAS_ES.get(prog.fecha.weekday(), "") if prog.fecha else ""
    val = getattr(prog, campo, None)
    return val.value if hasattr(val, "value") else val


def _header_fill(header: str) -> PatternFill:
    if header in _LIME_HEADERS:
        return PatternFill("solid", fgColor=HEADER_LIME)
    if header in _LBLUE_HEADERS:
        return PatternFill("solid", fgColor=HEADER_LBLUE)
    if header in _SLATE_HEADERS:
        return PatternFill("solid", fgColor=HEADER_SLATE)
    if header in _RED_HEADERS:
        return PatternFill("solid", fgColor=HEADER_RED)
    return PatternFill("solid", fgColor=HEADER_BLUE)


def _write_to_sheet(ws, registros: list, columnas=None, anchos=None):
    """Escribe headers y datos en un worksheet existente con todos los estilos."""
    if columnas is None:
        columnas, anchos = COLUMNAS_EXCEL, COL_WIDTHS
    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    h_font = Font(bold=True, color="FFFFFF", size=10)
    d_font = Font(size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(vertical="center", wrap_text=False)

    fill_cache: dict = {}

    def get_fill(hex_color: str) -> PatternFill:
        if hex_color not in fill_cache:
            fill_cache[hex_color] = PatternFill("solid", fgColor=hex_color)
        return fill_cache[hex_color]

    # ── Header row ────────────────────────────────────────────────────────────
    for col_idx, (header, _, _cell_color) in enumerate(columnas, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = h_font
        cell.fill      = _header_fill(header)
        cell.alignment = center
        cell.border    = border

    ws.row_dimensions[1].height = 40
    ws.freeze_panes = "A2"

    # ── Data rows ─────────────────────────────────────────────────────────────
    rojo_font = Font(size=10, bold=True, color="FFFFFF")
    for row_idx, prog in enumerate(registros, 2):
        # Domingo o festivo → DÍA y FECHA en rojo (igual que la plantilla del cliente)
        no_laborable = bool(prog.fecha) and es_no_laborable(prog.fecha)
        for col_idx, (header, campo, cell_color) in enumerate(columnas, 1):
            # Determine sub for date columns
            sub = None
            if campo == "fecha":
                if header == "MES":
                    sub = "mes"
                elif header == "DÍA":
                    sub = "dia"

            val  = _get_val(prog, campo, sub)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font      = d_font
            cell.border    = border
            cell.alignment = left

            if cell_color:
                cell.fill = get_fill(cell_color)

            if no_laborable and header in ("DÍA", "FECHA"):
                cell.fill = get_fill(CELL_NO_LABORABLE)
                cell.font = rojo_font

            if header in _MONEY_HEADERS and val is not None:
                cell.number_format = CURRENCY_FMT

    # ── Column widths ─────────────────────────────────────────────────────────
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _zona_sheet_name(zona_value: str) -> str:
    return f"P. {zona_value}"[:31]


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.get("/excel")
async def exportar_excel(
    zona: Optional[ZonaEnum] = None,
    mes: Optional[int] = None,
    anio: Optional[int] = None,
    unidad: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    uni = _unidad_efectiva(current_user, unidad)
    cols, anchos = cols_y_anchos(uni)
    filters = [Programacion.unidad == uni]
    if zona:
        filters.append(Programacion.zona == zona)
    if mes:
        filters.append(func.extract("month", Programacion.fecha) == mes)
    if anio:
        filters.append(func.extract("year", Programacion.fecha) == anio)

    q = select(Programacion).where(and_(*filters)).order_by(Programacion.fecha, Programacion.zona, Programacion.lider)
    registros = (await db.execute(q)).scalars().all()

    zona_label = zona.value if zona else "TODAS"
    sheet_name = _zona_sheet_name(zona_label) if zona else "P. UNIFICADA"
    mes_str    = MESES_ES.get(mes, "todos") if mes else "todos"
    filename   = f"programacion_{uni.lower()}_{zona_label}_{mes_str}_{anio or 'todos'}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    _write_to_sheet(ws, registros, cols, anchos)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/excel/unificado")
async def exportar_excel_unificado(
    mes: Optional[int] = None,
    anio: Optional[int] = None,
    unidad: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    """Excel con una hoja por zona (P. ALMEIDAS, etc.) + hoja P. UNIFICADA, de la unidad."""
    uni = _unidad_efectiva(current_user, unidad)
    cols, anchos = cols_y_anchos(uni)
    filters = [Programacion.unidad == uni]
    if mes:
        filters.append(func.extract("month", Programacion.fecha) == mes)
    if anio:
        filters.append(func.extract("year", Programacion.fecha) == anio)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Una hoja por zona
    for zona in ZonaEnum:
        zona_filters = filters + [Programacion.zona == zona]
        q = select(Programacion).where(and_(*zona_filters)).order_by(Programacion.fecha, Programacion.lider)
        registros = (await db.execute(q)).scalars().all()
        if not registros:
            continue
        ws = wb.create_sheet(title=_zona_sheet_name(zona.value))
        _write_to_sheet(ws, registros, cols, anchos)

    # Hoja unificada con todos los registros de la unidad
    q_all = select(Programacion).where(and_(*filters)).order_by(Programacion.zona, Programacion.fecha, Programacion.lider)
    todos = (await db.execute(q_all)).scalars().all()
    ws_uni = wb.create_sheet(title="P. UNIFICADA")
    _write_to_sheet(ws_uni, todos, cols, anchos)

    mes_str  = MESES_ES.get(mes, "todos") if mes else "todos"
    filename = f"programacion_unificada_{uni.lower()}_{mes_str}_{anio or ''}.xlsx"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ── Pre-planilla del día (el correo «Envío pre-planilla del día de mañana …») ─────
# Formato de la planilla que se manda a Enel (formato recibido el 2026-09-10): 15 columnas,
# una fila por cuadrilla, agrupada por zona y por trabajo (LCL). INGENIERO ZONA no existe en
# la Programación: sale del DIRECTORIO (profesional de apoyo de la zona, igual que el ANS).
PREPLANILLA_COLS = [
    # (encabezado, ancho)
    ("FECHA", 11), ("ZONA", 10), ("PROCESO", 16), ("INGENIERO ZONA", 18),
    ("SUPERVISOR TECNICO ENCARGADO", 22), ("GT / MOVIL", 9), ("JEFE CUADRILLA", 26),
    ("TIPO DE TRABAJO", 9), ("NUMERO DE TRABAJO", 24), ("LCL", 13), ("MUNICIPIO", 13),
    ("VEREDA", 18), ("CIRCUITO", 12), ("ALCANCE", 55), ("VALOR MO", 17),
]
PREPLANILLA_HEADER = "1F3864"   # azul marino de la planilla
PREPLANILLA_FILA   = "DDEBF7"   # fondo claro de todas las filas
# LCL coloreado por zona, alternando dos tonos cada vez que cambia el trabajo
PREPLANILLA_LCL = {
    "ALMEIDAS": ("FCE4D6", "F8CBAD"),
    "GUAVIO":   ("FFF2CC", "FFD966"),
    "UBATE":    ("E2EFDA", "C6E0B4"),
    "LV&TLC":   ("D9D9D9", "A6A6A6"),
}
_ORDEN_ZONA = {z.value: i for i, z in enumerate(ZonaEnum)}


@router.get("/preplanilla")
async def exportar_preplanilla(
    fecha: date,
    zona: Optional[ZonaEnum] = None,
    unidad: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    """Pre-planilla de UN día con lo cargado en la Programación, en el formato que se envía
    el día anterior. Omite las filas sin cuadrilla (líder «TODOS»: pausas activas, charlas)."""
    from app.api.ans import _apoyo_map, _zona_cuenta   # import tardío: ans.py es pesado
    uni = _unidad_efectiva(current_user, unidad)
    filters = [Programacion.unidad == uni, Programacion.fecha == fecha]
    if zona:
        filters.append(Programacion.zona == zona)
    regs = [p for p in (await db.execute(select(Programacion).where(and_(*filters)))).scalars().all()
            if (p.lider or "").strip().upper() not in ("", "TODOS")]
    # Zona → trabajo (LCL; los sin LCL al final) → orden de carga
    regs.sort(key=lambda p: (_ORDEN_ZONA.get(p.zona.value, 99), not p.lcl, p.lcl or "",
                             p.num_pdt or "", p.id))
    apoyos = {(_zona_cuenta(z) or ""): (n or "").upper() for z, n in (await _apoyo_map(db)).items()}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"PRE-PLANILLA {fecha.day}-{fecha.month}-{fecha.year}"
    thin   = Side(style="thin", color="8EA9DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    h_fill = PatternFill("solid", fgColor=PREPLANILLA_HEADER)
    f_fila = PatternFill("solid", fgColor=PREPLANILLA_FILA)
    h_font = Font(bold=True, color="FFFFFF", size=10)
    d_font = Font(size=9)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(vertical="center", wrap_text=True)

    for c, (h, w) in enumerate(PREPLANILLA_COLS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = h_font, h_fill, center, border
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    grupo, previo = 0, None
    for r, p in enumerate(regs, 2):
        z = p.zona.value
        clave = (z, p.lcl or "", p.num_pdt or "")
        if clave != previo:
            grupo, previo = grupo + 1, clave
        valores = [
            p.fecha, z, p.proceso or p.subproceso, apoyos.get(_zona_cuenta(z) or "", ""),
            p.supervisor, p.tipo_grupo, p.lider, p.tipo_pdt, p.num_pdt, p.lcl,
            p.municipio, p.vereda, p.circuito, p.alcance, p.mo_proyectada or 0,
        ]
        for c, v in enumerate(valores, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.fill, cell.border = d_font, f_fila, border
            cell.alignment = left
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=15).number_format = CURRENCY_FMT
        tonos = PREPLANILLA_LCL.get(z)
        if tonos and p.lcl:
            ws.cell(row=r, column=10).fill = PatternFill("solid", fgColor=tonos[grupo % 2])

    if regs:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(PREPLANILLA_COLS))}{len(regs) + 1}"

    filename = f"Pre-planilla {fecha.day}-{fecha.month}-{fecha.year}.xlsx"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ── Planilla compacta del día (la que Enel devuelve para validar los programados) ──
# Formato pedido por el cliente (planilla del 8-9-2026 y aclaraciones del 2026-09-16): la misma información de la pre-planilla pero compactada —
# UNA FILA POR TRABAJO, no por cuadrilla. Las cuadrillas del mismo trabajo se agrupan en
# JEFE CUADRILLA separadas por «/» y con su tipo entre paréntesis:
# «LIDER UNO (CAMP5)/LIDER DOS (ARRIME)/LIDER TRES (GRUA)».
# Las dos columnas de nodos del PMS no existen en la Programación y salen VACÍAS a propósito
# (textual del cliente: «esos datos no los tendríamos, sería dejar eso como vacío»).
PLANILLA_COLS = [
    # (encabezado, ancho)
    ("FECHA DE EJECUCION", 12), ("LCL", 13), ("PROCESO", 15), ("SUBPROCESO", 18),
    ("QR EN PMS", 28), ("NODOS PMS PROGRAMADOS", 16), ("CANT NODOS PROG", 11),
    ("TIPO DE PDT", 10), ("PDT", 20), ("ALCANCE TRABAJO", 70), ("VALOR MO", 15),
    ("VALOR MATERIALES", 15), ("MUNICIPIO", 14), ("VEREDA", 15), ("NOMBRE CIRCUITO", 14),
    ("JEFE CUADRILLA", 30), ("PROVINCIA", 11),
    # Requisito del coordinador (2026-09-23): «dos casillas a la derecha,
    # en la primera el supervisor encargado del trabajo y en la segunda el responsable de Enel».
    ("SUPERVISOR", 22), ("RESPONSABLE ENEL", 20),
]
PLANILLA_HEADER = "1F3864"   # azul marino del encabezado original
_VACIOS = ("", "N/A", "NA", "-")


def _clave_trabajo(p: Programacion):
    """Qué filas de la Programación son *el mismo trabajo* y por tanto se funden en una sola
    línea de la planilla. Sin LCL y sin PDT no hay trabajo que agrupar (cuadrilla disponible,
    mantenimiento interno, emergencias sin número): esas filas quedan cada una por su lado."""
    lcl = (p.lcl or "").strip().upper()
    pdt = (p.num_pdt or "").strip().upper()
    if lcl in _VACIOS and pdt in _VACIOS:
        return ("", "", "", p.id)          # irrepetible: nunca se agrupa con nadie
    return (p.zona.value, lcl, pdt, "")


def _jefes_cuadrilla(grupo: list) -> str:
    """«NOMBRE (TIPO)/NOMBRE (TIPO)…» — el formato exacto con el que el cliente escribe la
    casilla cuando varias cuadrillas están en el mismo trabajo. Se respeta el orden de carga
    de la Programación y se descartan los líderes repetidos."""
    vistos, partes = set(), []
    for p in grupo:
        nombre = (p.lider or "").strip()
        if not nombre or nombre.upper() in vistos:
            continue
        vistos.add(nombre.upper())
        tipo = (p.tipo_grupo or "").strip()
        partes.append(f"{nombre} ({tipo})" if tipo else nombre)
    return "/".join(partes)


def _valor_trabajo(grupo: list, campo: str):
    """Del grupo, el valor del campo que MÁS SE REPITE entre las filas que lo traen (empate: el
    que aparece primero). Los datos técnicos del trabajo suelen estar en una sola fila de las
    cuadrillas —ahí no hay competencia—, pero cuando varias filas los traen manda la mayoría.
    Antes se tomaba el primer valor no vacío y bastaba UNA fila corrida en el Excel para que el
    trabajo entero saliera con municipio, vereda y alcance ajenos (planilla del 28-09-2026,
    UBATE: la fila de Leandro Rios traía «M. SUSA VDA LLANO GRANDE / EN GESTION / 0.0001» en
    esas columnas mientras sus dos compañeros de trabajo traían lo correcto)."""
    conteo: dict = {}
    for p in grupo:
        v = getattr(p, campo, None)
        if v in (None, ""):
            continue
        k = v.strip().upper() if isinstance(v, str) else v
        conteo.setdefault(k, [0, v])[0] += 1
    if not conteo:
        return None
    return max(conteo.values(), key=lambda c: c[0])[1]


@router.get("/planilla")
async def exportar_planilla(
    fecha: date,
    zona: Optional[ZonaEnum] = None,
    unidad: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    """Planilla compacta de UN día: una fila por trabajo, con las cuadrillas agrupadas.
    Omite las filas sin cuadrilla (líder «TODOS»: pausas activas, charlas)."""
    uni = _unidad_efectiva(current_user, unidad)
    filters = [Programacion.unidad == uni, Programacion.fecha == fecha]
    if zona:
        filters.append(Programacion.zona == zona)
    regs = [p for p in (await db.execute(select(Programacion).where(and_(*filters)))).scalars().all()
            if (p.lider or "").strip().upper() not in ("", "TODOS")]
    # Zona → trabajo (LCL; los sin LCL al final) → orden de carga: el mismo orden de la pre-planilla
    regs.sort(key=lambda p: (_ORDEN_ZONA.get(p.zona.value, 99), not p.lcl, p.lcl or "",
                             p.num_pdt or "", p.id))

    # Agrupar conservando el orden de aparición
    grupos: dict = {}
    for p in regs:
        grupos.setdefault(_clave_trabajo(p), []).append(p)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"PLANILLA COMPACTA {fecha.day}-{fecha.month}-{fecha.year}"
    thin   = Side(style="thin", color="8EA9DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    h_fill = PatternFill("solid", fgColor=PLANILLA_HEADER)
    h_font = Font(bold=True, color="FFFFFF", size=9)
    d_font = Font(size=9)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(vertical="center", wrap_text=True)

    for c, (h, w) in enumerate(PLANILLA_COLS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = h_font, h_fill, center, border
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"

    for r, grupo in enumerate(grupos.values(), 2):
        p = grupo[0]
        valores = [
            fecha,                                       # FECHA DE EJECUCION
            _valor_trabajo(grupo, "lcl"),                 # LCL
            _valor_trabajo(grupo, "proceso"),             # PROCESO
            _valor_trabajo(grupo, "subproceso"),          # SUBPROCESO
            _valor_trabajo(grupo, "pms"),                 # QR EN PMS
            None,                                        # NODOS PMS PROGRAMADOS — no está en la plataforma
            None,                                        # CANT NODOS PROG — no está en la plataforma
            _valor_trabajo(grupo, "tipo_pdt"),            # TIPO DE PDT
            _valor_trabajo(grupo, "num_pdt"),             # PDT
            _valor_trabajo(grupo, "alcance"),             # ALCANCE TRABAJO
            sum(x.mo_proyectada or 0 for x in grupo),    # VALOR MO — el valor va en una sola fila del trabajo
            sum(x.valor_material or 0 for x in grupo),   # VALOR MATERIALES
            _valor_trabajo(grupo, "municipio"),           # MUNICIPIO
            _valor_trabajo(grupo, "vereda"),              # VEREDA
            _valor_trabajo(grupo, "circuito"),            # NOMBRE CIRCUITO
            _jefes_cuadrilla(grupo),                     # JEFE CUADRILLA
            p.zona.value,                                # PROVINCIA
            _valor_trabajo(grupo, "supervisor"),          # SUPERVISOR (col. SUPERVISOR de la Programación)
            _valor_trabajo(grupo, "gestor_cliente"),      # RESPONSABLE ENEL (col. GESTOR ENEL)
        ]
        for c, v in enumerate(valores, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = d_font, border
            cell.alignment = left
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=11).number_format = CURRENCY_FMT
        ws.cell(row=r, column=12).number_format = CURRENCY_FMT

    if grupos:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(PLANILLA_COLS))}{len(grupos) + 1}"

    filename = f"Planilla Compacta {fecha.day}-{fecha.month}-{fecha.year}.xlsx"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/plantilla")
async def descargar_plantilla(
    unidad: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    current_user: Usuario = Depends(get_current_active_user),
):
    """Plantilla Excel vacía (solo encabezados) con el formato exacto de la unidad,
    para llenar masivamente en Excel y luego importar."""
    uni = _unidad_efectiva(current_user, unidad)
    cols, anchos = cols_y_anchos(uni)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PROGRAMACIÓN"
    _write_to_sheet(ws, [], cols, anchos)   # sin registros → solo encabezados

    filename = f"plantilla_programacion_{uni.lower()}.xlsx"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
