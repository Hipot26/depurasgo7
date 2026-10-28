"""API del historial de auditoría global (solo consulta + respaldo, v1).

El llenado es automático (app/services/auditoria.py, listener before_flush); aquí solo se
lee/filtra/exporta. Gate: desarrollador + directivo (require_auditoria)."""
import io
from datetime import datetime, date, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, and_
import openpyxl
from openpyxl.styles import Font, PatternFill

from app.db.database import get_db
from app.api.deps import require_auditoria
from app.models.models import Auditoria

router = APIRouter(prefix="/auditoria", tags=["auditoria"], dependencies=[Depends(require_auditoria)])


def _filtros(usuario_id, modulo, tabla, accion, desde, hasta, q):
    f = []
    if usuario_id:
        f.append(Auditoria.usuario_id == usuario_id)
    if modulo:
        f.append(Auditoria.modulo == modulo)
    if tabla:
        f.append(Auditoria.tabla == tabla)
    if accion:
        f.append(Auditoria.accion == accion)
    if desde:
        f.append(Auditoria.creado_en >= datetime.fromisoformat(desde))
    if hasta:
        # 'hasta' inclusivo: hasta el final de ese día
        f.append(Auditoria.creado_en < datetime.fromisoformat(hasta) + timedelta(days=1))
    if q:
        like = f"%{q}%"
        f.append(func.coalesce(Auditoria.campo, "").ilike(like) |
                 func.coalesce(Auditoria.valor_anterior, "").ilike(like) |
                 func.coalesce(Auditoria.valor_nuevo, "").ilike(like) |
                 func.coalesce(Auditoria.resumen, "").ilike(like))
    return f


def _row(a: Auditoria) -> dict:
    return {
        "id": a.id, "creado_en": a.creado_en.isoformat() if a.creado_en else None,
        "usuario_id": a.usuario_id, "usuario_nombre": a.usuario_nombre, "usuario_rol": a.usuario_rol,
        "unidad": a.unidad, "modulo": a.modulo, "tabla": a.tabla, "registro_id": a.registro_id,
        "accion": a.accion, "campo": a.campo, "valor_anterior": a.valor_anterior,
        "valor_nuevo": a.valor_nuevo, "request_id": a.request_id, "ruta": a.ruta,
        "metodo": a.metodo, "ip": a.ip, "resumen": a.resumen,
    }


@router.get("")
async def listar(
    usuario_id: Optional[int] = None, modulo: Optional[str] = None, tabla: Optional[str] = None,
    accion: Optional[str] = None, desde: Optional[str] = None, hasta: Optional[str] = None,
    q: Optional[str] = None, limite: int = Query(200, le=2000), offset: int = 0,
    db: AsyncSession = Depends(get_db),
):
    """Lista plana de cambios (más recientes primero), con filtros."""
    f = _filtros(usuario_id, modulo, tabla, accion, desde, hasta, q)
    base = select(Auditoria).where(and_(*f)) if f else select(Auditoria)
    total = (await db.execute(
        (select(func.count()).select_from(Auditoria).where(and_(*f))) if f
        else select(func.count()).select_from(Auditoria)
    )).scalar_one()
    rows = (await db.execute(
        base.order_by(Auditoria.creado_en.desc(), Auditoria.id.desc()).limit(limite).offset(offset)
    )).scalars().all()
    return {"total": total, "items": [_row(a) for a in rows]}


@router.get("/operaciones")
async def operaciones(
    usuario_id: Optional[int] = None, modulo: Optional[str] = None, tabla: Optional[str] = None,
    accion: Optional[str] = None, desde: Optional[str] = None, hasta: Optional[str] = None,
    q: Optional[str] = None, limite: int = Query(200, le=1000),
    db: AsyncSession = Depends(get_db),
):
    """Vista agrupada por operación (request_id): una línea por acción de usuario, con el nº de
    filas afectadas. Ideal para ver el comportamiento y detectar cargas masivas."""
    f = _filtros(usuario_id, modulo, tabla, accion, desde, hasta, q)
    g = select(
        Auditoria.request_id,
        func.min(Auditoria.creado_en).label("inicio"),
        func.max(Auditoria.creado_en).label("fin"),
        func.max(Auditoria.usuario_nombre).label("usuario"),
        func.max(Auditoria.usuario_id).label("usuario_id"),
        func.max(Auditoria.usuario_rol).label("rol"),
        func.max(Auditoria.modulo).label("modulo"),
        func.max(Auditoria.accion).label("accion"),
        func.max(Auditoria.ruta).label("ruta"),
        func.count().label("filas"),
    )
    if f:
        g = g.where(and_(*f))
    g = g.group_by(Auditoria.request_id).order_by(func.min(Auditoria.creado_en).desc()).limit(limite)
    rows = (await db.execute(g)).all()
    return [{
        "request_id": r.request_id, "inicio": r.inicio.isoformat() if r.inicio else None,
        "fin": r.fin.isoformat() if r.fin else None, "usuario": r.usuario, "usuario_id": r.usuario_id,
        "rol": r.rol, "modulo": r.modulo, "accion": r.accion, "ruta": r.ruta, "filas": r.filas,
    } for r in rows]


@router.get("/usuarios")
async def usuarios_con_actividad(db: AsyncSession = Depends(get_db)):
    """Usuarios distintos que tienen actividad registrada (para el filtro de la UI)."""
    rows = (await db.execute(
        select(Auditoria.usuario_id, func.max(Auditoria.usuario_nombre), func.count())
        .where(Auditoria.usuario_id.isnot(None))
        .group_by(Auditoria.usuario_id).order_by(func.max(Auditoria.usuario_nombre))
    )).all()
    return [{"usuario_id": uid, "nombre": nom, "acciones": n} for uid, nom, n in rows]


@router.get("/modulos")
async def modulos(db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(
        select(Auditoria.modulo, func.count()).group_by(Auditoria.modulo).order_by(Auditoria.modulo)
    )).all()
    return [{"modulo": m, "acciones": n} for m, n in rows if m]


@router.get("/operacion/{request_id}")
async def detalle_operacion(request_id: str, db: AsyncSession = Depends(get_db)):
    """Todas las filas de una operación (request_id) — el detalle de una carga masiva."""
    rows = (await db.execute(
        select(Auditoria).where(Auditoria.request_id == request_id)
        .order_by(Auditoria.id.asc())
    )).scalars().all()
    return [_row(a) for a in rows]


@router.get("/export")
async def exportar(
    usuario_id: Optional[int] = None, modulo: Optional[str] = None, tabla: Optional[str] = None,
    accion: Optional[str] = None, desde: Optional[str] = None, hasta: Optional[str] = None,
    q: Optional[str] = None, limite: int = Query(20000, le=100000),
    db: AsyncSession = Depends(get_db),
):
    """Descarga el historial filtrado como .xlsx (respaldo offline)."""
    f = _filtros(usuario_id, modulo, tabla, accion, desde, hasta, q)
    base = select(Auditoria).where(and_(*f)) if f else select(Auditoria)
    rows = (await db.execute(
        base.order_by(Auditoria.creado_en.desc(), Auditoria.id.desc()).limit(limite)
    )).scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Auditoria"
    cab = ["Fecha", "Usuario", "Rol", "Unidad", "Módulo", "Tabla", "Registro", "Acción",
           "Campo", "Valor anterior", "Valor nuevo", "Operación", "Ruta", "Método", "IP"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
    for a in rows:
        ws.append([
            a.creado_en.strftime("%Y-%m-%d %H:%M:%S") if a.creado_en else "",
            a.usuario_nombre, a.usuario_rol, a.unidad, a.modulo, a.tabla, a.registro_id,
            a.accion, a.campo, a.valor_anterior, a.valor_nuevo, a.request_id, a.ruta, a.metodo, a.ip,
        ])
    widths = [19, 18, 13, 10, 14, 18, 9, 8, 16, 30, 30, 36, 24, 8, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"Auditoria_{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
