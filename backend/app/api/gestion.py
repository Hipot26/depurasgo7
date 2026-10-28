"""gestion.py — Gestión General de SGO7 (22 oct 2026): un caso por proyecto (LCL o ticket DF) con el
flujo FL-PS-E-03 de Dominion como columna vertebral. 39 pasos en 3 pilares; el paso actual lo decide la máquina
a partir de cuatro fuentes:

  correo  → el agente del PC (agente_gestion.py, Outlook por COM) entrega los correos que nombran un código.
            De qué proyecto es cada correo lo decide el CENTRO MO: si ya interpretó ese correo
            (mismo entry_id de Outlook) vale su código validado; si no, sus reglas: candidatos en orden
            asunto > cuerpo reciente > cadena vieja, solo códigos que existen en SGS, candado de supervisor.
            Los hilos citados se desglosan y cada mensaje se ubica en un paso (gestion_flujo.clasificar).
  sgs     → lectura en vivo por el Centro MO (/interno/gestion/sgs): todos los padres del código (regla 16),
            supervisor y departamento de SGS, y la MO liquidada de cada familia agrupada por el día en que se
            liquidó (create_date de los avances). Ya no se suman meses del Control Mensual.
  sonny   → casos del Balance Integrado (sgo7_casos) sin las reconsultas automáticas que no cambian nada, y (CHANGELOG 2026-10-23) el
            Historial de Planillas de Sonny (sgo7_gg_sonny, export que carga cargar_sonny_gestion.py desde el PC): cada
            cambio de estado de cada planilla con su fecha oficial, en los pasos 24-39.
  sgo7    → correos de MO del Centro MO (estado, completado) y Registros de los módulos (Liquidador, Balance).
  Cambio del 23 oct 2026 — el hilo del dinero, exportes manuales que suben los scripts del PC con el token del agente:
  sonny   → además la Lista de Planillas (sgo7_gg_sonny_planilla): MO/MAT real y fechas del supervisor e ingeniero (pasos 19, 25, 26)
  pms     → solicitudes de PMS y su historial (sgo7_gg_pms, sgo7_gg_pms_log): pasos 1, 4, 5, 9, 14 y 15, y la caja «Presupuestado»
  scm     → libretos de SCM con su MO en pesos y su material (sgo7_gg_scm_libreto, sgo7_gg_scm_mat): paso 37 y la caja «Libreto»;
            el enlace planilla ↔ libreto lo hace el servidor (gestion_flujo.enlazar_libretos)
  deuda   → el informe de deuda de facturación (sgo7_gg_deuda): un chip en la franja, sin nombres ni cédulas
  La franja del dinero (p.dinero, gestion_flujo.calcular_dinero) y 5 alertas se calculan en cada recálculo. Un caso puede nacer
  también de Sonny, PMS o SCM (origen) si su código existe en SGS (árbitro: Centro MO), con POST /agente/nacer.
  Cambio del 23 oct 2026 — cada cifra con su evidencia: cada carga del PC queda en sgo7_gg_cargas (archivo, tamaño, sha256, fecha del
  exporte, quién y cuándo cargó, filas) y sus filas llevan carga_id; SCM manda además cada prestación (sgo7_gg_scm_prest:
  cantidad × precio QP = pesos). «¿De dónde sale?» (GET /proyectos/{id}/evidencia, gestion_flujo.armar_evidencia) devuelve las filas
  que suman cada caja y cada alerta de dinero, con fórmula, fuente y confianza del cruce; el recálculo guarda si cada suma
  reproduce su cifra (p.dinero.evidencia). Las confirmaciones humanas del cruce planilla ↔ libreto («Es este» / «No es») van a
  sgo7_gg_decisiones y mandan en todo recálculo (enlazar_libretos). Panel «Fuentes y controles»: GET /fuentes.
  Cambio del 23 oct 2026 — SAP: el MB51 del SAP de Enel (llega por correo; lo sube cargar_sap_gestion.py) entra a sgo7_gg_sap (upsert por
  documento + posición, sin usuarios ni textos libres). Es solo MATERIAL: por caso, lo descargado contra su LCL (o su DF, o su LCL con la
  orden de uno de sus libretos: Orden de SAP = Trabajo de SCM) y lo recuperado. Al cerrar la carga se rehace el precio de Enel por código
  (sgo7_gg_precio_mat: mediana de Importe/cantidad de las 261), que valora todo el material de la franja; la fila «Material por código»
  (SGS → SCM → SAP), la caja «Facturado SAP» (sigue punteada: no hay factura, HES ni pago) y 3 alertas de material más.
  Cambio del 25 oct 2026 — material por liquidar: el Centro MO trae las bandejas aprobadas de SGS por código (despachado, devuelto, última salida
  BAN-) y la caja SGS MAT pasa a tres cifras: liquidado (consumido), por liquidar (despachado − devuelto − consumido) y proyectado
  (gestion_flujo.partes_mat), a precio de Enel. El proyectado del LCL se compara con el material de PMS (cuadra / excede / por debajo,
  ±GESTION_MAT_EXCESO_PCT de config.py) y hay 2 alertas de material: mat_excede_pms y mat_por_liquidar (informativa, desde
  GESTION_MAT_DIAS_SIN_LIQUIDAR días hábiles). Control nuevo en «Fuentes y controles»: liquidado + por liquidar = despachado neto.
  Cambio del 25 oct 2026 — Sonny por etapa: tres exportes más de Sonny que sube cargar_sonny_gestion.py — «Lista Cruzada de Nodos» (Contrato
  JA10185664: cada nodo con su estado, la fecha de ese estado y su planilla), «Nodos por gestion almacen» (el material del nodo que espera
  planilla) y el «Historial de Actividades» (creación, cambios de estado y firmas de cada planilla, sin nombres) — entran a sgo7_gg_sonny_nodo y
  sgo7_gg_sonny_plan_act (POST /agente/sonny-nodos). La caja Sonny muestra el plan al lado del real (su cifra no cambia) y debajo va la fila
  «Sonny por etapa» (gestion_flujo.calcular_etapas): sin planilla (nodos) → en curso → en Enel → libretos, con el techo de PMS en gris, y 4
  alertas: nodo_sin_planilla, planilla_en_curso_quieta, planilla_en_firmas y pms_finalizado_sin_planilla (días en config.py: GESTION_SONNY_DIAS_*).
  Cambio del 25 oct 2026 — el cierre en SCM: el exporte de LCL de SCM («lcl_export», Informe LCL/LCA: una fila por LCL × Trabajo) entra por
  cargar_scm_gestion.py a sgo7_gg_scm_lcl (POST /agente/scm-lcl; upsert por LCL + Trabajo; sin «Nombre técnico EDE», «Ejecutor» ni descripción).
  Un caso está cerrado (sgo7_gg_proyectos.cerrado, cerrado_en) cuando todas sus LCL figuran CONCLUSA (gestion_flujo.calcular_scm_lcl); la caja
  SCM lleva el chip «LCL concluida dd/mm» o «LCL abierta · <estado>», la fila «SCM · cierre» (LCL y libretos por estado) y 3 alertas:
  lcl_concluida_libreto_abierto, lcl_sin_libreto_vence e importe_lcl_vs_mo (config.py: GESTION_SCM_DIAS_VENCE, GESTION_SCM_IMPORTE_PCT). El
  Trabajo (= orden de SAP) también da la confianza alta del cruce LCL ↔ MB51 en las LCL sin libreto. Filtro «abiertos / cerrados» en la lista.
  Cambio del 25 oct 2026 — el cobro en DRAPE (portal de facturación de Enel, del lado del proveedor): las listas de facturas (Borrador, Enviadas,
  Rechazadas, Registradas, Pagadas/compensadas) y el detalle de cada factura entran desde el PC (POST /agente/drape) a
  sgo7_gg_drape_factura (upsert por DRAPE ID: estado, LCL de «Num. Pedido», EM, registro, pago, total; con el detalle, vencimiento, documento de pago y
  condiciones), sgo7_gg_drape_pos (el «Resumen de Posición»: LCL/posición e importe) y sgo7_gg_drape_em (EM pendientes de facturar, si DRAPE trae la
  lista). Sin textos libres (notas y mensajes pueden traer nombres) ni datos bancarios. La caja SAP pasa a «Facturado (DRAPE)»
  (gestion_flujo.calcular_cobro): sin factura / facturado / pagado / vencida, importe por LCL donde hay posiciones, y 4 alertas de dinero:
  lcl_concluida_sin_factura, factura_vencida_sin_pago, factura_rechazada y em_sin_facturar (config.py: GESTION_DRAPE_*). La orden de SAP no
  aparece en DRAPE: la llave es la LCL.
  Cambio del 25 oct 2026 — Gestión General para directivos: en cada recálculo se arman y guardan en p.dinero["ux"] el estado del trabajo (uno, con color y
  quién tiene la pelota), «qué falta para cobrar» (frases llanas con pesos y fechas), la cadena de identificadores (LCL → solicitud PMS → PDL →
  nodos → planillas → libretos → facturas → pago, con el tipo de cada enlace) y las 6 cifras de MO del trabajo; cada alerta lleva `pesos` y `frase`
  (gestion_flujo.vestir_alertas). GET /zona da la portada (MO por etapa, lo que frena el cobro por pesos, estados, lo más grande en juego y la
  garantía de los datos) con caché de 60 s; GET /proyectos/{id}/cadena lee lo guardado. config.GESTION_PAGO_CONFIRMADO (False): «pagado» se dice
  «sin pago registrado en DRAPE · por confirmar» y no cuenta como cobro ni como deuda; config.GESTION_CASO_EJEMPLO: el trabajo de «¿Cómo fluye?».
  Cambio del 25 oct 2026 — El motivo del rechazo: el texto de «Observaciones» del Historial de Planillas de Sonny (sgo7_gg_sonny, que no se toca) se
  desescapa en el evento y se clasifica (gestion_flujo.motivo_rechazo, regex): ref.motivo en el evento, motivos en la alerta rechazo_abierto,
  dinero.rechazos (abiertos con su texto y los anteriores) y dinero.ux.rechazos; GET /zona desglosa el freno por categoría con pesos y
  GET /proyectos acepta ?rechazo=<categoría> (alertas con esa categoría en motivo_cats).
  Cambio del 25 oct 2026 — «Descargar» por fuente: cada cargador del PC sube el archivo de su carga (POST /agente/carga/{id}/archivo, multipart; el
  servidor comprueba que su sha256 es el de la carga, lo guarda en /var/www/sgo7/registros/gestion/cargas/<id>/<nombre> y no lo duplica);
  GET /cargas/{id}/archivo lo entrega (solo quien ve Gestión) y la descarga queda en Registros (gestion_descarga_fuente, con el sha256);
  GET /agente/cargas da al relleno del PC las cargas y si ya tienen archivo. Con config.GESTION_MOSTRAR_DRAPE en False (valor actual), los
  datos del portal de facturación del cliente no aparecen en ninguna respuesta (ni cobro, ni sus alertas, ni su fila de fuentes, ni sus
  controles, ni la pieza de factura o pago de la cadena); la facturación sale del informe de deuda (gestion_flujo.aplicar_deuda). Las tablas se
  conservan.

Rutas:
  GET  /api/gestion/agente/estado      (X-Agente-Token)  cursor y hilos conocidos para el agente del PC
  POST /api/gestion/agente/ingesta     (X-Agente-Token)  correos nuevos → eventos → estado del proyecto
  POST /api/gestion/agente/sonny       (X-Agente-Token)  filas del Historial de Planillas de Sonny (upsert por id_cambio)
  POST /api/gestion/agente/sonny-planillas (X-Agente-Token) Lista de Planillas de Sonny (upsert por id_planilla)        CHANGELOG 2026-10-23
  POST /api/gestion/agente/pms         (X-Agente-Token)  solicitudes y cambios de estado de PMS (upsert por request_code / huella) CHANGELOG 2026-10-23
  POST /api/gestion/agente/scm         (X-Agente-Token)  libretos de SCM con su material (upsert por lcl + libreto)  CHANGELOG 2026-10-23
  POST /api/gestion/agente/deuda       (X-Agente-Token)  informe de deuda de facturación (upsert por corte + ítem) CHANGELOG 2026-10-23
  POST /api/gestion/agente/nacer       (X-Agente-Token)  casos que nacen de Sonny, PMS o SCM (por lotes)            CHANGELOG 2026-10-23
  POST /api/gestion/agente/carga       (X-Agente-Token)  abre / cierra el registro de una carga (archivo, sha256, filas) CHANGELOG 2026-10-23
  POST /api/gestion/agente/sap         (X-Agente-Token)  filas del MB51 de SAP (upsert por documento + posición); fin=true rehace precios CHANGELOG 2026-10-23
  POST /api/gestion/agente/sonny-nodos (X-Agente-Token)  nodos de la Lista Cruzada (upsert por Id_nodo) y actividad por planilla (upsert por Id) CHANGELOG 2026-10-25
  POST /api/gestion/agente/scm-lcl     (X-Agente-Token)  filas del exporte de LCL de SCM (upsert por LCL + Trabajo); fin=true recalcula       CHANGELOG 2026-10-25
  POST /api/gestion/agente/drape       (X-Agente-Token)  facturas de DRAPE (upsert por DRAPE ID), su detalle y posiciones, EM; fin=true recalcula CHANGELOG 2026-10-25
  GET  /api/gestion/fuentes           (puede_gestion)   panel «Fuentes y controles»: último exporte y carga por sistema    CHANGELOG 2026-10-23
  GET  /api/gestion/zona              (puede_gestion)   portada de la zona para directivos (caché 60 s)                     CHANGELOG 2026-10-25
  GET  /api/gestion/proyectos/{id}/cadena (puede_gestion) estado, qué falta, cadena y plata del trabajo (lo guardado)     CHANGELOG 2026-10-25
  GET  /api/gestion/proyectos/{id}/evidencia        (puede_gestion)  tarjetas «¿De dónde sale?» del caso                   CHANGELOG 2026-10-23
  GET  /api/gestion/proyectos/{id}/evidencia/excel  (puede_gestion)  Excel con una hoja por caja y =SUM al pie (Registros) CHANGELOG 2026-10-23
  POST /api/gestion/proyectos/{id}/recalcular       (dev)            recalcula un caso (SGS en vivo + enlace + franja)     CHANGELOG 2026-10-23
  POST /api/gestion/decisiones          (dev o facturación)  «Es este» / «No es» de un cruce planilla ↔ libreto (Registros) CHANGELOG 2026-10-23
  DELETE /api/gestion/decisiones/{id}   (dev o facturación)  deshace una decisión (Registros)                          CHANGELOG 2026-10-23
  GET  /api/gestion/resumen            (puede_gestion)   KPIs
  GET  /api/gestion/proyectos          (puede_gestion)   lista con filtros
  GET  /api/gestion/proyectos/{id}     (puede_gestion)   detalle + línea de tiempo + enlaces a los módulos
  PATCH /api/gestion/proyectos/{id}    (puede_gestion)   zona, supervisor, nota, activo
  GET  /api/gestion/eventos/{id}/correo (puede_gestion)  el correo (HTML saneado) para verlo como captura
  GET  /api/gestion/pasos              (puede_gestion)   el flujo FL-03
  POST /api/gestion/recalcular         (dev)             relee SGS y el Centro MO y recalcula todos los proyectos
  POST /api/gestion/reclasificar       (dev)             reaplica las reglas de pasos a los correos guardados
  DELETE /api/gestion/proyectos/{id}   (dev)             borrar (pruebas)
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
import re
import time
import urllib.error
import urllib.request
from datetime import datetime, timedelta, timezone
from typing import Optional
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from urllib.parse import quote
from pydantic import BaseModel, Field
from collections import Counter

from sqlalchemy import Text, and_, cast, delete, desc, func, or_, select, text, tuple_
from sqlalchemy.ext.asyncio import AsyncSession

from app import config as _config
from app.api.deps import get_current_active_user, require_dev
from app.config import puede_gestion, settings
from app.db.database import AsyncSessionLocal, get_db
from app.models.models import (CasoEvento, CasoFacturacion, GestionAgente, GestionCarga, GestionDecision, GestionDeuda, GestionDrapeEm, GestionDrapeFactura,
                               GestionDrapePos, GestionEvento, GestionPms, GestionPmsLog, GestionPrecioMat, GestionProyecto, GestionSap, GestionScmLibreto,
                               GestionScmMat, GestionScmPrest, GestionScmLcl, GestionSinCaso, GestionSonny, GestionSonnyNodo, GestionSonnyPlanAct,
                               GestionSonnyPlanilla, RegistroSGO7)
from app.services import gestion_flujo as gf

# Datos del portal de facturación del cliente (DRAPE): visibles solo con config.GESTION_MOSTRAR_DRAPE = True (hoy False, hasta definir la fuente oficial)
gf.MOSTRAR_DRAPE = bool(getattr(_config, "GESTION_MOSTRAR_DRAPE", False))
_DIR_CARGAS = "/var/www/sgo7/registros/gestion/cargas"      # CHANGELOG 2026-10-25: el archivo de cada carga, para «Descargar» (fuera de git: registros/)
_MAX_ARCHIVO = 80 * 1024 * 1024
_TIPO_FUENTE = {"drape": "pantalla", "sgs": "vivo", "correo": "buzon"}      # tipo de fuente que pinta la interfaz; el resto: «exporte»
_URL_PUBLICA = os.getenv("GESTION_PUBLIC_URL", "")      # URL pública del módulo para los textos exportados (vacía = no se nombra)

router = APIRouter(prefix="/gestion", tags=["gestion"])
_TZ = ZoneInfo("America/Bogota")
_CENTRO_MO = "http://127.0.0.1:5088"
_cache_cm: dict = {}          # caché corta por proceso de las respuestas del Centro MO
_cache_cargas: dict = {}      # CHANGELOG 2026-10-23: fecha de la última carga de cada exporte (60 s); CHANGELOG 2026-10-23: también el panel de fuentes
_FUENTES = ("correo", "sgs", "sonny", "sgo7", "pms", "scm")
# CHANGELOG 2026-10-23: la lista de precios QP con la que el Balance Integrado de producción pasa cantidades a pesos (la misma que usa el PC)
_PRECIOS_QP = "/var/www/sgo7/sony-sgs/nuevo/precios_mdo_mat.json"
_cache_precios: dict = {}
_SISTEMAS = {"pms": ("solicitudes", "cambios"), "scm": ("libretos", "prestaciones", "materiales", "lcl"),      # CHANGELOG 2026-10-25: el exporte de LCL
             "sonny": ("historial", "lista", "nodos", "almacen", "actividades", "estados"),     # CHANGELOG 2026-10-25: Lista Cruzada, gestión almacén, Actividades, Estados
             "deuda": ("informe",), "sap": ("mb51",),                        # CHANGELOG 2026-10-23
             "drape": ("listas", "detalle", "em")}                           # CHANGELOG 2026-10-25: listas de facturas, detalle (posiciones) y EM a facturar
_cache_sonny: dict = {}      # CHANGELOG 2026-10-25: QR con planilla y macroprocesos de PMS que usan Sonny (se valida con un sello de las tablas)
_TAREAS_SONNY: set = set()
_cache_sap: dict = {}         # CHANGELOG 2026-10-23: corte, catálogo y precios de Enel (60 s por proceso; cada carga de SAP lo limpia)
_TAREAS_SAP: set = set()
ROLES_DECIDEN = ("desarrollador", "facturacion")      # CHANGELOG 2026-10-23: quién confirma un cruce («Es este» / «No es»); el módulo sigue solo para dev (CHANGELOG 2026-10-22)


# ── Permisos ─────────────────────────────────────────────────────────────────────────────────────────
def _agente(x_agente_token: str = Header(default="")):
    tok = (settings.GESTION_AGENTE_TOKEN or "").strip()
    if not tok or x_agente_token.strip() != tok:
        raise HTTPException(status_code=401, detail="Token del agente inválido")


async def require_gestion(user=Depends(get_current_active_user)):
    if not puede_gestion(user.rol, user.email):
        raise HTTPException(status_code=403, detail="Sin acceso a Gestión General")
    return user


# ── Utilidades ───────────────────────────────────────────────────────────────────────────────────────
def _ahora():
    return datetime.now(timezone.utc)


def _local(d: Optional[datetime]) -> Optional[datetime]:
    if d is None:
        return None
    if d.tzinfo is None:
        return d
    return d.astimezone(_TZ).replace(tzinfo=None)


def _aware(d: Optional[datetime]) -> Optional[datetime]:
    if d is None:
        return None
    return d if d.tzinfo else d.replace(tzinfo=_TZ)


def _iso(s: Optional[str]) -> Optional[datetime]:
    if not s:
        return None
    try:
        d = datetime.fromisoformat(str(s).replace("Z", "+00:00"))
    except ValueError:
        try:
            d = datetime.strptime(str(s)[:16], "%Y-%m-%d %H:%M")
        except ValueError:
            return None
    return _aware(d)


def _sha(s: str) -> str:
    return hashlib.sha256(s.encode("utf-8")).hexdigest()


def _sanear_html(h: str) -> str:
    h = h or ""
    h = re.sub(r"(?is)<script.*?</script>", "", h)
    h = re.sub(r"(?is)<iframe.*?</iframe>", "", h)
    h = re.sub(r'(?i)\son\w+\s*=\s*("[^"]*"|\'[^\']*\'|[^\s>]+)', "", h)
    h = re.sub(r'(?i)src="cid:[^"]*"', 'src="" alt="[imagen incrustada]"', h)
    h = re.sub(r"(?i)javascript:", "", h)
    return h[:400_000]


def _limpiar_asunto(a: str) -> str:
    return re.sub(r"^\s*((RE|RV|FW|FWD|TR|AW)\s*:\s*)+", "", a or "", flags=re.I).strip()[:200]


def _union(a: dict, b: dict) -> dict:
    out = dict(a or {})
    for k, v in (b or {}).items():
        out[k] = sorted(set(out.get(k, []) or []) | set(v or []))[:40]
    return out


# ── El Centro MO como árbitro (canal interno, 127.0.0.1:5088) ────────────────────────────────────────
def _cm_http(metodo: str, ruta: str, body: Optional[dict] = None, timeout: int = 300) -> dict:
    req = urllib.request.Request(_CENTRO_MO + ruta, method=metodo,
                                 headers={"X-SGO7-Interno": settings.SGO7_INTERNAL_TOKEN or "", "Content-Type": "application/json"},
                                 data=json.dumps(body).encode() if body is not None else None)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return json.loads(r.read() or b"{}")
    except urllib.error.HTTPError as e:
        try:
            det = json.loads(e.read() or b"{}").get("error")
        except Exception:
            det = ""
        raise HTTPException(status_code=502, detail=f"Centro MO {e.code}: {det or 'error'}")
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Centro MO no responde: {str(e)[:120]}")


async def _cm_sgs(codigos: list[str], detalle: bool, mat: bool = False) -> dict:
    """codigo → {existe, metodo, padres[...]} (con MO por familia si detalle; con el material de cada familia si mat, CHANGELOG 2026-10-23).
    En lotes de 40. Lo que trae material no se guarda en la caché de este proceso (pesa; el Centro MO ya lo cachea)."""
    out: dict = {}
    pend = []
    for c in dict.fromkeys(x.upper() for x in codigos if x):
        v = None if mat else _cache_cm.get(("sgs", c, detalle))
        if v and time.time() - v[0] < (900 if detalle else 6 * 3600):
            out[c] = v[1]
        else:
            pend.append(c)
    for i in range(0, len(pend), 40):
        lote = pend[i:i + 40]
        r = await asyncio.to_thread(_cm_http, "POST", "/interno/gestion/sgs", {"codigos": lote, "detalle": detalle, "mat": mat})
        for c, v in (r.get("codigos") or {}).items():
            out[c] = v
            ligero = v if not mat else dict(v, padres=[{k: x for k, x in pa.items() if k != "mat"} for pa in v.get("padres") or []])
            _cache_cm[("sgs", c, detalle)] = (time.time(), ligero)
            if detalle:
                _cache_cm[("sgs", c, False)] = (time.time(), ligero)
    return out


async def _cm_registros() -> list[dict]:
    v = _cache_cm.get("registros")
    if v and time.time() - v[0] < 120:
        return v[1]
    r = await asyncio.to_thread(_cm_http, "GET", "/interno/gestion/registros", None, 60)
    regs = r.get("registros") or []
    _cache_cm["registros"] = (time.time(), regs)
    return regs


# ── Modelos de entrada ───────────────────────────────────────────────────────────────────────────────
class CorreoIn(BaseModel):
    entry_id: str = Field(max_length=300)
    carpeta: str = Field(default="", max_length=200)
    asunto: str = Field(default="", max_length=500)
    de: str = Field(default="", max_length=200)
    de_correo: str = Field(default="", max_length=200)
    para: str = Field(default="", max_length=2000)
    cc: str = Field(default="", max_length=2000)
    fecha: str = Field(max_length=25)                      # "YYYY-MM-DD HH:MM" hora Colombia
    categorias: str = Field(default="", max_length=200)
    adjuntos: list[str] = Field(default_factory=list)
    conversacion: str = Field(default="", max_length=500)
    cuerpo: str = Field(default="")
    html: Optional[str] = None


class IngestaIn(BaseModel):
    buzon: str = Field(max_length=150)
    version: str = Field(default="", max_length=40)
    cursor: Optional[str] = None
    correos: list[CorreoIn] = Field(default_factory=list)


class SonnyFila(BaseModel):
    """Un cambio de estado del Historial de Planillas de Sonny, ya ligado (en el PC) a la Lista de Planillas."""
    id_cambio: str = Field(max_length=80)                  # columna Id del historial (única)
    id_planilla: str = Field(max_length=40)
    codigo: Optional[str] = Field(default=None, max_length=40)   # LCL 63… (o el ticket si la planilla no trae LCL)
    ticket: Optional[str] = Field(default=None, max_length=40)   # ticket DF… (mayor impacto: en SGS el proyecto es el ticket)
    nodo: str = Field(default="", max_length=200)
    supervisor: str = Field(default="", max_length=150)
    estado: str = Field(max_length=60)
    fecha: str = Field(max_length=25)                      # "YYYY-MM-DD HH:MM:SS" hora Colombia
    observacion: str = Field(default="", max_length=2000)


class SonnyIn(BaseModel):
    archivo: str = Field(default="", max_length=200)
    filas: list[SonnyFila] = Field(default_factory=list, max_length=1000)
    carga_id: Optional[int] = None                         # CHANGELOG 2026-10-23: la carga (sgo7_gg_cargas) de la que salen estas filas


# ── Modelos de entrada del hilo del dinero (los arman los scripts del PC; sin nombres de personas salvo el supervisor) ──
class SonnyPlanillaIn(BaseModel):
    id_planilla: str = Field(max_length=40)
    lcl: Optional[str] = Field(default=None, max_length=40)
    ticket: Optional[str] = Field(default=None, max_length=40)
    pdl: str = Field(default="", max_length=80)
    qr: str = Field(default="", max_length=100)
    odm: str = Field(default="", max_length=200)
    nodo: str = Field(default="", max_length=200)
    supervisor: str = Field(default="", max_length=150)
    estado: str = Field(default="", max_length=60)
    libreto_sonny: str = Field(default="", max_length=20)
    mo_plan: float = 0
    mat_plan: float = 0
    mo_real: float = 0
    mat_real: float = 0
    fecha_planilla: str = Field(default="", max_length=25)
    fecha_supervisor: str = Field(default="", max_length=25)
    fecha_ingeniero: str = Field(default="", max_length=25)
    fecha_enel: str = Field(default="", max_length=25)
    estado_supervisor: str = Field(default="", max_length=30)
    estado_ingeniero: str = Field(default="", max_length=30)


class SonnyPlanillasIn(BaseModel):
    archivo: str = Field(default="", max_length=200)
    planillas: list[SonnyPlanillaIn] = Field(default_factory=list, max_length=1000)
    recalcular: bool = True
    carga_id: Optional[int] = None                         # CHANGELOG 2026-10-23


class PmsSolicitudIn(BaseModel):
    request_code: str = Field(max_length=100)
    subprojects: str = Field(default="", max_length=300)
    lcls: list[str] = Field(default_factory=list, max_length=40)
    dfs: list[str] = Field(default_factory=list, max_length=40)
    pdls: list[str] = Field(default_factory=list, max_length=80)
    odms: list[str] = Field(default_factory=list, max_length=80)
    agp: str = Field(default="", max_length=60)
    macroproceso: str = Field(default="", max_length=100)
    services: float = 0
    materials: float = 0
    total: float = 0
    status: str = Field(default="", max_length=60)
    substatus: str = Field(default="", max_length=60)
    request_date: str = Field(default="", max_length=25)
    fecha_estado: str = Field(default="", max_length=25)


class PmsLogIn(BaseModel):
    huella: str = Field(min_length=16, max_length=64)
    request_code: str = Field(max_length=100)
    estado_ini: str = Field(default="", max_length=60)
    estado_fin: str = Field(default="", max_length=60)
    substatus: str = Field(default="", max_length=60)
    fecha: str = Field(max_length=26)


class PmsIn(BaseModel):
    archivo: str = Field(default="", max_length=200)
    solicitudes: list[PmsSolicitudIn] = Field(default_factory=list, max_length=1000)
    logs: list[PmsLogIn] = Field(default_factory=list, max_length=8000)
    recalcular: bool = True
    carga_solicitudes: Optional[int] = None                # CHANGELOG 2026-10-23: cargas de request.csv y de logs.csv
    carga_cambios: Optional[int] = None


class ScmMatIn(BaseModel):
    material: str = Field(max_length=40)
    posada: float = 0
    recuperada: float = 0
    unidad: str = Field(default="", max_length=20)         # CHANGELOG 2026-10-23: «Unidad de Medida» del Informe Materiales


class ScmPrestIn(BaseModel):
    """CHANGELOG 2026-10-23: una prestación del libreto (Informe Prestaciones): cantidad × precio QP = pesos (lo calcula el PC; el servidor lo
    compara con la lista de precios de producción)."""
    n: int = 0
    prestacion: str = Field(default="", max_length=20)
    unidad: str = Field(default="", max_length=20)
    cantidad: float = 0
    puntos_unit: float = 0
    puntos: float = 0
    precio: Optional[float] = None                         # None = el código no está en la lista de precios QP (no suma)
    pesos: float = 0
    ticket: str = Field(default="", max_length=40)


class ScmLibretoIn(BaseModel):
    lcl: str = Field(max_length=20)
    libreto: str = Field(max_length=20)
    descripcion: str = Field(default="", max_length=300)
    estado_libreto: str = Field(default="", max_length=12)
    estado_lcl: str = Field(default="", max_length=20)
    fecha_autocert: str = Field(default="", max_length=25)
    mo_pesos: float = 0
    puntos: float = 0
    n_prest: int = 0
    sin_precio: int = 0
    tickets: list[str] = Field(default_factory=list, max_length=60)
    materiales: list[ScmMatIn] = Field(default_factory=list, max_length=3000)
    prestaciones: Optional[list[ScmPrestIn]] = Field(default=None, max_length=3000)   # CHANGELOG 2026-10-23: None = el PC no mandó el detalle (se conserva)
    trabajos: Optional[list[str]] = Field(default=None, max_length=60)   # CHANGELOG 2026-10-23: «Trabajo» del Informe Materiales (= Orden de SAP); None = se conserva


class ScmIn(BaseModel):
    archivo: str = Field(default="", max_length=200)
    corte: str = Field(max_length=25)                    # fecha y hora del exporte de libretos (hora Colombia)
    precios_sha: str = Field(default="", max_length=64)  # huella de la lista de precios QP con que el PC pasó cantidades a pesos
    libretos: list[ScmLibretoIn] = Field(default_factory=list, max_length=500)
    recalcular: bool = True
    carga_libretos: Optional[int] = None                 # CHANGELOG 2026-10-23: cargas del Informe LCL_LCA, de Prestaciones y de Materiales
    carga_prestaciones: Optional[int] = None
    carga_materiales: Optional[int] = None


class DeudaFilaIn(BaseModel):
    item: str = Field(max_length=40)
    lcl: Optional[str] = Field(default=None, max_length=40)
    qr: str = Field(default="", max_length=100)
    id_planilla: str = Field(default="", max_length=40)
    pdl_ticket: str = Field(default="", max_length=80)
    libreto: str = Field(default="", max_length=20)
    mo: float = 0
    mat: float = 0
    estado: str = Field(default="", max_length=80)
    estado_desc: str = Field(default="", max_length=80)
    responsabilidad: str = Field(default="", max_length=40)
    fecha_estado: str = Field(default="", max_length=25)


class DeudaIn(BaseModel):
    archivo: str = Field(default="", max_length=200)
    corte: str = Field(max_length=25)                    # fecha del informe (el nombre del archivo la trae: «17-10-2026»)
    filas: list[DeudaFilaIn] = Field(default_factory=list, max_length=2000)
    recalcular: bool = True
    carga_id: Optional[int] = None                       # CHANGELOG 2026-10-23


class CargaIn(BaseModel):
    """CHANGELOG 2026-10-23: el registro de una carga del PC. «abrir» antes de mandar las filas (devuelve carga_id); «cerrar» al terminar (con lo que
    respondió el servidor). Mismo archivo (sistema, parte, sha256) = la misma carga: se actualizan la fecha y quién cargó."""
    accion: str = Field(pattern="^(abrir|cerrar)$")
    carga_id: Optional[int] = None
    sistema: str = Field(default="", max_length=20)
    parte: str = Field(default="", max_length=30)
    archivo: str = Field(default="", max_length=300)
    tamano: int = Field(default=0, ge=0)
    sha256: str = Field(default="", max_length=64)
    modificado: str = Field(default="", max_length=25)   # fecha de modificación del archivo (hora Colombia): la del exporte
    corte: str = Field(default="", max_length=25)        # fecha que trae el nombre del archivo, si la trae
    usuario: str = Field(default="", max_length=100)
    equipo: str = Field(default="", max_length=100)
    filas_archivo: int = Field(default=0, ge=0)
    filas_enviadas: int = Field(default=0, ge=0)
    duplicados: int = Field(default=0, ge=0)             # filas repetidas por llave natural EN EL ARCHIVO (el PC las resuelve)
    fuera_rango: int = Field(default=0, ge=0)            # fechas antes de 2020 o en el futuro EN EL ARCHIVO
    detalle: dict = Field(default_factory=dict)
    resultado: dict = Field(default_factory=dict)


class SapFilaIn(BaseModel):
    """CHANGELOG 2026-10-23: una fila del MB51 de SAP tal como la lee el PC. NO trae «Nombre del usuario» (cédulas): el PC no la manda y aquí no existe.
    Los textos llegan ya filtrados por el PC y el servidor los vuelve a filtrar (gestion_flujo.sap_texto_seguro)."""
    doc: str = Field(min_length=1, max_length=20)
    pos: int = Field(ge=0, le=999999)
    clase: str = Field(min_length=1, max_length=4)
    material: str = Field(default="", max_length=20)
    descripcion: str = Field(default="", max_length=120)
    cantidad: float = 0
    um: str = Field(default="", max_length=10)
    importe: float = 0
    f_contab: str = Field(default="", max_length=10)      # AAAA-MM-DD
    lote: str = Field(default="", max_length=20)
    centro: str = Field(default="", max_length=10)
    almacen: str = Field(default="", max_length=10)
    centro_rec: str = Field(default="", max_length=10)
    almacen_rec: str = Field(default="", max_length=10)
    lote_rec: str = Field(default="", max_length=20)
    carta: str = Field(default="", max_length=60)          # «Carta Entrega Trabajos»: el LCL (o el DF de mayor impacto)
    orden: str = Field(default="", max_length=20)
    orden_2: str = Field(default="", max_length=20)
    pedido: str = Field(default="", max_length=20)
    referencia: str = Field(default="", max_length=20)
    cuenta: str = Field(default="", max_length=20)
    texto_cab: str = Field(default="", max_length=60)
    texto: str = Field(default="", max_length=60)
    texto_clase: str = Field(default="", max_length=40)
    dh: str = Field(default="", max_length=1)


class SapIn(BaseModel):
    archivo: str = Field(default="", max_length=200)
    filas: list[SapFilaIn] = Field(default_factory=list, max_length=1000)
    carga_id: Optional[int] = None
    fin: bool = False                                      # último lote: rehace el precio de Enel, mide los controles y recalcula
    recalcular: bool = True


class SonnyNodoIn(BaseModel):
    """CHANGELOG 2026-10-25: un nodo de la «Lista Cruzada de Nodos» de Sonny, con lo de «Nodos por gestion almacen» y la confirmación de almacén del Historial de
    Actividades. NO trae Lider_cuadrilla (correo) ni Nombre_corto (persona): el PC no los manda y aquí no existen."""
    id_nodo: str = Field(min_length=1, max_length=20)
    codigo: str = Field(default="", max_length=40)          # columna Lcl: LCL 63… o ticket DF…
    odm: str = Field(default="", max_length=200)
    pdl: str = Field(default="", max_length=300)
    qr: str = Field(default="", max_length=100)
    nodo: str = Field(default="", max_length=120)
    estado_nodo: str = Field(default="", max_length=20)
    id_planillas: list[str] = Field(default_factory=list, max_length=20)
    estado_planilla: str = Field(default="", max_length=200)
    f_creacion: str = Field(default="", max_length=25)
    f_estado: str = Field(default="", max_length=25)
    estado_material: str = Field(default="", max_length=20)
    movimiento: str = Field(default="", max_length=20)
    f_almacen: str = Field(default="", max_length=25)
    n_almacen: int = Field(default=0, ge=0, le=1000)
    id_gestion: str = Field(default="", max_length=20)
    f_confirmado: str = Field(default="", max_length=25)


class SonnyPlanActIn(BaseModel):
    """CHANGELOG 2026-10-25: lo que dice el Historial de Actividades de una planilla (sin personas): creación, último cambio de estado, primera asignación a un
    firmante después de ese cambio, cuántas firmas y por quién (enel / dominion / interventoria) y la última acción."""
    id_planilla: str = Field(min_length=1, max_length=40)
    f_creada: str = Field(default="", max_length=25)
    estado_act: str = Field(default="", max_length=60)
    f_estado_act: str = Field(default="", max_length=25)
    f_firma: str = Field(default="", max_length=25)
    n_firmas: int = Field(default=0, ge=0, le=100000)
    roles: dict = Field(default_factory=dict)
    f_ultima: str = Field(default="", max_length=25)
    que_ultima: str = Field(default="", max_length=20)
    n_act: int = Field(default=0, ge=0, le=1000000)


class SonnyNodosIn(BaseModel):
    archivo: str = Field(default="", max_length=200)
    carga_id: Optional[int] = None
    nodos: list[SonnyNodoIn] = Field(default_factory=list, max_length=1000)
    planillas: list[SonnyPlanActIn] = Field(default_factory=list, max_length=1000)
    fin: bool = False                                      # último lote de nodos: los que no vinieron en esta carga dejan de estar vigentes
    barrer: Optional[list[str]] = Field(default=None, max_length=50)   # con fin: solo estos códigos (pruebas); sin él, toda la tabla
    recalcular: bool = True


class ScmLclFilaIn(BaseModel):
    """CHANGELOG 2026-10-25: una fila del exporte de LCL de SCM (Informe LCL/LCA, pantalla «Informe Penalización y Bonificaciones»): una LCL × Trabajo. NO trae
    «Nombre técnico EDE» ni «Ejecutor» (personas y la empresa) ni la descripción: el PC no los lee y aquí no existen."""
    lcl: str = Field(min_length=1, max_length=20)
    trabajo: str = Field(min_length=1, max_length=24)        # orden de SAP (12 dígitos con ceros a la izquierda)
    estado: str = Field(default="", max_length=20)
    fecha_envio: str = Field(default="", max_length=25)
    inicio_validez: str = Field(default="", max_length=25)
    fin_validez: str = Field(default="", max_length=25)
    importe: float = 0
    obra: Optional[int] = Field(default=None, ge=0, le=100000)
    acumulativa: Optional[bool] = None
    tam: str = Field(default="", max_length=10)


class ScmLclIn(BaseModel):
    archivo: str = Field(default="", max_length=200)
    corte: str = Field(max_length=25)                        # fecha y hora del exporte (la trae el nombre: lcl_export20261023160236)
    carga_id: Optional[int] = None
    filas: list[ScmLclFilaIn] = Field(default_factory=list, max_length=3000)   # el PC manda cada LCL completa en un mismo lote
    fin: bool = False                                        # último lote: recalcula los casos de lo que cambió en esta carga
    recalcular: bool = True


class DrapeFacturaIn(BaseModel):
    """CHANGELOG 2026-10-25: una factura de una lista de DRAPE (Borrador, Enviadas, Rechazadas, Registradas, Pagadas/compensadas). NO trae «Nota para la
    Administración», «Mensaje para proveedor» ni la «Nota» de las rechazadas (texto libre: puede traer nombres), ni el emisor, ni códigos fiscales:
    el PC no los manda y aquí no existen. Los campos del detalle (vencimiento, documento de pago, condiciones) solo llegan cuando se abrió."""
    drape_id: str = Field(min_length=4, max_length=30)
    factura: str = Field(default="", max_length=30)
    estado: str = Field(pattern="^(borrador|enviada|rechazada|registrada|pagada)$")
    destinatario: str = Field(default="", max_length=80)
    f_factura: str = Field(default="", max_length=25)
    n_registro: str = Field(default="", max_length=20)
    f_registro: str = Field(default="", max_length=25)
    pedidos: list[str] = Field(default_factory=list, max_length=400)      # «Num. Pedido» separado por comas (casi siempre LCL 63…)
    ems: list[str] = Field(default_factory=list, max_length=600)          # «Numero EM/Conformidad» (HES, 5003…)
    f_pago: str = Field(default="", max_length=25)
    total: float = 0
    moneda: str = Field(default="", max_length=5)
    f_envio: str = Field(default="", max_length=25)
    f_creacion: str = Field(default="", max_length=25)


class DrapePosIn(BaseModel):
    """CHANGELOG 2026-10-25: una línea del «Resumen de Posición» del detalle de una factura (LCL/posición e importe)."""
    posicion: str = Field(min_length=1, max_length=12)
    producto: str = Field(default="", max_length=40)
    descripcion: str = Field(default="", max_length=200)
    pedido_pos: str = Field(default="", max_length=40)                  # «6301018885/20»
    unidad: str = Field(default="", max_length=10)
    cantidad: float = 0
    precio: float = 0
    impuesto: str = Field(default="", max_length=12)
    importe: float = 0


class DrapeDetalleIn(BaseModel):
    """CHANGELOG 2026-10-25: lo que se lee al abrir una factura: de la Cabecera solo fechas, registro, pago, condiciones y totales (NUNCA datos bancarios), los
    eventos (código y fecha) y el Resumen de Posición. pos_lcls = las LCL cuyas posiciones vienen completas (el PC puede leer solo las de unos casos)."""
    drape_id: str = Field(min_length=4, max_length=30)
    estado_detalle: str = Field(default="", max_length=40)
    f_vence: str = Field(default="", max_length=25)
    doc_pago: str = Field(default="", max_length=30)
    f_pago: str = Field(default="", max_length=25)
    condiciones: str = Field(default="", max_length=40)
    total_neto: Optional[float] = None
    total_impuesto: Optional[float] = None
    total: Optional[float] = None
    n_pos: int = Field(default=0, ge=0, le=5000)
    suma_pos: Optional[float] = None
    pos_lcls: list[str] = Field(default_factory=list, max_length=400)
    eventos: list[list[str]] = Field(default_factory=list, max_length=40)
    posiciones: list[DrapePosIn] = Field(default_factory=list, max_length=3000)


class DrapeEmIn(BaseModel):
    em: str = Field(min_length=4, max_length=20)
    fecha: str = Field(default="", max_length=25)
    importe: float = 0
    lcl: str = Field(default="", max_length=20)
    empresa: str = Field(default="", max_length=80)
    moneda: str = Field(default="", max_length=5)


class DrapeIn(BaseModel):
    archivo: str = Field(default="", max_length=200)
    corte: str = Field(max_length=25)                        # corte de la carga de DRAPE (la lista o el detalle)
    carga_id: Optional[int] = None
    facturas: list[DrapeFacturaIn] = Field(default_factory=list, max_length=1500)
    detalles: list[DrapeDetalleIn] = Field(default_factory=list, max_length=200)
    ems: list[DrapeEmIn] = Field(default_factory=list, max_length=3000)
    listas_completas: list[str] = Field(default_factory=list, max_length=6)   # con fin: las listas (y «em») que vinieron enteras en esta carga
    fin: bool = False
    recalcular: bool = True
    barrer: Optional[list[str]] = Field(default=None, max_length=50)          # con fin: solo estos DRAPE ID pueden dejar de estar vigentes (pruebas)


class DecisionIn(BaseModel):
    """CHANGELOG 2026-10-23: confirmación humana de un cruce planilla ↔ libreto (fase 3 del hilo del dinero)."""
    proyecto_id: Optional[int] = None
    tipo: str = Field(default="planilla_libreto", pattern="^planilla_libreto$")
    lcl: str = Field(max_length=20)
    libreto: str = Field(max_length=20)
    id_planilla: str = Field(max_length=40)
    decision: str = Field(pattern="^(si|no)$")
    nota: Optional[str] = Field(default=None, max_length=500)


class NacerIn(BaseModel):
    limite: int = Field(default=40, ge=1, le=200)        # cuántos códigos se revisan en esta llamada
    desde: Optional[str] = Field(default=None, max_length=10)   # «2026-03-25»: solo códigos con movimiento desde esa fecha
    simular: bool = False
    codigos: Optional[list[str]] = Field(default=None, max_length=50)   # solo estos códigos (pruebas)


class ProyectoPatch(BaseModel):
    zona: Optional[str] = Field(default=None, max_length=40)
    supervisor: Optional[str] = Field(default=None, max_length=150)
    nota: Optional[str] = None
    activo: Optional[bool] = None
    tipo: Optional[str] = Field(default=None, max_length=30)


# ── ¿De qué proyecto es este correo? (reglas del Centro MO) ──────────────────────────────────────────
def _tipo_de(codigo: str, padre: dict) -> str:
    if (codigo or "").startswith("DF"):
        return "mayor_impacto"
    return "linea_viva" if (padre or {}).get("zona") == "LV&TLC" else "programado"


async def _hilos_por_padre(db: AsyncSession, codigo: str) -> dict:
    """code del padre → hilos de correo que ya tiene (para no partir una conversación entre gemelos)."""
    filas = (await db.execute(select(GestionProyecto.clave, GestionEvento.asunto).join(GestionEvento, GestionEvento.proyecto_id == GestionProyecto.id)
                              .where(GestionProyecto.codigo == codigo, GestionEvento.fuente == "correo", GestionEvento.directo.is_(True)))).all()
    out: dict = {}
    for clave, asunto in filas:
        out.setdefault(clave, set()).add(gf.hilo_normalizado(asunto or ""))
    return out


async def _proyecto_para(db: AsyncSession, codigo: str, padre: dict, asunto: str, hilo: str) -> GestionProyecto:
    """El caso es el proyecto PADRE de SGS (clave = PRO-…); `codigo` es el LCL o ticket que comparte con sus gemelos."""
    clave = padre["code"]
    p = (await db.execute(select(GestionProyecto).where(GestionProyecto.clave == clave))).scalar_one_or_none()
    if p is None:
        p = GestionProyecto(clave=clave, codigo=codigo, padre={k: padre.get(k) for k in ("id", "code", "name", "order_number", "supervisor", "departamento", "zona", "codes")},
                            tipo=_tipo_de(codigo, padre), nombre=_limpiar_asunto(asunto) or codigo, hilo=hilo or "", codigos={}, fijados={}, alertas=[],
                            activo=True, n_eventos=0, n_correos=0, supervisor=padre.get("supervisor") or None, zona=padre.get("zona"))
        db.add(p)
        await db.flush()
    return p


async def _por_hilo(db: AsyncSession, hilo: str) -> Optional[GestionProyecto]:
    if not hilo:
        return None
    r = (await db.execute(select(GestionProyecto).where(GestionProyecto.hilo == hilo).limit(2))).scalars().all()
    return r[0] if len(r) == 1 else None


async def _hermanos(db: AsyncSession, p: GestionProyecto) -> list[GestionProyecto]:
    if not p.codigo:
        return []
    return list((await db.execute(select(GestionProyecto).where(GestionProyecto.codigo == p.codigo, GestionProyecto.id != p.id))).scalars().all())


# ── Enlace con las otras fuentes (SGS en vivo, Centro MO, Sonny, Registros) ──────────────────────────
_ACCIONES = {
    ("liquidador", "liquidador_subir_excel"): (21, "planilla en el Liquidador", "El Liquidador Automático recibió la planilla", "AUXSGS"),
    ("liquidador", "liquidador_resultado"): (21, "resultado del Liquidador", "Resultado del Liquidador Automático", "AUXSGS"),
    ("liquidador", "liquidador_crear_nodo"): (12, "nodo creado", "Nodo creado en SGS desde el Liquidador", "AUXSGS"),
    ("balance", "balance_integrado"): (24, None, "Balance Integrado: planilla Sonny contra SGS", "F"),
    ("balance", "materiales_cargue"): (13, None, "Materiales cargados desde el Balance", "L"),
}
_CASO_EVENTOS = {
    "consulta": (24, None, "Balance Integrado consultado (caso de facturación)", "F", "sonny"),
    "cuadro": (24, "cuadra", "El Balance Integrado ya cuadra", "F", "sonny"),
    "estado_manual": (24, "estado manual", "Facturación cambió el estado del caso", "F", "sgo7"),
    "excel_generado": (24, "excel", "Excel final del Balance generado", "F", "sgo7"),
    "orden_enviada": (21, "mandado a liquidar", "Facturación mandó el caso a liquidar", "F", "sgo7"),
    "orden_en_liquidacion": (21, "en liquidación", "El Liquidador tomó la orden", "AUXSGS", "sgo7"),
    "orden_liquidada": (21, "orden liquidada", "La orden se liquidó en el Liquidador", "AUXSGS", "sgo7"),
    "orden_con_fallos": (21, "orden con fallos", "La orden se liquidó con fallos", "AUXSGS", "sgo7"),
    "orden_cancelada": (21, "orden cancelada", "Orden de liquidación cancelada", "F", "sgo7"),
}
_ESTADO_CM = {"completado": "completado (Centro MO)", "descartado": "descartado por Centro MO", "planilla_lista": "planilla lista",
              "listo_auto": "listo para planilla", "falta_info": "falta información", "esperando_supervisor": "esperando al supervisor"}


def _sgs_resumen_padre(padres: list[dict]) -> dict:
    """El pilar SGS del caso: la MO acumulada (todas las liquidaciones de la familia) y cada liquidación por día. Si el
    caso junta varios padres del mismo supervisor (SGS partió el trabajo), se suman y cada liquidación dice de cuál es."""
    lotes = sorted((dict(l, padre=pa.get("code")) for pa in padres for l in pa.get("lotes") or []), key=lambda l: l.get("ts") or "")
    base = padres[0]
    return {
        "consultado_en": _ahora().isoformat(),
        "padre": {k: base.get(k) for k in ("code", "name", "order_number", "supervisor", "departamento", "zona")}
                 | {"codes": [pa.get("code") for pa in padres],
                    "detalle": [{k: pa.get(k) for k in ("code", "nodos", "mo_total", "n_avances", "por_mes")} for pa in padres]},
        "total_mo": round(sum(float(pa.get("mo_total") or 0) for pa in padres), 2), "nodos": sum(int(pa.get("nodos") or 0) for pa in padres),
        # CHANGELOG 2026-10-23: todas las liquidaciones (antes las últimas 80): son las filas que suman la caja de SGS en «¿De dónde sale?»
        "n_lotes": len(lotes), "lotes": lotes, "ultima_liq": lotes[-1] if lotes else None,
    }


def _sumar_mat(mats: list) -> Optional[dict]:
    """CHANGELOG 2026-10-23: suma el material de varios padres (lo que trae el Centro MO con "mat": true). leido = todos los padres leídos.
    CHANGELOG 2026-10-25: también las bandejas por código (mov_codigo: despachado y devuelto de las bandejas aprobadas, última salida); bandejas = todos
    los padres trajeron sus bandejas."""
    if not mats:
        return None
    pc: dict = {}
    vc: dict = {}
    nom: dict = {}
    mv: dict = {}
    sc = {"consumido_valor": 0.0, "pendiente_valor": 0.0, "n_lineas": 0, "neto_valor": 0.0, "por_liquidar_valor": 0.0, "sobre_valor": 0.0}
    tot = {"consumido_valor": 0.0, "pendiente_valor": 0.0, "propio_consumido_valor": 0.0}
    for m in mats:
        if not m:
            continue
        for cod, v in (m.get("por_codigo") or {}).items():
            x = pc.setdefault(cod, [0.0, 0.0, 0.0])
            x[0] += float(v[0] or 0)
            x[1] += float(v[1] or 0)
            x[2] = max(x[2], float(v[2] or 0))
        # CHANGELOG 2026-10-23 (Centro MO): valor consumido y despachado por código línea a línea, nombre y unidad de SGS, material sin código
        for cod, v in (m.get("valor_codigo") or {}).items():
            x = vc.setdefault(cod, [0.0, 0.0])
            x[0] += float(v[0] or 0)
            x[1] += float(v[1] or 0)
        for cod, v in (m.get("nombres") or {}).items():
            nom.setdefault(cod, v)
        for cod, v in (m.get("mov_codigo") or {}).items():          # CHANGELOG 2026-10-25
            y = mv.setdefault(cod, {"d": 0.0, "v": 0.0, "dc": 0.0, "f": None, "b": "", "f0": None, "nd": 0, "nv": 0})
            for k in ("d", "v", "dc"):
                y[k] += float(v.get(k) or 0)
            for k in ("nd", "nv"):
                y[k] += int(v.get(k) or 0)
            if v.get("f") and (y["f"] is None or v["f"] > y["f"] or (v["f"] == y["f"] and (v.get("b") or "") > y["b"])):
                y["f"], y["b"] = v["f"], v.get("b") or ""
            if v.get("f0") and (y["f0"] is None or v["f0"] < y["f0"]):
                y["f0"] = v["f0"]
        for k in sc:
            sc[k] += float((m.get("sin_codigo") or {}).get(k) or 0)
        for k in tot:
            tot[k] += float(m.get(k) or 0)
    out = {"por_codigo": {k: [round(a, 3), round(b, 3), c] for k, (a, b, c) in sorted(pc.items())}, **{k: round(v, 2) for k, v in tot.items()},
           "leido": all(bool(m and m.get("leido")) for m in mats)}
    if vc or nom:
        out.update(valor_codigo={k: [round(a, 2), round(b, 2)] for k, (a, b) in sorted(vc.items())}, nombres=nom,
                   sin_codigo={k: (int(v) if k == "n_lineas" else round(v, 2)) for k, v in sc.items()})
    # CHANGELOG 2026-10-25: las bandejas solo valen si TODOS los padres las trajeron (si no, el por liquidar del LCL saldría corto)
    bands = [m for m in mats if m]
    out["bandejas"] = bool(bands) and all(m.get("bandejas") and "mov_codigo" in m for m in bands)
    if out["bandejas"]:
        out["mov_codigo"] = {k: dict(v, d=round(v["d"], 3), v=round(v["v"], 3), dc=round(v["dc"], 3)) for k, v in sorted(mv.items())}
    errs = [m.get("bandejas_error") for m in bands if m.get("bandejas_error")]
    if errs:
        out["bandejas_error"] = errs[0]
    return out


def _cod_valido(c: Optional[str]) -> bool:
    return bool(c) and (bool(gf.lcl_norm(c)) or bool(re.fullmatch(r"DF20\d{10}", c or "")))


async def _cargas(db: AsyncSession) -> dict:
    """CHANGELOG 2026-10-23: hasta cuándo llega cada exporte cargado (caché de 60 s por proceso; las cargas la limpian)."""
    v = _cache_cargas.get("c")
    if v and time.time() - v[0] < 60:
        return v[1]
    scm = (await db.execute(select(func.max(GestionScmLibreto.corte)))).scalar()
    deu = (await db.execute(select(func.max(GestionDeuda.corte)))).scalar()
    son = (await db.execute(select(func.max(GestionSonnyPlanilla.cargado_en)))).scalar()
    pms = (await db.execute(select(func.max(GestionPms.cargado_en)))).scalar()
    sap = await _sap_meta(db)
    scm_lcl = (await db.execute(select(func.max(GestionScmLcl.corte)))).scalar()          # CHANGELOG 2026-10-25: hasta cuándo llega el exporte de LCL
    drape = (await db.execute(select(func.max(GestionDrapeFactura.corte)))).scalar()      # CHANGELOG 2026-10-25: última carga de DRAPE
    out = {"scm": _local(scm), "deuda": deu, "sonny": bool(son), "pms": bool(pms), "sonny_cargado": _local(son), "pms_cargado": _local(pms),
           "sap": sap.get("corte"), "scm_lcl": _local(scm_lcl), "drape": _local(drape) if gf.MOSTRAR_DRAPE else None}      # CHANGELOG 2026-10-25
    _cache_cargas["c"] = (time.time(), out)
    return out


async def _sap_meta(db: AsyncSession) -> dict:
    """CHANGELOG 2026-10-23: hasta dónde llega el MB51 cargado (última fecha de contabilización del consumo 261: «SAP al dd/mm»), su catálogo de códigos
    (material que maneja SAP: stock de Enel) y cuántas filas hay. Caché de 60 s por proceso (cada carga de SAP la limpia)."""
    # Con 2 workers de gunicorn una caché por tiempo dejaría a uno con el MB51 viejo tras una carga: se valida con un sello barato de la tabla
    (n, corte, desde, mx, act), = (await db.execute(select(func.count(), func.max(GestionSap.f_contab).filter(GestionSap.clase == "261"),
                                                            func.min(GestionSap.f_contab), func.max(GestionSap.id), func.max(GestionSap.actualizado_en)))).all()
    sello = (n, corte, mx, act)
    v = _cache_sap.get("meta")
    if v and v[0] == sello:
        return v[1]
    cat = set((await db.execute(select(GestionSap.material).distinct())).scalars().all()) if n else set()
    out = {"hay": bool(n and corte), "n": int(n or 0), "corte": datetime(corte.year, corte.month, corte.day) if corte else None,
           "desde": datetime(desde.year, desde.month, desde.day) if desde else None, "catalogo": cat}
    _cache_sap["meta"] = (sello, out)
    return out


async def _precios_enel(db: AsyncSession) -> tuple[dict, dict]:
    """CHANGELOG 2026-10-23: {código: {precio, um, n, fuente}} de sgo7_gg_precio_mat y de dónde salen (carga, corte, cuántos códigos). Caché de 60 s."""
    sello = tuple((await db.execute(select(func.count(), func.max(GestionPrecioMat.id), func.max(GestionPrecioMat.actualizado_en)))).all()[0])
    v = _cache_sap.get("precios")
    if v and v[0] == sello:
        return v[1], v[2]
    xs = (await db.execute(select(GestionPrecioMat))).scalars().all()
    out = {x.material: {"precio": float(x.precio or 0), "um": x.um or "", "n": int(x.n or 0), "fuente": x.fuente or "",
                        "fecha": x.fecha.isoformat() if x.fecha else None} for x in xs if float(x.precio or 0) > 0}
    cid = max((x.carga_id for x in xs if x.carga_id), default=None)
    carga = _carga_out(await db.get(GestionCarga, cid)) if cid else None
    info = {"n": len(out), "carga": carga, "fecha": max((x.fecha for x in xs if x.fecha), default=None),
            "tabla": "sgo7_gg_precio_mat", "regla": "mediana de Importe ML / cantidad de las descargas (clase 261) del MB51, por código de Enel"}
    if info["fecha"]:
        info["fecha"] = info["fecha"].isoformat()
    _cache_sap["precios"] = (sello, out, info)
    return out, info


_SAP_COLS = ("doc", "pos", "clase", "material", "descripcion", "cantidad", "um", "importe", "f_contab", "lote", "zona", "lcl", "df", "orden", "orden_2",
             "texto_tipo", "texto_cab", "texto", "carga_id")


async def _sap_del_caso(db: AsyncSession, cod: str, es_df: bool, lcls: list, libs: list, trab_lcl: Optional[set] = None) -> list[dict]:
    """CHANGELOG 2026-10-23: las filas de SAP (consumo y recuperado) de un caso con la confianza del cruce. LCL: «Carta Entrega Trabajos» = el LCL (llave
    exacta; alta si además la orden es el «Trabajo» de un libreto de SCM del LCL, media si no). Ticket DF: la carta es el DF (alta) o el LCL
    de sus planillas con la orden de uno de SUS libretos (Orden de SAP = Trabajo de SCM, alta). El ODM de PMS o de Sonny no se usa.
    CHANGELOG 2026-10-25: en un caso LCL también vale alta la orden que es un «Trabajo» de la LCL en el exporte de LCL de SCM (trab_lcl; cubre las LCL sin
    libreto). En un DF no: los Trabajo de la LCL son de todos los tickets que la comparten."""
    meta = await _sap_meta(db)
    if not meta["hay"] or not cod:
        return []
    trab = {str(t) for l in libs for t in (l.trabajos or [])}
    trab_l = set(trab_lcl or ()) if not es_df else set()
    if es_df:
        conds = [GestionSap.df == cod]
        if lcls and trab:
            conds.append(and_(GestionSap.lcl.in_(sorted(lcls)), GestionSap.orden.in_(sorted(trab))))
    else:
        conds = [GestionSap.lcl == cod]
    rows = (await db.execute(select(GestionSap).where(or_(*conds), GestionSap.clase.in_(gf.SAP_CONSUMO + gf.SAP_RECUPERADO))
                             .order_by(GestionSap.f_contab, GestionSap.doc, GestionSap.pos))).scalars().all()
    out = []
    for r in rows:
        d = {k: getattr(r, k) for k in _SAP_COLS}
        if es_df and r.df == cod:
            d["confianza"], d["regla"] = "alta", "la «Carta Entrega Trabajos» es el ticket del caso"
        elif es_df:
            d["confianza"], d["regla"] = "alta", "LCL de la planilla del ticket y orden = «Trabajo» de su libreto en SCM"
        elif r.orden and r.orden in trab:
            d["confianza"], d["regla"] = "alta", "LCL del caso y orden = «Trabajo» de un libreto de SCM"
        elif r.orden and r.orden in trab_l:
            d["confianza"], d["regla"] = "alta", "LCL del caso y orden = «Trabajo» de la LCL en el exporte de LCL de SCM"
        else:
            d["confianza"], d["regla"] = "media", "LCL del caso (la orden no está entre los «Trabajo» de sus libretos ni de la LCL en SCM)"
        out.append(d)
    return out


async def _rivales(db: AsyncSession, cods: list[str]) -> list[dict]:
    """Casos que comparten alguno de estos códigos (gemelos): para elegir de cuál es una planilla de Sonny."""
    if not cods:
        return []
    return [dict(id=x.id, tipo=x.tipo, codigo=x.codigo, supervisor=(x.padre or {}).get("supervisor") or x.supervisor or "")
            for x in (await db.execute(select(GestionProyecto).where(GestionProyecto.codigo.in_(cods)))).scalars().all()]


async def _fuentes_dinero(db: AsyncSession, p: GestionProyecto) -> dict:
    """CHANGELOG 2026-10-23: todo lo de Sonny (Lista), PMS, SCM y la deuda que es de este caso. LCL: lo del LCL (con gemelos, compartido).
    Ticket DF: sus planillas (ticket), sus solicitudes (subprojects) y los libretos de su LCL enlazados a sus planillas o que
    nombran el DF; un DF de mayor impacto no trae LCL en PMS: lo da su planilla (R3)."""
    fd = {"planillas": [], "hist": {}, "pms": [], "logs": [], "libretos": [], "mats": {}, "deuda": None, "lcls": [], "cargas": await _cargas(db), "sap": [],
          "nodos": [], "act": {}, "scm_lcl": [], "libs_lcl": {}, "drape": [], "drape_pos": [], "drape_em": []}     # CHANGELOG 2026-10-25, CHANGELOG 2026-10-25, CHANGELOG 2026-10-25
    cod = p.codigo
    if not cod:
        return fd
    es_df = cod.startswith("DF")
    pls = (await db.execute(select(GestionSonnyPlanilla).where(or_(GestionSonnyPlanilla.lcl == cod, GestionSonnyPlanilla.ticket == cod))
                            .order_by(GestionSonnyPlanilla.id_planilla))).scalars().all()
    ids = [x.id_planilla for x in pls]
    hist: dict = {}
    if ids:
        for idp, est, f in (await db.execute(select(GestionSonny.id_planilla, GestionSonny.estado, GestionSonny.fecha).where(GestionSonny.id_planilla.in_(ids))
                                              .order_by(GestionSonny.fecha, GestionSonny.id_cambio))).all():
            h = hist.setdefault(idp, {"ultimo": None, "fecha": None, "aprobado": None, "hitos": {}})
            h["ultimo"], h["fecha"] = est, f
            if est == "Libretos Aprobados":
                h["aprobado"] = f
            if f is not None:                       # CHANGELOG 2026-10-25: la primera vez en cada estado (la cadena y «¿Cómo fluye?»)
                h["hitos"].setdefault(gf.estado_sonny(est), _local(f).isoformat())
    pms = (await db.execute(select(GestionPms).where(or_(GestionPms.lcls.contains([cod]), GestionPms.dfs.contains([cod]))).order_by(GestionPms.request_code))).scalars().all()
    lcls = {cod} if not es_df else ({x.lcl for x in pls if gf.lcl_norm(x.lcl)} | {l for x in pms for l in (x.lcls or [])})
    libs = (await db.execute(select(GestionScmLibreto).where(GestionScmLibreto.lcl.in_(sorted(lcls))).order_by(GestionScmLibreto.lcl, GestionScmLibreto.libreto))).scalars().all() if lcls else []
    libs_lcl: dict = {}                       # CHANGELOG 2026-10-25: todos los libretos de cada LCL del caso (en DF, también los de otros tickets)
    for l in libs:
        libs_lcl.setdefault(l.lcl, []).append(l)
    if es_df:
        libs = [l for l in libs if (l.planilla and l.planilla in ids) or cod in (l.pdls or []) or cod in (l.tickets or [])]
    logs = (await db.execute(select(GestionPmsLog).where(GestionPmsLog.request_code.in_([x.request_code for x in pms])).order_by(GestionPmsLog.fecha))).scalars().all() if pms else []
    vivos = [(l.lcl, l.libreto) for l in libs if l.estado_libreto not in gf.LIBRETO_MUERTO]
    mats: dict = {}
    if vivos:
        # CHANGELOG 2026-10-23: con la fecha de autocertificación más nueva de los libretos que posan cada código (lo que SAP aún no ve no es un 0)
        for mat, pos, rec, uni, fa in (await db.execute(select(GestionScmMat.material, func.sum(GestionScmMat.posada), func.sum(GestionScmMat.recuperada),
                                                               func.max(GestionScmMat.unidad), func.max(GestionScmLibreto.fecha_autocert))
                                                        .join(GestionScmLibreto, and_(GestionScmLibreto.lcl == GestionScmMat.lcl, GestionScmLibreto.libreto == GestionScmMat.libreto))
                                                        .where(tuple_(GestionScmMat.lcl, GestionScmMat.libreto).in_(vivos)).group_by(GestionScmMat.material))).all():
            mats[mat] = [float(pos or 0), float(rec or 0), uni or "", _local(fa)]
    corte = fd["cargas"].get("deuda")
    if corte:
        conds = []
        if not es_df:
            conds.append(GestionDeuda.lcl == cod)
        else:
            conds.append(GestionDeuda.pdl_ticket == cod)
            if ids:
                conds.append(GestionDeuda.id_planilla.in_(ids))
        filas = (await db.execute(select(GestionDeuda).where(GestionDeuda.corte == corte, or_(*conds)))).scalars().all()
        fd["deuda"] = {"corte": _local(corte).date().isoformat(), "filas": [dict(libreto=f.libreto, id_planilla=f.id_planilla, mo=f.mo, mat=f.mat, estado=f.estado,
                                                                                    estado_desc=f.estado_desc, responsabilidad=f.responsabilidad) for f in filas]}
    fd.update(planillas=pls, hist=hist, pms=pms, logs=logs, libretos=libs, mats=mats, lcls=sorted(lcls))
    # CHANGELOG 2026-10-23: SAP (MB51). En LCL las órdenes de todos sus libretos (cualquier estado) dan la confianza; en DF solo los libretos del ticket.
    # CHANGELOG 2026-10-25: el exporte de LCL de SCM (filas vigentes de las LCL del caso)
    fd["scm_lcl"] = (await db.execute(select(GestionScmLcl).where(GestionScmLcl.lcl.in_(sorted(lcls)), GestionScmLcl.vigente.is_(True))
                                      .order_by(GestionScmLcl.lcl, GestionScmLcl.trabajo))).scalars().all() if lcls else []
    fd["libs_lcl"] = libs_lcl
    fd["sap"] = await _sap_del_caso(db, cod, es_df, sorted(lcls), libs, {x.trabajo for x in fd["scm_lcl"] if x.lcl == cod and x.trabajo})
    # CHANGELOG 2026-10-25: los nodos vigentes de Sonny del código (Lista Cruzada; la columna Lcl es el LCL o el ticket DF) y la actividad de sus planillas
    fd["nodos"] = (await db.execute(select(GestionSonnyNodo).where(or_(GestionSonnyNodo.codigo == cod, GestionSonnyNodo.codigos.contains([cod])), GestionSonnyNodo.vigente.is_(True))
                                    .order_by(GestionSonnyNodo.f_estado, GestionSonnyNodo.id_nodo))).scalars().all()
    fd["act"] = {a.id_planilla: a for a in (await db.execute(select(GestionSonnyPlanAct).where(GestionSonnyPlanAct.id_planilla.in_(ids)))).scalars().all()} if ids else {}
    # CHANGELOG 2026-10-25 DRAPE: las facturas vigentes que nombran alguna LCL del caso («Num. Pedido»), las posiciones de esas LCL y sus EM pendientes de facturar.
    # Si hay lectura de DRAPE se pregunta a la tabla (no a la caché de 60 s: con 2 workers uno diría «sin datos de DRAPE» justo tras una carga)
    fd["cargas"] = dict(fd["cargas"], drape=_local((await db.execute(select(func.max(GestionDrapeFactura.corte)))).scalar()) if gf.MOSTRAR_DRAPE else None)   # CHANGELOG 2026-10-25
    if lcls and fd["cargas"].get("drape"):
        ls_ = sorted(lcls)
        fd["drape"] = (await db.execute(select(GestionDrapeFactura).where(GestionDrapeFactura.vigente.is_(True), or_(*[GestionDrapeFactura.lcls.contains([l]) for l in ls_]))
                                        .order_by(GestionDrapeFactura.f_factura, GestionDrapeFactura.factura))).scalars().all()
        ids_d = [x.drape_id for x in fd["drape"]]
        fd["drape_pos"] = (await db.execute(select(GestionDrapePos).where(GestionDrapePos.drape_id.in_(ids_d), GestionDrapePos.lcl.in_(ls_))
                                            .order_by(GestionDrapePos.drape_id, GestionDrapePos.id))).scalars().all() if ids_d else []
        fd["drape_em"] = (await db.execute(select(GestionDrapeEm).where(GestionDrapeEm.lcl.in_(ls_), GestionDrapeEm.vigente.is_(True)))).scalars().all()
    return fd


async def _dinero(db: AsyncSession, p: GestionProyecto, fd: dict, ahora: datetime) -> dict:
    """CHANGELOG 2026-10-23: la franja del dinero del caso (gestion_flujo.calcular_dinero) con lo que trajo _fuentes_dinero y el pilar SGS guardado."""
    sgs = p.sgs or {}
    lotes = sgs.get("lotes") or []
    ts_l = [_local(_iso(l.get("ts"))) for l in lotes if _iso(l.get("ts"))]
    sgs_in = {"mo": sgs.get("total_mo"), "mo_lcl": sgs.get("mo_lcl"), "n_lotes": sgs.get("n_lotes") or 0,
              "ultima_liq": max(ts_l) if ts_l else None, "liq_desde_sonny": max((t for t in ts_l if t >= gf.INICIO_SONNY), default=None),
              "codes": (sgs.get("padre") or {}).get("codes") or [], "mat": sgs.get("mat"), "mat_lcl": sgs.get("mat_lcl")} if sgs else None
    hist = fd["hist"]
    planillas = []
    for x in fd["planillas"]:
        h = hist.get(x.id_planilla) or {}
        planillas.append(dict(id_planilla=x.id_planilla, estado=x.estado or h.get("ultimo"), mo_real=x.mo_real, mat_real=x.mat_real, mo_plan=x.mo_plan, mat_plan=x.mat_plan,
                              fecha_estado=_local(h.get("fecha")) or _local(x.fecha_ingeniero or x.fecha_supervisor), fecha_aprobado=_local(h.get("aprobado")),
                              lcl=x.lcl, ticket=x.ticket, libretos=[l.libreto for l in fd["libretos"] if l.planilla == x.id_planilla]))
    pms = [dict(request_code=x.request_code, status=x.status, services=x.services, materials=x.materials,
                fecha_estado=_local(x.fecha_estado) or _local(x.request_date), compartida=[c for c in (x.lcls or []) + (x.dfs or []) if c != p.codigo])
           for x in fd["pms"]]
    libretos = [dict(lcl=l.lcl, libreto=l.libreto, estado_libreto=l.estado_libreto, mo_pesos=l.mo_pesos, puntos=l.puntos, fecha_autocert=_local(l.fecha_autocert),
                     planilla=l.planilla, regla=l.regla_enlace, pdls=l.pdls or []) for l in fd["libretos"]]
    hs = await _hermanos(db, p)
    recibe_por_lcl = not (p.tipo == "linea_viva" and any(h.tipo != "linea_viva" for h in hs))     # la misma regla de _enlazar_fuentes
    meta = await _sap_meta(db)
    precios, _ = await _precios_enel(db)
    sap_in = gf.resumir_sap(fd.get("sap") or [], meta.get("corte"), meta.get("catalogo")) if meta.get("hay") else None
    d = gf.calcular_dinero(pms, sgs_in, planillas, libretos, fd["mats"], fd["deuda"], [h.clave for h in hs], ahora, fd["cargas"], alertas_lcl=recibe_por_lcl,
                           mat_minimo=_mat_minimo(), sap=sap_in, precios=precios, exceso_pct=_mat_exceso_pct(), dias_sin_liquidar=_mat_dias_sin_liquidar())
    d["lcls"] = fd["lcls"]
    # CHANGELOG 2026-10-25 Sonny por etapa: nodos sin planilla (Lista Cruzada), planillas en curso con sus días (Historial de Actividades), en Enel y libretos
    sv = await _sonny_globales(db)
    ult = (await _ultimas_cargas(db)).get(("sonny", "nodos")) or {}
    et = gf.calcular_etapas(
        [dict(id_planilla=x.id_planilla, estado=x.estado or (hist.get(x.id_planilla) or {}).get("ultimo"), mo_real=x.mo_real, mat_real=x.mat_real, mo_plan=x.mo_plan,
              mat_plan=x.mat_plan, fecha_planilla=_local(x.fecha_planilla), qr=x.qr, nodo=x.nodo) for x in fd["planillas"]],
        [_nodo_dict(n) for n in fd["nodos"]], {k: _act_dict(a) for k, a in fd["act"].items()},
        [dict(request_code=x.request_code, status=x.status, services=x.services, materials=x.materials, fecha_estado=_local(x.fecha_estado) or _local(x.request_date),
              macroproceso=x.macroproceso) for x in fd["pms"]],
        ahora, alertas_lcl=recibe_por_lcl, qr_con_planilla=sv["qrs"], macros=sv["macros"],
        sin_cobro=any(a.get("tipo") == "liquidado_sin_cobro" for a in d.get("alertas") or []), mat_minimo=_mat_minimo(), **_sonny_dias(),
        corte_nodos=ult.get("exporte"), nodos_cargados=sv["nodos"] > 0)
    d["etapas"] = {k: v for k, v in et.items() if k != "alertas"}
    d["alertas"] = (d.get("alertas") or []) + et["alertas"]
    if d.get("resumen") is not None:
        d["resumen"]["etapas"] = et["resumen"]
    # CHANGELOG 2026-10-25 el cierre en SCM: estado de las LCL del caso (exporte de LCL), cerrado = todas CONCLUSA, y 3 alertas
    sl = gf.calcular_scm_lcl([_scm_lcl_dict(x) for x in fd.get("scm_lcl") or []], fd["lcls"],
                             {k: [_lib_dict(l) for l in v] for k, v in (fd.get("libs_lcl") or {}).items()}, [_lib_dict(l) for l in fd["libretos"]], ahora,
                             corte=fd["cargas"].get("scm_lcl"), codigo=p.codigo or "", alertas_lcl=recibe_por_lcl, mat_minimo=_mat_minimo(), **_scm_cfg())
    d["scm_lcl"] = {k: v for k, v in sl.items() if k != "alertas"}
    d["alertas"] = d["alertas"] + sl["alertas"]
    cj = {c["k"]: c for c in d.get("cajas") or []}
    if cj.get("scm") is not None:
        cj["scm"]["lcl"] = {"estado": sl["estado"], "cerrado": sl["cerrado"], "cerrado_en": sl["cerrado_en"], "corte": sl["corte"], "n": len(sl["lcls"]),
                            "abiertas": sl["abiertas"][:12], "sin_scm": sl["sin_scm"][:12]}
        cj["scm"]["por_estado"] = gf.por_estado_libretos([_lib_dict(l) for l in fd["libretos"]])      # con el nombre en pantalla de cada estado
    if cj.get("sap") is not None:
        cj["sap"]["scm_lcl"] = sl["estado"]
        cj["sap"]["nota_scm"] = gf.NOTA_CONCLUSA + ". " + gf.nota_cobro_sap() + "."
    if d.get("resumen") is not None:
        d["resumen"]["scm_lcl"] = {"estado": sl["estado"], "cerrado_en": sl["cerrado_en"], "abiertas": [x["estado"] for x in sl["abiertas"]][:6]}
    # CHANGELOG 2026-10-25 el cobro en DRAPE: las facturas de las LCL del caso con su estado, el importe por posiciones, las vencidas y 4 alertas de dinero; la caja
    # SAP pasa a «Facturado (DRAPE)» (el material del MB51 sigue en ella)
    if not gf.MOSTRAR_DRAPE:                     # CHANGELOG 2026-10-25: sin los datos del portal; la facturación sale del informe de deuda de facturación
        gf.aplicar_deuda(d)
        return d
    co = gf.calcular_cobro([_drape_dict(x) for x in fd.get("drape") or []], [_pos_dict(x) for x in fd.get("drape_pos") or []],
                           [_em_dict(x) for x in fd.get("drape_em") or []], fd["lcls"], ahora, hay=bool(fd["cargas"].get("drape")),
                           corte=fd["cargas"].get("drape"), codigo=p.codigo or "", alertas_lcl=recibe_por_lcl, scm=sl,
                           libs_lcl={k: [_lib_dict(l) for l in v] for k, v in (fd.get("libs_lcl") or {}).items()}, mat_minimo=_mat_minimo(), **_drape_cfg())
    gf.aplicar_cobro(d, co)
    return d


def _drape_dict(x) -> dict:
    """CHANGELOG 2026-10-25: una factura de sgo7_gg_drape_factura para gestion_flujo.calcular_cobro (fechas en hora Colombia)."""
    return dict(drape_id=x.drape_id, factura=x.factura, estado=x.estado, destinatario=x.destinatario, f_factura=_local(x.f_factura), n_registro=x.n_registro,
                f_registro=_local(x.f_registro), f_pago=_local(x.f_pago), doc_pago=x.doc_pago, f_envio=_local(x.f_envio), f_vence=_local(x.f_vence),
                condiciones=x.condiciones, total=x.total, pedidos=x.pedidos or [], lcls=x.lcls or [], ems=x.ems or [], pos_lcls=x.pos_lcls or [],
                carga_id=x.carga_id, carga_det_id=x.carga_det_id)


def _pos_dict(x) -> dict:
    return dict(drape_id=x.drape_id, factura=x.factura, posicion=x.posicion, lcl=x.lcl, pedido_pos=x.pedido_pos, producto=x.producto, descripcion=x.descripcion,
                unidad=x.unidad, cantidad=x.cantidad, precio=x.precio, impuesto=x.impuesto, importe=x.importe, carga_id=x.carga_id)


def _em_dict(x) -> dict:
    return dict(em=x.em, fecha=_local(x.fecha), importe=x.importe, lcl=x.lcl, carga_id=x.carga_id)


def _drape_cfg() -> dict:
    """CHANGELOG 2026-10-25: umbrales del cobro en DRAPE (config.GESTION_DRAPE_*; recarga graciosa)."""
    out = {}
    for k, nombre, defecto in (("dias_sin_factura", "GESTION_DRAPE_DIAS_SIN_FACTURA", gf.DRAPE_DIAS_SIN_FACTURA), ("dias_mora", "GESTION_DRAPE_DIAS_MORA", gf.DRAPE_DIAS_MORA),
                               ("vence_dias", "GESTION_DRAPE_VENCE_DIAS", gf.DRAPE_VENCE_DIAS), ("dias_em", "GESTION_DRAPE_DIAS_EM", gf.DRAPE_DIAS_EM)):
        try:
            out[k] = int(getattr(_config, nombre, defecto))
        except (TypeError, ValueError):
            out[k] = int(defecto)
    return out


def _lib_dict(l) -> dict:
    """CHANGELOG 2026-10-25: un libreto de sgo7_gg_scm_libreto para gestion_flujo.calcular_scm_lcl (fechas en hora Colombia)."""
    return dict(lcl=l.lcl, libreto=l.libreto, estado_libreto=l.estado_libreto, mo_pesos=l.mo_pesos, planilla=l.planilla, pdls=l.pdls or [],
                fecha_autocert=_local(l.fecha_autocert))


def _scm_lcl_dict(x) -> dict:
    """CHANGELOG 2026-10-25: una fila de sgo7_gg_scm_lcl (fechas en hora Colombia, sin zona)."""
    return dict(lcl=x.lcl, trabajo=x.trabajo, estado=x.estado, estado_desde=_local(x.estado_desde), fecha_envio=_local(x.fecha_envio),
                inicio_validez=_local(x.inicio_validez), fin_validez=_local(x.fin_validez), importe=x.importe, obra=x.obra, acumulativa=x.acumulativa,
                tam=x.tam, carga_id=x.carga_id)


def _scm_cfg() -> dict:
    """CHANGELOG 2026-10-25: días al fin de validez y ± % del Importe contra la MO (config.GESTION_SCM_DIAS_VENCE, GESTION_SCM_IMPORTE_PCT; recarga graciosa)."""
    out = {}
    for k, nombre, defecto, tipo in (("dias_vence", "GESTION_SCM_DIAS_VENCE", gf.SCM_DIAS_VENCE, int), ("importe_pct", "GESTION_SCM_IMPORTE_PCT", gf.SCM_IMPORTE_PCT, float)):
        try:
            out[k] = tipo(getattr(_config, nombre, defecto))
        except (TypeError, ValueError):
            out[k] = tipo(defecto)
    return out


def _nodo_dict(n) -> dict:
    """CHANGELOG 2026-10-25: un nodo de sgo7_gg_sonny_nodo para gestion_flujo (fechas en hora Colombia, sin zona)."""
    return dict(id_nodo=n.id_nodo, codigo=n.codigo, nodo=n.nodo, estado_nodo=n.estado_nodo, id_planillas=n.id_planillas or [], estado_planilla=n.estado_planilla,
                qr=n.qr, f_estado=_local(n.f_estado), f_creacion=_local(n.f_creacion), estado_material=n.estado_material, movimiento=n.movimiento,
                f_almacen=_local(n.f_almacen), f_confirmado=_local(n.f_confirmado), carga_id=n.carga_id)


def _act_dict(a) -> dict:
    return dict(f_creada=_local(a.f_creada), estado_act=a.estado_act, f_estado_act=_local(a.f_estado_act), f_firma=_local(a.f_firma), n_firmas=a.n_firmas,
                f_ultima=_local(a.f_ultima), que_ultima=a.que_ultima, carga_id=a.carga_id)


def _sonny_dias() -> dict:
    """CHANGELOG 2026-10-25: los días de las 4 alertas de Sonny por etapa (config.GESTION_SONNY_DIAS_*; se cambian con una recarga graciosa)."""
    out = {}
    for k, nombre, defecto in (("dias_nodo", "GESTION_SONNY_DIAS_NODO", gf.SONNY_DIAS_NODO), ("dias_quieta", "GESTION_SONNY_DIAS_QUIETA", gf.SONNY_DIAS_QUIETA),
                               ("dias_firmas", "GESTION_SONNY_DIAS_FIRMAS", gf.SONNY_DIAS_FIRMAS), ("dias_pms", "GESTION_SONNY_DIAS_PMS5", gf.SONNY_DIAS_PMS5)):
        try:
            out[k] = int(getattr(_config, nombre, defecto))
        except (TypeError, ValueError):
            out[k] = int(defecto)
    return out


async def _sonny_globales(db: AsyncSession) -> dict:
    """CHANGELOG 2026-10-25: los QR (request_code) con alguna planilla en toda la Lista de Sonny, los macroprocesos de PMS que usan Sonny (≥ 20 % de sus
    solicitudes con planilla) y cuántos nodos hay cargados. Caché por proceso validada con un sello barato de las tablas (hay 2 workers)."""
    (n_p, mx_p, ac_p), = (await db.execute(select(func.count(), func.max(GestionSonnyPlanilla.id), func.max(GestionSonnyPlanilla.actualizado_en)))).all()
    (n_s, mx_s, ac_s), = (await db.execute(select(func.count(), func.max(GestionPms.id), func.max(GestionPms.actualizado_en)))).all()
    (n_n,), = (await db.execute(select(func.count()).select_from(GestionSonnyNodo).where(GestionSonnyNodo.vigente.is_(True)))).all()
    sello = (n_p, mx_p, ac_p, n_s, mx_s, ac_s, n_n)
    v = _cache_sonny.get("g")
    if v and v[0] == sello:
        return v[1]
    qrs = {q for (x,) in (await db.execute(select(GestionSonnyPlanilla.qr).where(GestionSonnyPlanilla.qr.isnot(None)))).all() for q in gf.qrs_de(x)}
    por_macro: dict = {}
    for rc, mp in (await db.execute(select(GestionPms.request_code, GestionPms.macroproceso))).all():
        x = por_macro.setdefault(mp or "", [0, 0])
        x[0] += 1
        x[1] += 1 if rc in qrs else 0
    macros = {m for m, (n, c) in por_macro.items() if m and n and 100 * c / n >= gf.MACRO_SONNY_PCT}
    out = {"qrs": qrs, "macros": macros, "nodos": int(n_n or 0), "por_macro": {m: {"n": n, "con_planilla": c} for m, (n, c) in por_macro.items()}}
    _cache_sonny["g"] = (sello, out)
    return out


def _pago_confirmado() -> bool:
    """CHANGELOG 2026-10-25: config.GESTION_PAGO_CONFIRMADO. Mientras sea False, «pagado» se dice «sin pago registrado en DRAPE · por confirmar» (gris) y no cuenta
    como cobro ni como deuda. Se pone True cuando facturación confirme que DRAPE marca los pagos de este contrato (+ recarga graciosa y recálculo)."""
    return bool(getattr(_config, "GESTION_PAGO_CONFIRMADO", gf.PAGO_CONFIRMADO))


def _caso_ejemplo() -> str:
    """CHANGELOG 2026-10-25: config.GESTION_CASO_EJEMPLO, el trabajo cuyos números reales muestra «¿Cómo fluye?» (hoy 6301018885)."""
    return str(getattr(_config, "GESTION_CASO_EJEMPLO", "6301018885") or "")


def _mat_minimo() -> float:
    """CHANGELOG 2026-10-23: diferencia de material (en pesos, a costo SGS) desde la que un caso tiene alerta; se cambia en config.py
    (GESTION_MAT_MINIMO) con una recarga graciosa, sin tocar código."""
    try:
        return float(getattr(_config, "GESTION_MAT_MINIMO", gf.MAT_MINIMO))
    except (TypeError, ValueError):
        return float(gf.MAT_MINIMO)


def _mat_exceso_pct() -> float:
    """CHANGELOG 2026-10-25: ± % en que el material proyectado de SGS «cuadra» con PMS (config.GESTION_MAT_EXCESO_PCT; se cambia con una recarga graciosa)."""
    try:
        return float(getattr(_config, "GESTION_MAT_EXCESO_PCT", gf.MAT_EXCESO_PCT))
    except (TypeError, ValueError):
        return float(gf.MAT_EXCESO_PCT)


def _mat_dias_sin_liquidar() -> int:
    """CHANGELOG 2026-10-25: días hábiles desde la última bandeja de salida desde los que el material por liquidar es alerta (config.GESTION_MAT_DIAS_SIN_LIQUIDAR)."""
    try:
        return int(getattr(_config, "GESTION_MAT_DIAS_SIN_LIQUIDAR", gf.MAT_DIAS_SIN_LIQUIDAR))
    except (TypeError, ValueError):
        return int(gf.MAT_DIAS_SIN_LIQUIDAR)


# ── Evidencia: de dónde sale cada cifra ───────────────────────────────────────────────────────────────────────────
def _precios_qp() -> tuple[dict, dict]:
    """La lista de precios QP de producción (la del Balance Integrado) y su huella; caché por fecha del archivo."""
    try:
        st = os.stat(_PRECIOS_QP)
    except OSError:
        return {}, {}
    v = _cache_precios.get("p")
    if v and v[0] == st.st_mtime:
        return v[1], v[2]
    try:
        with open(_PRECIOS_QP, "rb") as fh:
            crudo = fh.read()
        data = {str(k).upper(): v for k, v in json.loads(crudo.decode("utf-8")).items()}
    except (OSError, ValueError):
        return {}, {}
    info = {"archivo": os.path.basename(_PRECIOS_QP), "sha256": hashlib.sha256(crudo).hexdigest(), "codigos": len(data),
            "modificado": datetime.fromtimestamp(st.st_mtime, _TZ).replace(tzinfo=None).isoformat(timespec="seconds")}
    _cache_precios["p"] = (st.st_mtime, data, info)
    return data, info


def _carga_out(c) -> Optional[dict]:
    if c is None:
        return None
    hecho = c.cerrada_en or c.abierta_en
    return {"id": c.id, "sistema": c.sistema, "parte": c.parte, "archivo": c.archivo, "tamano": c.tamano, "sha256": c.sha256,
            "exporte": _local(c.modificado).isoformat() if c.modificado else None, "corte": _local(c.corte).isoformat() if c.corte else None,
            "cargado": _local(hecho).isoformat() if hecho else None, "usuario": c.usuario, "equipo": c.equipo, "filas": c.filas_archivo,
            "enviadas": c.filas_enviadas, "duplicados": c.duplicados, "fuera_rango": c.fuera_rango, "cerrada": bool(c.cerrada_en), "detalle": c.detalle or {},
            "resultado": c.resultado or {},
            "tiene_archivo": bool(getattr(c, "archivo_ruta", None)), "archivo_nombre": getattr(c, "archivo_nombre", None),         # CHANGELOG 2026-10-25
            "archivo_tamano": getattr(c, "archivo_tamano", None)}


async def _ultimas_cargas(db: AsyncSession) -> dict:
    """(sistema, parte) → la carga más reciente. Caché de 60 s (cada carga la limpia)."""
    v = _cache_cargas.get("ult")
    if v and time.time() - v[0] < 60:
        return v[1]
    out = {}
    for c in (await db.execute(select(GestionCarga).order_by(GestionCarga.abierta_en, GestionCarga.id))).scalars().all():
        if (c.archivo or "").startswith("prueba-"):        # las cargas de las pruebas automáticas nunca son «el último exporte»
            continue
        out[(c.sistema, c.parte)] = _carga_out(c)
    _cache_cargas["ult"] = (time.time(), out)
    return out


def _vieja(carga_id, ult: Optional[dict]) -> bool:
    """La fila no vino en la última carga de su sistema (se conserva de un exporte anterior)."""
    return bool(ult) and carga_id != ult["id"]


async def _fuentes_evidencia(db: AsyncSession, p: GestionProyecto, fd: dict, din: dict) -> dict:
    """Lo que necesita gestion_flujo.armar_evidencia: las filas reales de cada caja (con su carga), los libretos del alcance del caso
    (en mayor impacto también los que una planilla del ticket podría reclamar: pendientes y descartados), el detalle de prestaciones y
    de material por libreto, las decisiones humanas y la fuente de cada sistema."""
    cod = p.codigo or ""
    es_df = cod.startswith("DF")
    ult = await _ultimas_cargas(db)
    pls = fd["planillas"]
    ids = {x.id_planilla for x in pls}
    lcls = fd["lcls"]
    libs_all = (await db.execute(select(GestionScmLibreto).where(GestionScmLibreto.lcl.in_(lcls)).order_by(GestionScmLibreto.lcl, GestionScmLibreto.libreto))).scalars().all() if lcls else []
    decs = (await db.execute(select(GestionDecision).where(GestionDecision.lcl.in_(lcls)).order_by(GestionDecision.actualizado_en))).scalars().all() if lcls else []
    en_caso = {(l.lcl, l.libreto) for l in fd["libretos"]}
    miembros = {(l.lcl, l.libreto) for l in fd["libretos"] if l.estado_libreto not in gf.LIBRETO_MUERTO}
    neg_caso = {(d.lcl, d.libreto) for d in decs if d.decision == "no" and d.id_planilla in ids}
    alcance = [l for l in libs_all if not es_df or (l.lcl, l.libreto) in en_caso or (l.lcl, l.libreto) in neg_caso or (set(l.candidatas or []) & ids)]
    claves = [(l.lcl, l.libreto) for l in alcance if (l.lcl, l.libreto) in miembros]
    precios, info_precios = _precios_qp()
    prest: dict = {}
    mats: dict = {}
    prest_obj: dict = {}
    mats_obj: dict = {}
    if claves:
        for r in (await db.execute(select(GestionScmPrest).where(tuple_(GestionScmPrest.lcl, GestionScmPrest.libreto).in_(claves))
                                   .order_by(GestionScmPrest.lcl, GestionScmPrest.libreto, GestionScmPrest.n))).scalars().all():
            pl = precios.get((r.prestacion or "").upper()) or {}
            prest_obj.setdefault(f"{r.lcl}|{r.libreto}", []).append(r)
            prest.setdefault(f"{r.lcl}|{r.libreto}", []).append(dict(n=r.n, prestacion=r.prestacion, descripcion=(pl.get("descripcion") or "")[:160],
                                                                     unidad=r.unidad or pl.get("unidad") or "", cantidad=r.cantidad, precio=r.precio,
                                                                     pesos=r.pesos, precio_lista=pl.get("precio")))
        for m in (await db.execute(select(GestionScmMat).where(tuple_(GestionScmMat.lcl, GestionScmMat.libreto).in_(claves))
                                   .order_by(GestionScmMat.lcl, GestionScmMat.libreto, GestionScmMat.material))).scalars().all():
            mats_obj.setdefault(f"{m.lcl}|{m.libreto}", []).append(m)
            mats.setdefault(f"{m.lcl}|{m.libreto}", []).append(dict(material=m.material, posada=m.posada, recuperada=m.recuperada, unidad=m.unidad or ""))
    otras = ({c for l in alcance for c in (l.candidatas or [])} | {d.id_planilla for d in decs}) - ids
    ext = {}
    if otras:
        for x in (await db.execute(select(GestionSonnyPlanilla).where(GestionSonnyPlanilla.id_planilla.in_(sorted(otras))))).scalars().all():
            ext[x.id_planilla] = dict(lcl=x.lcl, ticket=x.ticket, mo_real=x.mo_real, estado=x.estado)
    qrs = sorted({x.qr for x in pls if x.qr})
    qr_pms = set((await db.execute(select(GestionPms.request_code).where(GestionPms.request_code.in_(qrs)))).scalars().all()) if qrs else set()
    hist = fd["hist"]
    u_pms, u_lis, u_lib, u_pre, u_mat = (ult.get(k) for k in (("pms", "solicitudes"), ("sonny", "lista"), ("scm", "libretos"), ("scm", "prestaciones"),
                                                               ("scm", "materiales")))
    car = fd["cargas"]
    sgs = p.sgs or {}
    # La fuente de cada tarjeta son las cargas de las que vinieron SUS filas (si alguna fila viene de un exporte anterior, se ve)
    ids_pms = [x.carga_id for x in fd["pms"]]
    ids_lis = [x.carga_id for x in pls]
    miem = [l for l in alcance if (l.lcl, l.libreto) in miembros]
    ids_lib = [l.carga_id for l in miem]
    ids_pre = [r.carga_id for rs in prest_obj.values() for r in rs]
    ids_mat = [m.carga_id for ms in mats_obj.values() for m in ms]
    sap_f = fd.get("sap") or []
    ids_sap = [r.get("carga_id") for r in sap_f]                     # CHANGELOG 2026-10-23
    u_sap = ult.get(("sap", "mb51"))
    ids_nod = [n.carga_id for n in fd.get("nodos") or []]              # CHANGELOG 2026-10-25
    ids_lcl = [x.carga_id for x in fd.get("scm_lcl") or []]            # CHANGELOG 2026-10-25
    u_lcl = ult.get(("scm", "lcl"))
    ids_dr = [x.carga_id for x in fd.get("drape") or []] + [x.carga_det_id for x in fd.get("drape") or [] if x.carga_det_id]     # CHANGELOG 2026-10-25: listas y detalle
    u_dr = ult.get(("drape", "listas"))
    u_deu = ult.get(("deuda", "informe"))                              # CHANGELOG 2026-10-25: la fuente de «En facturación (informe de deuda)»
    ids_act = [a.carga_id for a in (fd.get("act") or {}).values()]
    u_nod = ult.get(("sonny", "nodos"))
    u_his = ult.get(("sonny", "historial"))                            # CHANGELOG 2026-10-25: el motivo del rechazo sale del Historial de Planillas
    usados = ({i for i in ids_pms + ids_lis + ids_lib + ids_pre + ids_mat + ids_sap + ids_nod + ids_act + ids_lcl + ids_dr if i} | ({u_sap["id"]} if u_sap else set())
              | ({u_nod["id"]} if u_nod else set()) | ({u_lcl["id"]} if u_lcl else set()) | ({u_dr["id"]} if u_dr else set())
              | ({u_his["id"]} if u_his else set()) | ({u_deu["id"]} if u_deu else set()))
    por_id = {c.id: _carga_out(c) for c in (await db.execute(select(GestionCarga).where(GestionCarga.id.in_(sorted(usados))))).scalars().all()} if usados else {}
    ultimas = {c["id"] for c in ult.values() if c}

    def fuente(sistema, *listas, sin_registro=None, **kw):
        cs, sin = [], 0
        for ids in listas:
            for i in dict.fromkeys(x for x in ids if x):
                if i in por_id:
                    cs.append(dict(por_id[i], ultima=i in ultimas))
            sin += sum(1 for x in ids if not x)
        nota = (sin_registro if not cs else f"{sin} fila(s) cargadas antes del registro de cargas") if sin or not cs else None
        return {"sistema": sistema, "cargas": cs, "sin_registro": nota, **kw}
    fuentes = {
        "pms": fuente("PMS", ids_pms, sin_registro=f"cargado el {car['pms_cargado']:%d/%m %H:%M} (antes del registro de cargas)" if car.get("pms_cargado") else "sin carga"),
        "sonny_hist": fuente("Sonny · Historial de Planillas (columna «Observaciones»)", [u_his["id"]] if u_his else [],
                             sin_registro=f"cargado el {car['sonny_cargado']:%d/%m %H:%M} (antes del registro de cargas)" if car.get("sonny_cargado") else "sin carga"),   # CHANGELOG 2026-10-25
        "sonny": fuente("Sonny · Lista de Planillas", ids_lis,
                        sin_registro=f"cargado el {car['sonny_cargado']:%d/%m %H:%M} (antes del registro de cargas)" if car.get("sonny_cargado") else "sin carga"),
        "scm_mo": fuente("SCM · Informe LCL_LCA + Prestaciones", ids_lib, ids_pre, precios=info_precios,
                         sin_registro=f"exporte del {car['scm']:%d/%m %H:%M} (antes del registro de cargas)" if car.get("scm") else "sin carga"),
        "scm_mat": fuente("SCM · Informe Materiales", ids_mat,
                          sin_registro=f"exporte del {car['scm']:%d/%m %H:%M} (antes del registro de cargas)" if car.get("scm") else "sin carga"),
        "sgs": {"sistema": "SGS", "vivo": True, "consultado": _local(_iso(sgs.get("consultado_en"))).isoformat() if sgs.get("consultado_en") else None,
                "nota": "lectura en vivo por el Centro MO (solo lectura en SGS)"},
        # CHANGELOG 2026-10-23: el MB51; si el caso no tiene filas, la fuente es la última carga (consta que el archivo está cargado y hasta dónde llega)
        "sap": fuente("SAP · MB51 (material de Enel)", ids_sap or ([u_sap["id"]] if u_sap else []), sin_registro="sin carga de SAP",
                      corte=(car.get("sap").date().isoformat() if car.get("sap") else None)),
        # CHANGELOG 2026-10-25: la fila «Sonny por etapa» sale de la Lista de Planillas, la Lista Cruzada de Nodos (si el código no tiene nodos, la última carga:
        # consta que está cargada) y el Historial de Actividades
        "sonny_etapa": fuente("Sonny · Lista de Planillas + Lista Cruzada de Nodos + Historial de Actividades", ids_lis, ids_nod or ([u_nod["id"]] if u_nod else []), ids_act,
                              sin_registro="sin carga de la Lista Cruzada de Nodos ni del Historial de Actividades"),
        # CHANGELOG 2026-10-25: el exporte de LCL de SCM; si el caso no tiene filas, la última carga (consta que está cargado y hasta cuándo llega)
        "scm_lcl": fuente("SCM · Informe LCL/LCA (exporte de LCL: estado, validez, Importe y Trabajo)", ids_lcl or ([u_lcl["id"]] if u_lcl else []),
                          sin_registro="sin carga del exporte de LCL de SCM (lcl_export)"),
        # CHANGELOG 2026-10-25: DRAPE (listas de facturas y detalle de las abiertas); si el caso no tiene facturas, la última carga (consta que está cargado)
        "drape": fuente("DRAPE · portal de facturación de Enel (listas de facturas y detalle)", ids_dr or ([u_dr["id"]] if u_dr else []),
                        sin_registro="sin carga de DRAPE",
                        leido=(car.get("drape").isoformat(timespec="minutes") if car.get("drape") else None)),
        "deuda": fuente(gf.rotulo_deuda(_local(car["deuda"]) if car.get("deuda") else None), [u_deu["id"]] if u_deu else [],
                        sin_registro=f"corte del {_local(car['deuda']):%d/%m}, antes del registro de cargas" if car.get("deuda") else "sin carga"),     # CHANGELOG 2026-10-25
    }
    if not gf.MOSTRAR_DRAPE:                                            # CHANGELOG 2026-10-25
        fuentes.pop("drape", None)
    precios_mat, precios_info = await _precios_enel(db)
    for k in ("sgs", "scm_mat"):
        fuentes[k]["precios_enel"] = precios_info
    hs = await _hermanos(db, p)
    return dict(
        caso=dict(id=p.id, clave=p.clave, codigo=cod, es_df=es_df, codes=(p.padre or {}).get("codes") or [p.clave], compartido_con=[h.clave for h in hs]),
        dinero=din, sgs=sgs, fuentes=fuentes, precios=info_precios, alertas=din.get("alertas") or [], qr_en_pms=sorted(qr_pms), planillas_ext=ext,
        pms=[dict(request_code=x.request_code, subprojects=x.subprojects, lcls=x.lcls or [], dfs=x.dfs or [], status=x.status, services=x.services,
                  materials=x.materials, fecha_estado=_local(x.fecha_estado) or _local(x.request_date), anterior=_vieja(x.carga_id, u_pms)) for x in fd["pms"]],
        planillas=[dict(id_planilla=x.id_planilla, lcl=x.lcl, ticket=x.ticket, pdl=x.pdl, qr=x.qr, nodo=x.nodo, estado=x.estado or (hist.get(x.id_planilla) or {}).get("ultimo"),
                        mo_real=x.mo_real, mat_real=x.mat_real, fecha_aprobado=_local((hist.get(x.id_planilla) or {}).get("aprobado")), anterior=_vieja(x.carga_id, u_lis),
                        mo_plan=x.mo_plan, mat_plan=x.mat_plan, fecha_planilla=_local(x.fecha_planilla),                       # CHANGELOG 2026-10-25: el plan y el último cambio
                        hitos=(hist.get(x.id_planilla) or {}).get("hitos") or {})                                             # CHANGELOG 2026-10-25: primera fecha de cada estado
                   for x in pls],
        nodos=[_nodo_dict(n) for n in fd.get("nodos") or []], act={k: _act_dict(a) for k, a in (fd.get("act") or {}).items()},     # CHANGELOG 2026-10-25
        libretos=[dict(lcl=l.lcl, libreto=l.libreto, descripcion=l.descripcion, estado_libreto=l.estado_libreto, mo_pesos=l.mo_pesos, puntos=l.puntos,
                       n_prest=l.n_prest, sin_precio=l.sin_precio, planilla=l.planilla, regla=l.regla_enlace, candidatas=l.candidatas or [],
                       fecha_autocert=_local(l.fecha_autocert), miembro=(l.lcl, l.libreto) in miembros, anterior=_vieja(l.carga_id, u_lib),
                       en_caso=(l.lcl, l.libreto) in en_caso, pdls=l.pdls or []) for l in alcance],                                       # CHANGELOG 2026-10-25
        scm_lcl=[dict(_scm_lcl_dict(x), anterior=_vieja(x.carga_id, u_lcl)) for x in fd.get("scm_lcl") or []],                          # CHANGELOG 2026-10-25
        lcl_libs={k: [_lib_dict(l) for l in v] for k, v in (fd.get("libs_lcl") or {}).items()},
        prest=prest, mats=mats,
        sap_filas=[dict(r, anterior=_vieja(r.get("carga_id"), u_sap)) for r in sap_f], precios_mat=precios_mat, precios_enel=precios_info,
        drape_pos=[_pos_dict(x) for x in fd.get("drape_pos") or []],                                                                  # CHANGELOG 2026-10-25
        decisiones=[dict(id=d.id, lcl=d.lcl, libreto=d.libreto, id_planilla=d.id_planilla, decision=d.decision, usuario=d.usuario,
                         fecha=_local(d.actualizado_en or d.creado_en).isoformat() if (d.actualizado_en or d.creado_en) else None, nota=d.nota) for d in decs],
    )


async def _evidencia_resumen(db: AsyncSession, p: GestionProyecto, fd: dict, din: dict, ahora: datetime) -> Optional[dict]:
    """Guarda en la franja si cada suma reproduce su cifra (lo cuentan el panel de controles y la lista). Nunca tumba el recálculo.
    CHANGELOG 2026-10-25: devuelve lo que leyó (E de _fuentes_evidencia) para armar la cadena con lo mismo (None si falló)."""
    E = None
    try:
        E = await _fuentes_evidencia(db, p, fd, din)
        ev = gf.armar_evidencia(E)
        din["evidencia"] = {"ok": ev["ok"], "malas": ev["malas"], "cajas": ev["resumen"], "calculado_en": ahora.isoformat(timespec="seconds")}
    except Exception as e:                              # queda dicho en la franja y en el panel (control rojo); no se esconde
        din["evidencia"] = {"ok": False, "malas": ["error"], "error": f"{type(e).__name__}: {str(e)[:160]}", "calculado_en": ahora.isoformat(timespec="seconds")}
    if din.get("resumen") is not None:
        din["resumen"]["evidencia_ok"] = din["evidencia"]["ok"]
        din["resumen"]["evidencia_malas"] = din["evidencia"]["malas"]
    return E


async def _enlazar_fuentes(db: AsyncSession, p: GestionProyecto, regs_cm: list[dict], fd: Optional[dict] = None) -> None:
    """Rehace los eventos derivados (sgs, sonny, sgo7, pms, scm) del proyecto. Los de correo no se tocan."""
    if fd is None:
        fd = await _fuentes_dinero(db, p)
    await db.execute(delete(GestionEvento).where(GestionEvento.proyecto_id == p.id, GestionEvento.fuente.in_(("sgs", "sonny", "sgo7", "pms", "scm"))))
    clave = p.codigo or p.clave                      # LCL o ticket (lo que nombran los correos, los casos y los archivos)
    codes_padre = set((p.padre or {}).get("codes") or [(p.padre or {}).get("code") or p.clave])
    hermanos = await _hermanos(db, p)
    # Lo que viene por LCL sin decir el padre (Centro MO, casos del Balance, archivos del Liquidador) es del gemelo que NO es
    # línea viva; si este es LV y hay un gemelo que no lo es, no se le cuelga nada de eso.
    recibe_por_lcl = not (p.tipo == "linea_viva" and any(h.tipo != "linea_viva" for h in hermanos))
    n = 0

    def add(**kw):
        nonlocal n
        n += 1
        db.add(GestionEvento(proyecto_id=p.id, directo=False, huella=_sha(f"{p.id}|{kw.get('fuente')}|{n}|{kw.get('ts')}|{kw.get('prueba')}"), **kw))

    # 1. SGS en vivo (vía Centro MO): una liquidación = los avances con MO creados el mismo día en un padre
    for l in (p.sgs or {}).get("lotes") or []:
        ts = _iso(l.get("ts"))
        if not ts:
            continue
        usuarios = ", ".join(l.get("usuarios") or [])
        add(ts=ts, fuente="sgs", paso=21, cond="liquidado en SGS", actor="AUXSGS" if any(f in usuarios.upper() for f in gf.AUXSGS_USUARIOS_SGS) else "SGS",
            prueba=f"Liquidación en SGS: ${float(l.get('mo') or 0):,.0f} en {l.get('n')} avances ({l.get('padre')})",
            de=usuarios, asunto=f"SGS · {l.get('padre')} · {l.get('dia')}", texto="", ref={"padre": l.get("padre"), "dia": l.get("dia")})

    # 2. Correos de MO que el Centro MO ya interpretó (código validado = este LCL/ticket)
    descartados = set()
    for r in regs_cm:
        if gf.clave_de(r.get("codigo") or "") != clave or not r.get("codigo_validado") or not recibe_por_lcl:
            continue
        est = r.get("estado") or ""
        if est == "descartado" and r.get("entry_id"):
            descartados.add(r["entry_id"])
        ts = _iso(r.get("fecha")) or _iso(r.get("detectado_en"))
        if ts:
            add(ts=ts, fuente="sgo7", paso=21, cond="recibido por Centro MO", actor="SGO7",
                prueba=f"Centro MO interpretó el correo de MO: código {r.get('codigo')} (por {r.get('codigo_fuente') or '?'}), "
                       f"{r.get('n_filas_mo') or 0} prestaciones, ejecución {r.get('fecha_ejecucion') or '?'}",
                de=r.get("remitente") or "", asunto=(r.get("asunto") or "")[:300], texto=", ".join(r.get("lideres") or []),
                ref={"centro_mo": r.get("id"), "entry_id": r.get("entry_id"), "estado": est})
        ts2 = _iso(r.get("completado_en"))
        if est in _ESTADO_CM and (ts2 or est != "completado"):
            add(ts=ts2 or ts, fuente="sgo7", paso=21, cond=_ESTADO_CM[est], actor="AUXSGS",
                prueba=f"Centro MO: {_ESTADO_CM[est]}" + (f" · PDL {r['pdl']}" if r.get("pdl") else ""),
                de="Centro MO", asunto=(r.get("asunto") or "")[:300], texto="", ref={"centro_mo": r.get("id"), "estado": est})
    if descartados:
        evs = (await db.execute(select(GestionEvento).where(GestionEvento.proyecto_id == p.id, GestionEvento.entry_id.in_(list(descartados)),
                                                            GestionEvento.paso == 21))).scalars().all()
        for e in evs:
            e.cond = "descartado por Centro MO"

    # 3. Casos de la línea de facturación (Balance Integrado) — sin las reconsultas automáticas que no cambian nada
    casos = (await db.execute(select(CasoFacturacion).where(or_(CasoFacturacion.pro_code.in_(list(codes_padre)), CasoFacturacion.lcl == clave,
                                                                 CasoFacturacion.ticket.ilike(f"%{clave}%"), CasoFacturacion.proyecto_id.ilike(f"%{clave}%")))
                              .order_by(CasoFacturacion.id))).scalars().all()
    fact = []
    for c in casos:
        if c.pro_code and c.pro_code not in codes_padre and not any(c.pro_code.startswith(x) for x in codes_padre):
            continue                                  # el caso del Balance dice de qué padre es: el otro gemelo
        if not c.pro_code and not recibe_por_lcl:
            continue
        fact.append(dict(caso_id=c.id, id_planilla=c.id_planilla, estado=c.estado, semaforo=c.semaforo, orden_estado=c.orden_estado,
                         actualizado=(c.actualizado_en.isoformat() if c.actualizado_en else None), pro_code=c.pro_code))
        for e in (await db.execute(select(CasoEvento).where(CasoEvento.caso_id == c.id).order_by(CasoEvento.ts))).scalars().all():
            m = _CASO_EVENTOS.get(e.tipo)
            if not m:
                continue
            det = e.detalle or {}
            if e.tipo == "consulta" and det.get("origen") == "auto" and det.get("estado_previo") == det.get("estado"):
                continue
            paso, cond, prueba, actor, fuente = m
            if e.tipo == "consulta" and det.get("semaforo"):
                cond = f"semáforo {det['semaforo']}"
            add(ts=e.ts, fuente=fuente, paso=paso, cond=cond, actor=actor, prueba=prueba, de=e.usuario or "",
                asunto=f"Caso {c.id} · planilla {c.id_planilla}", texto="", ref={"caso_id": c.id, "tipo": e.tipo})
    p.facturacion = fact

    # 3b. Historial de Planillas de Sonny (CHANGELOG 2026-10-23): cada cambio de estado de cada planilla, con la fecha oficial de Sonny. La
    # planilla se liga por su LCL o por su ticket (una de mayor impacto trae los dos y en SGS el proyecto es el ticket). Entre
    # gemelos del mismo código va al caso del supervisor de la planilla; si no coincide con ninguno, al que no es línea viva
    # (gf.elegir_caso_sonny: lo mismo que recibe_por_lcl, con la excepción del supervisor).
    filas_sonny = (await db.execute(select(GestionSonny).where(or_(GestionSonny.codigo == clave, GestionSonny.ticket == clave))
                                    .order_by(GestionSonny.fecha, GestionSonny.id_cambio))).scalars().all() if p.codigo else []
    if filas_sonny:
        cods = sorted({c for s in filas_sonny for c in (s.codigo, s.ticket) if c})
        rivales = [dict(id=x.id, tipo=x.tipo, codigo=x.codigo, supervisor=(x.padre or {}).get("supervisor") or x.supervisor or "")
                   for x in (await db.execute(select(GestionProyecto).where(GestionProyecto.codigo.in_(cods)))).scalars().all()]
        elegido: dict = {}
        for s in filas_sonny:
            k = (s.codigo, s.ticket, s.supervisor or "")
            if k not in elegido:
                elegido[k] = gf.elegir_caso_sonny(s.supervisor or "", [c for c in rivales if c["codigo"] and c["codigo"] in (s.codigo, s.ticket)])
            if elegido[k] != p.id:
                continue
            paso, cond, actor, _ = gf.paso_sonny(s.estado)
            ref_s = {"id_planilla": s.id_planilla, "estado": s.estado, "sonny": True, "id_cambio": s.id_cambio}
            if cond == "no cumple":                                  # CHANGELOG 2026-10-25: el motivo del rechazo, clasificado (el texto de Sonny no se toca)
                ref_s["motivo"] = gf.motivo_corto(gf.motivo_rechazo(s.observacion))
            add(ts=s.fecha, fuente="sonny", paso=paso, cond=cond, actor=actor, prueba=f"Sonny: planilla {s.id_planilla} pasa a «{s.estado}»"[:300],
                de="Sonny", asunto=(f"Sonny · planilla {s.id_planilla}" + (f" · nodo {s.nodo}" if s.nodo else ""))[:500], texto=gf.texto_motivo(s.observacion),
                ref=ref_s)

    # 3c. Lista de Planillas de Sonny (CHANGELOG 2026-10-23): la aprobación del supervisor fecha los pasos 19 (planilla creada) y 25 (su firma); la
    # del ingeniero, el 26. Van al caso de la planilla con la misma regla de gemelos que el historial. No cuentan como estado de
    # la planilla para los rechazos (ref sin «sonny»).
    if fd["planillas"]:
        cods_pl = sorted({c for x in fd["planillas"] for c in (x.lcl, x.ticket) if c})
        riv = await _rivales(db, cods_pl)
        for x in fd["planillas"]:
            if gf.elegir_caso_sonny(x.supervisor or "", [c for c in riv if c["codigo"] in (x.lcl, x.ticket)]) != p.id:
                continue
            asunto = (f"Sonny · planilla {x.id_planilla}" + (f" · nodo {x.nodo}" if x.nodo else ""))[:500]
            if x.fecha_supervisor and (x.estado_supervisor or "").lower().startswith("aprob"):
                for paso, cond in ((19, "aprobada por el supervisor"), (25, "firma del supervisor")):
                    add(ts=x.fecha_supervisor, fuente="sonny", paso=paso, cond=cond, actor="ST", de="Sonny", asunto=asunto, texto="",
                        prueba=f"Sonny: el supervisor aprueba la planilla {x.id_planilla} (Lista de Planillas, fecha del supervisor)",
                        ref={"id_planilla": x.id_planilla, "lista": True, "campo": "Fecha_supervisor"})
            if x.fecha_ingeniero and (x.estado_ingeniero or "").lower().startswith("aprob"):
                add(ts=x.fecha_ingeniero, fuente="sonny", paso=26, cond="firma del ingeniero", actor="PA", de="Sonny", asunto=asunto, texto="",
                    prueba=f"Sonny: el ingeniero aprueba la planilla {x.id_planilla} (Lista de Planillas, fecha del ingeniero)",
                    ref={"id_planilla": x.id_planilla, "lista": True, "campo": "Fecha_ingeniero"})

    # 3d. PMS (CHANGELOG 2026-10-23): cada cambio de estado de sus solicitudes que mueve el FL-03 (pasos 1, 4, 5, 9, 14, 15). Fuente de fondo: no
    # retrocede el paso actual. Lo que llega por LCL va al gemelo que no es línea viva, como el resto.
    if fd["pms"] and recibe_por_lcl:
        for lg in fd["logs"]:
            paso, cond = gf.paso_pms(lg.estado_fin)
            if not paso or (lg.estado_ini or "") == (lg.estado_fin or ""):
                continue
            add(ts=lg.fecha, fuente="pms", paso=paso, cond=cond, actor="Enel", de="PMS", asunto=f"PMS · {lg.request_code}"[:500], texto=lg.substatus or "",
                prueba=f"PMS: la solicitud {lg.request_code} pasa a «{lg.estado_fin}»"[:300], ref={"request_code": lg.request_code, "estado": lg.estado_fin, "pms": True})

    # 3e. SCM (CHANGELOG 2026-10-23): cada libreto autocertificado es el paso 37 en su fecha, con su estado de hoy. SCM no fecha los cambios de estado
    # (RIF, ACQ, RIL, SAP): no se inventa una fecha para el paso 39; lo dicen la franja y las alertas.
    if fd["libretos"] and recibe_por_lcl:
        corte = fd["cargas"].get("scm")
        for l in fd["libretos"]:
            if not l.fecha_autocert:
                continue
            que = gf.ESTADOS_LIBRETO.get(l.estado_libreto or "", ("?", l.estado_libreto or "?"))[1]
            add(ts=l.fecha_autocert, fuente="scm", paso=37, cond=f"hoy {l.estado_libreto or '?'}"[:60], actor="F", de="SCM",
                asunto=(f"SCM · LCL {l.lcl} · libreto {l.libreto}" + (f" · planilla {l.planilla}" if l.planilla else ""))[:500], texto=(l.descripcion or "")[:300],
                prueba=(f"SCM: libreto {l.libreto} autocertificado · MO {gf._m(l.mo_pesos)} · en el exporte del "
                        f"{corte:%d/%m}: {l.estado_libreto} ({que})" if corte else f"SCM: libreto {l.libreto} autocertificado")[:300],
                ref={"lcl": l.lcl, "libreto": l.libreto, "estado": l.estado_libreto, "planilla": l.planilla, "regla": l.regla_enlace, "scm": True})

    # 4. Registros de los módulos (Liquidador, Balance) que nombran la clave en el archivo
    regs = (await db.execute(select(RegistroSGO7).where(RegistroSGO7.archivo.ilike(f"%{clave}%"), RegistroSGO7.modulo != "gestion")   # CHANGELOG 2026-10-23: sin las descargas de evidencia
                             .order_by(RegistroSGO7.ts).limit(300))).scalars().all() if recibe_por_lcl else []
    for r in regs:
        m = _ACCIONES.get((r.modulo, r.accion))
        if not m:
            continue
        paso, cond, prueba, actor = m
        det = r.detalle or {}
        if r.accion == "liquidador_resultado":
            prueba = f"Liquidador Automático: {det.get('ok', '?')} filas OK · {det.get('fallidos', '?')} fallidas"
        add(ts=r.ts, fuente="sgo7", paso=paso, cond=cond, actor=actor, prueba=prueba, de=r.usuario or "", asunto=(r.archivo or r.accion)[:300],
            texto="", ref={"registro_id": r.id, "modulo": r.modulo, "accion": r.accion})


async def _recalcular(db: AsyncSession, p: GestionProyecto, regs_cm: list[dict], enlazar: bool = True) -> None:
    fd = await _fuentes_dinero(db, p)
    if enlazar:
        await _enlazar_fuentes(db, p, regs_cm, fd)
        await db.flush()
    evs = (await db.execute(select(GestionEvento).where(GestionEvento.proyecto_id == p.id).order_by(GestionEvento.ts, GestionEvento.id))).scalars().all()
    lista = []
    for e in evs:
        son = e.fuente == "sonny" and bool((e.ref or {}).get("sonny"))           # Historial de Sonny: el rechazo se mide por planilla
        lista.append(dict(ts=_local(e.ts), paso=e.paso, cond=e.cond, fuente=e.fuente, actor=e.actor, prueba=e.prueba, de=e.de,
                          planilla=(e.ref or {}).get("id_planilla") if son else None, estado_sonny=(e.ref or {}).get("estado") if son else None,
                          obs=e.texto if son else None))                                   # CHANGELOG 2026-10-25: el motivo del rechazo
    ahora = datetime.now(_TZ).replace(tzinfo=None)
    est = gf.calcular_estado(lista, ahora)
    din = await _dinero(db, p, fd, ahora)                          # CHANGELOG 2026-10-23: la franja del dinero y sus 5 alertas
    mo_pl = {str(x.id_planilla): float(x.mo_real or 0) for x in fd["planillas"]}
    try:                                                           # CHANGELOG 2026-10-25: los rechazos del caso con su motivo (abiertos y anteriores)
        din["rechazos"] = gf.resumir_rechazos(est.get("rechazos"), mo_pl)
    except Exception as e:                                        # nunca tumba el recálculo
        din["rechazos"] = {"error": f"{type(e).__name__}: {str(e)[:200]}", "abiertos": [], "previos": []}
    E = await _evidencia_resumen(db, p, fd, din, ahora)            # CHANGELOG 2026-10-23: ¿cada suma reproduce su cifra?
    # CHANGELOG 2026-10-25: cada alerta con su plata y su frase llana; estado del trabajo, «qué falta para cobrar», la cadena y la plata (lo leen las páginas nuevas)
    alertas = est["alertas"] + [{k: v for k, v in a.items()} for a in din.get("alertas") or []]
    try:
        gf.vestir_alertas(alertas, din, mo_pl)
        ux = gf.armar_ux(din, alertas, E, est["pilar"], _pago_confirmado())
        ux["facturas"] = [{"factura": f.get("factura"), "estado": f.get("estado"), "vencida": bool(f.get("vencida")), "total": f.get("total")}
                          for f in (din.get("cobro") or {}).get("facturas") or []]
        ux["calculado_en"] = ahora.isoformat(timespec="seconds")
        din["ux"] = ux
        if din.get("resumen") is not None:
            din["resumen"]["ux"] = {"estado": ux["estado"], "valor": ux["valor"]}
    except Exception as e:                              # nunca tumba el recálculo: la página dice que falta y el control lo cuenta
        din["ux"] = {"version": gf.UX_VERSION, "error": f"{type(e).__name__}: {str(e)[:200]}", "calculado_en": ahora.isoformat(timespec="seconds")}
    p.dinero = din
    sl = din.get("scm_lcl") or {}                                  # CHANGELOG 2026-10-25: cerrado = todas sus LCL en CONCLUSA (None = sin LCL en SCM)
    ce = _aware(datetime.fromisoformat(sl["cerrado_en"])) if sl.get("cerrado_en") else None
    if (p.cerrado, p.cerrado_en) != (sl.get("cerrado"), ce):
        p.cerrado, p.cerrado_en = sl.get("cerrado"), ce
    fij = p.fijados or {}
    p.paso_actual, p.max_paso, p.pilar = est["paso_actual"], est["max_paso"], est["pilar"]
    p.paso_fecha = _aware(est["paso_fecha"])
    p.estado_texto, p.fuente_paso = est["estado_texto"], est["fuente_paso"]
    p.alertas = alertas
    p.n_eventos = len(evs)
    p.n_correos = sum(1 for e in evs if e.fuente == "correo" and e.directo)
    p.ultima_actividad = max((e.ts for e in evs), default=None)
    p.fuentes = {k: sum(1 for e in evs if e.fuente == k) for k in _FUENTES}
    padre = p.padre or {}
    if "supervisor" not in fij:
        s = padre.get("supervisor") or gf.supervisor_de(lista)
        if s:
            p.supervisor = re.sub(r"\s+", " ", s).strip()
    if "zona" not in fij:
        z = padre.get("zona") or gf.zona_de([p.nombre or ""] + [e.asunto or "" for e in evs[:40]])
        if z:
            p.zona = z
    if "tipo" not in fij:
        p.tipo = _tipo_de(p.codigo or p.clave, padre)
    p.actualizado_en = _ahora()


def _cod_cm(codigo: str) -> str:
    return codigo if codigo.startswith("DF") else "LCL" + codigo


async def _refrescar_sgs(proyectos: list[GestionProyecto]) -> None:
    """Relee en SGS (por el Centro MO) el padre de cada proyecto: supervisor, zona, MO acumulada y liquidaciones."""
    proyectos = [p for p in proyectos if p.codigo]
    if not proyectos:
        return
    info = await _cm_sgs([_cod_cm(p.codigo) for p in proyectos], detalle=True, mat=True)
    for p in proyectos:
        padres = (info.get(_cod_cm(p.codigo)) or {}).get("padres") or []
        codes = set((p.padre or {}).get("codes") or [(p.padre or {}).get("code")])
        mios = [x for x in padres if x.get("code") in codes]
        if not mios:
            continue
        base = next((x for x in mios if x.get("code") == (p.padre or {}).get("code")), mios[0])
        p.padre = {k: base.get(k) for k in ("id", "code", "name", "order_number", "supervisor", "departamento", "zona")} | {"codes": sorted(x.get("code") for x in mios)}
        if len(mios) > 1:
            p.padre["name"] = " + ".join((x.get("name") or x.get("code") or "")[:40] for x in sorted(mios, key=lambda x: x.get("code") or ""))
        p.sgs = _sgs_resumen_padre(sorted(mios, key=lambda x: x.get("code") or ""))
        # CHANGELOG 2026-10-23: el material de Enel de la familia (consumido y despachado a costo SGS, por código) y, si el LCL tiene otros padres
        # (gemelos), lo del LCL completo: PMS, Sonny y SCM son por LCL y la franja se compara contra eso.
        p.sgs["mat"] = _sumar_mat([x.get("mat") for x in mios])
        if len(padres) > len(mios):
            p.sgs["mo_lcl"] = round(sum(float(x.get("mo_total") or 0) for x in padres), 2)
            p.sgs["mat_lcl"] = _sumar_mat([x.get("mat") for x in padres])
            p.sgs["padres_lcl"] = [{k: x.get(k) for k in ("code", "supervisor", "zona", "mo_total")} for x in padres]


# ── Agente del PC ────────────────────────────────────────────────────────────────────────────────────
@router.get("/agente/estado", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_estado(buzon: str = Query(..., max_length=150), version: str = "", db: AsyncSession = Depends(get_db)):
    ag = (await db.execute(select(GestionAgente).where(GestionAgente.buzon == buzon))).scalar_one_or_none()
    if ag is None:
        ag = GestionAgente(buzon=buzon, version=version)
        db.add(ag)
    ag.version = version or ag.version
    ag.ultimo_contacto = _ahora()
    await db.commit()
    hilos = (await db.execute(select(GestionProyecto.hilo).where(GestionProyecto.activo.is_(True), GestionProyecto.hilo != "").limit(1500))).scalars().all()
    return {"cursor": ag.cursor.isoformat() if ag.cursor else None, "hilos": sorted(set(hilos)), "dias_atras_defecto": 240,
            "config": {"poll_minutos": 10, "lote": 15, "html_max": 200_000}}


@router.post("/agente/ingesta", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_ingesta(body: IngestaIn, db: AsyncSession = Depends(get_db)):
    ag = (await db.execute(select(GestionAgente).where(GestionAgente.buzon == body.buzon))).scalar_one_or_none()
    if ag is None:
        ag = GestionAgente(buzon=body.buzon, version=body.version)
        db.add(ag)
        await db.flush()
    regs_cm = await _cm_registros()
    por_entry = {r["entry_id"]: r for r in regs_cm if r.get("entry_id")}

    # 1. candidatos de cada correo (y los códigos que nombran sus mensajes citados); una sola consulta de existencia a SGS
    plan, todos = [], set()
    for c in body.correos:
        r = por_entry.get(c.entry_id)
        if r and r.get("codigo_validado") and r.get("codigo"):
            cands = [{"codigo": r["codigo"].upper(), "fuente": "centro_mo"}]
        else:
            cands = gf.candidatos_codigo(c.asunto, c.cuerpo)
        mensajes = gf.desglosar(dict(cuerpo=c.cuerpo, fecha=c.fecha, de=c.de, de_correo=c.de_correo, para=c.para, cc=c.cc, asunto=c.asunto))
        for m in mensajes:
            if not m["directo"]:
                m["codigo_propio"] = gf.codigo_propio(m.get("texto"))
                # Un correo viejo reutilizado con otro asunto (CHANGELOG 2026-10-23): si el texto citado no nombra código, vale el de SU asunto
                # (solo para ese mensaje; no se hereda a la cadena).
                if not m["codigo_propio"] and gf.otro_hilo(m.get("asunto") or "", c.asunto):
                    m["codigo_asunto"] = gf.codigo_propio(m.get("asunto"))
                for k in ("codigo_propio", "codigo_asunto"):
                    if m.get(k):
                        todos.add(m[k])
        plan.append((c, cands, mensajes))
        todos |= {x["codigo"] for x in cands}
    existencia = await _cm_sgs(sorted(todos), detalle=False) if todos else {}

    tocados: dict[int, GestionProyecto] = {}
    correos_ok = eventos_nuevos = sin_proyecto = citados_omitidos = citados_otro = 0
    decisiones = []
    for c, cands, mensajes in plan:
        hilo = gf.hilo_normalizado(c.conversacion or c.asunto)
        if not (cands and cands[0]["fuente"] == "centro_mo") and gf.es_correo_lista(c.asunto, c.cuerpo):
            decisiones.append({"asunto": c.asunto[:90], "fecha": c.fecha, "proyecto": None, "via": "lista (reporte de varios tickets; no se asigna)", "log": []})
            sin_proyecto += 1
            continue
        ganador, log = gf.elegir_codigo(cands, existencia, c.de)
        if ganador:
            codigo = gf.clave_de(ganador["codigo"])
            padre, como = gf.elegir_padre(ganador.get("padres") or [], c.de, hilo, await _hilos_por_padre(db, codigo), para=c.para)
            if padre is None:
                sin_proyecto += 1
                decisiones.append({"asunto": c.asunto[:90], "fecha": c.fecha, "proyecto": None, "via": "sin padre en SGS", "log": log[-3:]})
                continue
            p = await _proyecto_para(db, codigo, padre, c.asunto, hilo)
            via = f"{ganador['fuente']} · padre por {como}"
            log.append(f"padre {padre['code']} ({como})")
        else:
            p = await _por_hilo(db, hilo)
            via = "hilo" if p else None
        decisiones.append({"asunto": c.asunto[:90], "fecha": c.fecha, "proyecto": p.clave if p else None, "via": via, "log": log[-3:]})
        if p is None:
            sin_proyecto += 1
            continue
        correos_ok += 1
        arriba, _ = gf.partir_cuerpo(c.cuerpo)
        p.codigos = _union(p.codigos, gf.extraer_codigos(f"{c.asunto}\n{arriba}"))   # solo asunto + mensaje reciente
        if not p.hilo and hilo:
            p.hilo = hilo
        huellas_por_proy: dict[int, set] = {}

        async def huellas_de(proy):
            if proy.id not in huellas_por_proy:
                huellas_por_proy[proy.id] = set((await db.execute(select(GestionEvento.huella).where(GestionEvento.proyecto_id == proy.id))).scalars().all())
            return huellas_por_proy[proy.id]

        gemelos = gf.agrupar_padres(ganador.get("padres") or []) if ganador else []
        cod_ganador = ganador["codigo"].upper() if ganador else None
        # Los mensajes citados se recorren en orden de fecha: cada uno vale por el código que nombra su propio texto, y si no
        # nombra ninguno hereda el del mensaje anterior de la cadena (un hilo puede cambiar de LCL a mitad de camino: de un LCL
        # anterior que ya no existe en SGS al nuevo). Lo que es de OTRO código existente va al caso de
        # ese código; lo que es de un código que ya no existe en SGS solo se conserva si lo escribió el supervisor del caso.
        citados = sorted([m for m in mensajes if not m["directo"] and m.get("fecha")], key=lambda m: m["fecha"])
        vigente = None
        for m in citados:
            if m.get("codigo_propio"):
                vigente = m["codigo_propio"]
            m["codigo_vigente"] = vigente
        primero = next((m["codigo_propio"] for m in citados if m.get("codigo_propio")), None)
        for m in citados:
            if not m.get("codigo_vigente"):
                m["codigo_vigente"] = primero or cod_ganador
            if m.get("codigo_asunto"):
                m["codigo_vigente"] = m["codigo_asunto"]
        html_guardado = False
        for m in [x for x in mensajes if x["directo"]] + citados:
            if not m.get("fecha"):
                continue
            m["adjuntos"] = c.adjuntos if m["directo"] else []
            destino = p
            if not m["directo"]:
                cod_m = (m.get("codigo_vigente") or cod_ganador or "").upper()
                if cod_ganador and cod_m and cod_m != cod_ganador and gf.clave_de(cod_m) != gf.clave_de(cod_ganador):
                    ex = existencia.get(cod_m) or {}
                    if ex.get("existe"):
                        pm, como = gf.elegir_padre(ex.get("padres") or [], m.get("de") or "", "", None, para=m.get("para") or "")
                        if pm is None:
                            citados_omitidos += 1
                            continue
                        destino = await _proyecto_para(db, gf.clave_de(cod_m), pm, m.get("asunto") or c.asunto, gf.hilo_normalizado(m.get("asunto") or ""))
                        tocados[destino.id] = destino
                        citados_otro += 1
                    elif m.get("codigo_asunto") or not gf.es_del_supervisor(m.get("de") or "", (p.padre or {}).get("supervisor") or p.supervisor or ""):
                        citados_omitidos += 1          # historia de un código que ya no existe en SGS: de otra persona, o de otra conversación
                        continue
                elif len(gemelos) > 1:
                    # Mismo código, varios padres: si lo escribió el supervisor del OTRO gemelo (línea viva dentro del hilo del
                    # programado) o va dirigido a él (CMS le responde), va al caso de ese gemelo y no se mezcla aquí.
                    pm, como = gf.elegir_padre(gemelos, m.get("de") or "", "", None, para=m.get("para") or "")
                    if como in ("supervisor", "destinatario") and pm and pm["code"] != p.clave:
                        destino = await _proyecto_para(db, codigo, pm, m.get("asunto") or c.asunto, gf.hilo_normalizado(m.get("asunto") or ""))
                        tocados[destino.id] = destino
            huellas = await huellas_de(destino)
            h = gf.huella_mensaje(destino.clave, m)
            if h in huellas:
                continue
            cl = gf.clasificar(m)
            if not cl and not m["directo"]:
                continue
            actor = gf.actor_de(m)
            for i, x in enumerate(cl or [dict(paso=None, cond=None, prueba="Correo del hilo sin paso reconocido")]):
                guardar_html = bool(c.html) and m["directo"] and not html_guardado
                db.add(GestionEvento(
                    proyecto_id=destino.id, ts=_aware(m["fecha"]), fuente="correo", paso=x["paso"], cond=x["cond"], actor=actor, prueba=x["prueba"],
                    de=(m.get("de") or "")[:200], de_correo=(m.get("de_correo") or "")[:200], para=(m.get("para") or "")[:1000],
                    asunto=(m.get("asunto") or "")[:500], texto=(m.get("texto") or "")[:6000],
                    html=_sanear_html(c.html) if guardar_html else None, directo=m["directo"], entry_id=c.entry_id if m["directo"] else None,
                    carpeta=c.carpeta[:200], buzon=body.buzon, adjuntos=(c.adjuntos[:20] if m["directo"] else []),
                    huella=h if i == 0 else _sha(h + f"|{x['paso']}"),
                    ref={"categorias": c.categorias, "via": via, "codigo": ganador["codigo"] if ganador else None} if m["directo"] else None))
                html_guardado = html_guardado or guardar_html
                eventos_nuevos += 1
            huellas.add(h)
        tocados[p.id] = p
    await db.flush()
    await _refrescar_sgs(list(tocados.values()))
    for p in tocados.values():
        await _recalcular(db, p, regs_cm)
    ag.ultimo_scan = _ahora()
    ag.version = body.version or ag.version
    ag.correos_total = (ag.correos_total or 0) + correos_ok
    if body.cursor:
        cur = _iso(body.cursor)
        if cur and (ag.cursor is None or cur > ag.cursor):
            ag.cursor = cur
    await db.commit()
    return {"ok": True, "correos": correos_ok, "sin_proyecto": sin_proyecto, "eventos_nuevos": eventos_nuevos,
            "proyectos_tocados": len(tocados), "citados_otro_codigo": citados_otro, "citados_omitidos": citados_omitidos, "decisiones": decisiones}


def _fecha_sonny(s: str) -> Optional[datetime]:
    """«2026-04-08 08:57:11» (hora Colombia, como la exporta Sonny) → datetime con zona."""
    s = (s or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s[:19], fmt).replace(tzinfo=_TZ)
        except ValueError:
            continue
    return None


@router.post("/agente/sonny", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_sonny(body: SonnyIn, db: AsyncSession = Depends(get_db)):
    """Historial de Planillas de Sonny (CHANGELOG 2026-10-23; export manual, lo sube cargar_sonny_gestion.py en lotes de 500). Upsert por
    id_cambio: reenviar el mismo archivo no duplica. Recalcula los casos cuyo LCL o ticket nombran las filas nuevas o cambiadas."""
    filas, invalidas = {}, 0
    for f in body.filas:
        ts, k, est = _fecha_sonny(f.fecha), f.id_cambio.strip(), gf.estado_sonny(f.estado)
        if not ts or not k or not est or not f.id_planilla.strip():
            invalidas += 1
            continue
        filas[k] = dict(id_planilla=f.id_planilla.strip(), codigo=gf.clave_de(f.codigo or "") or None, ticket=gf.clave_de(f.ticket or "") or None,
                        nodo=re.sub(r"\s+", " ", f.nodo or "").strip()[:200], supervisor=re.sub(r"\s+", " ", f.supervisor or "").strip()[:150],
                        estado=est[:60], fecha=ts, observacion=(f.observacion or "").strip()[:2000] or None)
    existentes = {s.id_cambio: s for s in (await db.execute(select(GestionSonny).where(GestionSonny.id_cambio.in_(list(filas))))).scalars().all()} if filas else {}
    cid = await _carga_valida(db, body.carga_id, "sonny")
    nuevas = actualizadas = iguales = 0
    codigos_cambiados: set = set()
    for k, v in filas.items():
        s = existentes.get(k)
        if s is None:
            db.add(GestionSonny(id_cambio=k, carga_id=cid, **v))
            nuevas += 1
            codigos_cambiados |= {v["codigo"], v["ticket"]}
            continue
        if cid and s.carga_id != cid:               # CHANGELOG 2026-10-23: la fila vino en esta carga (aunque no cambie)
            s.carga_id = cid
        dif = {c: x for c, x in v.items() if getattr(s, c) != x}
        if not dif:
            iguales += 1
            continue
        codigos_cambiados |= {s.codigo, s.ticket, v["codigo"], v["ticket"]}
        for c, x in dif.items():
            setattr(s, c, x)
        s.actualizado_en = _ahora()
        actualizadas += 1
    await db.flush()

    # Cuántas planillas del lote casan con un caso de Gestión (por LCL o por ticket) y cuántas no
    planillas: dict = {}
    for v in filas.values():
        planillas.setdefault(v["id_planilla"], v)
    cods_lote = sorted({c for v in planillas.values() for c in (v["codigo"], v["ticket"]) if c})
    en_casos = set((await db.execute(select(GestionProyecto.codigo).where(GestionProyecto.codigo.in_(cods_lote)))).scalars().all()) if cods_lote else set()
    por_lcl = por_ticket = sin_codigo = 0
    sin_caso: set = set()
    for v in planillas.values():
        lcl = v["codigo"] if (v["codigo"] or "").startswith("63") else None
        df = v["ticket"] or (v["codigo"] if (v["codigo"] or "").startswith("DF") else None)
        if not lcl and not df:
            sin_codigo += 1
        elif lcl and lcl in en_casos:
            por_lcl += 1
        elif df and df in en_casos:
            por_ticket += 1
        else:
            sin_caso.add(lcl or df)
    desconocidos: dict = {}
    for v in filas.values():
        if not gf.paso_sonny(v["estado"])[3]:
            desconocidos[v["estado"]] = desconocidos.get(v["estado"], 0) + 1

    tocar = sorted(c for c in codigos_cambiados if c)
    ps = (await db.execute(select(GestionProyecto).where(GestionProyecto.codigo.in_(tocar)).order_by(GestionProyecto.id))).scalars().all() if tocar else []
    pendiente = False
    if ps:
        try:
            regs_cm = await _cm_registros()
        except HTTPException:
            regs_cm, pendiente = None, True       # el Centro MO no responde: las filas quedan; recalcular después (POST /recalcular)
        if regs_cm is not None:
            for p in ps:
                await _recalcular(db, p, regs_cm)
    await db.commit()
    return {"ok": True, "archivo": body.archivo, "filas": len(body.filas), "invalidas": invalidas, "nuevas": nuevas, "actualizadas": actualizadas,
            "ya_estaban": iguales, "planillas": len(planillas), "planillas_con_caso": por_lcl + por_ticket, "por_lcl": por_lcl, "por_ticket": por_ticket,
            "planillas_sin_caso": len(planillas) - por_lcl - por_ticket - sin_codigo, "planillas_sin_codigo": sin_codigo,
            "codigos_sin_caso": len(sin_caso), "estados_desconocidos": desconocidos, "casos_recalculados": 0 if pendiente else len(ps),
            "recalculo_pendiente": pendiente}


# ── El hilo del dinero: cargas de Sonny (Lista), PMS, SCM y deuda; casos que nacen sin correo ─────────────────────────
def _num(v) -> float:
    try:
        return round(float(v or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def _t(s: Optional[str], n: int) -> Optional[str]:
    s = re.sub(r"\s+", " ", (s or "")).strip()
    return s[:n] or None


def _dif(obj, nuevo: dict) -> dict:
    return {c: x for c, x in nuevo.items() if getattr(obj, c) != x}


async def _carga_valida(db: AsyncSession, cid: Optional[int], sistema: str) -> Optional[int]:
    """CHANGELOG 2026-10-23: el carga_id que manda el PC, si existe y es de ese sistema (si no, las filas entran sin carga, como antes)."""
    if not cid:
        return None
    c = await db.get(GestionCarga, cid)
    return c.id if c is not None and c.sistema == sistema else None


def _sello(obj, cid: Optional[int]) -> None:
    if cid and obj.carga_id != cid:
        obj.carga_id = cid


async def _recalcular_codigos(db: AsyncSession, codigos: set, recalcular: bool = True) -> tuple[int, bool]:
    """Recalcula los casos cuyo código (LCL o ticket) está en `codigos`. Si el Centro MO no responde, las filas quedan y el
    recálculo queda pendiente (POST /recalcular)."""
    tocar = sorted(c for c in codigos if c)
    if not recalcular or not tocar:
        return 0, False
    ps = (await db.execute(select(GestionProyecto).where(GestionProyecto.codigo.in_(tocar)).order_by(GestionProyecto.id))).scalars().all()
    if not ps:
        return 0, False
    try:
        regs_cm = await _cm_registros()
    except HTTPException:
        return 0, True
    for p in ps:
        await _recalcular(db, p, regs_cm)
    return len(ps), False


async def _reenlazar(db: AsyncSession, lcls: set) -> Counter:
    """R5: rehace el enlace planilla ↔ libreto de estos LCL (gestion_flujo.enlazar_libretos) y lo guarda en el libreto."""
    cnt: Counter = Counter()
    lcls = sorted(c for c in lcls if c)
    if not lcls:
        return cnt
    libs = (await db.execute(select(GestionScmLibreto).where(GestionScmLibreto.lcl.in_(lcls)))).scalars().all()
    if not libs:
        return cnt
    pls = {x.id_planilla: x for x in (await db.execute(select(GestionSonnyPlanilla).where(GestionSonnyPlanilla.lcl.in_(lcls)))).scalars().all()}
    toks = {t for l in libs for t in gf.RX_ID_PLANILLA.findall(re.sub(r"(?i)ITEM\s*\d+", " ", l.descripcion or ""))} - set(pls)
    if toks:
        for x in (await db.execute(select(GestionSonnyPlanilla).where(GestionSonnyPlanilla.id_planilla.in_(sorted(toks))))).scalars().all():
            pls[x.id_planilla] = x
    # CHANGELOG 2026-10-23: las confirmaciones humanas («Es este» / «No es») mandan sobre las reglas; la última «Es este» de un libreto gana
    decs: dict = {}
    for d in (await db.execute(select(GestionDecision).where(GestionDecision.tipo == "planilla_libreto", GestionDecision.lcl.in_(lcls))
                               .order_by(GestionDecision.actualizado_en, GestionDecision.id))).scalars().all():
        x = decs.setdefault((d.lcl, d.libreto), {"si": None, "no": set()})
        if d.decision == "si":
            x["si"] = (d.id_planilla, (d.usuario or "").split("@")[0], _local(d.actualizado_en or d.creado_en))
        else:
            x["no"].add(d.id_planilla)
    faltan = {d for v in decs.values() for d in ([v["si"][0]] if v["si"] else []) + sorted(v["no"])} - set(pls)
    if faltan:
        for x in (await db.execute(select(GestionSonnyPlanilla).where(GestionSonnyPlanilla.id_planilla.in_(sorted(faltan))))).scalars().all():
            pls[x.id_planilla] = x
    res = gf.enlazar_libretos([dict(lcl=l.lcl, libreto=l.libreto, descripcion=l.descripcion, pdls=l.pdls or [], estado_libreto=l.estado_libreto, mo_pesos=l.mo_pesos)
                               for l in libs],
                              [dict(id_planilla=x.id_planilla, lcl=x.lcl, pdl=x.pdl, ticket=x.ticket, mo_real=x.mo_real) for x in pls.values()], decs)
    for l in libs:
        pid, regla, cands = res.get((l.lcl, l.libreto), (None, "sin planilla", []))
        regla = regla[:60]
        if (l.planilla, l.regla_enlace) != (pid, regla):
            l.planilla, l.regla_enlace = pid, regla
        if (l.candidatas or []) != cands:
            l.candidatas = cands
        if l.estado_libreto not in gf.LIBRETO_MUERTO:
            cnt[regla.split(":")[0]] += 1
    return cnt


async def _codigos_de_lcls(db: AsyncSession, lcls: set) -> set:
    """Los LCL y los tickets DF de las planillas de esos LCL (los casos de mayor impacto ven SCM por el LCL de su planilla)."""
    lcls = {c for c in lcls if c}
    if not lcls:
        return set()
    tks = (await db.execute(select(GestionSonnyPlanilla.ticket).where(GestionSonnyPlanilla.lcl.in_(sorted(lcls)), GestionSonnyPlanilla.ticket.isnot(None)))).scalars().all()
    return lcls | {t for t in tks if t}


@router.post("/agente/sonny-planillas", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_sonny_planillas(body: SonnyPlanillasIn, db: AsyncSession = Depends(get_db)):
    """Lista de Planillas de Sonny (CHANGELOG 2026-10-23): MO/MAT real, QR, PDL y las fechas de supervisor e ingeniero. Upsert por id_planilla."""
    filas, invalidas = {}, 0
    for f in body.planillas:
        k = (f.id_planilla or "").strip()
        if not k:
            invalidas += 1
            continue
        lcl = gf.lcl_norm(f.lcl)
        tk = gf.pdl_norm(f.ticket) if (f.ticket or "").upper().startswith("DF") else None
        filas[k] = dict(lcl=lcl, ticket=tk, pdl=_t(f.pdl, 80), qr=_t(f.qr, 100), odm=_t(f.odm, 200), nodo=_t(f.nodo, 200), supervisor=_t(f.supervisor, 150),
                        estado=_t(gf.estado_sonny(f.estado), 60), libreto_sonny=_t(f.libreto_sonny, 20), mo_plan=_num(f.mo_plan), mat_plan=_num(f.mat_plan),
                        mo_real=_num(f.mo_real), mat_real=_num(f.mat_real), fecha_planilla=_fecha_sonny(f.fecha_planilla), fecha_supervisor=_fecha_sonny(f.fecha_supervisor),
                        fecha_ingeniero=_fecha_sonny(f.fecha_ingeniero), fecha_enel=_fecha_sonny(f.fecha_enel), estado_supervisor=_t(f.estado_supervisor, 30),
                        estado_ingeniero=_t(f.estado_ingeniero, 30))
    ex = {x.id_planilla: x for x in (await db.execute(select(GestionSonnyPlanilla).where(GestionSonnyPlanilla.id_planilla.in_(list(filas))))).scalars().all()} if filas else {}
    cid = await _carga_valida(db, body.carga_id, "sonny")
    nuevas = actualizadas = iguales = 0
    cods: set = set()
    lcls: set = set()
    for k, v in filas.items():
        x = ex.get(k)
        if x is None:
            db.add(GestionSonnyPlanilla(id_planilla=k, carga_id=cid, **v))
            nuevas += 1
        else:
            _sello(x, cid)
            d = _dif(x, v)
            if not d:
                iguales += 1
                continue
            cods |= {x.lcl, x.ticket}
            lcls.add(x.lcl)
            for c, val in d.items():
                setattr(x, c, val)
            x.actualizado_en = _ahora()
            actualizadas += 1
        cods |= {v["lcl"], v["ticket"]}
        lcls.add(v["lcl"])
    await db.flush()
    _cache_cargas.clear()
    enl = await _reenlazar(db, lcls)
    casos = set((await db.execute(select(GestionProyecto.codigo).where(GestionProyecto.codigo.in_(sorted({c for v in filas.values() for c in (v["lcl"], v["ticket"]) if c}))))).scalars().all()) if filas else set()
    con_caso = sum(1 for v in filas.values() if v["lcl"] in casos or v["ticket"] in casos)
    n, pend = await _recalcular_codigos(db, cods, body.recalcular)
    await db.commit()
    return {"ok": True, "archivo": body.archivo, "planillas": len(body.planillas), "invalidas": invalidas, "nuevas": nuevas, "actualizadas": actualizadas,
            "ya_estaban": iguales, "con_caso": con_caso, "sin_codigo": sum(1 for v in filas.values() if not v["lcl"] and not v["ticket"]),
            "enlace_libretos": dict(enl), "casos_recalculados": n, "recalculo_pendiente": pend}


def _nodo_fila(f: SonnyNodoIn) -> Optional[dict]:
    """CHANGELOG 2026-10-25: el nodo tal como se guarda (sin personas). None si el Lcl no es un LCL 63… ni un ticket DF…."""
    # La columna Lcl trae a veces dos LCL separados por «;»: el nodo va al primero y se encuentra por cualquiera (CHANGELOG 2026-10-25)
    cods = []
    for t in re.split(r"[;,/\s]+", f.codigo or ""):
        c = gf.lcl_norm(t) or (gf.pdl_norm(t) if t.strip().upper().startswith("DF") else None)
        if c and c not in cods:
            cods.append(c)
    if not cods:
        return None
    cod = cods[0]
    ids = sorted({re.sub(r"[^0-9A-Za-z]", "", x) for x in f.id_planillas if re.sub(r"[^0-9A-Za-z]", "", x)})[:20]       # Ids de Sonny (numéricos)
    return dict(codigo=cod, codigos=cods[:5], odm=_t(f.odm, 200), pdl=_t(f.pdl, 300), qr=_t(f.qr, 100), nodo=_t(f.nodo, 120), estado_nodo=_t(f.estado_nodo, 20), id_planillas=ids,
                estado_planilla=_t(f.estado_planilla, 200), f_creacion=_fecha_sonny(f.f_creacion), f_estado=_fecha_sonny(f.f_estado),
                estado_material=_t(f.estado_material, 20), movimiento=_t(f.movimiento, 20), f_almacen=_fecha_sonny(f.f_almacen), n_almacen=int(f.n_almacen or 0),
                id_gestion=_t(re.sub(r"\D", "", f.id_gestion or ""), 20), f_confirmado=_fecha_sonny(f.f_confirmado))


def _act_fila(f: SonnyPlanActIn) -> dict:
    roles = {k: int(v) for k, v in (f.roles or {}).items() if k in ("enel", "dominion", "interventoria", "otro") and isinstance(v, (int, float))}
    return dict(f_creada=_fecha_sonny(f.f_creada), estado_act=_t(gf.estado_sonny(f.estado_act) if f.estado_act else "", 60), f_estado_act=_fecha_sonny(f.f_estado_act),
                f_firma=_fecha_sonny(f.f_firma), n_firmas=int(f.n_firmas or 0), roles=roles, f_ultima=_fecha_sonny(f.f_ultima),
                que_ultima=_t(f.que_ultima if f.que_ultima in gf.QUE_ACT else "", 20), n_act=int(f.n_act or 0))


async def _recalcular_sonny(db: AsyncSession, cods: set, recalcular: bool) -> tuple[int, bool]:
    """CHANGELOG 2026-10-25: recalcula los casos de estos códigos: hasta 30 aquí mismo; más, en segundo plano (cientos de casos en la carga completa)."""
    if not recalcular or not cods:
        return 0, False
    casos = set((await db.execute(select(GestionProyecto.codigo).where(GestionProyecto.codigo.in_(sorted(c for c in cods if c))))).scalars().all())
    if not casos:
        return 0, False
    if len(casos) <= 30:
        n, _ = await _recalcular_codigos(db, casos)
        return n, False
    await db.commit()
    t = asyncio.create_task(_recalcular_codigos_fondo(casos))
    _TAREAS_SONNY.add(t)
    t.add_done_callback(_TAREAS_SONNY.discard)
    return 0, True


@router.post("/agente/sonny-nodos", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_sonny_nodos(body: SonnyNodosIn, db: AsyncSession = Depends(get_db)):
    """CHANGELOG 2026-10-25 Sonny por etapa (lo sube cargar_sonny_gestion.py en lotes de 500):
      nodos     — la «Lista Cruzada de Nodos» con lo de «Nodos por gestion almacen»: upsert por Id_nodo. En el último lote (fin=true) los nodos que
                  no vinieron en esta carga dejan de estar vigentes (Sonny los borró o la consulta ya no los trae); nunca se borran.
      planillas — lo que dice el Historial de Actividades de cada planilla (creación, estado, firmas, última acción): upsert por Id de planilla.
    Recalcula los casos de los códigos que cambian. Nada de esto toca la MO de SGS."""
    cid = await _carga_valida(db, body.carga_id, "sonny")
    out = {"ok": True, "archivo": body.archivo, "carga_id": cid}
    cods: set = set()
    # ── nodos
    filas, inval = {}, 0
    for f in body.nodos:
        v = _nodo_fila(f)
        k = re.sub(r"\D", "", f.id_nodo or "")[:20]
        if not k or v is None:
            inval += 1
            continue
        filas[k] = v
    ex = {x.id_nodo: x for x in (await db.execute(select(GestionSonnyNodo).where(GestionSonnyNodo.id_nodo.in_(list(filas))))).scalars().all()} if filas else {}
    nuevas = actualizadas = iguales = 0
    for k, v in filas.items():
        x = ex.get(k)
        if x is None:
            db.add(GestionSonnyNodo(id_nodo=k, carga_id=cid, vigente=True, **v))
            nuevas += 1
            cods.add(v["codigo"])
            continue
        _sello(x, cid)
        d = _dif(x, v)
        if not x.vigente:
            x.vigente = True
            d = d or {"vigente": True}
        if not d:
            iguales += 1
            continue
        cods |= {x.codigo, v["codigo"]}
        for c, val in d.items():
            setattr(x, c, val)
        x.actualizado_en = _ahora()
        actualizadas += 1
    await db.flush()
    no_vig = 0
    carga = await db.get(GestionCarga, cid) if cid else None
    # La Lista Cruzada es de todo el contrato: lo que no vino en esta carga ya no está en Sonny. Una carga de prueba solo barre sus códigos.
    barre_todo = body.barrer is None and carga is not None and not (carga.archivo or "").startswith("prueba-")
    if body.fin and cid and body.nodos and (barre_todo or body.barrer):
        conds = [GestionSonnyNodo.vigente.is_(True), or_(GestionSonnyNodo.carga_id.is_(None), GestionSonnyNodo.carga_id != cid)]
        if not barre_todo:
            conds.append(GestionSonnyNodo.codigo.in_([c for c in body.barrer if c]))
        viejos = (await db.execute(select(GestionSonnyNodo).where(*conds))).scalars().all()
        for x in viejos:
            x.vigente, x.actualizado_en = False, _ahora()
            cods.add(x.codigo)
        no_vig = len(viejos)
        await db.flush()
    # ── actividad por planilla
    acts = {}
    for f in body.planillas:
        k = re.sub(r"[^0-9A-Za-z]", "", f.id_planilla or "")[:40]
        if k:
            acts[k] = _act_fila(f)
    exa = {x.id_planilla: x for x in (await db.execute(select(GestionSonnyPlanAct).where(GestionSonnyPlanAct.id_planilla.in_(list(acts))))).scalars().all()} if acts else {}
    a_nuevas = a_act = a_igual = 0
    cambiadas = []
    for k, v in acts.items():
        x = exa.get(k)
        if x is None:
            db.add(GestionSonnyPlanAct(id_planilla=k, carga_id=cid, **v))
            a_nuevas += 1
            cambiadas.append(k)
            continue
        _sello(x, cid)
        d = _dif(x, v)
        if not d:
            a_igual += 1
            continue
        for c, val in d.items():
            setattr(x, c, val)
        x.actualizado_en = _ahora()
        a_act += 1
        cambiadas.append(k)
    await db.flush()
    if cambiadas:
        for lcl, tk in (await db.execute(select(GestionSonnyPlanilla.lcl, GestionSonnyPlanilla.ticket).where(GestionSonnyPlanilla.id_planilla.in_(cambiadas)))).all():
            cods |= {lcl, tk}
    _cache_sonny.clear()
    _cache_cargas.clear()
    cods = {c for c in cods if c}
    con_caso = set((await db.execute(select(GestionProyecto.codigo).where(GestionProyecto.codigo.in_(sorted({v["codigo"] for v in filas.values()}))))).scalars().all()) if filas else set()
    n, fondo = await _recalcular_sonny(db, cods, body.recalcular)
    await db.commit()
    out.update(nodos=len(body.nodos), invalidas=inval, nuevas=nuevas, actualizadas=actualizadas, ya_estaban=iguales, no_vigentes=no_vig,
               con_caso=sum(1 for v in filas.values() if v["codigo"] in con_caso), sin_planilla=sum(1 for v in filas.values() if not v["id_planillas"]),
               planillas=len(body.planillas), act_nuevas=a_nuevas, act_actualizadas=a_act, act_ya_estaban=a_igual,
               codigos_tocados=len(cods), casos_recalculados=n, recalculo_en_fondo=fondo)
    return out


@router.post("/agente/pms", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_pms(body: PmsIn, db: AsyncSession = Depends(get_db)):
    """PMS (CHANGELOG 2026-10-23): solicitudes (request.csv) y su historial de estados (logs.csv), sin nombres ni usuarios. Upsert por
    request_code y por huella del cambio; reenviar el mismo zip no duplica. R1: el caso se liga por LCL o por el DF de subprojects."""
    sol, invalidas = {}, 0
    for s in body.solicitudes:
        rc = (s.request_code or "").strip()
        if not rc:
            invalidas += 1
            continue
        sol[rc] = dict(subprojects=_t(s.subprojects, 300), lcls=sorted({c for c in (gf.lcl_norm(x) for x in s.lcls) if c}),
                       dfs=sorted({c for c in (gf.pdl_norm(x) for x in s.dfs) if c and c.startswith("DF")}),
                       pdls=sorted({_t(x, 60) for x in s.pdls if _t(x, 60)})[:80], odms=sorted({_t(x, 30) for x in s.odms if _t(x, 30)})[:80],
                       agp=_t(s.agp, 60), macroproceso=_t(s.macroproceso, 100), services=_num(s.services), materials=_num(s.materials), total=_num(s.total),
                       status=_t(s.status, 60), substatus=_t(s.substatus, 60), request_date=_fecha_sonny(s.request_date), fecha_estado=_fecha_sonny(s.fecha_estado))
    ex = {x.request_code: x for x in (await db.execute(select(GestionPms).where(GestionPms.request_code.in_(list(sol))))).scalars().all()} if sol else {}
    cid_s, cid_c = await _carga_valida(db, body.carga_solicitudes, "pms"), await _carga_valida(db, body.carga_cambios, "pms")
    nuevas = actualizadas = iguales = 0
    cods: set = set()
    for rc, v in sol.items():
        x = ex.get(rc)
        if x is None:
            db.add(GestionPms(request_code=rc, carga_id=cid_s, **v))
            nuevas += 1
            cods |= set(v["lcls"]) | set(v["dfs"])
            continue
        _sello(x, cid_s)
        d = _dif(x, v)
        if not d:
            iguales += 1
            continue
        cods |= set(x.lcls or []) | set(x.dfs or []) | set(v["lcls"]) | set(v["dfs"])
        for c, val in d.items():
            setattr(x, c, val)
        x.actualizado_en = _ahora()
        actualizadas += 1
    logs = {}
    for lg in body.logs:
        f = _fecha_sonny(lg.fecha)
        if not f or not lg.request_code.strip():
            invalidas += 1
            continue
        logs[lg.huella] = dict(request_code=lg.request_code.strip()[:100], estado_ini=_t(lg.estado_ini, 60), estado_fin=_t(lg.estado_fin, 60), substatus=_t(lg.substatus, 60), fecha=f)
    ya = set((await db.execute(select(GestionPmsLog.huella).where(GestionPmsLog.huella.in_(list(logs))))).scalars().all()) if logs else set()
    if cid_c and ya:                                 # CHANGELOG 2026-10-23: los cambios que ya estaban vinieron también en esta carga
        await db.execute(GestionPmsLog.__table__.update().where(GestionPmsLog.huella.in_(sorted(ya))).values(carga_id=cid_c))
    logs_nuevos = 0
    rc_log = set()
    for h, v in logs.items():
        if h in ya:
            continue
        db.add(GestionPmsLog(huella=h, carga_id=cid_c, **v))
        logs_nuevos += 1
        rc_log.add(v["request_code"])
    await db.flush()
    if rc_log - set(sol):
        for x in (await db.execute(select(GestionPms).where(GestionPms.request_code.in_(sorted(rc_log - set(sol)))))).scalars().all():
            cods |= set(x.lcls or []) | set(x.dfs or [])
    for rc in rc_log & set(sol):
        cods |= set(sol[rc]["lcls"]) | set(sol[rc]["dfs"])
    _cache_cargas.clear()
    casos = set((await db.execute(select(GestionProyecto.codigo).where(GestionProyecto.codigo.isnot(None)))).scalars().all())
    por_lcl = sum(1 for v in sol.values() if any(c in casos for c in v["lcls"]))
    por_df = sum(1 for v in sol.values() if not any(c in casos for c in v["lcls"]) and any(c in casos for c in v["dfs"]))
    sin_paso = Counter(v["status"] for v in sol.values() if v["status"] and not gf.paso_pms(v["status"])[0])
    n, pend = await _recalcular_codigos(db, cods, body.recalcular)
    await db.commit()
    return {"ok": True, "archivo": body.archivo, "solicitudes": len(sol), "invalidas": invalidas, "nuevas": nuevas, "actualizadas": actualizadas, "ya_estaban": iguales,
            "logs": len(logs), "logs_nuevos": logs_nuevos, "con_caso": por_lcl + por_df, "por_lcl": por_lcl, "por_df": por_df,
            "sin_codigo": sum(1 for v in sol.values() if not v["lcls"] and not v["dfs"]), "estados_sin_paso": dict(sin_paso),
            "casos_recalculados": n, "recalculo_pendiente": pend}


@router.post("/agente/scm", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_scm(body: ScmIn, db: AsyncSession = Depends(get_db)):
    """SCM (CHANGELOG 2026-10-23): libretos (Informe LCL_LCA) con su MO en pesos (prestaciones × precio QP, calculada en el PC) y su material
    posado (Informe Materiales). Upsert por (lcl, libreto); el material de cada libreto se reemplaza por el del exporte. Después
    se rehace el enlace planilla ↔ libreto de los LCL del lote (R5) y se recalculan sus casos."""
    corte = _fecha_sonny(body.corte)
    if not corte:
        raise HTTPException(status_code=422, detail="corte inválido (AAAA-MM-DD HH:MM:SS)")
    cid_l, cid_p, cid_m = (await _carga_valida(db, body.carga_libretos, "scm"), await _carga_valida(db, body.carga_prestaciones, "scm"),
                           await _carga_valida(db, body.carga_materiales, "scm"))
    libs, invalidos = {}, 0
    prests: dict = {}
    for l in body.libretos:
        lcl, lib = gf.lcl_norm(l.lcl), gf.libreto_norm(l.libreto)
        if not lcl or not lib:
            invalidos += 1
            continue
        desc = _t(l.descripcion, 300)
        mats: dict = {}
        for m in l.materiales:
            cod = re.sub(r"\s+", "", m.material or "")[:40]
            if cod:
                x = mats.setdefault(cod, [0.0, 0.0, ""])
                x[0] += _num(m.posada)
                x[1] += _num(m.recuperada)
                x[2] = x[2] or (_t(m.unidad, 20) or "")
        if l.prestaciones is not None:                # CHANGELOG 2026-10-23: el detalle que suma la MO del libreto (None = el PC no lo mandó)
            prests[(lcl, lib)] = [dict(n=int(r.n or i), prestacion=(_t(r.prestacion, 20) or "").upper(), unidad=_t(r.unidad, 20), cantidad=float(r.cantidad or 0),
                                       puntos_unit=float(r.puntos_unit or 0), puntos=float(r.puntos or 0),
                                       precio=None if r.precio is None else float(r.precio), pesos=round(float(r.pesos or 0), 2), ticket=_t(r.ticket, 40))
                                  for i, r in enumerate(l.prestaciones, 1)]
        libs[(lcl, lib)] = (dict(descripcion=desc, pdls=gf.pdls_de(desc or ""), tickets=sorted({c for c in (gf.pdl_norm(t) for t in l.tickets) if c})[:60],
                                 estado_libreto=_t(l.estado_libreto.upper(), 12), estado_lcl=_t(l.estado_lcl.upper(), 20), fecha_autocert=_fecha_sonny(l.fecha_autocert),
                                 mo_pesos=_num(l.mo_pesos), puntos=_num(l.puntos), n_prest=int(l.n_prest or 0), sin_precio=int(l.sin_precio or 0), corte=corte)
                     # CHANGELOG 2026-10-23: las órdenes «Trabajo» del libreto (Orden de SAP), sin ceros a la izquierda; si el PC no las manda, se conservan
                     | ({"trabajos": sorted({t for t in (re.sub(r"\D", "", x or "").lstrip("0") for x in l.trabajos) if t})[:60]} if l.trabajos is not None else {}),
                     mats)
    claves = list(libs)
    ex = {(x.lcl, x.libreto): x for x in (await db.execute(select(GestionScmLibreto).where(tuple_(GestionScmLibreto.lcl, GestionScmLibreto.libreto).in_(claves))))
          .scalars().all()} if claves else {}
    mex: dict = {}
    if claves:
        for m in (await db.execute(select(GestionScmMat).where(tuple_(GestionScmMat.lcl, GestionScmMat.libreto).in_(claves)))).scalars().all():
            mex.setdefault((m.lcl, m.libreto), {})[m.material] = m
    pex: dict = {}
    if prests:
        for r in (await db.execute(select(GestionScmPrest).where(tuple_(GestionScmPrest.lcl, GestionScmPrest.libreto).in_(list(prests))))).scalars().all():
            pex.setdefault((r.lcl, r.libreto), []).append(r)
    # CHANGELOG 2026-10-23: qué libretos traen un detalle de prestaciones distinto al guardado (se borra de una vez y se inserta el nuevo)
    p_cambia = {k for k, nuevo in prests.items()
                if sorted((r.n, r.prestacion, round(r.cantidad or 0, 4), r.precio, round(r.pesos or 0, 2)) for r in pex.get(k, []))
                != sorted((r["n"], r["prestacion"], round(r["cantidad"], 4), r["precio"], r["pesos"]) for r in nuevo)}
    if p_cambia & set(pex):
        await db.execute(delete(GestionScmPrest).where(tuple_(GestionScmPrest.lcl, GestionScmPrest.libreto).in_(sorted(p_cambia & set(pex))))
                         .execution_options(synchronize_session=False))
    nuevos = actualizados = iguales = mat_filas = prest_filas = 0
    for k, (v, mats) in libs.items():
        x = ex.get(k)
        cambio = False
        if x is None:
            db.add(GestionScmLibreto(lcl=k[0], libreto=k[1], carga_id=cid_l, **v))
            nuevos += 1
            cambio = True
        else:
            d = _dif(x, {c: val for c, val in v.items() if c != "corte"})
            if d:
                for c, val in d.items():
                    setattr(x, c, val)
                x.actualizado_en = _ahora()
                cambio = True
            x.corte = corte
            _sello(x, cid_l)
        viejos = mex.get(k, {})
        for cod, (pos, rec, uni) in mats.items():
            mat_filas += 1
            m = viejos.pop(cod, None)
            if m is None:
                db.add(GestionScmMat(lcl=k[0], libreto=k[1], material=cod, posada=round(pos, 3), recuperada=round(rec, 3), unidad=uni or None, carga_id=cid_m))
                cambio = True
            else:
                if (round(m.posada or 0, 3), round(m.recuperada or 0, 3)) != (round(pos, 3), round(rec, 3)):
                    m.posada, m.recuperada = round(pos, 3), round(rec, 3)
                    cambio = True
                if uni and m.unidad != uni:
                    m.unidad = uni
                _sello(m, cid_m)
        for m in viejos.values():              # material que ya no trae el exporte para este libreto
            await db.delete(m)
            cambio = True
        if k in prests:                        # CHANGELOG 2026-10-23: el detalle de prestaciones del libreto se reemplaza por el del exporte
            prest_filas += len(prests[k])
            if k in p_cambia:
                for r in prests[k]:
                    db.add(GestionScmPrest(lcl=k[0], libreto=k[1], carga_id=cid_p, **r))
                cambio = True
            else:
                for r in pex.get(k, []):
                    _sello(r, cid_p)
        if x is not None:
            if cambio:
                actualizados += 1
            else:
                iguales += 1
    await db.flush()
    _cache_cargas.clear()
    lcls = {k[0] for k in libs}
    enl = await _reenlazar(db, lcls)
    await db.flush()
    cods = await _codigos_de_lcls(db, lcls)
    casos = set((await db.execute(select(GestionProyecto.codigo).where(GestionProyecto.codigo.in_(sorted(lcls))))).scalars().all()) if lcls else set()
    n, pend = await _recalcular_codigos(db, cods, body.recalcular)
    await db.commit()
    return {"ok": True, "archivo": body.archivo, "corte": body.corte, "precios_sha": body.precios_sha[:12], "libretos": len(libs), "invalidos": invalidos,
            "nuevos": nuevos, "actualizados": actualizados, "ya_estaban": iguales, "materiales": mat_filas, "prestaciones": prest_filas, "lcls": len(lcls),
            "lcls_con_caso": len(casos), "cargas": {"libretos": cid_l, "prestaciones": cid_p, "materiales": cid_m},
            "enlace_vivos": dict(enl), "casos_recalculados": n, "recalculo_pendiente": pend}


def _scm_lcl_fila(f: ScmLclFilaIn) -> Optional[dict]:
    """CHANGELOG 2026-10-25: la fila del exporte de LCL tal como se guarda (sin personas). None si la LCL o el Trabajo no son válidos."""
    lcl, tr = gf.lcl_norm(f.lcl), gf.trabajo_norm(f.trabajo)
    if not lcl or not tr or not re.fullmatch(r"\s*\d{1,24}\s*", f.trabajo or ""):
        return None
    return dict(lcl=lcl, trabajo=tr, estado=_t((f.estado or "").upper(), 20), fecha_envio=_fecha_sonny(f.fecha_envio), inicio_validez=_fecha_sonny(f.inicio_validez),
                fin_validez=_fecha_sonny(f.fin_validez), importe=_num(f.importe), obra=f.obra, acumulativa=f.acumulativa, tam=_t(f.tam, 10))


@router.post("/agente/scm-lcl", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_scm_lcl(body: ScmLclIn, db: AsyncSession = Depends(get_db)):
    """CHANGELOG 2026-10-25 El exporte de LCL de SCM (lo sube cargar_scm_gestion.py en lotes con cada LCL completa): upsert por (LCL, Trabajo). Si el estado de
    la LCL cambia, estado_desde = el corte de este exporte (SCM no fecha el cambio). Un Trabajo que una LCL del lote ya no trae deja de estar
    vigente (nunca se borra). Una fila de un exporte más viejo que el guardado no lo pisa. En el último lote (fin=true) recalcula los casos de lo
    que cambió en esta carga: hasta 30 aquí mismo, más en segundo plano. Solo lectura de SCM; nada de esto toca la MO de SGS."""
    corte = _fecha_sonny(body.corte)
    if not corte:
        raise HTTPException(status_code=422, detail="corte inválido (AAAA-MM-DD HH:MM:SS)")
    cid = await _carga_valida(db, body.carga_id, "scm")
    filas, inval, dup = {}, 0, 0
    for f in body.filas:
        v = _scm_lcl_fila(f)
        if v is None:
            inval += 1
            continue
        k = (v["lcl"], v["trabajo"])
        dup += 1 if k in filas else 0
        filas[k] = v
    lcls = sorted({k[0] for k in filas})
    ex = {(x.lcl, x.trabajo): x for x in (await db.execute(select(GestionScmLcl).where(GestionScmLcl.lcl.in_(lcls)))).scalars().all()} if lcls else {}
    nuevas = actualizadas = iguales = viejas = cambia_estado = 0
    tocados: set = set()
    for k, v in filas.items():
        x = ex.get(k)
        if x is None:
            db.add(GestionScmLcl(carga_id=cid, corte=corte, estado_desde=corte, vigente=True, **v))
            nuevas += 1
            tocados.add(k[0])
            continue
        if x.corte and x.corte > corte:                 # un exporte más viejo no pisa lo guardado
            viejas += 1
            continue
        _sello(x, cid)
        x.corte = corte
        d = _dif(x, v)
        if not x.vigente:
            x.vigente = True
            d = d or {"vigente": True}
        if not d:
            iguales += 1
            continue
        if "estado" in d:
            x.estado_anterior, x.estado_desde = x.estado, corte
            cambia_estado += 1
        for c, val in d.items():
            setattr(x, c, val)
        x.actualizado_en = _ahora()
        actualizadas += 1
        tocados.add(k[0])
    # la LCL del lote vino completa: el Trabajo que ya no trae deja de estar vigente
    no_vig = 0
    for k, x in ex.items():
        if k not in filas and x.vigente and not (x.corte and x.corte > corte):
            x.vigente, x.actualizado_en = False, _ahora()
            no_vig += 1
            tocados.add(k[0])
    await db.flush()
    _cache_cargas.clear()
    out = {"ok": True, "archivo": body.archivo, "corte": body.corte, "carga_id": cid, "filas": len(body.filas), "invalidas": inval, "repetidas_en_lote": dup,
           "lcls": len(lcls), "nuevas": nuevas, "actualizadas": actualizadas, "ya_estaban": iguales, "de_un_exporte_anterior": viejas,
           "cambian_estado": cambia_estado, "no_vigentes": no_vig}
    # LCL con caso: la LCL es el código de un caso, o es la LCL de una planilla de un caso de mayor impacto (su ticket es el caso)
    tks: dict = {}
    if lcls:
        for l_, t_ in (await db.execute(select(GestionSonnyPlanilla.lcl, GestionSonnyPlanilla.ticket).where(GestionSonnyPlanilla.lcl.in_(lcls),
                                                                                                         GestionSonnyPlanilla.ticket.isnot(None)))).all():
            tks.setdefault(l_, set()).add(t_)
    cods_all = set(lcls) | {t for v in tks.values() for t in v}
    casos_l = set((await db.execute(select(GestionProyecto.codigo).where(GestionProyecto.codigo.in_(sorted(cods_all))))).scalars().all()) if cods_all else set()
    out["lcls_con_caso"] = sum(1 for l in lcls if l in casos_l or (tks.get(l, set()) & casos_l))
    if body.fin:
        if cid:
            carga = await db.get(GestionCarga, cid)
            desde = carga.abierta_en if carga is not None else None
            q = select(GestionScmLcl.lcl).where(GestionScmLcl.carga_id == cid)
            if desde is not None:
                q = q.where(or_(GestionScmLcl.cargado_en >= desde, GestionScmLcl.actualizado_en >= desde))
            tocados |= set((await db.execute(q.distinct())).scalars().all())
            if desde is not None:        # lo que dejó de estar vigente en esta carga (en lotes anteriores)
                tocados |= set((await db.execute(select(GestionScmLcl.lcl).where(GestionScmLcl.vigente.is_(False), GestionScmLcl.actualizado_en >= desde).distinct())).scalars().all())
        cods = await _codigos_de_lcls(db, tocados)
        n, fondo = await _recalcular_sonny(db, cods, body.recalcular)
        out.update(lcls_tocadas=len(tocados), casos_recalculados=n, recalculo_en_fondo=fondo)
    await db.commit()
    return out


# ── DRAPE: el portal de facturación de Enel (lado proveedor) ─────────────────────────────────────────────────────────────────
def _fecha_drape(s: str) -> Optional[datetime]:
    """«22/09/2026» o «2026-09-22[ HH:MM:SS]» (hora Colombia) → datetime con zona; «01/01/1900» (DRAPE lo usa como vacío) → None."""
    s = (s or "").strip()
    m = re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})(?:\s+(\d{2}):(\d{2})(?::(\d{2}))?)?", s)
    if m:
        if m.group(3) < "2000":
            return None
        try:
            return datetime(int(m.group(3)), int(m.group(2)), int(m.group(1)), int(m.group(4) or 0), int(m.group(5) or 0), int(m.group(6) or 0), tzinfo=_TZ)
        except ValueError:
            return None
    d = _fecha_sonny(s)
    return d if d is None or d.year >= 2000 else None


def _digitos(xs, n: int = 20) -> list[str]:
    out = []
    for x in xs or []:
        d = re.sub(r"\D", "", str(x or ""))[:n]
        if d and d not in out:
            out.append(d)
    return out


def _drape_id(s: str) -> Optional[str]:
    d = re.sub(r"\s+", "", s or "").upper()
    return d if re.fullmatch(r"[A-Z0-9][A-Z0-9\-]{3,29}", d) else None


def _drape_fila(f: DrapeFacturaIn) -> Optional[dict]:
    """CHANGELOG 2026-10-25: la factura tal como se guarda (sin textos libres, emisor ni códigos fiscales). None si el DRAPE ID no es válido."""
    did = _drape_id(f.drape_id)
    if not did:
        return None
    pedidos = _digitos(f.pedidos)
    return dict(drape_id=did, factura=_t((f.factura or "").upper(), 30), estado=f.estado, destinatario=_t(f.destinatario, 80),
                f_factura=_fecha_drape(f.f_factura), n_registro=_t(re.sub(r"\D", "", f.n_registro or ""), 20), f_registro=_fecha_drape(f.f_registro),
                pedidos=pedidos, lcls=[l for l in dict.fromkeys(gf.lcl_norm(p) for p in pedidos) if l], ems=_digitos(f.ems),
                f_pago=_fecha_drape(f.f_pago), total=_num(f.total), moneda=_t((f.moneda or "").upper(), 5), f_envio=_fecha_drape(f.f_envio),
                f_creacion=_fecha_drape(f.f_creacion))


async def _codigos_de_lcls_lotes(db: AsyncSession, lcls: set) -> set:
    """CHANGELOG 2026-10-25: como _codigos_de_lcls pero solo con los códigos que son caso, por lotes (DRAPE nombra decenas de miles de LCL de todos los contratos:
    un IN con más de ~32 mil valores no cabe en una consulta)."""
    ls_ = sorted(c for c in lcls if c)
    out: set = set()
    for i in range(0, len(ls_), 5000):
        lote = ls_[i:i + 5000]
        out |= set((await db.execute(select(GestionProyecto.codigo).where(GestionProyecto.codigo.in_(lote)))).scalars().all())
        tks = (await db.execute(select(GestionSonnyPlanilla.ticket).where(GestionSonnyPlanilla.lcl.in_(lote), GestionSonnyPlanilla.ticket.isnot(None)))).scalars().all()
        out |= {t for t in tks if t}
    return out


@router.post("/agente/drape", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_drape(body: DrapeIn, db: AsyncSession = Depends(get_db)):
    """CHANGELOG 2026-10-25 DRAPE (datos cargados desde el PC): las facturas de las listas (upsert por DRAPE ID; un cambio de lista —enviada → registrada →
    pagada— actualiza la fila y guarda la lista anterior), el detalle de cada factura (vencimiento, documento de pago, condiciones, totales,
    eventos y el Resumen de Posición, upsert por DRAPE ID + posición) y las EM pendientes de facturar (upsert por EM). Una fila de una carga más
    vieja no pisa la guardada. En el último lote (fin=true), las facturas de las listas que vinieron enteras (listas_completas) que esta carga no
    trajo dejan de estar vigentes (nunca se borran) y se recalculan los casos de las LCL tocadas: hasta 30 aquí mismo, más en segundo plano.
    Solo lectura de DRAPE; nada toca la MO de SGS."""
    corte = _fecha_sonny(body.corte)
    if not corte:
        raise HTTPException(status_code=422, detail="corte inválido (AAAA-MM-DD HH:MM:SS)")
    cid = await _carga_valida(db, body.carga_id, "drape")
    ahora = _ahora()
    filas, inval, dup = {}, 0, 0
    for f in body.facturas:
        v = _drape_fila(f)
        if v is None:
            inval += 1
            continue
        dup += 1 if v["drape_id"] in filas else 0
        filas[v["drape_id"]] = v
    ex = {x.drape_id: x for x in (await db.execute(select(GestionDrapeFactura).where(GestionDrapeFactura.drape_id.in_(sorted(filas))))).scalars().all()} if filas else {}
    nuevas = actualizadas = iguales = viejas = cambia_estado = 0
    tocadas: set = set()
    for did, v in filas.items():
        x = ex.get(did)
        if x is None:
            db.add(GestionDrapeFactura(carga_id=cid, corte=corte, estado_desde=corte, vigente=True, **v))
            nuevas += 1
            tocadas |= set(v["lcls"])
            continue
        if x.corte and x.corte > corte:                 # una lectura más vieja no pisa lo guardado
            viejas += 1
            continue
        _sello(x, cid)
        x.corte = corte
        d = _dif(x, v)
        if not x.vigente:
            x.vigente = True
            d = d or {"vigente": True}
        if not d:
            iguales += 1
            continue
        tocadas |= set(x.lcls or []) | set(v["lcls"])
        if "estado" in d:
            x.estado_anterior, x.estado_desde = x.estado, corte
            cambia_estado += 1
        for c, val in d.items():
            setattr(x, c, val)
        x.actualizado_en = ahora
        actualizadas += 1
    await db.flush()
    # detalle de las facturas abiertas: cabecera (sin banco), eventos y Resumen de Posición
    det_ok = det_sin = pos_nuevas = pos_act = pos_iguales = pos_borradas = 0
    if body.detalles:
        dids = sorted({d_ for d_ in (_drape_id(dd.drape_id) for dd in body.detalles) if d_})
        exf = {x.drape_id: x for x in (await db.execute(select(GestionDrapeFactura).where(GestionDrapeFactura.drape_id.in_(dids))))
               .scalars().all()} if dids else {}
        exp: dict = {}
        if dids:
            for p_ in (await db.execute(select(GestionDrapePos).where(GestionDrapePos.drape_id.in_(dids)))).scalars().all():
                exp[(p_.drape_id, p_.posicion)] = p_
        for dd in body.detalles:
            did = _drape_id(dd.drape_id)
            x = exf.get(did) if did else None
            if x is None:                              # el detalle de una factura que ninguna lista trajo: no se guarda
                det_sin += 1
                continue
            if x.detalle_en and x.detalle_en > corte:
                continue
            pos_l = [l for l in dict.fromkeys(gf.lcl_norm(c) for c in dd.pos_lcls) if l]
            upd = dict(f_vence=_fecha_drape(dd.f_vence), doc_pago=_t(re.sub(r"[^\w\-]", "", dd.doc_pago or ""), 30), condiciones=_t(dd.condiciones, 40),
                       total_neto=None if dd.total_neto is None else _num(dd.total_neto), total_impuesto=None if dd.total_impuesto is None else _num(dd.total_impuesto),
                       n_pos=dd.n_pos, suma_pos=None if dd.suma_pos is None else _num(dd.suma_pos), pos_lcls=pos_l, estado_detalle=_t(dd.estado_detalle, 40),
                       eventos=[[re.sub(r"\D", "", str(e[1] if len(e) > 1 else ""))[:4], (_fecha_drape(e[0]) or _fecha_sonny(e[0]) or corte).date().isoformat()]
                                for e in dd.eventos if e][:40],
                       detalle_en=corte, carga_det_id=cid)
            fp = _fecha_drape(dd.f_pago)
            if fp and not x.f_pago:
                upd["f_pago"] = fp
            dx = _dif(x, upd)
            for c, val in dx.items():
                setattr(x, c, val)
            if dx:
                x.actualizado_en = ahora
            det_ok += 1
            tocadas |= set(x.lcls or [])
            vistos = set()
            for pp in dd.posiciones:
                pos_ = _t(re.sub(r"\s+", "", pp.posicion or ""), 12)
                if not pos_:
                    continue
                lcl_ = gf.lcl_norm((pp.pedido_pos or "").split("/")[0])
                if pos_l and lcl_ not in pos_l:        # solo las LCL cuyas posiciones vienen completas
                    continue
                vistos.add(pos_)
                nv = dict(factura=x.factura, lcl=lcl_, pedido_pos=_t(pp.pedido_pos, 40), producto=_t(pp.producto, 40), descripcion=_t(pp.descripcion, 200),
                          unidad=_t(pp.unidad, 10), cantidad=_num(pp.cantidad), precio=_num(pp.precio), impuesto=_t(pp.impuesto, 12), importe=_num(pp.importe))
                y = exp.get((did, pos_))
                if y is None:
                    db.add(GestionDrapePos(drape_id=did, posicion=pos_, carga_id=cid, **nv))
                    pos_nuevas += 1
                    continue
                _sello(y, cid)
                dy = _dif(y, nv)
                if not dy:
                    pos_iguales += 1
                    continue
                for c, val in dy.items():
                    setattr(y, c, val)
                y.actualizado_en = ahora
                pos_act += 1
            # las posiciones guardadas de estas LCL que la lectura nueva ya no trae (el detalle viene completo para pos_lcls) se quitan
            for (d2, p2), y in exp.items():
                if d2 == did and p2 not in vistos and (not pos_l or y.lcl in pos_l):
                    await db.delete(y)
                    pos_borradas += 1
    # EM pendientes de facturar (si DRAPE trae la lista)
    em_nuevas = em_act = em_iguales = 0
    if body.ems:
        ems_ = {}
        for e in body.ems:
            k = re.sub(r"\D", "", e.em or "")[:20]
            if k:
                ems_[k] = dict(fecha=_fecha_drape(e.fecha), importe=_num(e.importe), lcl=gf.lcl_norm(e.lcl), empresa=_t(e.empresa, 80), moneda=_t((e.moneda or "").upper(), 5))
        exe = {x.em: x for x in (await db.execute(select(GestionDrapeEm).where(GestionDrapeEm.em.in_(sorted(ems_))))).scalars().all()}
        for k, nv in ems_.items():
            y = exe.get(k)
            if y is None:
                db.add(GestionDrapeEm(em=k, corte=corte, vigente=True, carga_id=cid, **nv))
                em_nuevas += 1
            else:
                if y.corte and y.corte > corte:
                    continue
                _sello(y, cid)
                y.corte = corte
                dy = _dif(y, nv) | ({"vigente": True} if not y.vigente else {})
                if not dy:
                    em_iguales += 1
                    continue
                for c, val in dy.items():
                    setattr(y, c, val)
                y.actualizado_en = ahora
                em_act += 1
            if nv["lcl"]:
                tocadas.add(nv["lcl"])
    await db.flush()
    _cache_cargas.clear()
    out = {"ok": True, "archivo": body.archivo, "corte": body.corte, "carga_id": cid, "facturas": len(body.facturas), "invalidas": inval, "repetidas_en_lote": dup,
           "nuevas": nuevas, "actualizadas": actualizadas, "ya_estaban": iguales, "de_una_lectura_anterior": viejas, "cambian_de_lista": cambia_estado,
           "detalles": det_ok, "detalles_sin_factura": det_sin, "posiciones_nuevas": pos_nuevas, "posiciones_actualizadas": pos_act,
           "posiciones_ya_estaban": pos_iguales, "posiciones_quitadas": pos_borradas, "em_nuevas": em_nuevas, "em_actualizadas": em_act, "em_ya_estaban": em_iguales,
           "lcls": len({l for v in filas.values() for l in v["lcls"]})}
    lc_ = {l for v in filas.values() for l in v["lcls"]}
    out["lcls_con_caso"] = len(lc_ & await _codigos_de_lcls_lotes(db, lc_)) if lc_ else 0
    if body.fin:
        no_vig = 0
        listas = [x for x in body.listas_completas if x in gf.DRAPE_LISTAS]
        barrer = {d_ for d_ in (_drape_id(b) for b in body.barrer or []) if d_} if body.barrer is not None else None
        if listas and cid:
            q = select(GestionDrapeFactura).where(GestionDrapeFactura.estado.in_(listas), GestionDrapeFactura.vigente.is_(True),
                                                  or_(GestionDrapeFactura.carga_id != cid, GestionDrapeFactura.carga_id.is_(None)))
            if barrer is not None:
                q = q.where(GestionDrapeFactura.drape_id.in_(sorted(barrer) or ["-"]))
            for x in (await db.execute(q)).scalars().all():
                if x.corte and x.corte > corte:
                    continue
                x.vigente, x.actualizado_en = False, ahora
                no_vig += 1
                tocadas |= set(x.lcls or [])
        em_no = 0
        if "em" in body.listas_completas and cid:
            for y in (await db.execute(select(GestionDrapeEm).where(GestionDrapeEm.vigente.is_(True), or_(GestionDrapeEm.carga_id != cid,
                                                                                                             GestionDrapeEm.carga_id.is_(None))))).scalars().all():
                if barrer is not None and y.em not in barrer:
                    continue
                y.vigente, y.actualizado_en = False, ahora
                em_no += 1
                if y.lcl:
                    tocadas.add(y.lcl)
        if cid:
            carga = await db.get(GestionCarga, cid)
            desde = carga.abierta_en if carga is not None else None
            if desde is not None:
                for ls_ in (await db.execute(select(GestionDrapeFactura.lcls).where(or_(GestionDrapeFactura.carga_id == cid, GestionDrapeFactura.carga_det_id == cid),
                                                                                    or_(GestionDrapeFactura.cargado_en >= desde, GestionDrapeFactura.actualizado_en >= desde)))).scalars().all():
                    tocadas |= set(ls_ or [])
        await db.flush()
        cods = await _codigos_de_lcls_lotes(db, tocadas)
        n, fondo = await _recalcular_sonny(db, cods, body.recalcular)
        out.update(no_vigentes=no_vig, em_no_vigentes=em_no, lcls_tocadas=len(tocadas), casos_tocados=len(cods), casos_recalculados=n, recalculo_en_fondo=fondo)
    await db.commit()
    _cache_cargas.clear()
    return out


@router.post("/agente/deuda", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_deuda(body: DeudaIn, db: AsyncSession = Depends(get_db)):
    """Informe de deuda de facturación (CHANGELOG 2026-10-23, hoja Data): solo códigos, montos y estado. Upsert por (corte, ítem)."""
    corte = _fecha_sonny(body.corte)
    if not corte:
        raise HTTPException(status_code=422, detail="corte inválido (AAAA-MM-DD)")
    filas, invalidas = {}, 0
    for f in body.filas:
        item = (f.item or "").strip()
        if not item:
            invalidas += 1
            continue
        h = _sha(f"{corte.date().isoformat()}|{item}")
        filas[h] = dict(corte=corte, lcl=gf.lcl_norm(f.lcl), qr=_t(f.qr, 100), id_planilla=_t(f.id_planilla, 40), pdl_ticket=gf.pdl_norm(f.pdl_ticket) or _t(f.pdl_ticket, 80),
                        libreto=gf.libreto_norm(f.libreto), mo=_num(f.mo), mat=_num(f.mat), estado=_t(f.estado, 80), estado_desc=_t(f.estado_desc, 80),
                        responsabilidad=_t(f.responsabilidad, 40), fecha_estado=_fecha_sonny(f.fecha_estado))
    ex = {x.huella: x for x in (await db.execute(select(GestionDeuda).where(GestionDeuda.huella.in_(list(filas))))).scalars().all()} if filas else {}
    cid = await _carga_valida(db, body.carga_id, "deuda")
    nuevas = actualizadas = iguales = 0
    cods: set = set()
    for h, v in filas.items():
        x = ex.get(h)
        if x is None:
            db.add(GestionDeuda(huella=h, carga_id=cid, **v))
            nuevas += 1
        else:
            _sello(x, cid)
            d = _dif(x, v)
            if not d:
                iguales += 1
                continue
            for c, val in d.items():
                setattr(x, c, val)
            actualizadas += 1
        cods |= {v["lcl"], v["pdl_ticket"]}
    await db.flush()
    _cache_cargas.clear()
    n, pend = await _recalcular_codigos(db, {c for c in cods if _cod_valido(c)}, body.recalcular)
    await db.commit()
    return {"ok": True, "archivo": body.archivo, "corte": body.corte, "filas": len(filas), "invalidas": invalidas, "nuevas": nuevas, "actualizadas": actualizadas,
            "ya_estaban": iguales, "con_lcl": sum(1 for v in filas.values() if v["lcl"]), "casos_recalculados": n, "recalculo_pendiente": pend}


# ── SAP (MB51 de Enel): solo material ─────────────────────────────────────────────────────────────────────────────────
def _fecha_sap(s: str):
    try:
        return datetime.strptime((s or "")[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _sap_fila(f: SapFilaIn) -> dict:
    """Normaliza una fila del MB51: LCL o DF desde «Carta Entrega Trabajos», orden sin ceros, lote → zona y textos filtrados."""
    carta = re.sub(r"\s+", "", f.carta or "").upper()
    lcl = gf.lcl_norm(carta) if re.fullmatch(r"63\d{8,9}", carta) else None
    df = carta if re.fullmatch(r"DF20\d{10}", carta) else None
    orden = lambda x: (re.sub(r"\D", "", x or "").lstrip("0") or None) if re.fullmatch(r"\s*\d{1,20}\s*", x or "") else None   # noqa: E731
    lote = _t(f.lote, 20)
    return dict(clase=f.clase.strip()[:4], material=re.sub(r"\s+", "", f.material or "")[:20], descripcion=_t(f.descripcion, 120), cantidad=round(float(f.cantidad or 0), 3),
                um=_t(f.um, 10), importe=round(float(f.importe or 0), 2), f_contab=_fecha_sap(f.f_contab), lote=lote, zona=gf.ZONA_LOTE.get((lote or "")[:2]),
                centro=_t(f.centro, 10), almacen=_t(f.almacen, 10), centro_rec=_t(f.centro_rec, 10), almacen_rec=_t(f.almacen_rec, 10), lote_rec=_t(f.lote_rec, 20),
                lcl=lcl, df=df, orden=orden(f.orden), orden_2=orden(f.orden_2), pedido=_t(re.sub(r"\D", "", f.pedido or ""), 20),
                referencia=_t(re.sub(r"\D", "", f.referencia or ""), 20), cuenta=_t(f.cuenta if re.fullmatch(r"[A-Z0-9]{1,20}", (f.cuenta or "").strip()) else "", 20),
                texto_cab=gf.sap_texto_seguro(f.texto_cab), texto=gf.sap_texto_seguro(f.texto, posicion=True), texto_clase=_t(f.texto_clase, 40),
                texto_tipo=gf.sap_tipo_texto(f.texto_cab, f.texto) or None, dh=(f.dh or "").strip().upper()[:1] or None)


async def _recalcular_precios(db: AsyncSession, cid: Optional[int]) -> dict:
    """El precio de Enel por código: mediana de Importe ML / cantidad de las descargas 261 de TODO lo cargado (las filas son únicas por
    documento + posición). Se reemplaza la tabla entera: un código que ya no tiene 261 deja de tener precio de SAP."""
    meta = await _sap_meta(db)
    corte = meta.get("corte").date() if meta.get("corte") else None
    filas = await _sql(db, """select material, percentile_cont(0.5) within group (order by importe / cantidad), count(*), min(importe / cantidad),
                                     max(importe / cantidad), max(um), min(f_contab), max(f_contab),
                                     percentile_cont(0.1) within group (order by importe / cantidad), percentile_cont(0.9) within group (order by importe / cantidad),
                                     max(carga_id)
                              from sgo7_gg_sap where clase = '261' and cantidad < 0 and importe < 0 group by material""")
    ex = {x.material: x for x in (await db.execute(select(GestionPrecioMat))).scalars().all()}
    nuevos = cambian = 0
    vistos = set()
    for mat, med, n, mn, mx, um, desde, hasta, p10, p90, cid_mat in filas:
        vistos.add(mat)
        fuente = (f"MB51 de SAP (Enel): mediana de Importe ML / cantidad de {int(n)} descarga(s) 261 entre el {desde:%d/%m/%Y} y el {hasta:%d/%m/%Y}")[:300]
        # la carga de la que vienen sus descargas (la más nueva), no la que disparó el recálculo
        v = dict(precio=round(float(med), 2), um=um, n=int(n), fuente=fuente, fecha=corte, desde=desde, hasta=hasta, carga_id=cid_mat or cid,
                 detalle={"min": round(float(mn), 2), "max": round(float(mx), 2), "p10": round(float(p10), 2), "p90": round(float(p90), 2)})
        x = ex.get(mat)
        if x is None:
            db.add(GestionPrecioMat(material=mat, actualizado_en=_ahora(), **v))
            nuevos += 1
            continue
        if round(float(x.precio or 0), 2) != v["precio"]:
            cambian += 1
        for c, val in v.items():
            setattr(x, c, val)
        x.actualizado_en = _ahora()
    quitados = [m for m in ex if m not in vistos]
    if quitados:
        await db.execute(delete(GestionPrecioMat).where(GestionPrecioMat.material.in_(quitados)))
    await db.flush()
    _cache_sap.pop("precios", None)
    return {"codigos": len(vistos), "nuevos": nuevos, "cambian": cambian, "quitados": len(quitados)}


async def _controles_sap(db: AsyncSession) -> dict:
    """Los 3 números de la carga de SAP (también en el panel «Fuentes y controles»):
    · LCL de SAP (con consumo) que están en SCM · pares (LCL, código) con SAP descargado = SCM posado (libretos vivos; códigos que SAP maneja;
      libreto autocertificado hasta el corte; LCL en los dos) · recuperado 965 − 966 = SCM «Cant. Recuperada» por (LCL, código)."""
    (l_sap, l_en_scm), = await _sql(db, """select count(distinct lcl), count(distinct lcl) filter (where lcl in (select lcl from sgo7_gg_scm_libreto))
                                           from sgo7_gg_sap where clase in ('261', '262', '281') and lcl is not null""")
    (pares, iguales, lcls), = await _sql(db, """
        with corte as (select max(f_contab) c from sgo7_gg_sap where clase = '261'),
        sap as (select coalesce(lcl, df) k, material, -sum(cantidad) q from sgo7_gg_sap where clase in ('261', '262', '281') and coalesce(lcl, df) is not null group by 1, 2),
        scm as (select m.lcl k, m.material, sum(m.posada) q, max(l.fecha_autocert) fa from sgo7_gg_scm_mat m join sgo7_gg_scm_libreto l on l.lcl = m.lcl and l.libreto = m.libreto
                where coalesce(l.estado_libreto, '') not in ('ANN', 'RIF') group by 1, 2),
        cat as (select distinct material from sgo7_gg_sap),
        comun as (select k from sap intersect select k from scm),
        p as (select coalesce(s.k, c.k) k, coalesce(s.material, c.material) material, coalesce(s.q, 0) sap, coalesce(c.q, 0) scm, c.fa
              from sap s full join scm c on c.k = s.k and c.material = s.material)
        select count(*), count(*) filter (where abs(sap - scm) < 0.000001), count(distinct k) from p
        where k in (select k from comun) and material in (select material from cat) and (scm > 0 or sap <> 0)
          and (fa is null or fa < (select c from corte) + interval '1 day')""")
    (r_pares, r_iguales), = await _sql(db, """
        with r1 as (select coalesce(lcl, df) k, material, sum(cantidad) q from sgo7_gg_sap where clase in ('965', '966') and coalesce(lcl, df) is not null group by 1, 2),
        r2 as (select lcl k, material, sum(recuperada) q from sgo7_gg_scm_mat where recuperada > 0 group by 1, 2)
        select count(*), count(*) filter (where abs(r1.q - r2.q) < 0.000001) from r1 join r2 using (k, material)""")
    pct = lambda a, b: round(100 * a / b, 1) if b else None   # noqa: E731
    return {"lcl_sap": int(l_sap or 0), "lcl_en_scm": int(l_en_scm or 0), "pct_lcl_en_scm": pct(l_en_scm or 0, l_sap or 0),
            "pares": int(pares or 0), "pares_iguales": int(iguales or 0), "pct_pares_iguales": pct(iguales or 0, pares or 0), "lcl_pares": int(lcls or 0),
            "rec_pares": int(r_pares or 0), "rec_iguales": int(r_iguales or 0), "pct_rec_iguales": pct(r_iguales or 0, r_pares or 0)}


async def _recalcular_codigos_fondo(codigos: set) -> None:
    """Recalcula en segundo plano los casos de estos códigos (con su propia sesión), de 40 en 40."""
    try:
        async with AsyncSessionLocal() as db:
            ps = (await db.execute(select(GestionProyecto).where(GestionProyecto.codigo.in_(sorted(codigos))).order_by(GestionProyecto.id))).scalars().all()
            regs = await _cm_registros()
            for i in range(0, len(ps), 40):
                for p in ps[i:i + 40]:
                    await _recalcular(db, p, regs)
                await db.commit()
            _cache_sap["fondo"] = {"fin": _ahora().isoformat(), "casos": len(ps)}
    except Exception as e:                              # queda dicho para quien consulte; no tumba el proceso
        _cache_sap["fondo"] = {"error": f"{type(e).__name__}: {str(e)[:160]}"}


@router.post("/agente/sap", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_sap(body: SapIn, db: AsyncSession = Depends(get_db)):
    """CHANGELOG 2026-10-23 El MB51 del SAP de Enel (lo sube cargar_sap_gestion.py en lotes). Upsert por (documento, posición): reenviar el mismo archivo no
    duplica. En el último lote (fin=true): rehace el precio de Enel por código, mide los controles de la carga y recalcula en segundo plano
    los casos de los LCL y DF que trae esta carga. Solo MATERIAL: nada de esto toca la MO."""
    cid = await _carga_valida(db, body.carga_id, "sap")
    filas, invalidas, dup = {}, 0, 0
    for f in body.filas:
        k = ((f.doc or "").strip()[:20], int(f.pos))
        v = _sap_fila(f)
        if not k[0] or not v["material"] or not v["f_contab"]:
            invalidas += 1
            continue
        dup += 1 if k in filas else 0
        filas[k] = v
    ex = {(x.doc, x.pos): x for x in (await db.execute(select(GestionSap).where(tuple_(GestionSap.doc, GestionSap.pos).in_(list(filas))))).scalars().all()} if filas else {}
    nuevas = actualizadas = iguales = 0
    for k, v in filas.items():
        x = ex.get(k)
        if x is None:
            db.add(GestionSap(doc=k[0], pos=k[1], carga_id=cid, **v))
            nuevas += 1
            continue
        _sello(x, cid)
        d = _dif(x, v)
        if not d:
            iguales += 1
            continue
        for c, val in d.items():
            setattr(x, c, val)
        x.actualizado_en = _ahora()
        actualizadas += 1
    await db.flush()
    _cache_sap.clear()
    _cache_cargas.clear()
    out = {"ok": True, "archivo": body.archivo, "filas": len(body.filas), "invalidas": invalidas, "repetidas_en_lote": dup, "nuevas": nuevas,
           "actualizadas": actualizadas, "ya_estaban": iguales, "carga_id": cid}
    if body.fin:
        precios = await _recalcular_precios(db, cid)
        meta = await _sap_meta(db)
        ctl = await _controles_sap(db)
        if cid:
            cods = {c for (a, b) in (await db.execute(select(GestionSap.lcl, GestionSap.df).where(GestionSap.carga_id == cid).distinct())).all() for c in (a, b) if c}
        else:
            cods = {c for v in filas.values() for c in (v["lcl"], v["df"]) if c}
        cods = await _codigos_de_lcls(db, {c for c in cods if not c.startswith("DF")}) | {c for c in cods if c.startswith("DF")}
        casos = set((await db.execute(select(GestionProyecto.codigo).where(GestionProyecto.codigo.in_(sorted(cods))))).scalars().all()) if cods else set()
        en_linea = 0
        if body.recalcular and casos and len(casos) <= 30:          # pocos casos: aquí mismo (la respuesta ya los trae recalculados)
            en_linea, _ = await _recalcular_codigos(db, casos)
        await db.commit()
        if body.recalcular and len(casos) > 30:                       # la carga completa: en segundo plano (cientos de casos)
            t = asyncio.create_task(_recalcular_codigos_fondo(casos))
            _TAREAS_SAP.add(t)
            t.add_done_callback(_TAREAS_SAP.discard)
        alerta = [k for k, umbral in (("pct_lcl_en_scm", 90), ("pct_pares_iguales", 90), ("pct_rec_iguales", 95)) if ctl.get(k) is not None and ctl[k] < umbral]
        out.update(precios=precios, controles=ctl, por_debajo=alerta, corte=meta["corte"].date().isoformat() if meta.get("corte") else None,
                   filas_tabla=meta["n"], casos_con_sap=len(casos), casos_recalculados=en_linea, recalculo_en_fondo=bool(body.recalcular and len(casos) > 30))
        return out
    await db.commit()
    return out


async def _candidatos_nacer(db: AsyncSession) -> dict:
    """Códigos de Sonny, PMS y SCM que no son caso: {tupla de códigos en orden de preferencia: {fuentes, fecha, supervisores}}.
    En mayor impacto el caso es el ticket (DF) y, si el DF no existe en SGS, el LCL."""
    casos = set((await db.execute(select(GestionProyecto.codigo).where(GestionProyecto.codigo.isnot(None)))).scalars().all())
    pares: dict = {}

    def add(cods, fuente, fecha, sup=None):
        cods = tuple(dict.fromkeys(c for c in cods if _cod_valido(c)))
        if not cods or any(c in casos for c in cods):
            return
        d = pares.setdefault(cods, {"fuentes": set(), "fecha": None, "supervisores": set()})
        d["fuentes"].add(fuente)
        if fecha and (d["fecha"] is None or fecha > d["fecha"]):
            d["fecha"] = fecha
        if sup:
            d["supervisores"].add(sup)

    ult = dict((await db.execute(select(GestionSonny.id_planilla, func.max(GestionSonny.fecha)).group_by(GestionSonny.id_planilla))).all())
    for x in (await db.execute(select(GestionSonnyPlanilla))).scalars().all():
        f = max([d for d in (ult.get(x.id_planilla), x.fecha_supervisor, x.fecha_planilla) if d], default=None)
        add([x.ticket, x.lcl], "sonny", f, x.supervisor)
    for x in (await db.execute(select(GestionPms))).scalars().all():
        add(list(x.dfs or []) + list(x.lcls or []), "pms", x.fecha_estado or x.request_date)
    for lcl, f in (await db.execute(select(GestionScmLibreto.lcl, func.max(GestionScmLibreto.fecha_autocert)).group_by(GestionScmLibreto.lcl))).all():
        add([lcl], "scm", f)
    return pares


@router.post("/agente/nacer", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_nacer(body: NacerIn, db: AsyncSession = Depends(get_db)):
    """CHANGELOG 2026-10-23: casos que nacen de Sonny, PMS o SCM (no del correo). Por lotes: toma hasta `limite` grupos de códigos sin caso (los de
    movimiento más reciente primero), pregunta al Centro MO si existen en SGS y crea el caso del padre (mismo elegir_padre: el
    supervisor de la planilla, si no el que no es línea viva). Lo que no existe en SGS o queda fuera del corte `desde` se anota en
    sgo7_gg_sin_caso (se revisa otra vez pasados 3 días). Repetir hasta `quedan` = 0."""
    pares = await _candidatos_nacer(db)
    if body.codigos is not None:
        solo = {gf.clave_de(c) for c in body.codigos}
        pares = {k: v for k, v in pares.items() if set(k) & solo}
    desde = _fecha_sonny(body.desde) if body.desde else None
    revisados = {x.codigo: x for x in (await db.execute(select(GestionSinCaso))).scalars().all()}
    fresco = _ahora() - timedelta(days=3)

    def resuelto(c):
        r = revisados.get(c)
        # «el caso nace con DF…»: el LCL de una planilla de mayor impacto no abre un segundo caso (en Gestión el caso es el ticket)
        return bool(r) and ((r.motivo or "").startswith("el caso nace con") or bool(r.revisado_en and r.revisado_en > fresco))
    pend = [(k, v) for k, v in pares.items() if not all(resuelto(c) for c in k)]
    pend.sort(key=lambda kv: kv[1]["fecha"] or datetime(2000, 1, 1, tzinfo=_TZ), reverse=True)
    con_mov = sum(1 for _, v in pend if desde is None or (v["fecha"] and v["fecha"] >= desde))
    resumen = {"grupos_sin_caso": len(pares), "pendientes": len(pend), "con_movimiento": con_mov,
               "por_fuente": dict(Counter(",".join(sorted(v["fuentes"])) for _, v in pend))}
    if body.simular:
        return {"ok": True, "simulado": True, **resumen}

    def anotar(cod, v, motivo):
        r = revisados.get(cod)
        if r is None:
            r = GestionSinCaso(codigo=cod, fuentes=sorted(v["fuentes"]), motivo=motivo[:80], fecha_mov=v["fecha"], revisado_en=_ahora())
            db.add(r)
            revisados[cod] = r
        else:
            r.fuentes, r.motivo, r.fecha_mov, r.revisado_en = sorted(set(r.fuentes or []) | v["fuentes"]), motivo[:80], v["fecha"], _ahora()

    fuera = 0
    lote = []
    for k, v in pend:
        if desde is not None and not (v["fecha"] and v["fecha"] >= desde):
            for c in k:
                anotar(c, v, f"sin movimiento desde {body.desde} (corte de la primera carga)")
            fuera += 1
            continue
        if len(lote) < body.limite:
            lote.append((k, v))
    existencia = await _cm_sgs([_cod_cm(c) for k, _ in lote for c in k], detalle=False) if lote else {}
    creados: list[GestionProyecto] = []
    por_origen: Counter = Counter()
    no_existe = 0
    claves = set((await db.execute(select(GestionProyecto.clave))).scalars().all())
    for k, v in lote:
        cod = next((c for c in k if (existencia.get(_cod_cm(c)) or {}).get("existe")), None)
        if not cod:
            for c in k:
                anotar(c, v, "no existe en SGS")
            no_existe += 1
            continue
        for c in k:
            if c != cod:
                anotar(c, v, f"el caso nace con {cod}")
        padres = (existencia.get(_cod_cm(cod)) or {}).get("padres") or []
        elegidos = {}
        for sup in (sorted(v["supervisores"]) or [""]):
            pa, como = gf.elegir_padre(padres, sup, "", None)
            if pa:
                elegidos[pa["code"]] = pa
        origen = next(f for f in ("sonny", "pms", "scm") if f in v["fuentes"])
        if all(code in claves for code in elegidos):        # el padre ya es un caso con otro código: no se repite, se anota
            for c in k:
                anotar(c, v, f"su padre en SGS ya es un caso ({', '.join(sorted(elegidos))})")
            continue
        for code, pa in elegidos.items():
            if code in claves:
                continue
            p = GestionProyecto(clave=code, codigo=cod, padre={kk: pa.get(kk) for kk in ("id", "code", "name", "order_number", "supervisor", "departamento", "zona", "codes")},
                                tipo=_tipo_de(cod, pa), nombre=(pa.get("name") or cod)[:220], hilo="", codigos={}, fijados={}, alertas=[], activo=True, n_eventos=0,
                                n_correos=0, supervisor=pa.get("supervisor") or None, zona=pa.get("zona"), origen=origen)
            db.add(p)
            claves.add(code)
            creados.append(p)
            por_origen[origen] += 1
    await db.flush()
    pendiente_recalculo = False
    if creados:
        try:
            await _refrescar_sgs(creados)
            regs_cm = await _cm_registros()
            for p in creados:
                await _recalcular(db, p, regs_cm)
        except HTTPException:
            pendiente_recalculo = True
    await db.commit()
    quedan = len(pend) - fuera - len(lote)
    return {"ok": True, **resumen, "revisados": len(lote), "creados": len(creados), "por_origen": dict(por_origen), "no_existe_en_sgs": no_existe,
            "fuera_de_corte": fuera, "quedan": max(0, quedan), "recalculo_pendiente": pendiente_recalculo,
            "casos": [dict(id=p.id, clave=p.clave, codigo=p.codigo, origen=p.origen) for p in creados][:60]}


# ── Consulta ─────────────────────────────────────────────────────────────────────────────────────────
def _proy_out(p: GestionProyecto, dinero_completo: bool = False) -> dict:
    sgs = p.sgs
    if sgs:
        sgs = {k: v for k, v in sgs.items() if k not in ("lotes", "mat", "mat_lcl")}
    din = p.dinero or {}
    ux = din.get("ux") or {}
    return dict(
        origen=p.origen or "correo", dinero_resumen=din.get("resumen"), dinero=(din if dinero_completo else None),
        estado_cobro=ux.get("estado"), que_falta=(ux.get("que_falta") or [])[:3], cadena_resumen=(ux.get("cadena") or {}).get("resumen"),   # CHANGELOG 2026-10-25
        id=p.id, clave=p.clave, codigo=p.codigo, padre=p.padre, tipo=p.tipo, nombre=p.nombre, hilo=p.hilo, zona=p.zona, supervisor=p.supervisor,
        paso_actual=p.paso_actual, max_paso=p.max_paso, pilar=p.pilar, estado_texto=p.estado_texto, fuente_paso=p.fuente_paso,
        paso_fecha=p.paso_fecha.isoformat() if p.paso_fecha else None, alertas=p.alertas or [],
        ultima_actividad=p.ultima_actividad.isoformat() if p.ultima_actividad else None, n_eventos=p.n_eventos, n_correos=p.n_correos,
        fuentes=p.fuentes or {}, sgs=sgs, facturacion=p.facturacion or [], codigos=p.codigos or {}, activo=p.activo, nota=p.nota,
        fijados=p.fijados or {}, creado_en=p.creado_en.isoformat() if p.creado_en else None,
        cerrado=p.cerrado, cerrado_en=p.cerrado_en.isoformat() if p.cerrado_en else None,          # CHANGELOG 2026-10-25: todas sus LCL en CONCLUSA (SCM)
    )


_SQL_LCL_GESTION = """with c as (select codigo from sgo7_gg_proyectos where activo and codigo is not null),
    l as (select codigo lcl from c where codigo like '63%' union select p.lcl from sgo7_gg_sonny_planilla p where p.ticket in (select codigo from c) and p.lcl is not null)"""


async def _cobro_resumen(db: AsyncSession) -> dict:
    """CHANGELOG 2026-10-25: totales de DRAPE de las facturas que nombran al menos una LCL con caso en Gestión (el total de cada factura entera: DRAPE no lo reparte
    por LCL sin el detalle). Caché de 60 s (cada carga la limpia)."""
    v = _cache_cargas.get("cobro")
    if v and time.time() - v[0] < 60:
        return v[1]
    vd = _drape_cfg()["vence_dias"]
    filas = await _sql(db, _SQL_LCL_GESTION + """, f as (select * from sgo7_gg_drape_factura x where vigente
                                and exists (select 1 from jsonb_array_elements_text(coalesce(x.lcls, '[]'::jsonb)) e where e in (select lcl from l)))
        select estado, count(*), coalesce(sum(total), 0), count(*) filter (where f_factura >= '2026-01-01'), coalesce(sum(total) filter (where f_factura >= '2026-01-01'), 0),
               count(*) filter (where f_pago >= '2026-01-01'), coalesce(sum(total) filter (where f_pago >= '2026-01-01'), 0),
               count(*) filter (where estado = 'registrada' and f_pago is null
                                and coalesce(f_vence, f_factura + make_interval(days => cast(:vd as int))) < date_trunc('day', now())),
               coalesce(sum(total) filter (where estado = 'registrada' and f_pago is null
                                and coalesce(f_vence, f_factura + make_interval(days => cast(:vd as int))) < date_trunc('day', now())), 0)
        from f group by estado""", vd=vd)
    por = {e: {"n": int(n), "total": float(t), "n_2026": int(n6), "total_2026": float(t6), "n_pago_2026": int(np6), "pagado_2026": float(tp6),
               "n_vencidas": int(nv), "vencido": float(tv)} for e, n, t, n6, t6, np6, tp6, nv, tv in filas}

    def g(e, k):
        return (por.get(e) or {}).get(k) or 0
    (n_all, t_all), = await _sql(db, "select count(*), coalesce(sum(total), 0) from sgo7_gg_drape_factura where vigente")
    out = {"hay": True, "corte": (_local((await _cargas(db)).get("drape")) or None), "por_estado": por, "vence_dias": vd,
           "facturado_2026": round(g("registrada", "total_2026") + g("pagada", "total_2026"), 2), "n_facturado_2026": g("registrada", "n_2026") + g("pagada", "n_2026"),
           "pagado_2026": round(sum(x["pagado_2026"] for x in por.values()), 2), "n_pagado_2026": sum(x["n_pago_2026"] for x in por.values()),
           "pendiente": round(g("registrada", "total"), 2), "n_pendiente": g("registrada", "n"), "vencido": round(g("registrada", "vencido"), 2),
           "n_vencido": g("registrada", "n_vencidas"), "rechazado_2026": round(g("rechazada", "total_2026"), 2), "n_rechazado_2026": g("rechazada", "n_2026"),
           "enviado": round(g("enviada", "total"), 2), "n_enviado": g("enviada", "n"), "todas": {"n": int(n_all), "total": float(t_all)},
           "regla": "facturas vigentes de DRAPE que nombran al menos una LCL con caso en Gestión; total de la factura entera"}
    if out["corte"]:
        out["corte"] = out["corte"].isoformat(timespec="minutes")
    _cache_cargas["cobro"] = (time.time(), out)
    return out


@router.get("/pasos")
async def pasos(user=Depends(require_gestion)):
    return {"pilares": gf.PILARES, "pasos": gf.PASOS, "actores": gf.ACTORES}


@router.get("/resumen")
async def resumen(db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    ps = (await db.execute(select(GestionProyecto).where(GestionProyecto.activo.is_(True)))).scalars().all()
    por_pilar = {1: 0, 2: 0, 3: 0, 0: 0}
    alertas: dict[str, int] = {}
    zonas: dict[str, int] = {}
    sups: dict[str, int] = {}
    por_origen: dict[str, int] = {}
    con_alerta = 0
    rechazos: dict[str, int] = {}                       # CHANGELOG 2026-10-25: trabajos por categoría del motivo del rechazo
    cierre = {"cerrados": 0, "abiertos": 0, "sin_scm": 0, "cerrados_con_alertas": 0, "cerrados_con_alertas_mat": 0}      # CHANGELOG 2026-10-25
    mat75 = {"mat_cobrado_sin_consumir", *gf.ALERTAS_SAP, *gf.ALERTAS_MAT73}
    for p in ps:
        if p.cerrado is True:
            cierre["cerrados"] += 1
            tps = {a.get("tipo") for a in p.alertas or []}
            cierre["cerrados_con_alertas"] += 1 if tps else 0
            cierre["cerrados_con_alertas_mat"] += 1 if tps & mat75 else 0
        elif p.cerrado is False:
            cierre["abiertos"] += 1
        else:
            cierre["sin_scm"] += 1
        por_origen[p.origen or "correo"] = por_origen.get(p.origen or "correo", 0) + 1
        por_pilar[p.pilar or 0] = por_pilar.get(p.pilar or 0, 0) + 1
        if p.alertas:
            con_alerta += 1
            for a in p.alertas:
                alertas[a["tipo"]] = alertas.get(a["tipo"], 0) + 1
            for k in {c for a in p.alertas for c in (a.get("motivo_cats") or [])}:      # CHANGELOG 2026-10-25
                rechazos[k] = rechazos.get(k, 0) + 1
        zonas[p.zona or "—"] = zonas.get(p.zona or "—", 0) + 1
        if p.supervisor:
            sups[p.supervisor] = sups.get(p.supervisor, 0) + 1
    ags = (await db.execute(select(GestionAgente))).scalars().all()
    car = await _cargas(db)
    cargas = {"pms": car.get("pms_cargado").isoformat() if car.get("pms_cargado") else None, "scm": car["scm"].isoformat() if car.get("scm") else None,
              "sonny": car.get("sonny_cargado").isoformat() if car.get("sonny_cargado") else None,
              "deuda": _local(car["deuda"]).date().isoformat() if car.get("deuda") else None,
              "sap": car["sap"].date().isoformat() if car.get("sap") else None,          # CHANGELOG 2026-10-23: «SAP al dd/mm»
              "scm_lcl": car["scm_lcl"].isoformat() if car.get("scm_lcl") else None}     # CHANGELOG 2026-10-25: el exporte de LCL de SCM
    if gf.MOSTRAR_DRAPE:                                                                     # última carga de DRAPE (solo con la bandera)
        cargas["drape"] = car["drape"].isoformat() if car.get("drape") else None
    cobro = await _cobro_resumen(db) if car.get("drape") else {"hay": False}          # CHANGELOG 2026-10-25: totales de DRAPE y casos por estado del cobro
    if cobro.get("hay"):
        ce: Counter = Counter()
        for p in ps:
            r = ((p.dinero or {}).get("resumen") or {}).get("cobro") or {}
            if r.get("estado"):
                ce[r["estado"]] += 1
            if r.get("vencida"):
                ce["vencida"] += 1
        cobro["casos"] = dict(ce)
    return dict(total=len(ps), por_pilar=por_pilar, con_alerta=con_alerta, alertas=alertas, zonas=zonas, por_origen=por_origen, cargas=cargas, cierre=cierre, cobro=cobro,
                rechazos=rechazos, motivos_rechazo=[{"k": k, "t": t} for k, t, _ in gf.CATEGORIAS_MOTIVO if gf.MOSTRAR_DRAPE or k != "sin_motivo_drape"],   # CHANGELOG 2026-10-25
                sin_correo=sum(n for k, n in por_origen.items() if k != "correo"),
                supervisores=sorted(sups.items(), key=lambda kv: -kv[1]),
                agentes=[dict(buzon=a.buzon, ultimo_scan=a.ultimo_scan.isoformat() if a.ultimo_scan else None,
                              ultimo_contacto=a.ultimo_contacto.isoformat() if a.ultimo_contacto else None,
                              cursor=a.cursor.isoformat() if a.cursor else None, correos_total=a.correos_total, version=a.version) for a in ags],
                # «con libretos aprobados»: lo último del caso es el paso 39 cumplido (con Sonny, «Libretos Rechazados» también es paso 39)
                terminados=sum(1 for p in ps if p.paso_actual == 39 and (p.estado_texto or "").endswith("· cumple")), sin_sgs=sum(1 for p in ps if not p.sgs))


@router.get("/proyectos")
async def listar(db: AsyncSession = Depends(get_db), user=Depends(require_gestion), q: Optional[str] = None, zona: Optional[str] = None,
                 supervisor: Optional[str] = None, pilar: Optional[int] = None, alerta: Optional[str] = None, tipo: Optional[str] = None,
                 origen: Optional[str] = Query(None, max_length=12), activo: bool = True, limit: int = Query(500, ge=1, le=5000), offset: int = Query(0, ge=0),
                 cerrado: Optional[str] = Query(None, pattern="^(si|no|sin)$"),
                 rechazo: Optional[str] = Query(None, max_length=30, pattern="^[a-z_]+$")):
    stmt = select(GestionProyecto).where(GestionProyecto.activo.is_(activo))
    if cerrado == "si":                                 # CHANGELOG 2026-10-25: cerrado en SCM (todas sus LCL en CONCLUSA)
        stmt = stmt.where(GestionProyecto.cerrado.is_(True))
    elif cerrado == "no":
        stmt = stmt.where(GestionProyecto.cerrado.is_(False))
    elif cerrado == "sin":
        stmt = stmt.where(GestionProyecto.cerrado.is_(None))
    if origen == "sin_correo":                          # CHANGELOG 2026-10-23: casos que nacieron de Sonny, PMS o SCM
        stmt = stmt.where(GestionProyecto.origen != "correo")
    elif origen:
        stmt = stmt.where(GestionProyecto.origen == origen)
    if q:
        like = f"%{q.strip()}%"
        stmt = stmt.where(or_(GestionProyecto.clave.ilike(like), GestionProyecto.codigo.ilike(like), GestionProyecto.nombre.ilike(like),
                              GestionProyecto.supervisor.ilike(like), cast(GestionProyecto.codigos, Text).ilike(like)))
    if zona:
        stmt = stmt.where(GestionProyecto.zona == zona)
    if supervisor:
        stmt = stmt.where(GestionProyecto.supervisor.ilike(f"%{supervisor}%"))
    if pilar:
        stmt = stmt.where(GestionProyecto.pilar == pilar)
    if tipo:
        stmt = stmt.where(GestionProyecto.tipo == tipo)
    if alerta:
        stmt = stmt.where(cast(GestionProyecto.alertas, Text).like(f'%"tipo": "{alerta}"%'))
    if rechazo:                                         # CHANGELOG 2026-10-25: alguna alerta de rechazo con esa categoría de motivo
        stmt = stmt.where(GestionProyecto.alertas.contains([{"motivo_cats": [rechazo]}]))
    stmt = stmt.order_by(desc(GestionProyecto.ultima_actividad).nulls_last()).offset(offset).limit(limit)
    ps = (await db.execute(stmt)).scalars().all()
    gemelos = dict((await db.execute(select(GestionProyecto.codigo, func.count()).where(GestionProyecto.codigo.isnot(None)).group_by(GestionProyecto.codigo))).all())
    out = []
    for p in ps:
        d = _proy_out(p)
        d["hermanos_n"] = max(0, (gemelos.get(p.codigo) or 1) - 1)
        out.append(d)
    return out


def _enlaces(p: GestionProyecto) -> list[dict]:
    k = p.codigo or p.clave
    out = []
    if p.facturacion:
        out.append(dict(label=f"Balance de Proyectos · {len(p.facturacion)} caso(s)", url="/facturacion"))
    out += [dict(label="Centro MO", url="/modulos/centro-mo"), dict(label="Control Mensual de Liquidación", url="/modulos/control"),
            dict(label="Liquidador Automático", url="/liquidador"), dict(label="Registros de SGO7", url=f"/registros?q={k}")]
    return out


@router.get("/proyectos/{pid}")
async def ver(pid: int, db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    p = await db.get(GestionProyecto, pid)
    if not p:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")
    evs = (await db.execute(select(GestionEvento).where(GestionEvento.proyecto_id == pid).order_by(GestionEvento.ts, GestionEvento.id))).scalars().all()
    d = _proy_out(p, dinero_completo=True)
    d["sgs_lotes"] = list(reversed((p.sgs or {}).get("lotes") or []))
    d["hermanos"] = [dict(id=h.id, clave=h.clave, tipo=h.tipo, zona=h.zona, supervisor=h.supervisor, paso_actual=h.paso_actual, estado_texto=h.estado_texto,
                          total_mo=(h.sgs or {}).get("total_mo"), n_correos=h.n_correos) for h in await _hermanos(db, p)]
    d["eventos"] = [dict(id=e.id, ts=e.ts.isoformat() if e.ts else None, fuente=e.fuente, paso=e.paso, cond=e.cond, actor=e.actor, prueba=e.prueba,
                         de=e.de, de_correo=e.de_correo, para=(e.para or "")[:160], asunto=e.asunto, texto=(e.texto or "")[:700], directo=e.directo,
                         adjuntos=e.adjuntos or [], tiene_html=bool(e.html), ref=e.ref, carpeta=e.carpeta, pilar=gf.pilar_de(e.paso)) for e in evs]
    d["enlaces"] = _enlaces(p)
    return d


@router.get("/eventos/{eid}/correo")
async def correo(eid: int, db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    e = await db.get(GestionEvento, eid)
    if not e:
        raise HTTPException(status_code=404, detail="Evento no encontrado")
    return dict(id=e.id, asunto=e.asunto, de=e.de, de_correo=e.de_correo, para=e.para, fecha=e.ts.isoformat() if e.ts else None,
                carpeta=e.carpeta, adjuntos=e.adjuntos or [], html=e.html or "", texto=e.texto or "", directo=e.directo)


@router.patch("/proyectos/{pid}")
async def editar(pid: int, body: ProyectoPatch, db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    p = await db.get(GestionProyecto, pid)
    if not p:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")
    fij = dict(p.fijados or {})
    for campo in ("zona", "supervisor", "tipo"):
        v = getattr(body, campo)
        if v is not None:
            setattr(p, campo, v.strip() or None)
            fij[campo] = user.email
    if body.nota is not None:
        p.nota = body.nota.strip() or None
    if body.activo is not None:
        p.activo = body.activo
    p.fijados = fij
    p.actualizado_en = _ahora()
    await db.commit()
    return _proy_out(p)


_RECALCULO: dict = {"en_curso": False}
_TAREAS: set = set()


async def recalcular_todos(db: AsyncSession, limite: int = 5000, desde_id: int = 0) -> dict:
    """Relee SGS (MO y material, por el Centro MO) y recalcula los casos de 40 en 40, con commit por lote (CHANGELOG 2026-10-23: con más de mil
    casos, la lectura de material hace que la vuelta completa dure minutos)."""
    ps = (await db.execute(select(GestionProyecto).where(GestionProyecto.id >= desde_id).order_by(GestionProyecto.id).limit(limite))).scalars().all()
    regs_cm = await _cm_registros()
    _RECALCULO.update(total=len(ps), hechos=0, inicio=_ahora().isoformat())
    for i in range(0, len(ps), 40):
        lote = ps[i:i + 40]
        await _refrescar_sgs(lote)
        for p in lote:
            await _recalcular(db, p, regs_cm)
        await db.commit()
        _RECALCULO["hechos"] = i + len(lote)
    return {"ok": True, "proyectos": len(ps), "con_sgs": sum(1 for p in ps if p.sgs)}


async def _recalcular_en_fondo(limite: int) -> None:
    _RECALCULO.update(en_curso=True, error=None)
    try:
        async with AsyncSessionLocal() as db:
            r = await recalcular_todos(db, limite)
        _RECALCULO.update(resultado=r)
    except Exception as e:                              # queda en el estado para quien consulte; no tumba el proceso
        _RECALCULO.update(error=str(e)[:200])
    finally:
        _RECALCULO.update(en_curso=False, fin=_ahora().isoformat())


@router.post("/recalcular", dependencies=[Depends(require_dev)])
async def recalcular_todo(db: AsyncSession = Depends(get_db), limite: int = Query(5000, ge=1, le=20000), fondo: bool = False):
    """Relee SGS y el Centro MO y recalcula todos los casos. Con fondo=true responde enseguida y sigue en segundo plano (CHANGELOG 2026-10-23)."""
    if fondo:
        if not _RECALCULO.get("en_curso"):
            t = asyncio.create_task(_recalcular_en_fondo(limite))
            _TAREAS.add(t)
            t.add_done_callback(_TAREAS.discard)
        total = (await db.execute(select(func.count()).select_from(GestionProyecto))).scalar()
        return {"ok": True, "en_fondo": True, "proyectos": min(total or 0, limite), "estado": {k: v for k, v in _RECALCULO.items() if k != "resultado"}}
    return await recalcular_todos(db, limite)


@router.post("/reclasificar", dependencies=[Depends(require_dev)])
async def reclasificar(db: AsyncSession = Depends(get_db), limite: int = Query(5000, ge=1, le=20000)):
    """Reaplica gestion_flujo.clasificar a los correos guardados (cuando cambian las reglas) y recalcula."""
    ps = (await db.execute(select(GestionProyecto).order_by(GestionProyecto.id).limit(limite))).scalars().all()
    regs_cm = await _cm_registros()
    cambiados = creados = borrados = 0
    for p in ps:
        evs = (await db.execute(select(GestionEvento).where(GestionEvento.proyecto_id == p.id, GestionEvento.fuente == "correo")
                                .order_by(GestionEvento.ts, GestionEvento.id))).scalars().all()
        grupos: dict[tuple, list] = {}
        for e in evs:
            k = ((e.de or "")[:30].lower(), e.ts.isoformat() if e.ts else "", re.sub(r"\W+", "", (e.texto or "")[:160]).lower())
            grupos.setdefault(k, []).append(e)
        huellas = set((await db.execute(select(GestionEvento.huella).where(GestionEvento.proyecto_id == p.id))).scalars().all())
        for lista in grupos.values():
            base = lista[0]
            m = dict(texto=base.texto or "", asunto=base.asunto or "", de=base.de or "", de_correo=base.de_correo or "", para=base.para or "",
                     cc="", adjuntos=base.adjuntos or [])
            cl = gf.clasificar(m)
            nuevos = cl or ([dict(paso=None, cond=None, prueba="Correo del hilo sin paso reconocido")] if base.directo else [])
            actor = gf.actor_de(m)
            for i, x in enumerate(nuevos):
                if i < len(lista):
                    e = lista[i]
                    if (e.paso, e.cond, e.prueba, e.actor) != (x["paso"], x["cond"], x["prueba"], actor):
                        e.paso, e.cond, e.prueba, e.actor = x["paso"], x["cond"], x["prueba"], actor
                        cambiados += 1
                else:
                    h = _sha(f"{base.huella}|re|{x['paso']}")
                    if h in huellas:
                        continue
                    db.add(GestionEvento(proyecto_id=p.id, ts=base.ts, fuente="correo", paso=x["paso"], cond=x["cond"], actor=actor, prueba=x["prueba"],
                                         de=base.de, de_correo=base.de_correo, para=base.para, asunto=base.asunto, texto=base.texto, html=None,
                                         directo=base.directo, entry_id=base.entry_id, carpeta=base.carpeta, buzon=base.buzon, adjuntos=base.adjuntos,
                                         huella=h, ref=base.ref))
                    huellas.add(h)
                    creados += 1
            sobrevive = lista[0] if nuevos else None
            for e in lista[len(nuevos):]:
                if sobrevive is not None and e.html and not sobrevive.html:
                    sobrevive.html = e.html
                await db.delete(e)
                borrados += 1
        await db.flush()
        await _recalcular(db, p, regs_cm)
    await db.commit()
    return {"ok": True, "proyectos": len(ps), "cambiados": cambiados, "creados": creados, "borrados": borrados}


@router.delete("/proyectos/{pid}", dependencies=[Depends(require_dev)])
async def borrar(pid: int, db: AsyncSession = Depends(get_db)):
    p = await db.get(GestionProyecto, pid)
    if not p:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")
    sonny_borradas = 0
    otras: dict = {}
    # Sus filas del Historial de Sonny (CHANGELOG 2026-10-23) y de las tablas del hilo del dinero (CHANGELOG 2026-10-23) se van con él, salvo que un gemelo comparta
    # el código (vuelven con la próxima carga). De PMS solo las solicitudes que no nombran otro código.
    if p.codigo and not (await db.execute(select(GestionProyecto.id).where(GestionProyecto.codigo == p.codigo, GestionProyecto.id != p.id).limit(1))).first():
        c = p.codigo
        ids_pl = list((await db.execute(select(GestionSonnyPlanilla.id_planilla).where(or_(GestionSonnyPlanilla.lcl == c, GestionSonnyPlanilla.ticket == c)))).scalars().all())
        otras["sonny_nodos"] = (await db.execute(delete(GestionSonnyNodo).where(GestionSonnyNodo.codigo == c))).rowcount or 0          # CHANGELOG 2026-10-25
        otras["sonny_act"] = ((await db.execute(delete(GestionSonnyPlanAct).where(GestionSonnyPlanAct.id_planilla.in_(ids_pl)))).rowcount or 0) if ids_pl else 0
        _cache_sonny.clear()
        r = await db.execute(delete(GestionSonny).where(or_(GestionSonny.codigo == c, GestionSonny.ticket == c)))
        sonny_borradas = r.rowcount or 0
        otras["sonny_planillas"] = (await db.execute(delete(GestionSonnyPlanilla).where(or_(GestionSonnyPlanilla.lcl == c, GestionSonnyPlanilla.ticket == c)))).rowcount or 0
        rcs = [x.request_code for x in (await db.execute(select(GestionPms).where(or_(GestionPms.lcls.contains([c]), GestionPms.dfs.contains([c]))))).scalars().all()
               if set((x.lcls or []) + (x.dfs or [])) <= {c}]
        if rcs:
            otras["pms_logs"] = (await db.execute(delete(GestionPmsLog).where(GestionPmsLog.request_code.in_(rcs)))).rowcount or 0
            otras["pms"] = (await db.execute(delete(GestionPms).where(GestionPms.request_code.in_(rcs)))).rowcount or 0
        otras["scm_mat"] = (await db.execute(delete(GestionScmMat).where(GestionScmMat.lcl == c))).rowcount or 0
        otras["scm_prest"] = (await db.execute(delete(GestionScmPrest).where(GestionScmPrest.lcl == c))).rowcount or 0          # CHANGELOG 2026-10-23
        otras["decisiones"] = (await db.execute(delete(GestionDecision).where(GestionDecision.lcl == c))).rowcount or 0       # CHANGELOG 2026-10-23
        otras["scm_libretos"] = (await db.execute(delete(GestionScmLibreto).where(GestionScmLibreto.lcl == c))).rowcount or 0
        otras["deuda"] = (await db.execute(delete(GestionDeuda).where(or_(GestionDeuda.lcl == c, GestionDeuda.pdl_ticket == c)))).rowcount or 0
        otras["sin_caso"] = (await db.execute(delete(GestionSinCaso).where(GestionSinCaso.codigo == c))).rowcount or 0
        otras["sap"] = (await db.execute(delete(GestionSap).where(or_(GestionSap.lcl == c, GestionSap.df == c)))).rowcount or 0     # CHANGELOG 2026-10-23
        otras["scm_lcl"] = (await db.execute(delete(GestionScmLcl).where(GestionScmLcl.lcl == c))).rowcount or 0                     # CHANGELOG 2026-10-25
        # CHANGELOG 2026-10-25 DRAPE: una factura nombra muchas LCL; solo se van las que nombran SOLO este código (con sus posiciones) y las posiciones y EM de esta LCL
        solas = [x.drape_id for x in (await db.execute(select(GestionDrapeFactura).where(GestionDrapeFactura.lcls.contains([c])))).scalars().all()
                 if set(x.pedidos or []) <= {c}]
        otras["drape_pos"] = (await db.execute(delete(GestionDrapePos).where(or_(GestionDrapePos.lcl == c, GestionDrapePos.drape_id.in_(solas or ["-"]))))).rowcount or 0
        otras["drape_facturas"] = ((await db.execute(delete(GestionDrapeFactura).where(GestionDrapeFactura.drape_id.in_(solas)))).rowcount or 0) if solas else 0
        otras["drape_em"] = (await db.execute(delete(GestionDrapeEm).where(GestionDrapeEm.lcl == c))).rowcount or 0
        _cache_sap.clear()
    await db.delete(p)
    await db.commit()
    _cache_cargas.clear()
    return {"ok": True, "sonny_borradas": sonny_borradas, "otras_borradas": otras}


# ── Cargas, controles, evidencia y decisiones ─────────────────────────────────────────────────────────────────────
def _ip(request: Request) -> Optional[str]:
    return (request.headers.get("x-real-ip") or (request.headers.get("x-forwarded-for") or "").split(",")[0].strip()
            or (request.client.host if request.client else None) or None)


def _rol(user) -> str:
    return getattr(user.rol, "value", str(user.rol))


def _puede_decidir(user) -> bool:
    return _rol(user) in ROLES_DECIDEN


@router.post("/agente/carga", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_carga(body: CargaIn, db: AsyncSession = Depends(get_db)):
    """CHANGELOG 2026-10-23: abre o cierra el registro de una carga del PC. Abrir devuelve el carga_id que llevan las filas; cerrar guarda lo que
    respondió el servidor y una foto de los controles automáticos (panel «Fuentes y controles»)."""
    ahora = _ahora()
    if body.accion == "abrir":
        if body.sistema not in _SISTEMAS or body.parte not in _SISTEMAS[body.sistema]:
            raise HTTPException(status_code=422, detail=f"sistema/parte desconocidos: {body.sistema}/{body.parte}")
        if not re.fullmatch(r"[0-9a-f]{64}", body.sha256 or ""):
            raise HTTPException(status_code=422, detail="sha256 inválido")
        if len(json.dumps(body.detalle, default=str)) > 8000:
            raise HTTPException(status_code=422, detail="detalle demasiado grande")
        c = (await db.execute(select(GestionCarga).where(GestionCarga.sistema == body.sistema, GestionCarga.parte == body.parte,
                                                          GestionCarga.sha256 == body.sha256))).scalar_one_or_none()
        repetida = c is not None
        if c is None:
            c = GestionCarga(sistema=body.sistema, parte=body.parte, sha256=body.sha256, primera_en=ahora)
            db.add(c)
        c.archivo, c.tamano = (body.archivo or "")[:300], body.tamano
        c.modificado, c.corte = _fecha_sonny(body.modificado), (_fecha_sonny(body.corte) if body.corte else None)
        c.usuario, c.equipo = (body.usuario or "")[:100] or None, (body.equipo or "")[:100] or None
        c.filas_archivo, c.duplicados, c.fuera_rango, c.detalle = body.filas_archivo, body.duplicados, body.fuera_rango, body.detalle or {}
        c.abierta_en, c.cerrada_en, c.resultado, c.filas_enviadas = ahora, None, None, 0
        await db.commit()
        _cache_cargas.clear()
        return {"ok": True, "carga_id": c.id, "repetida": repetida}
    c = await db.get(GestionCarga, body.carga_id or 0)
    if c is None:
        raise HTTPException(status_code=404, detail="Carga no encontrada")
    if len(json.dumps(body.resultado, default=str)) > 20000:
        raise HTTPException(status_code=422, detail="resultado demasiado grande")
    c.cerrada_en, c.filas_enviadas, c.resultado = ahora, body.filas_enviadas, body.resultado or {}
    await db.flush()
    _cache_cargas.clear()
    ctl = await _controles(db)
    c.controles = {"calculado_en": _local(ahora).isoformat(timespec="seconds"), "controles": ctl}
    await db.commit()
    return {"ok": True, "carga_id": c.id, "controles": [{k: x[k] for k in ("k", "titulo", "valor", "estado")} for x in ctl]}


def _nombre_seguro(n: str) -> str:
    """CHANGELOG 2026-10-25: el nombre del archivo subido, sin rutas ni caracteres raros (se guarda con él y se descarga con él)."""
    n = os.path.basename(str(n or "").replace("\\", "/")).strip()
    n = re.sub(r"[^\w .,()+\-]", "_", n)[:200].strip(" .")
    return n or "archivo"


def _visible_carga(c) -> bool:
    return gf.MOSTRAR_DRAPE or c.sistema != "drape"


def _escribir(ruta: str, datos: bytes) -> None:
    os.makedirs(os.path.dirname(ruta), mode=0o750, exist_ok=True)
    tmp = ruta + ".subiendo"
    with open(tmp, "wb") as fh:
        fh.write(datos)
    os.chmod(tmp, 0o640)
    os.replace(tmp, ruta)


def _leer(ruta: str) -> bytes:
    with open(ruta, "rb") as fh:
        return fh.read()


@router.post("/agente/carga/{cid}/archivo", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_carga_archivo(cid: int, sha256: str = Form(..., max_length=64), archivo: UploadFile = File(...), db: AsyncSession = Depends(get_db)):
    """CHANGELOG 2026-10-25: el archivo de una carga, desde el PC. Su sha256 tiene que ser el declarado (422 si no) y el de la carga (409 si no); se guarda en
    registros/gestion/cargas/<id>/<nombre> y, si ya estaba con la misma huella, no se duplica."""
    c = await db.get(GestionCarga, cid)
    if c is None:
        raise HTTPException(status_code=404, detail="Carga no encontrada")
    datos = await archivo.read(_MAX_ARCHIVO + 1)
    if len(datos) > _MAX_ARCHIVO:
        raise HTTPException(status_code=413, detail="archivo demasiado grande")
    h = hashlib.sha256(datos).hexdigest()
    if h != (sha256 or "").strip().lower():
        raise HTTPException(status_code=422, detail="el archivo recibido no tiene el sha256 declarado")
    if h != c.sha256:
        raise HTTPException(status_code=409, detail="el sha256 del archivo no es el de la carga: no se guarda")
    if c.archivo_ruta and c.archivo_sha256 == h and os.path.isfile(c.archivo_ruta):
        return {"ok": True, "ya_estaba": True, "carga_id": c.id, "nombre": c.archivo_nombre, "tamano": c.archivo_tamano}
    nombre = _nombre_seguro(archivo.filename or c.archivo or "archivo")
    ruta = os.path.join(_DIR_CARGAS, str(c.id), nombre)
    await asyncio.to_thread(_escribir, ruta, datos)
    c.archivo_ruta, c.archivo_nombre, c.archivo_sha256, c.archivo_tamano, c.archivo_en = ruta, nombre, h, len(datos), _ahora()
    await db.commit()
    _cache_cargas.clear()
    return {"ok": True, "ya_estaba": False, "carga_id": c.id, "nombre": nombre, "tamano": len(datos)}


@router.get("/agente/cargas", dependencies=[Depends(_agente)], include_in_schema=False)
async def agente_cargas(db: AsyncSession = Depends(get_db), pruebas: bool = False):
    """CHANGELOG 2026-10-25: las cargas registradas (para el relleno de archivos del PC): sha256, tamaño, si es la última de su sistema y parte y si ya tiene archivo."""
    cs = (await db.execute(select(GestionCarga).order_by(GestionCarga.abierta_en, GestionCarga.id))).scalars().all()
    ult = {}
    for c in cs:
        if not (c.archivo or "").startswith("prueba-"):
            ult[(c.sistema, c.parte)] = c.id
    return {"cargas": [dict(id=c.id, sistema=c.sistema, parte=c.parte, archivo=c.archivo, tamano=c.tamano, sha256=c.sha256, ultima=ult.get((c.sistema, c.parte)) == c.id,
                            tiene_archivo=bool(c.archivo_ruta and os.path.isfile(c.archivo_ruta)))
                       for c in cs if (pruebas or not (c.archivo or "").startswith("prueba-")) and _visible_carga(c)]}


@router.get("/cargas/{cid}/archivo")
async def descargar_carga(cid: int, request: Request, db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    """CHANGELOG 2026-10-25 «Descargar»: el archivo de una fuente tal como se cargó (comprobado contra su sha256). La descarga queda en Registros."""
    c = await db.get(GestionCarga, cid)
    if c is None or not _visible_carga(c) or not c.archivo_ruta or not os.path.isfile(c.archivo_ruta):
        raise HTTPException(status_code=404, detail="Esta fuente no tiene su archivo en el servidor")
    datos = await asyncio.to_thread(_leer, c.archivo_ruta)
    h = hashlib.sha256(datos).hexdigest()
    if h != c.archivo_sha256:
        raise HTTPException(status_code=409, detail="el archivo guardado no coincide con su huella: no se entrega")
    nombre = c.archivo_nombre or "archivo"
    reg = RegistroSGO7(ts=_ahora(), modulo="gestion", usuario=user.email, rol=_rol(user), accion="gestion_descarga_fuente", metodo="GET",
                       ruta=f"/api/gestion/cargas/{cid}/archivo", archivo=nombre[:300], archivo_hash=h, archivo_bytes=len(datos), estado="ok", ip=_ip(request),
                       detalle={"carga_id": c.id, "sistema": c.sistema, "parte": c.parte, "archivo_registrado": c.archivo,
                                "exporte": _local(c.modificado).isoformat() if c.modificado else None})
    db.add(reg)
    await db.commit()
    ascii_ = re.sub(r"[^A-Za-z0-9 ._()+-]", "_", nombre)
    tipo = {".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xls": "application/vnd.ms-excel", ".csv": "text/csv",
            ".json": "application/json", ".zip": "application/zip"}.get(os.path.splitext(nombre)[1].lower(), "application/octet-stream")
    return Response(content=datos, media_type=tipo, headers={"Content-Disposition": f"attachment; filename=\"{ascii_}\"; filename*=UTF-8''{quote(nombre)}",
                                                             "X-Registro-Id": str(reg.id), "Cache-Control": "no-store"})


async def _sql(db: AsyncSession, q: str, **kw):
    return (await db.execute(text(q), kw)).all()


def _m_txt(v) -> str:
    """CHANGELOG 2026-10-25: pesos legibles para los textos de los controles («$3.841,8 M»)."""
    v = float(v or 0)
    if abs(v) < 1e6:
        return "$" + f"{v:,.0f}".replace(",", ".")
    return "$" + f"{v / 1e6:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".") + " M"


async def _controles(db: AsyncSession) -> list[dict]:
    """Controles automáticos del hilo del dinero, calculados sobre todo lo cargado (cada carga guarda una foto al cerrarse).
    estado: verde (bien) · rojo (algo no reproduce o se rompió el cruce) · ambar (pendiente de una persona) · gris (sin datos)."""
    out = []

    def ctl(k, titulo, valor, estado, detalle="", **kw):
        out.append({"k": k, "titulo": titulo, "valor": valor, "estado": estado, "detalle": detalle, **kw})

    # 1. la suma de las filas reproduce cada cifra de la franja, en todos los casos
    (bien, mal, sin), = await _sql(db, """select count(*) filter (where dinero->'evidencia'->>'ok' = 'true'), count(*) filter (where dinero->'evidencia'->>'ok' = 'false'),
                                            count(*) filter (where dinero->'evidencia' is null) from sgo7_gg_proyectos where activo and dinero is not null""")
    malos = await _sql(db, """select id, coalesce(codigo, clave), dinero->'evidencia'->'malas' from sgo7_gg_proyectos
                              where activo and dinero->'evidencia'->>'ok' = 'false' order by id limit 40""")
    ctl("cajas_filas", "Suma de las filas = cifra de cada caja", f"{bien} de {bien + mal + sin} casos", "rojo" if mal else ("ambar" if sin else "verde"),
        (f"{mal} caso(s) donde alguna suma no reproduce su cifra" if mal else "todas las cajas se reproducen al peso con sus filas")
        + (f" · {sin} caso(s) sin verificar (recalcular)" if sin else ""),
        casos=[{"id": i, "codigo": c, "malas": m} for i, c, m in malos])
    # 2. Sonny → PMS: el QR de cada planilla es una solicitud de PMS (medido el 23 oct: 98 %)
    (con_qr, en_pms, tot), = await _sql(db, """select count(*) filter (where coalesce(qr,'') <> ''), count(*) filter (where qr in (select request_code from sgo7_gg_pms)),
                                                  count(*) from sgo7_gg_sonny_planilla""")
    pct = 100 * en_pms / con_qr if con_qr else None
    ctl("sonny_pms", "Planillas de Sonny cuyo QR está en PMS", f"{pct:.1f} % ({en_pms} de {con_qr})" if pct is not None else "—",
        "gris" if pct is None else ("verde" if pct >= 95 else "rojo"),
        "QR de la Lista de Planillas = request_code de PMS (llave exacta). Por debajo de 95 % algo cambió en un exporte.")
    # 3. planillas de la Lista que son de un caso (medido el 23 oct: 93,8 %)
    (pl_caso, pl_tot), = await _sql(db, """select count(*) filter (where lcl in (select codigo from sgo7_gg_proyectos where codigo is not null)
                                                   or ticket in (select codigo from sgo7_gg_proyectos where codigo is not null)), count(*) from sgo7_gg_sonny_planilla""")
    pct2 = 100 * pl_caso / pl_tot if pl_tot else None
    ctl("planillas_caso", "Planillas de Sonny con caso en Gestión", f"{pct2:.1f} % ({pl_caso} de {pl_tot})" if pct2 is not None else "—",
        "gris" if pct2 is None else ("verde" if pct2 >= 90 else "rojo"),
        "Por su Lcl o su ticket. Las que faltan son de códigos que no existen en SGS o sin código.")
    # 4. libretos vivos de SCM enlazados a su planilla, por confianza (alta / confirmada / media / ambigua / sin planilla), con su MO
    enl: dict = {}
    for regla, n, mo in await _sql(db, """select split_part(coalesce(regla_enlace, 'sin planilla'), ':', 1), count(*), coalesce(sum(mo_pesos), 0) from sgo7_gg_scm_libreto
                                         where coalesce(estado_libreto, '') not in ('ANN', 'RIF') group by 1"""):
        k = {"a": "alta", "b": "alta", "d": "media", "ambiguo": "ambigua", "confirmada": "confirmada"}.get(regla, "sin planilla")
        x = enl.setdefault(k, {"n": 0, "mo": 0.0})
        x["n"] += int(n)
        x["mo"] += float(mo)
    orden = ("alta", "confirmada", "media", "ambigua", "sin planilla")
    ctl("enlaces", "Libretos vivos enlazados a su planilla, por confianza",
        " · ".join(f"{k} {enl[k]['n']}" for k in orden if k in enl) or "—",
        "ambar" if (enl.get("ambigua") or enl.get("media")) else "verde",
        "Alta: Id de planilla o PDL exactos en la descripción del libreto (el PDL sin ceros de más tras la «R»). Media: misma MO (±1 %). Ambigua: varias planillas candidatas. "
        "Las media y ambiguas esperan «Es este» / «No es» de una persona.",
        por_confianza={k: {"n": enl[k]["n"], "mo": round(enl[k]["mo"], 2)} for k in orden if k in enl})
    # 5. duplicados por la llave natural con que se guarda cada exporte (0 por construcción; si sale otro número, algo se rompió)
    llaves = [("sgo7_gg_pms", "request_code"), ("sgo7_gg_pms_log", "huella"), ("sgo7_gg_sonny", "id_cambio"), ("sgo7_gg_sonny_planilla", "id_planilla"),
              ("sgo7_gg_scm_libreto", "lcl, libreto"), ("sgo7_gg_scm_mat", "lcl, libreto, material"), ("sgo7_gg_scm_prest", "lcl, libreto, n"), ("sgo7_gg_deuda", "huella"),
              ("sgo7_gg_sap", "doc, pos"), ("sgo7_gg_precio_mat", "material"), ("sgo7_gg_sonny_nodo", "id_nodo"), ("sgo7_gg_sonny_plan_act", "id_planilla"),
              ("sgo7_gg_scm_lcl", "lcl, trabajo"),                                                                                  # CHANGELOG 2026-10-25
              ("sgo7_gg_drape_factura", "drape_id"), ("sgo7_gg_drape_pos", "drape_id, posicion"), ("sgo7_gg_drape_em", "em")]       # CHANGELOG 2026-10-25
    dup = {}
    for t, k in llaves:
        (n,), = await _sql(db, f"select count(*) from (select 1 from {t} group by {k} having count(*) > 1) x")
        if n:
            dup[t] = int(n)
    (pms_sem,), = await _sql(db, "select count(*) from (select 1 from sgo7_gg_pms_log group by request_code, fecha, estado_ini, estado_fin having count(*) > 1) x")
    ult = await _ultimas_cargas(db)
    dup_arch = {f"{s}/{p}": c["duplicados"] for (s, p), c in ult.items() if c.get("duplicados")}
    ctl("duplicados", "Duplicados por llave natural", str(sum(dup.values())), "rojo" if dup else "verde",
        ("por tabla: " + ", ".join(f"{t} {n}" for t, n in dup.items()) if dup else "ninguna fila repetida por su llave en las tablas")
        + (" · en los archivos (resueltos al cargar): " + ", ".join(f"{k} {n}" for k, n in dup_arch.items()) if dup_arch else "")
        + (f" · {pms_sem} cambio(s) de PMS repetidos con otro sub-estado (no mueven dinero)" if pms_sem else ""))
    # 6. fechas fuera de rango (antes de 2020 o más de 2 días en el futuro)
    cols = [("sgo7_gg_pms", "request_date"), ("sgo7_gg_pms", "fecha_estado"), ("sgo7_gg_pms_log", "fecha"), ("sgo7_gg_sonny", "fecha"),
            ("sgo7_gg_sonny_planilla", "fecha_planilla"), ("sgo7_gg_sonny_planilla", "fecha_supervisor"), ("sgo7_gg_sonny_planilla", "fecha_ingeniero"),
            ("sgo7_gg_scm_libreto", "fecha_autocert"), ("sgo7_gg_deuda", "fecha_estado"), ("sgo7_gg_sap", "f_contab"), ("sgo7_gg_sonny_nodo", "f_estado"),
            ("sgo7_gg_sonny_plan_act", "f_ultima"), ("sgo7_gg_scm_lcl", "fecha_envio"),          # CHANGELOG 2026-10-25: el fin de validez es futuro por naturaleza
            ("sgo7_gg_drape_factura", "f_factura"), ("sgo7_gg_drape_factura", "f_pago")]          # CHANGELOG 2026-10-25: el vencimiento es futuro por naturaleza
    fuera = {}
    for t, c in cols:
        (n,), = await _sql(db, f"select count(*) from {t} where {c} < '2020-01-01' or {c} > now() + interval '2 days'")
        if n:
            fuera[f"{t}.{c}"] = int(n)
    ctl("fechas", "Filas con fecha fuera de rango", str(sum(fuera.values())), "rojo" if fuera else "verde",
        ", ".join(f"{k} {n}" for k, n in fuera.items()) if fuera else "todas las fechas entre 2020 y hoy")
    # 7. la lista de precios QP con que el PC pasó cantidades a pesos es la de producción
    _, info = _precios_qp()
    u_pre = ult.get(("scm", "prestaciones")) or {}
    sha_pc = (u_pre.get("detalle") or {}).get("precios_sha") or ""
    ctl("precios", "Precios QP del PC = lista de producción", ("sí" if sha_pc == info.get("sha256") else "no") if sha_pc and info else "—",
        "gris" if not (sha_pc and info) else ("verde" if sha_pc == info.get("sha256") else "rojo"),
        f"{info.get('archivo', '?')} de producción (huella {str(info.get('sha256', ''))[:8]}, del {str(info.get('modificado', ''))[:10]})"
        + (f" · el PC usó {sha_pc[:8]}" if sha_pc else " · la carga de prestaciones aún no dice con qué lista"))
    # 8. la MO de cada libreto = la suma de sus prestaciones (cantidad × precio QP)
    (mal_l, con_det), = await _sql(db, """select count(*) filter (where abs(coalesce(l.mo_pesos, 0) - p.s) >= 1), count(*) from sgo7_gg_scm_libreto l
                                          join (select lcl, libreto, sum(pesos) s from sgo7_gg_scm_prest group by 1, 2) p using (lcl, libreto)""")
    (sin_det,), = await _sql(db, """select count(*) from sgo7_gg_scm_libreto l where coalesce(l.n_prest, 0) > 0
                                    and not exists (select 1 from sgo7_gg_scm_prest p where p.lcl = l.lcl and p.libreto = l.libreto)""")
    ctl("prestaciones", "MO de cada libreto = suma de sus prestaciones", f"{con_det - mal_l} de {con_det}" + (f" · {sin_det} sin detalle" if sin_det else ""),
        "rojo" if mal_l else ("ambar" if sin_det else ("verde" if con_det else "gris")),
        (f"{mal_l} libreto(s) donde no cuadra" if mal_l else "cuadran al peso")
        + (f" · {sin_det} libreto(s) sin el detalle de prestaciones (carga anterior)" if sin_det else ""))
    # 9. CHANGELOG 2026-10-25 SGS: liquidado + por liquidar − consumido de más = despachado neto (bandejas aprobadas), por código y en pesos, en todos los casos;
    #    y de paso, en cuántos casos las bandejas no dan lo mismo que los campos «despachado − devuelto» de project.material (manda la bandeja)
    (n_ok, n_mal, n_sinb, n_tot, n_campos, v_liq, v_por, v_sob), = await _sql(db, """
        select count(*) filter (where c->>'mat_identidad_ok' = 'true'), count(*) filter (where c->>'mat_identidad_ok' = 'false'),
               count(*) filter (where coalesce(c->>'mat_bandejas', 'false') <> 'true'), count(*),
               count(*) filter (where coalesce((c->>'mat_campos_dif')::int, 0) > 0),
               coalesce(sum((c->>'mat')::numeric), 0), coalesce(sum((c->>'mat_por_liquidar')::numeric), 0), coalesce(sum((c->>'mat_sobre')::numeric), 0)
        from sgo7_gg_proyectos p, jsonb_array_elements(p.dinero->'cajas') c where p.activo and c->>'k' = 'sgs' and c->>'mat' is not null""")
    m_neto = await _sql(db, """select p.id, coalesce(p.codigo, p.clave), c->>'mat_identidad_ok', coalesce((c->>'mat_campos_dif')::int, 0)
                               from sgo7_gg_proyectos p, jsonb_array_elements(p.dinero->'cajas') c where p.activo and c->>'k' = 'sgs' and c->>'mat' is not null
                               and (c->>'mat_identidad_ok' = 'false' or coalesce(c->>'mat_bandejas', 'false') <> 'true' or coalesce((c->>'mat_campos_dif')::int, 0) > 0)
                               order by (c->>'mat_identidad_ok' = 'false') desc, p.id limit 40""")
    ctl("sgs_neto", "SGS: liquidado + por liquidar = despachado neto", f"{n_ok} de {n_tot} casos" if n_tot else "—",
        "gris" if not n_tot else ("rojo" if n_mal else ("ambar" if n_sinb else "verde")),
        "Por código, con despachado y devuelto de las bandejas aprobadas de SGS (BAN-): consumido + por liquidar − consumido de más = despachado neto, "
        f"en cantidades y en pesos. En todos los casos: liquidado {_m_txt(v_liq)} · por liquidar {_m_txt(v_por)} · consumido de más (aparte) {_m_txt(v_sob)}."
        + (f" · {n_mal} caso(s) donde NO cuadra" if n_mal else "") + (f" · {n_sinb} caso(s) sin las bandejas de SGS (recalcular)" if n_sinb else "")
        + (f" · en {n_campos} caso(s) las bandejas no dan lo mismo que los campos «despachado − devuelto» de project.material (manda la bandeja)" if n_campos else ""),
        casos=[{"id": i, "codigo": c, "identidad": ok_, "campos_dif": cd} for i, c, ok_, cd in m_neto])
    # 13-14. CHANGELOG 2026-10-25 Sonny por etapa: los nodos de la Lista Cruzada son de un caso de Gestión (por su LCL o su ticket) y las planillas que nombran
    # existen en la Lista de Planillas (una celda puede traer dos o más Ids separados por coma: se cuentan por separado)
    (n_v, n_c), = await _sql(db, """select count(*), count(*) filter (where codigo in (select codigo from sgo7_gg_proyectos where codigo is not null)
                                          or exists (select 1 from jsonb_array_elements_text(coalesce(codigos, '[]'::jsonb)) x
                                                     where x in (select codigo from sgo7_gg_proyectos where codigo is not null)))
                                   from sgo7_gg_sonny_nodo where vigente""")
    if n_v:
        pc = 100 * n_c / n_v
        (c_sin,), = await _sql(db, """select count(distinct codigo) from sgo7_gg_sonny_nodo where vigente
                                      and codigo not in (select codigo from sgo7_gg_proyectos where codigo is not null)""")
        ctl("nodos_caso", "Nodos de la Lista Cruzada con caso en Gestión", f"{pc:.1f} % ({n_c} de {n_v})", "verde" if pc >= 95 else "ambar",
            f"Por la columna Lcl (LCL o ticket DF). Los {n_v - n_c} que faltan son de {c_sin} código(s) que no son caso (no existen en SGS o aún no nacen). "
            "Meta 95 %.")
        (p_t, p_ok), = await _sql(db, """select count(*), count(*) filter (where p in (select id_planilla from sgo7_gg_sonny_planilla))
                                         from (select distinct jsonb_array_elements_text(id_planillas) p from sgo7_gg_sonny_nodo where vigente) x""")
        pc2 = 100 * p_ok / p_t if p_t else None
        ctl("nodos_planillas", "Planillas de la Lista Cruzada que existen en la Lista de Planillas", f"{pc2:.1f} % ({p_ok} de {p_t})" if pc2 is not None else "—",
            "gris" if pc2 is None else ("verde" if pc2 >= 98 else "rojo"),
            "Id_planilla de cada nodo (separando las celdas con coma) contra la Lista de Planillas cargada. Medido el 25 oct: 1.183 de 1.183. Meta 98 %.")
    else:
        for k, t in (("nodos_caso", "Nodos de la Lista Cruzada con caso en Gestión"), ("nodos_planillas", "Planillas de la Lista Cruzada que existen en la Lista de Planillas")):
            ctl(k, t, "—", "gris", "No hay una Lista Cruzada de Nodos cargada.")
    # 15-16. CHANGELOG 2026-10-25 SCM · LCL: las LCL del exporte de LCL son de un caso (por su código o por la LCL de las planillas de un caso DF) y, en las CONCLUSA,
    # el Importe da la MO de sus libretos vivos con ±GESTION_SCM_IMPORTE_PCT (medido el 25 oct: 676 de 814 = 83 %)
    (l_tot, l_caso, l_sinlib, l_sinlib_caso), = await _sql(db, """
        with l as (select distinct lcl from sgo7_gg_scm_lcl where vigente), c as (select codigo from sgo7_gg_proyectos where codigo is not null),
        k as (select lcl, (lcl in (select codigo from c) or lcl in (select p.lcl from sgo7_gg_sonny_planilla p where p.ticket in (select codigo from c))) caso,
                     lcl not in (select lcl from sgo7_gg_scm_libreto) sinlib from l)
        select count(*), count(*) filter (where caso), count(*) filter (where sinlib), count(*) filter (where sinlib and caso) from k""")
    if l_tot:
        pc5 = 100 * l_caso / l_tot
        ctl("scm_lcl_caso", "LCL del exporte de LCL de SCM con caso en Gestión", f"{pc5:.1f} % ({l_caso} de {l_tot})", "verde" if pc5 >= 90 else "ambar",
            f"Por el código del caso (LCL) o por la LCL de las planillas de un caso de mayor impacto. De las {l_tot - l_caso} sin caso, "
            f"{l_sinlib - l_sinlib_caso} no tienen ningún libreto en SCM (un caso nace de Sonny, PMS, SCM-libretos o el correo; la LCL sola no lo hace nacer). "
            f"LCL sin libreto: {l_sinlib} (con caso {l_sinlib_caso}). Meta 90 %.")
        tol = _scm_cfg()["importe_pct"]
        (i_tot, i_ok, i_raras), = await _sql(db, """
            with i as (select lcl, sum(importe) imp from sgo7_gg_scm_lcl where vigente group by lcl having bool_and(estado = 'CONCLUSA')),
            m as (select lcl, sum(mo_pesos) mo from sgo7_gg_scm_libreto where coalesce(estado_libreto, '') not in ('ANN', 'RIF') group by lcl)
            select count(*) filter (where i.imp > 0 and coalesce(m.mo, 0) > 0),
                   count(*) filter (where i.imp > 0 and coalesce(m.mo, 0) > 0 and abs(m.mo - i.imp) <= cast(:t as double precision) / 100.0 * i.imp),
                   count(*) filter (where not (i.imp > 0 and coalesce(m.mo, 0) > 0)) from i left join m using (lcl)""", t=tol)
        pc6 = 100 * i_ok / i_tot if i_tot else None
        ctl("scm_importe_mo", f"Importe de la LCL ≈ MO de sus libretos en las CONCLUSA (±{tol:g} %)", f"{pc6:.1f} % ({i_ok} de {i_tot})" if pc6 is not None else "—",
            "gris" if pc6 is None else ("verde" if pc6 >= 80 else "rojo"),
            "Suma del Importe de las filas de cada LCL CONCLUSA contra la MO de sus libretos vivos (cantidad × precio QP). Medido el 25 oct: 676 de 814 (83 %). "
            "Por debajo de 80 % algo cambió en el exporte o en los precios." + (f" · {i_raras} LCL concluida(s) con Importe o MO en $0 (no entran a la cuenta)" if i_raras else ""))
    else:
        for k, t in (("scm_lcl_caso", "LCL del exporte de LCL de SCM con caso en Gestión"), ("scm_importe_mo", "Importe de la LCL ≈ MO de sus libretos en las CONCLUSA")):
            ctl(k, t, "—", "gris", "No hay un exporte de LCL de SCM cargado (lcl_export).")
    # 17-20. CHANGELOG 2026-10-25 DRAPE: las LCL facturadas del contrato (las del exporte de LCL de SCM) son de un caso; las facturas traen «Num. Pedido»; las LCL
    # CONCLUSA tienen factura; en el detalle leído, la suma de las posiciones da el total de la factura
    (d_tot,), = await _sql(db, "select count(*) from sgo7_gg_drape_factura where vigente")
    if d_tot:
        (l_scm, l_caso, l_drape), = await _sql(db, _SQL_LCL_GESTION + """, d as (select distinct e lcl from sgo7_gg_drape_factura x,
                                                 jsonb_array_elements_text(coalesce(x.lcls, '[]'::jsonb)) e where x.vigente and x.estado in ('enviada', 'registrada', 'pagada'))
            select count(*) filter (where d.lcl in (select lcl from sgo7_gg_scm_lcl)),
                   count(*) filter (where d.lcl in (select lcl from sgo7_gg_scm_lcl) and d.lcl in (select lcl from l)), count(*) from d""")
        pc7 = 100 * l_caso / l_scm if l_scm else None
        ctl("drape_lcl_caso", "LCL facturadas en DRAPE con caso en Gestión", f"{pc7:.1f} % ({l_caso} de {l_scm})" if pc7 is not None else "—",
            "gris" if pc7 is None else ("verde" if pc7 >= 90 else "ambar"),
            "LCL de las facturas vivas de DRAPE (enviadas, registradas o pagadas) que están en el exporte de LCL de SCM (el contrato de la Zona VII), con caso "
            f"en Gestión por su código o por la LCL de la planilla de un caso DF. DRAPE nombra {l_drape} LCL vivas en total: las demás son de otros contratos o años. "
            "Meta 90 %.")
        por_e = await _sql(db, """select estado, count(*), count(*) filter (where jsonb_array_length(coalesce(pedidos, '[]'::jsonb)) > 0)
                                  from sgo7_gg_drape_factura where vigente group by 1 order by 1""")
        n_t, n_p = sum(int(r[1]) for r in por_e), sum(int(r[2]) for r in por_e)
        n_v, n_vp = sum(int(r[1]) for r in por_e if r[0] in gf.DRAPE_VIVAS), sum(int(r[2]) for r in por_e if r[0] in gf.DRAPE_VIVAS)
        pc8 = 100 * n_vp / n_v if n_v else None
        ctl("drape_pedido", "Facturas de DRAPE con «Num. Pedido»", f"{100 * n_p / n_t:.1f} % ({n_p} de {n_t})", "gris" if pc8 is None else ("verde" if pc8 >= 95 else "rojo"),
            f"Sin «Num. Pedido» una factura no se liga a ninguna LCL. Vivas (enviadas, registradas, pagadas): {n_vp} de {n_v}. Por lista: "
            + " · ".join(f"{e} {p} de {n}" for e, n, p in por_e) + ". Meta 95 % en las vivas.")
        (c_tot, c_fac, c_pag), = await _sql(db, """with c as (select lcl from sgo7_gg_scm_lcl where vigente group by lcl having bool_and(estado = 'CONCLUSA')),
              d as (select e lcl, bool_or(x.estado = 'pagada' or x.f_pago is not null) pag from sgo7_gg_drape_factura x, jsonb_array_elements_text(coalesce(x.lcls, '[]'::jsonb)) e
                    where x.vigente and x.estado in ('enviada', 'registrada', 'pagada') group by e)
            select count(*), count(*) filter (where c.lcl in (select lcl from d)), count(*) filter (where c.lcl in (select lcl from d where pag)) from c""")
        pc9 = 100 * c_fac / c_tot if c_tot else None
        ctl("drape_concluida", "LCL CONCLUSA en SCM con factura en DRAPE", f"{pc9:.1f} % ({c_fac} de {c_tot})" if pc9 is not None else "—",
            "gris" if pc9 is None else ("verde" if pc9 >= 95 else "ambar"),
            f"LCL concluidas según SCM que están en una factura enviada, registrada o pagada de DRAPE. Con pago en DRAPE: {c_pag} de {c_tot}: CONCLUSA no quiere "
            "decir pagado. Meta 95 %.")
        (p_ok, p_tot), = await _sql(db, "select count(*) filter (where abs(coalesce(suma_pos, 0) - coalesce(total, 0)) <= 5), count(*) from sgo7_gg_drape_factura where suma_pos is not null")
        ctl("drape_posiciones", "DRAPE: suma de las posiciones = total de la factura", f"{p_ok} de {p_tot}" if p_tot else "—",
            "gris" if not p_tot else ("verde" if p_ok == p_tot else "rojo"),
            "En cada factura cuyo detalle está cargado, todas las líneas del «Resumen de Posición» contra el «Total Factura» de la cabecera (±$5 de redondeo). "
            "Medido el 25 oct: 18 de 18.")
    else:
        for k, t in (("drape_lcl_caso", "LCL facturadas en DRAPE con caso en Gestión"), ("drape_pedido", "Facturas de DRAPE con «Num. Pedido»"),
                     ("drape_concluida", "LCL CONCLUSA en SCM con factura en DRAPE"), ("drape_posiciones", "DRAPE: suma de las posiciones = total de la factura")):
            ctl(k, t, "—", "gris", "No hay una lectura de DRAPE cargada.")
    # 10-12. CHANGELOG 2026-10-23 SAP (MB51): las filas de SAP reproducen la caja en todos los casos; SAP descargado = SCM posado por (LCL, código); recuperado
    meta = await _sap_meta(db)
    if not meta.get("hay"):
        for k, t in (("sap_filas", "Filas de SAP = cifra de la caja SAP"), ("sap_scm", "SAP descargado = SCM posado por (LCL, código)"),
                     ("sap_rec", "Recuperado SAP (965) = SCM recuperado")):
            ctl(k, t, "—", "gris", "No hay un MB51 de SAP cargado.")
        return out
    (s_bien, s_mal, s_sin), = await _sql(db, """select count(*) filter (where dinero->'evidencia'->'cajas'->'sap_mat'->>'ok' = 'true'),
                                                   count(*) filter (where dinero->'evidencia'->'cajas'->'sap_mat'->>'ok' = 'false'),
                                                   count(*) filter (where dinero->'evidencia'->'cajas'->'sap_mat' is null)
                                            from sgo7_gg_proyectos where activo and dinero is not null""")
    s_malos = await _sql(db, """select id, coalesce(codigo, clave) from sgo7_gg_proyectos where activo and dinero->'evidencia'->'cajas'->'sap_mat'->>'ok' = 'false' order by id limit 40""")
    ctl("sap_filas", "Filas de SAP = cifra de la caja SAP", f"{s_bien} de {s_bien + s_mal + s_sin} casos", "rojo" if s_mal else ("ambar" if s_sin else "verde"),
        (f"{s_mal} caso(s) donde las filas del MB51 no reproducen la caja" if s_mal else "en todos los casos las filas de SAP suman la caja «MAT descargado» al peso")
        + (f" · {s_sin} caso(s) sin verificar (recalcular)" if s_sin else ""), casos=[{"id": i, "codigo": c} for i, c in s_malos])
    cs = await _controles_sap(db)
    p1, p2 = cs.get("pct_pares_iguales"), cs.get("pct_rec_iguales")
    ctl("sap_scm", "SAP descargado = SCM posado por (LCL, código)", f"{p1:.1f} % ({cs['pares_iguales']} de {cs['pares']})" if p1 is not None else "—",
        "gris" if p1 is None else ("verde" if p1 >= 90 else ("ambar" if p1 >= 80 else "rojo")),
        f"Pares (LCL, código de Enel) de {cs['lcl_pares']} LCL que están en SAP y en SCM, con el libreto autocertificado hasta el corte del MB51 "
        f"({meta['corte']:%d/%m}); iguales al milésimo. Medido el 23 oct: 92,0 %. LCL de SAP que están en SCM: "
        + (f"{cs['pct_lcl_en_scm']:.1f} % ({cs['lcl_en_scm']} de {cs['lcl_sap']})." if cs.get("pct_lcl_en_scm") is not None else "—."), **cs)
    ctl("sap_rec", "Recuperado SAP (965) = SCM recuperado", f"{p2:.1f} % ({cs['rec_iguales']} de {cs['rec_pares']})" if p2 is not None else "—",
        "gris" if p2 is None else ("verde" if p2 >= 95 else ("ambar" if p2 >= 90 else "rojo")),
        "Pares (LCL, código) con chatarra en SAP (965 − 966, KG) y «Cant. Recuperada» en SCM. Medido el 23 oct: 98,1 %. Confirma que 965 es el material "
        "retirado de la obra (clase propia de Enel).")
    return out


@router.get("/fuentes")
async def fuentes(db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    """CHANGELOG 2026-10-23 Panel «Fuentes y controles»: una fila por sistema (último exporte, última carga, filas) y los controles automáticos."""
    v = _cache_cargas.get("fuentes")
    if v and time.time() - v[0] < 60:
        return v[1]
    ult = await _ultimas_cargas(db)
    n = {t: (await _sql(db, f"select count(*) from {t}"))[0][0] for t in ("sgo7_gg_pms", "sgo7_gg_pms_log", "sgo7_gg_scm_libreto", "sgo7_gg_scm_prest",
                                                                            "sgo7_gg_scm_mat", "sgo7_gg_sonny", "sgo7_gg_sonny_planilla", "sgo7_gg_deuda",
                                                                            "sgo7_gg_sap", "sgo7_gg_precio_mat", "sgo7_gg_sonny_nodo", "sgo7_gg_sonny_plan_act",
                                                                            "sgo7_gg_scm_lcl", "sgo7_gg_drape_factura", "sgo7_gg_drape_pos", "sgo7_gg_drape_em")}
    # CHANGELOG 2026-10-25: DRAPE por lista, cuántas con el detalle leído y la tabla «Entradas de bienes a facturar» de Inicio (la guarda la carga en su detalle)
    dr_e = await _sql(db, "select estado, count(*), coalesce(sum(total), 0) from sgo7_gg_drape_factura where vigente group by 1 order by 1")
    (dr_det,), = await _sql(db, "select count(*) from sgo7_gg_drape_factura where detalle_en is not null")
    u_dr = ult.get(("drape", "listas")) or {}
    em_res = (u_dr.get("detalle") or {}).get("entradas_a_facturar") or {}
    dr_nota = None
    if dr_e:
        dr_nota = (" · ".join(f"{gf.ESTADOS_DRAPE.get(e, e).split(' (')[0]} {int(nn)} ({_m_txt(t)})" for e, nn, t in dr_e)
                   + f" · detalle leído de {dr_det} factura(s), {n['sgo7_gg_drape_pos']} posiciones"
                   + (f" · EM pendientes de facturar: {em_res.get('em')} por {_m_txt(em_res.get('importe'))} (total de Inicio; la lista de EM no se ve sin «Nueva factura»)"
                      if em_res.get("em") else "") + ". " + gf.NOTA_DRAPE + ".")
    meta_sap = await _sap_meta(db)
    car = await _cargas(db)
    (nodos_v, nodos_sp), = await _sql(db, "select count(*), count(*) filter (where jsonb_array_length(coalesce(id_planillas, '[]'::jsonb)) = 0) from sgo7_gg_sonny_nodo where vigente")
    (sgs_min, sgs_max, sgs_n), = await _sql(db, "select min(sgs->>'consultado_en'), max(sgs->>'consultado_en'), count(*) from sgo7_gg_proyectos where sgs is not null")
    lcl_v = (await _sql(db, """select count(*), count(*) filter (where c), count(*) filter (where s) from (select lcl, bool_and(estado = 'CONCLUSA') c,
                                 lcl not in (select lcl from sgo7_gg_scm_libreto) s from sgo7_gg_scm_lcl where vigente group by lcl) x"""))[0]      # CHANGELOG 2026-10-25
    ags = (await db.execute(select(GestionAgente))).scalars().all()
    ag = max(ags, key=lambda a: a.ultimo_contacto or datetime(2000, 1, 1, tzinfo=timezone.utc)) if ags else None

    def fila(k, titulo, partes, filas, sin_registro):
        cs = [ult.get(x) for x in partes]
        return {"k": k, "titulo": titulo, "cargas": [c for c in cs if c], "filas_tabla": filas, "sin_registro": None if any(cs) else sin_registro,
                "tipo": _TIPO_FUENTE.get(k, "exporte")}                         # CHANGELOG 2026-10-25: exporte · pantalla · vivo · buzon
    sistemas = [
        fila("pms", "PMS · solicitudes y cambios", [("pms", "solicitudes"), ("pms", "cambios")], n["sgo7_gg_pms"],
             f"cargado el {car['pms_cargado']:%d/%m %H:%M}, antes del registro de cargas" if car.get("pms_cargado") else "sin carga"),
        fila("scm", "SCM · libretos, prestaciones y materiales", [("scm", "libretos"), ("scm", "prestaciones"), ("scm", "materiales")], n["sgo7_gg_scm_libreto"],
             f"exporte del {car['scm']:%d/%m %H:%M}, antes del registro de cargas" if car.get("scm") else "sin carga"),
        # CHANGELOG 2026-10-25: el exporte de LCL de SCM (estado de cada LCL, validez, Importe y Trabajo = orden de SAP)
        fila("scm_lcl", "SCM · LCL (estado, validez, Importe y Trabajo)", [("scm", "lcl")], n["sgo7_gg_scm_lcl"],
             "sin carga (Informe > Informe LCL/LCA, pantalla «Informe Penalización y Bonificaciones»: lcl_export)")
        | {"nota": (f"{lcl_v[0]} LCL vigentes · CONCLUSA {lcl_v[1]} · sin libreto {lcl_v[2]} · " + gf.NOTA_CONCLUSA) if lcl_v[0] else None},
        fila("sonny_historial", "Sonny · Historial de Planillas", [("sonny", "historial")], n["sgo7_gg_sonny"],
             "antes del registro de cargas" if n["sgo7_gg_sonny"] else "sin carga"),
        fila("sonny_lista", "Sonny · Lista de Planillas", [("sonny", "lista")], n["sgo7_gg_sonny_planilla"],
             f"cargado el {car['sonny_cargado']:%d/%m %H:%M}, antes del registro de cargas" if car.get("sonny_cargado") else "sin carga"),
        # CHANGELOG 2026-10-25: Lista Cruzada de Nodos (Contrato JA10185664) + Nodos por gestion almacen, y Historial de Actividades + Lista de Estados
        fila("sonny_nodos", "Sonny · Lista Cruzada de Nodos y Nodos por gestión almacén", [("sonny", "nodos"), ("sonny", "almacen")], n["sgo7_gg_sonny_nodo"],
             "sin carga (Reportes > Consultas > «Lista Cruzada de Nodos», Contrato JA10185664)")
        | {"nota": (f"{nodos_v} nodos vigentes · {nodos_sp} sin planilla (Asignado, Finalizado o Pendiente)" if nodos_v else None)},
        fila("sonny_actividades", "Sonny · Historial de Actividades y Lista de Estados", [("sonny", "actividades"), ("sonny", "estados")], n["sgo7_gg_sonny_plan_act"],
             "sin carga (Reportes > Consultas > «Historial de Actividades», 2026-01-01 → hoy)")
        | {"nota": "solo la fecha y el tipo de cada acción por planilla (creó, cambió el estado, firma, editó): sin personas ni correos"},
        fila("deuda", "Informe de deuda de facturación", [("deuda", "informe")], n["sgo7_gg_deuda"],
             f"corte del {_local(car['deuda']):%d/%m}, antes del registro de cargas" if car.get("deuda") else "sin carga"),
        # CHANGELOG 2026-10-23: el MB51 (llega por correo) y hasta dónde llega el consumo; solo material, sin factura, HES ni pago
        fila("sap", "SAP · MB51 (material de Enel)", [("sap", "mb51")], n["sgo7_gg_sap"], "sin carga (el MB51 llega por correo)")
        | {"consumo_hasta": meta_sap["corte"].date().isoformat() if meta_sap.get("corte") else None, "precios": n["sgo7_gg_precio_mat"],
           "nota": ((f"consumo (261) hasta el {meta_sap['corte']:%d/%m/%Y} · desde el {meta_sap['desde']:%d/%m/%Y} · precio de Enel en "
                     f"{n['sgo7_gg_precio_mat']} códigos · solo material: falta " + gf.FALTA_SAP + ". " + gf.nota_cobro_sap()) if meta_sap.get("hay") else None)},
        # CHANGELOG 2026-10-25: DRAPE, el portal de facturación de Enel (lado proveedor): listas de facturas, detalle de las abiertas y EM a facturar
        fila("drape", "DRAPE · facturas (portal de facturación de Enel)", [("drape", "listas"), ("drape", "detalle"), ("drape", "em")], n["sgo7_gg_drape_factura"],
             "sin carga (portal de facturación del cliente: listas de facturas por estado)") | {"nota": dr_nota},
        {"k": "sgs", "titulo": "SGS · en vivo por el Centro MO", "vivo": True, "casos": sgs_n, "tipo": "vivo",
         "desde": _local(_iso(sgs_min)).isoformat() if sgs_min else None, "hasta": _local(_iso(sgs_max)).isoformat() if sgs_max else None},
        {"k": "correo", "titulo": "Correo · agente del PC", "agente": ag.buzon if ag else None, "correos": ag.correos_total if ag else 0, "tipo": "buzon",
         "leido_hasta": _local(ag.cursor).isoformat() if ag and ag.cursor else None,
         "ultimo_contacto": _local(ag.ultimo_contacto).isoformat() if ag and ag.ultimo_contacto else None},
    ]
    ctl = await _controles(db)
    if not gf.MOSTRAR_DRAPE:                         # CHANGELOG 2026-10-25: ni la fila ni los controles del portal de facturación
        sistemas = [x for x in sistemas if x["k"] != "drape"]
        ctl = [c for c in ctl if not str(c.get("k") or "").startswith("drape") and "DRAPE" not in json.dumps(c, default=str).upper()]
        n = {k: v for k, v in n.items() if "drape" not in k}
    out = {"calculado_en": datetime.now(_TZ).replace(tzinfo=None).isoformat(timespec="seconds"), "sistemas": sistemas, "controles": ctl,
           "verdes": sum(1 for c in ctl if c["estado"] == "verde"), "rojos": sum(1 for c in ctl if c["estado"] == "rojo"),
           "filas": {k.replace("sgo7_gg_", ""): v for k, v in n.items()}}
    _cache_cargas["fuentes"] = (time.time(), out)
    return out


# ── Para directivos: la portada de la zona y la cadena de cada trabajo ─────────────────────────────────────────────
_CTL_LLANO = {"cajas_filas": "Suma de las filas = cifra de cada trabajo", "sonny_pms": "Planillas de Sonny cuya solicitud está en PMS",
              "planillas_caso": "Planillas de Sonny con trabajo en Gestión", "enlaces": "Libretos vigentes enlazados a su planilla",
              "scm_importe_mo": "Importe de la LCL ≈ mano de obra de sus libretos, en las LCL concluidas", "drape_concluida": "LCL concluidas en SCM con factura en DRAPE",
              "precios": "Tarifas de Enel usadas = lista vigente", "prestaciones": "Mano de obra de cada libreto = suma de sus actividades",
              "sap_filas": "Filas de SAP = cifra de material en SAP", "sap_scm": "Material descargado en SAP = instalado en SCM (por LCL y código)",
              "sap_rec": "Material recuperado: SAP = SCM"}
_FUENTE_LLANA = {"sap": "SAP · material de Enel", "sgs": "SGS · lectura en vivo", "correo": "Correo · lectura automática del buzón"}


@router.get("/zona")
async def zona(db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    """CHANGELOG 2026-10-25 La portada de la zona: MO por etapa (una figura; los gemelos de una LCL cuentan una vez), lo que frena el cobro (frases ordenadas por
    pesos, las facturas vencidas una vez por factura), los trabajos por estado, los 10 más grandes en juego y la garantía de los datos. Todo
    calculado aquí con lo guardado en el recálculo (dinero.resumen y dinero.ux); caché de 60 s que limpia cada carga y cada recálculo."""
    v = _cache_cargas.get("zona")
    if v and time.time() - v[0] < 60:
        return v[1]
    D = GestionProyecto.dinero
    filas = (await db.execute(select(GestionProyecto.id, GestionProyecto.codigo, GestionProyecto.clave, GestionProyecto.nombre, GestionProyecto.zona,
                                     GestionProyecto.tipo, GestionProyecto.pilar, GestionProyecto.alertas, D["resumen"], D["ux"]["estado"], D["ux"]["valor"],
                                     D["ux"]["facturas"], D["ux"]["error"])
                              .where(GestionProyecto.activo.is_(True)).order_by(GestionProyecto.id))).all()
    casos = [dict(id=i, codigo=c, clave=k, nombre=n, zona=z, tipo=t, pilar=pi, alertas=al or [], resumen=r or {},
                  ux=({"estado": e, "valor": vv} if e else None), facturas=fs or []) for i, c, k, n, z, t, pi, al, r, e, vv, fs, _err in filas]
    car = await _cargas(db)
    cobro = await _cobro_resumen(db) if car.get("drape") else {"hay": False}
    fu = await fuentes(db=db, user=user)
    ctl = fu.get("controles") or []
    garantia = {"verdes": sum(1 for c in ctl if c.get("estado") == "verde"), "ambar": sum(1 for c in ctl if c.get("estado") == "ambar"),
                "rojos": sum(1 for c in ctl if c.get("estado") == "rojo"), "calculado": fu.get("calculado_en"),
                "controles": [{"k": c.get("k"), "estado": c.get("estado"), "titulo": _CTL_LLANO.get(c.get("k")) or c.get("titulo"), "valor": c.get("valor")} for c in ctl],
                "fuentes": [{"k": s_.get("k"), "titulo": _FUENTE_LLANA.get(s_.get("k")) or s_.get("titulo"), "vivo": bool(s_.get("vivo")),
                             "exporte": ((s_.get("cargas") or [{}])[0] or {}).get("exporte"), "cargado": ((s_.get("cargas") or [{}])[0] or {}).get("cargado"),
                             "filas": ((s_.get("cargas") or [{}])[0] or {}).get("filas") or s_.get("filas_tabla"),
                             "leido_hasta": s_.get("leido_hasta") or s_.get("hasta"), "tipo": s_.get("tipo") or "exporte",                # CHANGELOG 2026-10-25
                             "cargas": [{k_: c_.get(k_) for k_ in ("id", "parte", "archivo", "exporte", "cargado", "tamano", "tiene_archivo", "archivo_nombre",
                                                                    "archivo_tamano")} for c_ in s_.get("cargas") or []]}
                            for s_ in fu.get("sistemas") or []]}
    out = gf.armar_zona(casos, cobro, garantia, _pago_confirmado())
    iso = lambda d: (d.isoformat() if hasattr(d, "isoformat") else d) if d else None   # noqa: E731
    out["corte"] = {"sonny": iso(car.get("sonny_cargado")), "scm": iso(car.get("scm")), "pms": iso(car.get("pms_cargado")),
                    "sap": car["sap"].date().isoformat() if car.get("sap") else None, "scm_lcl": iso(car.get("scm_lcl")),
                    "deuda": iso(_local(car.get("deuda")).date()) if car.get("deuda") else None}                    # CHANGELOG 2026-10-25: el informe de deuda
    if gf.MOSTRAR_DRAPE:
        out["corte"]["drape"] = iso(car.get("drape"))
    ej = _caso_ejemplo()
    ej_id = next((c["id"] for c in casos if c["codigo"] == ej and c["tipo"] != "linea_viva"), None) or next((c["id"] for c in casos if c["codigo"] == ej), None)
    out["ejemplo"] = {"codigo": ej, "id": ej_id}
    out["sin_ux"] = sum(1 for f in filas if not f[9])
    out["con_error_ux"] = sum(1 for f in filas if f[12])
    out["calculado_en"] = datetime.now(_TZ).replace(tzinfo=None).isoformat(timespec="seconds")
    _cache_cargas["zona"] = (time.time(), out)
    return out


@router.get("/proyectos/{pid}/cadena")
async def cadena(pid: int, db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    """CHANGELOG 2026-10-25 El trabajo para el directivo: estado (y quién tiene la pelota), «qué falta para cobrar», la cadena de identificadores y la plata (6 cifras
    de MO y la línea del material). Lee lo guardado en el último recálculo: no calcula en la petición."""
    p = await db.get(GestionProyecto, pid)
    if not p:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")
    din = p.dinero or {}
    ux = din.get("ux") or {}
    hs = await _hermanos(db, p)
    return {"id": p.id, "clave": p.clave, "codigo": p.codigo, "nombre": p.nombre, "tipo": p.tipo, "zona": p.zona, "supervisor": p.supervisor,
            "codes": (p.padre or {}).get("codes") or [p.clave], "cerrado": p.cerrado, "pilar": p.pilar,
            "hermanos": [dict(id=h.id, clave=h.clave, tipo=h.tipo, zona=h.zona, supervisor=h.supervisor) for h in hs],
            "estado": ux.get("estado"), "que_falta": ux.get("que_falta") or [], "cadena": ux.get("cadena"), "plata": ux.get("plata"),
            "valor": ux.get("valor"), "pago_confirmado": _pago_confirmado() and gf.MOSTRAR_DRAPE, "cobro_visible": gf.MOSTRAR_DRAPE,           # CHANGELOG 2026-10-25
            "calculado_en": ux.get("calculado_en") or din.get("calculado_en"),
            "error": ux.get("error"), "sin_calcular": not ux, "n_eventos": p.n_eventos, "n_alertas": len(p.alertas or []),
            "alertas": [{k: a.get(k) for k in ("tipo", "frase", "color", "pelota", "pesos", "grupo", "texto", "dias", "desde", "planillas", "motivo_cat", "motivo_cats")}
                        for a in p.alertas or []]}                                                                  # CHANGELOG 2026-10-25: planillas y categoría del motivo


async def _evidencia_de(db: AsyncSession, p: GestionProyecto) -> dict:
    fd = await _fuentes_dinero(db, p)
    din = p.dinero or {}
    if not din.get("cajas"):
        return {"tarjetas": {}, "orden": [], "sin_franja": True, "ok": None, "malas": [], "decisiones": []}
    E = await _fuentes_evidencia(db, p, fd, din)
    ev = gf.armar_evidencia(E)
    ev.update(calculado_en=din.get("calculado_en"), decisiones=E["decisiones"], confianza=gf.CONFIANZA, fuentes=E["fuentes"],
              caso={"id": p.id, "clave": p.clave, "codigo": p.codigo, "nombre": p.nombre})
    return ev


@router.get("/proyectos/{pid}/evidencia")
async def evidencia(pid: int, db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    """CHANGELOG 2026-10-23 «¿De dónde sale?»: las tarjetas de evidencia del caso (filas que suman cada caja y cada alerta de dinero)."""
    p = await db.get(GestionProyecto, pid)
    if not p:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")
    ev = await _evidencia_de(db, p)
    ev["puede_decidir"] = _puede_decidir(user)
    return ev


_HOJA = {"pms": "PMS", "sgs_mo": "SGS MO", "sgs_mat": "SGS MAT", "sonny": "Sonny", "scm_mo": "SCM MO", "scm_mat": "SCM MAT",
         "alerta:mat_cobrado_sin_consumir": "Alerta material", "alerta:aprobado_sin_libreto": "Alerta sin libreto", "alerta:libreto_en_enel": "Alerta libreto Enel",
         "alerta:libreto_devuelto": "Alerta devuelto", "alerta:liquidado_sin_cobro": "Alerta sin cobro",
         # CHANGELOG 2026-10-23
         "sap_mat": "SAP", "material": "Material por código", "alerta:scm_sin_sap": "Alerta SCM sin SAP", "alerta:sap_mas_que_scm": "Alerta SAP más que SCM",
         "alerta:sgs_sin_sap": "Alerta SGS sin SAP",
         # CHANGELOG 2026-10-25
         "alerta:mat_excede_pms": "Alerta exceso PMS", "alerta:mat_por_liquidar": "Alerta por liquidar",
         # CHANGELOG 2026-10-25
         "sonny_etapa": "Sonny por etapa", "alerta:nodo_sin_planilla": "Alerta nodo sin planilla", "alerta:planilla_en_curso_quieta": "Alerta planilla quieta",
         "alerta:planilla_en_firmas": "Alerta en firmas", "alerta:pms_finalizado_sin_planilla": "Alerta PMS sin planilla",
         # CHANGELOG 2026-10-25
         "scm_estado": "SCM por estado", "scm_lcl": "SCM LCL", "alerta:lcl_concluida_libreto_abierto": "Alerta LCL concluida",
         "alerta:lcl_sin_libreto_vence": "Alerta LCL vence", "alerta:importe_lcl_vs_mo": "Alerta importe LCL",
         # CHANGELOG 2026-10-25
         "drape": "DRAPE", "alerta:lcl_concluida_sin_factura": "Alerta concluida sin factura", "alerta:factura_vencida_sin_pago": "Alerta factura vencida",
         "alerta:factura_rechazada": "Alerta factura rechazada", "alerta:em_sin_facturar": "Alerta EM sin facturar",
         # CHANGELOG 2026-10-25
         "alerta:rechazo_abierto": "Alerta rechazo Enel",
         # CHANGELOG 2026-10-25
         "deuda": "Informe de deuda"}


def _fecha_txt_local(s) -> str:
    if not s:
        return ""
    try:
        d = datetime.fromisoformat(str(s).replace("Z", "+00:00"))
    except ValueError:
        return str(s)
    if d.tzinfo:
        d = d.astimezone(_TZ).replace(tzinfo=None)
    return d.strftime("%d/%m/%Y %H:%M")


def _fuente_txt(f: Optional[dict]) -> str:
    if not f:
        return ""
    if f.get("vivo"):
        return f"{f.get('sistema')}: {f.get('nota')} · {_fecha_txt_local(f.get('consultado'))}"
    partes = [f"{c['archivo']} (exporte {_fecha_txt_local(c.get('exporte'))}; cargado {_fecha_txt_local(c.get('cargado'))} por {c.get('usuario') or '?'}; "
              f"{c.get('filas') or 0} filas; sha256 {str(c.get('sha256'))[:12]}…)" for c in f.get("cargas") or []]
    return (f"{f.get('sistema')}: " + (" · ".join(partes) if partes else (f.get("sin_registro") or ""))
            + (f" · consumo hasta {f['corte']}" if f.get("corte") else ""))


def _excel_evidencia(p: GestionProyecto, ev: dict) -> bytes:
    """Una hoja por caja con sus filas y =SUM() al pie (más la cifra de la caja y «¿cuadra?»), hoja «Fuentes» con los exportes y sus
    sha256 y hoja «Decisiones». Patrón de trabajo/rechazos-facturacion/evidencia/construir_evidencia.py."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    CAB, FONDO, NEG = Font(bold=True, color="F6F4EF"), PatternFill("solid", fgColor="13202E"), Font(bold=True)
    wb = Workbook()
    res = wb.active
    res.title = "Resumen"
    res.append(["Evidencia de la franja del dinero", f"{p.codigo or ''} · {p.clave} · {p.nombre or ''}"])
    res.append(["Generado", datetime.now(_TZ).strftime("%d/%m/%Y %H:%M") + " (hora Colombia)" + (f" en {_URL_PUBLICA}/gestion" if _URL_PUBLICA else "")])
    res.append(["Franja calculada", _fecha_txt_local(ev.get("calculado_en"))])
    res.append([])
    res.append(["Tarjeta", "Fórmula", "Cifra de la caja", "Suma de las filas", "¿Cuadra al peso?", "Filas"])
    for c in res[5]:
        c.font, c.fill = CAB, FONDO
    res["A1"].font = NEG

    def val(v):
        if v is None:
            return None
        if isinstance(v, bool):
            return "sí" if v else "no"
        if isinstance(v, (int, float)):
            return v
        if isinstance(v, (list, dict)):
            return json.dumps(v, ensure_ascii=False, default=str)[:3000]
        return str(v)[:3000]
    for k in ev.get("orden") or []:
        t = ev["tarjetas"][k]
        ws = wb.create_sheet(_HOJA.get(k, k.replace(":", " ")[:28]))
        ws.append([t["titulo"]])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append(["Fórmula", t.get("formula") or ""])
        ws.append(["Fuente", _fuente_txt(t.get("fuente"))])
        ws.append([])
        cols = list(t["columnas"])
        ws.append([c["t"] for c in cols])
        r0 = ws.max_row
        for c in ws[r0]:
            c.font, c.fill = CAB, FONDO
        for f in (f for f in t["filas"] if not f.get("_cab")):
            ws.append([val(f.get(c["k"])) for c in cols])
        r1 = ws.max_row
        pie = r1 + 1
        ws.cell(row=pie, column=1, value="Suma de las filas").font = NEG
        ws.cell(row=pie + 1, column=1, value="Cifra de la caja").font = NEG
        ws.cell(row=pie + 2, column=1, value="¿Cuadra al peso?").font = NEG
        celda_suma = celda_cifra = None
        for j, c in enumerate(cols, 1):
            L = get_column_letter(j)
            if c["tipo"] in ("pesos", "num"):
                for r in range(r0 + 1, r1 + 1):
                    ws.cell(row=r, column=j).number_format = "#,##0" if c["tipo"] == "pesos" else "#,##0.###"
            if not c.get("suma"):
                continue
            ws.cell(row=pie, column=j, value=f"=SUM({L}{r0 + 1}:{L}{max(r1, r0 + 1)})").number_format = "#,##0"
            s = next((x for x in t["sumas"] if x["k"] == c["k"]), None)
            if s is not None:
                ws.cell(row=pie + 1, column=j, value=s["cifra"]).number_format = "#,##0"
                ws.cell(row=pie + 2, column=j, value=(f'=IF(ABS({L}{pie}-{L}{pie + 1})<1,"sí","no")' if s["cifra"] is not None else "sin cifra"))
                if celda_suma is None and s["cifra"] is not None:      # CHANGELOG 2026-10-23: la columna que se compara es la que tiene cifra de caja
                    celda_suma, celda_cifra = f"{L}{pie}", f"{L}{pie + 1}"
        for titulo, clave in (("Pendientes de confirmar (no suman)", "pendientes"), ("Descartadas por una persona", "descartadas"),
                              ("No suman (anulados o devueltos)", "no_suman"), ("Revisar: SGS no reconoce el código en este LCL", "revisar"),
                              ("Diferencias menores (por debajo del mínimo en el caso: no son alerta)", "menores")):
            xs = t.get(clave) or []
            if not xs:
                continue
            ws.append([])
            ws.append([titulo])
            ws.cell(row=ws.max_row, column=1).font = NEG
            ks = [k2 for k2 in xs[0].keys() if k2 not in ("acciones", "copiar", "en_sgs")]
            ws.append(ks)
            ra = ws.max_row + 1
            for x in xs:
                ws.append([val(x.get(k2)) for k2 in ks])
            if clave == "menores" and "valor" in ks:      # CHANGELOG 2026-10-23: la suma de las diferencias menores, con fórmula
                jv = ks.index("valor") + 1
                L = get_column_letter(jv)
                ws.cell(row=ws.max_row + 1, column=1, value="Suma de las diferencias menores").font = NEG
                ws.cell(row=ws.max_row, column=jv, value=f"=SUM({L}{ra}:{L}{ws.max_row - 1})").number_format = "#,##0"
        for x in t.get("extra") or []:                 # CHANGELOG 2026-10-23: tablas de apoyo (filas de SAP de los códigos, recuperado) con =SUM en lo que suma
            xc = list(x.get("columnas") or [])
            if not xc:
                continue
            ws.append([])
            ws.append([x.get("titulo") or ""])
            ws.cell(row=ws.max_row, column=1).font = NEG
            ws.append([c["t"] for c in xc])
            ra = ws.max_row + 1
            for f in x.get("filas") or []:
                ws.append([val(f.get(c["k"])) for c in xc])
            rb = ws.max_row
            if rb >= ra:
                for j, c in enumerate(xc, 1):
                    if c.get("suma"):
                        L = get_column_letter(j)
                        ws.cell(row=rb + 1, column=j, value=f"=SUM({L}{ra}:{L}{rb})").number_format = "#,##0" if c["tipo"] == "pesos" else "#,##0.###"
                ws.cell(row=rb + 1, column=1, value="Suma").font = NEG
        for nta in t.get("notas") or []:
            ws.append([nta])
        for j in range(1, len(cols) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 18 if j > 1 else 22
        ws.freeze_panes = f"A{r0 + 1}"
        nombre = ws.title.replace("'", "")
        r = res.max_row + 1
        res.append([t["titulo"], t.get("formula") or "", f"='{nombre}'!{celda_cifra}" if celda_cifra else "sin cifra",
                    f"='{nombre}'!{celda_suma}" if celda_suma else "", (f'=IF(ABS(D{r}-C{r})<1,"sí","no")' if celda_suma and celda_cifra else "—"),
                    t["n_filas"]])
        res.cell(row=r, column=3).number_format = res.cell(row=r, column=4).number_format = "#,##0"
        res.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in (("A", 34), ("B", 90), ("C", 18), ("D", 18), ("E", 16), ("F", 8)):
        res.column_dimensions[col].width = w
    fu = wb.create_sheet("Fuentes")
    fu.append(["Sistema", "Parte", "Archivo", "Fecha del exporte", "Cargado", "Usuario", "Equipo", "Filas del archivo", "Tamaño (bytes)", "sha256"])
    for c in fu[1]:
        c.font, c.fill = CAB, FONDO
    vistos = set()
    for k in ev.get("orden") or []:
        f = ev["tarjetas"][k].get("fuente") or {}
        for c in f.get("cargas") or []:
            if c["id"] in vistos:
                continue
            vistos.add(c["id"])
            fu.append([c["sistema"], c["parte"], c["archivo"], _fecha_txt_local(c.get("exporte")), _fecha_txt_local(c.get("cargado")), c.get("usuario") or "",
                       c.get("equipo") or "", c.get("filas"), c.get("tamano"), c.get("sha256")])
        if f.get("sin_registro") and f.get("sistema") not in vistos:
            vistos.add(f.get("sistema"))
            fu.append([f.get("sistema"), "", "(sin registro de archivo)", "", f.get("sin_registro"), "", "", None, None, ""])
        if f.get("precios") and "precios" not in vistos:
            vistos.add("precios")
            pr = f["precios"]
            fu.append(["Precios QP", "lista", pr.get("archivo"), _fecha_txt_local(pr.get("modificado")), "producción", "", "", pr.get("codigos"), None, pr.get("sha256")])
        if f.get("vivo") and "sgs" not in vistos:
            vistos.add("sgs")
            fu.append(["SGS", "en vivo", "lectura por el Centro MO (solo lectura)", "", _fecha_txt_local(f.get("consultado")), "", "", None, None, ""])
        if f.get("precios_enel") and "precios_enel" not in vistos:     # CHANGELOG 2026-10-23
            vistos.add("precios_enel")
            pe = f["precios_enel"]
            cg = pe.get("carga") or {}
            fu.append(["Precio de Enel", "sgo7_gg_precio_mat", f"{pe.get('regla')} · de {cg.get('archivo') or 'MB51'}", _fecha_txt_local(pe.get("fecha")),
                       _fecha_txt_local(cg.get("cargado")), cg.get("usuario") or "", cg.get("equipo") or "", pe.get("n"), None, cg.get("sha256") or ""])
    for col, w in zip("ABCDEFGHIJ", (14, 14, 46, 18, 18, 16, 18, 10, 14, 68)):
        fu.column_dimensions[col].width = w
    de = wb.create_sheet("Decisiones")
    de.append(["Id", "Fecha", "Usuario", "LCL", "Libreto", "Planilla", "Decisión", "Nota"])
    for c in de[1]:
        c.font, c.fill = CAB, FONDO
    for d in ev.get("decisiones") or []:
        de.append([d["id"], _fecha_txt_local(d.get("fecha")), d.get("usuario"), d["lcl"], d["libreto"], d["id_planilla"],
                   "Es este" if d["decision"] == "si" else "No es", d.get("nota") or ""])
    if not ev.get("decisiones"):
        de.append(["—", "", "", "", "", "", "sin decisiones de personas en este caso", ""])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


@router.get("/proyectos/{pid}/evidencia/excel")
async def evidencia_excel(pid: int, request: Request, db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    """CHANGELOG 2026-10-23 Excel de evidencia del caso. Queda en Registros como toda descarga (usuario, IP, archivo y su sha256)."""
    p = await db.get(GestionProyecto, pid)
    if not p:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")
    ev = await _evidencia_de(db, p)
    if ev.get("sin_franja"):
        raise HTTPException(status_code=409, detail="El caso aún no tiene franja del dinero (recalcular)")
    datos = await asyncio.to_thread(_excel_evidencia, p, ev)
    nombre = f"evidencia_{re.sub(r'[^A-Za-z0-9_-]', '', p.codigo or p.clave)}_{datetime.now(_TZ):%Y%m%d-%H%M}.xlsx"
    reg = RegistroSGO7(ts=_ahora(), modulo="gestion", usuario=user.email, rol=_rol(user), accion="gestion_evidencia_excel", metodo="GET",
                       ruta=f"/api/gestion/proyectos/{pid}/evidencia/excel", archivo=nombre, archivo_hash=hashlib.sha256(datos).hexdigest(),
                       archivo_bytes=len(datos), estado="ok", ip=_ip(request),
                       detalle={"proyecto_id": p.id, "clave": p.clave, "codigo": p.codigo, "hojas": len(ev.get("orden") or []) + 3, "no_cuadran": ev.get("malas") or []})
    db.add(reg)
    await db.commit()
    return Response(content=datos, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{nombre}"', "X-Registro-Id": str(reg.id), "Cache-Control": "no-store"})


@router.post("/proyectos/{pid}/recalcular", dependencies=[Depends(require_dev)])
async def recalcular_caso(pid: int, db: AsyncSession = Depends(get_db)):
    """CHANGELOG 2026-10-23: recalcula un solo caso (relee SGS por el Centro MO, rehace el enlace planilla ↔ libreto de sus LCL con las decisiones y la franja)."""
    p = await db.get(GestionProyecto, pid)
    if not p:
        raise HTTPException(status_code=404, detail="Proyecto no encontrado")
    await _refrescar_sgs([p])
    fd = await _fuentes_dinero(db, p)
    enl = await _reenlazar(db, set(fd["lcls"]))
    await db.flush()
    await _recalcular(db, p, await _cm_registros())
    await db.commit()
    _cache_cargas.clear()
    return {"ok": True, "enlace": dict(enl), "evidencia": (p.dinero or {}).get("evidencia"), "alertas": [a["tipo"] for a in p.alertas or []]}


async def _aplicar_decision(db: AsyncSession, lcl: str) -> tuple:
    enl = await _reenlazar(db, {lcl})
    await db.flush()
    n, pend = await _recalcular_codigos(db, await _codigos_de_lcls(db, {lcl}))
    return enl, n, pend


@router.post("/decisiones")
async def decidir(body: DecisionIn, request: Request, db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    """CHANGELOG 2026-10-23 «Es este» / «No es» sobre un cruce planilla ↔ libreto. Se guarda en sgo7_gg_decisiones, manda en todo recálculo y queda en
    Registros. Lo puede hacer el desarrollador (y facturación cuando se le abra el módulo)."""
    if not _puede_decidir(user):
        raise HTTPException(status_code=403, detail="Sin permiso para confirmar cruces")
    lcl, lib, pid = gf.lcl_norm(body.lcl), gf.libreto_norm(body.libreto), (body.id_planilla or "").strip()
    if not lcl or not lib or not pid:
        raise HTTPException(status_code=422, detail="LCL, libreto o planilla inválidos")
    l = (await db.execute(select(GestionScmLibreto).where(GestionScmLibreto.lcl == lcl, GestionScmLibreto.libreto == lib))).scalar_one_or_none()
    if l is None:
        raise HTTPException(status_code=404, detail="Libreto no encontrado")
    pl = (await db.execute(select(GestionSonnyPlanilla).where(GestionSonnyPlanilla.id_planilla == pid))).scalar_one_or_none()
    if pl is None:
        raise HTTPException(status_code=404, detail="Planilla no encontrada")
    if pid not in (l.candidatas or []) and pid != l.planilla and pl.lcl != lcl:
        raise HTTPException(status_code=422, detail="Esa planilla no es candidata de ese libreto")
    d = (await db.execute(select(GestionDecision).where(GestionDecision.tipo == body.tipo, GestionDecision.lcl == lcl, GestionDecision.libreto == lib,
                                                         GestionDecision.id_planilla == pid))).scalar_one_or_none()
    antes = {"planilla": l.planilla, "regla": l.regla_enlace, "decision": d.decision if d else None}
    ahora = _ahora()
    if d is None:
        d = GestionDecision(tipo=body.tipo, lcl=lcl, libreto=lib, id_planilla=pid, creado_en=ahora)
        db.add(d)
    d.decision, d.usuario, d.rol, d.nota, d.proyecto_id, d.actualizado_en = body.decision, user.email, _rol(user), body.nota, body.proyecto_id, ahora
    d.llaves = {"lcl": lcl, "libreto": lib, "id_planilla": pid, "regla_antes": antes["regla"], "planilla_antes": antes["planilla"]}
    d.monto = l.mo_pesos
    await db.flush()
    enl, n, pend = await _aplicar_decision(db, lcl)
    reg = RegistroSGO7(ts=ahora, modulo="gestion", usuario=user.email, rol=_rol(user), accion="gestion_decision", metodo="POST", ruta="/api/gestion/decisiones",
                       archivo=f"{lcl}-{lib}", estado="ok", ip=_ip(request),
                       detalle={"decision_id": d.id, "proyecto_id": body.proyecto_id, "lcl": lcl, "libreto": lib, "id_planilla": pid, "decision": body.decision,
                                "antes": antes, "despues": {"planilla": l.planilla, "regla": l.regla_enlace}, "casos_recalculados": n})
    db.add(reg)
    await db.flush()
    out = {"ok": True, "decision": {"id": d.id, "decision": d.decision, "usuario": d.usuario}, "registro_id": reg.id,
           "libreto": {"planilla": l.planilla, "regla": l.regla_enlace}, "casos_recalculados": n, "recalculo_pendiente": pend}
    await db.commit()
    _cache_cargas.clear()
    return out


@router.delete("/decisiones/{did}")
async def deshacer_decision(did: int, request: Request, db: AsyncSession = Depends(get_db), user=Depends(require_gestion)):
    """CHANGELOG 2026-10-23: deshace una decisión (vuelven a mandar las reglas). Queda en Registros."""
    if not _puede_decidir(user):
        raise HTTPException(status_code=403, detail="Sin permiso para confirmar cruces")
    d = await db.get(GestionDecision, did)
    if d is None:
        raise HTTPException(status_code=404, detail="Decisión no encontrada")
    info = {"decision_id": d.id, "lcl": d.lcl, "libreto": d.libreto, "id_planilla": d.id_planilla, "decision": d.decision, "de": d.usuario}
    lcl = d.lcl
    await db.delete(d)
    await db.flush()
    enl, n, pend = await _aplicar_decision(db, lcl)
    l = (await db.execute(select(GestionScmLibreto).where(GestionScmLibreto.lcl == lcl, GestionScmLibreto.libreto == info["libreto"]))).scalar_one_or_none()
    reg = RegistroSGO7(ts=_ahora(), modulo="gestion", usuario=user.email, rol=_rol(user), accion="gestion_decision_deshacer", metodo="DELETE",
                       ruta=f"/api/gestion/decisiones/{did}", archivo=f"{lcl}-{info['libreto']}", estado="ok", ip=_ip(request),
                       detalle={**info, "despues": {"planilla": l.planilla if l else None, "regla": l.regla_enlace if l else None}, "casos_recalculados": n})
    db.add(reg)
    await db.flush()
    out = {"ok": True, "registro_id": reg.id, "libreto": {"planilla": l.planilla if l else None, "regla": l.regla_enlace if l else None},
           "casos_recalculados": n, "recalculo_pendiente": pend}
    await db.commit()
    _cache_cargas.clear()
    return out
